#!/usr/bin/env python3
"""K7 生成器：设计做法对应浙江2025清单编码参考表（待造价员签核），外加本地 xlsx 审签件。

受控源：
  SKILL.md 映射表（内部技能 building-cost-estimator 的方法论文件「常见映射参考」表；不在仓内）
  clause_llm.py 的 _MAPPING_REF（离线出清单工具的源码文件；不在仓内）
      —— 只用正则取出 _MAPPING_REF 这一个字符串变量；同包 app/config/ 下的任何文件一律不读。
  engine-service/app/engine/data/quantity_rules.yaml   @commit（默认 33601ca8）改正依据 + 逐行回放
  engine-service/app/engine/data/gb50854_zj2025.yaml   @commit  编码、名称、单位核对
  engine-service/app/engine/quantity_rules.py          @commit  回放直接执行引擎的 check_consistency

两个外部源（SKILL.md、clause_llm.py）不在仓里：默认从 sources/k7_sources_snapshot.json 读取摘录
（映射表原文行、_MAPPING_REF 各行，连同原文件 sha256）。在有原文件的机器上传 --skill / --clause，
脚本改从原文件读取并刷新快照（快照变了会在 stderr 提示）；产物与用快照生成的逐字节相同。

用法（默认 --repo 仓根、--out out/、--local local_only/）：
  python3 gen_map.py [--repo R] [--commit C] [--out DIR] [--local DIR] [--skill SKILL.md --clause clause_llm.py]
  生成时间取环境变量 SOURCE_DATE_EPOCH（秒），没有则取当前时间；固定它即可让 md 逐字节复现。
自检（任何一项不过都非零退出）：每个编码在浙江2025码表存在且名称、单位一致；5 处改正都已落实且写明原值；
逐行回放 quantity_rules 一致性闸 0 冲突；SKILL.md 映射表每行、_MAPPING_REF 每一项都有去向；不含金额。
"""
from __future__ import annotations

import argparse
import hashlib
import itertools
import json
import os
import re
import subprocess
import sys
import types
from datetime import datetime, timedelta, timezone
from pathlib import Path

import yaml

from kbsrc import SNAPSHOT_REPO, full_commit, git_show_bytes

SCRIPT = "gen_map.py"
RULES_PATH = "engine-service/app/engine/data/quantity_rules.yaml"
ZJ_PATH = "engine-service/app/engine/data/gb50854_zj2025.yaml"
ENGINE_PY = "engine-service/app/engine/quantity_rules.py"
HERE = Path(__file__).resolve().parent
# 两个外部源在 K7 文件头与快照里的显示名：只写文件名，不写任何本机路径 —— K7 要上传到租户知识库，
# 文件头第 1 行会被 chatbot 的 GREP / DOC 读到（PR #788 评审）。溯源靠后面跟着的原文件 sha256。
# 产物逐字节复现依赖这两串，改了要重登 expected_manifest.json。
SKILL_SHOW = "SKILL.md"
CLAUSE_SHOW = "clause_llm.py"
GB2013_DEFAULT = str(HERE / "sources" / "GB50854-2013.md")
GB2013_MD5 = "26504c410ebb4c458f53d8db3377103e"          # 原件 md5（方案登记）
GB2013_REPO_COPY_SHA256 = "2b8bdc7f1afcc96b3aeb2956df74bc8092081b3d128c59ebbcb69007253a3cff"  # 入仓副本
FILE_NAME = "【清单·房屋建筑】设计做法对应浙江2025清单编码参考表（待造价员签核）.md"
XLSX_NAME = "【清单·房屋建筑】设计做法对应浙江2025清单编码参考表（待造价员签核）-审签件.xlsx"
TITLE = "设计做法对应浙江2025清单编码参考表（待造价员签核）"
K6_TITLE = "浙江2025清单易混淆编码与列项校核规则（系统口径）"

CODE_RE = re.compile(r"(?<![0-9A-Za-z])Z?\d{9}(?!\d)")
DIGITS_RE = re.compile(r"\d{9,}")
MONEY_RE = re.compile(r"\d+(\.\d+)?\s*(元|万元)")
CN_TZ = timezone(timedelta(hours=8))

# ----------------------------------------------------------------------------- 人工口径（全部可回链）
# SKILL.md 映射表逐行的去向。键 = SKILL.md 行号；值 = 该行产出的一行或多行。
#   part：部位列（SKILL 原文没写部位的记「—」）；desc 缺省 = SKILL 原文描述；code 缺省 = SKILL 原编码。
#   fix：改正编号（见 CORRECTIONS），orig 为原值；note：适用或禁用说明（其中写的编码名称一律取码表原名，自检核对）。
SKILL_ROWS: dict[int, list[dict]] = {
    107: [dict(part="墙体")],
    108: [dict(part="墙体")],
    109: [dict(part="基础/砖胎膜",
               note="【待核】浙江2025码表另有 Z010401012「预制胎膜」（m²），砖胎膜是否改按该码列项待造价员核定")],
    110: [dict(part="柱")],
    111: [dict(part="梁")],
    112: [dict(part="墙/剪力墙")],
    113: [dict(part="板/楼板")],
    114: [dict(part="楼梯")],
    115: [dict(part="后浇带")],
    116: [dict(part="圈梁")],
    117: [dict(part="过梁")],
    118: [dict(part="构造柱")],
    119: [dict(part="—")],
    120: [dict(part="不同基体交接处")],
    121: [dict(part="—")],
    122: [dict(part="—")],
    123: [dict(part="—")],
    124: [dict(part="—")],
    125: [dict(part="—")],
    126: [dict(part="—")],
    127: [dict(part="屋面/顶板",
               note="隔离层用卷材时不按本行（roof_membrane 规则排除含「隔离」的原文），按 Z010902003「屋面隔离层及表面处理」")],
    128: [dict(part="屋面/顶板", note="天沟、檐沟涂膜防水按 010902009「天沟、檐沟涂膜防水」（roof_coating）")],
    129: [dict(part="屋面",
               note="楼地面细石混凝土面层用 011101002「细石混凝土楼地面」，找平层用 Z011101007「细石混凝土找平层」（confusables[3]）")],
    130: [dict(part="墙面")],
    131: [dict(part="墙面/外墙")],
    132: [dict(part="墙面")],
    133: [
        dict(key="a", desc="防水卷材 + 楼地面", part="楼地面",
             note="仅限楼地面；基础底板卷材按 Z010905001「基础卷材防水」（下一行）"),
        dict(key="b", desc="防水卷材 + 基础底板", part="基础底板/地下室底板", code="Z010905001", fix="C1",
             basis="quantity_rules.yaml#slab_membrane",
             note="基础（地下室）底板卷材防水按本行，不用 Z010904001「楼(地)面卷材防水」或屋面卷材编码"),
    ],
    134: [
        dict(key="a", desc="防水涂料 + 楼地面", part="楼地面",
             note="仅限楼地面；基础底板涂膜按 010905002「基础涂膜防水」（下一行）"),
        dict(key="b", desc="防水涂料 + 基础底板", part="基础底板/地下室底板", code="010905002", fix="C1",
             basis="quantity_rules.yaml#confusables[2]",
             note="基础（地下室）底板涂膜防水按本行，不用 Z010904002「楼(地)面涂膜防水」"),
        dict(key="c", desc="防水涂料 + 天棚", part="天棚",
             note="【待核】保留 SKILL.md 原口径；浙江2025码表另设 Z010902011「天棚面涂膜防水」（m²），"
                  "quantity_rules 无对应规则，待造价员核定"),
    ],
    135: [dict(part="屋面")],
    136: [dict(part="墙面")],
    137: [
        dict(key="a", desc="水泥砂浆 + 楼地面面层", part="楼地面",
             note="只用于水泥砂浆面层；平面找平层按 Z011101006「平面砂浆找平层」，立面找平层按 Z011201004「立面砂浆找平层」"
                  "（confusables[4]）"),
        dict(key="b", desc="水泥砂浆 + 平面找平层", part="楼地面/屋面", code="Z011101006", fix="C2",
             basis="quantity_rules.yaml#roof_leveling_not_floor 及 #confusables[4]",
             note="楼地面、屋面的水泥砂浆平面找平层按本行；屋面找平不得套 011101001「水泥砂浆楼地面」"),
        dict(key="c", desc="水泥砂浆 + 立面找平层", part="墙面/柱面", code="Z011201004", fix="C2",
             basis="quantity_rules.yaml#confusables[4]", note="墙、柱面的水泥砂浆立面找平层按本行"),
    ],
    138: [dict(part="楼地面",
               note="只用于细石混凝土面层；找平层用 Z011101007「细石混凝土找平层」，屋面保护层用 Z010902004「屋面刚性层」"
                    "（confusables[3]）")],
    139: [dict(part="—")],
    140: [dict(part="—")],
    141: [dict(part="楼地面")],
    142: [dict(part="楼地面")],
    143: [dict(part="—")],
    144: [dict(part="踢脚", note="踢脚按 0111-05 踢脚线列项，单位 m，不用墙面抹灰 Z011201001「墙、柱面一般抹灰」（confusables[0]）")],
    145: [dict(part="踢脚", note="踢脚按 0111-05 踢脚线列项，单位 m，不用墙面抹灰 Z011201001「墙、柱面一般抹灰」（confusables[0]）")],
    146: [dict(part="楼梯")],
    147: [dict(part="楼梯")],
    148: [dict(part="楼梯")],
    149: [dict(part="墙面/柱面", note="墙面刷涂料不用本码，按 011404001「墙面喷刷涂料」")],
    150: [dict(part="墙面/柱面")],
    151: [dict(part="墙面/柱面")],
    152: [dict(part="天棚",
               note="只用于天棚抹灰层本身；天棚刷涂料按 011404002「天棚喷刷涂料」，吊顶按 011302 节（confusables[1]）")],
    153: [dict(part="天棚/吊顶", note="吊顶不是天棚抹灰 Z011301001「天棚抹灰」")],
    154: [dict(part="墙面/内墙/外墙", note="clause_llm.py 第22行另列「无机涂料」，同按本码")],
    155: [dict(part="天棚/顶棚", note="顶棚刷涂料不是天棚抹灰 Z011301001「天棚抹灰」")],
}

# _MAPPING_REF 的补充行（SKILL.md 没有的）。键 = (clause_llm.py 行号, 该行第几项)。after = 排在 SKILL.md 哪一行之后。
CLAUSE_ROWS: dict[tuple[int, int], list[dict]] = {
    (13, 1): [dict(after=126, desc="木质窗", part="—")],
    (14, 0): [dict(after=126, desc="水泥瓦/瓦屋面", part="屋面")],
    (16, 1): [dict(after=129, desc="屋面隔离层（水泥砂浆）", part="屋面", code="Z010902003", fix="C4",
                   basis="zj2025.yaml codes.Z010902003",
                   note="屋面隔离层按本行（confusables[2]：屋面防水类用 Z010902xxx）")],
    (16, 2): [dict(after=129, desc="屋面变形缝密封", part="屋面", code="010901008", fix="C3",
                   basis="zj2025.yaml codes.010901008",
                   note="010903004「墙面变形缝」只用于墙面；原文另写「分仓缝」，若指刚性层分格缝，其嵌缝材料在 "
                        "Z010902004「屋面刚性层」项目特征「嵌缝材料种类」中描述，是否另列待造价员核定（【待核】）")],
    (19, 2): [
        dict(after=142, key="a", desc="防滑地砖 + 楼面/露台", part="楼面/露台", code="011102003",
             note="整片铺贴的楼面、露台防滑地砖按块料楼地面（原文「011108001 或 011102003」取后者）；零星部位见下一行"),
        dict(after=142, key="b", desc="防滑地砖 + 零星部位", part="楼地面零星部位/楼梯/台阶侧面", code="011108003", fix="C5",
             basis="zj2025.yaml codes.011108003",
             note="零星部位按本行，不是 011108001「石材零星项目」；零星范围参见国标2013 表 L.8 注（楼梯、台阶牵边和侧面等）"),
    ],
}

# _MAPPING_REF 里与 SKILL.md 重复的项：(行号, 项) → SKILL.md 行号。生成器核对两边编码相同。
CLAUSE_DUP: dict[tuple[int, int], int] = {
    (10, 0): 108, (10, 1): 107,
    (11, 0): 110, (11, 1): 111, (11, 2): 112, (11, 3): 113, (11, 4): 114,
    (12, 0): 119,
    (13, 0): 125,
    (15, 0): 127, (15, 1): 128,
    (16, 0): 129,
    (17, 0): 135, (17, 1): 136,
    (18, 0): 137, (18, 1): 138,
    (19, 0): 142, (19, 1): 141,
    (20, 0): 149, (20, 1): 151, (20, 2): 150,
    (21, 0): 152, (21, 1): 153,
    (22, 0): 154, (22, 1): 155,
    (23, 0): 133, (23, 1): 134,
}
# 重复项挂到 SKILL 拆分行的哪一行（其余拆分行不挂）
DUP_SPLIT_TARGET = {(18, 0): "a", (23, 0): "a", (23, 1): "a"}

CORRECTIONS = {
    "C1": "基础底板防水：SKILL.md 把底板防水并入楼(地)面防水",
    "C2": "水泥砂浆找平层：SKILL.md 把找平层并入水泥砂浆楼地面",
    "C3": "屋面变形缝：_MAPPING_REF 写成墙面变形缝的编码",
    "C4": "屋面隔离层：_MAPPING_REF 写了浙江2025码表没有的编码",
    "C5": "防滑地砖零星部位：_MAPPING_REF 写成石材零星项目的编码",
}
# 每处改正的原值（与源文件核对，不一致即报错）
EXPECTED_ORIG = {("skill", 133, "b"): "Z010904001", ("skill", 134, "b"): "Z010904002",
                 ("skill", 137, "b"): "011101001", ("skill", 137, "c"): "011101001",
                 ("clause", 16, 2, ""): "010903004", ("clause", 16, 1, ""): "010902003",
                 ("clause", 19, 2, "b"): "011108001"}


# ----------------------------------------------------------------------------- 工具

git_show = git_show_bytes


def sha256(b: bytes) -> str:
    return hashlib.sha256(b).hexdigest()


def gen_time() -> str:
    epoch = os.environ.get("SOURCE_DATE_EPOCH")
    dt = datetime.fromtimestamp(int(epoch), CN_TZ) if epoch else datetime.now(CN_TZ).replace(microsecond=0)
    return dt.isoformat()


def cell(s: str) -> str:
    s = str(s).replace("\r", "").replace("\n", "；")
    return s.replace("|", "\\|")


def entry(code: str, codes: dict) -> str:
    return f"{code}「{codes[code]['name']}」（{codes[code]['unit']}）"


# ----------------------------------------------------------------------------- 读源（或快照）

def read_sources(skill_path: str | None, clause_path: str | None, gb_path: str, snapshot: Path) -> dict:
    if bool(skill_path) != bool(clause_path):
        raise SystemExit("--skill 与 --clause 要么都给（从原文件读并刷新快照），要么都不给（读快照）")
    sp, cp = (Path(skill_path), Path(clause_path)) if skill_path else (None, None)
    if sp is not None and not (sp.is_file() and cp.is_file()):
        raise SystemExit(f"--skill / --clause 指向的文件不存在：{sp} / {cp}")
    if sp is not None:
        skill_b = sp.read_bytes()
        clause_b = cp.read_bytes()
        s_lines = skill_b.decode("utf-8").split("\n")
        hdr = next(i for i, l in enumerate(s_lines) if l.startswith("| 输入文件中的描述特征 |"))
        if not re.match(r"^\|(\s*-+\s*\|)+\s*$", s_lines[hdr + 1]):
            raise SystemExit("SKILL.md 映射表表头下不是分隔行")
        rows, i = [], hdr + 2
        while i < len(s_lines) and s_lines[i].startswith("|"):
            cells = [c.strip() for c in s_lines[i].strip().split("|")[1:-1]]
            if len(cells) != 3:
                raise SystemExit(f"SKILL.md 第{i + 1}行不是 3 列")
            rows.append({"line": i + 1, "desc": cells[0], "code": cells[1], "name": cells[2]})
            i += 1
        # 映射表原文（表头、分隔行、数据行、表后一行）逐行留底，供干净环境下校验器按行号回源
        table_lines = [{"line": n + 1, "text": s_lines[n]} for n in range(hdr, min(i + 1, len(s_lines)))]
        text = clause_b.decode("utf-8")
        m = re.search(r'^_MAPPING_REF = """\\' + "\n" + r'(.*?)"""', text, re.S | re.M)
        if not m:
            raise SystemExit("clause_llm.py 里找不到 _MAPPING_REF")
        ref = m.group(1)
        first = text[:m.start(1)].count("\n") + 1  # _MAPPING_REF 内容第一行的文件行号
        ref_lines = [{"line": first + k, "text": t} for k, t in enumerate(ref.split("\n")) if t.strip()]
        decl_line = text[:m.start()].count("\n") + 1
        end_line = text[:m.end()].count("\n") + 1
        gb_line = None
        gp = Path(gb_path)
        if gp.is_file():
            gb_b = gp.read_bytes()
            gb_md5 = hashlib.md5(gb_b).hexdigest()
            if hashlib.sha256(gb_b).hexdigest() == GB2013_REPO_COPY_SHA256:
                gb_md5 = GB2013_MD5   # 入仓副本只改了第 43 行，出处按原件登记
            for n, l in enumerate(gb_b.decode("utf-8").split("\n"), 1):
                if l.startswith("| 010902003 |"):
                    gb_line = {"line": n, "name": l.split("|")[2].strip(), "md5": gb_md5}
                    break
        snap = {
            "skill": {"label": SKILL_SHOW, "sha256": sha256(skill_b), "table_header_line": hdr + 1, "rows": rows,
                      "table_lines": table_lines},
            "clause": {"label": CLAUSE_SHOW, "sha256": sha256(clause_b), "mapping_ref_sha256": sha256(ref.encode()),
                       "decl_line": decl_line, "end_line": end_line, "lines": ref_lines},
            "gb2013_010902003": gb_line,
        }
        body = json.dumps(snap, ensure_ascii=False, indent=1, sort_keys=True) + "\n"
        if snapshot.is_file() and snapshot.read_text(encoding="utf-8") != body:
            print(f"注意：源文件与快照 {snapshot.name} 不一致，已按源文件刷新快照", file=sys.stderr)
        snapshot.write_text(body, encoding="utf-8")
        snap["_from"] = "source"
        return snap
    if not snapshot.is_file():
        raise SystemExit("源文件不在，快照也不在，无法生成 K7")
    snap = json.loads(snapshot.read_text(encoding="utf-8"))
    ref = "\n".join(x["text"] for x in snap["clause"]["lines"])
    snap["_from"] = "snapshot"
    return snap


def parse_clause(snap: dict) -> dict[tuple[int, int], dict]:
    items = {}
    for ln in snap["clause"]["lines"]:
        for k, it in enumerate(ln["text"].split("；")):
            if "→" not in it:
                raise SystemExit(f"clause_llm.py 第{ln['line']}行第{k}项没有「→」：{it}")
            lhs, rhs = it.split("→", 1)
            items[(ln["line"], k)] = {"text": it, "lhs": lhs, "rhs": rhs, "codes": CODE_RE.findall(rhs)}
    return items


# ----------------------------------------------------------------------------- 引擎回放

def load_engine(repo: str, commit: str, rules_obj: dict):
    src = git_show(repo, commit, ENGINE_PY)
    name = f"qr_engine_{commit}"
    mod = types.ModuleType(name)
    mod.__file__ = f"/nonexistent/{commit}/quantity_rules.py"
    sys.modules[name] = mod
    exec(compile(src.decode("utf-8"), f"quantity_rules.py@{commit}", "exec"), mod.__dict__)
    mod.load_rules = lambda: rules_obj  # 规则数据用同一 commit 的 yaml
    return mod, sha256(src)


def combos(desc: str, part: str) -> list[tuple[str, str]]:
    ctx = "" if part == "—" else part
    d_tok = [t for t in re.split(r"\s*\+\s*|/", desc) if t.strip()]
    p_tok = [t for t in ctx.split("/") if t.strip()] if ctx else []
    out = {(desc, ctx), (desc, ""), (f"{desc} {ctx}".strip(), ctx), (desc.replace(" + ", ""), ctx)}
    for dt, pt in itertools.product(d_tok, p_tok + [""]):
        out.add((dt, pt))
        out.add((f"{dt} {pt}".strip(), pt))
    for pt in p_tok:
        out.add((desc, pt))
    return sorted(out)


def replay(eng, code: str, desc: str, part: str) -> dict:
    viol, applied, n = [], set(), 0
    for text, ctx in combos(desc, part):
        n += 1
        t, c = eng._norm(text), eng._norm(ctx)
        for rule in eng.load_rules()["consistency_rules"]:
            if eng._rule_applies(rule, t, c):
                applied.add(rule["id"])
        for v in eng.check_consistency(code, text, ctx):
            viol.append({"text": text, "ctx": ctx, "rule": v.rule_id})
    return {"combos": n, "applied": sorted(applied), "violations": viol}


# ----------------------------------------------------------------------------- 组装

def build(a) -> tuple[str, list[dict], list[dict], dict]:
    commit_full = full_commit(a.repo, a.commit)
    rules_b = git_show(a.repo, a.commit, RULES_PATH)
    zj_b = git_show(a.repo, a.commit, ZJ_PATH)
    rules = yaml.safe_load(rules_b.decode("utf-8"))
    codes = yaml.safe_load(zj_b.decode("utf-8"))["codes"]
    eng, eng_sha = load_engine(a.repo, a.commit, rules)
    snap = read_sources(a.skill, a.clause, a.gb2013, Path(a.snapshot))
    skill_rows = {r["line"]: r for r in snap["skill"]["rows"]}
    clause_items = parse_clause(snap)
    problems: list[str] = []

    # 覆盖性：SKILL 表每行都有去向；_MAPPING_REF 每项都有去向且只一个
    if set(skill_rows) != set(SKILL_ROWS):
        problems.append(f"SKILL.md 映射表行号与口径表不一致：源 {sorted(skill_rows)} vs 口径 {sorted(SKILL_ROWS)}")
    dispo = set(CLAUSE_DUP) | set(CLAUSE_ROWS)
    if set(clause_items) != dispo or set(CLAUSE_DUP) & set(CLAUSE_ROWS):
        problems.append(f"_MAPPING_REF 项与去向表不一致：源 {sorted(clause_items)} vs 去向 {sorted(dispo)}")
    for key, sl in CLAUSE_DUP.items():
        ci, sr = clause_items.get(key), skill_rows.get(sl)
        if ci and sr and ci["codes"] != [sr["code"]]:
            problems.append(f"_MAPPING_REF {key} 编码 {ci['codes']} 与 SKILL.md 第{sl}行 {sr['code']} 不同")
    for key, rows in CLAUSE_ROWS.items():
        for r in rows:
            if r.get("fix"):
                ek = ("clause", key[0], key[1], r.get("key", ""))
                ci = clause_items.get(key)
                if not ci or EXPECTED_ORIG.get(ek) not in ci["codes"]:
                    problems.append(f"改正 {r['fix']} 的原值与 _MAPPING_REF {key} 不符：{ci and ci['codes']}")

    dup_of_skill: dict[tuple[int, str], list[int]] = {}
    for key, sl in CLAUSE_DUP.items():
        tgt = DUP_SPLIT_TARGET.get(key, "")
        dup_of_skill.setdefault((sl, tgt), []).append(key[0])

    out_rows: list[dict] = []

    def add(row: dict, origin: str, orig_code: str | None, orig_desc: str):
        code = row["code"]
        if code not in codes:
            problems.append(f"{origin}：编码 {code} 不在浙江2025码表")
            return
        rp = replay(eng, code, row["desc"], row["part"])
        notes = [row["note"]] if row.get("note") else []
        if rp["applied"]:
            notes.append("校核规则：" + "、".join(rp["applied"]) + "（逐行回放通过）")
        src = origin
        corr = None
        if row.get("fix"):
            if not orig_code:
                problems.append(f"{origin}：改正 {row['fix']} 缺原值")
            else:
                why = (f"该码浙江2025名称为「{codes[orig_code]['name']}」" if orig_code in codes
                       else "浙江2025码表无此编码")
                src += f"；按 {row['basis']} 改正（原写 {orig_code}，{why}）"
                corr = {"fix": row["fix"], "orig": orig_code, "orig_desc": orig_desc, "new": code,
                        "basis": row["basis"], "origin": origin, "orig_in_zj2025": orig_code in codes,
                        "orig_name_zj2025": codes.get(orig_code, {}).get("name"),
                        "negative_control": replay(eng, orig_code, row["desc"], row["part"])}
        out_rows.append({
            "desc": row["desc"], "part": row["part"], "code": code, "name": codes[code]["name"],
            "unit": codes[code]["unit"], "note": "；".join(notes) if notes else "—", "src": src,
            "replay": rp, "correction": corr, "flag": "【待核】" in (row.get("note") or ""),
        })

    clause_by_after: dict[int, list] = {}
    for key, rows in CLAUSE_ROWS.items():
        for r in rows:
            clause_by_after.setdefault(r["after"], []).append((key, r))

    for sl in sorted(SKILL_ROWS):
        sr = skill_rows.get(sl)
        if not sr:
            continue
        if sr["code"] in codes and codes[sr["code"]]["name"] != sr["name"]:
            problems.append(f"SKILL.md 第{sl}行名称「{sr['name']}」与码表「{codes[sr['code']]['name']}」不一致")
        for spec in SKILL_ROWS[sl]:
            row = dict(spec)
            row.setdefault("desc", sr["desc"])
            row.setdefault("code", sr["code"])
            k = spec.get("key", "")
            origin = f"SKILL.md 第{sl}行"
            if k:
                origin += f"（原「{sr['desc']}」拆分）"
            dups = dup_of_skill.get((sl, k), []) if k else dup_of_skill.get((sl, ""), [])
            if dups:
                origin += "；clause_llm.py " + "、".join(f"第{d}行" for d in sorted(set(dups))) + "同"
            orig = None
            if spec.get("fix"):
                orig = sr["code"]
                if EXPECTED_ORIG.get(("skill", sl, k)) != orig:
                    problems.append(f"SKILL.md 第{sl}行原值 {orig} 与预期不符")
            elif row["code"] != sr["code"]:
                problems.append(f"SKILL.md 第{sl}行未登记改正却改了编码")
            add(row, origin, orig, sr["desc"])
        for key, r in clause_by_after.get(sl, []):
            ci = clause_items[key]
            row = dict(r)
            if "code" not in row:
                if len(ci["codes"]) != 1:
                    problems.append(f"_MAPPING_REF {key} 多个编码却未指定取哪个")
                    continue
                row["code"] = ci["codes"][0]
            origin = f"clause_llm.py 第{key[0]}行（_MAPPING_REF）"
            orig = EXPECTED_ORIG.get(("clause", key[0], key[1], r.get("key", ""))) if r.get("fix") else None
            add(row, origin, orig, ci["text"])

    meta = {
        "commit": commit_full, "generated_at": gen_time(), "engine_sha256": eng_sha,
        "rules_sha256": sha256(rules_b), "zj_sha256": sha256(zj_b), "snap": snap, "codes": codes,
        "clause_items": {f"{k[0]}-{k[1]}": v["text"] for k, v in clause_items.items()},
        "problems": problems,
    }
    corrections = sorted((r for r in out_rows if r["correction"]),
                         key=lambda r: r["correction"]["fix"])  # 稳定排序：同一处改正内保持表内顺序
    md = render_md(out_rows, corrections, meta, codes)
    return md, out_rows, corrections, meta


def render_md(rows: list[dict], corrs: list[dict], meta: dict, codes: dict) -> str:
    s, c = meta["snap"]["skill"], meta["snap"]["clause"]
    cm = meta["commit"][:8]
    n_fix = len({x["correction"]["fix"] for x in corrs})
    n_flag = sum(1 for r in rows if r["flag"])
    L = [
        f"> 来源：内部方法论 {s['label']} 的「常见映射参考」表（第{s['table_header_line']}–{s['rows'][-1]['line']}行："
        f"表头、分隔行和 {len(s['rows'])} 个数据行；sha256 {s['sha256']}）；离线出清单工具 {c['label']} 中的 `_MAPPING_REF` 映射提示"
        f"（第{c['decl_line']}–{c['end_line']}行，只取此变量；文件 sha256 {c['sha256']}）。",
        f"> 改正与核对依据：zhitong-cost-ai 仓库 `{RULES_PATH}`（commit {cm}；sha256 {meta['rules_sha256']}）、"
        f"`{ZJ_PATH}`（同一 commit；sha256 {meta['zj_sha256']}）；逐行回放用同一 commit 的 `{ENGINE_PY}` 中 "
        f"check_consistency（sha256 {meta['engine_sha256']}）。",
        f"> 生成：{SCRIPT}，生成时间 {meta['generated_at']}",
        "> 声明：本文件由脚本从受控源确定性生成，不含任何金额；引用编码必须注明版本（浙江2025 / 国标2013）。",
        f"> 口径：系统口径参考表，待造价员签核。与引擎 quantity_rules 冲突处一律以 quantity_rules 为准，共改正 {n_fix} 处，"
        f"出处列写明原值；标【待核】的 {n_flag} 行保留原口径并给出码表中的备选编码，由造价员核定。签核前只作参考。",
        "> 版本：本文件所有编码、名称、计量单位均为浙江2025（GB50854 浙江省2025版），已与浙江2025码表逐一核对；"
        "「Z」开头的是浙江补充编码。同一 9 位编码在国标2013 中可能是另一个项目，不能混用。",
        "> 读法：每行一条映射，列为 描述特征｜部位｜编码｜浙江2025名称｜单位｜适用或禁用说明｜出处。"
        "「校核规则」是生成器用引擎一致性闸逐行回放时命中的规则，全部通过；规则口径见"
        f"《【清单·房屋建筑】{K6_TITLE}》。",
        "",
        f"# {TITLE}",
        "",
        f"## 映射参考表（{len(rows)} 行）",
        "",
        "| 描述特征 | 部位 | 编码 | 浙江2025名称 | 单位 | 适用或禁用说明 | 出处 |",
        "|---|---|---|---|---|---|---|",
    ]
    for r in rows:
        L.append("| " + " | ".join(cell(x) for x in (r["desc"], r["part"], r["code"], r["name"], r["unit"],
                                                      r["note"], r["src"])) + " |")
    L += [
        "",
        f"## 改正记录（{n_fix} 处）",
        "",
        "| 序号 | 改正事项 | 原文描述 | 原写编码 | 原写编码在浙江2025的名称 | 改为（浙江2025） | 依据 | 出处 |",
        "|---|---|---|---|---|---|---|---|",
    ]
    counter: dict[str, int] = {}
    for r in corrs:
        x = r["correction"]
        counter[x["fix"]] = counter.get(x["fix"], 0) + 1
    seen: dict[str, int] = {}
    for r in corrs:
        x = r["correction"]
        seen[x["fix"]] = seen.get(x["fix"], 0) + 1
        no = x["fix"][1:] + (f"-{seen[x['fix']]}" if counter[x["fix"]] > 1 else "")
        desc = re.sub(r"→.*$", "", x["orig_desc"])  # 原文描述只留左边的描述，不带原编码和原名称
        oname = x["orig_name_zj2025"] or "浙江2025码表无此编码"
        if x["orig"] == "010902003" and meta["snap"].get("gb2013_010902003"):
            g = meta["snap"]["gb2013_010902003"]
            oname += f"（国标2013 同号为{g['name']}，GB50854-2013.md 第{g['line']}行）"
        L.append("| " + " | ".join(cell(v) for v in (
            no, CORRECTIONS[x["fix"]], desc, x["orig"], oname, entry(r["code"], codes), x["basis"],
            x["origin"])) + " |")
    return "\n".join(L) + "\n"


# ----------------------------------------------------------------------------- 自检

def parse_tables(md: str) -> list[dict]:
    lines = md.split("\n")
    tables, i = [], 0
    split = lambda s: [c.strip() for c in re.split(r"(?<!\\)\|", s.strip())[1:-1]]  # noqa: E731
    while i < len(lines):
        if lines[i].startswith("|") and i + 1 < len(lines) and re.match(r"^\|(\s*-+\s*\|)+\s*$", lines[i + 1]):
            header = split(lines[i])
            rows, j = [], i + 2
            while j < len(lines) and lines[j].startswith("|"):
                rows.append((j + 1, split(lines[j])))
                j += 1
            tables.append({"header": header, "rows": rows})
            i = j
        else:
            i += 1
    return tables


def self_check(md: str, rows: list[dict], corrs: list[dict], meta: dict) -> dict:
    codes = meta["codes"]
    problems = list(meta["problems"])
    raw = md.encode("utf-8")
    if raw.startswith(b"\xef\xbb\xbf") or b"\r" in raw:
        problems.append("有 BOM 或 CR")
    heads = [l for l in md.split("\n") if l.startswith("#")]
    if len(heads) > 200:
        problems.append("标题超过 200")
    if MONEY_RE.search(md):
        problems.append(f"金额正则命中：{MONEY_RE.search(md).group(0)}")
    first_h = next(i for i, l in enumerate(md.split("\n")) if l.startswith("#"))
    pre = [l for l in md.split("\n")[:first_h] if l.strip()]
    if not pre or not all(l.startswith(">") for l in pre):
        problems.append("文件头不是首个标题前的引用块")
    # 1) 映射表逐行：编码在码表、名称单位一致
    tables = parse_tables(md)
    if len(tables) != 2:
        problems.append(f"表格数 {len(tables)} ≠ 2")
    t1 = tables[0]
    for ln, cs in t1["rows"]:
        if len(cs) != 7:
            problems.append(f"第{ln}行 {len(cs)} 列")
            continue
        c, n, u = cs[2], cs[3], cs[4]
        if c not in codes or codes[c]["name"] != n or codes[c]["unit"] != u:
            problems.append(f"第{ln}行 {c} {n} {u} 与码表不一致")
    for t in tables:
        for ln, cs in t["rows"]:
            if len(cs) != len(t["header"]):
                problems.append(f"第{ln}行单元格数不对（疑似未转义 |）")
            if t["header"][-1] != "出处" or not cs[-1] or cs[-1] == "—":
                problems.append(f"第{ln}行出处为空")
    # 2) 全文编码：除「原写 X」语境外，都要在码表；编码「名称」必须与码表一致
    orig_ok = {x["correction"]["orig"] for x in corrs}
    not_in = []
    for m in DIGITS_RE.finditer(md):
        s = m.group(0)
        z = m.start() > 0 and md[m.start() - 1] == "Z"
        code = ("Z" if z else "") + s
        if len(s) != 9 or code not in codes:
            not_in.append(code)
    bad_unknown = sorted(set(x for x in not_in if x not in orig_ok))
    if bad_unknown:
        problems.append(f"不在浙江2025码表的编码：{bad_unknown}")
    exempt_unknown = sorted(set(x for x in not_in if x in orig_ok))
    for m in re.finditer(r"(Z?\d{9})「([^」]+)」", md):
        c, n = m.groups()
        if c not in codes or codes[c]["name"] != n:
            problems.append(f"编码名称不一致：{c}「{n}」")
    # 3) 5 处改正落实 + 写明原值
    fixes = {}
    for r in corrs:
        fixes.setdefault(r["correction"]["fix"], []).append(r)
    if sorted(fixes) != ["C1", "C2", "C3", "C4", "C5"]:
        problems.append(f"改正处数不对：{sorted(fixes)}")
    want_new = {("C1", "Z010905001"), ("C1", "010905002"), ("C2", "Z011101006"), ("C2", "Z011201004"),
                ("C3", "010901008"), ("C4", "Z010902003"), ("C5", "011108003")}
    got_new = {(r["correction"]["fix"], r["code"]) for r in corrs}
    if got_new != want_new:
        problems.append(f"改正结果与方案口径不符：{sorted(got_new ^ want_new)}")
    for r in corrs:
        if f"原写 {r['correction']['orig']}" not in r["src"]:
            problems.append(f"改正行未写原值：{r['code']}")
    # 4) 回放 0 冲突
    viol = [(r["code"], r["desc"], v) for r in rows for v in r["replay"]["violations"]]
    if viol:
        problems.append(f"回放冲突 {len(viol)} 处：{viol[:5]}")
    # 5) 不得残留被改正的错误映射：改正后的描述不再映射到原值
    for r in rows:
        for x in corrs:
            if r["desc"] == x["desc"] and r["code"] == x["correction"]["orig"]:
                problems.append(f"残留错误映射：{r['desc']} → {r['code']}")
    fn = FILE_NAME
    if len(fn) > 60 or len(fn.encode("utf-8")) > 255 - 19 or re.search(r'[|/\\<>:"?*;,]', fn) or "待造价员签核" not in fn:
        problems.append("文件名不合规")
    return {
        "file": fn, "chars": len(md), "bytes": len(raw), "lines": md.count("\n"), "headings": len(heads),
        "rows": len(rows), "flags_pending_review": [(r["desc"], r["code"]) for r in rows if r["flag"]],
        "replay_combos_total": sum(r["replay"]["combos"] for r in rows),
        "replay_violations": len(viol),
        "rows_with_rules_applied": sum(1 for r in rows if r["replay"]["applied"]),
        "codes_not_in_zj2025_exempted_as_orig": exempt_unknown,
        "source_mode": meta["snap"]["_from"],
        "over_16000_chars": len(md) > 16000, "problems": problems,
    }


# ----------------------------------------------------------------------------- xlsx 审签件（只放本地）

def write_xlsx(path: Path, md: str, rows: list[dict], corrs: list[dict], meta: dict, codes: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    wrap = Alignment(wrap_text=True, vertical="top")
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="FFF2CC")

    def sheet(ws, header, data, widths):
        ws.append(header)
        for c in ws[1]:
            c.font = bold
            c.alignment = wrap
        for d in data:
            ws.append(d)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = wrap
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "映射参考表"
    hdr = ["序号", "描述特征", "部位", "编码", "浙江2025名称", "单位", "适用或禁用说明", "出处",
           "造价员意见（同意/修改/删除）", "修改后编码", "修改说明", "签核人", "签核日期"]
    sheet(ws, hdr, [[i + 1, r["desc"], r["part"], r["code"], r["name"], r["unit"], r["note"], r["src"],
                     "", "", "", "", ""] for i, r in enumerate(rows)],
          [6, 26, 18, 13, 20, 7, 50, 50, 16, 13, 24, 10, 12])
    dv = DataValidation(type="list", formula1='"同意,修改,删除"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"I2:I{len(rows) + 1}")
    for i, r in enumerate(rows, 2):
        if r["flag"] or r["correction"]:
            for c in ws[i]:
                c.fill = fill
    ws2 = wb.create_sheet("改正记录")
    tbl = parse_tables(md)[1]
    sheet(ws2, tbl["header"] + ["造价员意见", "签核人", "签核日期"],
          [[c.replace("\\|", "|") for c in cs] + ["", "", ""] for _, cs in tbl["rows"]],
          [6, 30, 30, 13, 26, 26, 30, 30, 16, 10, 12])
    ws3 = wb.create_sheet("说明")
    ws3.column_dimensions["A"].width = 160
    for l in md.split("\n"):
        if l.startswith(">"):
            ws3.append([l[1:].strip()])
    ws3.append([""])
    ws3.append(["黄色底纹行 = 已按 quantity_rules 改正或标【待核】的行，请重点签核。本审签件只在本地流转，不上传知识库。"])
    for row in ws3.iter_rows():
        for c in row:
            c.alignment = wrap
    wb.properties.creator = SCRIPT
    wb.properties.title = TITLE
    try:
        dt = datetime.fromisoformat(meta["generated_at"]).astimezone(timezone.utc).replace(tzinfo=None)
        wb.properties.created = dt
        wb.properties.modified = dt
    except Exception:
        pass
    wb.save(path)


def main() -> int:
    here = Path(__file__).resolve().parent
    ap = argparse.ArgumentParser()
    ap.add_argument("--repo", default=str(SNAPSHOT_REPO))   # 源快照根（见 kbsrc.py）
    ap.add_argument("--commit", default="33601ca8")
    ap.add_argument("--skill", default=None, help="SKILL.md 原文件（与 --clause 同给；不给则读快照）")
    ap.add_argument("--clause", default=None, help="clause_llm.py 原文件（只取 _MAPPING_REF；同包 app/config/ 一律不读）")
    ap.add_argument("--gb2013", default=GB2013_DEFAULT)
    ap.add_argument("--snapshot", default=str(here / "sources" / "k7_sources_snapshot.json"))
    ap.add_argument("--out", default=str(here / "out"))
    ap.add_argument("--local", default=str(here / "local_only"))
    a = ap.parse_args()
    md, rows, corrs, meta = build(a)
    rep = self_check(md, rows, corrs, meta)
    out = Path(a.out) / FILE_NAME
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_bytes(md.encode("utf-8"))
    rep["sha256"] = sha256(md.encode("utf-8"))
    rep["generated_at"] = meta["generated_at"]
    rep["commit"] = meta["commit"]
    loc = Path(a.local)
    loc.mkdir(parents=True, exist_ok=True)
    write_xlsx(loc / XLSX_NAME, md, rows, corrs, meta, meta["codes"])
    rep["xlsx"] = str(loc / XLSX_NAME)
    (loc / "K7_replay.json").write_text(json.dumps(
        [{"desc": r["desc"], "part": r["part"], "code": r["code"], "src": r["src"], **r["replay"]} for r in rows],
        ensure_ascii=False, indent=1), encoding="utf-8")
    (loc / "K7_corrections.json").write_text(json.dumps(
        [{"code": r["code"], "name": r["name"], "desc": r["desc"], **{k: v for k, v in r["correction"].items()}}
         for r in corrs], ensure_ascii=False, indent=1), encoding="utf-8")
    (loc / "K7_check.json").write_text(json.dumps(rep, ensure_ascii=False, indent=1), encoding="utf-8")
    print(json.dumps({k: rep[k] for k in ("file", "chars", "lines", "rows", "replay_combos_total",
                                          "replay_violations", "flags_pending_review",
                                          "codes_not_in_zj2025_exempted_as_orig", "source_mode", "sha256",
                                          "problems")}, ensure_ascii=False, indent=1))
    return 1 if rep["problems"] or rep["over_16000_chars"] else 0


if __name__ == "__main__":
    sys.exit(main())
