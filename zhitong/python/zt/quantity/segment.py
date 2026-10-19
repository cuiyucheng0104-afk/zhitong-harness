# 来源：engine-service/app/engine/quantity_segment.py（zhitong-cost-ai@991ff5cd，W4 搬入 zt；改动见本文件内「zt 改动」注释）
"""工程量清单生成：确定性切片（WS11 QTY-03 / QTY-13，issue #763）。零 LLM。

改前整份原文拼成一条 prompt 交给模型「自己切、自己归类、自己写位置」：模型漏掉的层在 items
和 unmatched 里都不出现（QTY-03），位置「Sheet1 第1行」是模型编的（QTY-13）。本模块把原文切成
片段（Segment），**原文和位置全部由代码产生**，LLM 只对片段编号作答（统一约定 C4：定位只能由
代码回填）。每个 practice 片段之后必须落入且只落入一种去向，由 quantity_list 做不变式核算。

三种输入形态：
- 网格（xlsx / xls）：openpyxl 读单元格与合并区域（经 WS08-A 的容错读取件，C6），跳过隐藏的表、
  行、列。识别含「做法」的表头行，按列跨度建立做法列组，列头决定类别；部位 / 名称 / 编号列出现
  新值或新合并块就开始一个新做法；列组内纯数字格是层号，文字格是层正文，无层号的续行并入上一层。
  表外的格按文本模式切，没有表头的表逐行处理。位置写真实坐标：`Sheet1!D6（第6行·楼地面做法第4层）`。
- docx：WS04 的块级 API（docx_reader.docx_blocks，C6）—— 段落按文本模式，表格按网格。
  位置写「表k 第r行第c列」「段落#k」。
- 文本（pdf / txt / OCR）：逐行处理，识别页锚（`[pN]`、旁车 `## pN`）和编号标题；同一行有多个
  「厚度+材料」时拆成多层；标记样式与本块首层不同的带号片（「2. 防水层：」下的「①」）是子项。
  位置写「pN 第k行（抽取文本行）」。

范围排除的顺序（WS29-f，issue #835）：**先切层、后判排除**。一行或一格先按层号切开；成层的（带层号的
多层行、带厚度与材料的散行、做法名称下的编号层、表格做法列里的格）逐层产出 practice 片段，层正文里的
「用户自理」「为二次装修」「详精装修设计」只是这一层的标记（`Segment.inline`），不产生全局排除句。
只有不成层的句子（说明段落、表注、独立文字行）才有资格当全局排除句。改前整行先判排除，命中就整行
变成一个 exclusion 片段：一层写了「自理」，同一行的抗裂砂浆、保温板、防水涂料跟着从覆盖核算里消失。
"""
from __future__ import annotations

import logging
import re
import unicodedata
from dataclasses import dataclass, field
from typing import Iterable, Iterator, Optional

from zt.quantity import rules

log = logging.getLogger("zt.quantity.segment")

PRACTICE, HEADING, EXCLUSION, OTHER = "practice", "heading", "exclusion", "other"
# 同一单元格 / 段落 / 文本行（「所在格」）里按层号切出的各片的角色（R4F-d fix14 → fix15，存疑闸的同格作用域）：
# HOST_LAYER = 构造层自己的文字（成片的做法层、排除层、带层号的片 —— 带层号却没认成做法层的「4)清理基层，除锈…」
# 也是这一格里的一层）；HOST_FREE = 不属于任何层的文字（做法名称 / 编号行与批注、行首空框、层号前的「以下做法取消：」
# 「（取消）」、并进上一层的悬空续句）
HOST_LAYER, HOST_FREE = "layer", "free"


@dataclass
class HostPiece:
    start: int                    # 在所在格原文里的 [start, end)；定位不到时 (-1, -1)
    end: int
    role: str = HOST_FREE


@dataclass
class Host:
    """一个所在格（单元格 / 段落 / 文本行）切层之前的原文，以及这一格切出的各片。只记录，不改切片结果。

    R4F-d fix15（用户 2026-10-12 决定 1）：xlsx 单元格里的多行共用一个 Host（格内换行不切断所在格，只是句界）；
    docx 单元格的 text 是与格文字逐字等长的格内原文（格内另起一段 / 软回车处是换行，即合同句界）；做法名称行与层
    同段 / 同格时，名称、编号、批注与行首空框记成这一格里的 HOST_FREE 片；无表头表格一行拼成一句时每格一个 Host
    （fix16：片段挂在带材料词的那一格上，见 _own_hosts）。"""
    text: str
    where: str                    # 单元格 / 段落 / 行
    pieces: list[HostPiece] = field(default_factory=list)

    def add(self, start: int, end: int, role: str = HOST_FREE) -> int:
        self.pieces.append(HostPiece(start, end, role))
        return len(self.pieces) - 1

    def mark(self, k: int, role: str) -> None:
        if 0 <= k < len(self.pieces):
            self.pieces[k].role = role

    def attach(self, seg: "Segment", k: int, role: str) -> None:
        if 0 <= k < len(self.pieces):
            self.pieces[k].role = role
            seg.host, seg.host_piece = self, k

    def covers(self, a: int, b: int) -> bool:
        return any(p.start >= 0 and p.start < b and a < p.end for p in self.pieces)


_HOST_SPLIT = re.compile(r"[，,；;。]")


def _locate(t: str, body: str, cur: int) -> tuple[list[tuple[int, int, int]], int]:
    """片正文 body 按 ，,；;。 切成的各段（去首尾空白）在 t 里从 cur 起依次定位 → ([(t 起点, body 起点, 长度)], 新游标)。
    有一段找不到就返回 ([], cur)。逐段定位：按厚度切开的片在正文里用「，」重新连接，原文里的分隔符可能不同。"""
    out: list[tuple[int, int, int]] = []
    pos, bpos = cur, 0
    for c in _HOST_SPLIT.split(body or ""):
        cs = c.strip()
        if not cs:
            continue
        bat = (body or "").find(cs, bpos)
        tat = t.find(cs, pos)
        if bat < 0 or tat < 0:
            return [], cur
        out.append((tat, bat, len(cs)))
        bpos, pos = bat + len(cs), tat + len(cs)
    return out, (pos if out else cur)


@dataclass
class _Placement:
    """一片在所在格里的落点：hosts = [(host, 该片在 host.pieces 里的下标)]，第一个是锚（片正文开头所在的格）；
    chunks = 片正文逐段落在哪一格 [(host, 正文起点, 长度)]（无表头表格一行拼成一句时，一片可能跨格）。"""
    hosts: list[tuple[Host, int]]
    chunks: list[tuple[Host, int, int]]

    def set_role(self, role: str) -> None:
        for h, k in self.hosts:
            h.mark(k, role)


class _Mapper:
    """行坐标 → 所在格坐标。spans = [(行起, 行止, host, 偏移)]：行坐标 x ∈ [行起, 行止) 落在 host 的 x + 偏移 处。
    普通段落 / 文本行 / 单元格是一个 span；无表头表格一行拼成的一句是每格一个 span（格间连接空格不属于任何格）。"""

    def __init__(self, spans: list[tuple[int, int, Host, int]]):
        self.spans = spans

    def at(self, x: int) -> tuple[Optional[Host], int]:
        for a, b, h, off in self.spans:
            if a <= x < b:
                return h, x + off
        return None, -1

    def _runs(self, x: int, n: int, body_at: int) -> list[list]:
        """行坐标 [x, x+n) → [[host, host 起点, 长度, 正文起点]]（按所在格切开）。"""
        out: list[list] = []
        for k in range(max(0, n)):
            h, hx = self.at(x + k)
            if h is None:
                continue
            if out and out[-1][0] is h and out[-1][1] + out[-1][2] == hx:
                out[-1][2] += 1
            else:
                out.append([h, hx, 1, body_at + k])
        return out

    def place(self, t: str, base: int, pieces, start: int = 0) -> list[Optional[_Placement]]:
        """t 是行里从行坐标 base 起的一段；pieces = split_layers_marked(…) 的片 → 每片的落点（定位不到为 None）。
        各片按顺序从 t[start:] 里依次定位。"""
        out: list[Optional[_Placement]] = []
        cur = start
        for _no, body, _style in pieces:
            chunks, cur2 = _locate(t, body or "", cur)
            if not chunks:
                out.append(None)
                continue
            cur = cur2
            ext: dict[int, list[int]] = {}
            order: list[Host] = []
            pcs: list[tuple[Host, int, int]] = []
            for tat, bat, n in chunks:
                for h, hx, ln, b0 in self._runs(base + tat, n, bat):
                    if id(h) not in ext:
                        ext[id(h)] = [hx, hx + ln]
                        order.append(h)
                    else:
                        ext[id(h)][0] = min(ext[id(h)][0], hx)
                        ext[id(h)][1] = max(ext[id(h)][1], hx + ln)
                    pcs.append((h, b0, ln))
            if not order:
                out.append(None)
                continue
            out.append(_Placement([(h, h.add(ext[id(h)][0], ext[id(h)][1])) for h in order], pcs))
        return out

    def free(self, a: int, b: int) -> None:
        """行坐标 [a, b) 这段文字（做法名称 / 编号、批注、行首空框）记成所在格里的 HOST_FREE 片（按格切开、去首尾空白）。"""
        for h, hx, ln, _b0 in self._runs(a, b - a, 0):
            part = h.text[hx:hx + ln]
            if part.strip():
                lead = len(part) - len(part.lstrip())
                h.add(hx + lead, hx + len(part.rstrip()))


def _layer_role_words() -> tuple[str, ...]:
    g = rules.load_rules().get("granularity") or {}
    return tuple(g.get("standalone") or ()) + tuple(g.get("aux") or ())


def _own_hosts(pl: _Placement) -> list[tuple[Host, int]]:
    """一片跨了几格（无表头表格一行拼成一句）时，哪几格是本层自己的格（R4F-d fix16 → fix17，决定 1：「它所在单元格」是本层
    描述所在的格）。fix17（协调者裁定 5，review17b 反向探针 rv-x03 的口径）：按「完整层描述」认——锚格先认带厚度 / 强度的格，
    没有再认带材料词的格；锚格还缺材料词、或缺层名词而下一格以层名词（保温层 / 面层 …）开头时，把后面相邻的格并进来
    （「1｜60厚｜岩棉板保温层（取消）」）；往后并完还缺材料词时再看前面相邻的格。备注格只是提到了某种材料（「砂浆禁止暴晒」
    「墙面砖不做」）不再算本层格（review17b R17B-02）。首格只是层号「1」「（1）」、字母「A」、层名「第1层」时不当锚
    （review16a R16A-01）。一格都认不出时整片各格都算（拿不准时取宽）。"""
    if len(pl.hosts) <= 1:
        return list(pl.hosts)
    parts = []
    for h, k in pl.hosts:
        piece = h.pieces[k] if 0 <= k < len(h.pieces) else None
        parts.append(h.text[piece.start:piece.end] if piece is not None and piece.start >= 0 else "")
    sigs = [rules.practice_signal(p) for p in parts]
    anchor = next((i for i, s in enumerate(sigs) if s.thickness or s.strength), None)
    if anchor is None:
        anchor = next((i for i, s in enumerate(sigs) if s.materials), None)
    if anchor is None:
        return list(pl.hosts)
    roles = _layer_role_words()
    own = [anchor]
    acc = parts[anchor]
    j = anchor + 1
    while j < len(parts):
        if not rules.practice_signal(acc).materials or (
                not any(w in acc for w in roles) and any(parts[j].strip().startswith(w) for w in roles)):
            own.append(j)
            acc += parts[j]
            j += 1
        else:
            break
    j = anchor - 1
    while j >= 0 and not rules.practice_signal(acc).materials:
        own.insert(0, j)
        acc = parts[j] + acc
        j -= 1
    # R4F-d fix18（裁定 5「按完整层描述认本层格」，review18c R18C-02）：锚格（规格格）前面紧挨着的一格带材料词时，那是本层的
    # 主材格（「1｜加气混凝土砌块墙（取消）｜200厚，M5混合砂浆砌筑」：锚格里的「混合砂浆」只是砌筑砂浆），并进本层格；备注格
    # 在锚格后面（「1｜60厚岩棉板保温层｜砂浆禁止暴晒」），不受影响
    j = own[0] - 1
    if j >= 0 and rules.practice_signal(parts[j]).materials:
        own.insert(0, j)
    return [pl.hosts[i] for i in own]


def _attach(seg: "Segment", pl: Optional[_Placement], role: str, fallback: Optional[Host] = None) -> None:
    """片段挂到它的所在格（锚格 = 本层材料提及所在的格）；片正文里落在别的格的部分（无表头表格一行拼成一句时跨格的片：
    层号格、部位格、备注格）记进 out_ranges。定位不到时挂在 fallback 格上、作用域取整格（拿不准时取宽）。"""
    if pl is None:
        if fallback is not None:
            seg.host, seg.host_piece = fallback, fallback.add(0, len(fallback.text), HOST_LAYER)
        return
    pl.set_role(role)
    own = _own_hosts(pl)
    h, k = own[0]
    seg.host, seg.host_piece = h, k
    keep = {id(x) for x, _k in own}
    seg.out_ranges = [(b, b + n) for hh, b, n in pl.chunks if id(hh) not in keep]


def _merge(last: "Segment", old_len: int, pl: Optional[_Placement]) -> None:
    """无层号的续行并进上一层（last.text = f"{旧} {续}"）：续行里来自另一个**单元格**的部分不进 last 的作用域
    （R4F-d fix15 决定 1：另一单元格的文字不得因「续行」并进本层作用域，保留来源单元格边界），记进 out_ranges；
    同一单元格里的续行（xlsx 格内多行、docx 格内另起一段）仍在作用域内。段落 / 抽取文本行的续行（引擎认定是本层
    描述的下一段、PDF 折行）不是单元格，照 fix14 留在本层原文里（拿不准时取宽）。"""
    if pl is None:
        return
    pl.set_role(HOST_FREE)
    off = old_len + 1
    for hh, b, n in pl.chunks:
        if hh is not last.host and hh.where == "单元格":
            last.out_ranges.append((off + b, off + b + n))


@dataclass
class Segment:
    kind: str
    text: str
    file_index: int
    file_name: str
    location: str
    sheet: Optional[str] = None
    cell: Optional[str] = None
    row: Optional[int] = None
    page: Optional[int] = None
    paragraph: Optional[int] = None
    line: Optional[int] = None
    category: str = ""
    group: str = ""               # 做法列组的列头原文（「楼地面做法」）
    part: str = ""
    # 无表头做法表的「适用部位」栏（WS29-k1）：取自哪一格、未截断的原文。part 按约定截到 60 字给人看，
    # 精装修引用和「范围相关原文」要看全文，截掉的部分不能悄悄丢（铁律 8）
    part_location: str = ""
    part_full: str = ""
    practice: str = ""
    practice_key: str = ""
    layer_no: Optional[int] = None
    heading: str = ""
    section_domain: Optional[str] = None   # 所在章节只含非房建词时的领域（QTY-07 章节级）
    exclusions: list = field(default_factory=list)   # kind=exclusion：[(kind, objects, lot)]
    # 层内标记（只管这一层）：自理 / 本层为二次装修 / 精装修引用。mark_inline 在续行并完之后统一填
    inline: Optional[rules.InlineMarker] = None
    # kind=other 且落在做法块里（做法名称行 / 标题行开头的做法、表格的做法列组）：所在做法的名称。
    # 这些行不是做法层、不进覆盖核算，改前哪一页都查不到（外露铁件的「除锈」、整组漆层，E31）；
    # 清单把它们列进「说明」页的「块内未识别行」（WS29-k2，铁律 8）。None = 不是做法块里的行
    block_practice: Optional[str] = None
    id: str = ""
    order: int = 0
    # 层正文所在的格 / 段 / 行与本片是其中第几片（R4F-d fix14：存疑闸看同格里不属于任何层的文字、带明示指代的
    # 另一层）。不参与切片、去向与任何输出字段；其他路径造出的片段为 None
    host: Optional[Host] = None
    host_piece: int = -1
    # R4F-d fix15：片段原文里来自另一单元格的部分 [起, 止)（续行并入、无表头表格一行拼成一句时跨格）：存疑闸不看
    out_ranges: list = field(default_factory=list)

    def scope_text(self) -> str:
        """存疑闸看的本行原文：来自另一单元格的部分换成等长空白（位置不变，原因栏「原文第 N 字」仍按全文计）。"""
        if not self.out_ranges:
            return self.text
        chars = list(self.text)
        for a, b in self.out_ranges:
            for i in range(max(0, a), min(len(chars), b)):
                chars[i] = " "
        return "".join(chars)

    def context(self) -> str:
        return " ".join(x for x in (self.category, self.group, self.part, self.practice, self.heading) if x)

    def evidence(self, rel_path: Optional[str] = None, file_id: Optional[str] = None) -> dict:
        ev = {"file_id": file_id, "rel_path": rel_path or self.file_name, "file_name": self.file_name,
              "sheet": self.sheet, "cell": self.cell, "row": self.row, "page": self.page,
              "paragraph": self.paragraph, "heading": self.heading or None, "quote": self.text[:200]}
        return {k: v for k, v in ev.items() if v is not None}


@dataclass
class FileSegments:
    file_index: int
    file_name: str
    mode: str
    segments: list[Segment]
    warnings: list[tuple[str, str]] = field(default_factory=list)   # (code, message)
    hidden_sheets: list[str] = field(default_factory=list)


# ============================================================ 通用识别

_CATEGORY_WORDS: list[tuple[str, tuple[str, ...]]] = [
    ("踢脚", ("踢脚",)),
    ("屋面", ("屋面", "屋顶")),
    ("顶棚", ("顶棚", "天棚", "吊顶", "平顶")),
    ("墙面", ("内墙", "外墙", "墙面", "墙裙", "墙身")),
    ("楼梯", ("楼梯",)),
    # 台阶、坡道、散水各是各的类别（WS29-k2）。改前并成一个「台阶坡道」：坡道的每一层上下文里都带着「台阶」，
    # 一致性闸认不出「坡道给了台阶的码」，台阶的每一层又都带着「坡道」；模型看到的类别也是「台阶坡道」（E30）
    ("台阶", ("台阶",)),
    ("坡道", ("坡道",)),
    ("散水", ("散水",)),
    ("楼地面", ("楼地面", "地面", "楼面", "地坪")),
]


def category_of(text: str) -> str:
    t = text or ""
    for name, words in _CATEGORY_WORDS:
        if any(w in t for w in words):
            return name
    return ""


_NAME_HEAD = r"(?:楼地面|地面|楼面|屋面|内墙面?|外墙面?|墙面|墙裙|顶棚|天棚|吊顶|平顶|踢脚线?|台阶|散水|坡道|楼梯|雨篷|女儿墙|地防)"
# 做法编号前的项目符号（「■ 楼面9」）：做法表常用它标做法起头，名称识别要先跳过它（WS29-k1）
_LEAD_SYMBOLS = "■□●◆▲◇○·"
_PRACTICE_NAME = re.compile(
    r"^[\s" + _LEAD_SYMBOLS + r"]*(?P<name>" + _NAME_HEAD
    + r"\s*[-_]?\s*[0-9一二三四五六七八九十]{1,3}[A-Za-z]?)(?![0-9])\s*(?P<rest>.*)$")
_COLON_NAME = re.compile(r"^\s*(?P<name>[^：:，,。；;]{2,24})[：:]\s*(?P<rest>.*)$")
_PART = re.compile(r"(?:用于|适用于|使用部位|适用部位|部位)\s*[:：]?\s*(?P<part>[^，,。；;）)]{1,30})")
_PAREN_PART = re.compile(r"[（(](?P<p>[^）)]{1,24})[）)]")

_HEADING = re.compile(
    r"^\s*(?:第[一二三四五六七八九十0-9]+[章节条部分篇]|[（(][一二三四五六七八九十]{1,3}[)）]|[一二三四五六七八九十]{1,3}\s*[、.．]"
    r"|\d{1,2}(?:\.\d{1,2}){1,3}(?![\d.])\s*[、.．]?)\s*(?P<title>.{1,40})$")
_ARABIC_HEAD = re.compile(r"^\s*(?P<n>\d{1,2})\s*[、.．]\s*(?P<title>\D.{0,24})$")
# 做法块里「N.xxx」是层还是节标题：节名尾缀是强信号（「外墙保温做法」含材料词也是节名）；
# 引用句、施工动作是层的写法（「详见05J1工程做法」「刷基层处理剂一遍」）
_SECTION_SUFFIX = re.compile(r"(?:做法|工程|构造)\s*(?:[（(][^）)]{0,16}[）)])?\s*[:：]?\s*$")
_LAYER_WORDING = re.compile(r"详见|参见|参照|按照|[一二两三四]\s*[道遍]|夯实")
# 表注的开头：「注：」单独一行，或后面直接跟编号条款（「注：1. …」）。「注：面层为二次装修」这种
# 夹在两层之间的一句话不算 —— 它后面还是这个做法的层
_NOTE_LIST = re.compile(
    r"^\s*(?:注|附注|备注|说明)\s*[:：]\s*(?:$|\d{1,2}\s*[、.．)）]|[（(]\d{1,2}[)）]|[①-⑳])")


def _section_title(title: str) -> bool:
    t = (title or "").strip()
    if not t or len(t) > 24:
        return False
    sig = rules.practice_signal(t)
    if sig.thickness or sig.strength or _LAYER_WORDING.search(t):
        return False
    if _SECTION_SUFFIX.search(t):
        return True
    lw = rules.layer_words(t)
    return not (sig.materials or lw.standalone or lw.aux or lw.base)

_PAGE_ANCHOR = re.compile(r"^\s*(?:\[p(?P<a>\d+)(?:[–\-]p?\d+)?\]|##\s*p(?P<b>\d+)\b|[-—=]*\s*第\s*(?P<c>\d+)\s*页\s*[-—=]*$)")

# 编号层标记：「1、」「1.」「1）」「①」；「1.5厚」这类小数不是标记（点后只有一位数字接「厚」）
_THK = r"\s*(?:mm|MM|毫米)?\s*厚"
_MARK = re.compile(
    r"(?:^|(?<=[\s，,；;。]))(?:(?P<n>\d{1,2})\s*(?:[、)）]|[.．](?=\s|[^\d\s]|\d{2,}" + _THK
    + r"|\d+\.\d+" + _THK + r"))|(?P<c>[①-⑳]))")
_THICK_CHUNK = re.compile(r"^\s*\d+(?:\.\d+)?\s*(?:mm|MM|毫米)?\s*厚")
_LAYER_NUM_CELL = re.compile(r"^\s*\d{1,2}\s*$")


def _norm_text(s: str) -> str:
    # 只去首尾空白：片段原文要逐字进 quote 与「原文」列（C4），全角标点不能被 NFKC 改掉
    return (s or "").strip()


def _forced_mark(s: str, expect_next: Optional[int]) -> Optional[re.Match]:
    """「3.3厚SBS」这类歧义写法：做法块里正等着第 3 层时，开头的「3.」按层号读（否则是 3.3 厚）。"""
    if expect_next is None:
        return None
    return re.match(rf"\s*(?P<n>{int(expect_next)})\s*[.．](?=\d+(?:\.\d+)?" + _THK + ")", s)


# 层号标记的样式：「1.」「1．」→ dot，「1、」→ dun，「1)」「1）」→ paren，「①」→ circle。
# 做法说明里一级层和它下面的子项用不同样式（「2. 防水层：」下接「①…②…」），文本模式靠它分辨子项
DOT, DUN, PAREN, CIRCLE = "dot", "dun", "paren", "circle"


def _mark_style(m: re.Match) -> str:
    if m.group("c"):
        return CIRCLE
    sep = m.group(0)[-1]          # _MARK 的分隔符是匹配的最后一个字符（「.」后面只是前瞻）
    return DUN if sep == "、" else PAREN if sep in ")）" else DOT


def split_layers_marked(text: str, expect_next: Optional[int] = None
                        ) -> list[tuple[Optional[int], str, Optional[str]]]:
    """split_layers 带上标记样式：[(层号, 层正文, 样式)]，无层号的片样式为 None。"""
    s = (text or "").strip()
    if not s:
        return []
    marks = list(_MARK.finditer(s))
    forced = _forced_mark(s, expect_next)
    if forced is not None and (not marks or marks[0].start() > 0):
        rest = split_layers_marked(s[forced.end():], None)
        if rest and rest[0][0] is None:
            rest[0] = (int(forced.group("n")), rest[0][1], DOT)
        elif not rest:
            return []
        else:
            rest.insert(0, (int(forced.group("n")), "", DOT))
            rest = [x for x in rest if x[1]]
        return rest
    pieces: list[tuple[Optional[int], str, Optional[str]]] = []
    if marks:
        head = s[: marks[0].start()].strip(" ，,；;。")
        if head:
            pieces.append((None, head, None))
        for i, m in enumerate(marks):
            end = marks[i + 1].start() if i + 1 < len(marks) else len(s)
            no = int(m.group("n")) if m.group("n") else (ord(m.group("c")) - ord("①") + 1)
            body = s[m.end():end].strip(" ，,；;。")
            if body:
                pieces.append((no, body, _mark_style(m)))
    else:
        pieces.append((None, s, None))
    out: list[tuple[Optional[int], str, Optional[str]]] = []
    for no, body, style in pieces:
        chunks = [c.strip() for c in re.split(r"[，,；;。]", body) if c.strip()]
        thick = [i for i, c in enumerate(chunks) if _THICK_CHUNK.match(c)]
        if len(thick) < 2:
            out.append((no, body, style))
            continue
        cur: list[str] = []
        first = True
        for i, c in enumerate(chunks):
            if i in thick and cur:
                out.append((no, "，".join(cur), style) if first else (None, "，".join(cur), None))
                first = False
                cur = []
            cur.append(c)
        if cur:
            out.append((no, "，".join(cur), style) if first else (None, "，".join(cur), None))
    return out


def split_layers(text: str, expect_next: Optional[int] = None) -> list[tuple[Optional[int], str]]:
    """一行原文 → [(层号, 层正文)]。先按编号标记切，再把含多个「厚度+材料」的片按厚度切开。

    expect_next：所在做法块下一层的层号（没有块时 None），用来消解「3.3厚」这类歧义开头。
    """
    return [(no, body) for no, body, _style in split_layers_marked(text, expect_next)]


# ============================================================ 文本模式

@dataclass
class Line:
    text: str
    location: str = ""
    sheet: Optional[str] = None
    cell: Optional[str] = None
    row: Optional[int] = None
    paragraph: Optional[int] = None
    page_aware: bool = False
    force_heading: bool = False
    # R4F-d fix15（只给存疑闸的作用域用，不参与切片）：host = 本行所在格的共用 Host（xlsx 格内多行、表外格的多行），
    # host_offset = 本行 text 在 host.text 里的起点；scope_text = 与 text 逐字等长的格内原文（docx 格内硬换行处是换行）；
    # cells = 无表头表格一行拼成一句时各格在 text 里的 [起, 止)
    host: Optional[Host] = None
    host_offset: int = 0
    scope_text: Optional[str] = None
    cells: Optional[list] = None


def _host_where(line: Line) -> str:
    """所在格的称呼（存疑闸原因栏用）：docx 段落 → 段落；表格单元格 → 单元格；抽取文本 → 行。"""
    if line.paragraph is not None:
        return "段落"
    if line.cell or "列" in (line.location or ""):
        return "单元格"
    return "行"


def _line_mapper(line: Line, raw: str, lead: int, s: str) -> tuple[_Mapper, bool]:
    """一行的行坐标 → 所在格坐标（见 _Mapper）。返回 (mapper, 是否共用格)。"""
    where = _host_where(line)
    src = (line.scope_text if line.scope_text is not None else line.text).replace("\r", "")
    if len(src) != len(raw):
        src = raw
    if line.host is not None:
        return _Mapper([(0, len(raw), line.host, line.host_offset)]), True
    if line.cells:
        return _Mapper([(a, b, Host(src[a:b], where), -a) for a, b in line.cells if b > a]), False
    return _Mapper([(lead, lead + len(s), Host(src[lead:lead + len(s)], where), -lead)]), False


def _name_tail(rest: str) -> tuple[str, str]:
    """做法名称行名称后面的部分 → (标题, 直接跟着的层)。「屋面1：40厚C20细石混凝土，10厚…」的层从第一个
    层号或厚度起算；没有就全是标题。"""
    title = rest.strip().lstrip("：: ")
    mk = _MARK.search(title)
    th = re.search(r"\d+(?:\.\d+)?\s*(?:mm|MM|毫米)?\s*厚", title)
    cut = min([x for x in (mk.start() if mk else None, th.start() if th else None) if x is not None],
              default=None)
    if cut is None:
        return title, ""
    return title[:cut].strip(" ：:，,"), title[cut:]


_R_HEADING, _R_NAME, _R_LAYERS = "heading", "name", "layers"


@dataclass(frozen=True)
class _Route:
    """一行的去向（_TextSegmenter._route）。layers：这一行（或名称行名称后面的部分）会进 _layers。"""
    kind: str
    title: str = ""
    name: str = ""
    rest: str = ""
    layers: bool = False


class _TextSegmenter:
    def __init__(self, file_index: int, file_name: str, out: list[Segment]):
        self.fi = file_index
        self.fname = file_name
        self.out = out
        self.page: Optional[int] = None
        self.line_no = 0
        self.heading = ""
        self.heading_domain: Optional[str] = None
        self.heading_category = ""
        self.in_block = False
        self.block_numbered = False
        self.last_no: Optional[int] = None
        # 本块一级层号的标记样式（首个带号的层定）。样式不同的带号片是当前层的子项（「2. 防水层：」下的
        # 「①…」「1、…」「1）…」）：不推进层号，也不因「回到 1」另起做法（PR #766 第二轮评审阻断项）
        self.block_style: Optional[str] = None
        # 本块由做法名称行开头（「楼面2：…」）。名称下面的编号行是这个做法的层，不是说明条款
        self.block_named = False
        self.practice = ""
        self.part = ""
        self.category = ""
        self.practice_key = ""
        self.last: Optional[Segment] = None
        # 紧挨着的上一片是「N.外墙面」这类按层号收下、但读着像节名的片段：(片段, 整行, 标题)。
        # 只收 OTHER —— 事后改判成标题时，不会有做法片段从覆盖核算里消失（QTY-03，铁律 8）
        self.head_cand: Optional[tuple[Segment, str, str]] = None
        self._auto = 0
        # 正在处理的这一行：行坐标 → 所在格坐标（R4F-d fix15，存疑闸作用域）
        self._map: Optional[_Mapper] = None
        self._s = ""
        self._lead = 0

    # —— 位置 ——
    def _loc(self, line: Line) -> dict:
        if line.page_aware:
            where = f"第{self.line_no}行（抽取文本行）"
            loc = f"p{self.page} {where}" if self.page is not None else where
            return {"location": loc, "page": self.page, "line": self.line_no}
        return {"location": line.location, "sheet": line.sheet, "cell": line.cell, "row": line.row,
                "paragraph": line.paragraph}

    def _emit(self, kind: str, text: str, line: Line, *, layer_no: Optional[int] = None,
              exclusions: Optional[list] = None) -> Segment:
        self.head_cand = None
        seg = Segment(kind=kind, text=text, file_index=self.fi, file_name=self.fname,
                      heading=self.heading, section_domain=self.heading_domain,
                      exclusions=list(exclusions or []), **self._loc(line))
        if kind == PRACTICE:
            seg.category = self.category or self.heading_category
            seg.part = self.part
            seg.practice = self.practice
            seg.practice_key = self.practice_key
            seg.layer_no = layer_no
        self.out.append(seg)
        return seg

    def _start_practice(self, name: str, category: str, part: str, key_hint: str, *,
                        named: bool = False) -> None:
        self.in_block = True
        self.block_named = named
        self.block_numbered = False
        self.last_no = None
        self.block_style = None
        self.practice = name
        self.category = category
        self.part = part
        self._auto += 1
        self.practice_key = f"f{self.fi}:{key_hint}:{self._auto}"
        self.last = None
        self.head_cand = None

    def end_block(self) -> None:
        self.in_block = False
        self.block_named = False
        self.block_numbered = False
        self.last_no = None
        self.block_style = None
        self.last = None
        self.head_cand = None

    def expect_next(self) -> Optional[int]:
        if not self.in_block:
            return None
        return (self.last_no or 0) + 1

    def set_heading(self, title: str) -> None:
        self.heading = title
        self.heading_domain = rules.heading_scope(title)
        self.heading_category = category_of(title)
        self.end_block()

    # —— 入口 ——
    def feed(self, line: Line) -> None:
        raw = line.text.replace("\r", "")
        if line.page_aware:
            m = _PAGE_ANCHOR.match(raw)
            if m:
                self.page = int(m.group("a") or m.group("b") or m.group("c"))
                self.line_no = 0
                return
            self.line_no += 1
        s = raw.strip()
        if not s:
            return
        lead = len(raw) - len(raw.lstrip())
        mapper, shared = _line_mapper(line, raw, lead, s)
        self._map, self._s, self._lead = mapper, s, lead
        try:
            self._feed(s, line)
        finally:
            if shared:
                h, x = mapper.at(lead)
                if h is not None and not h.covers(x, x + len(s)):
                    # 本行没切出任何片（名称行、说明句、排除句、标题）：整行是这一格里不属于任何层的文字
                    h.add(x, x + len(s))
            self._map = None

    def _feed(self, s: str, line: Line) -> None:
        if self.block_named and _NOTE_LIST.match(s):
            # 表注开始：后面的「1. …」「2. …」是表注条款，不再是做法名称下面的层，层号恰好接得上也不是
            #（PR #839 评审阻断项）。不结束做法块：结束了，表注后面真正的层会丢掉做法名称和类别
            self.block_named = False
        route = self._route(s, line)
        layer_row, numbered, own = self._layer_row(s)
        # 先切层、后判排除（WS29-f 第零条）：成层的行不整行判，整行判会让同一行的其它层一起消失。
        # 「不整行判」必须同时满足「这一行真的进了 _layers」：判成层却落进标题 / 名称行分支的句子，
        # 既不成全局排除句、也不成层，哪一页都查不到（QTY-11 回归 + 铁律 8，PR #839 评审阻断项）
        if not (own and route.layers):
            scan = rules.scan_exclusions(s)
            if scan.found:
                self._emit(EXCLUSION, s, line, exclusions=list(scan.found))
                self.last = None
                return
            if scan.rejected and not numbered and not line.force_heading and not _PRACTICE_NAME.match(s):
                # 句式是排除句、对象词却不成立（只剩「面层」这类泛指词）：它是一句说明，不是上一层的续行。
                # 放下去会被并进上一层，上一层就凭空多出一个「为二次装修」的标记
                self._emit(OTHER, s, line)
                self.last = None
                return
        if route.kind == _R_HEADING:
            self._emit(HEADING, s, line)
            self.set_heading(route.title)
            return
        if route.kind == _R_NAME:
            self._name_line(route.name, route.rest, line)
            return
        self._layers(s, line, layer_row)

    def _route(self, s: str, line: Line) -> _Route:
        """这一行走哪条分支。只读状态、不改状态：feed 先拿它定「要不要整行判全局排除句」，再照它分派，
        两处用的是同一个结论。"""
        if line.force_heading:
            return _Route(_R_HEADING, title=s)
        m = _PRACTICE_NAME.match(s)
        if m:
            return _Route(_R_NAME, name=m.group("name"), rest=m.group("rest"),
                          layers=bool(_name_tail(m.group("rest"))[1]))
        nxt = self.expect_next()
        am = _ARABIC_HEAD.match(s)
        if am and self.in_block and len(s) <= 26 and _section_title(am.group("title")) and (
                int(am.group("n")) != nxt or _SECTION_SUFFIX.search(am.group("title").strip())):
            # 做法块里的「7.楼面做法」：层号对不上下一层、正文又不像层，是节标题。改前块内一律按层读，
            # 后面各节的层全并进第一个做法，同码跨节合并、类别错配（QTY-03，PR #766 评审）
            return _Route(_R_HEADING, title=am.group("title").strip())
        if nxt is not None and re.match(
                rf"\s*{nxt}\s*(?:[、)）]|[.．](?!\d)|[.．](?=\d+(?:\.\d+)?" + _THK + "))", s):
            # 做法块里等着的下一层：「1.9.5厚纸面石膏板」「3.3厚SBS」不是多级标题
            return _Route(_R_LAYERS, layers=True)
        hm = _HEADING.match(s)
        if hm and not rules.thickness_starts(hm.group("title")) and not rules.practice_signal(s).thickness:
            return _Route(_R_HEADING, title=hm.group("title").strip())
        if (am and not self.in_block and not rules.practice_signal(s).thickness
                and not rules.practice_signal(s).strength and len(s) <= 26):
            return _Route(_R_HEADING, title=am.group("title").strip())
        cm = _COLON_NAME.match(s)
        if cm and len(s) <= 30 and not cm.group("rest").strip() and category_of(cm.group("name")):
            return _Route(_R_NAME, name=cm.group("name"), rest="")
        return _Route(_R_LAYERS, layers=True)

    def _layer_row(self, s: str) -> tuple[bool, bool, bool]:
        """这一行成不成层 → (成层, 带层号, 层里的自理 / 二次装修只管这一层)。

        成层：一行切出两个以上带号的层（做法表的「1)… 2)… 3)…」）；或带厚度（强度）与材料词的散行；
        或做法名称行下面的编号行。只带一个层号、又不像做法层的行（「3、卫生间洁具由业主另行发包」）
        是说明条款，仍按全局排除句判 —— 标段句和发包句通常就这么写（QTY-11 不回归）。

        名称行下面的编号行，只有层号正是块里等着的下一层时才「只管这一层」。层号对不上的（做法块后面
        从 1 重新编号的表注、「1.」块里的「1、」）先按全局排除句判：是排除句就照改前的口径管全局，
        不是排除句才照常按层读（PR #839 评审阻断项）。
        """
        nos = [no for no, _body, _style in split_layers_marked(s, self.expect_next()) if no is not None]
        if len(nos) >= 2 or rules.looks_like_practice(s, in_block=False):
            return True, bool(nos), True
        named = len(nos) == 1 and self.in_block and self.block_named
        return named, bool(nos), named and nos[0] == self.expect_next()

    def _name_line(self, name: str, rest: str, line: Line) -> None:
        rest = rest.strip()
        title = rest.lstrip("：: ")
        part = ""
        pm = _PART.search(title)
        if pm:
            part = pm.group("part").strip()
        else:
            pp = _PAREN_PART.search(title)
            if pp and not rules.practice_signal(pp.group("p")).thickness:
                part = pp.group("p").strip()
        title, layers_text = _name_tail(rest)
        full = (name + (" " + title if title else "")).strip()
        self._start_practice(full, category_of(name) or category_of(title) or self.heading_category,
                             part, f"p{line.row or line.paragraph or self.line_no}", named=True)
        if layers_text:
            # R4F-d fix15（决定 1）：名称行与层同段 / 同格时，名称、编号、批注（「（取消）」）与行首空框（□ ○）和层在
            # 同一句里——记成所在格里的 HOST_FREE 片；层从名称后面那一段里依次定位
            s = self._s
            off = len(s) - len(layers_text) if s.endswith(layers_text) else max(0, s.find(layers_text))
            if off > 0 and self._map is not None:
                self._map.free(self._lead, self._lead + off)
            self._layers(layers_text, line, True, search_from=off)

    def _is_sub(self, no: Optional[int], style: Optional[str]) -> bool:
        """带号、但标记样式与本块一级层不同：是当前层的子项。"""
        return no is not None and self.block_style is not None and style != self.block_style

    def _global_in_layer(self, body: str, line: Line, layer_ctx: bool) -> bool:
        """标段句、「不在本次范围」句写在层里也管全局（不变）：这一层单独成一个 exclusion 片段。"""
        gl = rules.find_exclusions(body, rules.GLOBAL_IN_LAYER) if layer_ctx else []
        if gl:
            self._emit(EXCLUSION, body, line, exclusions=gl)
            self.last = None
        return bool(gl)

    def _layers(self, s: str, line: Line, layer_ctx: bool = False, search_from: int = 0) -> None:
        """layer_ctx：这一行已判为成层（见 _layer_row）。不成层的行在 feed 里整行判过全局排除句了。
        search_from：s 是本行从第几个字起的一段（做法名称行名称后面的层），各片从那里起在整行里定位。"""
        pieces = split_layers_marked(s, self.expect_next())
        # 只记录所在格原文与各片角色（R4F-d fix14 → fix15：按切层之前的整行 / 整格原文定位），不改切片
        mp = self._map if self._map is not None else _line_mapper(line, line.text, 0, line.text)[0]
        placed = mp.place(self._s if self._map is not None else line.text, self._lead if self._map is not None else 0,
                          pieces, search_from)
        fallback = mp.spans[0][2] if mp.spans else None
        numbered = any(n is not None for n, _, _ in pieces)
        if not self.in_block and not numbered:
            # 块外散行（总说明里的「屋面防水：2.0厚…」）：严格判定，自成一个做法
            if not rules.looks_like_practice(s, in_block=False):
                self._emit(OTHER, s, line)
                return
            label = ""
            cm = _COLON_NAME.match(s)
            if cm and category_of(cm.group("name")):
                label = cm.group("name").strip()
            self._start_practice(label or self.heading or "", category_of(label) or self.heading_category,
                                 "", f"l{line.row or line.paragraph or self.line_no}")
            for k, (no, body, _style) in enumerate(pieces):
                if self._global_in_layer(body, line, layer_ctx):
                    if placed[k] is not None:
                        placed[k].set_role(HOST_LAYER)
                    continue
                kind = PRACTICE if rules.looks_like_practice(body, in_block=True) else OTHER
                self.last = self._emit(kind, body, line, layer_no=no)
                _attach(self.last, placed[k], HOST_LAYER if kind == PRACTICE or no is not None else HOST_FREE, fallback)
            self.end_block()
            return
        if numbered and not self.in_block:
            self._start_practice(self.heading or "", self.heading_category, "",
                                 f"n{line.row or line.paragraph or self.line_no}")
        elif (pieces and pieces[0][0] == 1 and (self.last_no or 0) >= 1
              and not self._is_sub(pieces[0][0], pieces[0][2])):
            # 行首层号回到 1、且标记样式与本块首层相同（「1.」对「1.」）：新的做法（与网格模式对齐）。
            # 只认行首 —— 同一行「面层：1）…」是子项；接在「2. 防水层：」后面的「①」「1、」「1）」也是子项
            self._restart_practice(line)
        for k, (no, body, style) in enumerate(pieces):
            sub = self._is_sub(no, style)
            if no is not None and not sub:
                self.block_style = self.block_style or style
                self.block_numbered = True
                self.last_no = no
            # 先于续行合并：并进上一层之后就再也认不出这是一句标段句
            if self._global_in_layer(body, line, layer_ctx):
                if placed[k] is not None:
                    placed[k].set_role(HOST_LAYER)
                continue
            if (no is None and self.last is not None and self.block_numbered
                    and not rules.thickness_starts(body)
                    and not rules.looks_like_practice(body, in_block=False)):
                # 无层号的续行并入上一层（来自另一单元格的续行不进上一层的作用域，见 _merge）
                old = self.last.text
                self.last.text = f"{old} {body}".strip()
                _merge(self.last, len(old), placed[k])
                self.head_cand = None
                continue
            # 带层内标记的编号层仍是做法层（「装修吊顶（详二次装修图纸）」词表里没有材料词）：
            # 要进覆盖核算，不能落成 other 之后哪一页都查不到（铁律 8）
            marked = layer_ctx and no is not None and rules.find_inline_marker(body) is not None
            kind = PRACTICE if marked or rules.looks_like_practice(body, in_block=True) else OTHER
            # 子项挂在父层号下：「2. 防水层：①3厚SBS…」的卷材是第 2 层的组成，不是第 1 层
            seg = self._emit(kind, body, line, layer_no=self.last_no if sub else no)
            _attach(seg, placed[k], HOST_LAYER if kind == PRACTICE or no is not None else HOST_FREE, fallback)
            if kind == OTHER and self.block_named:
                # 只认有名称的做法块（名称行、标题行开头）：说明文字里一串「1、…2、…」也会开一个块，
                # 那里的条款不是没认出来的层，列出来只是噪声（建施说明实测 45 行规范名称）
                seg.block_practice = self.practice
                seg.practice_key = self.practice_key
            if kind == PRACTICE:
                self.last = seg
            elif no is None and not numbered:
                self.end_block()
            elif no is not None and len(pieces) == 1 and category_of(body) and _section_title(body):
                self.head_cand = (seg, s.strip(), body)

    def _restart_practice(self, line: Line) -> None:
        cand = self.head_cand
        if cand is not None:
            # 「3.外墙面」层号恰好接得上、先按层收下了；紧跟着层号回到 1，才确认它是节标题
            seg, raw, title = cand
            seg.kind, seg.text = HEADING, raw
            seg.category = seg.part = seg.practice = seg.practice_key = ""
            seg.layer_no = None
            self.set_heading(title)
            name, category, part, named = self.heading or "", self.heading_category, "", False
        else:
            # 没有节标题：做法名、类别、部位沿用当前做法，只换 practice_key（与网格模式层号重起时
            # 沿用列组类别和行上的名称 / 部位一致）。清空会让后面的层失去上下文、被静默编成别类的码
            #（屋面卷材 → 楼(地)面卷材防水，status=ok 不进待核；PR #766 第二轮评审阻断项，铁律 8）
            name, category, part, named = self.practice, self.category, self.part, self.block_named
        self._start_practice(name, category, part, f"n{line.row or line.paragraph or self.line_no}",
                             named=named)


def mark_inline(segments: list[Segment]) -> list[Segment]:
    """层内标记统一在这里识别 —— 文本、表格、docx 三条路径共用这一处。

    必须等整份文件切完：标记常常落在续行上（一层分两行，第二行只有「（用户自理）」），续行并进上一层
    之前识别会漏掉。做法名称、适用部位上的精装修引用下传到这个做法的各层；做法名称只是沿用章节标题时
    不下传（「17.3 厨房设备见二次装修。」不是哪个做法的名称）。
    """
    for s in segments:
        if s.kind == PRACTICE:
            name = s.practice if s.practice != s.heading else ""
            s.inline = rules.find_inline_marker(s.text, name, s.part_full or s.part)
    return segments


def _text_lines(text: str) -> Iterator[Line]:
    for raw in (text or "").split("\n"):
        yield Line(text=raw, page_aware=True)


def segment_text(text: str, file_index: int, file_name: str) -> list[Segment]:
    out: list[Segment] = []
    seg = _TextSegmenter(file_index, file_name, out)
    for line in _text_lines(text):
        seg.feed(line)
    return mark_inline(out)


# ============================================================ 网格模式

@dataclass
class Grid:
    name: str                     # 表名（xlsx 工作表名 / docx「表k」）
    kind: str                     # xlsx / docx
    values: dict[tuple[int, int], str]
    merges: dict[tuple[int, int], tuple[int, int, int, int]] = field(default_factory=dict)
    rows: list[int] = field(default_factory=list)
    cols: list[int] = field(default_factory=list)
    docx_table_no: Optional[int] = None
    # docx 每行的格数（含空格）。docx 块级 API 不给横向合并，同一张表的标题行 4 格、构造层行 2 格，
    # 「适用部位」栏只能按「从行尾数第几格」对齐（WS29-k1）。xlsx 的合并格保留真实列号，不用它
    widths: dict[int, int] = field(default_factory=dict)
    # docx 每格的格内原文（与 values 逐字等长，格内硬换行是换行；R4F-d fix15，只给存疑闸的作用域用）
    raw: dict[tuple[int, int], str] = field(default_factory=dict)

    def cell(self, r: int, c: int) -> str:
        return self.values.get((r, c), "")

    def scope_cell(self, r: int, c: int, text: str) -> str:
        """与 text（该格文字或其去首尾空白后的样子）逐字等长的格内原文；对不上就用 text 本身。"""
        raw = self.raw.get((r, c))
        if raw is None:
            return text
        if len(raw) == len(text):
            return raw
        st = raw.strip()
        return st if len(st) == len(text) else text

    def col_span(self, r: int, c: int) -> tuple[int, int]:
        m = self.merges.get((r, c))
        return (m[1], m[3]) if m else (c, c)

    def coord(self, r: int, c: int) -> str:
        return f"{_col_letter(c)}{r}"

    def where(self, r: int, c: int, note: str = "") -> dict:
        if self.kind == "docx":
            loc = f"表{self.docx_table_no} 第{r}行第{c}列" + (f"（{note}）" if note else "")
            return {"location": loc, "row": r}
        loc = f"{self.name}!{self.coord(r, c)}（第{r}行" + (f"·{note}" if note else "") + "）"
        return {"location": loc, "sheet": self.name, "cell": self.coord(r, c), "row": r}


def _col_letter(n: int) -> str:
    s = ""
    while n > 0:
        n, rem = divmod(n - 1, 26)
        s = chr(65 + rem) + s
    return s


_ID_WORDS = {
    "code": ("编号", "代号", "编码"),
    "name": ("名称",),
    "part": ("部位", "房间", "区域", "楼层", "位置"),
    "category": ("类别", "分类", "类型"),
}
_LAYER_HEAD = ("层次", "层号", "序号", "层")


def _is_group_header(t: str) -> bool:
    return ("做法" in t or t in ("构造", "构造层次")) and len(t) <= 16 \
        and not any(w in t for w in ("名称", "编号", "代号", "编码"))


def _id_kind(t: str) -> Optional[str]:
    if len(t) > 12:
        return None
    for kind, words in _ID_WORDS.items():
        if any(w in t for w in words):
            return kind
    return None


@dataclass
class _Group:
    c1: int
    c2: int
    header: str
    category: str


@dataclass
class _Table:
    header_row: int
    groups: list[_Group]
    ids: dict[str, int]       # kind → col
    covered: set[int]         # 列组与标识列占用的列


def _detect_header(g: Grid, r: int) -> Optional[_Table]:
    texts = [(c, _norm_text(g.cell(r, c))) for c in g.cols if g.cell(r, c).strip()]
    if not texts:
        return None
    group_cells = [(c, t) for c, t in texts if _is_group_header(t)]
    if not group_cells:
        return None
    ids: dict[str, int] = {}
    for c, t in texts:
        if _is_group_header(t):
            continue
        k = _id_kind(t)
        if k and k not in ids:
            ids[k] = c
    merged = any(g.col_span(r, c)[1] > g.col_span(r, c)[0] for c, _ in group_cells)
    if not (len(group_cells) >= 2 or ids or merged):
        return None
    header_of = {c: t for c, t in texts}
    groups: list[_Group] = []
    for c, t in group_cells:
        c1, c2 = g.col_span(r, c)
        if c1 == c2:
            left = header_of.get(c - 1, "")
            if left and any(left == w or left.startswith(w) for w in _LAYER_HEAD) and not _is_group_header(left):
                c1 = c - 1
        cat = category_of(t.replace("做法", ""))
        groups.append(_Group(c1, c2, t, cat))
    groups.sort(key=lambda x: x.c1)
    covered = {c for grp in groups for c in range(grp.c1, grp.c2 + 1)} | set(ids.values())
    ids = {k: v for k, v in ids.items() if not any(grp.c1 <= v <= grp.c2 for grp in groups)}
    return _Table(r, groups, ids, covered)


def _is_subheader(g: Grid, r: int) -> bool:
    texts = [_norm_text(g.cell(r, c)) for c in g.cols if g.cell(r, c).strip()]
    if not texts:
        return False
    return all(len(t) <= 6 and not t.isdigit() and not rules.practice_signal(t).materials for t in texts)


@dataclass
class _PracticeState:
    key: str
    name: str
    part: str
    category: str
    numbered: bool = False
    last_no: Optional[int] = None
    last: Optional[Segment] = None


# ------------------------------------------------ 无表头做法表：标题行 + 构造层行（WS29-k1，issue #891）
# 每个做法两行或多行：标题行「■ 楼面9｜花岗岩楼面1（燃烧性能A级）｜50厚｜适用部位」，下面的行「1)… 2)…｜
# 适用部位的文字」。没有「做法」字样的表头，改前整张表逐行拼成一句走文本模式：标题行带「50厚」成了一个做法层
#（清单里 26 条「■ 楼面9 …」重复项），部位文字接在最后一层后面（部位 141/142 待补），不带厚度的标题行又按续行
# 并进上一个做法（内墙9 的条目指向内墙10 的标题行）。

_PART_LABELS = ("适用部位", "使用部位", "部位", "适用范围")
# 标题行首格是做法编号（「■ 楼面9」「■其它3」「■ 漆1」「■ 地防1」），编号后面可以直接跟名称。
# 不限于 _NAME_HEAD：油漆、配件、其它这几类的编号头不是部位词，标题行照样要认
_TITLE_CODE = re.compile(r"^[\s" + _LEAD_SYMBOLS
                         + r"]*(?P<code>[\u4e00-\u9fff]{1,6}\s*[-_]?\s*\d{1,3}[A-Za-z]?)(?![\d.．])\s*(?P<rest>.*)$")
_PART_MAX = 60


@dataclass
class _Titled:
    """由标题行开头的做法。标题行本身不出片段，它定做法名、类别，和「适用部位」栏在哪一格。"""
    name: str
    category: str
    label_col: int          # 标题行里「适用部位」那一格的列号：xlsx 按它取部位
    from_right: int         # 那一格离行尾几格：docx 按它取部位（见 Grid.widths）
    start: int              # 这个做法的第一个片段在 out 里的下标：部位跨行时回填前面的层
    key_hint: str
    parts: list[str] = field(default_factory=list)
    part_location: str = ""
    raw: str = ""           # 标题行原文（各格以空格连接）与位置：这个做法一层都没切出来时据此留痕（WS29-k2）
    where: dict = field(default_factory=dict)

    def full(self) -> str:
        return "、".join(self.parts)

    def part(self) -> str:
        return self.full()[:_PART_MAX]


def _part_label(t: str) -> bool:
    return re.sub(r"\s+", "", t or "").rstrip("：:") in _PART_LABELS


def _has_layer_no(t: str) -> bool:
    return any(no is not None for no, _b, _s in split_layers_marked(t))


class _GridSegmenter:
    def __init__(self, g: Grid, file_index: int, file_name: str, out: list[Segment],
                 text: _TextSegmenter):
        self.g = g
        self.fi = file_index
        self.fname = file_name
        self.out = out
        self.text = text
        self.table: Optional[_Table] = None
        self.states: dict[int, _PracticeState] = {}
        self.last_id: dict[str, str] = {}
        self.titled: Optional[_Titled] = None
        self._n = 0

    def run(self) -> None:
        g = self.g
        skip_sub = False
        for r in g.rows:
            t = _detect_header(g, r)
            if t is not None:
                self.table = t
                self.states = {}
                self.last_id = {}
                self._close_titled()
                self.text.end_block()
                skip_sub = True
                continue
            if self.table is None:
                if self._start_titled(r):
                    continue
                if self.titled is not None:
                    if self._titled_row(r):
                        continue
                    # 不是构造层行（整行一格的章节标题、配件标题）：做法到此结束，这一行照原来的逐行处理。
                    # 先结束文本块，否则这一行会按「无层号的续行」并进上一个做法的最后一层
                    self._close_titled()
                    self.text.end_block()
                self._text_row(r, set())
                continue
            if skip_sub and _is_subheader(g, r):
                continue
            skip_sub = False
            self._table_row(r)
        self._close_titled()

    def _title_row(self, r: int) -> Optional[_Titled]:
        """标题行：有一格恰好是「适用部位」（当列头用）、全行没有层号、首格是做法编号。缺一条都不是。

        首格这一条是防误认：「编号｜名称｜部位」这种普通列头行也有一格写着「部位」，但它首格不是编号。
        """
        g = self.g
        cells = [(c, _norm_text(g.cell(r, c)).replace("\n", " ")) for c in g.cols if g.cell(r, c).strip()]
        labels = [c for c, t in cells if _part_label(t)]
        if not labels or any(_has_layer_no(t) for _c, t in cells):
            return None
        head = [t for c, t in cells if c < labels[0]]
        m = _TITLE_CODE.match(head[0]) if head else None
        if m is None:
            return None
        code = re.sub(r"\s+", "", m.group("code"))
        title = m.group("rest").strip() or (head[1].strip() if len(head) > 1 else "")
        # 类别先看编号再看名称：「楼面5 面砖地面(含墙面）」按名称会先命中「墙面」
        category = category_of(code) or category_of(title) or self.text.heading_category
        from_right = g.widths.get(r, 0) - labels[0] if g.kind == "docx" else 0
        return _Titled(f"{code} {title}".strip(), category, labels[0], from_right, len(self.out), f"t{r}",
                       raw=" ".join(t for _c, t in cells), where=g.where(r, cells[0][0]))

    def _start_titled(self, r: int) -> bool:
        t = self._title_row(r)
        if t is None:
            return False
        self._close_titled()
        self.text.end_block()
        self.titled = t
        self.text._start_practice(t.name, t.category, "", t.key_hint, named=True)
        return True

    def _close_titled(self) -> None:
        """标题行开头的做法到此结束。它下面一行都没切出来（连着两个标题行、标题行在表尾、构造层行全空）时，
        改前这个做法整个消失、不出片段也没有提示；现在留一行 other 带着标题行原文，进「块内未识别行」（铁律 8）。"""
        t, self.titled = self.titled, None
        if t is None or any(s.kind == PRACTICE or s.block_practice is not None for s in self.out[t.start:]):
            return
        self.out.append(Segment(kind=OTHER, text=t.raw, file_index=self.fi, file_name=self.fname,
                                heading=self.text.heading, practice_key=f"f{self.fi}:{t.key_hint}:empty",
                                block_practice=t.name, **t.where))

    def _part_col(self, r: int) -> Optional[int]:
        t = self.titled
        assert t is not None
        if self.g.kind != "docx":
            return t.label_col
        c = self.g.widths.get(r, 0) - t.from_right
        return c if c >= 2 else None

    def _titled_row(self, r: int) -> bool:
        """标题行下面的行：「适用部位」栏的文字是这个做法的部位，不拼进层正文；其余格按层切，位置写真实列号。

        返回 False：这一行不是构造层行，做法到此结束。docx 按格数判（整行一格的是章节标题）；xlsx 看不出
        整行合并，只有一格、又没有层号的行按章节标题处理。
        """
        g, t = self.g, self.titled
        assert t is not None
        cells = [(c, g.cell(r, c)) for c in g.cols if g.cell(r, c).strip()]
        if not cells:
            return True
        if g.kind == "docx" and g.widths.get(r, 0) < 2:
            return False
        pc = self._part_col(r)
        layer_cells = [(c, x) for c, x in cells if c != pc]
        if g.kind != "docx" and len(cells) == 1 and layer_cells and not _has_layer_no(layer_cells[0][1]):
            return False
        part_cell = g.cell(r, pc) if pc is not None else ""
        bits = [b for b in (_norm_text(x) for x in part_cell.split("\n")) if b]
        if bits:
            if not t.parts:
                t.part_location = g.where(r, pc, "适用部位")["location"]
            t.parts.extend(bits)
        # 部位是这个做法的属性，不是句子：不送 feed，所以「…等不需二次装修的非配水房间」不参与任何排除判断
        self.text.part = t.part()
        for c, x in layer_cells:
            # R4F-d fix15：一格一个所在格（xlsx 格内多行不切断；docx 用格内原文，格内硬换行是句界）
            host = Host(g.scope_cell(r, c, x), "单元格")
            pos = 0
            for piece in x.split("\n"):
                at = pos + (len(piece) - len(piece.lstrip()))
                pos += len(piece) + 1
                if not self.text.in_block:
                    # 层正文里一句不成层的话（落成 other / 章节标题）会结束文本块；后面的层仍属这个做法
                    self.text._start_practice(t.name, t.category, t.part(), t.key_hint, named=True)
                self.text.feed(Line(text=_norm_text(piece), **g.where(r, c), host=host, host_offset=at))
        for s in self.out[t.start:]:
            if s.kind == PRACTICE:
                s.part, s.part_location, s.part_full = t.part(), t.part_location, t.full()
        return True

    def _text_row(self, r: int, covered: set[int]) -> None:
        g = self.g
        cells = [(c, g.cell(r, c)) for c in g.cols if c not in covered and g.cell(r, c).strip()]
        if not cells:
            return
        if not covered:
            # 没有表头的表：一行一行处理（CAD 导出的一行常被拆进相邻几格）。R4F-d fix15：切片照旧按拼成的一句，
            # 存疑闸的作用域按来源单元格分开（cells），格内原文（scope_text）里格内换行仍是句界
            texts, raws, spans, pos = [], [], [], 0
            for k, (c, t) in enumerate(cells):
                nt = _norm_text(t)
                txt = nt.replace("\n", " ")
                if k:
                    pos += 1
                spans.append((pos, pos + len(txt)))
                pos += len(txt)
                texts.append(txt)
                raws.append(g.scope_cell(r, c, nt))
            self.text.feed(Line(text=" ".join(texts), **g.where(r, cells[0][0]), scope_text=" ".join(raws),
                                cells=spans))
            return
        for c, t in cells:   # 表外的格（备注列等）
            host = Host(g.scope_cell(r, c, t), "单元格")
            pos = 0
            for piece in t.split("\n"):
                at = pos + (len(piece) - len(piece.lstrip()))
                pos += len(piece) + 1
                self.text.feed(Line(text=_norm_text(piece), **g.where(r, c), host=host, host_offset=at))

    def _id_values(self, r: int) -> tuple[dict[str, str], bool]:
        g, t = self.g, self.table
        assert t is not None
        vals: dict[str, str] = {}
        starts = False
        for kind, c in t.ids.items():
            v = _norm_text(g.cell(r, c)).replace("\n", " ")
            if v:
                starts = True
                self.last_id[kind] = v
            vals[kind] = self.last_id.get(kind, "")
        return vals, starts

    def _table_row(self, r: int) -> None:
        g, t = self.g, self.table
        assert t is not None
        vals, starts = self._id_values(r)
        if starts:
            self.states = {}
        name = " ".join(x for x in (vals.get("code", ""), vals.get("name", "")) if x)
        for gi, grp in enumerate(t.groups):
            cells = [(c, g.cell(r, c)) for c in range(grp.c1, grp.c2 + 1) if c in g.cols and g.cell(r, c).strip()]
            nums = [int(x) for _, x in cells if _LAYER_NUM_CELL.match(x)]
            texts = [(c, x) for c, x in cells if not _LAYER_NUM_CELL.match(x)]
            if not texts:
                continue
            layer_no = nums[0] if nums else None
            st = self.states.get(gi)
            if st is not None and layer_no == 1 and st.last_no is not None and st.last_no >= 1:
                st = None     # 同一部位下层号重新从 1 起：新的做法
            if st is None:
                self._n += 1
                cat = grp.category or category_of(vals.get("category", "")) or category_of(name)
                label = name or vals.get("part", "") or grp.header
                if name and grp.category and category_of(name) not in ("", grp.category):
                    # 按房间排的做法表：编号列是楼地面的做法编号，别拿它给踢脚 / 顶棚列组命名
                    label = "·".join(x for x in (vals.get("part", ""), grp.header) if x)
                st = _PracticeState(key=f"f{self.fi}:{g.name}:g{gi}:r{r}:{self._n}", name=label,
                                    part=vals.get("part", ""), category=cat)
                self.states[gi] = st
            for c, x in texts:
                # 一格一个所在格（R4F-d fix15：格内多行共用，docx 用格内原文），只记录，不改切片
                host = Host(g.scope_cell(r, c, x), "单元格")
                mp = _Mapper([(0, len(x), host, 0)])
                pos = 0
                for raw in x.split("\n"):
                    at = pos + (len(raw) - len(raw.lstrip()))
                    pos += len(raw) + 1
                    nxt = (st.last_no + 1) if st.last_no is not None else (1 if layer_no is None else None)
                    s = _norm_text(raw)
                    marked = split_layers_marked(s, nxt)        # 与 split_layers 同一结果，另留层号样式
                    placed = mp.place(s, at, marked)
                    for k, (no, body, _style) in enumerate(marked):
                        self._layer(st, grp, r, c, no if no is not None else layer_no, body,
                                    explicit=(no is not None or layer_no is not None), placement=placed[k],
                                    fallback=host)
        self._text_row(r, t.covered)

    def _layer(self, st: _PracticeState, grp: _Group, r: int, c: int, no: Optional[int], body: str,
               *, explicit: bool, placement: Optional[_Placement] = None, fallback: Optional[Host] = None) -> None:
        if not body:
            return
        if explicit:
            st.numbered = True
        if (not explicit and st.numbered and st.last is not None
                and not rules.thickness_starts(body) and not rules.looks_like_practice(body, in_block=False)):
            old = st.last.text
            st.last.text = f"{old} {body}".strip()      # 无层号的续行并入上一层（另一单元格的续行不进它的作用域）
            _merge(st.last, len(old), placement)
            return
        # 做法列里的一格就是一层。标段句、「不在本次范围」句写在层里也管全局（不变）；自理、二次装修
        # 只管这一层（WS29-f 第一条），由 mark_inline 标在片段上，不再把整层变成排除句
        ex = rules.find_exclusions(body, rules.GLOBAL_IN_LAYER)
        note = f"{grp.header}第{no}层" if no is not None else grp.header
        if ex:
            seg = Segment(kind=EXCLUSION, text=body, file_index=self.fi, file_name=self.fname,
                          exclusions=ex, heading=self.text.heading, **self.g.where(r, c, note))
            self.out.append(seg)
            _attach(seg, placement, HOST_LAYER, fallback)
            return
        marked = rules.find_inline_marker(body) is not None
        kind = PRACTICE if marked or rules.looks_like_practice(body, in_block=True) else OTHER
        seg = Segment(kind=kind, text=body, file_index=self.fi, file_name=self.fname,
                      heading=self.text.heading, section_domain=self.text.heading_domain,
                      category=st.category, group=grp.header, part=st.part, practice=st.name,
                      practice_key=st.key, layer_no=no, **self.g.where(r, c, note),
                      block_practice=st.name if kind == OTHER else None)
        self.out.append(seg)
        _attach(seg, placement, HOST_LAYER if kind == PRACTICE or no is not None else HOST_FREE, fallback)
        if kind == PRACTICE:
            st.last = seg
            st.last_no = no if no is not None else st.last_no


def segment_grids(grids: Iterable[Grid], file_index: int, file_name: str,
                  text_seg: Optional[_TextSegmenter] = None,
                  out: Optional[list[Segment]] = None) -> list[Segment]:
    out = out if out is not None else []
    ts = text_seg or _TextSegmenter(file_index, file_name, out)
    for g in grids:
        _GridSegmenter(g, file_index, file_name, out, ts).run()
        ts.end_block()
    return mark_inline(out)


def _cell_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).replace("\r\n", "\n").replace("\r", "\n")
    return s.strip()


def xlsx_grids(content: bytes) -> tuple[list[Grid], list[str], list[dict]]:
    """xlsx → (可见工作表的 Grid, 隐藏工作表名, 容错读取件的告警)。读不出来抛 XlsxUnreadable。"""
    from openpyxl.cell.cell import MergedCell
    from openpyxl.utils import column_index_from_string

    # zt 改动：容错读取件来自 zt.costsheet（同一份旧 xlsx_tolerant.py）。
    from zt.costsheet.xlsx_tolerant import load_workbook_tolerant

    wb, notes = load_workbook_tolerant(content, read_only=False, data_only=True)
    grids: list[Grid] = []
    hidden: list[str] = []
    try:
        for ws in wb.worksheets:
            if (getattr(ws, "sheet_state", "visible") or "visible") != "visible":
                hidden.append(ws.title)
                continue
            hidden_rows = {i for i, d in ws.row_dimensions.items() if getattr(d, "hidden", False)}
            hidden_cols: set[int] = set()
            for key, d in ws.column_dimensions.items():
                if getattr(d, "hidden", False):
                    lo = d.min or column_index_from_string(key)
                    hi = d.max or lo
                    hidden_cols.update(range(lo, hi + 1))
            merges: dict[tuple[int, int], tuple[int, int, int, int]] = {}
            for mr in ws.merged_cells.ranges:
                box = (mr.min_row, mr.min_col, mr.max_row, mr.max_col)
                for rr in range(mr.min_row, mr.max_row + 1):
                    for cc in range(mr.min_col, mr.max_col + 1):
                        merges[(rr, cc)] = box
            values: dict[tuple[int, int], str] = {}
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell, MergedCell) or cell.value is None:
                        continue
                    t = _cell_text(cell.value)
                    if t:
                        values[(cell.row, cell.column)] = t
            rows = sorted({r for r, _ in values} - hidden_rows)
            cols = sorted({c for _, c in values} - hidden_cols)
            grids.append(Grid(ws.title, "xlsx", values, merges, rows, cols))
    finally:
        try:
            wb.close()
        except Exception:  # noqa: BLE001
            pass
    return grids, hidden, notes


# ============================================================ docx

def segment_docx(content: bytes, file_index: int, file_name: str) -> list[Segment]:
    """WS04 块级 API：段落按文本模式、表格按网格。解析失败抛 DocxParseError（调用方退回逐行）。"""
    from zt.textextract.docx_reader import docx_blocks

    blocks = docx_blocks(content)
    out: list[Segment] = []
    ts = _TextSegmenter(file_index, file_name, out)
    table_rows: list = []
    table_no = 0

    def flush_table() -> None:
        nonlocal table_rows
        if not table_rows:
            return
        values: dict[tuple[int, int], str] = {}
        raws: dict[tuple[int, int], str] = {}
        widths: dict[int, int] = {}
        for rr, b in enumerate(table_rows, start=1):
            widths[rr] = len(b.cells)
            braw = list(getattr(b, "cells_raw", None) or [])
            for cc, txt in enumerate(b.cells, start=1):
                t = (txt or "").strip()
                if t:
                    values[(rr, cc)] = t
                    rw = (braw[cc - 1] if cc - 1 < len(braw) else "").strip()
                    if len(rw) == len(t) and rw.replace("\n", " ") == t:
                        raws[(rr, cc)] = rw
        rows = sorted({r for r, _ in values})
        cols = sorted({c for _, c in values})
        grid = Grid(f"表{table_no}", "docx", values, {}, rows, cols, docx_table_no=table_no, widths=widths,
                    raw=raws)
        _GridSegmenter(grid, file_index, file_name, out, ts).run()
        ts.end_block()
        table_rows = []

    last_table: Optional[int] = None
    for b in blocks:
        if b.kind == "row":
            if b.table_index != last_table:
                flush_table()
                table_no += 1
                last_table = b.table_index
            table_rows.append(b)
            continue
        flush_table()
        last_table = None
        if b.is_toc:
            continue
        ts.feed(Line(text=(b.num_prefix or "") + b.text, location=f"段落#{b.no}", paragraph=b.no,
                     force_heading=bool(b.heading_level)))
    flush_table()
    return mark_inline(out)


# ============================================================ 分派

def segment_file(content: Optional[bytes], file_name: str, ext: str, text: str,
                 file_index: int) -> FileSegments:
    """按格式切片。结构化读取失败时退回文本模式并记一条告警（铁律 8：退化必须可见）。"""
    ext = (ext or "").lower().lstrip(".")
    warnings: list[tuple[str, str]] = []
    if content and ext in ("xlsx", "xlsm", "xls"):
        raw = content
        if ext == "xls":
            from zt.textextract.office import libreoffice_convert_ex

            raw, why = libreoffice_convert_ex(content, "xls", "xlsx")
            if raw is None:
                log.warning("xls→xlsx 转换失败，逐行处理：%s（%s）", file_name, why)
        if raw is not None:
            try:
                grids, hidden, notes = xlsx_grids(raw)
            except Exception as e:  # noqa: BLE001 —— 坏文件抛什么的都有，退文本模式并说明
                log.warning("xlsx 结构化读取失败，逐行处理：%s（%s）", file_name, e)
            else:
                segs = segment_grids(grids, file_index, file_name)
                for n in notes:
                    warnings.append(("quantity_xlsx_repaired", f"「{file_name}」{n.get('message')}"))
                if hidden:
                    warnings.append(("quantity_hidden_sheets",
                                     f"「{file_name}」隐藏工作表已跳过：" + "、".join(hidden)))
                return FileSegments(file_index, file_name, "grid", segs, warnings, hidden)
    if content and ext in ("docx", "docm", "doc", "rtf"):
        raw = content
        if ext in ("doc", "rtf"):
            from zt.textextract.office import libreoffice_convert_ex

            raw, _why = libreoffice_convert_ex(content, ext, "docx")
        if raw is not None:
            try:
                return FileSegments(file_index, file_name, "docx", segment_docx(raw, file_index, file_name))
            except Exception as e:  # noqa: BLE001
                log.warning("docx 块级读取失败，逐行处理：%s（%s）", file_name, e)
        warnings.append(("quantity_docx_line_mode",
                         f"「{file_name}」表格结构读取失败，已按文本逐行处理（位置按抽取文本行）"))
    return FileSegments(file_index, file_name, "text", segment_text(text, file_index, file_name), warnings)


def assign_ids(files: list[FileSegments]) -> list[Segment]:
    """全局编号（s1、s2 …），按文件顺序、文件内原文顺序 —— 同一输入永远同一编号（QTY-04）。"""
    all_segs: list[Segment] = []
    n = 0
    for f in files:
        for i, s in enumerate(f.segments):
            n += 1
            s.id = f"s{n}"
            s.order = n
            all_segs.append(s)
    return all_segs
