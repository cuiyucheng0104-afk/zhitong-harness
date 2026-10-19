# 来源：engine-service/app/engine/quantity_list.py 的 ⑩ render 部分（zhitong-cost-ai@991ff5cd，W4 搬入 zt；改动见本文件内「zt 改动」注释）
"""清单 xlsx：清单 / 待核与未匹配 / 覆盖核算 / 说明 四个工作表（QTY-08/10）。

输入是 :func:`zt.quantity.pipeline.finalize` 的结果（``render`` 子命令读 finalize 的 result.json）。
工程量列刻意留空：无施工图纸无法精确计算工程量，这是既定设计不是漏填。
"""
from __future__ import annotations

from io import BytesIO
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from zt.quantity import rules
from zt.quantity.pipeline import DEST_AUX, DEST_BASE, DEST_CN, DEST_ITEM, SCOPE_NOTE
from zt.quantity.stamp import stamp_ai_draft_xlsx

SHEET_LIST = "分部分项工程量清单"
SHEET_PENDING = "待核与未匹配"
SHEET_COVERAGE = "覆盖核算"
SHEET_NOTES = "说明"
TITLE = "分部分项工程量清单（框架 · AI 初稿 · 未含工程量）"
HEADER_ROW = 3

_COL_WIDTHS = {"A": 8, "B": 15.67, "C": 26.17, "D": 60, "E": 8, "F": 12.33, "M": 60}
_HEADERS = {1: "序号", 2: "项目编码", 3: "项目名称", 4: "项目特征描述", 5: "计量单位", 6: "工程量", 13: "来源依据"}


def safe_cell(v: Any) -> Any:
    """防 xlsx 公式注入：来源与原文是外部文档逐字原文（不可信），= + - @ 制表符/回车 开头加单引号。

    :param v: 单元格值。
    :returns: 处理后的值。
    """
    if isinstance(v, str) and v.startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + v
    return v


def code_basis(meta: dict) -> str:
    """表头「编码依据」一行。

    :param meta: finalize 结果的 meta。
    :returns: 文字。
    """
    fp = (meta.get("registry_fingerprint") or "")[:12]
    return (f"编码依据：GB50854 受控码表 {meta.get('code_version')}（共 {meta.get('registry_count', '?')} 条，"
            f"指纹 sha256:{fp}）。注意：本码表与 2013 版清单编码存在同号异义，不可与 2013 口径清单混用。")


def render_xlsx(result: dict) -> bytes:
    """finalize 结果 → xlsx 字节。

    zt 改动：旧版签名是 ``render_xlsx(items, ocr_files, *, unmatched, ledger, practice, coverage, meta, warnings)``，
    台账与片段是内存对象；现在全部取自 finalize 结果（``ledger`` 已按行序列化）。

    :param result: :func:`zt.quantity.pipeline.finalize` 的结果。
    :returns: xlsx 字节。
    """
    items = result.get("items") or []
    meta = dict(result.get("meta") or {})
    coverage = result.get("coverage_detail") or {}
    warnings = [w.get("message", "") for w in result.get("warning_items") or []] or list(result.get("warnings") or [])
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_LIST
    thin = Border(left=Side("thin"), right=Side("thin"), top=Side("thin"), bottom=Side("thin"))
    bold = Font(name="宋体", size=10, bold=True)
    normal = Font(name="宋体", size=10)
    small = Font(name="宋体", size=8, color="666666")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    for col, w in _COL_WIDTHS.items():
        ws.column_dimensions[col].width = w

    ws.cell(row=1, column=1, value=TITLE).font = Font(name="宋体", size=14, bold=True)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=13)
    ws.cell(row=2, column=1, value=code_basis(meta)).font = Font(name="宋体", size=9, color="CC0000")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=13)
    ws.cell(row=2, column=1).alignment = left
    for col, val in _HEADERS.items():
        c = ws.cell(row=HEADER_ROW, column=col, value=val)
        c.font, c.alignment, c.border = bold, center, thin

    row, seq, current_section = HEADER_ROW + 1, 1, None
    for item in items:
        if item["section"] != current_section:
            current_section = item["section"]
            ws.cell(row=row, column=3, value=current_section).font = bold
            for col in range(1, 14):
                ws.cell(row=row, column=col).border = thin
            row += 1
        values = {1: seq, 2: item["code"], 3: item["name"], 4: item["features"], 5: item["unit"],
                  13: safe_cell(item.get("source", ""))}
        for col, val in values.items():
            c = ws.cell(row=row, column=col, value=val)
            c.font = small if col == 13 else normal
            c.alignment = center if col in (1, 2, 5) else left
        # F 列（工程量）刻意不写 —— 无施工图纸无法精确计算，这是既定设计不是漏填
        for col in range(1, 14):
            ws.cell(row=row, column=col).border = thin
        row += 1
        seq += 1

    tail = row + 1
    red = Font(name="宋体", size=9, bold=True, color="CC0000")
    if not items:
        c = ws.cell(row=tail, column=3, value="注：本次没有形成任何正式清单项——全部做法片段已降为待核 / 未匹配 / 范围排除"
                                              "（或并入主项、基层），明细见「待核与未匹配」页与「覆盖核算」页。")
        c.font, c.alignment = red, left
        tail += 1
    if meta.get("status") == "partial":
        c = ws.cell(row=tail, column=3, value="注：本清单为部分完成结果——部分做法片段没有分类回答，"
                                              "明细见「覆盖核算」页（去向「模型未完成」）。")
        c.font = red
        tail += 1
    ocr_files = meta.get("ocr_files") or []
    if ocr_files:
        note = ("注：以下输入文件为扫描件，正文来自旁车识别，本表相关条目与原文摘录"
                "请人工核对原件：" + "、".join(f"「{n}」" for n in ocr_files))
        c = ws.cell(row=tail, column=3, value=note)
        c.font, c.alignment = red, left

    _render_pending(wb, result.get("unmatched") or [], bold, normal, left, thin)
    _render_coverage(wb, result.get("ledger") or [], coverage, bold, normal, left, thin)
    _render_notes(wb, meta, coverage, warnings, bold, normal, left)
    stamp_ai_draft_xlsx(wb)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _render_pending(wb, rows, bold, normal, left, thin) -> None:
    ws = wb.create_sheet(SHEET_PENDING)
    heads = ["序号", "去向", "原因码", "原因", "原文", "文件", "位置", "模型给的编码", "建议 / 规则期望"]
    widths = [6, 12, 22, 50, 50, 24, 36, 14, 50]
    for i, (h, w) in enumerate(zip(heads, widths), start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font, c.border = bold, thin
        ws.column_dimensions[chr(64 + i)].width = w
    for n, r in enumerate(rows, start=1):
        vals = [n, DEST_CN.get(r["kind"], r["kind"]), r["reason_code"], r["reason"], r["text"], r["file_name"],
                r["location"], r.get("llm_guessed_code") or "", r.get("suggested") or ""]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=n + 1, column=i, value=safe_cell(v))
            c.font, c.alignment, c.border = normal, left, thin
    if not rows:
        ws.cell(row=2, column=1, value="（无：全部做法片段都已进入正式清单或并入主项 / 基层）").font = normal


def _render_coverage(wb, ledger: list[dict], coverage: dict, bold, normal, left, thin) -> None:
    ws = wb.create_sheet(SHEET_COVERAGE)
    by = coverage.get("by_destination") or {}
    ws.cell(row=1, column=1, value=f"做法片段共 {coverage.get('practice_segments', len(ledger))} 个；去向合计 "
                                   f"{sum(by.values())} 个（两者必须相等）").font = bold
    ws.cell(row=2, column=1, value="；".join(f"{DEST_CN[k]} {v}" for k, v in by.items())).font = normal
    heads = ["片段", "文件", "位置", "类别", "部位", "做法", "层", "原文", "去向", "清单编码 / 模型编码", "说明"]
    widths = [7, 22, 36, 10, 16, 24, 5, 50, 14, 20, 50]
    for i, (h, w) in enumerate(zip(heads, widths), start=1):
        c = ws.cell(row=4, column=i, value=h)
        c.font, c.border = bold, thin
        ws.column_dimensions[chr(64 + i)].width = w
    fill = PatternFill("solid", fgColor="FFF2CC")
    for n, r in enumerate(ledger, start=5):
        dest = r.get("dest")
        vals = [r["id"], r["file_name"], r["location"], r.get("category"), r.get("part"), r.get("practice"),
                r.get("layer_no"), r["text"], DEST_CN.get(dest, dest) if dest else "（缺去向）",
                r.get("code") or "", r.get("reason") or ""]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=n, column=i, value=safe_cell(v))
            c.font, c.alignment, c.border = normal, left, thin
            if dest is not None and dest not in (DEST_ITEM, DEST_AUX, DEST_BASE):
                c.fill = fill


def _render_notes(wb, meta, coverage, warnings, bold, normal, left) -> None:
    ws = wb.create_sheet(SHEET_NOTES)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 100
    rows: list[tuple[str, Optional[str]]] = [
        ("产物", TITLE),
        ("编码依据", code_basis(meta)),
        ("码表版本", meta.get("code_version")),
        ("码表指纹", meta.get("registry_fingerprint")),
        ("提示词版本", meta.get("prompt_version")),
        # zt 改动：分类由 agent 作答，不再记模型名、调用次数与缓存命中
        ("片段指纹", meta.get("prepare_fingerprint")),
        ("分类回答", f"{meta.get('answered', 0)} / {meta.get('to_classify', 0)} 个待分类片段有回答"),
        ("结果状态", "部分完成（partial）" if meta.get("status") == "partial" else "完成（ok）"),
        ("项目名称", meta.get("project_name") or "（未提供）"),
        ("识别出的标段", f"{meta['project_lot']} 标" if meta.get("project_lot") is not None else "未识别"),
        ("适用范围", SCOPE_NOTE),
        ("工程量列", "留空：无施工图纸无法精确计算工程量，这是既定设计不是漏填"),
        ("待核的含义", "模型给的编码与原文语义不一致、角色冲突或排除未确认的条目一律降为待核，不进正式清单、不自动改码"),
        ("复现性", "切片、编号、闸、聚合与流水号全部由脚本确定性完成；同一输入的片段编号与片段指纹不变。"
                   "分类回答保存在工作目录的 answers 文件里，用同一份 answers 重跑 finalize 结果逐字相同"),
    ]
    rows.append(("识别出的做法", f"识别出做法 {coverage.get('practices', 0)} 个，"
                                  f"做法层（片段）{coverage.get('practice_segments', 0)} 个"))
    rows.append(("做法暂定的层",
                 f"{coverage.get('provisional_segments', 0)} 个做法层{rules.PROVISIONAL_NOTE}：照常列项，"
                 "清单页「来源依据」列与「覆盖核算」页带标记，是否属于本次范围请人工确认"))
    exs = coverage.get("exclusions") or []
    glob = [e for e in exs if e.get("scope") == "global"]
    inline_n = sum(1 for e in exs if e.get("scope") == "inline")
    rows.append(("排除句清单",
                 f"全局排除句 {len(glob)} 条（作用到所有所选文件，逐条列在下面，排掉 0 个片段的也列）；"
                 f"层内标记（自理 / 本层为二次装修）排掉 {inline_n} 层，逐层见「待核与未匹配」页"))
    for e in glob:
        rows.append((f"排除句 {e['id']}",
                     f"{e.get('file_name', '')} {e.get('location', '')} | 原句：{e['sentence']} | "
                     f"对象词：{'、'.join(e.get('objects') or [])} | 排掉 {e.get('hits', 0)} 个片段"))
    rel = coverage.get("scope_related") or []
    rows.append(("范围相关原文",
                 f"{len(rel)} 句：原文里读到了「二次装修 / 精装修 / 自理」字样、或句式像排除句，"
                 "但没有据此排除或标记任何做法，逐句列在下面"))
    for i, x in enumerate(rel, start=1):
        rows.append((f"范围相关原文 {i}",
                     f"{x.get('file_name', '')} {x.get('location', '')} | 原句：{x['sentence']} | {x['note']}"))
    blk = coverage.get("block_lines_unrecognized") or []
    empty = sorted({x["practice"] for x in blk if not x.get("practice_has_layers")})
    rows.append(("块内未识别行",
                 f"{len(blk)} 行：做法块里没有认成构造层的行（不是做法片段，不进覆盖核算的去向），逐行列在下面，"
                 "请人工确认有没有漏项" + (f"；其中 {len(empty)} 个做法一层都没有认出：{'、'.join(empty)}" if empty else "")))
    for i, x in enumerate(blk, start=1):
        rows.append((f"块内未识别行 {i}",
                     f"{x.get('file_name', '')} {x.get('location', '')} | 做法：{x.get('practice') or '（无名称）'}"
                     + ("" if x.get("practice_has_layers") else "（这个做法一层都没有认出）")
                     + f" | 原文：{x['text']}"))
    for f in coverage.get("by_file") or []:
        if f.get("duplicate_of"):
            rows.append((f"输入文件：{f['file_name']}", f["notes"][0]))
            continue
        rows.append((f"输入文件：{f['file_name']}",
                     f"来源 {f['text_source']} · 形态 {f['mode']} · 专业 {rules.domain_name(f['scope'])} · "
                     f"做法片段 {f['practice_segments']} 个"))
    for i, w in enumerate(warnings, start=1):
        rows.append((f"告警 {i}", w))
    for n, (k, v) in enumerate(rows, start=1):
        ws.cell(row=n, column=1, value=k).font = bold
        c = ws.cell(row=n, column=2, value=safe_cell("" if v is None else str(v)))
        c.font, c.alignment = normal, left
