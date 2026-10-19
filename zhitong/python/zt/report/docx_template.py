"""正式报告 docx 模板：扫描占位符、按项目信息替换、生成登记清单。

模板在 ``zhitong/skills/zt-report/data/docx-templates/``（旧仓 ``docs/报告模板扩充包/`` 16 份与
``docs/报告模板起步包/`` 3 份，原样迁入）。占位符写法 ``{名称}``，名称来自旧系统项目「汇总信息」的种子名。

替换口径与旧服务端 ``DocxPlaceholderRenderer`` 相同：

- 正文、表格（含嵌套）、页眉页脚里的 ``{名称}`` 换成给定的值；
- 没给值的占位符换成 ``###名称###``，并列入 ``missing``（「数据没填」，不是模板坏了）；
- ``{#…}`` ``{/…}`` ``{^…}`` ``{@…}`` 循环 / 条件标签不支持：原样保留，列入 ``unsupported``；
- 文本框、内容控件、脚注里的占位符不替换（python-docx 读不到这些部件），扫描时也不计。

一个占位符被 Word 拆在几个 run 里时，替换值写进第一个 run，其余 run 里属于它的字删掉；其余字与格式不动。
本模块不写任何金额：值全部来自调用方（项目说明里人工确认过的信息），结果带 ``needs_manual``。
"""
from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable, Iterator, Optional

import yaml
from docx import Document
from docx.document import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph

from zt.common import ZtError

PLACEHOLDER_RE = re.compile(r"\{([^{}\n]{1,40})\}")
_UNSUPPORTED_LEAD = ("#", "/", "^", "@")
MISSING_FORMAT = "###{name}###"
TODO_MARK = "【待补"

MANIFEST_FILE = "manifest.json"
MAPPING_FILE = "mapping.yaml"
SOURCE_MANIFEST_FILE = "source-manifest-扩充包.json"
PACK_EXPANSION = "扩充包"
PACK_STARTER = "起步包"
STARTER_SUFFIX = "（起步模板）"
# 从旧扩充包登记清单带过来的字段
_CARRIED = ("no", "folder_name", "folder_id_suffix", "source", "headings", "amount_cells")


def _iter_tables(tables: Iterable[Table]) -> Iterator[Paragraph]:
    for t in tables:
        for row in t.rows:
            seen: set[int] = set()
            for cell in row.cells:
                # 合并单元格在 python-docx 里每个格都返回同一个 _tc：只走一次
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))
                yield from cell.paragraphs
                yield from _iter_tables(cell.tables)


def iter_paragraphs(doc: DocxDocument) -> Iterator[Paragraph]:
    """正文、表格（含嵌套）、各节页眉页脚（含首页与偶数页）的全部段落。

    :param doc: python-docx 文档。
    :returns: 段落迭代器。
    """
    yield from doc.paragraphs
    yield from _iter_tables(doc.tables)
    seen: set[int] = set()
    for sec in doc.sections:
        for part in (sec.header, sec.footer, sec.first_page_header, sec.first_page_footer,
                     sec.even_page_header, sec.even_page_footer):
            # 沿用上一节（第一节则是没有定义）的不读：读 paragraphs 会替它新建一个空部件
            if part.is_linked_to_previous:
                continue
            key = id(part._element)
            if key in seen:
                continue
            seen.add(key)
            yield from part.paragraphs
            yield from _iter_tables(part.tables)


def _supported(name: str) -> bool:
    return not name.startswith(_UNSUPPORTED_LEAD)


@dataclass
class ScanResult:
    """一份模板的扫描结果。

    :ivar placeholders: 去重后的占位符名，按首次出现顺序。
    :ivar unsupported: 循环 / 条件标签（原样保留）。
    :ivar todo_count: 「【待补」处数。
    """

    placeholders: list[str] = field(default_factory=list)
    unsupported: list[str] = field(default_factory=list)
    todo_count: int = 0


def _open(path: Path) -> DocxDocument:
    try:
        return Document(str(path))
    except Exception as e:  # noqa: BLE001 - python-docx 对坏文件抛的异常类型不一
        raise ZtError("docx_unreadable", f"读不了 docx 模板：{path.name}（{type(e).__name__}）") from None


def scan(path: Path) -> ScanResult:
    """扫描一份 docx 模板。

    :param path: docx 路径。
    :returns: 扫描结果。
    :raises ZtError: 文件读不了（``docx_unreadable``）。
    """
    doc = _open(path)
    out = ScanResult()
    for p in iter_paragraphs(doc):
        text = "".join(r.text for r in p.runs)
        out.todo_count += text.count(TODO_MARK)
        for m in PLACEHOLDER_RE.finditer(text):
            name = m.group(1)
            bucket = out.placeholders if _supported(name) else out.unsupported
            if name not in bucket:
                bucket.append(name)
    return out


@dataclass
class RenderResult:
    """替换结果。

    :ivar replaced: 换上了值的占位符名。
    :ivar missing: 没给值、换成 ``###名称###`` 的占位符名。
    :ivar unsupported: 原样保留的循环 / 条件标签。
    :ivar unused_values: 给了值、模板里却没有的名称。
    :ivar occurrences: 替换的处数（含 missing）。
    """

    replaced: list[str] = field(default_factory=list)
    missing: list[str] = field(default_factory=list)
    unsupported: list[str] = field(default_factory=list)
    unused_values: list[str] = field(default_factory=list)
    occurrences: int = 0


def _replace_in_paragraph(p: Paragraph, values: dict[str, str], res: RenderResult) -> None:
    runs = list(p.runs)
    if not runs:
        return
    texts = [r.text for r in runs]
    joined = "".join(texts)
    matches = [m for m in PLACEHOLDER_RE.finditer(joined)]
    if not matches:
        return
    starts: list[int] = []
    pos = 0
    for t in texts:
        starts.append(pos)
        pos += len(t)
    # 按文档顺序登记，再从右往左换（前面的位置不受影响）
    for m in matches:
        name = m.group(1)
        bucket = res.unsupported if not _supported(name) else (res.replaced if name in values else res.missing)
        if name not in bucket:
            bucket.append(name)
    for m in reversed(matches):
        name = m.group(1)
        if not _supported(name):
            continue
        new = values[name] if name in values else MISSING_FORMAT.format(name=name)
        res.occurrences += 1
        a, b = m.span()
        first = True
        for i, s in enumerate(starts):
            e = s + len(texts[i])
            if e <= a or s >= b:
                continue
            lo, hi = max(a, s) - s, min(b, e) - s
            t = texts[i]
            texts[i] = t[:lo] + (new if first else "") + t[hi:]
            first = False
        # 下一轮按新文本重算起点
        starts = []
        pos = 0
        for t in texts:
            starts.append(pos)
            pos += len(t)
    for r, t in zip(runs, texts):
        if r.text != t:
            r.text = t


def render(src: Path, values: dict[str, Any], dst: Path) -> RenderResult:
    """按项目信息替换一份模板的占位符，写到 ``dst``。

    :param src: 模板路径。
    :param values: 名称 → 值（值转成字符串；None 视为没给）。
    :param dst: 输出路径（父目录不存在时创建）。
    :returns: 替换结果。
    :raises ZtError: 模板读不了（``docx_unreadable``）。
    """
    vals = {str(k): str(v) for k, v in (values or {}).items() if v is not None}
    doc = _open(src)
    res = RenderResult()
    for p in iter_paragraphs(doc):
        _replace_in_paragraph(p, vals, res)
    res.unused_values = [k for k in vals if k not in res.replaced]
    dst.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dst))
    return res


# ================================================================ 登记清单


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def build_manifest(directory: Path) -> dict:
    """从目录里的 docx 重算登记清单（不手工登记）。

    每份：文件名、包（扩充包 / 起步包）、sha256、字节数、占位符集合与个数、「【待补」处数、对应的 AI 初稿模板码
    （``mapping.yaml``）；扩充包的份另带旧登记清单里的分类、来源、章节标题、金额格数。

    :param directory: ``data/docx-templates``。
    :returns: 登记清单。
    :raises ZtError: 目录里没有 docx（``no_docx_templates``）或 mapping.yaml 非法（``data_file_invalid``）。
    """
    files = sorted(directory.glob("*.docx"), key=lambda p: p.name)
    if not files:
        raise ZtError("no_docx_templates", f"目录里没有 docx 模板：{directory}")
    source = {}
    src_path = directory / SOURCE_MANIFEST_FILE
    if src_path.is_file():
        source = {t["upload_name"]: t for t in json.loads(src_path.read_text(encoding="utf-8")).get("templates") or []}
    mapping: dict[str, list[str]] = {}
    map_path = directory / MAPPING_FILE
    if map_path.is_file():
        try:
            raw = yaml.safe_load(map_path.read_text(encoding="utf-8")) or {}
        except yaml.YAMLError as e:
            raise ZtError("data_file_invalid", f"{MAPPING_FILE} 不是合法 YAML：{e}") from None
        mapping = {str(k): [str(x) for x in (v or [])] for k, v in (raw.get("report_templates") or {}).items()}
    entries = []
    union: list[str] = []
    for p in files:
        sc = scan(p)
        pack = PACK_STARTER if p.stem.endswith(STARTER_SUFFIX) else PACK_EXPANSION
        entry: dict[str, Any] = {
            "file": p.name, "pack": pack, "sha256": _sha256(p), "size_bytes": p.stat().st_size,
            "placeholder_count": len(sc.placeholders), "placeholders": sorted(sc.placeholders),
            "todo_count": sc.todo_count, "report_templates": mapping.get(p.name, []),
        }
        old = source.get(p.name)
        if old is not None:
            entry.update({k: old[k] for k in _CARRIED if k in old})
        entries.append(entry)
        union.extend(n for n in sc.placeholders if n not in union)
    return {
        "_说明": ["由 zt.report.docx_template.build_manifest 从本目录的 docx 重算写出，不手工登记；",
                  "tests/report 核对本文件与 docx 一致。report_templates 来自 mapping.yaml（待签核）。",
                  "扩充包的 no / folder_name / folder_id_suffix / source / headings / amount_cells 取自旧仓登记清单 "
                  + SOURCE_MANIFEST_FILE + "（原样迁入）。"],
        "count": len(entries),
        "placeholder_union": sorted(union),
        "templates": entries,
    }


def load_manifest(directory: Path) -> Optional[dict]:
    """读登记清单。

    :param directory: ``data/docx-templates``。
    :returns: 登记清单；文件不存在时为 None。
    :raises ZtError: 不是合法 JSON（``data_file_invalid``）。
    """
    path = directory / MANIFEST_FILE
    if not path.is_file():
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as e:
        raise ZtError("data_file_invalid", f"{MANIFEST_FILE} 不是合法 JSON：{e}") from None
