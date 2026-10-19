# 来源：engine-service/app/engine/costsheet.py（zhitong-cost-ai@991ff5cd，W2 搬入 zt；改动见本文件内「zt 改动」注释）
"""计价表（广联达 GBQ 导出 Excel 等）解析 —— 对账(reconcile)/ 指标(indicator)/ 投控检测(tk)
共用的取数层。纯 Python + openpyxl，不联网、不调模型，结果确定可复算。

WS08-A（issue #724）重写要点。旧版的根因是**不知道每一列是什么意思**：列词表只有
「工程量/合价/金额」，没有期次（本期/上期/累计/合同）也没有当事方（送审/审核）；
同长关键词取最左列；看不懂的表静默跳过；只在「解析出 0 行」时告警。于是进度款表里
「投标量 × 合同合价」被当成本期上报、门窗三段表的审定侧取成了「原合同」块、工日表
和标题行一起进了合计 —— 结果全错，warnings 却是 null。

现在的模型：
1. **表头块**（_detect_header_block）：前 30 行找锚行（序号/编码/名称任两个），向下最多
   吸收 2 行续行；组标题横向前填，每列得到一条列路径，如「施工单位/本期产值」。
2. **列角色**（_col_role）：量度 qty/price/amount × 期次 current/previous/cumulative/
   contract/dated × 当事方 submit/audit/delta × 量纲（万元）；「其中/人工费/暂估价」等分项列
   不参与绑定。
3. **绑定**（_bind）：side 按当事方过滤，basis 按期次过滤（auto：有本期列取本期，否则取
   合同或不分期列）。候选唯一才绑定，没有记缺失（None，绝不写 0.0），多于一个 = 歧义：
   该表不产出行，出 column_ambiguous（铁律 8：看不懂就说出来，不猜）。
   「一列不分期合价 + 多组分期工程量」（A1 市政形）用逐行 合价≈量×价 投票定期次。
4. **sheet 状态**：used / skipped_by_title / skipped_non_list / no_header / ambiguous /
   summary / unpriced，每张表都有原因；有清单却被跳过的表出 sheet_skipped_has_items；
   合计行勾稽偏差 >0.5% 出 sheet_total_mismatch；文件级覆盖率 ok / partial / poor。
5. **行分类**（_row_kind）：item / lump_sum 才是叶子；标题行（S.1、一、0104 砌筑工程…）
   只记作后续行的 section；编码列里写「分部分项合计/规费」的行是小计。
6. **汇总表与明细表二选一**：同簿有明细表时汇总表只用来勾稽，不计入。
7. **章节**：一律走 code_registry.classify_chapter（三个端点同一口径）。

审核汇总版式（进度款审核对比表：序号|工程名称|报审|审核|核减，没有编码/单位/量价列）
并入同一套列角色：side=None 保持旧语义（金额取报审），side=audit 取审核金额。
明细区以首个聚合行（累计/本期完成/合计…）截止 —— 其后的派生段行名会与明细区重名
（30002 实测「钢结构」），逐行过滤分不开，只能按段截断。完整保留合计/派生行、
金额用 Decimal 字符串、带单元格坐标的读法见 parse_audit_summary（WS12 报告取数用）。

去重口径（#22 D-12）：(code,name,amount) 只做**跨工作表**去重，同表完全相同行如实保留
（那是最典型的「重复计取」，留给 tk duplicate 规则判）；去重条数用 cross_sheet_dedup 披露。
"""
from __future__ import annotations

import datetime
import re
import unicodedata
from dataclasses import dataclass, field, replace
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from typing import Optional

from openpyxl.utils import get_column_letter

from zt.code_registry import registry as cr
from zt.costsheet.warning_codes import warn
from zt.costsheet.xlsx_tolerant import XlsxUnreadable, load_workbook_tolerant

_MAX_ROWS = 8000          # 单表安全上限
_HEADER_SCAN = 30         # 锚行只在前 30 行里找
_TOTAL_TOL = 0.005        # 合计勾稽容差 0.5%
_VOTE_HIT = 0.9           # 期次投票命中率门槛
_VOTE_TOL = 0.01          # 合价 ≈ 量×价 的相对容差 1%

# ---------------------------------------------------------------- 表级词表

# 工作表标题黑名单：综合单价分析/人材机/材料价/费用表 等非「分部分项清单」表，列头相似会污染合计与去重。
_SHEET_SKIP = ["综合单价分析", "单价分析", "综合单价计算", "人材机", "工料机", "主要材料", "材料价", "主要工料",
               "封面", "扉页", "总说明", "编制说明", "取费", "费用表", "费用汇总", "工程费用汇总", "投标总价", "封-",
               # 真实广联达导出的主要污染源:费用表(表-13)、措施综合单价计算表(表-17)、其他项目/规费税金
               "其他项目", "规费", "税金", "暂估价", "暂列金额", "计日工", "总承包服务",
               # WS08-03：市政控制价里的「主要工日一览表」「主要机械台班一览表」曾被计入，
               # 总造价比招标控制价还高 2,253,617.75（QA:indicator-02）
               "工日", "机械台班", "一览表", "人材机汇总", "材料汇总"]
# sheet 名看不出、要看表内前几行标题的非清单表（sheet 名常是「表-09」这类编号）
_TITLE_SKIP_IN_ROWS = ["主要工日", "机械台班", "主要材料", "人材机", "综合单价分析", "费用汇总", "投标报价费用表"]
# 标题像计价表：没认出表头时据此判断「这张表本该有清单」（与 ≥3 个清单编码二选一）
_LISTY_TITLE = ["清单与计价表", "分部分项", "措施项目", "计价表"]

# ---------------------------------------------------------------- 列词表

# 身份列：(字段, 关键词按优先级排)。排位越靠前越优先；同排位取最左列。
_IDENT_KW = [
    ("seq", ["序号"]),
    ("code", ["项目编码", "清单编码", "项目编号", "清单编号", "编码", "编号"]),
    ("name", ["项目名称", "清单名称", "子目名称", "费用名称", "工程名称", "名称"]),
    ("feature", ["项目特征描述", "项目特征", "特征描述", "特征"]),
]
# 「单位」只认整格是单位的列：施工单位/建设单位/单位工程 都含「单位」二字
_UNIT_RE = re.compile(r"^(?:计量)?单位(?:\(.*\))?$")

# 量度：同一段里先判单价、再判金额、最后判工程量（「本期完成工程产值」是金额不是量）。
#   qty   工程量/数量/面积：通用；「完成工程」：A1 市政表头被截成「本期完成工程」（无「量」字）
#   price 单价：综合单价/全费用单价
#   amount 合价/金额：广联达表-08；产值：进度款（浙大「本期产值」、浙四「本月上报产值」）；总价
_MEASURE_KW = [
    ("price", ["单价"]),
    ("amount", ["合价", "金额", "产值", "总价"]),
    ("qty", ["工程量", "工程数量", "数量", "面积", "完成工程", "完成量"]),
]
# 期次：按这个顺序判，先命中先得 ——
#   「上期累计已完工程量」「截止上月累计已审核产值」是上期（先判上期）；
#   「截止本月审核累计」「累计（含本次）」是累计（累计先于本期/本月/本次判）。
_PERIOD_KW = [
    ("previous", ["上期", "截止上月", "上月"]),
    ("cumulative", ["累计", "至今", "截止本期", "截止本月"]),
    ("current", ["本期", "本月", "当期", "本次"]),
]
# 带具体月份/期号的列组：30003「2023年6月上报产值」「24年1月审核产值」、「第5期上报」
_DATED_RE = re.compile(r"\d{1,2}月|第\d+期")
# 合同/投标口径：浙大「投标价(元)」、门窗「原合同」、结算「合同金额」「控制价」
_CONTRACT_KW = ["合同", "投标", "中标", "控制价", "招标"]
# 当事方：先判增减（「核增核减」），再判审核，最后判送审。
#   submit 浙大组标题「施工单位」、浙四「本月上报产值」、审定单「送审金额」、30002「报审金额」
#   audit  浙四「本月审核产值」、门窗「审核」、审定单「审定金额」、30003「咨询审定产值」
_PARTY_KW = [
    ("delta", ["核增", "核减", "增减", "审减", "审增"]),
    ("audit", ["审核", "审定", "核定", "监理审核", "投控审核", "咨询审定"]),
    ("submit", ["送审", "报审", "上报", "申报", "施工单位", "施工方", "承包人"]),
]
# 分项列：浙大 R6「人工费|机械费|暂估价」、表-08「其中：暂估价」—— 是合价的组成，不是合价
_COMPONENT_KW = ["其中", "人工费", "机械费", "暂估价", "人机"]
# 量纲：金额（万元）/单价（万）→ ×10000（WS08-15 长兴设备估算表算出「单方 0.04 元/㎡」的根因）
_WAN_KW = ["万元", "(万)"]
# 组标题不横向延伸的：备注/比例/说明 下面的空格不属于它
_NO_EXTEND = ["备注", "比例", "说明", "占比", "%"]

# ---------------------------------------------------------------- 行词表

_SUBTOTAL_ANY = ["小计", "合计", "总计"]
# 整格就是这些名字的行是小计/费用汇总行（「安全文明施工费」这类总价措施项不在此列）
_SUBTOTAL_EXACT = {"规费", "税金", "措施费", "施工组织措施费", "施工技术措施费", "组织措施费", "技术措施费",
                   "分部分项合计", "分部分项工程合计", "分部分项工程费", "措施项目合计", "其他项目合计"}
_FEE_SUBTOTAL = ("规费", "税金", "措施费")
_EXPLICIT_TOTAL = ("分部分项合计", "分部分项工程合计", "清单合计")
# 审核汇总表明细区终止词:名称命中即进入 累计/派生段,该表后续行全部不取。
_AUDIT_STOP_KW = ["累计", "总产值", "本期完成", "本期支付", "截止上期", "之和", "合计", "总计", "备注"]
_HAN = re.compile(r"[一-鿿]")
_CN_NUM = "一二三四五六七八九十"
_TITLE_SEQ_RE = re.compile(rf"^(?:[{_CN_NUM}]+[、.．]?|[（(][{_CN_NUM}]+[）)]|[A-Z]\.\d+(?:\.\d+)*)$")
_TITLE_NAME_RE = re.compile(
    rf"^(?:[{_CN_NUM}]+[、.．]|[（(][{_CN_NUM}]+[）)]|[A-Z]\.\d+|0[1-9]\d{{2}}(?:\d{{2}})?(?![\d.]))")
_NUM_SEQ_RE = re.compile(r"^\d+(?:\.\d+)*$")

ITEM, LUMP_SUM, TITLE, SUBTOTAL, SKIP = "item", "lump_sum", "title", "subtotal", "skip"

_PERIOD_LABEL = {"current": "本期", "previous": "上期", "cumulative": "累计", "contract": "合同",
                 "dated": "指定月份", None: "不分期"}
_PARTY_LABEL = {"submit": "送审", "audit": "审核", "delta": "核增减"}

# WS26-e（07-009，P10）：文件 meta.basis_available —— 清单表金额列里出现过的期次。plain（不分期）与 contract
# 分开报：TK-F01 那份表头是「工程量 / 综合单价 / 合价」，是不分期的合价列，不是「合同价」列（评审 S7）
BASIS_KEYS = ("plain", "contract", "current", "cumulative", "previous", "dated")
_BASIS_COL_TEXT = {"plain": "不分期的合价列", "contract": "合同价列", "current": "本期列", "cumulative": "累计列",
                   "previous": "上期列", "dated": "按月份分组的列"}
# 对比口径下拉的名字（前端 BASIS_OPTIONS 的 label 前半句）
BASIS_NAME = {"auto": "自动", "current": "本期", "cumulative": "累计", "contract": "合同价", "previous": "上期"}


# ================================================================ 数值

def _to_float(v) -> Optional[float]:
    """单元格值 → float;无法解析回 None(不抛)。日期/布尔/百分比均视为无效返回 None。"""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (datetime.datetime, datetime.date, datetime.time)):
        return None  # 数字列被误设日期格式 → 不臆测年份数字
    if isinstance(v, (int, float)):
        try:
            f = float(v)
        except (TypeError, ValueError):
            return None
        return None if f != f else f  # NaN → None
    s = str(v).strip()
    if not s or s in ("-", "—", "－", "/", "\\", "N/A", "na", "无"):
        return None
    if "%" in s or "％" in s:
        return None  # 金额/工程量列出现百分比=异常数据,不还原量纲臆测
    neg = s.startswith("(") and s.endswith(")")  # 会计负数 (123)
    s = s.translate(str.maketrans("０１２３４５６７８９．，", "0123456789.,"))
    s = s.replace(",", "").replace("¥", "").replace("￥", "")
    s = s.replace("元", "").strip()
    m = re.search(r"-?\d+(?:\.\d+)?", s)
    if not m:
        return None
    f = float(m.group())
    return -f if neg else f


def _to_decimal(v) -> Optional[Decimal]:
    """单元格值 → Decimal（parse_audit_summary 用）。float 走 repr，不引入二进制尾差。"""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, int):
        return Decimal(v)
    if isinstance(v, float):
        if v != v:
            return None
        return Decimal(repr(v))
    f = _to_float(v)
    if f is None:
        return None
    s = str(v).strip().translate(str.maketrans("０１２３４５６７８９．，", "0123456789.,"))
    m = re.search(r"-?\d+(?:\.\d+)?", s.replace(",", ""))
    try:
        d = Decimal(m.group()) if m else Decimal(repr(f))
    except InvalidOperation:
        return None
    return -d if (s.startswith("(") and s.endswith(")")) else d


_CENT = Decimal("0.01")


def to_cents(v) -> Optional[int]:
    """金额 → 整数分（四舍五入 ROUND_HALF_UP）。float 走 repr，不带二进制尾差。

    WS08-19：对账比较与解析取整共用这一个函数。旧版解析用 round(x, 2)（二进制浮点上的
    round，4316.765 可能舍成 4316.76），对账再用 abs(a-b) <= 0.01 比，两处口径不一，
    1 分钱的舍入噪声被判成「调整」。"""
    if v is None or isinstance(v, bool):
        return None
    try:
        d = v if isinstance(v, Decimal) else Decimal(repr(float(v)))
    except (TypeError, ValueError, InvalidOperation):
        return None
    if d != d:      # NaN
        return None
    return int(d.quantize(_CENT, rounding=ROUND_HALF_UP) * 100)


def round_money(v) -> Optional[float]:
    """金额取到分（ROUND_HALF_UP），与 to_cents 同一口径；取不到回 None，不写 0。"""
    c = to_cents(v)
    return None if c is None else c / 100


_NUMBER_CELL_RE = re.compile(r"^[-+(（]?[\d,，.．]+[)）]?$")


def _is_number_cell(v) -> bool:
    if isinstance(v, bool):
        return False
    if isinstance(v, (int, float)):
        return True
    return bool(_NUMBER_CELL_RE.match(re.sub(r"\s+", "", str(v))))


def _norm(s) -> str:
    return re.sub(r"\s+", "", str(s or ""))


def _txt(v) -> str:
    return "" if v is None else str(v).strip()


# ================================================================ 表头块与列角色

@dataclass
class _Col:
    idx: int
    path: tuple
    ident: Optional[str] = None
    ident_rank: int = 99
    measure: Optional[str] = None
    period: Optional[str] = None
    party: Optional[str] = None
    scale: int = 1
    component: bool = False

    @property
    def text(self) -> str:
        return "/".join(self.path)

    @property
    def key(self) -> str:
        return ".".join(p for p in (self.period, self.party, self.measure) if p)


@dataclass
class _HeaderBlock:
    rows: list          # 0 起的行下标
    cols: list          # [_Col]
    data_start: int


def _hdr_cell(v) -> str:
    """表头单元格 → 规范文本；数值、「工程名称：xx」信息栏、超长文本一律视为空。"""
    if v is None or isinstance(v, (bool, int, float, datetime.date, datetime.time)):
        return ""
    s = unicodedata.normalize("NFKC", str(v))
    s = re.sub(r"\s+", "", s)
    if not s or len(s) > 40:
        return ""
    m = re.search(r"[:：]", s)
    if m and m.end() < len(s):
        return ""
    return s


def _ident_of(seg: str) -> tuple[Optional[str], int]:
    if _UNIT_RE.match(seg):
        return "unit", 0
    for ident, kws in _IDENT_KW:
        for rank, kw in enumerate(kws):
            if kw in seg:
                return ident, rank
    return None, 99


def _measure_of(seg: str) -> Optional[str]:
    for m, kws in _MEASURE_KW:
        if any(k in seg for k in kws):
            return m
    return None


def _period_of(text: str) -> Optional[str]:
    for p, kws in _PERIOD_KW:
        if any(k in text for k in kws):
            return p
    if _DATED_RE.search(text):
        return "dated"
    if any(k in text for k in _CONTRACT_KW):
        return "contract"
    return None


def _party_of(text: str) -> Optional[str]:
    for p, kws in _PARTY_KW:
        if any(k in text for k in kws):
            return p
    return None


def _col_role(idx: int, path: tuple) -> _Col:
    """列路径 → 角色。量度与身份看最深的一段（叶子优先），期次/当事方/量纲看整条路径。"""
    c = _Col(idx=idx, path=path)
    if not path:
        return c
    joined = "".join(path)
    for seg in reversed(path):
        m = _measure_of(seg)
        if m:
            c.measure = m
            break
    ident, rank = _ident_of(path[-1])
    if ident and not _measure_of(path[-1]):
        c.ident, c.ident_rank, c.measure = ident, rank, None
    if c.measure:
        c.component = any(k in joined for k in _COMPONENT_KW)
        c.period = _period_of(joined)
        c.party = _party_of(joined)
        c.scale = 10000 if any(k in joined for k in _WAN_KW) else 1
    return c


def _has_header_kw(s: str) -> bool:
    if not s:
        return False
    if _UNIT_RE.match(s) or _ident_of(s)[0] or _measure_of(s) or _party_of(s):
        return True
    return any(k in s for k in _COMPONENT_KW + _WAN_KW + [k for _p, kws in _PERIOD_KW for k in kws])


def _is_anchor(row: list) -> bool:
    idents, measure = set(), False
    for v in row:
        h = _hdr_cell(v)
        if not h:
            continue
        ident, _r = _ident_of(h)
        if ident:
            idents.add(ident)
        if _measure_of(h):
            measure = True
    if len(idents & {"seq", "code", "name"}) >= 2:
        return True
    return bool(idents & {"code", "name"}) and measure


def _is_continuation(row: list, anchor_row: list) -> bool:
    """表头续行：没有清单编码、非空格以文字为主、至少一格是表头词；
    并且不在锚行的 序号/编码/名称 列里写别的字。

    最后一条挡的是紧贴表头的分部标题行：「二 | 单价措施项目」含「单价」二字，
    若被吸进表头，名称列的路径就成了「项目名称/单价措施项目」，被当成单价列。
    续行只会在这几列重复同一个表头词（30003 的 R4 就是把「序号|项目编码|项目名称」重抄一遍）。
    """
    vals = [v for v in row if v is not None and str(v).strip()]
    if not vals:
        return False
    for v in vals:
        if cr.normalize_code(v).kind in (cr.GB, cr.SUPPLEMENT):
            return False
    # 「6月上报产值（元）」是表头文字：_to_float 会从里面抠出 6，这里只认整格是数的
    texts = [v for v in vals if not _is_number_cell(v)]
    if len(texts) / len(vals) < 0.6:
        return False
    for c, av in enumerate(anchor_row):
        h = _hdr_cell(av)
        if _ident_of(h)[0] in ("seq", "code", "name") and c < len(row):
            v = row[c]
            if v is not None and str(v).strip() and _hdr_cell(v) != h:
                return False
    return any(_has_header_kw(_hdr_cell(v)) for v in vals)


def _detect_header_block(rows: list) -> Optional[_HeaderBlock]:
    """前 30 行找锚行（序号/编码/名称任两个；或编码/名称 + 量度词），向下最多吸收 2 行续行，
    组标题横向前填（只覆盖下一行有字的列），得到每列的列路径。"""
    anchor = next((ri for ri in range(min(_HEADER_SCAN, len(rows))) if _is_anchor(rows[ri])), None)
    if anchor is None:
        return None
    block = [anchor]
    for nxt in (anchor + 1, anchor + 2):
        if nxt < len(rows) and _is_continuation(rows[nxt], rows[anchor]):
            block.append(nxt)
        else:
            break
    width = max(len(rows[r]) for r in block)
    grid = [[_hdr_cell(rows[r][c]) if c < len(rows[r]) else "" for c in range(width)] for r in block]
    filled = [row[:] for row in grid]
    # 自下而上填：上一层组标题要看「已填好的下一层」有没有字（浙大三行表头：施工单位 → 本期 → 工程量/产值）
    for bi in range(len(block) - 2, -1, -1):
        below = filled[bi + 1]
        for c in range(width):
            t = grid[bi][c]
            if not t or any(k in t for k in _NO_EXTEND):
                continue
            c2 = c + 1
            while c2 < width and not grid[bi][c2] and below[c2]:
                filled[bi][c2] = t
                c2 += 1
    cols = []
    for c in range(width):
        segs: list[str] = []
        for bi in range(len(block)):
            s = filled[bi][c]
            if s and (not segs or segs[-1] != s):
                segs.append(s)
        cols.append(_col_role(c, tuple(segs)))
    return _HeaderBlock(rows=block, cols=cols, data_start=block[-1] + 1)


# ================================================================ 绑定

@dataclass
class _Binding:
    idents: dict = field(default_factory=dict)       # seq/code/name/feature/unit → _Col
    qty: Optional[_Col] = None
    price: Optional[_Col] = None
    amount: Optional[_Col] = None
    audit_amount: Optional[_Col] = None               # 送审侧附带审核金额（投控保留 audit_amount）
    delta_amount: Optional[_Col] = None
    side: Optional[str] = None                        # 实际生效的当事方
    want: str = "contract"                            # 实际生效的期次口径
    ambiguous: list = field(default_factory=list)     # [(measure, [列路径…])]
    party_missing: bool = False
    has_measure: bool = False

    @property
    def basis_used(self) -> str:
        col = self.amount or self.qty
        if col is None:
            return self.want
        return ".".join(p for p in (col.period or "plain", col.party) if p)

    @property
    def basis_label(self) -> str:
        col = self.amount or self.qty
        if col is None:
            return _PERIOD_LABEL.get(self.want, self.want)
        lab = _PERIOD_LABEL.get(col.period, col.period or "")
        return lab + (f"·{_PARTY_LABEL[col.party]}" if col.party else "")


def _period_ok(period: Optional[str], want: str) -> bool:
    if want == "contract":
        return period in ("contract", None)
    return period == want


def _measure_cols(cols: list) -> list:
    return [c for c in cols if c.measure and not c.component and c.party != "delta"]


def _pick(ms: list, measure: str, want: str, side: Optional[str]) -> list:
    cands = [c for c in ms if c.measure == measure and _period_ok(c.period, want)]
    mine = [c for c in cands if side and c.party == side]
    if mine:
        return mine
    return [c for c in cands if c.party is None]


def _narrow(cands: list, want: str, explicit: bool = True) -> list:
    """调用方**明确**要合同口径（basis=contract，如指标）时，「合同」列与不分期列并存取合同那一列。
    auto 落到合同/不分期时两者并存就是歧义（「合同金额 | 结算金额」不知道比哪列），不替用户选。"""
    if len(cands) > 1 and want == "contract" and explicit:
        exact = [c for c in cands if c.period == "contract"]
        if len(exact) == 1:
            return exact
    return cands


def _eff_side(cols: list, side: Optional[str]) -> Optional[str]:
    parties = {c.party for c in _measure_cols(cols) if c.party}
    if side in ("submit", "audit"):
        return side
    # side=None：表头分了送审/审核时按送审取（与审核汇总表的旧语义、tk/indicator 的 submit 语义一致）
    return "submit" if parties else None


def _vote_amount_period(rows: list, data_start: int, amount: _Col, qty_by_period: dict,
                        price: Optional[_Col]) -> Optional[str]:
    """A1 市政形：只有一列不分期的合价，却有多组分期的工程量。逐行检验 合价≈该期量×单价
    （容差 1%），命中率 ≥90% 的期次胜出；本期与累计因上期为空而同时命中时取本期。
    → 期次 / "ambiguous" / None（没有可投票的行，如无价清单）。"""
    if price is None:
        return None
    hits = {p: 0 for p in qty_by_period}
    n = 0
    for row in rows[data_start:]:
        a = _to_float(_cell(row, amount))
        p = _to_float(_cell(row, price))
        if not a or not p:
            continue
        n += 1
        for per, qc in qty_by_period.items():
            q = _to_float(_cell(row, qc))
            if q is not None and abs(a - q * p) <= _VOTE_TOL * abs(a):
                hits[per] += 1
    if n == 0:
        return None
    winners = [per for per, h in hits.items() if h / n >= _VOTE_HIT]
    if len(winners) == 1:
        return winners[0]
    if winners and "current" in winners and set(winners) <= {"current", "cumulative"}:
        return "current"
    return "ambiguous"


def _bind(cols: list, side: Optional[str], basis: str, rows: list, data_start: int) -> _Binding:
    b = _Binding()
    b.side = _eff_side(cols, side)
    ms = _measure_cols(cols)
    b.has_measure = bool(ms)
    parties = {c.party for c in ms if c.party}
    if side in ("submit", "audit") and parties and side not in parties:
        b.party_missing = True

    # 期次投票：一列不分期合价 + ≥2 组分期工程量（只看本侧/无当事方的列）
    in_scope = [c for c in ms if c.party in (b.side, None)]
    amts = [c for c in in_scope if c.measure == "amount"]
    qtys = {}
    for c in in_scope:
        if c.measure == "qty" and c.period and c.period != "dated":
            qtys.setdefault(c.period, c)
    if len(amts) == 1 and amts[0].period is None and len(qtys) >= 2 \
            and {"current", "cumulative"} & set(qtys):
        prices = [c for c in in_scope if c.measure == "price" and c.period in ("contract", None)]
        voted = _vote_amount_period(rows, data_start, amts[0], qtys,
                                    prices[0] if len(prices) == 1 else None)
        if voted == "ambiguous":
            b.ambiguous.append(("amount", [amts[0].text] + [q.text for q in qtys.values()]))
            return b
        if voted:
            ms = [replace(c, period=voted) if c is amts[0] else c for c in ms]

    if basis in ("current", "cumulative", "contract", "previous"):
        b.want = basis
    else:   # auto：有本期列就取本期，否则取合同或不分期列（D45）
        b.want = "current" if (_pick(ms, "amount", "current", b.side)
                               or _pick(ms, "qty", "current", b.side)) else "contract"

    explicit = basis == "contract"
    for measure in ("qty", "price", "amount"):
        cands = _narrow(_pick(ms, measure, b.want, b.side), b.want, explicit)
        if not cands and measure == "price" and b.want != "contract":
            # 单价跨期共用：本期表里只有合同综合单价是常态（浙大「投标价/综合单价」）
            cands = _narrow(_pick(ms, "price", "contract", b.side), "contract")
        if len(cands) == 1:
            setattr(b, measure, cands[0])
        elif len(cands) > 1:
            b.ambiguous.append((measure, [c.text for c in cands]))
    if b.side == "submit" and "audit" in parties:
        aud = _narrow(_pick(ms, "amount", b.want, "audit"), b.want, explicit)
        if len(aud) == 1 and aud[0].party == "audit":
            b.audit_amount = aud[0]
    deltas = [c for c in cols if c.measure == "amount" and c.party == "delta" and not c.component
              and _period_ok(c.period, b.want)]
    if len(deltas) == 1:
        b.delta_amount = deltas[0]

    for ident in ("seq", "code", "name", "feature", "unit"):
        cands = [c for c in cols if c.ident == ident]
        if cands:
            b.idents[ident] = min(cands, key=lambda c: (c.ident_rank, c.idx))
    return b


def _cell(row: list, col: Optional[_Col]):
    if col is None or col.idx >= len(row):
        return None
    return row[col.idx]


def _scaled(row: list, col: Optional[_Col]) -> Optional[float]:
    """该列的数值换算到元。凡是落进 item 字段或合计勾稽的金额/单价都走这里，
    漏一处就是一侧万元一侧元、差 4 个数量级且不报错（PR #747 评审：审定单 amount 取成 100）。"""
    v = _to_float(_cell(row, col))
    return v * col.scale if v is not None else None


def _warn_scaled(sh, cols, file) -> None:
    """只对真正换算后落进字段的列说「已换算」；同一列只说一次。"""
    seen = set()
    for col in cols:
        if col is not None and col.scale != 1 and col.idx not in seen:
            seen.add(col.idx)
            sh.warnings.append(warn("amount_scaled_wan",
                                    f"工作表「{sh.title}」列「{col.text}」单位为万元，已乘以 10000 换算为元",
                                    file=file, sheet=sh.title))


# ================================================================ 行分类

def _core_label(t: str) -> str:
    """「1、规费」「（一）税金（元）」→「规费」「税金」。"""
    t = unicodedata.normalize("NFKC", t)
    t = re.sub(r"^(?:\d+(?:\.\d+)*|[一二三四五六七八九十]+)[、.]?", "", t)
    t = re.sub(r"^\([一二三四五六七八九十\d]+\)", "", t)
    t = re.sub(r"\(.*?\)$", "", t)
    return t.strip("：:、.")


def _is_subtotal(seq: str, code_s: str, name: str) -> bool:
    if code_s and _HAN.search(code_s):
        return True        # 浙四：编码列里写「分部分项合计 / 规费 / 施工组织措施费」
    for t in (seq, code_s, _norm(name)):
        if not t:
            continue
        if any(k in t for k in _SUBTOTAL_ANY) or _core_label(t) in _SUBTOTAL_EXACT:
            return True
    return False


def _row_kind(seq: str, code_raw, name: str, unit: str, qty, price, amount,
              *, measure_sheet: bool) -> str:
    """item（清单项）/ lump_sum（总价措施项）/ title（标题，记作后续行的 section）/ subtotal / skip。

    只有 item 与 lump_sum 是叶子。旧版 is_leaf 接受 qty=0 的「S.1脚手架工程」标题行、
    unit=项 的合计行，二者都进了合计（QA:tk-08、reconcile-05）。
    """
    code_s = _norm(code_raw)
    nm = unicodedata.normalize("NFKC", _norm(name))
    sq = unicodedata.normalize("NFKC", _norm(seq))
    if _is_subtotal(sq, code_s, name):
        return SUBTOTAL
    if (sq and _TITLE_SEQ_RE.match(sq)) or (nm and _TITLE_NAME_RE.match(nm)) \
            or (code_s and _TITLE_SEQ_RE.match(code_s)):
        return TITLE
    nc = cr.normalize_code(code_raw)
    has_code = nc.is_item_code
    if measure_sheet and not has_code and _NUM_SEQ_RE.match(sq or "") and nm and amount is not None:
        return LUMP_SUM
    if has_code and (unit or qty is not None or amount is not None):
        return ITEM
    # 单位、名称必须带字：挡住表头下面那行「1|2|3|4|5|6|7」列号行
    if not has_code and unit and not _NUM_SEQ_RE.match(unit) and nm and not _NUM_SEQ_RE.match(nm) \
            and (qty or 0) > 0 and (price is not None or amount is not None):
        return ITEM
    if not has_code and not unit:
        return TITLE if nm else SKIP
    return SKIP


# ================================================================ sheet 解析

@dataclass
class _Sheet:
    title: str
    detail: dict
    items: list = field(default_factory=list)
    warnings: list = field(default_factory=list)
    summary_candidate: bool = False
    listy: bool = False
    skipped_estimate: Optional[float] = None
    has_priced_values: bool = False


def _title_text(rows: list, upto: int) -> str:
    parts = []
    for row in rows[:upto]:
        for v in row:
            if isinstance(v, str) and v.strip():
                parts.append(_norm(v))
    return "".join(parts)


# 表头上方写单位工程的标签（GB50500 表-08 抬头「单位（专业）工程名称：6#楼【建筑】」）。
# 刻意不认裸「工程名称」：它在多数导出件里写的是整个项目名，拿它分组会把各栋楼并成一组。
_UNIT_PROJECT_LABELS = ("单位及专业工程名称", "单位(专业)工程名称", "单位工程名称", "专业工程名称")
_UNIT_PROJECT_CUT = re.compile(r"\s{2,}|标段[:：]|第\s*\d+\s*页")


def _unit_project(rows: list, upto: int, labels: tuple = _UNIT_PROJECT_LABELS) -> Optional[str]:
    """表头块上方的「单位工程名称：xx」→ xx；值在同格冒号后或右侧第一个非空格。认不出回 None。"""
    for row in rows[:upto]:
        cells = [_txt(v) for v in row]
        for i, t in enumerate(cells):
            if not t:
                continue
            tn = unicodedata.normalize("NFKC", t)
            for lab in labels:
                if lab not in tn:
                    continue
                after = tn.split(lab, 1)[1].lstrip(":： ")
                if not after:
                    after = next((unicodedata.normalize("NFKC", c) for c in cells[i + 1:] if c), "")
                after = _UNIT_PROJECT_CUT.split(after, 1)[0].strip()
                if after:
                    return after[:120]
    return None


# ---------------------------------------------------------------- 费用表合计（WS08-16；WS26-h 分层取数）

# 费用表合计的读取只在这里（总纲 K-R18 / K-R44）：指标、字段提取候选、报告金额事实都调这一份。
# 行名**整格相等**才算（NFKC、去空白、去前导序号后比较）——「其中：安全文明施工费合计」
# 「分部分项合计」都不许被当成总造价（铁律 3：不模糊取数）。
_FEE_SHEET_KW = ("费用表", "费用汇总", "汇总", "投标总价", "招标控制价", "工程造价")
_FEE_TOTAL_LABELS = ("招标控制价合计", "投标报价合计", "工程造价", "合计(含税)")
_FEE_SEQ_PREFIX = re.compile(r"^(?:\d+(?:\.\d+)*|[一二三四五六七八九十]+|[(][一二三四五六七八九十\d]+[)])[、.]?")
# 泛名：单位工程费用表、清单表、其他项目表里到处有「合计」，所以它不进白名单，只在项目级表里认，
# 而且必须与各单位工程费用表之和勾稽通过才采用（见 resolve_fee_total）。
_FEE_GENERIC_TOTAL = "合计"
# 金额列的表头（整格相等）。项目级费用表「金额(元)」右边还有暂估价、安全文明施工费、规费、税金四列，
# 取「最靠右的数值格」会拿到规费——必须按表头定列。
_FEE_AMOUNT_HEADERS = ("金额(元)", "金额", "合计(元)", "造价(元)")
# 表名或表内标题带这些词的不是造价汇总表（GB50500 表-12「其他项目清单与计价汇总表」的表头也是
# 「金额(元)」、末行也叫「合计」），不取任何合计。
_FEE_EXCLUDE_KW = ("其他项目", "措施", "规费", "税金", "暂估", "暂列", "计日工", "总承包服务")
# 项目级合计与各单位工程费用表之和的勾稽容差（D-R13）：项目级合计常取整到元（最多差 0.5 元），
# 单位工程表到分。按比例（0.5%）放会漏掉一整张小单位工程表还显示「勾稽一致」，所以用绝对值 1 元。
FEE_TOTAL_TOL_CENTS = 100
FEE_UNIT, FEE_SINGLE, FEE_PROJECT, FEE_SPECIALTY = "unit", "single", "project", "specialty"
# 分层短语：「X工程…费用表 / 汇总表」必须连在**同一格**里（表内标题格或表名），中间不许有冒号。
# 抬头标签不算数——广联达的专业工程费用表抬头写的是「单位工程名称：一期02标-1#-地上工程」，把它算进来
# 专业工程表就成了 unit 层，和它上面那张单位工程表相加，总造价正好翻倍（WS26-h 评审，样本库 78 份）。
# 「单位(专业)工程…费用表」是浙江计价规则表10.2.2-13 的原名，不分两级的软件导出件就用这个名字，归 unit。
_FEE_TIER_PHRASE = re.compile(r"(单位\(专业\)|单位及专业|单位|专业|单项)工程[^:]*?(?:费用表|汇总表)")
_FEE_TIER_OF = {"单位(专业)": FEE_UNIT, "单位及专业": FEE_UNIT, "单位": FEE_UNIT,
                "专业": FEE_SPECIALTY, "单项": FEE_SINGLE}
_FEE_HEAD_LABEL = "名称:"


def _fee_label(v) -> str:
    t = re.sub(r"\s+", "", unicodedata.normalize("NFKC", _txt(v)))
    return _FEE_SEQ_PREFIX.sub("", t)


def _fee_plain(v) -> str:
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", _txt(v)))


def _fee_amount(v) -> Optional[float]:
    """合计行的金额格 → 数。数值格直接取；文本格只有整格是数字（可带千分位）才取 ——
    广联达导出件的单位工程费用表金额全是文本「15,869,272.04」，左边一格是公式文字「1+2+3+4+5」，
    后者通不过 _is_number_cell，不会被拆出一个 1 来。"""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, str) and not _is_number_cell(v):
        return None
    return _to_float(v)


def _fee_header(rows: list) -> tuple[Optional[int], Optional[int], bool]:
    """费用表的表头行与金额列 → (行下标, 列下标, 是否有歧义)。

    前 _HEADER_SCAN 行里第一行含「金额(元) / 金额 / 合计(元) / 造价(元)」整格的就是表头行；
    该行恰 1 个这样的格才给列下标，≥2 个（并列两组金额）回 (行, None, True)：不知道取哪一列，不猜。
    找不到表头时退而找「序号」行，只用来划出表头上方的标题区，列下标为 None。"""
    seq_row = None
    for ri, row in enumerate(rows[:_HEADER_SCAN]):
        cols = [ci for ci, v in enumerate(row) if isinstance(v, str) and _fee_plain(v) in _FEE_AMOUNT_HEADERS]
        if cols:
            return (ri, cols[0], False) if len(cols) == 1 else (ri, None, True)
        if seq_row is None and any(isinstance(v, str) and _fee_plain(v) == "序号" for v in row):
            seq_row = ri
    return seq_row, None, False


def _fee_tier_phrases(cells) -> set:
    """这些格子里写着哪几层的表名 → {unit / specialty / single}。每格只看第一个「…名称:」标签之前的部分。"""
    found = set()
    for c in cells:
        m = _FEE_TIER_PHRASE.search(_fee_plain(c).split(_FEE_HEAD_LABEL, 1)[0])
        if m:
            found.add(_FEE_TIER_OF[m.group(1)])
    return found


def _fee_tier(title: str, rows: list, header_row: Optional[int]) -> tuple[bool, Optional[str]]:
    """这张表是哪一层 → (是否排除, unit / specialty / single / project / None)。按顺序判，命中即止（正向判定）。

    只看表名与**表头上方**的标题行（最多 5 行）。不能把表头行、数据行算进来：项目级费用表的表头
    第二行就写着「暂估价 | 规费 | 税金」，单位工程费用表第 3、4 行就是「措施项目费」「其他项目费」，
    整段前 5 行都拿来比排除词，三层表会全部被排除。认不出表头时没法区分标题与正文，前 5 行全算标题
    （宁可不取，铁律 3）。

    单位 / 专业 / 单项按表内标题格里的短语判（_FEE_TIER_PHRASE）；标题格里没有才看表名——表名是人可以
    随手改的，表内标题是表式本身。标题格里同时写着两层的短语 → 判不出，按未分层处理，不猜。"""
    upto = min(5, len(rows)) if header_row is None else min(5, header_row)
    text = unicodedata.normalize("NFKC", _norm(title) + _title_text(rows, upto))
    if any(k in text for k in _FEE_EXCLUDE_KW):
        return True, None
    found = (_fee_tier_phrases(v for row in rows[:upto] for v in row if isinstance(v, str))
             or _fee_tier_phrases([title]))
    if found:
        return False, next(iter(found)) if len(found) == 1 else None
    if "招标控制价费用表" in text or "投标报价费用表" in text or ("建设项目" in text and "汇总表" in text):
        return False, FEE_PROJECT
    return False, None


def _fee_row_cell(row: list, ci: int, amount_col: Optional[int]) -> Optional[int]:
    """费用表某一行（行名在第 ci 格）取金额的那一格的列下标（0 起）：有金额表头就只认那一列；找不到表头时
    只有行名右侧恰好 1 个数值格才是它。取不到回 None。"""
    if amount_col is not None:
        return amount_col if ci < amount_col < len(row) and _fee_amount(row[amount_col]) is not None else None
    nums = [k for k in range(ci + 1, len(row)) if _fee_amount(row[k]) is not None]
    return nums[0] if len(nums) == 1 else None


def _fee_row_amount(row: list, ci: int, amount_col: Optional[int]) -> Optional[float]:
    """费用表某一行（行名在第 ci 格）的金额（取哪一格见 _fee_row_cell）。scan_fee_sheet 与
    scan_measure_fee_row 共用这一个读法。"""
    k = _fee_row_cell(row, ci, amount_col)
    return None if k is None else _fee_amount(row[k])


def _fee_amount_str(v) -> Optional[str]:
    """金额格 → 到分的十进制串（Decimal，不经 float：文本格「15,869,272.21」、数值格 318375872 → 「318375872.00」）。"""
    d = _to_decimal(v)
    return None if d is None else str(d.quantize(_CENT, rounding=ROUND_HALF_UP))


def _fee_display(v) -> str:
    """金额格的原样写法：文本格照抄（含千分位）；数值格按「常规」格式写（整数不带小数点）。
    只读载入拿不到 number_format，设了千分位或两位小数格式的数值格写出来与 Excel 里看到的不同（数值相同）。"""
    if isinstance(v, str):
        return v.strip()
    d = _to_decimal(v)
    if d is None:
        return _txt(v)
    return format(d.quantize(Decimal(1)), "f") if d == d.to_integral_value() else format(d, "f")


def scan_fee_sheet(title: str, rows: list) -> dict:
    """一张表 → {sheet, tier, hits}。tier：unit / specialty / single / project / None（未分层或不是费用表）。

    hits：[{sheet, row, label, amount, tier, generic, cell, amount_str, display}]。白名单行名在各层与未分层的表里都认；
    泛名「合计」只在 project 层认（generic=True）。金额取「金额」表头所在的那一列；找不到表头时
    只有合计行右侧恰好 1 个数值格才取它，否则这一行不出命中。"""
    out = {"sheet": title, "tier": None, "hits": []}
    head = _title_text(rows, min(5, len(rows)))
    if not any(k in title or k in head for k in _FEE_SHEET_KW):
        return out
    header_row, amount_col, ambiguous = _fee_header(rows)
    excluded, tier = _fee_tier(title, rows, header_row)
    if excluded:
        return out
    out["tier"] = tier
    if ambiguous:
        return out
    labels = {unicodedata.normalize("NFKC", lab) for lab in _FEE_TOTAL_LABELS}
    for ri, row in enumerate(rows):
        for ci, v in enumerate(row):
            if not isinstance(v, str):
                continue
            label = _fee_label(v)
            generic = tier == FEE_PROJECT and label == _FEE_GENERIC_TOTAL
            if label not in labels and not generic:
                continue
            k = _fee_row_cell(row, ci, amount_col)
            if k is not None:
                # cell / amount_str / display 是 WS24-j（报告金额事实，K-R18）只增的三个字段：报告要回链到「哪一格、
                # 原样写的是什么」，金额要 Decimal 串（不经 float）。取哪一格、采不采用都没变；指标不读这三个字段。
                # rows 从 A1 起逐行读（_read_rows，缺行补空），所以 ri+1、k+1 就是 Excel 的行号、列号
                out["hits"].append({"sheet": title, "row": ri + 1, "label": label,
                                    "amount": round_money(_fee_amount(row[k])), "tier": tier, "generic": generic,
                                    "cell": f"{get_column_letter(k + 1)}{ri + 1}",
                                    "amount_str": _fee_amount_str(row[k]), "display": _fee_display(row[k])})
            break
    return out


def extract_fee_totals(title: str, rows: list) -> list[dict]:
    """一张表里的「总造价」合计行 → [{sheet, row, label, amount, tier, generic, cell, amount_str, display}]
    （见 scan_fee_sheet）。
    要不要采用、采用哪一个，由 resolve_fee_total 按整份文件判。"""
    return scan_fee_sheet(title, rows)["hits"]


# WS26-p：单位工程费用表的「施工技术措施项目」行，用来和同一单位工程的措施综合单价计算表逐表勾稽。
_FEE_MEASURE_LABEL = "施工技术措施项目"
# 单位工程费用表的抬头写的是裸「工程名称：地下室-土石方工程」——在单位工程这一层它就是单位工程名
# （_unit_project 默认不认裸「工程名称」，是因为清单表抬头里它常写整个项目名）。
_FEE_UNIT_NAME_LABELS = _UNIT_PROJECT_LABELS + ("工程名称",)


def scan_measure_fee_row(title: str, rows: list) -> Optional[dict]:
    """单位工程费用表 → {sheet, unit, rows: [{row, amount}]}；不是单位工程层的费用表回 None。

    读法与 scan_fee_sheet 相同（K-R18：费用表只有这一种读法）：同样的分层判定、金额取表头定下的那一列、
    行名整格相等（「其中：…」不算）、金额格认带千分位的文本。行在、金额格空着时 amount 为 None。"""
    head = _title_text(rows, min(5, len(rows)))
    if not any(k in title or k in head for k in _FEE_SHEET_KW):
        return None
    header_row, amount_col, ambiguous = _fee_header(rows)
    if header_row is None or ambiguous:
        return None
    excluded, tier = _fee_tier(title, rows, header_row)
    if excluded or tier != FEE_UNIT:
        return None
    hits = []
    for ri in range(header_row + 1, len(rows)):
        row = rows[ri]
        for ci, v in enumerate(row):
            if isinstance(v, str) and _fee_label(v) == _FEE_MEASURE_LABEL:
                amount = _fee_row_amount(row, ci, amount_col)
                hits.append({"row": ri + 1, "amount": None if amount is None else round_money(amount)})
                break
    return {"sheet": title, "unit": _unit_project(rows, header_row, _FEE_UNIT_NAME_LABELS), "rows": hits}


def _fee_money(cents: int) -> str:
    return f"{cents / 100:,.2f}"


def _fee_layer_sum(scans: list, tier: str) -> tuple[Optional[int], int, bool]:
    """某一层各表合计之和 → (分, 张数, 这一层有没有表)。每张表恰 1 个白名单命中、各表行名相同、
    表名互不相同才可用；任何一张 0 个或 ≥2 个命中就整层不可用（少一张表的和不是和）。"""
    layer = [s for s in scans if s["tier"] == tier]
    if not layer:
        return None, 0, False
    picked = []
    for s in layer:
        white = [h for h in s["hits"] if not h["generic"]]
        if len(white) != 1:
            return None, len(layer), True
        picked.append(white[0])
    if len({h["label"] for h in picked}) != 1 or len({h["sheet"] for h in picked}) != len(picked):
        return None, len(layer), True
    return sum(to_cents(h["amount"]) for h in picked), len(layer), True


_FEE_LAYER_NAMES = {FEE_UNIT: "单位工程", FEE_SPECIALTY: "专业工程", FEE_SINGLE: "单项工程"}


def resolve_fee_total(scans: list) -> dict:
    """整份文件的费用表合计采用规则（WS26-h，spec §3.8 第 5 条；D-R13）。

    P = project 层的命中；U / S / Z = unit / single / specialty（专业工程）层各表合计之和。
    层级不混加：单项 ⊃ 单位 ⊃ 专业 是包含关系，相加会翻倍。对不上就不采用（铁律 3）。
    B（核对层）= 有单位工程表时是 U，一张单位工程表都没有时是 Z。

      · P 恰 1 个、B 可用：|P − B| ≤ 1 元采用 P，否则不采用并写出 P、B、差额；
      · P 恰 1 个、B 不可用：P 是泛名「合计」→ 不采用（泛名不单独成立）；
        P 是白名单行名且全文件白名单命中只有它 → 采用（与旧规则一致）；
      · 没有 P：全文件白名单命中恰 1 个、且单位 / 专业两层合起来（或单项层）只有它这一张表 → 采用它
        （与旧规则一致，表-13 金标）；
        否则 U 可用，且 S、Z 各自「算不出，或与 U 相差 ≤ 1 元」→ 采用 U；
        只有专业工程表（≥2 张）、没有单位工程表也没有 P → **不采用**：上一层的表不在，没有任何数可以
        核对这几张是不是全部（样本库里这种文件都是手工拼的进度款 / 报价工作簿）。原因里写出张数与和。

    S、Z 不作为取数来源，只核对、只写进提示行。

    返回 {amount, source, hit, has_fee_sheet, project_total, unit_sum, unit_count, single_sum,
    single_count, specialty_sum, specialty_count, diff, notes, reason}。amount 为 None 时 reason 写为什么没采用。"""
    hits = [h for s in scans for h in s["hits"]]
    white = [h for h in hits if not h["generic"]]
    proj = [h for h in hits if h["tier"] == FEE_PROJECT]
    u_cents, u_count, u_exists = _fee_layer_sum(scans, FEE_UNIT)
    s_cents, s_count, _ = _fee_layer_sum(scans, FEE_SINGLE)
    z_cents, z_count, z_exists = _fee_layer_sum(scans, FEE_SPECIALTY)
    out = {"amount": None, "source": None, "hit": None,
           # 「有费用表」= 认出了分层的造价汇总表，或有任何合计行命中。只有其他项目表 / 清单汇总表不算
           "has_fee_sheet": bool(hits) or any(s["tier"] for s in scans),
           "project_total": None, "unit_sum": None if u_cents is None else u_cents / 100,
           "unit_count": u_count, "single_sum": None if s_cents is None else s_cents / 100,
           "single_count": s_count, "specialty_sum": None if z_cents is None else z_cents / 100,
           "specialty_count": z_count, "diff": None, "notes": [], "reason": None}

    def adopt(amount: float, source: str, hit: Optional[dict], notes: list) -> dict:
        if not amount or amount <= 0:
            out["reason"] = "费用表合计为 0 或负数"
            return out
        out.update(amount=round_money(amount), source=source, hit=hit, notes=notes)
        return out

    def layer(tier: str, cents: int, count: int) -> str:
        return f"{count} 张{_FEE_LAYER_NAMES[tier]}费用表之和 {_fee_money(cents)}"

    check_only = "（只用于核对，不参与取数）"
    single_note = [layer(FEE_SINGLE, s_cents, s_count) + check_only] if s_cents is not None else []
    specialty_note = [layer(FEE_SPECIALTY, z_cents, z_count) + check_only] if z_cents is not None else []
    # 核对层：有单位工程表就只认单位层（哪怕它读不全——读不全是「不知道」，不拿下一层顶替）
    if u_exists:
        b_tier, b_cents, b_count, b_exists = FEE_UNIT, u_cents, u_count, True
    else:
        b_tier, b_cents, b_count, b_exists = FEE_SPECIALTY, z_cents, z_count, z_exists
    b_name = _FEE_LAYER_NAMES[b_tier]
    if len(proj) == 1:
        p = proj[0]
        p_cents = to_cents(p["amount"])
        out["project_total"] = p["amount"]
        if b_cents is not None:
            diff = p_cents - b_cents
            out["diff"] = diff / 100
            three = f"项目级合计 {_fee_money(p_cents)}，{layer(b_tier, b_cents, b_count)}，差 {_fee_money(abs(diff))} 元"
            if abs(diff) <= FEE_TOTAL_TOL_CENTS:
                return adopt(p["amount"], FEE_PROJECT, p,
                             [three] + single_note + (specialty_note if b_tier == FEE_UNIT else []))
            out["reason"] = three + "（超过 1 元，对不上）"
            return out
        if p["generic"]:
            out["reason"] = ("项目级费用表只有泛名「合计」行，"
                             + (f"{b_name}费用表的合计读不全" if b_exists else "没有单位工程费用表")
                             + "，无法核对")
            return out
        if len(white) == 1:
            return adopt(p["amount"], "single_hit", p, [])
        out["reason"] = f"有 {len(white)} 个合计行，{b_name}费用表的合计读不全，无法核对取哪一个"
        return out
    if len(proj) > 1:
        out["reason"] = f"项目级费用表里有 {len(proj)} 个合计行，不知道取哪一个"
        return out
    if len(white) == 1:
        h = white[0]
        # 「唯一命中」只在它所在那一层没有别的表时成立：两张单位工程表里只读到一张的合计，
        # 那是「另一张没读到」，不是「全文件只有这一个合计」。单位 / 专业两层合起来数——
        # 一张读到合计的专业工程表上面还有一张没读到合计的单位工程表，同样是没读全
        lower = u_count + z_count
        if {FEE_UNIT: lower, FEE_SPECIALTY: lower, FEE_SINGLE: s_count}.get(h["tier"], 1) > 1:
            out["reason"] = "同一层有多张费用表，只读到其中一张的合计"
            return out
        return adopt(h["amount"], "single_hit", h, [])
    if u_cents is not None:
        unit_note = layer(FEE_UNIT, u_cents, u_count)
        off = [(t, c, n) for t, c, n in ((FEE_SINGLE, s_cents, s_count), (FEE_SPECIALTY, z_cents, z_count))
               if c is not None and abs(c - u_cents) > FEE_TOTAL_TOL_CENTS]
        if not off:
            return adopt(u_cents / 100, "unit_sum", None, [unit_note] + single_note + specialty_note)
        out["reason"] = unit_note + "，" + "；".join(
            f"{layer(t, c, n)}，差 {_fee_money(abs(c - u_cents))} 元" for t, c, n in off) + "（超过 1 元，对不上）"
        return out
    if not u_exists and z_cents is not None and z_count > 1:
        out["reason"] = (f"只有 {layer(FEE_SPECIALTY, z_cents, z_count)}，没有单位工程费用表或项目级费用表可以核对，"
                         "不知道这几张是不是全部")
        return out
    if white:
        out["reason"] = f"有 {len(white)} 个合计行，不知道取哪一个"
    elif out["has_fee_sheet"]:
        out["reason"] = "费用表里没有读到合计行的金额"
    return out


# ---------------------------------------------------------------- 费用构成表（WS26-h2，07-005 / 07-010 / 07-012）

# 单位工程费用表（浙江表10.2.2-13）里构成「招标控制价合计」的六项。读法同 scan_fee_sheet（K-R18 / K-R44：
# 费用表只有一种读法）：行名整格相等（_fee_label），金额取「金额（元）」表头那一列。
# 「措施项目费」是技术 + 组织的小计，不进构成表（进了就重复计）；「其中：…」「专业工程暂估价」整格不等于任何一项。
# 「税金」是 GB 2013 费用表里同一项的名称，是第二个整格名，不是包含匹配（铁律 3）。
FEE_BREAKDOWN_ITEMS = (
    ("fbfx", "分部分项工程费", ("分部分项工程费",)),
    ("measure_tech", "施工技术措施项目", ("施工技术措施项目",)),
    ("measure_org", "施工组织措施项目", ("施工组织措施项目",)),
    ("other", "其他项目费", ("其他项目费",)),
    ("regulatory", "规费", ("规费",)),
    ("tax", "增值税", ("增值税", "税金")),
)
# 其他项目费的「其中」：只展示，不参与相加（已含在其他项目费里）。07-005 问的「真正的暂估价」就在这里
FEE_BREAKDOWN_SUBITEMS = (
    ("provisional", "暂估价", ("暂估价",)),
    ("contingency", "暂列金额", ("暂列金额",)),
)
_FEE_BD_LABEL_TO_KEY = {unicodedata.normalize("NFKC", lab): (key, name)
                        for key, name, labs in FEE_BREAKDOWN_ITEMS + FEE_BREAKDOWN_SUBITEMS for lab in labs}
_FEE_BD_MAIN = tuple(k for k, _n, _l in FEE_BREAKDOWN_ITEMS)
_FEE_BD_SUBS = tuple(k for k, _n, _l in FEE_BREAKDOWN_SUBITEMS)
# 措施项目费小计：不是构成项，只用来在「没拆技术 / 组织」的表上把原因说准（GB 2013 表-04 一类）
_FEE_BD_MEASURE_TOTAL = ("措施项目费", "措施项目")


def scan_fee_breakdown(title: str, rows: list, scan: dict) -> Optional[dict]:
    """单位工程费用表的分项行 → {sheet, total, items, subs, labels, problem}；不是 unit 层的表回 None。

    scan 是同一张表的 scan_fee_sheet 结果：层级与合计都从那里拿，不另读一遍（K-R44）。
    items / subs：{key: 分}，行在、金额格空着记 0（软件导出件「其他项目费」没有发生时就是空格）；行不在就没有这个键。
    problem：这张表不能采用的原因（某一项出现多行、没有「分部分项工程费」、合计读不到、六项相加与合计相差超过 1 元）。"""
    if scan.get("tier") != FEE_UNIT:
        return None
    out = {"sheet": title, "total": None, "items": {}, "subs": {}, "labels": {}, "problem": None}
    header_row, amount_col, ambiguous = _fee_header(rows)
    if ambiguous:
        out["problem"] = f"「{title}」有两组金额表头，不知道取哪一列"
        return out
    found: dict = {}
    measure_total = False
    extra: list = []        # 六项以外、序号是一级整数的行（「6 让利」「6 下浮」）：只用来把对不上的原因说准
    for ri in range(0 if header_row is None else header_row + 1, len(rows)):
        row = rows[ri]
        _note_extra_row(row, amount_col, extra)
        for ci, v in enumerate(row):
            if not isinstance(v, str):
                continue
            hit = _FEE_BD_LABEL_TO_KEY.get(_fee_label(v))
            if hit is None:
                measure_total = measure_total or _fee_label(v) in _FEE_BD_MEASURE_TOTAL
                continue
            amount = _fee_row_amount(row, ci, amount_col)
            found.setdefault(hit[0], []).append(0 if amount is None else to_cents(amount))
            out["labels"].setdefault(hit[0], _fee_label(v))
            break
    names = {k: n for k, n, _l in FEE_BREAKDOWN_ITEMS + FEE_BREAKDOWN_SUBITEMS}
    dup = [k for k, v in found.items() if len(v) > 1]
    white = [h for h in scan["hits"] if not h["generic"]]
    if dup:
        out["problem"] = "、".join(f"「{names[k]}」有 {len(found[k])} 行" for k in dup) + f"（「{title}」），不知道取哪一行"
    elif "fbfx" not in found:
        out["problem"] = f"「{title}」没有「分部分项工程费」行，不是能拆分的单位工程费用表"
    elif measure_total and "measure_tech" not in found and "measure_org" not in found:
        out["problem"] = f"「{title}」的措施项目费没有拆成「施工技术措施项目」「施工组织措施项目」两行，构成表拆不开"
    elif len(white) != 1:
        out["problem"] = f"「{title}」" + ("没有读到合计行的金额" if not white else f"有 {len(white)} 个合计行")
    for k, v in found.items():
        (out["items"] if k in _FEE_BD_MAIN else out["subs"])[k] = v[0]
    if out["problem"]:
        return out
    total = to_cents(white[0]["amount"])
    out["total"] = total
    got = sum(out["items"].values())
    if abs(got - total) > FEE_TOTAL_TOL_CENTS:
        out["problem"] = (f"「{title}」各项相加 {_fee_money(got)}，与{white[0]['label']} {_fee_money(total)} "
                          f"相差 {_fee_money(abs(got - total))} 元"
                          + ("；表里还有六项以外的" + "、".join(f"「{lab}」{_fee_money(c)}" for lab, c in extra)
                             + "，构成表只拆六项" if extra else ""))
    return out


def _note_extra_row(row: list, amount_col: Optional[int], extra: list) -> None:
    """序号格是一级整数（「6」）、行名不是六项里的任何一项、金额格有数 → 记下（行名, 分）。"""
    cells = [(ci, v) for ci, v in enumerate(row) if v is not None and str(v).strip()]
    if len(cells) < 2 or not re.fullmatch(r"\d+", str(cells[0][1]).strip()):
        return
    ci, v = cells[1]
    if not isinstance(v, str):
        return
    lab = _fee_label(v)
    if not lab or lab in _FEE_BD_LABEL_TO_KEY or lab in _FEE_BD_MEASURE_TOTAL:
        return
    amount = _fee_row_amount(row, ci, amount_col)
    if amount:
        extra.append((lab, to_cents(amount)))


def resolve_fee_breakdown(tables: list, fee_check: dict) -> dict:
    """整份文件的费用构成（WS26-h2，spec §3.17 第 5 条）。tables 是各 unit 层表的 scan_fee_breakdown 结果。

    采用条件（铁律 3：对不上就不显示，写原因）：
      1. 每一张单位工程费用表都能拆：六项各一行、相加与该表合计相差 ≤ 1 元——任一张不过，整份不出构成；
      2. 各表合计之和与本文件采用的费用表合计（resolve_fee_total）相差 ≤ 1 元——总造价没有取自这些表时
         （项目级合计对不上、没有采用），这张表不能说成总造价的构成。

    返回 {available, reason, unreconciled, unit_count, total, items, subs, labels}（金额是元；labels 是各项在表里
    实际写的行名，几种写法用「/」连；unreconciled 为 True 时原因是「表自己对不上」，调用方据此出
    fee_breakdown_unreconciled）。"""
    out = {"available": False, "reason": None, "unreconciled": False, "unit_count": len(tables),
           "total": None, "items": {}, "subs": {}, "labels": {}}
    if not tables:
        out["reason"] = "没有单位工程费用表"
        return out
    bad = [t for t in tables if t["problem"]]
    if bad:
        out["unreconciled"] = True
        out["reason"] = bad[0]["problem"] + (f"（{len(tables)} 张单位工程费用表里有 {len(bad)} 张拆不开）"
                                             if len(bad) > 1 else "")
        return out
    total = sum(t["total"] for t in tables)
    adopted = fee_check.get("amount") if fee_check else None
    if adopted is None:
        # 原因照抄采用规则的：手填了总造价时引擎不出 fee_total_unreadable，不能指人去看那条告警
        out["reason"] = f"费用表合计没有采用（{(fee_check or {}).get('reason') or '原因不明'}），构成表不显示"
        return out
    if abs(total - to_cents(adopted)) > FEE_TOTAL_TOL_CENTS:
        out["reason"] = (f"{len(tables)} 张单位工程费用表合计之和 {_fee_money(total)}，与采用的费用表合计 "
                         f"{_fee_money(to_cents(adopted))} 相差超过 1 元，构成表不显示")
        return out

    def add(keys: tuple, part: str) -> dict:
        present = [k for k in keys if any(k in t[part] for t in tables)]
        return {k: sum(t[part].get(k, 0) for t in tables) / 100 for k in present}

    labels: dict = {}
    for t in tables:
        for k, lab in t["labels"].items():
            if lab not in labels.setdefault(k, []):
                labels[k].append(lab)
    out.update(available=True, total=total / 100, items=add(_FEE_BD_MAIN, "items"), subs=add(_FEE_BD_SUBS, "subs"),
               labels={k: "/".join(v) for k, v in labels.items()})
    return out


def _code_like_cells(rows: list) -> int:
    """前 300 行里像 9–12 位清单编码的格子数。数值格只认 11 位（丢了前导 0 的 12 位编码）——
    8 位的金额补 0 后也像 9 位编码，不能算。"""
    n = 0
    for row in rows[:300]:
        for v in row:
            if isinstance(v, bool):
                continue
            if isinstance(v, str) or (isinstance(v, int) and len(str(v)) == 11):
                if cr.normalize_code(v).kind == cr.GB:
                    n += 1
    return n


def _looks_listy(title: str, rows: list) -> bool:
    return any(k in title for k in _LISTY_TITLE) or _code_like_cells(rows) >= 3


def _estimate_total(rows: list) -> Optional[float]:
    """跳过的表大约有多少钱：取最后一个「合计/总计」行里绝对值最大的数。估不出返回 None。"""
    for row in reversed(rows):
        texts = [_norm(v) for v in row if isinstance(v, str)]
        if any(("合计" in t or "总计" in t) and "本页" not in t for t in texts):
            nums = [abs(f) for f in (_to_float(v) for v in row if not isinstance(v, str)) if f]
            return max(nums) if nums else None
    return None


def _round(v: Optional[float], nd: int) -> Optional[float]:
    return None if v is None else round(v, nd)


def _row_values(row: list, cols: list) -> dict:
    """该行所有带角色的数值（不含分项列），键如 contract.amount / current.submit.amount。"""
    out: dict = {}
    for c in cols:
        if not c.measure or c.component:
            continue
        v = _to_float(_cell(row, c))
        if v is None:
            continue
        v *= c.scale
        v = round_money(v) if c.measure == "amount" else round(v, 6)
        k = c.key or c.measure
        if k in out:
            k = f"{k}[{c.text}]"
        out[k] = v
    return out


# ---------------------------------------------------------------- 口径指路与跳表分类（WS26-e）

def _basis_available(cols: list, side: Optional[str]) -> list[str]:
    """本表（本侧或不分当事方的）金额列出现过哪些期次；有单价列时工程量列的期次也算（量 × 价能出金额）。"""
    ms = [c for c in _measure_cols(cols) if c.party in (side, None)]
    src = [c for c in ms if c.measure == "amount"]
    if any(c.measure == "price" for c in ms):
        src += [c for c in ms if c.measure == "qty"]
    got = {c.period or "plain" for c in src}
    return [k for k in BASIS_KEYS if k in got]


def basis_columns_text(available) -> str:
    """["plain"] → 「不分期的合价列」；["contract", "current"] → 「合同价列、本期列」。"""
    return "、".join(_BASIS_COL_TEXT[k] for k in BASIS_KEYS if k in set(available or ()))


def basis_switch_text(requested: str, available) -> str:
    """改口径的方向。只有不分期列时只指「自动」（评审 S7：不分期 ≠ 合同价，虽然选合同价也读得出）；
    有合同价 / 本期 / 累计列时把它们列上。所选口径本身是「自动」时不再叫人选「自动」。"""
    have = set(available or ())
    choices = [] if requested == "auto" else ["自动"]
    choices += [BASIS_NAME[k] for k in ("contract", "current", "cumulative") if k in have and k != requested]
    if not choices:
        return "请确认选的是带金额列的计价表。"
    if len(choices) == 1:
        return f"请改用『{choices[0]}』。"
    return "可改用" + "或".join(f"『{c}』" for c in choices) + "。"


_AMBIGUOUS_TEXT = "表头存在多组金额列，无法确定口径，请选择对比口径"


def basis_rejection(requested: str, missing: list[dict], ambiguous: list[dict], *, lead: str,
                    ambiguous_tail: str = "") -> Optional[dict]:
    """一份可用的清单都没读出来、原因在口径上时，422 的 detail 对象（WS26-e，07-009 第 1、2 项）。

    missing / ambiguous：[{file, file_id, basis_available}] —— 表状态是 basis_missing / ambiguous 的文件。
    lead：「送审文件全部没有读出清单」「审定文件「x.xlsx」解析失败」这类开头；ambiguous_tail 接在「请选择对比口径」
    后面（投控是「：a、b」，对账是括号里的表级原因）。三种写法（spec §3.4 第 2 条）：
      全是 basis_missing：「所选口径『本期』在以下文件的表头里不存在：…。这些文件的金额列是不分期的合价列。请改用『自动』。」
      全是 ambiguous：保持原句（只把面向用户的 submission_files 换成「送审文件」）；
      混合：两句都写。
    → {message, reason_code, hint, basis_requested, basis_available, files}；不是口径问题返回 None。"""
    if not missing and not ambiguous:
        return None
    parts = []
    avail_all: list[str] = []
    if missing:
        avail = {k for f in missing for k in f.get("basis_available") or ()}
        avail_all = [k for k in BASIS_KEYS if k in avail]
        names = "、".join(f["file"] for f in missing)
        if requested == "auto":
            head = f"按『自动』口径（有本期列取本期，否则取合同价或不分期列）在以下文件里取不到金额列：{names}。"
        else:
            head = f"所选口径『{BASIS_NAME.get(requested, requested)}』在以下文件的表头里不存在：{names}。"
        desc = f"这些文件的金额列是{basis_columns_text(avail_all)}。" if avail_all else ""
        parts.append(head + desc + basis_switch_text(requested, avail_all))
    if ambiguous and missing:
        parts.insert(0, f"{lead}。")
        parts.append(f"另有文件{_AMBIGUOUS_TEXT}{ambiguous_tail}")
    elif ambiguous:
        parts.append(f"{lead}：{_AMBIGUOUS_TEXT}{ambiguous_tail}")
    files = [{"file": f["file"], "file_id": f.get("file_id"), "status": "basis_missing",
              "basis_available": list(f.get("basis_available") or [])} for f in missing]
    files += [{"file": f["file"], "file_id": f.get("file_id"), "status": "ambiguous",
               "basis_available": list(f.get("basis_available") or [])} for f in ambiguous]
    if missing:
        switch = basis_switch_text(requested, avail_all)
        hint = "勾选的文件不用改，只把对比口径换掉再开始。" if switch.startswith(("请改用", "可改用")) else switch
        # 前端只对 basis_not_found 给「改用自动口径重新选择」。选的本来就是「自动」时这个按钮没有用（再选一次自动
        # 还是读不出），换一个码，前端不给按钮，照 message 里指的口径（如『累计』）自己改
        code = "auto_basis_not_found" if requested == "auto" else "basis_not_found"
    else:
        hint = "表头有多组金额列，请在「对比口径」里明确选择合同价、本期或累计"
        code = "column_ambiguous"
    return {"message": "".join(parts), "reason_code": code,
            "hint": hint, "basis_requested": requested, "basis_available": avail_all, "files": files}


def _mark_title_skip(sh: _Sheet, title: str, rows: list, word: str, *, layout_ok: bool) -> None:
    """记下按哪个词跳过的（跳表告警按它计数），并认出「读不了的技术措施表」：表名或前 5 行标题含「措施」、
    命中的是综合单价计算 / 分析这一类词。它们带清单编码行、金额没算进去，不是「非清单表」（D-R87）。"""
    d = sh.detail
    d["skip_word"] = word
    if word not in _MEASURE_CALC_SKIP:
        return
    text = unicodedata.normalize("NFKC", _norm(title) + _title_text(rows, min(5, len(rows))))
    if "措施" not in text:
        return
    d["measure_sheet"] = True
    d["measure_skip"] = "basis" if layout_ok else "layout"
    d["has_codes"] = _code_like_cells(rows) > 0


def _parse_sheet(title: str, rows: list, *, side, basis, require_gb_code, file, file_id) -> _Sheet:
    detail = {"sheet": title, "status": None, "reason": "", "header_rows": [], "columns": {},
              "leaf_rows": 0, "priced_rows": 0, "amount_sum": 0.0, "total_row_amount": None,
              "total_ratio": None, "basis_used": None, "basis_label": None, "format": None,
              "unit_project": None, "measure_sheet": False,
              # 本表实际取金额的列下标（WS26-b）：没取数的表是空列表。对账靠它判断「同一份文件放两侧时
              # 两侧是不是取到了同一列」—— 比下标，不比 columns 里的表头文本（文本不保证唯一）
              "amount_col_idx": []}
    sh = _Sheet(title=title, detail=detail, summary_candidate="汇总" in title)

    def done(status: str, reason: str) -> _Sheet:
        detail["status"], detail["reason"] = status, reason
        return sh

    # WS26-p：措施综合单价计算表先认版式、只读清单项行；勾稽在整份文件读完后做（_settle_measure_calc）。
    # 表里只有不分期的数量与单价，调用方明确要本期 / 累计 / 上期时没有那个口径，照旧按标题跳过。
    lay = _measure_calc_layout(title, rows)
    if lay is not None and basis not in ("current", "cumulative", "previous"):
        return _measure_calc_rows(sh, rows, lay, require_gb_code=require_gb_code, file=file, file_id=file_id)
    hit = next((k for k in _SHEET_SKIP if k in title), None)
    if hit:
        _mark_title_skip(sh, title, rows, hit, layout_ok=lay is not None)
        return done("skipped_by_title", "表名属于非清单表（封面/说明/费用表/工日台班材料一览表等）")
    hb = _detect_header_block(rows)
    if hb is not None:
        detail["unit_project"] = _unit_project(rows, hb.rows[0])
    head_text = _title_text(rows, hb.rows[0] if hb else min(5, len(rows)))
    hit = next((k for k in _TITLE_SKIP_IN_ROWS if k in head_text), None)
    if hit:
        _mark_title_skip(sh, title, rows, hit, layout_ok=lay is not None)
        return done("skipped_by_title", "表内标题属于非清单表（工日/台班/材料/单价分析/费用汇总）")
    sh.listy = _looks_listy(title + head_text, rows)
    if hb is None:
        sh.skipped_estimate = _estimate_total(rows) if sh.listy else None
        return done("no_header", "前 30 行没有找到表头（需要 序号/编码/名称 中任两个）")
    detail["header_rows"] = [r + 1 for r in hb.rows]
    cols = hb.cols
    b = _bind(cols, side, basis, rows, hb.data_start)
    detail["basis_used"], detail["basis_label"] = b.basis_used, b.basis_label
    # 表头里有几组金额列：对账「两侧金额逐分相同」护栏据此判断是不是取成了同一列
    detail["amount_columns"] = sum(1 for c in cols if c.measure == "amount" and not c.component)
    detail["columns"] = {k: v.text for k, v in (("qty", b.qty), ("price", b.price), ("amount", b.amount),
                                                ("audit_amount", b.audit_amount)) if v is not None}
    detail["columns"].update({k: v.text for k, v in b.idents.items()})
    detail["basis_available"] = _basis_available(cols, b.side)
    money_cols = any(c.measure in ("amount", "price") for c in _measure_cols(cols))
    # 表头有金额列、只是没有本次要的口径（例如明确要本期而表里只有不分期的合价）。
    # 表头压根没有金额/单价列（只有工程量的清单框架）不算，按行解析后落 unpriced（WS08-04）。
    missing = b.amount is None and not (b.qty and b.price) and money_cols
    # 只有「单价」一项有歧义（本期没有单价列，退到合同 / 不分期单价时两列并存）而金额、工程量在所选期次下
    # 根本没有候选 —— 真正的原因是口径不存在，不是歧义（WS26-e：两件事指的方向相反）
    if b.ambiguous and not (missing and all(m == "price" for m, _ in b.ambiguous)):
        sh.skipped_estimate = _estimate_total(rows)
        paths = "；".join(f"{m}：{' | '.join(ps)}" for m, ps in b.ambiguous)
        sh.warnings.append(warn("column_ambiguous",
                                f"工作表「{title}」表头存在多组可取的列，无法确定口径（{paths}），该表未计入",
                                file=file, sheet=title))
        return done("ambiguous", f"列口径有歧义：{paths}")
    if not b.has_measure:
        sh.skipped_estimate = _estimate_total(rows) if sh.listy else None
        return done("no_header", "表头里没有工程量/单价/金额列")
    if missing:
        # WS26-e（07-009）：不拿别的列冒充，也不再说成「口径有歧义」—— 状态 basis_missing、码 basis_not_found，
        # 写出这张表实际有哪几种金额列，提示改口径的方向
        sh.skipped_estimate = _estimate_total(rows)
        want = "合同或不分期" if (b.want == "contract" and basis != "contract") else \
            _PERIOD_LABEL.get(b.want, b.want)
        avail = detail["basis_available"]
        reason = f"表头没有「{want}」口径的金额列" + (f"（只有{basis_columns_text(avail)}）" if avail else "")
        sh.warnings.append(warn("basis_not_found", f"工作表「{title}」{reason}，该表未计入",
                                file=file, sheet=title,
                                hint="这张表没有所选口径的金额列。" + basis_switch_text(basis, avail)))
        return done("basis_missing", reason)
    if b.party_missing:
        sh.warnings.append(warn("party_columns_missing",
                                f"工作表「{title}」表头分了送审/审核列，但没有「{_PARTY_LABEL[side]}」列，"
                                f"已按不带当事方的列（{(b.amount or b.qty).text}）取数", file=file, sheet=title))

    # 审核汇总版式：≥2 组当事方金额列，且没有编码/单位/量/价列（是「各专业一行」的汇总，不是清单）
    money_parties = {c.party for c in cols if c.measure == "amount" and not c.component and c.party}
    audit_shape = (len(money_parties) >= 2
                   and "code" not in b.idents and "unit" not in b.idents
                   and not any(c.measure in ("qty", "price") for c in cols))
    if audit_shape:
        return _audit_rows(sh, rows, hb, b, side=side, require_gb_code=require_gb_code,
                           file=file, file_id=file_id)
    # 只把已有的判断记进 detail（WS26-h）：指标回落到清单明细合计时，口径标签要按「实际读没读到
    # 措施项目清单表」生成，不能无论如何都写「分部分项+措施」。措施表怎么读不在这里改（WS26-p）。
    detail["measure_sheet"] = "措施" in title or "措施项目" in head_text
    return _detail_rows(sh, rows, hb, b, require_gb_code=require_gb_code, file=file, file_id=file_id,
                        measure_sheet=detail["measure_sheet"])


def _periods(b: _Binding) -> dict:
    """绑定列各自的期次：投控 self_check 只在工程量与金额同一期次组时才有意义（WS08-01）。"""
    return {m: (getattr(b, m).period or "plain") if getattr(b, m) is not None else None
            for m in ("qty", "price", "amount")}


def _item(sh: _Sheet, *, code_raw, name, feature, unit, qty, price, amount, values, basis_used,
          section, row, kind, file, file_id, periods=None, extra=None) -> dict:
    nc = cr.normalize_code(code_raw)
    ch = cr.classify_chapter(nc, section)
    it = {
        "code": _norm(code_raw), "code_norm": nc.display if nc.is_item_code else "",
        "name": name, "feature": feature[:500], "unit": unit,
        "qty": _round(qty, 6), "price": _round(price, 6), "amount": round_money(amount),
        "values": values, "basis_used": basis_used, "periods": periods or {},
        "chapter": ch.key, "chapter_name": ch.name, "chapter_source": ch.source,
        "appendix": ch.appendix, "appendix_name": ch.appendix_name,
        "section": section or "", "sheet": sh.title, "row": row, "row_kind": kind,
        # WS08-08/09：单位工程随行输出 —— 投控重复计取按它分组、对账无编码行按它择优配对，
        # 界面据它区分「5 行铝合金门窗各是哪栋楼」。表头上方没写就退回 sheet 名。
        "unit_project": sh.detail.get("unit_project") or sh.title,
        "file": file, "file_id": file_id,
    }
    if extra:
        it.update(extra)
    return it


def _provisional_col(cols: list, b: _Binding) -> Optional[_Col]:
    """「金额 / 其中 / 暂估价」分项列（WS26-n：投控暂估规则的第一档）。只认与绑定的合价列同期次、同当事方的那一列；
    没有或不唯一就不读（铁律 3）—— 分期表里「合同 · 其中暂估价」不能当成本期这一行的暂估价。"""
    if b.amount is None:
        return None
    cands = [c for c in cols if c.measure == "amount" and c.component and "暂估" in c.text
             and c.period == b.amount.period and c.party == b.amount.party]
    return cands[0] if len(cands) == 1 else None


def _detail_rows(sh: _Sheet, rows, hb, b: _Binding, *, require_gb_code, file, file_id, measure_sheet) -> _Sheet:
    d = sh.detail
    d["format"] = "detail"
    _warn_scaled(sh, (b.amount, b.price, b.audit_amount), file)
    prov_col = _provisional_col(hb.cols, b)
    section = None
    explicit_total = last_total = None
    fee_seen = False
    gb_rows = 0
    derived = False     # 有没有哪一行的金额是 量×价 补出来的
    for ri in range(hb.data_start, len(rows)):
        row = rows[ri]
        if not any(v is not None and str(v).strip() for v in row):
            continue
        seq = _txt(_cell(row, b.idents.get("seq")))
        code_raw = _cell(row, b.idents.get("code"))
        name = _txt(_cell(row, b.idents.get("name")))
        unit = _txt(_cell(row, b.idents.get("unit")))
        qty = _to_float(_cell(row, b.qty))
        price = _scaled(row, b.price)
        amount = _scaled(row, b.amount)
        if not name and not _norm(code_raw) and not seq:
            continue
        kind = _row_kind(seq, code_raw, name, unit, qty, price, amount, measure_sheet=measure_sheet)
        if kind == TITLE:
            # 编码列里的 4 位章码（「0101 | 土石方工程」分两格写）也要进 section，归章靠它
            section = " ".join(t for t in (seq, _txt(code_raw), name) if t)
            continue
        if kind == SUBTOTAL:
            label = "".join((seq, _norm(code_raw), _norm(name)))
            if any(k in label for k in _EXPLICIT_TOTAL):
                explicit_total = (ri, amount)
            elif any(k in label for k in ("合计", "总计")) and "本页" not in label:
                last_total = (ri, amount)
            if any(k in _core_label(label) for k in _FEE_SUBTOTAL):
                fee_seen = True
            continue
        if kind not in (ITEM, LUMP_SUM):
            continue
        if amount is None and qty is not None and price is not None:
            amount = qty * price     # 合价格为空（公式未缓存）时按 量×价 补，与旧版一致
            derived = True
        values = _row_values(row, hb.cols)
        if any(v for k, v in values.items() if k.endswith("amount") or k.endswith("price")
               or "amount[" in k or "price[" in k):
            sh.has_priced_values = True
        nc = cr.normalize_code(code_raw)
        if nc.kind == cr.GB:
            gb_rows += 1
        extra = {"audit_amount": round_money(_scaled(row, b.audit_amount))} if b.audit_amount else None
        prov = _scaled(row, prov_col)
        if prov is not None:
            # 空格不写 0：「这一列这行没填」和「填了 0」不是一回事
            extra = dict(extra or {}, provisional_amount=round_money(prov))
        sh.items.append(_item(sh, code_raw=code_raw, name=name, feature=_txt(_cell(row, b.idents.get("feature"))),
                              unit=unit, qty=qty, price=price, amount=amount, values=values,
                              basis_used=b.basis_used, section=section, row=ri + 1, kind=kind,
                              file=file, file_id=file_id, periods=_periods(b), extra=extra))
    # 绑定的金额列；有行按 量×价 补过金额的，量、价两列也是金额的来源（两侧合价列相同、
    # 工程量列不同时，补出来的金额并不相同）
    d["amount_col_idx"] = ([b.amount.idx] if b.amount is not None else []) + \
        ([b.qty.idx, b.price.idx] if derived else [])
    return _finish_sheet(sh, rows, total=explicit_total or (None if fee_seen else last_total),
                         require_gb_code=require_gb_code, gb_rows=gb_rows, file=file)


def _audit_rows(sh: _Sheet, rows, hb, b: _Binding, *, side, require_gb_code, file, file_id) -> _Sheet:
    """审核汇总版式：金额取本侧（side=None 保持旧语义：报审优先、缺则审核），
    审核/核减金额以 audit_amount/deduct_amount 附带；首个聚合行起按段截断。"""
    d = sh.detail
    d["format"] = "audit_summary"
    ms = _measure_cols(hb.cols)
    sub = _pick(ms, "amount", b.want, "submit")
    aud = _pick(ms, "amount", b.want, "audit")
    sub_col = sub[0] if len(sub) == 1 and sub[0].party == "submit" else None
    aud_col = aud[0] if len(aud) == 1 and aud[0].party == "audit" else None
    main = aud_col if side == "audit" else sub_col
    basis_used = ".".join(p for p in ((main.period or "plain") if main else b.want,
                                      main.party if main else None) if p)
    d["basis_used"] = basis_used
    d["basis_label"] = (_PERIOD_LABEL.get(main.period, main.period or "") + f"·{_PARTY_LABEL[main.party]}"
                        if main else _PERIOD_LABEL.get(b.want, b.want))
    d["columns"].update({k: v.text for k, v in (("submit_amount", sub_col), ("audit_amount", aud_col),
                                               ("deduct_amount", b.delta_amount)) if v is not None})
    # 本侧取的那一列；side=None 且没有报审列时逐行退到审核金额（见下面 amount 的取法）
    taken = main if main is not None else (aud_col if side is None else None)
    d["amount_col_idx"] = [taken.idx] if taken is not None else []
    _warn_scaled(sh, (main, aud_col, b.delta_amount), file)
    sh.warnings.append(warn("audit_summary_format",
                            f"工作表「{sh.title}」按审核汇总表（报审/审核/核减）口径解析，"
                            f"金额取{'审核' if side == 'audit' else '报审'}金额", file=file, sheet=sh.title))
    total = None
    for ri in range(hb.data_start, len(rows)):
        row = rows[ri]
        if not any(v is not None and str(v).strip() for v in row):
            continue
        name = _txt(_cell(row, b.idents.get("name")))
        n = _norm(name)
        if n and any(kw in n for kw in _AUDIT_STOP_KW):
            if any(k in n for k in ("合计", "总计", "之和", "总产值")):
                # 合计行与明细同一口径换算，否则万元表的勾稽拿未换算的 150 去比 1500000
                total = (ri, _scaled(row, main or sub_col or aud_col))
            break
        if not name or any(kw in n for kw in _SUBTOTAL_ANY):
            continue
        submit = _scaled(row, sub_col)
        audit = _scaled(row, aud_col)
        deduct = _scaled(row, b.delta_amount)
        if submit is None and audit is None:
            continue   # 无金额的名称/说明行
        if side == "audit":
            amount = audit
        elif side == "submit":
            amount = submit
        else:
            amount = submit if submit is not None else audit
        values = _row_values(row, hb.cols)
        sh.has_priced_values = sh.has_priced_values or any(values.values())
        sh.items.append(_item(sh, code_raw=_cell(row, b.idents.get("code")), name=name,
                              feature=_txt(_cell(row, b.idents.get("feature"))), unit="",
                              qty=None, price=None, amount=amount, values=values, basis_used=basis_used,
                              section=None, row=ri + 1, kind=ITEM, file=file, file_id=file_id,
                              extra={"audit_amount": round_money(audit), "deduct_amount": round_money(deduct)}))
    return _finish_sheet(sh, rows, total=total, require_gb_code=require_gb_code, gb_rows=0, file=file)


# ---------------------------------------------------------------- 施工技术措施项目综合单价计算表（WS26-p）

# 浙江表10.2.2-17：「清单序号 | 项目编码(定额编码) | 清单(定额)项目名称 | 计量单位 | 数量 | 综合单价(元)
# → 人工费…小计 | 合计(元)」。清单项行有清单序号，下面挂的定额子目行清单序号为空（编码写「补」「1-93 换」）。
# 照通用读法，「综合单价」下 6 个子列全是 price → ambiguous，一行都读不出，所以单开一条读法。
_MEASURE_CALC_SKIP = ("综合单价计算", "综合单价分析", "单价分析")
_MEASURE_CALC_SEQ = "清单序号"
_MEASURE_CALC_TOTAL = ("合计(元)", "合计")
_MEASURE_SECTION = "技术措施"
# 与同一单位工程费用表「施工技术措施项目」行的勾稽容差（spec §3.15）。两边都到分：原件 108 张有清单项的表
# 逐行按 ROUND_HALF_UP 取分后与费用表逐分相同；用二进制浮点 round() 取分会有 9 行少 1 分、单表最多差 0.02 元
MEASURE_FEE_TOL_CENTS = 100


def _measure_calc_layout(title: str, rows: list) -> Optional[dict]:
    """三条都满足才回列位置（spec §3.15）：表名或表内前 5 行含「措施」；命中的是「综合单价计算 / 综合单价分析 /
    单价分析」；表头有「清单序号」列，且「综合单价」下有「小计」子列。分部分项的综合单价分析表过不了第一条，
    照旧跳过 —— 它的单价已经在清单与计价表里计过一次，再读就是重复计。"""
    text = unicodedata.normalize("NFKC", _norm(title) + _title_text(rows, min(5, len(rows))))
    if "措施" not in text or not any(k in text for k in _MEASURE_CALC_SKIP):
        return None
    hb = _detect_header_block(rows)
    if hb is None:
        return None

    def only(pred) -> Optional[_Col]:
        cands = [c for c in hb.cols if pred(c)]
        return cands[0] if len(cands) == 1 else None

    def ident(name: str) -> Optional[_Col]:
        cands = [c for c in hb.cols if c.ident == name]
        return min(cands, key=lambda c: (c.ident_rank, c.idx)) if cands else None

    lay = {
        "hb": hb,
        "seq": only(lambda c: c.path[:1] == (_MEASURE_CALC_SEQ,)),
        "price": only(lambda c: len(c.path) >= 2 and c.path[-1] == "小计"
                      and any("综合单价" in p for p in c.path[:-1])),
        "qty": only(lambda c: c.measure == "qty" and not c.component),
        "total": only(lambda c: len(c.path) == 1 and c.path[0] in _MEASURE_CALC_TOTAL),
        "code": ident("code"), "name": ident("name"), "unit": ident("unit"),
    }
    if any(lay[k] is None for k in ("seq", "price", "qty", "code", "name")):
        return None
    return lay


def _seq_number(v) -> bool:
    if v is None or isinstance(v, bool):
        return False
    if isinstance(v, (int, float)):
        return True
    return bool(_NUM_SEQ_RE.match(unicodedata.normalize("NFKC", _norm(v))))


def _measure_calc_rows(sh: _Sheet, rows: list, lay: dict, *, require_gb_code, file, file_id) -> _Sheet:
    """只读「清单序号」是数字的行（清单项）；清单序号为空的定额子目行不读、不计 —— 它们也有单位、数量、单价，
    读进来就是重复计。金额 = 数量 × 综合单价小计，四舍五入到分；「合计(元)」列是取整到元的数，只用末行
    合计做没有费用表时的勾稽。状态先记 used，整份文件读完后由 _settle_measure_calc 决定计不计入。"""
    d = sh.detail
    hb = lay["hb"]
    seq_c, qty_c, price_c, total_c = lay["seq"], lay["qty"], lay["price"], lay["total"]
    d.update(format="measure_calc", measure_sheet=True, header_rows=[r + 1 for r in hb.rows],
             unit_project=_unit_project(rows, hb.rows[0]), basis_used="plain", basis_label=_PERIOD_LABEL[None],
             amount_columns=1, amount_col_idx=[qty_c.idx, price_c.idx], measure_check=None,
             columns={k: lay[k].text for k in ("seq", "code", "name", "unit", "qty", "price", "total") if lay[k]})
    section = None
    total = None
    gb_rows = 0
    periods = {"qty": "plain", "price": "plain", "amount": "plain"}
    for ri in range(hb.data_start, len(rows)):
        row = rows[ri]
        seq_v = _cell(row, seq_c)
        seq_t = _norm(seq_v)
        code_raw = _cell(row, lay["code"])
        name = _txt(_cell(row, lay["name"]))
        if seq_t and any(k in seq_t for k in ("合计", "总计")):
            if "本页" not in seq_t:
                total = _to_float(_cell(row, total_c))
            continue
        if not _seq_number(seq_v):
            # 清单序号为空：没有数量和单价的是分部标题（「011701 脚手架工程」「降排水费」），其余是定额子目
            if not seq_t and _cell(row, qty_c) is None and _cell(row, price_c) is None and (name or _norm(code_raw)):
                section = " ".join(t for t in (_txt(code_raw), name) if t)
            continue
        qty = _to_float(_cell(row, qty_c))
        price = _to_float(_cell(row, price_c))
        qd, pd = _to_decimal(_cell(row, qty_c)), _to_decimal(_cell(row, price_c))
        amount = float((qd * pd).quantize(_CENT, rounding=ROUND_HALF_UP)) if qd is not None and pd is not None \
            else None
        values = {k: v for k, v in (("qty", _round(qty, 6)), ("price", _round(price, 6)), ("amount", amount))
                  if v is not None}
        sh.has_priced_values = sh.has_priced_values or price is not None
        if cr.normalize_code(code_raw).kind == cr.GB:
            gb_rows += 1
        # section 带上表内的分部标题：补充码（01B004）只能靠它里面的章码（011702）归到措施项目那一章，
        # 与分部分项表同一口径（code_registry.classify_chapter）
        sh.items.append(_item(sh, code_raw=code_raw, name=name, feature="", unit=_txt(_cell(row, lay["unit"])),
                              qty=qty, price=price, amount=amount, values=values, basis_used="plain",
                              section=_MEASURE_SECTION + (f" {section}" if section else ""), row=ri + 1, kind=ITEM,
                              file=file, file_id=file_id, periods=periods, extra={"measure": True}))
    leaf = len(sh.items)
    d["total_row_amount"] = None if total is None else round_money(total)
    if leaf == 0:
        d["status"], d["reason"] = "measure_empty", "措施综合单价计算表里没有一个清单序号（只有分部标题或定额子目），不计入"
        return sh
    if require_gb_code and gb_rows / leaf < 0.5:
        sh.items = []
        d["status"], d["reason"] = "skipped_non_list", f"国标编码行 {gb_rows}/{leaf} 不到一半"
        return sh
    d.update(leaf_rows=leaf, priced_rows=sum(1 for it in sh.items if it["amount"]),
             amount_sum=round(sum(it["amount"] for it in sh.items if it["amount"]), 2),
             status="used", reason="")
    return sh


def _unit_key(name: Optional[str]) -> Optional[str]:
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", name)) if name else None


def _drop_measure_calc(sh: _Sheet, status: str, reason: str) -> None:
    sh.items = []
    sh.detail.update(status=status, reason=reason, leaf_rows=0, priced_rows=0, amount_sum=0.0, amount_col_idx=[])


def _settle_measure_calc(sheets: list, fee_rows: list, *, file) -> None:
    """措施综合单价计算表计不计入（铁律 3：勾稽通过才计入，任一条不过这张表整张不计入）。按顺序判：

      1. 同一单位工程已有被计入的「措施项目清单与计价表」→ 照旧跳过，不告警（两张表是同一批措施项，防重复计）；
      2. 工作簿里有这个单位工程的费用表「施工技术措施项目」行 → 清单项之和与它相差 ≤ 1 元；
      3. 没有对应的费用表行 → 与本表合计行相差 ≤ 0.5 × 行数 + 0.5 元（「合计(元)」逐行取整的最大误差）；
      2、3 不过 → status=measure_unreconciled，error 级 measure_sheet_unreconciled 写出两个数。
    一个清单序号都没有的表在读的时候已经是 measure_empty，不计入、不告警。"""
    calc = [s for s in sheets if s.detail["format"] == "measure_calc" and s.detail["status"] == "used"]
    if not calc:
        return
    lists = {_unit_key(s.detail["unit_project"]) for s in sheets
             if s.detail["status"] == "used" and s.detail["format"] == "detail" and s.detail["measure_sheet"]}
    by_unit: dict = {}
    for f in fee_rows:
        if f is not None and f["unit"]:
            by_unit.setdefault(_unit_key(f["unit"]), []).append(f)
    for s in calc:
        unit = _unit_key(s.detail["unit_project"])
        # 单位工程认不出（任一边）就不能证明两张表不是同一批措施项 → 按重复处理
        if lists and (unit is None or unit in lists or None in lists):
            s.detail["measure_check"] = {"against": "measure_list"}
            _drop_measure_calc(s, "skipped_by_title", "同一单位工程已有计入的措施项目清单与计价表，综合单价计算表不重复计入")
            continue
        n = len(s.items)
        got = sum(to_cents(it["amount"]) for it in s.items if it["amount"] is not None)
        fees = by_unit.get(unit, []) if unit else []
        refs = [(f, r) for f in fees for r in f["rows"] if r["amount"] is not None]
        check = {"against": None, "items": n, "items_sum": got / 100, "reference": None, "diff": None,
                 "tolerance": None, "fee_sheet": None, "fee_row": None}
        s.detail["measure_check"] = check
        if len(fees) > 1 or len(refs) > 1:
            msg = (f"工作表「{s.title}」的单位工程「{s.detail['unit_project']}」在工作簿里对应 {len(fees)} 张单位工程"
                   f"费用表、{len(refs)} 行「施工技术措施项目」，不知道和哪一行勾稽，这张表没有计入")
        else:
            if refs:
                f, r = refs[0]
                ref, tol = to_cents(r["amount"]), MEASURE_FEE_TOL_CENTS
                where = f"费用表「{f['sheet']}」第 {r['row']} 行「施工技术措施项目」"
                check.update(against="fee_table", fee_sheet=f["sheet"], fee_row=r["row"])
            elif s.detail["total_row_amount"] is not None:
                ref, tol = to_cents(s.detail["total_row_amount"]), 50 * n + 50
                where = "本表合计行"
                check["against"] = "sheet_total"
            else:
                ref = tol = None
                where = None
            if ref is None:
                msg = (f"工作表「{s.title}」没有对应的费用表「施工技术措施项目」行，本表也没有合计行，"
                       "清单项之和无从核对，这张表没有计入")
            else:
                check.update(reference=ref / 100, diff=(got - ref) / 100, tolerance=tol / 100)
                if abs(got - ref) <= tol:
                    continue
                msg = (f"工作表「{s.title}」技术措施清单项 {n} 行（数量 × 综合单价小计）之和 {_fee_money(got)}，"
                       f"与{where} {_fee_money(ref)} 相差 {_fee_money(abs(got - ref))} 元，超过 {_fee_money(tol)} 元，"
                       "这张表没有计入")
        _drop_measure_calc(s, "measure_unreconciled", msg)
        s.warnings.append(warn("measure_sheet_unreconciled", msg, file=file, sheet=s.title))


def _finish_sheet(sh: _Sheet, rows: list, *, total, require_gb_code: bool, gb_rows: int, file) -> _Sheet:
    d = sh.detail
    leaf = len(sh.items)
    priced = sum(1 for it in sh.items if it["amount"])
    amount_sum = round(sum(it["amount"] for it in sh.items if it["amount"]), 2)
    d.update(leaf_rows=leaf, priced_rows=priced, amount_sum=amount_sum)
    if leaf == 0:
        # 空模板不报警；表里明明有清单编码却一行没收下，才算「有清单的表被跳过」
        sh.listy = _code_like_cells(rows) >= 3
        sh.skipped_estimate = _estimate_total(rows) if sh.listy else None
        d["status"], d["reason"] = "no_header", "认出了表头，但没有一行是清单项"
        return sh
    if require_gb_code and gb_rows / leaf < 0.5:
        # WS08-15：指标只认国标清单；万元设备估算表、审核汇总表等不参与
        sh.items = []
        d["status"], d["reason"] = "skipped_non_list", f"国标编码行 {gb_rows}/{leaf} 不到一半"
        return sh
    if total is not None and total[1] is not None:
        tv = round(total[1], 2)
        d["total_row_amount"] = tv
        d["total_ratio"] = round(amount_sum / tv, 6) if tv else None
        if abs(amount_sum - tv) > _TOTAL_TOL * max(abs(tv), 1.0):
            sh.warnings.append(warn("sheet_total_mismatch",
                                    f"工作表「{sh.title}」明细行之和 {amount_sum:,.2f} 与第 {total[0] + 1} 行合计 "
                                    f"{tv:,.2f} 偏差超过 0.5%", file=file, sheet=sh.title))
    if priced == 0 and not sh.has_priced_values:
        d["status"], d["reason"] = "unpriced", "清单行只有工程量，没有单价和金额（疑似招标工程量清单）"
    else:
        d["status"], d["reason"] = "used", ""
    return sh


# ================================================================ 文件入口

def chapter_of(code: str, name: str = "", section: Optional[str] = None) -> tuple[str, str]:
    """清单编码 → (章键, 章名)。薄包装：口径与 code_registry.classify_chapter 完全一致（WS08-06）。"""
    ch = cr.classify_chapter(code, section)
    return ch.key, ch.name


def _read_rows(ws) -> tuple[list, bool]:
    rows: list = []
    for r in ws.iter_rows(values_only=True):
        rows.append(list(r) if r else [])
        if len(rows) >= _MAX_ROWS:
            return rows, True
    return rows, False


# 旧的字符串 warnings 通道（前端 tkControl.formatWarningNotice 把每一条都说成「未能纳入比对」）
# 只放 warn / error 级，外加两条旧版本来就有的 info；其余 info 只进 warning_items，
# 等 WS08-10 的分级展示接上再露出（否则「已补样式表名」也会被说成「未纳入」）。
_LEGACY_INFO = {"xls_converted", "audit_summary_format"}


def _legacy(wi: list) -> list[str]:
    return [w["message"] for w in wi if w["level"] != "info" or w["code"] in _LEGACY_INFO]


def _empty(filename: str, file_id, items_warnings: list) -> dict:
    meta = {"filename": filename, "file_id": file_id, "sheets": [], "sheets_detail": [],
            "parsed_total": 0, "item_count": 0, "leaf_rows": 0, "priced_rows": 0, "unpriced": False,
            "coverage": {"file": filename, "file_id": file_id, "status": "unparsed", "sheets": 0,
                         "used_sheets": [],
                         "amount_used": 0, "amount_skipped_estimate": None, "skipped_sheets": [],
                         "level": "poor", "ratio": None, "checked_rows": 0, "leaf_rows": 0,
                         "measure_used": 0, "measure_skipped": {"sheets": 0, "sheets_with_codes": 0,
                                                                "unreconciled": 0},
                         "scope_label": None},
            "basis_available": [],
            "chapter_coverage": None, "fee_totals": [], "fee_total": None, "fee_total_check": None,
            "fee_breakdown": None}
    return {"items": [], "meta": meta, "warnings": _legacy(items_warnings), "warning_items": items_warnings}


def parse_costsheet(raw: bytes, filename: str = "", *, side: Optional[str] = None,
                    basis: str = "auto", require_gb_code: bool = False,
                    file_id: Optional[str] = None) -> dict:
    """解析计价 Excel → {items, meta, warnings, warning_items}。

    side:  None / "submit" / "audit" —— 表头分了送审与审核列时取哪一侧（None 按送审）。
    basis: "auto" / "current" / "cumulative" / "contract" —— 期次口径（auto：有本期列取本期，
           否则取合同或不分期列）。
    require_gb_code: 国标编码行不到一半的表记 skipped_non_list（指标用）。

    items: [{code, code_norm, name, feature, unit, qty, price, amount, values, basis_used,
             chapter, chapter_name, chapter_source, appendix, appendix_name, section,
             sheet, row, row_kind, file, file_id, (audit_amount, deduct_amount)}]
           qty/price/amount 取不到时是 None，不写 0。
    meta:  {sheets（只列 used 的表）, sheets_detail, parsed_total, item_count, leaf_rows, priced_rows,
            unpriced, coverage, chapter_coverage, basis_used, format, filename, file_id}
           sheets_detail[].amount_col_idx：该表实际取金额的列下标（0 起；没取数的表是 []）。
           fee_totals / fee_total / fee_total_check：费用表合计的命中、采用值、采用依据（resolve_fee_total）。
           fee_breakdown：各单位工程费用表的分项之和（resolve_fee_breakdown，WS26-h2）。
    warnings: [str]（兼容旧调用方）；warning_items: [{code, level, message, file, sheet}]
    """
    wi: list[dict] = []
    ext = (filename.rsplit(".", 1)[-1].lower() if "." in filename else "")
    if ext and ext not in ("xlsx", "xlsm", "xls"):
        return _empty(filename, file_id, [warn("not_excel", f"非 Excel 文件:.{ext}", file=filename)])
    # zt 改动：.xls(旧二进制)openpyxl 不认。旧仓在这里起 soffice 转 xlsx；本机不装 LibreOffice，
    # 不转换，直接给 xls_unsupported 并提示另存为 xlsx（调用方据此回 status=insufficient）。
    if ext == "xls":
        return _empty(filename, file_id, [warn(
            "xls_unsupported", "旧版 .xls 不能直接解析，请用 Excel/WPS 另存为 .xlsx 后再试", file=filename)])
    try:
        wb, notes = load_workbook_tolerant(raw)
    except XlsxUnreadable as e:
        return _empty(filename, file_id, wi + [warn("xlsx_unreadable", e.message, file=filename)])
    wi.extend(warn(n["code"], n["message"], file=filename) for n in notes)

    sheets: list[_Sheet] = []
    fee_scans: list[dict] = []
    fee_parts: list[dict] = []
    measure_fee: list[Optional[dict]] = []
    try:
        for ws in wb.worksheets:
            rows, truncated = _read_rows(ws)
            if truncated:
                wi.append(warn("sheet_truncated", f"工作表「{ws.title}」超 {_MAX_ROWS} 行,已截断",
                               file=filename, sheet=ws.title))
            if not rows:
                continue
            fee_scan = scan_fee_sheet(str(ws.title), rows)
            fee_scans.append(fee_scan)
            fee_part = scan_fee_breakdown(str(ws.title), rows, fee_scan)
            if fee_part is not None:
                fee_parts.append(fee_part)
            measure_fee.append(scan_measure_fee_row(str(ws.title), rows))
            sheets.append(_parse_sheet(str(ws.title), rows, side=side, basis=basis,
                                       require_gb_code=require_gb_code, file=filename, file_id=file_id))
    finally:
        wb.close()
    # WS26-p：费用表在措施表前后都可能出现，整份文件读完才能逐单位工程勾稽
    _settle_measure_calc(sheets, measure_fee, file=filename)
    out = _assemble(sheets, wi, filename=filename, file_id=file_id, basis=basis)
    # WS08-16 / WS26-h：按整份文件分层勾稽后才采用（resolve_fee_total）；对不上、不知道取哪个都不猜
    check = resolve_fee_total(fee_scans)
    out["meta"]["fee_totals"] = [h for sc in fee_scans for h in sc["hits"]]
    out["meta"]["fee_total"] = check["amount"]
    out["meta"]["fee_total_check"] = check
    # WS26-h2：费用构成（各单位工程费用表的分项行之和），接在采用规则后面——总造价没取自这些表就不出
    out["meta"]["fee_breakdown"] = resolve_fee_breakdown(fee_parts, check)
    return out


# 跳过词 → 计数时给人看的说法（只改写不像话的那几个；其余就用跳过词本身，如「费用表 129」）
_SKIP_WORD_LABEL = {"封-": "封面"}
_SKIP_GROUPS_SHOWN = 3


def _measure_disclosure(sheets: list) -> dict:
    """技术措施表读了几张、哪些没读进来（WS26-e，D-R87 / D-R162）。没读进来 = 按标题跳过的措施类表
    （所选口径读不出 / 版式认不出）+ 勾稽不过的（measure_unreconciled，已有自己的 error 告警）。
    一个清单序号都没有的（measure_empty）与「同一单位工程已有措施清单与计价表」而不重复计的，都不算没读进来。"""
    not_read = [s for s in sheets if s.detail["status"] == "skipped_by_title" and s.detail.get("measure_skip")]
    unrec = [s for s in sheets if s.detail["status"] == "measure_unreconciled"]
    return {"used": sum(1 for s in sheets if s.detail["status"] == "used" and s.detail.get("measure_sheet")),
            "not_read": not_read, "unreconciled": len(unrec), "skipped": len(not_read) + len(unrec),
            "with_codes": sum(1 for s in not_read if s.detail.get("has_codes")) + len(unrec)}


def scope_label(coverages: list) -> Optional[str]:
    """结果的范围标签（C-R3 固定说法「仅分部分项清单」）。coverages 是计入了的文件的 file_coverage 项。
    一张技术措施表都没有（也没有跳过的）就不贴标签 —— 没有可披露的东西，不替用户下「只含分部分项」的结论。"""
    used = sum(int(c.get("measure_used") or 0) for c in coverages)
    skipped = sum(int((c.get("measure_skipped") or {}).get("sheets") or 0) for c in coverages)
    if skipped:
        head = "分部分项 + 技术措施清单" if used else "仅分部分项清单"
        return f"{head}（施工技术措施项目 {skipped} 张表未参与）"
    return "分部分项 + 技术措施清单" if used else None


def _title_skip_warnings(sheets: list, measure: dict, *, filename: str, basis: str) -> list[dict]:
    """按标题跳过的表 → 至多两条短句（WS26-e，07-004 第 1 项、07-003 披露）。表名进 sheets，不拼进 message：
    改前一句话顿号拼 242 个表名、7448 字，页面上高 943px，把异常表推到 1.6 屏以下。"""
    out = []
    not_read = measure["not_read"]
    if not_read:
        n = len(not_read)
        k = sum(1 for s in not_read if s.detail.get("has_codes"))
        kinds = {s.detail.get("measure_skip") for s in not_read}
        if kinds == {"basis"}:
            why = f"所选口径「{BASIS_NAME.get(basis, basis)}」读不出（这类表只有不分期的量价）"
        elif kinds == {"layout"}:
            why = "版式认不出（缺「清单序号」或「综合单价 / 小计」列）"
        else:
            why = "版式认不出或所选口径读不出"
        tail = f"技术措施只含读进来的 {measure['used']} 张表" if measure["used"] else "结果只含分部分项清单"
        codes = f"（其中 {k} 张有清单编码行）" if k else ""
        out.append(warn("measure_sheets_skipped", f"施工技术措施项目 {n} 张表没有参与{codes}：{why}；{tail}",
                        file=filename, sheets=[s.title for s in not_read],
                        hint="措施项目上的问题请人工核对"
                             + ("；口径改为『自动』可按不分期的量价读入这类表" if "basis" in kinds else "")))
    others = [s for s in sheets if s.detail["status"] == "skipped_by_title" and not s.detail.get("measure_skip")]
    if others:
        groups: dict[str, int] = {}
        for s in others:
            word = s.detail.get("skip_word")
            lab = _SKIP_WORD_LABEL.get(word, word) if word else (
                "措施项目清单已计入" if s.detail.get("format") == "measure_calc" else "其他")
            groups[lab] = groups.get(lab, 0) + 1
        parts = sorted(groups.items(), key=lambda kv: -kv[1])
        shown = "、".join(f"{lab} {cnt}" for lab, cnt in parts[:_SKIP_GROUPS_SHOWN])
        if len(parts) > _SKIP_GROUPS_SHOWN:
            shown += f" 等 {len(parts)} 类"
        out.append(warn("sheets_skipped_by_title", f"按标题跳过 {len(others)} 张表（{shown}）",
                        file=filename, sheets=[s.title for s in others]))
    return out


def _assemble(sheets: list, wi: list, *, filename: str, file_id, basis: str = "auto") -> dict:
    # —— 汇总表与明细表二选一（WS08-03）——
    has_detail = any(s.items and not s.summary_candidate and s.detail["status"] in ("used", "unpriced")
                     for s in sheets)
    summaries = [s for s in sheets if s.summary_candidate and s.items
                 and s.detail["status"] in ("used", "unpriced")]
    if summaries and has_detail:
        detail_sum = round(sum(s.detail["amount_sum"] for s in sheets
                               if not s.summary_candidate and s.detail["status"] == "used"), 2)
        detail_basis = {s.detail["basis_used"] for s in sheets
                        if not s.summary_candidate and s.detail["status"] == "used"}
        for s in summaries:
            s.detail["status"], s.detail["reason"] = "summary", "同簿有明细表，汇总表只用于勾稽、不计入"
            ssum = s.detail["amount_sum"]
            # 只在口径可比时勾稽：进度款汇总表比本期、明细表按合同价取时，两个数本来就不该相等
            comparable = s.detail["format"] == "detail" or detail_basis == {s.detail["basis_used"]}
            if comparable and detail_sum and abs(ssum - detail_sum) > _TOTAL_TOL * abs(detail_sum):
                wi.append(warn("summary_detail_mismatch",
                               f"汇总表「{s.title}」合计 {ssum:,.2f} 与明细表之和 {detail_sum:,.2f} 偏差超过 0.5%",
                               file=filename, sheet=s.title))
            s.items = []
            s.warnings = [w for w in s.warnings if w["code"] != "audit_summary_format"]
    elif summaries:
        for s in summaries:
            wi.append(warn("summary_only", f"工作簿里只有汇总表「{s.title}」，已按汇总表计入",
                           file=filename, sheet=s.title))
    if has_detail:
        # 汇总表本来就不计入；它的表头口径认不出（30003「土建造价汇总表」：金额 | 合同金额 并列）
        # 不影响结果，不该报成「有清单的表被跳过 / 口径歧义」
        for s in sheets:
            if s.summary_candidate and s.detail["status"] in ("ambiguous", "basis_missing", "no_header"):
                s.detail["status"] = "summary"
                s.detail["reason"] = f"同簿有明细表，汇总表不计入（表头未能识别，未勾稽：{s.detail['reason']}）"
                s.warnings = [w for w in s.warnings if w["code"] not in ("column_ambiguous", "basis_not_found")]
                s.listy = False

    items: list[dict] = []
    seen: set = set()
    dedup = 0
    for s in sheets:
        wi.extend(s.warnings)
        sheet_keys: set = set()
        kept = []
        for it in s.items:
            key = (it["code"], it["name"], it["amount"])
            if key in seen:
                dedup += 1
                continue       # 只挡"之前工作表"已出现的行;同表相同行如实保留(#22)
            sheet_keys.add(key)
            kept.append(it)
        seen |= sheet_keys
        items.extend(kept)
    if dedup:
        wi.append(warn("cross_sheet_dedup", f"不同工作表里编码、名称、金额完全相同的 {dedup} 行只计一次",
                       file=filename))

    measure = _measure_disclosure(sheets)
    wi.extend(_title_skip_warnings(sheets, measure, filename=filename, basis=basis))
    non_list = [s.title for s in sheets if s.detail["status"] == "skipped_non_list"]
    if non_list:
        wi.append(warn("sheets_skipped_non_list", "国标编码行不到一半、不参与计算的表：" + "、".join(non_list),
                       file=filename))
    skipped_listy = [s for s in sheets if s.listy and s.detail["status"] in ("no_header", "ambiguous", "basis_missing")]
    for s in skipped_listy:
        wi.append(warn("sheet_skipped_has_items",
                       f"工作表「{s.title}」含清单编码或标题像计价表，但未计入：{s.detail['reason']}",
                       file=filename, sheet=s.title))

    used = [s for s in sheets if s.detail["status"] == "used"]
    parsed_total = round(sum(i["amount"] for i in items if i["amount"]), 2)
    leaf_rows = len(items)
    priced_rows = sum(1 for i in items if i["amount"])
    unpriced = leaf_rows > 0 and not any(s.has_priced_values for s in sheets if s.items)
    coverage = _coverage(sheets, used, skipped_listy, parsed_total, leaf_rows, priced_rows, unpriced,
                         filename=filename, file_id=file_id)
    coverage["measure_used"] = measure["used"]
    coverage["measure_skipped"] = {"sheets": measure["skipped"], "sheets_with_codes": measure["with_codes"],
                                   "unreconciled": measure["unreconciled"]}
    coverage["scope_label"] = scope_label([coverage])
    if coverage["level"] == "poor" and items:
        wi.append(warn("coverage_poor", f"可识别清单表的金额计入不到一半（{coverage['ratio']}），结果不可作为结论",
                       file=filename))
    elif coverage["level"] == "partial":
        wi.append(warn("coverage_partial", f"部分清单表未计入（覆盖率 {coverage['ratio']}）", file=filename))
    if not items:
        # WS26-e（TK-G01）：改前是「未识别到清单明细(请确认是广联达分部分项/综合单价分析表(…))」括号套括号，
        # 还把「综合单价分析表」列成可选项 —— 那正是按标题跳过的表
        wi.append(warn("no_items", "没有读出清单行。请确认选的是分部分项清单与计价表（表头有项目编码、名称、合价），"
                                   "或审核汇总表（表头有报审、审核、核减金额）。", file=filename))
    # 章名覆盖率只在 meta 里给数；按章汇总的调用方（指标）在汇总层出 chapter_coverage_low
    ch_cov = cr.chapter_coverage(items)

    bases = {s.detail["basis_used"] for s in used}
    labels = {s.detail["basis_label"] for s in used}
    fmts = {s.detail["format"] for s in used}
    # WS26-e：这份文件的清单表里出现过哪些期次的金额列（汇总表不算）—— 口径选错时 422 据此指路
    avail = {k for s in sheets if s.detail["status"] != "summary" for k in s.detail.get("basis_available") or ()}
    meta = {
        "sheets": [s.title for s in used],
        "basis_available": [k for k in BASIS_KEYS if k in avail],
        "sheets_detail": [s.detail for s in sheets],
        "parsed_total": parsed_total, "item_count": leaf_rows,
        "leaf_rows": leaf_rows, "priced_rows": priced_rows, "unpriced": unpriced,
        "coverage": coverage, "chapter_coverage": ch_cov,
        "basis_used": (next(iter(bases)) if len(bases) == 1 else ("mixed" if bases else None)),
        "basis_label": (next(iter(labels)) if len(labels) == 1 else ("多种口径" if labels else None)),
        "filename": filename, "file_id": file_id,
    }
    if fmts:
        meta["format"] = ("audit_summary" if fmts == {"audit_summary"}
                          else "detail" if fmts == {"detail"} else "mixed")
    return {"items": items, "meta": meta, "warnings": _legacy(wi), "warning_items": wi}


def _coverage(sheets, used, skipped_listy, parsed_total, leaf_rows, priced_rows, unpriced,
              *, filename, file_id) -> dict:
    """文件级覆盖率（WS08-02 / D42）：可识别清单表的金额有多少计入了。"""
    if unpriced:
        status = "unpriced"
    elif priced_rows:
        status = "parsed"
    elif sheets and all(s.detail["status"] in ("skipped_by_title", "skipped_non_list") for s in sheets):
        status = "non_list"
    else:
        status = "unparsed"
    est = [s.skipped_estimate for s in skipped_listy]
    known = all(e is not None for e in est)
    amount_skipped = round(sum(e for e in est if e is not None), 2) if est and known else None
    # 无价表也是「认出来了」：它缺的是价，不是覆盖，由调用方的无价闸（file_unpriced / 422）处理
    recognized = [s for s in sheets if s.detail["status"] in ("used", "unpriced")]
    if not skipped_listy:
        level, ratio = ("ok", 1.0) if recognized else ("poor", None)
    elif known and (parsed_total + (amount_skipped or 0)) > 0:
        ratio = round(parsed_total / (parsed_total + amount_skipped), 4)
        level = "poor" if ratio < 0.5 else "partial"
    else:
        ratio = round(len(recognized) / (len(recognized) + len(skipped_listy)), 4)
        level = "poor" if ratio < 0.5 else "partial"
    # used_sheets（WS26-c）：本轮真的计入的表名，与 meta.sheets 同源。Java 清扫旧的待复核异常时
    # 按它判断「这张表这一轮看过没有」—— skipped_sheets 只含「像清单」的被跳表，拿它反推会漏
    # （上一轮用过、这一轮口径歧义且不够 3 个编码的表不在里面）。
    return {"file": filename, "file_id": file_id, "status": status, "sheets": len(used),
            "used_sheets": [s.title for s in used],
            "amount_used": parsed_total, "amount_skipped_estimate": amount_skipped,
            "skipped_sheets": [s.title for s in skipped_listy], "level": level, "ratio": ratio,
            "checked_rows": priced_rows, "leaf_rows": leaf_rows}


# ================================================================ 审核汇总表（带溯源）

def parse_audit_summary(raw: bytes, filename: str = "", *, file_id: Optional[str] = None) -> dict:
    """多行多列审核汇总表 → 带溯源的 Decimal 数据（总纲 K27：汇总表只在这里解析）。

    与 parse_costsheet 的审核汇总路径共用表头块与列角色，但目的不同：WS12 报告取数、
    WS07 审定单遇到汇总表时要的是**原样的每一格**，所以这里
      · 保留合计 / 累计 / 派生段的行（kind = total / derived），不截断丢弃；
      · 金额是 Decimal 字符串（float 走 repr，不带二进制尾差），万元列已 ×10000；
      · 每个值带单元格坐标（sheet!C5）与列路径，可回链到原格。
    只处理表头里有 ≥2 组送审/审核/核减金额列的表；其它表不出现在 sheets 里。

    → {file, file_id, sheets: [{sheet, header_rows, columns: [{col, index, path, key}],
        rows: [{row, seq, code, name, kind, values: {key: str}, cells: {key: "C5"}}],
        checks: [{key, items_sum, total, diff, total_row}]}],
       warnings, warning_items}
    """
    wi: list[dict] = []
    try:
        wb, notes = load_workbook_tolerant(raw)
    except XlsxUnreadable as e:
        w = warn("xlsx_unreadable", e.message, file=filename)
        return {"file": filename, "file_id": file_id, "sheets": [], "warnings": [w["message"]],
                "warning_items": [w]}
    wi.extend(warn(n["code"], n["message"], file=filename) for n in notes)
    out_sheets = []
    try:
        for ws in wb.worksheets:
            rows, _trunc = _read_rows(ws)
            hb = _detect_header_block(rows) if rows else None
            if hb is None:
                continue
            money = [c for c in hb.cols if c.measure == "amount" and not c.component and c.party]
            if len({c.party for c in money}) < 2:
                continue
            idents = {}
            for ident in ("seq", "code", "name"):
                cands = [c for c in hb.cols if c.ident == ident]
                if cands:
                    idents[ident] = min(cands, key=lambda c: (c.ident_rank, c.idx))
            vcols = [c for c in hb.cols if c.measure and not c.component]
            keys = {}
            for c in vcols:
                k = c.key
                keys[c.idx] = k if k not in keys.values() else f"{k}[{c.text}]"
            out_rows, in_derived = [], False
            for ri in range(hb.data_start, len(rows)):
                row = rows[ri]
                name = _txt(_cell(row, idents.get("name")))
                seq = _txt(_cell(row, idents.get("seq")))
                vals, cells = {}, {}
                for c in vcols:
                    dv = _to_decimal(_cell(row, c))
                    if dv is None:
                        continue
                    vals[keys[c.idx]] = str(dv * c.scale if c.scale != 1 else dv)
                    cells[keys[c.idx]] = f"{get_column_letter(c.idx + 1)}{ri + 1}"
                if not name and not vals:
                    continue
                n = _norm(name)
                if any(k in n for k in ("合计", "总计", "之和", "总产值")):
                    kind = "total"
                    in_derived = True
                elif in_derived or (n and any(kw in n for kw in _AUDIT_STOP_KW)):
                    kind, in_derived = "derived", True
                elif any(k in n for k in _SUBTOTAL_ANY):
                    kind = "subtotal"
                else:
                    kind = "item"
                out_rows.append({"row": ri + 1, "seq": seq, "code": _norm(_cell(row, idents.get("code"))),
                                 "name": name, "kind": kind, "values": vals, "cells": cells})
            checks = []
            first_total = next((r for r in out_rows if r["kind"] == "total"), None)
            if first_total:
                for k in sorted({k for r in out_rows for k in r["values"]}):
                    if k not in first_total["values"]:
                        continue
                    s = sum((Decimal(r["values"][k]) for r in out_rows
                             if r["kind"] == "item" and k in r["values"]), Decimal(0))
                    t = Decimal(first_total["values"][k])
                    checks.append({"key": k, "items_sum": str(s), "total": str(t), "diff": str(s - t),
                                   "total_row": first_total["row"]})
            out_sheets.append({
                "sheet": ws.title, "header_rows": [r + 1 for r in hb.rows],
                "columns": [{"col": get_column_letter(c.idx + 1), "index": c.idx, "path": c.text,
                             "key": keys[c.idx]} for c in vcols],
                "rows": out_rows, "checks": checks,
            })
    finally:
        wb.close()
    return {"file": filename, "file_id": file_id, "sheets": out_sheets,
            "warnings": _legacy(wi), "warning_items": wi}
