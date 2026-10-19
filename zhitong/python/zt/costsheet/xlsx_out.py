"""把核对、指标、投控的结果写成 xlsx 工作底稿（zt 新写；旧仓由 Java 侧导出）。

只做排版：每张表一个工作表，首行表头冻结，列宽按内容估算。数值原样写入单元格（不经过字符串），
金额列用两位小数格式。写入采用「同目录临时文件 + 替换」，中途失败不留下半个文件。
"""
from __future__ import annotations

import os
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Iterable, Optional, Union

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MONEY_FORMAT = "#,##0.00"
RATIO_FORMAT = "0.00%"
_HEADER_FILL = PatternFill("solid", fgColor="DDE7F3")
_MAX_WIDTH = 60


@dataclass(frozen=True)
class Column:
    """一列：表头、取值方式与数字格式。

    :ivar header: 表头文字。
    :ivar key: 行字典的键，或从行字典取值的函数。
    :ivar number_format: 单元格数字格式；None 为常规。
    """

    header: str
    key: Union[str, Callable[[dict], Any]]
    number_format: Optional[str] = None

    def value(self, row: dict) -> Any:
        """取这一列在某行的值；列表与字典写成文本。

        :param row: 行字典。
        :returns: 单元格值。
        """
        v = self.key(row) if callable(self.key) else row.get(self.key)
        if isinstance(v, (list, tuple)):
            return "；".join(str(x) for x in v) or None
        if isinstance(v, dict):
            return "；".join(f"{k}={x}" for k, x in v.items()) or None
        return v


@dataclass(frozen=True)
class Sheet:
    """一个工作表。

    :ivar title: 表名（≤31 字符，Excel 限制）。
    :ivar columns: 列定义。
    :ivar rows: 行字典。
    """

    title: str
    columns: list[Column]
    rows: list[dict]


def _width(values: Iterable[Any]) -> int:
    best = 6
    for v in values:
        if v is None:
            continue
        text = str(v)
        # 中文按两个字符宽估算
        w = sum(2 if ord(ch) > 0x2E80 else 1 for ch in text)
        best = max(best, w)
    return min(best + 2, _MAX_WIDTH)


def write_workbook(path: Path, sheets: list[Sheet]) -> Path:
    """写出 xlsx（原子替换）。

    :param path: 目标文件；父目录不存在时创建。
    :param sheets: 工作表，至少一个。
    :returns: 写入的路径。
    """
    wb = Workbook()
    wb.remove(wb.active)
    for sh in sheets:
        ws = wb.create_sheet(title=sh.title[:31])
        ws.append([c.header for c in sh.columns])
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = _HEADER_FILL
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row in sh.rows:
            ws.append([c.value(row) for c in sh.columns])
        for idx, col in enumerate(sh.columns, start=1):
            letter = get_column_letter(idx)
            if col.number_format:
                for (cell,) in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                    cell.number_format = col.number_format
            ws.column_dimensions[letter].width = _width(
                [col.header] + [c.value(r) for c in (col,) for r in sh.rows[:200]])
        ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=path.parent)
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    except BaseException:
        Path(tmp).unlink(missing_ok=True)
        raise
    return path


def kv_rows(pairs: Iterable[tuple[str, Any]]) -> list[dict]:
    """把 ``(名称, 值)`` 对转成两列表的行。

    :param pairs: 名称与值。
    :returns: ``[{"k": 名称, "v": 值}]``。
    """
    return [{"k": k, "v": v} for k, v in pairs]


KV_COLUMNS = [Column("项目", "k"), Column("值", "v")]
