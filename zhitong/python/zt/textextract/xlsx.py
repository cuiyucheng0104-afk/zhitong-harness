"""xlsx 抽取：保留列位置（列号坐标式）、跳过隐藏工作表并回报。

来源：/home/user/zhitong-cost-ai/engine-service/app/engine/textextract.py（xlsx 部分，旧 WS04-P6 / WS22-l）。
改动：只改包名与模块引用；``_xlsx_doc`` 保留在此，供统一出口调用。
"""
from __future__ import annotations

import io
import logging
import math
import re
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from decimal import Decimal
from typing import Optional

from zt.textextract.model import DocText, exc_detail as _exc_detail, failed_doc

log = logging.getLogger("zt.textextract.xlsx")

class XlsxReadError(RuntimeError):
    pass


@dataclass
class XlsxSheet:
    name: str
    state: str = "visible"  # visible / hidden / veryHidden
    rows: list[tuple[int, list[tuple[str, str]]]] = field(default_factory=list)


@dataclass
class XlsxBook:
    sheets: list[XlsxSheet]
    flags: list[str] = field(default_factory=list)

    @property
    def hidden_sheets(self) -> list[str]:
        return [s.name for s in self.sheets if s.state != "visible"]


def _col_letter(n: int) -> str:
    s = ""
    while n > 0:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


# 二进制累加误差的尾巴：15 位有效数字用满了，末尾是「一长串 0 或 9 + 1–3 位」（1737.66000000007、25.3099999999977）
_FLOAT_NOISE_TAIL = re.compile(r"^(\d*?)(?:0{6,}|9{6,})\d{1,3}$")


def _float_str(v: float) -> str:
    """浮点格写成 Excel 里看到的样子（整改 WS22-l，E31 第 9 点）。

    Excel 把公式结果按二进制浮点存，XML 里是 17 位的 `179159.09499999997`，界面上显示 179159.095；原样
    `str()` 进正文，模型照抄，候选就成了一长串小数（生产 E13）。规则：
      1. 15 位有效数字（Excel 自己的精度）—— 15 位以内的数一位都不丢；
      2. 15 位用满、末尾是一长串 0 / 9 再跟 1–3 位的，是求和累下的误差，舍到那串 0 / 9 之前
         （本机原件「X_面积统计表.xlsx」里有 1737.66000000007、25.3099999999977 这类格，只按第 1 条它们照样拖尾）；
      3. 整数值不带「.0」。"""
    if not math.isfinite(v):
        return str(v)
    if v.is_integer() and abs(v) < 1e15:
        return str(int(v))  # 含 0.0 / -0.0
    s = format(v, ".15g")
    digits = format(abs(v), ".14e").split("e")[0].replace(".", "").rstrip("0")
    if len(digits) >= 13:  # 误差只会出现在第 13–15 位；2.0000005 这种本来就短的数不碰
        m = _FLOAT_NOISE_TAIL.match(digits)
        if m:
            s = format(v, f".{max(1, len(m.group(1)))}g")
    if "e" in s and "e" not in repr(v):
        # 1e15–1e16 这一段（以及舍入后进位成整十整百的）.15g / .Ng 会写成科学计数而 repr 仍是定点：照定点写
        s = format(Decimal(s), "f")
    return s


def _cell_str(v) -> str:
    if v is None:
        return ""
    s = _float_str(v) if isinstance(v, float) else str(v)
    # 格内换行换成空格：生产 R04 清单第 1 条被格内换行拆成 5 行，编码和数量行分家
    return " ".join(s.replace("\r", "\n").split("\n")).strip()


def _raw_number(raw: str) -> "str | float":
    """zipfile 兜底读到的数值格（t 缺省或 n）：与 openpyxl 同口径 —— 有小数点或指数的才是浮点，交给 `_cell_str`
    按同一规则写；整数串原样（不经 int()，免得改掉 XML 里的写法）。两条读法得到同一段正文。"""
    if not any(ch in raw for ch in ".eE"):
        return raw
    try:
        return float(raw)
    except ValueError:
        return raw


def read_xlsx(content: bytes) -> XlsxBook:
    """结构化读 xlsx：每张表的 (行号, [(列字母, 值)])，只留非空格。失败抛 XlsxReadError。

    openpyxl 抛错（如 styles.xml 损坏）时退到 zipfile 直读 sharedStrings + sheet XML，
    打 `xlsx_raw_fallback`（WS04-P3）。只读模式拿不到合并区域与打印区域（本期不做）。
    """
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        log.warning("openpyxl 未装，xlsx 走 zipfile 兜底：%s", e)
        book = _read_xlsx_raw(content)
        book.flags.append("xlsx_raw_fallback")
        return book
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            sheets: list[XlsxSheet] = []
            for ws in wb.worksheets:
                sheet = XlsxSheet(name=ws.title, state=getattr(ws, "sheet_state", "visible") or "visible")
                if hasattr(ws, "reset_dimensions"):
                    # 只读模式信 <dimension> 标签；生成器写错它时会静默截掉列/行
                    ws.reset_dimensions()
                for rno, row in enumerate(ws.iter_rows(values_only=True), start=1):
                    cells = [(_col_letter(i), _cell_str(v)) for i, v in enumerate(row, start=1)]
                    cells = [(c, v) for c, v in cells if v]
                    if cells:
                        sheet.rows.append((rno, cells))
                sheets.append(sheet)
        finally:
            wb.close()
        return XlsxBook(sheets)
    except Exception as e:  # noqa: BLE001
        first = e
    try:
        book = _read_xlsx_raw(content)
    except Exception as e2:  # noqa: BLE001
        raise XlsxReadError(f"xlsx 解析失败：{_exc_detail(first)}；zipfile 兜底也失败：{_exc_detail(e2)}")
    log.warning("openpyxl 读 xlsx 失败，已走 zipfile 兜底：%s", _exc_detail(first))
    book.flags.append("xlsx_raw_fallback")
    return book


_NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_NS_PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
_CELL_REF_RE = re.compile(r"^([A-Z]+)(\d+)$")


def safe_xml(z: zipfile.ZipFile, name: str) -> Optional[ET.Element]:
    try:
        raw = z.read(name)
    except KeyError:
        return None
    if re.search(rb"<!DOCTYPE|<!ENTITY", raw, re.IGNORECASE):
        raise XlsxReadError(f"{name} 带 DOCTYPE/ENTITY，拒绝解析")
    return ET.fromstring(raw)


def _read_xlsx_raw(content: bytes) -> XlsxBook:
    m = f"{{{_NS_MAIN}}}"
    with zipfile.ZipFile(io.BytesIO(content)) as z:
        wbx = safe_xml(z, "xl/workbook.xml")
        if wbx is None:
            raise XlsxReadError("缺少 xl/workbook.xml")
        rels = safe_xml(z, "xl/_rels/workbook.xml.rels")
        targets: dict[str, str] = {}
        if rels is not None:
            for r in rels.iter(f"{{{_NS_PKG_REL}}}Relationship"):
                t = r.get("Target") or ""
                t = t.lstrip("/")
                targets[r.get("Id") or ""] = t if t.startswith("xl/") else f"xl/{t}"
        shared: list[str] = []
        sst = safe_xml(z, "xl/sharedStrings.xml")
        if sst is not None:
            for si in sst.iter(f"{m}si"):
                # 拼音注音 rPh 不是正文
                phon = {id(t) for rph in si.iter(f"{m}rPh") for t in rph.iter(f"{m}t")}
                parts = [t.text or "" for t in si.iter(f"{m}t") if id(t) not in phon]
                shared.append("".join(parts))
        sheets: list[XlsxSheet] = []
        for sh in wbx.iter(f"{m}sheet"):
            name = sh.get("name") or "?"
            state = sh.get("state") or "visible"
            rid = sh.get(f"{{{_NS_REL}}}id") or ""
            target = targets.get(rid)
            sheet = XlsxSheet(name=name, state=state)
            root = safe_xml(z, target) if target else None
            if root is not None:
                rows: dict[int, list[tuple[int, str]]] = {}
                for c in root.iter(f"{m}c"):
                    ref = c.get("r") or ""
                    mm = _CELL_REF_RE.match(ref)
                    if not mm:
                        continue
                    col = _col_index(mm.group(1))
                    rno = int(mm.group(2))
                    t = c.get("t")
                    if t == "inlineStr":
                        val = "".join(x.text or "" for x in c.iter(f"{m}t"))
                    else:
                        v = c.find(f"{m}v")
                        raw = v.text if v is not None else None
                        if raw is None:
                            continue
                        if t == "s":
                            try:
                                val = shared[int(raw)]
                            except (ValueError, IndexError):
                                val = ""
                        elif t == "b":
                            val = "TRUE" if raw == "1" else "FALSE"
                        elif t in (None, "n"):
                            val = _raw_number(raw)
                        else:
                            val = raw
                    val = _cell_str(val)
                    if val:
                        rows.setdefault(rno, []).append((col, val))
                for rno in sorted(rows):
                    cells = sorted(rows[rno])
                    sheet.rows.append((rno, [(_col_letter(ci), v) for ci, v in cells]))
            sheets.append(sheet)
    return XlsxBook(sheets)


def _col_index(letters: str) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + (ord(ch) - 64)
    return n


def xlsx_lines(content_or_book: "bytes | XlsxBook") -> list[str]:
    """xlsx 渲染成行（filetree 复用，WS04-P7）。决策 D11：列号坐标式。

        [工作表] 表名                 ← 分隔行保持不变：field_extract._SHEET_MARK 靠它分片
        [R8] A:土建<TAB>G:14253902<TAB>H:0

    只输出非空格，每格带列字母：生产 R16 第 8 行 D 列（审定）为空，旧格式把空格塌掉后
    「审定 / 核减 / 核增」三列语义整体左移、读反（铁律 2、3）。隐藏工作表跳过（C6），
    名单见 XlsxBook.hidden_sheets。
    """
    book = content_or_book if isinstance(content_or_book, XlsxBook) else read_xlsx(content_or_book)
    lines: list[str] = []
    for sheet in book.sheets:
        if sheet.state != "visible":
            continue
        lines.append(f"[工作表] {sheet.name}")
        for rno, cells in sheet.rows:
            lines.append(f"[R{rno}] " + "\t".join(f"{c}:{v}" for c, v in cells))
    return lines


def xlsx_text(content: bytes) -> str:
    """兼容出口：抽不出返回空串并告警（失败原因走 extract_document）。"""
    try:
        return "\n".join(xlsx_lines(content)).strip()
    except Exception as e:  # noqa: BLE001
        log.warning("xlsx 抽取失败:%s", e)
        return ""


def _xlsx_doc(content: bytes) -> DocText:
    try:
        book = read_xlsx(content)
    except Exception as e:  # noqa: BLE001
        return failed_doc("parse_failed", str(e))
    text = "\n".join(xlsx_lines(book)).strip()
    visible_rows = sum(len(s.rows) for s in book.sheets if s.state == "visible")
    if not visible_rows:
        detail = "xlsx 可见工作表里没有任何非空单元格"
        if book.hidden_sheets:
            detail += f"（已跳过隐藏工作表：{'、'.join(book.hidden_sheets)}）"
        return failed_doc("no_text", detail, hidden_sheets=book.hidden_sheets,
                          quality_flags=list(book.flags))
    return DocText(text=text, text_source="text_layer", hidden_sheets=book.hidden_sheets,
                   quality_flags=list(book.flags))
