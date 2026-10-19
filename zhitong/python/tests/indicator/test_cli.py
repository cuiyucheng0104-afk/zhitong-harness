"""python -m zt.indicator.cli run：面积、总造价、对标输入与工作底稿。"""
from __future__ import annotations

import json

from openpyxl import load_workbook

from tests.costsheet.cli_helpers import run_cli
from tests.costsheet.fixtures.golden import FIXTURES


def test_run_with_area_and_baseline_file(tmp_path):
    baseline = tmp_path / "对标.json"
    baseline.write_text(json.dumps({"indicator_baselines": {"unit_cost": 25.0}}), encoding="utf-8")
    out = tmp_path / "指标"
    code, summary = run_cli("zt.indicator.cli", "run", "--files", str(FIXTURES / "指标-计价表.xlsx"),
                            "--gross-area", "5000", "--area-source", "项目说明", "--baseline", str(baseline),
                            "--out", f"{out}/")
    assert code == 0 and summary["status"] == "ok"
    data = json.loads((out / "result.json").read_text(encoding="utf-8"))["data"]
    by = {r["indicator_code"]: r for r in data["indicators"]}
    assert by["unit_cost"]["baseline_value"] == 25.0 and by["unit_cost"]["needs_manual"] is True
    assert data["summary"]["area_source"] == "项目说明"
    wb = load_workbook(out / "指标结果.xlsx", read_only=True)
    assert wb.sheetnames == ["汇总", "指标", "按章", "造价驱动项", "未计入文件"]


def test_total_amount_flag_takes_priority(tmp_path):
    code, env = run_cli("zt.indicator.cli", "run", "--files", str(FIXTURES / "指标-计价表.xlsx"),
                        "--gross-area", "1000", "--total-amount", "500000")
    assert code == 0 and env["data"]["summary"]["cost_basis"] == "user_input"
    assert env["data"]["summary"]["unit_cost"] == 500.0


def test_no_files_is_a_usage_error():
    code, out = run_cli("zt.indicator.cli", "run", "--gross-area", "1000")
    assert code == 2 and out["error"]["code"] == "invalid_request"
