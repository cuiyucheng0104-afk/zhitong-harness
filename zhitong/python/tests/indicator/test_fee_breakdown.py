# 来源：engine-service/tests/test_fee_breakdown.py（zhitong-cost-ai@991ff5cd）
"""费用构成表（WS26-h2，issue #1006；spec WS26 §3.17「指标抽屉」第 5 条、§4 WS26-h2；总纲 K-R44）。

现场（IND-05，07-005 / 07-012）：总造价 318,375,872.00 取自费用表，但抽屉里看不到它由什么构成——
清单明细合计只有它的 58.2%（p 上线后 70.5%），「少的四成是什么」「真正的暂估价 5600 万在哪」无处可答。
113 张单位工程费用表本地回放：分部分项 185,405,919.90、技术措施 39,071,850.34、组织措施 6,354,845.40、
其他项目费 56,000,000.00（其中暂估价 56,000,000.00）、规费 5,255,339.95、增值税 26,287,916.06，
相加 318,375,871.65，等于 113 张表合计之和。

取数规矩同 §3.8（铁律 3）：行名整格相等、金额取「金额（元）」表头那一列、只取 unit 层；各项相加与该表
「招标控制价合计」相差 ≤ 1 元才采用这张表，任一张不过整张构成表不显示并写原因。

夹具 shapes.fee_breakdown_units 照原件的单位工程费用表形状造（「措施项目费」小计行、「其中：…」行、
「专业工程暂估价」行都在，当成一项就对不上合计）。期望值逐项手写，不从解析器输出反抄。
"""
from __future__ import annotations

import pytest

from tests.costsheet.fixtures.endpoint import HEADERS, client
from zt.costsheet import parser as cs
from zt.indicator import engine as indicator_engine
from zt.costsheet.warning_codes import LEVELS
from tests.costsheet.fixtures import costsheet_shapes as shapes

A, B = shapes.FEE_BD_A, shapes.FEE_BD_B
MAIN = ["fbfx", "measure_tech", "measure_org", "other", "regulatory", "tax"]


def _run(*files, total_amount=None):
    price = [(raw, name, f"F{i}") for i, (raw, name) in enumerate(files, 1)]
    return indicator_engine.run(price, None, gross_area=1000.0, total_amount=total_amount)


def _rows(res) -> dict:
    return {r["key"]: r for r in res["fee_breakdown"]["rows"]}


def _codes(res, code):
    return [w for w in res["warning_items"] if w["code"] == code]


# ---------------------------------------------------------------- 构成表六行

def test_two_unit_tables_reconcile_six_rows_sum_to_total():
    res = _run((shapes.fee_breakdown_units(), "一标段.xlsx"))
    fb = res["fee_breakdown"]
    assert fb["available"] is True and fb["reason"] is None
    assert fb["unit_count"] == 2
    rows = _rows(res)
    # 六项 = 两张单位工程表对应行之和（手算）
    want = {"fbfx": 1280000.00, "measure_tech": 104000.25, "measure_org": 52000.25, "other": 510000.00,
            "regulatory": 32400.50, "tax": 178964.37}
    assert {k: rows[k]["amount"] for k in MAIN} == want
    assert [r["label"] for r in fb["rows"] if r["level"] == 0] == \
        ["分部分项工程费", "施工技术措施项目", "施工组织措施项目", "其他项目费", "规费", "增值税"]
    assert fb["total"] == shapes.FEE_BD_TOTAL == 2157365.37
    assert cs.to_cents(sum(rows[k]["amount"] for k in MAIN)) == cs.to_cents(fb["total"])
    # 合计 = 两张表「招标控制价合计」之和（不是哪一张的、也不是项目级取整后的数）
    assert cs.to_cents(fb["total"]) == cs.to_cents(shapes.FEE_BD_TOTAL_A) + cs.to_cents(shapes.FEE_BD_TOTAL_B)
    # 占比的分母是构成表合计：六项相加 100%
    assert rows["fbfx"]["pct"] == round(1280000.00 / 2157365.37, 4)
    assert abs(sum(rows[k]["pct"] for k in MAIN) - 1) < 0.0005
    assert not _codes(res, "fee_breakdown_unreconciled")


def test_other_fee_carries_provisional_and_contingency_as_sub_rows():
    """07-005「真正的暂估价在哪」：暂估价、暂列金额是其他项目费的「其中」，只展示、不参与相加。
    「3.2.2 专业工程暂估价」整格不等于「暂估价」—— 包含匹配会把 500,000 数成 1,000,000。"""
    res = _run((shapes.fee_breakdown_units(), "一标段.xlsx"))
    fb = res["fee_breakdown"]
    keys = [r["key"] for r in fb["rows"]]
    assert keys == ["fbfx", "measure_tech", "measure_org", "other", "provisional", "contingency", "regulatory", "tax"]
    rows = _rows(res)
    assert (rows["provisional"]["label"], rows["provisional"]["amount"], rows["provisional"]["level"]) == \
        ("其中：暂估价", 500000.00, 1)
    assert (rows["contingency"]["label"], rows["contingency"]["amount"]) == ("其中：暂列金额", 10000.00)


def test_qizhong_and_subtotal_rows_are_not_items():
    """「其中：安全文明施工基本费」「其中：人工费+机械费」「措施项目费（技术 + 组织的小计）」都不是构成项。"""
    res = _run((shapes.fee_breakdown_units(), "一标段.xlsx"))
    labels = [r["label"] for r in res["fee_breakdown"]["rows"]]
    assert not any("安全文明" in x or "人工费" in x for x in labels)
    assert "措施项目费" not in labels
    scan = _scan_unit(shapes.fee_breakdown_units(), "地下室土建")
    assert set(scan["items"]) == set(MAIN)
    assert scan["items"]["measure_tech"] == cs.to_cents(A["measure_tech"])
    assert scan["problem"] is None


def test_one_table_two_yuan_short_hides_whole_breakdown_with_reason():
    """室外道路那张的增值税少 2 元（合计不变）→ 那张各项相加比合计少 2.00 元 → 整张构成表不显示，写原因。"""
    res = _run((shapes.fee_breakdown_units(b_tax_delta=-2.0), "一标段.xlsx"))
    fb = res["fee_breakdown"]
    assert fb["available"] is False
    assert fb["rows"] == [] and fb["total"] is None
    assert "室外道路" in fb["reason"] and "相差 2.00 元" in fb["reason"]
    assert "107,362.44" in fb["reason"] and "107,364.44" in fb["reason"], "写出相加的数与合计两个数"
    w = _codes(res, "fee_breakdown_unreconciled")
    assert len(w) == 1 and w[0]["level"] == "info" and LEVELS["fee_breakdown_unreconciled"] == "info"
    assert "室外道路" in w[0]["message"]
    # 总造价的取数与采用规则不归本片：构成表对不上不影响总造价（WS26-h 的口径照旧）
    assert res["summary"]["cost_basis"] == "fee_table"


@pytest.mark.parametrize("delta", [-1.0, 1.0])
def test_within_one_yuan_still_shown(delta):
    res = _run((shapes.fee_breakdown_units(b_tax_delta=delta), "一标段.xlsx"))
    assert res["fee_breakdown"]["available"] is True


@pytest.mark.parametrize("delta", [-1.01, 1.01])
def test_just_over_one_yuan_hidden(delta):
    res = _run((shapes.fee_breakdown_units(b_tax_delta=delta), "一标段.xlsx"))
    assert res["fee_breakdown"]["available"] is False


def test_tax_row_named_shuijin_is_the_tax_item():
    """GB 2013 的费用表把第 5 项写作「税金」：同一项的另一个整格名称，不是包含匹配。"""
    res = _run((shapes.fee_breakdown_units(tax_label="税金"), "一标段.xlsx"))
    fb = res["fee_breakdown"]
    assert fb["available"] is True
    assert (_rows(res)["tax"]["label"], _rows(res)["tax"]["amount"]) == ("税金", 178964.37), "行名照表里写的"


# ---------------------------------------------------------------- 与总造价的关系

def test_b8_provisional_column_facts_and_fee_breakdown_share_one_workbook():
    """B8 n→j→h2：清单暂估列、项目级单格事实、单位表六项构成互不覆盖；所有金额期望按夹具手算。"""
    import io
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(shapes.fee_breakdown_units(project_total=2157365)))
    ws = next(s for s in wb.worksheets if "分部分项" in s.title and "地下室" in s.title)
    ws["D7"] = "钢筋含量暂定按100kg/m3"
    ws["K7"] = 12000   # 三行表头的其中/暂估价列；不计入清单合价，也不等于费用表暂估价
    buf = io.BytesIO()
    wb.save(buf)
    raw = buf.getvalue()
    parsed = cs.parse_costsheet(raw, "合成.xlsx")
    column = next(i for i in parsed["items"] if i["code"] == "010502001001")
    assert column["provisional_amount"] == 12000.0 and column["amount"] == 600000.0
    assert "暂定" in column["feature"]
    assert all("provisional_amount" not in i for i in parsed["items"] if i["code"] != "010502001001")
    project_hit = next(h for h in parsed["meta"]["fee_totals"] if h["tier"] == "project")
    assert (project_hit["cell"], project_hit["amount_str"], project_hit["display"]) == ("C8", "2157365.00", "2157365")
    # zt 改动：旧用例在这里还断言 report_facts.collect_facts 的项目级事实（C8 = 2157365.00、差 -0.37）；
    # report_facts 归 W3（zt.report），那三行断言随 W3 迁移
    res = _run((raw, "合成.xlsx"))
    assert res["summary"]["total_amount"] == 2157365.0
    assert res["summary"]["parsed_total"] == 1384000.25
    assert res["fee_breakdown"]["available"] is True and res["fee_breakdown"]["total"] == 2157365.37
    assert _rows(res)["provisional"]["amount"] == 500000.0
    assert _rows(res)["other"]["amount"] == 510000.0


def test_project_total_adopted_breakdown_notes_the_rounding_gap():
    """项目级合计 2,157,365（取整到元）与两张单位表之和差 0.37：总造价取项目级合计（WS26-h），
    构成表合计仍是单位表之和，并写明与总造价差多少。"""
    res = _run((shapes.fee_breakdown_units(project_total=2157365), "一标段.xlsx"))
    assert res["summary"]["total_amount"] == 2157365.0
    fb = res["fee_breakdown"]
    assert fb["available"] is True and fb["total"] == 2157365.37
    assert any("2,157,365.00" in n and "0.37" in n for n in fb["notes"])


def test_fee_total_not_adopted_hides_breakdown():
    """项目级合计比单位表之和多 2 元 → 总造价不采用费用表（WS26-h）→ 构成表也不能说成总造价的构成。"""
    res = _run((shapes.fee_breakdown_units(project_total=2157367.37), "一标段.xlsx"))
    assert res["summary"]["cost_basis"] == "itemized"
    fb = res["fee_breakdown"]
    assert fb["available"] is False and "没有采用" in fb["reason"]
    assert "差 2.00 元" in fb["reason"], "原因照抄采用规则的（手填总造价时没有 fee_total_unreadable 可指）"
    assert not _codes(res, "fee_breakdown_unreconciled"), "这不是构成表自己对不上，原因已在 fee_total_unreadable 里"


def test_user_input_total_breakdown_still_shown_and_gap_written():
    res = _run((shapes.fee_breakdown_units(), "一标段.xlsx"), total_amount=2200000)
    assert res["summary"]["cost_basis"] == "user_input"
    fb = res["fee_breakdown"]
    assert fb["available"] is True and fb["total"] == 2157365.37
    assert any("手填" in n and "2,200,000.00" in n and "42,634.63" in n for n in fb["notes"])


def test_no_fee_table_no_breakdown_no_warning():
    res = _run((shapes.fee_table_04_with_12(with_04=False), "清单.xlsx"))
    fb = res["fee_breakdown"]
    assert fb["available"] is False and fb["rows"] == []
    assert "单位工程费用表" in fb["reason"]
    assert not _codes(res, "fee_breakdown_unreconciled")


def test_two_files_breakdowns_add_up_one_missing_hides_with_file_name():
    a = shapes.fee_breakdown_units()
    b = shapes.resaved(shapes.fee_breakdown_units(), "另一标段")
    both = _run((a, "一标段.xlsx"), (b, "二标段.xlsx"))
    fb = both["fee_breakdown"]
    assert fb["available"] is True and fb["unit_count"] == 4
    assert cs.to_cents(fb["total"]) == 2 * cs.to_cents(shapes.FEE_BD_TOTAL)
    assert _rows(both)["provisional"]["amount"] == 1000000.00

    bad = _run((a, "一标段.xlsx"), (shapes.fee_breakdown_units(b_tax_delta=-2.0), "二标段.xlsx"))
    fb2 = bad["fee_breakdown"]
    assert fb2["available"] is False and "「二标段.xlsx」" in fb2["reason"]


# ---------------------------------------------------------------- 与清单明细勾稽（WS26-h 留给本片）

def test_itemized_matches_fee_table_no_mismatch_warning():
    res = _run((shapes.fee_breakdown_units(), "一标段.xlsx"))
    assert not _codes(res, "fee_itemized_mismatch")
    chk = res["fee_breakdown"]["itemized_check"]
    assert chk == [{"file": "一标段.xlsx", "fee_fbfx": 1280000.0, "list_fbfx": 1280000.0, "fbfx_diff": 0.0,
                    "fee_measure_tech": 104000.25, "list_measure_tech": 104000.25, "measure_tech_diff": 0.0}]


def test_list_changed_fee_table_not_recomputed_warns_with_both_numbers():
    """测试件 B 的形状：只改清单行、费用表没重算。总造价仍取自费用表（如实），但费用表的分部分项工程费与
    读入的清单对不上要说出来（WS26-h 评审留给本片的「费用表合计不与清单明细勾稽」）。"""
    res = _run((shapes.fee_breakdown_units(list_delta=-3000), "送审改价.xlsx"))
    w = _codes(res, "fee_itemized_mismatch")
    assert len(w) == 1 and w[0]["level"] == "warn"
    assert "1,280,000.00" in w[0]["message"] and "1,277,000.00" in w[0]["message"] and "3,000.00" in w[0]["message"]
    assert res["summary"]["total_amount"] == 2157365.37, "总造价的取数不归本片，照旧取自费用表"


def test_measure_not_read_skips_measure_comparison():
    """清单里没有技术措施表时，清单明细本来就不含措施（口径标签已写），不拿费用表的技术措施去比。"""
    res = _run((shapes.fee_breakdown_units(with_measure_lists=False), "一标段.xlsx"))
    assert not _codes(res, "fee_itemized_mismatch")
    chk = res["fee_breakdown"]["itemized_check"][0]
    assert chk["fbfx_diff"] == 0.0 and chk["list_measure_tech"] is None and chk["measure_tech_diff"] is None


# ---------------------------------------------------------------- 一张表的读法

def _scan_unit(raw: bytes, short: str) -> dict:
    import io

    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = next(w for w in wb.worksheets if "单位工程" in w.title and short in w.title)
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    return cs.scan_fee_breakdown(ws.title, rows, cs.scan_fee_sheet(ws.title, rows))


def test_scan_reads_amount_column_by_header_not_rightmost_number():
    """金额取「金额（元）」那一列：「1.1 其中：人工费+机械费」右边没有别的数，但「1 分部分项工程费」的
    金额格是文本、右边是「见表…」——都按表头那一列取。"""
    scan = _scan_unit(shapes.fee_breakdown_units(), "地下室土建")
    assert scan["items"]["fbfx"] == cs.to_cents(1200000.00)
    assert scan["total"] == cs.to_cents(shapes.FEE_BD_TOTAL_A)
    assert scan["subs"] == {"provisional": cs.to_cents(500000), "contingency": 0}


def test_scan_non_unit_tables_return_none():
    import io

    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(shapes.fee_breakdown_units(project_total=2157365)), read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        got = cs.scan_fee_breakdown(ws.title, rows, cs.scan_fee_sheet(ws.title, rows))
        assert (got is not None) == ("单位工程招标控制价费用表" in ws.title), ws.title


def test_scan_duplicate_item_row_is_a_problem():
    rows = [["单位工程招标控制价费用表"], ["序号", "费用名称", "计算公式", "金额（元）"],
            ["1", "分部分项工程费", None, 100], ["4", "规费", None, 5], ["4", "规费", None, 5],
            ["5", "增值税", None, 10], ["招标控制价合计", None, None, 120]]
    scan = cs.scan_fee_breakdown("单位工程费用表", rows, cs.scan_fee_sheet("单位工程费用表", rows))
    assert scan["problem"] and "「规费」" in scan["problem"] and "2 行" in scan["problem"]


def test_scan_without_fbfx_row_is_a_problem():
    rows = [["单位工程招标控制价汇总表"], ["序号", "汇总内容", "金额(元)"],
            ["1", "分部分项工程", 100], ["4", "规费", 5], ["5", "税金", 10], [None, "招标控制价合计", 115]]
    scan = cs.scan_fee_breakdown("单位工程费用表", rows, cs.scan_fee_sheet("单位工程费用表", rows))
    assert scan["problem"] and "分部分项工程费" in scan["problem"]


# ---------------------------------------------------------------- 端点：结果里带出 fee_breakdown

def test_endpoint_carries_fee_breakdown():
    """抽屉读的是 engine_task.result（引擎响应原文）—— IndicatorOut 不声明这个键，响应里就没有它。"""
    headers = HEADERS
    ref = shapes.put(shapes.fee_breakdown_units(project_total=2157365), "ws26h2-breakdown.xlsx", "T")
    r = client.post("/engine/indicator/calc", headers=headers, json={"gross_area": 1000, "price_files": [ref]})
    assert r.status_code == 200, r.text[:300]
    fb = r.json()["fee_breakdown"]
    assert fb["available"] is True and fb["total"] == 2157365.37
    assert [x["amount"] for x in fb["rows"] if x["key"] in ("other", "provisional")] == [510000.0, 500000.0]


def test_measure_total_not_split_is_named_as_the_reason():
    """GB 2013 一类的费用表只有一行「措施项目费」：构成表按技术、组织两项拆，拆不开要把原因说准，
    不是只报一个「相差 812,345.67 元」。"""
    res = _run((shapes.fee_table_13(), "房建.xlsx"))
    fb = res["fee_breakdown"]
    assert fb["available"] is False
    assert "措施项目费没有拆成" in fb["reason"]
    assert _codes(res, "fee_breakdown_unreconciled"), "单位工程表拆不开属于「构成表对不上」"
    assert res["summary"]["total_amount"] == 8715564.88, "总造价的取数不受影响"


def test_extra_discount_row_is_named_in_the_reason():
    """样本库 1058 份带费用表的工作簿里，拆不开的多半是表尾多一行「6 让利 / 6 下浮」：只报「相差 x 元」说不清，
    原因里点出那一行。构成表只拆六项，不把让利当成第七项塞进去（不认识的行不猜）。"""
    res = _run((shapes.fee_breakdown_units(b_discount=-5000.0), "一标段.xlsx"))
    fb = res["fee_breakdown"]
    assert fb["available"] is False
    assert "「让利」-5,000.00" in fb["reason"] and "相差 5,000.00 元" in fb["reason"]


@pytest.mark.parametrize("delta,warned", [(-2.0, False), (-3.0, True)])
def test_whole_yuan_list_gets_rounding_tolerance(delta, warned):
    """清单合价全是整元（软件逐行取整导出）：费用表按到分算，逐行最多差 0.5 元。3 行 → 允许 0.5 × 3 + 0.5 = 2 元。"""
    res = _run((shapes.fee_breakdown_units(list_delta=delta), "一标段.xlsx"))
    assert bool(_codes(res, "fee_itemized_mismatch")) is warned
    assert res["fee_breakdown"]["itemized_check"][0]["fbfx_diff"] == delta


def test_list_with_cents_gets_one_yuan_tolerance():
    """有一行带角分就不是取整导出：差 1.50 元就要说出来（测试件 C 审定核减 3,345.70 也是这一类）。"""
    res = _run((shapes.fee_breakdown_units(list_delta=-1.5), "一标段.xlsx"))
    w = _codes(res, "fee_itemized_mismatch")
    assert len(w) == 1 and "相差 1.50 元" in w[0]["message"]
