# 来源：engine-service/tests/test_indicator_scope_basis.py（zhitong-cost-ai@991ff5cd）
"""WS08-14 / 15 / 16（issue #760）：指标的面积入口、送检范围与降级门槛、跨文件查重、总造价口径。

修的现场：
  · 部分文件失败静默降级：长兴 16 份控制价一份没读进，只用一张万元估算表算出「单方 0.04 元/㎡」，返回 200；
  · 同一份文件的 xls 与 xlsx、清单版与控制价版重复累计；
  · 只取主单元面积；按面积「分摊」的近似行每个单体都是同一个值（D41 删）；
  · 结果标「单方造价（含税总造价口径）」，实际只是明细合计 —— A2 引擎 1631.32，真实 1859.38。
"""
from __future__ import annotations

import pytest

from tests.costsheet.fixtures.endpoint import HEADERS, client
from zt.indicator import engine as indicator_engine
from tests.costsheet.fixtures import costsheet_shapes as S
from tests.costsheet.fixtures import costsheet_shapes as shapes

BAD = b"not a zip archive at all"


def _codes(res) -> list[str]:
    return [w["code"] for w in res["warning_items"]]


def _bill(sheet: str, rows: list[list]) -> bytes:
    import io
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(["序号", "项目编码", "项目名称", "项目特征描述", "计量单位", "工程量", "综合单价", "合价"])
    for i, r in enumerate(rows, 1):
        ws.append([i] + list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- WS08-15 降级门槛（D40）

def test_two_of_three_unreadable_is_422():
    # 两份坏文件字节必须不同 —— 相同的话会先被 sha 查重合并成一份（1/2 = 50%，不触发门槛）
    files = [(S.priced_bill(), "控制价.xlsx", "F1"), (BAD + b"1", "坏1.xlsx", "F2"), (BAD + b"2", "坏2.xlsx", "F3")]
    with pytest.raises(indicator_engine.ParseFailed) as e:
        indicator_engine.run(files, {}, gross_area=1000)
    msg = str(e.value)
    assert "3 份" in msg and "1 份可计入" in msg and "坏1.xlsx" in msg and "坏2.xlsx" in msg


def test_one_of_two_unreadable_is_200_with_files_failed():
    res = indicator_engine.run([(S.priced_bill(), "控制价.xlsx", "F1"), (BAD, "坏.xlsx", "F2")], {}, gross_area=1000)
    assert res["files_used"] == ["控制价.xlsx"]
    assert [f["file"] for f in res["files_failed"]] == ["坏.xlsx"]
    assert "files_failed" in _codes(res)
    assert res["summary"]["total_amount"] == 48000


def test_estimate_table_plus_control_price_counts_only_the_control_price():
    """万元估算表不是国标清单（require_gb_code）→ 只计控制价，单方不再是 0.04 元/㎡。"""
    res = indicator_engine.run([(S.wan_estimate(), "设备估算表.xlsx", "F1"), (S.priced_bill(), "控制价.xlsx", "F2")],
                               {}, gross_area=1000)
    assert res["files_used"] == ["控制价.xlsx"]
    assert res["files_failed"][0]["status"] == "non_list"
    assert res["summary"]["unit_cost"] == 48.0


def test_all_selected_non_list_is_422():
    with pytest.raises(indicator_engine.ParseFailed):
        indicator_engine.run([(S.wan_estimate(), "设备估算表.xlsx", "F1")], {}, gross_area=1000)


# ---------------------------------------------------------------- WS08-15 跨文件查重

def test_same_file_twice_counts_once():
    raw = S.priced_bill()
    res = indicator_engine.run([(raw, "控制价.xlsx", "F1"), (raw, "控制价(1).xlsx", "F2")], {}, gross_area=1000)
    assert res["summary"]["total_amount"] == 48000, "同一份文件放两次，总额翻倍了"
    assert "file_duplicate" in _codes(res)
    assert res["files_used"] == ["控制价.xlsx"] and res["files_selected"] == 1


def test_identical_row_sets_in_two_different_files_are_flagged_not_dropped():
    rows = [["010101001001", "平整场地", "", "m2", 1000, 12, 12000], ["010401003001", "实心砖墙", "", "m3", 50, 480, 24000]]
    a, b = _bill("6#楼", rows), _bill("7#楼", rows)
    assert a != b, "夹具失效：两份文件字节不能相同（否则走的是 sha 查重）"
    res = indicator_engine.run([(a, "6#楼.xlsx", "F1"), (b, "7#楼.xlsx", "F2")], {}, gross_area=1000)
    assert "file_overlap_suspect" in _codes(res)
    assert res["summary"]["total_amount"] == 72000, "只提示不剔除：同构楼栋可能确实同清单"


# ---------------------------------------------------------------- WS08-14 面积

def test_no_area_gives_null_unit_cost_and_area_missing():
    res = indicator_engine.run([(S.priced_bill(), "控制价.xlsx", "F1")], {})
    unit = next(r for r in res["indicators"] if r["indicator_code"] == "unit_cost")
    assert unit["value"] is None and unit["reason"] == "未提供总建筑面积"
    assert "area_missing" in _codes(res)
    assert all(c["unit_cost"] is None for c in res["chapters"])


def test_area_share_rows_are_gone_even_with_wbs_units():
    res = indicator_engine.run([(S.priced_bill(), "控制价.xlsx", "F1")], {},
                               [{"id": "W1", "gross_area": 600}, {"id": "W2", "gross_area": 400}],
                               gross_area=1000, area_source="wbs_sum")
    assert not [r for r in res["indicators"] if r["indicator_code"] == "unit_cost_area_share"]
    assert res["summary"]["gross_area"] == 1000 and res["summary"]["area_source"] == "wbs_sum"


# ---------------------------------------------------------------- WS08-16 总造价口径

def test_fee_table_basis_gives_1859_38():
    res = indicator_engine.run([(S.fee_table_13(), "房建控制价.xlsx", "F1")], {}, gross_area=4687.35,
                               area_source="user_input")
    sm = res["summary"]
    assert sm["cost_basis"] == "fee_table" and sm["total_amount"] == 8715564.88
    assert sm["unit_cost"] == 1859.38
    assert sm["cost_basis_label"] == "含规费税金，取自费用表"
    assert sm["fee_totals"] == [{"file": "房建控制价.xlsx", "file_id": "F1", "fee_total": 8715564.88}]
    total = next(r for r in res["indicators"] if r["indicator_code"] == "total_amount")
    assert "取自费用表" in total["label"]


def test_distractor_row_is_not_taken_as_fee_total():
    """「其中：安全文明施工费合计」也是「…合计」，但不是白名单行名 —— 整格相等才算。"""
    from zt.costsheet.parser import parse_costsheet
    meta = parse_costsheet(S.fee_table_13(), "x.xlsx")["meta"]
    assert [h["label"] for h in meta["fee_totals"]] == ["招标控制价合计"]


def test_itemized_label_never_says_tax_inclusive():
    res = indicator_engine.run([(S.priced_bill(), "控制价.xlsx", "F1")], {}, gross_area=1000)
    assert res["summary"]["cost_basis"] == "itemized"
    labels = [r.get("label") or "" for r in res["indicators"]]
    assert not [lb for lb in labels if "含税" in lb], labels
    assert "fee_total_unavailable" in _codes(res)
    assert "fee_total_unreadable" not in _codes(res), "压根没有费用表：info 级说明即可，不是「有却取不到」"


def test_fee_table_needs_every_file_to_have_one():
    """两份文件只有一份有费用表 → 不能拿一份的含税价加另一份的明细价，退回 itemized。"""
    res = indicator_engine.run([(S.fee_table_13(), "房建.xlsx", "F1"), (S.priced_bill(), "市政.xlsx", "F2")],
                               {}, gross_area=1000)
    assert res["summary"]["cost_basis"] == "itemized"


def test_user_input_total_wins():
    res = indicator_engine.run([(S.fee_table_13(), "房建.xlsx", "F1")], {}, gross_area=4687.35, total_amount=9000000)
    assert res["summary"]["cost_basis"] == "user_input" and res["summary"]["total_amount"] == 9000000


# ---------------------------------------------------------------- WS26-h 费用表分层取数 + 口径如实（07-012）

def _row(res, code: str) -> dict:
    return next(r for r in res["indicators"] if r["indicator_code"] == code)


def _all_text(res) -> str:
    """结果里所有给人看的字：指标行名、口径标签、提示行、告警。"""
    sm = res["summary"]
    parts = [r.get("label") or "" for r in res["indicators"]] + [w["message"] for w in res["warning_items"]]
    parts += [sm["cost_basis_label"], sm["itemized_basis_label"], sm.get("cost_basis_missing") or ""]
    parts += sm["cost_basis_notes"]
    return "\n".join(parts)


def test_three_tier_workbook_takes_the_project_total():
    """生产 IND-05 的形状：改前总造价 = 清单明细合计（控制价的 58.2%），口径标签「分部分项+措施」。
    夹具真值：项目级合计 15,873,168；单位工程之和 15,873,167.65；清单明细 12,797,000（= 80.6%）。"""
    res = indicator_engine.run([(shapes.fee_three_tier(with_single=True), "一标段.xlsx", "F1")], {},
                               gross_area=4319.5, area_source="user_input")
    sm = res["summary"]
    assert sm["cost_basis"] == "fee_table" and sm["total_amount"] == 15873168.0
    assert sm["cost_basis_label"] == "含规费税金，取自费用表"
    assert sm["parsed_total"] == 12797000.0 and sm["parsed_total"] < sm["total_amount"]
    assert sm["unit_cost"] == 3674.77          # 15,873,168 ÷ 4319.5
    assert sm["cost_basis_missing"] is None and _codes(res).count("fee_total_unreadable") == 0
    # 提示行：P、U、差额三个数；单项工程之和单独一句
    assert sm["cost_basis_notes"] == [
        "项目级合计 15,873,168.00，2 张单位工程费用表之和 15,873,167.65，差 0.35 元",
        "1 张单项工程费用表之和 15,873,167.65（只用于核对，不参与取数）"]
    assert sm["fee_total_checks"] == [{
        "file": "一标段.xlsx", "file_id": "F1", "source": "project", "has_fee_sheet": True,
        "project_total": 15873168.0, "unit_sum": 15873167.65, "unit_count": 2,
        "single_sum": 15873167.65, "single_count": 1, "specialty_sum": None, "specialty_count": 0,
        "diff": 0.35, "reason": None}]
    assert _row(res, "total_amount")["label"] == "总造价(元,含规费税金，取自费用表)"
    assert _row(res, "total_amount")["value"] == 15873168.0
    assert _row(res, "parsed_total")["value"] == 12797000.0


def test_mismatched_fee_table_falls_back_and_is_not_called_total_cost():
    """项目级合计比单位工程之和大 2 元 → 不采用；卡片上的数是清单明细合计，哪儿都不许叫「总造价」。"""
    res = indicator_engine.run([(shapes.fee_three_tier(project_total=15873169.65), "一标段.xlsx", "F1")], {},
                               gross_area=1000)
    sm = res["summary"]
    assert sm["cost_basis"] == "itemized" and sm["total_amount"] == 12797000.0
    w = [x for x in res["warning_items"] if x["code"] == "fee_total_unreadable"]
    assert len(w) == 1 and w[0]["level"] == "warn", "有费用表却没采用：warn 级，不是灰字"
    for num in ("15,873,169.65", "15,873,167.65", "2.00"):
        assert num in w[0]["message"], w[0]["message"]
    assert "含税总造价" in w[0]["message"] and "手填" in w[0]["message"]
    assert "fee_total_unavailable" not in _codes(res)
    total, unit = _row(res, "total_amount"), _row(res, "unit_cost")
    assert total["label"].startswith("清单明细合计(元") and "不是总造价" in total["label"]
    assert "不是单方造价指标" in unit["label"] and not unit["label"].startswith("单方造价")
    assert not [r["label"] for r in res["indicators"]
                if "总造价" in r["label"].replace("不是总造价", "")], "回落后仍有行叫「总造价」"
    assert sm["cost_basis_missing"] == "措施项目费、其他项目费、规费、税金"
    assert sm["cost_basis_notes"] == []


def test_missing_small_unit_sheet_falls_back():
    res = indicator_engine.run([(shapes.fee_three_tier(with_small_unit=False), "一标段.xlsx", "F1")], {})
    assert res["summary"]["cost_basis"] == "itemized"
    assert "3,895.79" in next(w["message"] for w in res["warning_items"] if w["code"] == "fee_total_unreadable")


def test_generic_project_total_without_unit_sheets_falls_back():
    res = indicator_engine.run([(shapes.fee_three_tier(with_units=False), "一标段.xlsx", "F1")], {})
    assert res["summary"]["cost_basis"] == "itemized" and res["summary"]["total_amount"] == 12797000.0
    assert "fee_total_unreadable" in _codes(res)


def test_unit_sum_basis_label_says_it_is_a_sum():
    res = indicator_engine.run([(shapes.fee_three_tier(project_total=None), "一标段.xlsx", "F1")], {})
    sm = res["summary"]
    assert sm["cost_basis"] == "fee_table" and sm["total_amount"] == 15873167.65
    assert sm["cost_basis_label"] == "各单位工程费用表合计之和(2 张)"
    assert sm["cost_basis_notes"] == ["2 张单位工程费用表之和 15,873,167.65"]


def test_unit_sheet_with_specialty_sheets_is_not_doubled():
    """评审抓到的翻倍：1 张单位工程费用表（7,516,640.82）+ 它下面 2 张专业工程费用表（抬头也写
    「单位工程名称：…」）。第一版给出 15,033,281.64、口径「各单位工程费用表合计之和(3 张)」、没有任何告警。"""
    res = indicator_engine.run([(shapes.fee_unit_with_specialty(), "一期02标-土建.xlsx", "F1")], {},
                               gross_area=1000)
    sm = res["summary"]
    assert sm["cost_basis"] == "fee_table" and sm["total_amount"] == 7516640.82
    assert sm["total_amount"] != 15033281.64, "专业工程表与单位工程表相加了（翻倍）"
    assert sm["cost_basis_label"] == "各单位工程费用表合计之和(1 张)"
    assert sm["unit_cost"] == 7516.64 and sm["parsed_total"] == 6200000.5
    assert sm["cost_basis_notes"] == ["1 张单位工程费用表之和 7,516,640.82",
                                      "2 张专业工程费用表之和 7,516,640.82（只用于核对，不参与取数）"]
    chk = sm["fee_total_checks"][0]
    assert (chk["source"], chk["unit_count"], chk["specialty_sum"], chk["specialty_count"]) == \
        ("unit_sum", 1, 7516640.82, 2)
    assert not [c for c in _codes(res) if c.startswith("fee_total")]


def test_project_total_with_unit_and_specialty_sheets_is_adopted():
    """同上再加项目级表：第一版因为 U = 2P 报「对不上」而回落，卡片成了「清单明细合计」。"""
    res = indicator_engine.run([(shapes.fee_unit_with_specialty(with_project=True), "一期02标-土建.xlsx", "F1")], {})
    sm = res["summary"]
    assert (sm["cost_basis"], sm["total_amount"], sm["cost_basis_label"]) == \
        ("fee_table", 7516641.0, "含规费税金，取自费用表")
    assert sm["cost_basis_notes"][0] == "项目级合计 7,516,641.00，1 张单位工程费用表之和 7,516,640.82，差 0.18 元"
    assert "fee_total_unreadable" not in _codes(res)


def test_unit_and_specialty_sums_that_disagree_fall_back_with_both_numbers():
    res = indicator_engine.run([(shapes.fee_unit_with_specialty(b_total="1445343.82"), "一期02标-土建.xlsx", "F1")], {})
    sm = res["summary"]
    assert sm["cost_basis"] == "itemized" and sm["total_amount"] == 6200000.5
    w = [x for x in res["warning_items"] if x["code"] == "fee_total_unreadable"]
    assert len(w) == 1 and w[0]["level"] == "warn"
    for part in ("7,516,640.82", "7,516,643.82", "3.00"):
        assert part in w[0]["message"], w[0]["message"]


def test_only_specialty_sheets_fall_back_and_the_warning_carries_their_sum():
    """只有专业工程表（2 张）：不采用，卡片是清单明细合计；黄色告警里写出这两张的和，用户可以手填。"""
    res = indicator_engine.run([(shapes.fee_unit_with_specialty(with_unit=False), "一期02标-土建.xlsx", "F1")], {})
    sm = res["summary"]
    assert sm["cost_basis"] == "itemized" and sm["total_amount"] == 6200000.5
    assert _row(res, "total_amount")["label"].startswith("清单明细合计(元")
    msg = next(w["message"] for w in res["warning_items"] if w["code"] == "fee_total_unreadable")
    assert "只有 2 张专业工程费用表之和 7,516,640.82" in msg and "手填" in msg


def test_table_12_total_never_shows_up_as_total_cost():
    """只有表-12（其他项目清单与计价汇总表，合计 20,000）→ 回落；结果里没有 20,000 这个数。"""
    res = indicator_engine.run([(shapes.fee_table_04_with_12(with_04=False), "清单.xlsx", "F1")], {})
    sm = res["summary"]
    assert sm["cost_basis"] == "itemized" and sm["total_amount"] == 126000.0
    assert sm["fee_totals"] == [{"file": "清单.xlsx", "file_id": "F1", "fee_total": None}]
    assert 20000 not in [r["value"] for r in res["indicators"]]
    assert "fee_total_unavailable" in _codes(res) and "fee_total_unreadable" not in _codes(res)
    both = indicator_engine.run([(shapes.fee_table_04_with_12(), "清单.xlsx", "F1")], {})
    assert both["summary"]["cost_basis"] == "fee_table" and both["summary"]["total_amount"] == 172220.0


def test_itemized_label_without_measure_sheet_does_not_claim_measures():
    """没读到措施项目清单表：标签是「分部分项清单合计口径(未含措施项目、规费、税金)」，
    任何一处都不出现「分部分项+措施」。"""
    res = indicator_engine.run([(S.priced_bill(), "控制价.xlsx", "F1")], {}, gross_area=1000)
    sm = res["summary"]
    assert sm["cost_basis_label"] == "分部分项清单合计口径(未含措施项目、规费、税金)"
    assert sm["itemized_basis_label"] == sm["cost_basis_label"]
    assert "分部分项+措施" not in _all_text(res), _all_text(res)
    # 采用了费用表时，「按清单明细合计算的单方」那一行同样不许说「+措施」
    fee = indicator_engine.run([(shapes.fee_three_tier(), "一标段.xlsx", "F1")], {}, gross_area=1000)
    assert "分部分项+措施" not in _all_text(fee), _all_text(fee)
    assert _row(fee, "unit_cost_itemized")["label"] == \
        "按清单明细合计算的单方(元/㎡,分部分项清单合计口径(未含措施项目、规费、税金))"


def test_itemized_label_with_measure_sheet_says_so():
    """读到了措施项目清单表（国标编码的单价措施项）→ 标签才写「分部分项+措施」。"""
    import io
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "分部分项工程量清单"
    head = ["序号", "项目编码", "项目名称", "项目特征描述", "计量单位", "工程量", "综合单价", "合价"]
    ws.append(head)
    ws.append([1, "010401003001", "实心砖墙", "墙厚240", "m3", 100, 480, 48000])
    ws2 = wb.create_sheet("单价措施项目清单与计价表")
    ws2.append(head)
    ws2.append([1, "011701001001", "综合脚手架", "檐高 30m 内", "m2", 1000, 12, 12000])
    buf = io.BytesIO()
    wb.save(buf)
    res = indicator_engine.run([(buf.getvalue(), "控制价.xlsx", "F1")], {}, gross_area=1000)
    sm = res["summary"]
    assert sm["cost_basis"] == "itemized" and sm["total_amount"] == 60000
    assert sm["cost_basis_label"] == "分部分项+措施清单合计口径"
    assert sm["cost_basis_missing"] == "措施费里没读入的部分、其他项目费、规费、税金"


def test_user_input_path_is_unchanged_by_fee_table_reading():
    """IND-07 手填那条路：口径「用户录入」，不出费用表的提示行，也不出回落告警。"""
    res = indicator_engine.run([(shapes.fee_three_tier(project_total=15873169.65), "一标段.xlsx", "F1")], {},
                               gross_area=1000, total_amount=15873168)
    sm = res["summary"]
    assert (sm["cost_basis"], sm["cost_basis_label"], sm["total_amount"]) == ("user_input", "用户录入", 15873168)
    assert sm["unit_cost"] == 15873.17
    assert sm["cost_basis_notes"] == [] and sm["cost_basis_missing"] is None
    assert not [c for c in _codes(res) if c.startswith("fee_total")]
    assert _row(res, "total_amount")["label"] == "总造价(元,用户录入)"


def test_two_files_notes_carry_the_file_name():
    a = shapes.fee_three_tier()
    b = shapes.fee_table_13()
    res = indicator_engine.run([(a, "一标段.xlsx", "F1"), (b, "房建.xlsx", "F2")], {})
    sm = res["summary"]
    assert sm["cost_basis"] == "fee_table" and sm["total_amount"] == round(15873168 + 8715564.88, 2)
    assert sm["cost_basis_notes"] == ["「一标段.xlsx」项目级合计 15,873,168.00，2 张单位工程费用表之和 15,873,167.65，差 0.35 元"]


def test_endpoint_carries_the_new_summary_fields():
    headers = HEADERS
    ref = S.put(shapes.fee_three_tier(), "ws26h-three-tier.xlsx", "T")
    r = client.post("/engine/indicator/calc", headers=headers, json={"gross_area": 4319.5, "price_files": [ref]})
    assert r.status_code == 200, r.text[:300]
    sm = r.json()["summary"]
    assert (sm["cost_basis"], sm["total_amount"], sm["unit_cost"]) == ("fee_table", 15873168.0, 3674.77)
    assert sm["cost_basis_notes"] and sm["fee_total_checks"][0]["diff"] == 0.35
    bad = S.put(shapes.fee_three_tier(project_total=15873169.65), "ws26h-three-tier-bad.xlsx", "B")
    r2 = client.post("/engine/indicator/calc", headers=headers, json={"price_files": [bad]})
    body = r2.json()
    assert r2.status_code == 200 and body["summary"]["cost_basis"] == "itemized"
    item = next(w for w in body["warning_items"] if w["code"] == "fee_total_unreadable")
    assert item["level"] == "warn" and item["file"] == "ws26h-three-tier-bad.xlsx"


# ---------------------------------------------------------------- 端点：C1 状态

def test_endpoint_partial_status_when_a_file_is_not_counted():
    headers = HEADERS
    good = S.put(S.priced_bill(), "ws08b-ind-good.xlsx", "G")
    bad = S.put(S.wan_estimate(), "ws08b-ind-est.xlsx", "E")
    r = client.post("/engine/indicator/calc", headers=headers,
                    json={"gross_area": 1000, "area_source": "user_input", "price_files": [good, bad]})
    assert r.status_code == 200, r.text[:300]
    body = r.json()
    assert body["status"] == "partial" and body["status_reason"] == "items_not_counted"
    assert body["summary"]["area_source"] == "user_input"
    r2 = client.post("/engine/indicator/calc", headers=headers, json={"price_files": [bad]})
    assert r2.status_code == 422
