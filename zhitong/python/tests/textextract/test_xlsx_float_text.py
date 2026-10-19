"""xlsx 正文里的浮点格（整改 WS22-l，spec §3.14 ⑥，E31 第 9 点；缺陷 01-010 第 9 点）。

生产 E13：`X_面积统计表.xlsx` 汇总表「总计」那一格存的是 179159.09499999997（Excel 公式算出来的二进制浮点），
Excel 里看到的是 179159.095；xlsx 正文原样 `str()`，模型照抄，候选就成了 `179159.09499999997`。

`textextract._cell_str` 是共享函数：字段提取、文件树问答（`filetree.tools` 经 `read_xlsx`）、报告取材的 xlsx 正文
都经过它，所以这里同时钉住「不该变的不变」。
"""
from __future__ import annotations

import io
import zipfile

import pytest
from openpyxl import Workbook

from zt.textextract.xlsx import _cell_str, _read_xlsx_raw, read_xlsx, xlsx_text


@pytest.mark.parametrize("value,want", [
    (179159.09499999997, "179159.095"),     # E13 原样
    (0.1 + 0.2, "0.3"),
    (12.34, "12.34"),                       # 本来就是 2 位小数
    (10613893.4628165, "10613893.4628165"),  # 正好 15 位有效数字：一位都不丢
    (11894327.9609491, "11894327.9609491"),
    (-1418635.5, "-1418635.5"),
    (318375872.0, "318375872"),             # 浮点格里的整数值：Excel 里看到的也没有「.0」
    (0.0, "0"),
    (-0.0, "0"),
    (1234567890123456.0, "1234567890123460"),  # 1e15–1e16 这一段不写成科学计数（金额不能变成 1.2e+15）
    (2.5e-05, "2.5e-05"),                   # 很小的数与改前同样写法
    (1e16, "1e+16"),
    # 求和累下的误差尾巴（形状照本机原件「X_面积统计表.xlsx」里读到的格）：15 位用满、末尾一串 0 / 9
    (1737.66000000007, "1737.66"),
    (379.860000000001, "379.86"),
    (25.3099999999977, "25.31"),
    (-25.3099999999977, "-25.31"),
    (0.409999999999997, "0.41"),
    (437.000000000015, "437"),
    (9.99999999999997, "10"),
    # 不是误差的不碰：本来就短的数、15 位但中间没有一长串 0 / 9 的数
    (2.0000005, "2.0000005"),
    (2.000005, "2.000005"),
    (123456.789012345, "123456.789012345"),
])
def test_float_cells_use_15_significant_digits(value, want):
    assert _cell_str(value) == want


@pytest.mark.parametrize("value,want", [
    (1234567, "1234567"),                   # 整数格
    (0, "0"),
    (True, "True"),
    ("179159.09499999997", "179159.09499999997"),  # 文本格：原样，不当数处理
    ("合计\n（元）", "合计 （元）"),          # 格内换行照旧换成空格
    (None, ""),
])
def test_non_float_cells_are_unchanged(value, want):
    assert _cell_str(value) == want


def _patch_sheet(content: bytes, swaps: dict[bytes, bytes]) -> bytes:
    src = io.BytesIO(content)
    out = io.BytesIO()
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(out, "w") as zout:
        for n in zin.namelist():
            b = zin.read(n)
            if n.startswith("xl/worksheets/sheet"):
                for a, c in swaps.items():
                    assert a in b, (n, a)
                    b = b.replace(a, c)
            zout.writestr(n, b)
    return out.getvalue()


def _book() -> bytes:
    """openpyxl 写浮点用短写法（179159.095），Excel 写的是 17 位的 `179159.09499999997` —— 夹具照 Excel 的写法
    改 sheet XML，否则读回来的本来就是短的，测不到要修的那个形状。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "汇总"
    ws["B9"], ws["C9"] = "总计", 179159.095
    ws["B10"], ws["C10"] = "文本格", "179159.09499999997"
    ws["B11"], ws["C11"] = "整数", 172149
    ws["B12"], ws["C12"] = "两位小数", 172149.31
    ws["B13"], ws["C13"] = "公式结果", 0.3
    buf = io.BytesIO()
    wb.save(buf)
    return _patch_sheet(buf.getvalue(), {b"<v>179159.095</v>": b"<v>179159.09499999997</v>",
                                         b"<v>0.3</v>": b"<v>0.30000000000000004</v>"})


def test_fixture_really_has_the_long_float():
    """夹具自检：openpyxl 读到的是 Excel 存的那个浮点（179159.09499999997），不是短写法。"""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(_book()), read_only=True, data_only=True)
    vals = [c for row in wb.active.iter_rows(values_only=True) for c in row if isinstance(c, float)]
    assert 179159.09499999997 in vals and repr(179159.09499999997) == "179159.09499999997"
    assert 0.1 + 0.2 in vals


def test_xlsx_text_writes_what_excel_shows():
    text = xlsx_text(_book())
    assert "[R9] B:总计\tC:179159.095" in text
    assert "[R10] B:文本格\tC:179159.09499999997" in text
    assert "[R11] B:整数\tC:172149" in text
    assert "[R12] B:两位小数\tC:172149.31" in text
    assert "[R13] B:公式结果\tC:0.3" in text
    assert "179159.09499999997" not in text.replace("[R10] B:文本格\tC:179159.09499999997", "")
    assert "0.30000000000000004" not in text


def test_raw_xml_fallback_reader_normalizes_the_same_way():
    """openpyxl 打不开（样式表坏）时退到 zipfile 直读：数值格在 XML 里是 `179159.09499999997` 这个串，
    要与 openpyxl 那条路得到同一段正文，否则同一本工作簿换个读法数就变了。"""
    raw = _read_xlsx_raw(_book())
    rows = {rno: dict(cells) for rno, cells in raw.sheets[0].rows}
    assert rows[9]["C"] == "179159.095"
    assert rows[10]["C"] == "179159.09499999997"   # 共享字符串（文本格）不动
    assert rows[11]["C"] == "172149"
    assert rows[12]["C"] == "172149.31"
    assert rows[13]["C"] == "0.3"
    assert {r: dict(c) for r, c in read_xlsx(_book()).sheets[0].rows} == rows


def test_raw_reader_keeps_integer_strings_as_written():
    """XML 里没有小数点、没有指数的数值格不经浮点：照 openpyxl 的口径，它本来就是整数，原样写。"""
    content = _patch_sheet(_book(), {b"<v>172149</v>": b"<v>00172149</v>"})
    rows = {rno: dict(cells) for rno, cells in _read_xlsx_raw(content).sheets[0].rows}
    assert rows[11]["C"] == "00172149"
