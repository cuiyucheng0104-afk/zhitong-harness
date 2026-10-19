# 来源：engine-service/tests/test_costsheet_dedup_scope.py（zhitong-cost-ai@991ff5cd）
"""#22 [D-12] 剩余项:解析层去重盲区。

背景:`parse_costsheet` 的去重键 `(code, name, amount)` 原本作用域是**整个工作簿**
(跨表 + 同表)。跨表去重的本意是防"同一份清单被 汇总表+明细表 重复计入合计";
但同表内的完全相同行——**最典型的重复计取形态(同编码同名同金额两行)——也被吃掉**,
根本到不了 tk 的 duplicate 规则(实测:单文件两条相同行 → parse items=1, duplicate=0)。

修法:去重作用域收窄为**只跨工作表**。同一工作表内的相同行如实保留:
- 送审文件写了两行,parsed_total 就该含两行(文件自己的合计行也含两行)——
  "如实呈现送审内容"是解析层的职责,判"重复计取"是检测层(tk duplicate)的职责;
- 跨表去重原语义保留:第二张表里再出现同 (code,name,amount) 行仍然去掉,
  防汇总/重复导出表污染合计。

跑法: cd engine-service && ./.venv/bin/pytest tests/test_costsheet_dedup_scope.py -v
"""
from __future__ import annotations

import io

from openpyxl import Workbook

from zt.costsheet.parser import parse_costsheet

_HEADER = ["项目编码", "项目名称", "项目特征", "计量单位", "工程量", "综合单价", "综合合价"]
_ROW = ["010101001001", "挖一般土方", "三类土", "m3", 100, 10, 1000]


def test_same_sheet_identical_rows_are_kept():
    """同一工作表内两条完全相同的行 → 2 条 items(留给下游 duplicate 规则判)。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "分部分项工程量清单计价表"
    ws.append(_HEADER)
    ws.append(list(_ROW))
    ws.append(list(_ROW))
    buf = io.BytesIO()
    wb.save(buf)
    out = parse_costsheet(buf.getvalue(), "4月进度款.xlsx")
    assert len(out["items"]) == 2
    # 合计如实含两行(送审文件自己的合计也会含两行,对账口径一致)
    assert out["meta"]["parsed_total"] == 2000.0


def test_cross_sheet_identical_rows_still_deduped():
    """跨工作表的同 (code,name,amount) 行仍去重 —— 原"防多表重复计入合计"语义不变。"""
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "分部分项清单1"
    ws1.append(_HEADER)
    ws1.append(list(_ROW))
    ws2 = wb.create_sheet("分部分项清单2")
    ws2.append(_HEADER)
    ws2.append(list(_ROW))
    buf = io.BytesIO()
    wb.save(buf)
    out = parse_costsheet(buf.getvalue(), "送审-计价表.xlsx")
    assert len(out["items"]) == 1
    assert out["meta"]["parsed_total"] == 1000.0
