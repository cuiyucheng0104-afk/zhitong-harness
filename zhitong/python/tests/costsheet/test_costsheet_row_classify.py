# 来源：engine-service/tests/test_costsheet_row_classify.py（zhitong-cost-ai@991ff5cd）
"""WS08-03 行与表分类（issue #724）。

旧版 _is_subtotal 只看名称（注释写着「无编码无单位多为标题行」，代码却直接 return False）；
叶子判定接受 qty=0 的标题行、unit=项 的合计行；跳表名单没有工日/台班/一览表；汇总表与
明细表同时计入。于是：
  · 「S.1脚手架工程（编码：011701）」被当成清单项，还被判为跨单体「重复计取」（QA:tk-08）；
  · 编码列里的「分部分项合计 / 规费」行被当成清单项，R2 总额约翻倍（QA:reconcile-05）；
  · 主要工日一览表、主要机械台班一览表被计入，市政总造价比控制价高 2,253,617.75（QA:indicator-02）；
  · 汇总表和明细表同时计入，R1 两侧各多算 266,375.49（QA:reconcile-04）。

变异验收：分别回退 标题行判定 / 编码列汉字判小计 / 工日台班跳表 / 汇总表二选一 —— 对应用例转红
（汇总表回退后总额翻倍）。
"""
from __future__ import annotations

from zt.tk import engine as tk_engine
from zt.costsheet.parser import parse_costsheet
from tests.costsheet.fixtures import costsheet_shapes as shapes


def test_s1_title_row_is_not_an_item():
    out = parse_costsheet(shapes.s1_title_rows(), "B2投控送审.xlsx")
    names = [i["name"] for i in out["items"]]
    assert not any(n.startswith("S.1") for n in names)
    assert out["meta"]["parsed_total"] == 29000.0
    scaffold = next(i for i in out["items"] if i["code"] == "011701001001")
    assert scaffold["section"].startswith("S.1脚手架工程"), "标题行记作后续行的 section"


def test_tk_no_longer_reports_s1_title_as_duplicate():
    res = tk_engine.run([(shapes.s1_title_rows(), "B2投控送审.xlsx")], None, project_id="P1")
    dup_names = {a["item_name"] for a in res["anomalies"] if a["subtype"] == "duplicate"}
    assert not any(n.startswith("S.1") for n in dup_names)


def test_code_column_subtotals_are_not_items_and_total_equals_leaf_sum():
    out = parse_costsheet(shapes.zhesi_code_column_totals(), "浙四第25期.xlsx", side="submit")
    codes = [i["code"] for i in out["items"]]
    for bad in ("分部分项合计", "安全文明施工费", "施工组织措施费", "规费", "税金"):
        assert bad not in codes
    assert codes == ["010101002001", "010501001001", "010502001001"]
    assert out["meta"]["parsed_total"] == 69000.0
    d = out["meta"]["sheets_detail"][0]
    assert d["total_row_amount"] == 69000.0, "勾稽用编码列里的「分部分项合计」，不用含规费税金的「合计」"


def test_code_column_fee_rows_excluded_even_when_they_look_like_items():
    """合同价口径下「安全文明施工费」行有名称、单位「项」、工程量 1、金额 6000 —— 形状上就是一条无编码
    清单项。只有「编码列含汉字 → 小计」这条能把它挡住（浙四原件就是把费用名写在编码列里）。"""
    out = parse_costsheet(shapes.zhesi_code_column_totals(), "浙四第25期.xlsx", basis="contract")
    assert [i["code"] for i in out["items"]] == ["010101002001", "010501001001", "010502001001"]
    assert out["meta"]["parsed_total"] == 206000.0
    assert out["meta"]["sheets_detail"][0]["total_row_amount"] == 206000.0


def test_labor_and_machine_tables_are_skipped_total_equals_bill_total():
    out = parse_costsheet(shapes.labor_machine_tables(), "市政控制价.xlsx")
    assert out["meta"]["parsed_total"] == 8000.0
    names = {i["name"] for i in out["items"]}
    assert not names & {"二类人工", "机械人工", "履带式挖掘机", "柴油"}
    status = {d["sheet"]: d["status"] for d in out["meta"]["sheets_detail"]}
    assert status == {"分部分项清单与计价表": "used", "主要工日一览表": "skipped_by_title",
                      "主要机械台班一览表": "skipped_by_title", "表-09": "skipped_by_title"}


def test_summary_plus_detail_counts_detail_only():
    out = parse_costsheet(shapes.door_three_block(), "门窗.xlsx", side="submit")
    assert out["meta"]["parsed_total"] == 158400.0, "汇总表不计入（计入则翻倍为 316800）"
    status = {d["sheet"]: d["status"] for d in out["meta"]["sheets_detail"]}
    assert status["汇总"] == "summary"
    assert "summary_detail_mismatch" not in [w["code"] for w in out["warning_items"]]


def test_door_group_rows_are_not_items():
    out = parse_costsheet(shapes.door_three_block(), "门窗.xlsx", side="audit")
    names = [i["name"] for i in out["items"]]
    assert "铝合金门窗" not in names and "防火门" not in names
    assert [i["code"] for i in out["items"]] == ["LC2020", "LM1521", "FM1021", "LC2020", "LM1521"]
    assert out["items"][0]["section"] == "一、 铝合金门窗"


def test_summary_only_workbook_is_counted_with_info():
    """工作簿里只有汇总表 → 照常计入，出 summary_only（info）。"""
    from openpyxl import load_workbook
    import io
    wb = load_workbook(io.BytesIO(shapes.door_three_block()))
    del wb["6#楼"]
    del wb["7#楼"]
    buf = io.BytesIO()
    wb.save(buf)
    out = parse_costsheet(buf.getvalue(), "门窗汇总.xlsx", side="submit")
    assert out["meta"]["parsed_total"] == 158400.0
    info = [w for w in out["warning_items"] if w["code"] == "summary_only"]
    assert info and info[0]["level"] == "info"


def test_summary_detail_mismatch_is_flagged():
    """汇总表只用来勾稽：与明细之和偏差 >0.5% 出 summary_detail_mismatch（error）。"""
    from openpyxl import load_workbook
    import io
    wb = load_workbook(io.BytesIO(shapes.door_three_block()))
    wb["汇总"]["D3"] = 209400      # 6#楼 送审金额改大 100000
    buf = io.BytesIO()
    wb.save(buf)
    out = parse_costsheet(buf.getvalue(), "门窗.xlsx", side="submit")
    assert out["meta"]["parsed_total"] == 158400.0
    assert "summary_detail_mismatch" in [w["code"] for w in out["warning_items"]]


def test_seven_lump_sum_measure_items_are_counted():
    out = parse_costsheet(shapes.measure_lump_sum(), "R1措施费.xlsx")
    assert [i["row_kind"] for i in out["items"]] == ["lump_sum"] * 7
    assert out["meta"]["parsed_total"] == 24000.0
    assert out["meta"]["sheets_detail"][0]["total_row_amount"] == 24000.0


def test_cross_sheet_dedup_is_disclosed():
    """跨表去重保留（防同一份清单重复导出），但去了几条要说出来。"""
    from openpyxl import Workbook
    import io
    wb = Workbook()
    for i, title in enumerate(("分部分项清单1", "分部分项清单2")):
        ws = wb.active if i == 0 else wb.create_sheet(title)
        ws.title = title
        ws.append(["项目编码", "项目名称", "项目特征", "计量单位", "工程量", "综合单价", "综合合价"])
        ws.append(["010101001001", "挖一般土方", "三类土", "m3", 100, 10, 1000])
    buf = io.BytesIO()
    wb.save(buf)
    out = parse_costsheet(buf.getvalue(), "x.xlsx")
    assert out["meta"]["parsed_total"] == 1000.0
    w = [w for w in out["warning_items"] if w["code"] == "cross_sheet_dedup"]
    assert w and "1 行" in w[0]["message"]


def test_column_number_row_under_header_is_not_an_item():
    """表头下一行「1|2|3|4|5|6|7」列号行：没有编码、单位是数字 —— 不是清单项。"""
    from openpyxl import Workbook
    import io
    wb = Workbook()
    ws = wb.active
    ws.title = "分部分项"
    ws.append(["序号", "项目编码", "项目名称", "计量单位", "工程量", "综合单价", "合价"])
    ws.append([1, 2, 3, 4, 5, 6, 7])
    ws.append([1, "010101001001", "平整场地", "m2", 100, 2, 200])
    buf = io.BytesIO()
    wb.save(buf)
    out = parse_costsheet(buf.getvalue(), "x.xlsx")
    assert [i["code"] for i in out["items"]] == ["010101001001"]
