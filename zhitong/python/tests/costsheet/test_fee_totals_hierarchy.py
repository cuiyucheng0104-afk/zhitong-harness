# 来源：engine-service/tests/test_fee_totals_hierarchy.py（zhitong-cost-ai@991ff5cd）
"""WS26-h（issue #831，实测缺陷 07-012）：费用表合计分层取数与采用规则。

修的现场：指标计算只勾招标控制价工作簿（1 张项目级费用表 + 15 张单项 + 113 张单位工程费用表），
「总造价」卡片显示 185,405,919.90，控制价是 318,375,872.00 —— 只有 58.2%，单方低四成。
旧读取器在这份文件上命中数是 0：
  · 单位 / 单项工程表的「招标控制价合计」金额格是**文本**（带千分位），旧版只取数值格；
  · 项目级表的合计行叫泛名「合计」，不在白名单；
  · 就算认了，「取最靠右的数值格」拿到的是规费列；128 个命中也过不了「恰好 1 个」。

夹具 shapes.fee_three_tier 照生产原件的形状脱敏造（CLAUDE.md §5.2：夹具要够得到会错的分支）。
期望值全部由夹具构造手算，不从解析器输出反抄。

评审补的一类（同一个 PR 第一版引入的错数）：广联达另一种导出版式是 1 张「单位工程…费用表」下面挂
若干张「专业工程…费用表」，专业工程表的抬头写的也是「单位工程名称：…」。第一版只要标题区里同时有
「单位工程」和「费用表」就判 unit，于是专业工程表与它上面那张单位工程表相加，总造价正好翻倍
（样本库 78 份文件）。fee_three_tier 没有专业工程表，够不到这条分支；补 shapes.fee_unit_with_specialty。
"""
from __future__ import annotations

import pytest

from zt.costsheet import parser as cs
from zt.costsheet.parser import parse_costsheet, resolve_fee_total, scan_fee_sheet
from tests.costsheet.fixtures import costsheet_shapes as shapes

SPEC_A, SPEC_B, SPEC_U = 6071300.0, 1445340.82, 7516640.82    # fee_unit_with_specialty：两张专业工程表、单位工程表
P = 15873168.0            # 项目级合计（取整到元）
U = 15873167.65           # 15,869,272.21 + 3,895.44
BIG, SMALL = 15869272.21, 3895.44


def _meta(raw: bytes) -> dict:
    return parse_costsheet(raw, "控制价.xlsx", side=None, basis="contract", require_gb_code=True)["meta"]


# ---------------------------------------------------------------- 三层：项目级 + 单位工程

def test_project_total_adopted_when_within_one_yuan_of_unit_sum():
    """项目级合计（整数）比 2 张单位工程表之和大 0.35 → 采用项目级合计，写出 P、U、差额三个数。"""
    meta = _meta(shapes.fee_three_tier())
    chk = meta["fee_total_check"]
    assert meta["fee_total"] == P
    assert chk["source"] == "project" and chk["reason"] is None
    assert (chk["project_total"], chk["unit_sum"], chk["unit_count"], chk["diff"]) == (P, U, 2, 0.35)
    note = chk["notes"][0]
    assert "15,873,168.00" in note and "2 张单位工程费用表之和 15,873,167.65" in note and "差 0.35 元" in note
    # 采用的那一格可回链（K-R44：指标、字段提取候选、报告金额事实拿同一格对）
    assert (chk["hit"]["sheet"], chk["hit"]["row"], chk["hit"]["label"]) == ("表10.2.2-12 招标控制价费用表", 9, "合计")


def test_amount_comes_from_the_amount_header_column_not_the_rightmost_cell():
    """合计行右边还有暂估价 100,000、安文费 852,285.06、规费 721,892.55（数值）和税金（文本）四格。
    取的必须是「金额(元)」那一列，不是其中任何一个。"""
    hits = [h for h in _meta(shapes.fee_three_tier())["fee_totals"] if h["tier"] == "project"]
    assert [h["amount"] for h in hits] == [P]
    assert hits[0]["amount"] not in (100000, 852285.06, 721892.55, 1210689.87)


def test_text_amounts_with_thousands_separators_are_read():
    """单位工程表合计是文本「15,869,272.21」，左边一格是公式文字「1+2+3+4+5」。"""
    units = [h for h in _meta(shapes.fee_three_tier())["fee_totals"] if h["tier"] == "unit"]
    assert sorted(h["amount"] for h in units) == [SMALL, BIG]
    assert {h["label"] for h in units} == {"招标控制价合计"}, "干扰行「其中：安全文明施工费合计」被取了"
    assert len(units) == 2, "每张单位工程表恰 1 个命中（公式文字、干扰行都不算）"


def test_project_total_two_yuan_off_is_not_adopted():
    meta = _meta(shapes.fee_three_tier(project_total=U + 2))
    chk = meta["fee_total_check"]
    assert meta["fee_total"] is None and chk["source"] is None
    assert chk["has_fee_sheet"] is True
    for num in ("15,873,169.65", "15,873,167.65", "2.00"):
        assert num in chk["reason"], chk["reason"]


def test_missing_one_small_unit_sheet_is_not_adopted():
    """少一张 3,895.44 的单位工程表：P − U = 3,895.79。**这条钉的就是容差**——按 0.5%（约 7.9 万元）
    它会通过，然后界面写「勾稽一致」。"""
    meta = _meta(shapes.fee_three_tier(with_small_unit=False))
    chk = meta["fee_total_check"]
    assert meta["fee_total"] is None
    assert chk["unit_count"] == 1 and chk["unit_sum"] == BIG and chk["diff"] == 3895.79
    assert "3,895.79" in chk["reason"]
    # 负控：这个差额确实落在旧的 0.5% 容差之内，否则本用例钉不住容差
    assert 3895.79 < 0.005 * P


@pytest.mark.parametrize("delta,adopted", [(1.0, True), (-1.0, True), (1.01, False), (-1.01, False)])
def test_tolerance_is_one_yuan_inclusive(delta, adopted):
    meta = _meta(shapes.fee_three_tier(project_total=round(U + delta, 2)))
    assert (meta["fee_total"] is not None) is adopted


def test_generic_total_alone_is_not_adopted():
    """项目级表只有泛名「合计」、没有任何单位工程表 → 不采用。泛名不单独成立，也不走「唯一命中」。"""
    meta = _meta(shapes.fee_three_tier(with_units=False))
    chk = meta["fee_total_check"]
    assert [h["generic"] for h in meta["fee_totals"]] == [True], "夹具失效：该有 1 个泛名命中"
    assert meta["fee_total"] is None and chk["has_fee_sheet"] is True
    assert "泛名" in chk["reason"]


# ---------------------------------------------------------------- 没有项目级表 / 单项工程表

def test_unit_sum_adopted_when_there_is_no_project_sheet():
    meta = _meta(shapes.fee_three_tier(project_total=None))
    chk = meta["fee_total_check"]
    assert meta["fee_total"] == U
    assert chk["source"] == "unit_sum" and chk["unit_count"] == 2 and chk["hit"] is None
    assert chk["notes"] == ["2 张单位工程费用表之和 15,873,167.65"]


def test_single_layer_sheet_does_not_change_the_number():
    """再加 1 张单项工程表（合计 = 两张单位工程表之和）：结果不变，提示行多一句。
    单项与单位是包含关系 —— 相加会翻倍。"""
    base = _meta(shapes.fee_three_tier())
    meta = _meta(shapes.fee_three_tier(with_single=True))
    chk = meta["fee_total_check"]
    assert meta["fee_total"] == base["fee_total"] == P
    assert (chk["unit_sum"], chk["single_sum"], chk["single_count"]) == (U, U, 1)
    assert len(chk["notes"]) == len(base["fee_total_check"]["notes"]) + 1
    assert "1 张单项工程费用表之和 15,873,167.65" in chk["notes"][-1]


def test_single_layer_is_not_added_to_unit_sum_without_project_sheet():
    meta = _meta(shapes.fee_three_tier(project_total=None, with_single=True))
    assert meta["fee_total"] == U, "单项与单位相加了（翻倍）"
    assert meta["fee_total_check"]["source"] == "unit_sum"


# ---------------------------------------------------------------- 单位工程表 + 专业工程表（评审）

def test_specialty_sheets_are_not_added_to_the_unit_sheet():
    """1 张单位工程表（7,516,640.82）+ 2 张专业工程表（6,071,300.00 + 1,445,340.82，抬头「单位工程名称：…」）。
    专业 ⊂ 单位：总造价是单位工程那一张的数，不是 15,033,281.64。"""
    meta = _meta(shapes.fee_unit_with_specialty())
    chk = meta["fee_total_check"]
    by_tier = sorted((h["tier"], h["amount"]) for h in meta["fee_totals"])
    assert by_tier == [("specialty", SPEC_B), ("specialty", SPEC_A), ("unit", SPEC_U)], \
        "夹具或分层失效：该是 1 张 unit + 2 张 specialty，其他项目汇总表的「合计」不该进来"
    assert meta["fee_total"] == SPEC_U, "15,033,281.64 = 专业工程表并进了单位层（翻倍）"
    assert chk["source"] == "unit_sum" and (chk["unit_sum"], chk["unit_count"]) == (SPEC_U, 1)
    assert (chk["specialty_sum"], chk["specialty_count"]) == (SPEC_U, 2)
    assert chk["notes"] == ["1 张单位工程费用表之和 7,516,640.82",
                            "2 张专业工程费用表之和 7,516,640.82（只用于核对，不参与取数）"]


def test_fixture_specialty_sheet_really_carries_the_unit_label():
    """负控：这条分支成立的前提是专业工程表的标题区里同时有「单位工程」和「费用表」（来自抬头标签）。
    夹具要是没有这两个词，上一条在「标题区含单位工程就算 unit」的旧判法下也会过，钉不住。"""
    from openpyxl import load_workbook
    import io
    wb = load_workbook(io.BytesIO(shapes.fee_unit_with_specialty()), read_only=True)
    rows = [list(r) for r in wb["2.表10.2.2-13专业工程招标控制价费用表"].iter_rows(values_only=True)]
    head = cs._title_text(rows, 3)
    assert "单位工程" in head and "费用表" in head and rows[3][3] == "金额（元）"
    assert rows[1][0] == "专业工程招标控制价费用表" and rows[2][0].startswith("单位工程名称：")


def test_project_total_is_adopted_when_unit_sheet_has_specialty_sheets_under_it():
    """再加项目级表（泛名「合计」7,516,641）：P 与**单位层** 7,516,640.82 核对，差 0.18 → 采用 P。
    专业工程表并进单位层的话 U = 15,033,281.64 = 2P，会报「对不上」而回落（样本库 9 份）。"""
    meta = _meta(shapes.fee_unit_with_specialty(with_project=True))
    chk = meta["fee_total_check"]
    assert meta["fee_total"] == 7516641.0 and chk["source"] == "project"
    assert (chk["project_total"], chk["unit_sum"], chk["unit_count"], chk["diff"]) == (7516641.0, SPEC_U, 1, 0.18)
    assert chk["notes"] == ["项目级合计 7,516,641.00，1 张单位工程费用表之和 7,516,640.82，差 0.18 元",
                            "2 张专业工程费用表之和 7,516,640.82（只用于核对，不参与取数）"]
    assert (chk["hit"]["sheet"], chk["hit"]["label"]) == ("1.表10.2.2-12招标控制价费用表", "合计")


@pytest.mark.parametrize("b_total,adopted", [("1445341.82", True), ("1445339.82", True),
                                             ("1445341.83", False), ("1445339.81", False)])
def test_unit_sum_needs_the_specialty_sum_within_one_yuan(b_total, adopted):
    """单位层与专业层都算得出：相差 ≤ 1 元才采用单位工程之和（±1.00 采用，±1.01 不采用）。"""
    meta = _meta(shapes.fee_unit_with_specialty(b_total=b_total))
    assert (meta["fee_total"] == SPEC_U) is adopted and (meta["fee_total"] is None) is not adopted


def test_unit_and_specialty_sums_three_yuan_apart_are_not_adopted_and_both_numbers_are_written():
    meta = _meta(shapes.fee_unit_with_specialty(b_total="1445343.82"))
    chk = meta["fee_total_check"]
    assert meta["fee_total"] is None and chk["source"] is None and chk["has_fee_sheet"] is True
    assert (chk["unit_sum"], chk["specialty_sum"]) == (SPEC_U, 7516643.82)
    for part in ("1 张单位工程费用表之和 7,516,640.82", "2 张专业工程费用表之和 7,516,643.82", "差 3.00 元"):
        assert part in chk["reason"], chk["reason"]


def test_only_specialty_sheets_are_not_adopted():
    """没有单位工程表、没有项目级表，只有 2 张专业工程表：不采用（上面那张不在，不知道这两张是不是全部）。
    原因里写出张数与和，用户可以据此手填。"""
    meta = _meta(shapes.fee_unit_with_specialty(with_unit=False))
    chk = meta["fee_total_check"]
    assert meta["fee_total"] is None and chk["source"] is None and chk["has_fee_sheet"] is True
    assert (chk["unit_count"], chk["specialty_sum"], chk["specialty_count"]) == (0, SPEC_U, 2)
    assert "只有 2 张专业工程费用表之和 7,516,640.82" in chk["reason"] and "不知道这几张是不是全部" in chk["reason"]


def test_single_specialty_sheet_is_still_adopted_as_the_only_hit():
    """全文件只有 1 张专业工程表、1 个白名单命中：与旧规则一致（样本库 300 多份单专业文件）。"""
    meta = _meta(shapes.fee_unit_with_specialty(with_unit=False, with_b=False))
    assert meta["fee_total"] == SPEC_A and meta["fee_total_check"]["source"] == "single_hit"


def test_project_total_is_checked_against_specialty_sum_when_there_is_no_unit_sheet():
    """项目级表下面直接是专业工程表（没有任何单位工程表）：核对层退到专业层。"""
    meta = _meta(shapes.fee_unit_with_specialty(with_project=True, with_unit=False))
    chk = meta["fee_total_check"]
    assert meta["fee_total"] == 7516641.0 and chk["source"] == "project" and chk["diff"] == 0.18
    assert chk["notes"] == ["项目级合计 7,516,641.00，2 张专业工程费用表之和 7,516,640.82，差 0.18 元"]
    off = _meta(shapes.fee_unit_with_specialty(with_project=True, with_unit=False, project_total=7516643))
    assert off["fee_total"] is None
    for part in ("项目级合计 7,516,643.00", "2 张专业工程费用表之和 7,516,640.82", "差 2.18 元"):
        assert part in off["fee_total_check"]["reason"]


def test_specialty_sheet_with_the_combined_label_is_not_a_project_sheet():
    """样本库里专业工程表的另一种抬头是「单位及专业工程名称：…」：标题区没有「单位工程」四个字，
    第一版把它判成 project 层（表内标题含「招标控制价费用表」），两张就是「项目级有 2 个合计行」。"""
    meta = _meta(shapes.fee_unit_with_specialty(label="单位及专业工程名称"))
    assert sorted(h["tier"] for h in meta["fee_totals"]) == ["specialty", "specialty", "unit"]
    assert meta["fee_total"] == SPEC_U and meta["fee_total_check"]["source"] == "unit_sum"


# ---------------------------------------------------------------- 表-04 + 表-12

def test_table_04_wins_and_table_12_is_ignored():
    """表-12「其他项目清单与计价汇总表」表头也是「金额(元)」、末行也叫「合计」，工程名称里还有
    「建设项目」—— 不排除它，它就成了项目级汇总表，20,000 成了总造价候选。"""
    meta = _meta(shapes.fee_table_04_with_12())
    assert meta["fee_total"] == 172220.0
    assert [(h["sheet"], h["label"], h["amount"]) for h in meta["fee_totals"]] == \
        [("表-04 单位工程招标控制价汇总表", "招标控制价合计", 172220.0)]


def test_table_12_alone_gives_no_hit():
    meta = _meta(shapes.fee_table_04_with_12(with_04=False))
    assert meta["fee_totals"] == [] and meta["fee_total"] is None
    assert meta["fee_total_check"]["has_fee_sheet"] is False, "其他项目汇总表不是造价汇总表"


# ---------------------------------------------------------------- 旧金标不变

def test_single_whitelist_hit_is_still_adopted():
    """表-13 形（单张单位工程费用表、数值金额）：与旧规则一致，8,715,564.88。"""
    meta = _meta(shapes.fee_table_13())
    assert meta["fee_total"] == 8715564.88
    assert meta["fee_total_check"]["source"] == "single_hit"
    assert [h["label"] for h in meta["fee_totals"]] == ["招标控制价合计"]


# ---------------------------------------------------------------- 一张表内的取数（scan_fee_sheet）

def _unit_rows(total_row: list, header: list | None = None) -> list:
    rows = [["单位工程招标控制价费用表"], ["工程名称：某楼"]]
    if header:
        rows.append(header)
    rows += [["1", "分部分项工程费", "∑", 100.0], total_row]
    return rows


def test_tier_is_read_from_title_rows_not_from_header_or_body():
    """项目级表的表头第二行就写着「暂估价 | 规费 | 税金」，单位工程表正文里有「措施项目费」——
    排除词只比表名与表头上方的标题行，否则三层表全被排除。"""
    rows = [["招标控制价费用表"], ["工程名称：某项目"],
            ["序号", "工程名称", "金额(元)", "其中：（元）", None, None],
            [None, None, None, "暂估价", "规费", "税金"],
            ["1", "地下室", 100.0, 1.0, 2.0, 3.0], ["合计", None, 100.0, 1.0, 2.0, 3.0]]
    sc = scan_fee_sheet("表10.2.2-12 招标控制价费用表", rows)
    assert sc["tier"] == cs.FEE_PROJECT and [h["amount"] for h in sc["hits"]] == [100.0]
    sc = scan_fee_sheet("x", _unit_rows(["招标控制价合计", None, "1+2", 100.0],
                                        ["序号", "费用名称", "计算公式", "金额（元）"]))
    assert sc["tier"] == cs.FEE_UNIT and len(sc["hits"]) == 1


def _tier(sheet: str, title_cell: str | None, label: str = "工程名称：某楼") -> str | None:
    rows = ([[title_cell]] if title_cell else []) + [[label, None, "标段：", "第 1 页 共 1 页"],
                                                      ["序号", "费用名称", "计算公式", "金额（元）"],
                                                      ["招标控制价合计", None, "1+2", 100.0]]
    return scan_fee_sheet(sheet, rows)["tier"]


@pytest.mark.parametrize("sheet,title_cell,label,tier", [
    # 表内标题格定层；抬头标签写什么都不算
    ("x", "专业工程招标控制价费用表", "单位工程名称：一期02标-1#-地上工程", cs.FEE_SPECIALTY),
    ("x", "专业工程投标报价费用表", "单位及专业工程名称：市政工程 - 交通标志标线", cs.FEE_SPECIALTY),
    ("x", "单位工程招标控制价费用表", "单位工程名称：一期02标", cs.FEE_UNIT),
    ("x", "单项工程招标控制价费用表", "工程名称：地下室", cs.FEE_SINGLE),
    ("x", "招标控制价费用表", "工程名称：某项目", cs.FEE_PROJECT),
    # 浙江表10.2.2-13 的原名「单位（专业）工程…费用表」（全角括号）：不分两级的导出件，归 unit，不是 project
    ("表10_2_2-13 【不含单列】单位(专业)工程投标报", "单位（专业）工程投标报价费用表", "工程名称:净化-净化装修", cs.FEE_UNIT),
    # 表内标题与表名不一致时听表内标题的（样本库「12变更-洽商06.xlsx」）
    ("1.表10.2.2-13专业工程投标报价费用表", "单位工程投标报价费用表", "工程名称：变更-洽商06", cs.FEE_UNIT),
    # 标题区没有分层短语才看表名
    ("2.表10.2.2-13专业工程招标控制价费用表", None, "单位工程名称：一期02标-1#-地上工程", cs.FEE_SPECIALTY),
    ("表10.2.2-13 单位工程投标报价费用表", None, "工程名称：某楼", cs.FEE_UNIT),
    # 抬头标签冒号后面的工程名里碰巧有分层短语：那是工程名，不是表式
    ("x", "招标控制价费用表", "工程名称：某小区幕墙专业工程费用表复核", cs.FEE_PROJECT),
    # 只有抬头标签、没有任何分层短语：不靠标签定层
    ("费用表", None, "单位工程名称：某楼", None),
    ("汇总表", "工程造价汇总表", "专业工程名称：某楼", None),
])
def test_tier_comes_from_the_form_title_phrase_not_from_the_header_label(sheet, title_cell, label, tier):
    assert _tier(sheet, title_cell, label) == tier


def test_two_different_tier_phrases_in_the_title_rows_give_no_tier():
    """标题区同时写着「单位工程…费用表」和「专业工程…费用表」：判不出，按未分层处理，不猜。"""
    rows = [["单位工程招标控制价费用表"], ["专业工程招标控制价费用表"], ["工程名称：某楼"],
            ["序号", "费用名称", "计算公式", "金额（元）"], ["招标控制价合计", None, "1+2", 100.0]]
    sc = scan_fee_sheet("x", rows)
    assert sc["tier"] is None and [h["amount"] for h in sc["hits"]] == [100.0]


@pytest.mark.parametrize("title", ["措施项目费用表", "规费、税金项目计价表", "专业工程暂估价及结算价表",
                                   "暂列金额明细表", "计日工表", "总承包服务费计价表"])
def test_excluded_sheets_give_no_hit_even_with_a_whitelisted_label(title):
    rows = [[f"单位工程{title}（汇总表）"], ["序号", "名称", "金额(元)"], ["1", "某项", 5.0],
            [None, "招标控制价合计", 5.0]]
    assert scan_fee_sheet(title, rows)["hits"] == []
    # 负控：同样的行换成不带排除词的表名就取得到（证明上面是被排除词挡住的）
    rows[0] = ["单位工程招标控制价汇总表"]
    assert [h["amount"] for h in scan_fee_sheet("表-04", rows)["hits"]] == [5.0]


def test_generic_total_is_only_recognised_in_the_project_layer():
    rows = _unit_rows(["合计", None, None, 100.0], ["序号", "费用名称", "计算公式", "金额（元）"])
    assert scan_fee_sheet("单位工程招标控制价费用表", rows)["hits"] == []


def test_empty_amount_cell_gives_no_hit_instead_of_falling_back_to_another_cell():
    """表头定了金额列，那一格是空的 → 不出命中；不拿同一行别的数顶上。"""
    rows = _unit_rows(["招标控制价合计", None, "1+2+3+4+5", None, 999.0],
                      ["序号", "费用名称", "计算公式", "金额（元）", "备注"])
    assert scan_fee_sheet("单位工程招标控制价费用表", rows)["hits"] == []


def test_formula_text_is_never_parsed_as_a_number():
    """没有表头时也一样：「1+2+3+4+5」不是数值格。右侧恰 1 个数值格才取。"""
    sc = scan_fee_sheet("单位工程招标控制价费用表", _unit_rows(["招标控制价合计", None, "1+2+3+4+5", "8,000.50"]))
    assert [h["amount"] for h in sc["hits"]] == [8000.5]


def test_without_header_two_numeric_cells_give_no_hit():
    sc = scan_fee_sheet("单位工程招标控制价费用表", _unit_rows(["招标控制价合计", None, 8000.5, 300.0]))
    assert sc["hits"] == [], "找不到表头、右侧有 2 个数值格：不知道哪个是合计，不猜"


def test_two_amount_headers_give_no_hit():
    rows = _unit_rows(["招标控制价合计", None, 8000.5, 7000.0], ["序号", "费用名称", "金额(元)", "金额(元)"])
    assert scan_fee_sheet("单位工程招标控制价费用表", rows)["hits"] == []


def test_non_fee_sheet_is_untouched():
    """清单表里的「合   计」行不是费用表合计：不是费用表 / 汇总表直接回空。"""
    rows = [["分部分项工程清单与计价表"], ["序号", "项目编码", "项目名称", "金额（元）"],
            [1, "010101002001", "挖一般土方", 100.0], ["合   计", None, None, 100.0]]
    assert scan_fee_sheet("表10.2.2-16 分部分项工程清单与计价表【土石方工程】", rows) == \
        {"sheet": "表10.2.2-16 分部分项工程清单与计价表【土石方工程】", "tier": None, "hits": []}


# ---------------------------------------------------------------- 采用规则（resolve_fee_total）

def _scan(sheet: str, tier, *amounts, label="招标控制价合计", generic=False) -> dict:
    return {"sheet": sheet, "tier": tier,
            "hits": [{"sheet": sheet, "row": 9, "label": label, "amount": a, "tier": tier, "generic": generic}
                     for a in amounts]}


def test_unit_layer_is_unusable_when_any_sheet_has_no_hit_or_two_hits():
    proj = _scan("项目", cs.FEE_PROJECT, 300.0, label="合计", generic=True)
    ok = resolve_fee_total([proj, _scan("u1", cs.FEE_UNIT, 100.0), _scan("u2", cs.FEE_UNIT, 200.0)])
    assert ok["amount"] == 300.0
    none = resolve_fee_total([proj, _scan("u1", cs.FEE_UNIT, 300.0), _scan("u2", cs.FEE_UNIT)])
    assert none["amount"] is None, "有一张单位工程表读不到合计：剩下的和不是和，不能拿来勾稽"
    two = resolve_fee_total([proj, _scan("u1", cs.FEE_UNIT, 100.0, 200.0)])
    assert two["amount"] is None


def test_unit_layer_needs_the_same_label_on_every_sheet():
    r = resolve_fee_total([_scan("u1", cs.FEE_UNIT, 100.0), _scan("u2", cs.FEE_UNIT, 200.0, label="投标报价合计")])
    assert r["amount"] is None and r["unit_sum"] is None


def test_unit_sum_is_rejected_when_single_layer_disagrees():
    r = resolve_fee_total([_scan("u1", cs.FEE_UNIT, 100.0), _scan("u2", cs.FEE_UNIT, 200.0),
                           _scan("s1", cs.FEE_SINGLE, 250.0)])
    assert r["amount"] is None and "300.00" in r["reason"] and "250.00" in r["reason"]


def test_unit_sum_is_adopted_when_single_layer_sum_cannot_be_computed():
    """GB50500 表-03（单项工程汇总表）的末行叫泛名「合计」，single 层不认泛名 → S 算不出。
    这不该挡住单位工程之和（spec §3.8：S 不存在或 |S − U| ≤ 1 元）。"""
    r = resolve_fee_total([_scan("u1", cs.FEE_UNIT, 100.0), _scan("u2", cs.FEE_UNIT, 200.0),
                           _scan("s1", cs.FEE_SINGLE)])
    assert (r["amount"], r["source"], r["single_sum"]) == (300.0, "unit_sum", None)
    assert r["notes"] == ["2 张单位工程费用表之和 300.00"]


def test_specialty_layer_is_never_added_to_the_unit_layer():
    r = resolve_fee_total([_scan("u1", cs.FEE_UNIT, 300.0), _scan("z1", cs.FEE_SPECIALTY, 100.0),
                           _scan("z2", cs.FEE_SPECIALTY, 200.0)])
    assert (r["amount"], r["source"], r["unit_count"], r["specialty_count"]) == (300.0, "unit_sum", 1, 2)


def test_unit_sum_is_adopted_when_specialty_layer_sum_cannot_be_computed():
    """有一张专业工程表读不到合计：专业层的和算不出，核对不了，但不挡单位工程之和（与单项层同样对待）。"""
    r = resolve_fee_total([_scan("u1", cs.FEE_UNIT, 300.0), _scan("z1", cs.FEE_SPECIALTY, 100.0),
                           _scan("z2", cs.FEE_SPECIALTY)])
    assert (r["amount"], r["source"], r["specialty_sum"], r["specialty_count"]) == (300.0, "unit_sum", None, 2)
    assert r["notes"] == ["1 张单位工程费用表之和 300.00"]


def test_project_total_is_not_checked_against_specialty_when_unit_sheets_exist_but_are_unreadable():
    """有单位工程表但读不全：核对层仍是单位层（读不全 = 不知道），不拿专业层顶替。"""
    proj = _scan("项目", cs.FEE_PROJECT, 300.0, label="合计", generic=True)
    r = resolve_fee_total([proj, _scan("u1", cs.FEE_UNIT), _scan("z1", cs.FEE_SPECIALTY, 100.0),
                           _scan("z2", cs.FEE_SPECIALTY, 200.0)])
    assert r["amount"] is None and "单位工程费用表的合计读不全" in r["reason"]
    # 负控：去掉那张读不到合计的单位工程表，P 与专业层核对通过
    ok = resolve_fee_total([proj, _scan("z1", cs.FEE_SPECIALTY, 100.0), _scan("z2", cs.FEE_SPECIALTY, 200.0)])
    assert (ok["amount"], ok["source"]) == (300.0, "project")


def test_project_total_adopted_against_unit_sum_even_if_specialty_sum_differs():
    """P 与 U 对上了就采用 P；专业层的和只写进提示行（有的单位工程下面没有专业工程表，Z < U 是正常的）。"""
    proj = _scan("项目", cs.FEE_PROJECT, 300.0, label="合计", generic=True)
    r = resolve_fee_total([proj, _scan("u1", cs.FEE_UNIT, 100.0), _scan("u2", cs.FEE_UNIT, 200.0),
                           _scan("z1", cs.FEE_SPECIALTY, 60.0), _scan("z2", cs.FEE_SPECIALTY, 40.0)])
    assert (r["amount"], r["source"]) == (300.0, "project")
    assert r["notes"][-1] == "2 张专业工程费用表之和 100.00（只用于核对，不参与取数）"


def test_single_hit_rule_counts_unit_and_specialty_sheets_together():
    """1 张读到合计的专业工程表 + 1 张读不到合计的单位工程表：不是「全文件唯一命中」，是上面那张没读到。"""
    r = resolve_fee_total([_scan("u1", cs.FEE_UNIT), _scan("z1", cs.FEE_SPECIALTY, 100.0)])
    assert r["amount"] is None and "只读到其中一张" in r["reason"]
    assert resolve_fee_total([_scan("z1", cs.FEE_SPECIALTY, 100.0)])["amount"] == 100.0


def test_single_hit_rule_needs_its_layer_to_have_no_other_sheet():
    """两张单位工程表只读到一张的合计：不是「全文件唯一命中」，是另一张没读到。"""
    r = resolve_fee_total([_scan("u1", cs.FEE_UNIT, 100.0), _scan("u2", cs.FEE_UNIT)])
    assert r["amount"] is None and r["has_fee_sheet"] is True and "只读到其中一张" in r["reason"]
    # 负控：只有这一张单位工程表时照旧采用
    assert resolve_fee_total([_scan("u1", cs.FEE_UNIT, 100.0)])["amount"] == 100.0


def test_two_project_totals_are_not_adopted():
    r = resolve_fee_total([_scan("项目", cs.FEE_PROJECT, 300.0, 300.0, label="合计", generic=True),
                           _scan("u1", cs.FEE_UNIT, 300.0)])
    assert r["amount"] is None


def test_zero_total_is_not_adopted():
    r = resolve_fee_total([_scan("u1", cs.FEE_UNIT, 0.0)])
    assert r["amount"] is None and r["has_fee_sheet"] is True


def test_no_fee_sheet_at_all():
    r = resolve_fee_total([{"sheet": "清单", "tier": None, "hits": []}])
    assert (r["amount"], r["has_fee_sheet"], r["reason"]) == (None, False, None)
