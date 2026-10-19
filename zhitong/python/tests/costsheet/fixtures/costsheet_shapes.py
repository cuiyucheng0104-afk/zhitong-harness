# 来源：engine-service/tests/fixtures/costsheet_shapes.py（zhitong-cost-ai@991ff5cd）
"""真实版式计价表夹具 —— 测试时用 openpyxl 现场生成，不入库任何客户文件（WS08-20，issue #724，D7）。

为什么要有它（CLAUDE.md §5.2「夹具够不到失败分支」）：此前 tests/fixtures 里只有 4 份自造
计价表，全是单 sheet、单行表头、3–6 行、全带编码。它们够不到任何真实版式的失败分支，
所以同一个版本在生产文件上全错（进度款取成投标量、门窗取成原合同块、工日表入账……），
测试却全绿。

这里的每个夹具只保留**版式结构**（表头层数、组标题位置、列顺序、标题/小计行的写法），
名称与金额全部改写，数字按「由夹具构造可直接算出真值」来编 —— 用例里的期望值逐项手写，
不从解析器输出反抄。docstring 注明原型出处（生产 QA 证据里的文件、sheet、行号）。

新增夹具请用 @shape 登记：tests/costsheet/test_fixture_shapes_used.py 保证每个夹具至少被一个用例引用
（防孤儿：造了形状却没人测，等于没造）。
"""
from __future__ import annotations

import io
import re
import zipfile
from typing import Callable

from openpyxl import Workbook

SHAPES: dict[str, Callable] = {}


def shape(fn: Callable) -> Callable:
    SHAPES[fn.__name__] = fn
    return fn


def _save(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def resaved(raw: bytes, marker: str = "另存") -> bytes:
    """同样的单元格、不同的字节：另存一次并改文档属性 —— 「内容相同的另一份文件」。

    对账把 sha256 相同的两侧当成同一份文件（WS26-b，issue #800）。openpyxl 的输出带存盘时间，
    连调两次同一个夹具得到的字节「同一秒内相同、跨秒不同」；要表达「两份不同的文件」必须显式造。"""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(raw))
    wb.properties.title = marker
    out = _save(wb)
    assert out != raw
    return out


# ---------------------------------------------------------------- 1. 浙大三行表头

@shape
def zheda_three_row_header() -> bytes:
    """原型：投控 C1 / 对账 R3「浙大第五期」送审件，sheet「地下室建筑--打印」。
    R4 组标题「投标价(元) … 施工单位」、R5 叶子「工程量|综合单价|合价|其中|上期累计已完工程量|
    本期完成工程量|本期产值|累计完成工程量|累计产值」、R6「人工费|机械费|暂估价」三行表头。
    旧版把投标量当工程量、合同合价当金额 → self_check 701 条、big_amount 416 条（QA:tk-01）。

    真值（basis=current）：
      010401003001 本期 30 × 480 = 14400
      010401004001 本期 0（本期没报）→ 金额 0
      010501001001 本期 60 × 520 = 31200
      010502001001 本期 40 × 600 = 24000
      本期产值合计 69600；合同合价合计 327000；累计产值合计 38400+45000+31200+30000=144600
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "地下室建筑--打印"
    ws.append(["某住宅项目第五期进度款产值申报表"])
    ws.append(["单位工程名称：地下室建筑"])
    ws.append(["工程名称：某住宅项目", None, None, None, None, None, None, None, None, None, None, None,
               None, "第1页 共1页"])
    ws.append(["序号", "项目编码", "项目名称", "项目特征描述", "计量单位", "投标价(元)", None, None, None, None,
               None, "施工单位", None, None, None, None])
    ws.append([None, None, None, None, None, "工程量", "综合单价", "合价", "其中", None, None,
               "上期累计已完工程量", "本期完成工程量", "本期产值", "累计完成工程量", "累计产值"])
    ws.append([None, None, None, None, None, None, None, None, "人工费", "机械费", "暂估价",
               None, None, None, None, None])
    ws.append(["", "", "0104 砌筑工程"])
    ws.append([1, "010401003001", "实心砖墙", "墙厚240", "m3", 200, 480, 96000, 20000, 1000, 0,
               50, 30, 14400, 80, 38400])
    ws.append([2, "010401004001", "多孔砖墙", "墙厚200", "m3", 100, 450, 45000, 9000, 500, 0,
               100, 0, 0, 100, 45000])
    ws.append(["", "", "0105 混凝土及钢筋混凝土工程"])
    ws.append([3, "010501001001", "现浇混凝土基础", "C30", "m3", 300, 520, 156000, 30000, 8000, 0,
               0, 60, 31200, 60, 31200])
    ws.append([4, "010502001001", "矩形柱", "C30", "m3", 50, 600, 30000, 6000, 1500, 0,
               10, 40, 24000, 50, 30000])
    ws.append([None, None, "本页小计", None, None, None, None, 327000, None, None, None,
               None, None, 69600, None, 144600])
    ws.append([None, None, "合计", None, None, None, None, 327000, None, None, None,
               None, None, 69600, None, 144600])
    return _save(wb)


# ---------------------------------------------------------------- 2. 浙大安装两行表头

@shape
def zheda_install_two_row() -> bytes:
    """原型：浙大第五期安装件「给排水--打印」，GB50500 表-08 两行表头：
    R1「…|工程量|金额(元)（合并跨 3 列）」、R2「综合单价|合价|其中：暂估价」。
    真值：031001006001 120×45.5=5460、030404017001 2×1500=3000、030901001001 80×60=4800，合计 13260。
    章：0310 给排水、采暖、燃气工程 / 0304 电气设备安装工程 / 0309 消防工程（旧码表三者都叫「其他」）。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "给排水--打印"
    ws.append(["分部分项工程和单价措施项目清单与计价表"])
    ws.append(["序号", "项目编码", "项目名称", "项目特征描述", "计量单位", "工程量", "金额(元)", None, None])
    ws.append([None, None, None, None, None, None, "综合单价", "合价", "其中：暂估价"])
    ws.append([1, "031001006001", "塑料管", "De110", "m", 120, 45.5, 5460, None])
    ws.append([2, "030404017001", "配电箱", "XL-21", "台", 2, 1500, 3000, None])
    ws.append([3, "030901001001", "水喷淋钢管", "DN100", "m", 80, 60, 4800, None])
    ws.append([None, None, "合计", None, None, None, None, 13260, None])
    return _save(wb)


# ---------------------------------------------------------------- 3. A1 市政：分期工程量 + 一列不分期合价

@shape
def a1_period_qty(mode: str = "current") -> bytes:
    """原型：投控 A1 市政进度款「分部分项」。表头
    「投标工程量|上期完成工程量|本期完成工程|累计完成工程量|综合单价|合价」—— 合价只有一列、
    不写期次，工程量却有四组；「本期完成工程」原件就是这么截的（没有「量」字）。
    旧版取了投标工程量当 qty → self_check 15 条假阳（QA:tk-01）。

    mode="current"：每行 合价 = 本期 × 单价 → 投票出本期。真值 qty=本期，合计 47000：
      040101001001 200×30=6000、040101002001 50×40=2000、040201004001 300×60=18000、
      040501001001 100×150=15000、40504001004（丢了前导 0 的 11 位编码）3×2000=6000
    mode="split"：两行按本期算、两行按累计算（上期都不为 0）→ 本期/累计各 50%，判歧义。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "分部分项"
    ws.append(["序号", "项目编码", "项目名称", "计量单位", "投标工程量", "上期完成工程量", "本期完成工程",
               "累计完成工程量", "综合单价", "合价"])
    if mode == "current":
        ws.append([1, "040101001001", "挖一般土方", "m3", 1000, 100, 200, 300, 30, 6000])
        ws.append([2, "040101002001", "挖沟槽土方", "m3", 500, 0, 50, 50, 40, 2000])
        ws.append([3, "040201004001", "道路基层", "m2", 2000, 500, 300, 800, 60, 18000])
        ws.append([4, "040501001001", "管道铺设", "m", 800, 0, 100, 100, 150, 15000])
        ws.append([5, 40504001004, "检查井", "座", 10, 2, 3, 5, 2000, 6000])
    else:
        ws.append([1, "040101001001", "挖一般土方", "m3", 1000, 100, 200, 300, 30, 6000])     # 本期
        ws.append([2, "040101002001", "挖沟槽土方", "m3", 500, 20, 50, 70, 40, 2000])        # 本期
        ws.append([3, "040201004001", "道路基层", "m2", 2000, 500, 300, 800, 60, 48000])     # 累计
        ws.append([4, "040501001001", "管道铺设", "m", 800, 40, 100, 140, 150, 21000])       # 累计
    return _save(wb)


# ---------------------------------------------------------------- 4. 浙四：编码列合计 + 本月上报/审核

@shape
def zhesi_code_column_totals() -> bytes:
    """原型：对账 R2「浙四第 25 期」sheet「地下」。R3「金额(元)|本月上报产值|本月审核产值|…」、
    R4「综合单价|合价」；编码列里写着「分部分项合计 / 施工组织措施费 / 规费 / 税金」，
    旧版把它们当清单项，R2 总额约翻倍（QA:reconcile-05）。

    真值：side=submit 本月上报 5000+52000+12000=69000；side=audit 本月审核 4800+41600+12000=58400。
    「分部分项合计」行：上报 69000 / 审核 58400（与明细之和相等）。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "地下"
    ws.append(["某项目第25期进度款审核表"])
    ws.append(["工程名称：地下室"])
    ws.append(["序号", "项目编码", "项目名称", "项目特征", "计量单位", "工程量", "金额(元)", None,
               "本月上报产值", "本月审核产值", "备注"])
    ws.append([None, None, None, None, None, None, "综合单价", "合价", None, None, None])
    ws.append([1, "010101002001", "挖一般土方", "三类土", "m3", 1000, 20, 20000, 5000, 4800, None])
    ws.append([2, "010501001001", "现浇混凝土基础", "C30", "m3", 300, 520, 156000, 52000, 41600, None])
    ws.append([3, "010502001001", "矩形柱", "C30", "m3", 50, 600, 30000, 12000, 12000, None])
    ws.append([None, "分部分项合计", None, None, None, None, None, 206000, 69000, 58400, None])
    ws.append([None, "安全文明施工费", "安全文明施工费", None, "项", 1, None, 6000, 2000, 1800, None])
    ws.append([None, "施工组织措施费", None, None, "项", 1, None, 8000, 3000, 2500, None])
    ws.append([None, "规费", None, None, "项", 1, None, 4000, 1000, 900, None])
    ws.append([None, "税金", None, None, "项", 1, None, 19800, 6570, 5598, None])
    ws.append([None, None, "合计", None, None, None, None, 243800, 81570, 69198, None])
    return _save(wb)


# ---------------------------------------------------------------- 5. 门窗三段对比 + 汇总表

def _door_sheet(ws, rows: list) -> None:
    ws.append(["门窗工程审核对比表"])
    ws.append(["序号", "编号", "名称", "单位", "原合同", None, None, "送审", None, None, "审核", None, None,
               "核增", "核减"])
    ws.append([None, None, None, None, "工程量", "单价", "金额", "工程量", "单价", "金额", "工程量", "单价", "金额",
               "金额", "金额"])
    for r in rows:
        ws.append(r)


@shape
def door_three_block() -> bytes:
    """原型：对账 R1「门窗一期」审核件。R2「原合同|送审|审核|核增|核减」三段组标题，R3 每段
    「工程量|单价|金额」；编号是字母数字（LC2020）；「一、铝合金门窗」分组行带单位与分组小计
    （没有编码、有单位有量有额 —— 不认「一、」形态就会被当成一条无编码清单项）；另有「汇总」表。
    旧版审定侧取成「原合同」块、名称与金额来自不同列块，34/41 行差额算错（QA:reconcile-06/08）；
    汇总表与明细表一起计入，两侧各多算 266,375.49（QA:reconcile-04）。

    真值（6#楼 + 7#楼，只计明细）：
      送审 55000+40000+14400 + 33000+16000 = 158400
      审核 50400+40000+11500 + 31500+16000 = 149400
      原合同 50000+40000+12000 + 30000+16000 = 148000
      LC2020（6#楼）side=audit：qty 105、price 480、amount 50400；side=submit：110、500、55000
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "汇总"
    ws.append(["门窗工程审核汇总表"])
    ws.append(["序号", "工程名称", "原合同金额", "送审金额", "审核金额", "核减金额"])
    ws.append([1, "6#楼", 102000, 109400, 101900, 7500])
    ws.append([2, "7#楼", 46000, 49000, 47500, 1500])
    ws.append([None, "合计", 148000, 158400, 149400, 9000])
    _door_sheet(wb.create_sheet("6#楼"), [
        ["一、", None, "铝合金门窗", "m2", 150, None, 90000, 160, None, 95000, 155, None, 90400, 0, 4600],
        [1, "LC2020", "铝合金推拉窗", "m2", 100, 500, 50000, 110, 500, 55000, 105, 480, 50400, 0, 4600],
        [2, "LM1521", "铝合金平开门", "m2", 50, 800, 40000, 50, 800, 40000, 50, 800, 40000, 0, 0],
        ["二、", None, "防火门", "樘", 10, None, 12000, 12, None, 14400, 10, None, 11500, 0, 2900],
        [3, "FM1021", "钢质防火门", "樘", 10, 1200, 12000, 12, 1200, 14400, 10, 1150, 11500, 0, 2900],
        [None, None, "合计", None, None, None, 102000, None, None, 109400, None, None, 101900, 0, 7500],
    ])
    _door_sheet(wb.create_sheet("7#楼"), [
        ["一、", None, "铝合金门窗", "m2", 80, None, 46000, 86, None, 49000, 83, None, 47500, 0, 1500],
        [1, "LC2020", "铝合金推拉窗", "m2", 60, 500, 30000, 66, 500, 33000, 63, 500, 31500, 0, 1500],
        [2, "LM1521", "铝合金平开门", "m2", 20, 800, 16000, 20, 800, 16000, 20, 800, 16000, 0, 0],
        [None, None, "合计", None, None, None, 46000, None, None, 49000, None, None, 47500, 0, 1500],
    ])
    return _save(wb)


# ---------------------------------------------------------------- 6. 审定单

@shape
def approval_form() -> bytes:
    """原型：对账审定侧常见的「工程结算审定单」—— 没有编码/单位/量价，只有 送审|审定|核减 三列。
    旧版审核汇总路径 amount 恒取报审金额，审定单放在审定侧时取到的是报审金额（QA:reconcile-01）。
    真值：side=audit 950000 + 480000 = 1430000；side=None/submit 1000000 + 500000 = 1500000；核减 70000。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "审定单"
    ws.append(["工程结算审定单"])
    ws.append(["工程名称：某学校改造工程"])
    ws.append(["序号", "项目名称", "送审金额（元）", "审定金额（元）", "核减金额（元）", "备注"])
    ws.append([1, "土建工程", 1000000, 950000, 50000, None])
    ws.append([2, "安装工程", 500000, 480000, 20000, None])
    ws.append([None, "合计", 1500000, 1430000, 70000, None])
    return _save(wb)


@shape
def approval_form_wan() -> bytes:
    """审定单的万元版：与 approval_form 同结构，三列金额表头改成「（万元）」、数字 ÷10000。
    审核汇总版式走 _audit_rows，不经过明细路径的换算；旧实现只有 values 换了算，
    amount / audit_amount / deduct_amount 与合计行仍是万元数，却照样出「已乘以 10000 换算为元」，
    对账核减算成 7.0（PR #747 评审）。
    真值与 approval_form 相同：side=None/submit 1000000 + 500000 = 1500000；side=audit 950000 + 480000
    = 1430000；核减 50000 + 20000 = 70000；合计行 ×10000 后 1500000 / 1430000 / 70000。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "审定单"
    ws.append(["工程结算审定单"])
    ws.append(["序号", "项目名称", "送审金额（万元）", "审定金额（万元）", "核减金额（万元）", "备注"])
    ws.append([1, "土建工程", 100, 95, 5, None])
    ws.append([2, "安装工程", 50, 48, 2, None])
    ws.append([None, "合计", 150, 143, 7, None])
    return _save(wb)


# ---------------------------------------------------------------- 7. 工日 / 台班 / 材料一览表

@shape
def labor_machine_tables() -> bytes:
    """原型：指标 run05 市政控制价。「主要工日一览表」「主要机械台班一览表」列头与清单表相似
    （编码|名称|单位|数量|单价|合价），旧版计入，总造价比招标控制价高 2,253,617.75（QA:indicator-02）。
    第三张 sheet 名是「表-09」，标题写在表内第 1 行（主要材料价格表）。
    真值：只计「分部分项清单与计价表」：6000 + 2000 = 8000（= 该表合计行）。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "分部分项清单与计价表"
    ws.append(["序号", "项目编码", "项目名称", "项目特征描述", "计量单位", "工程量", "综合单价", "合价"])
    ws.append([1, "040101001001", "挖一般土方", "三类土", "m3", 200, 30, 6000])
    ws.append([2, "040101002001", "挖沟槽土方", "三类土", "m3", 50, 40, 2000])
    ws.append([None, None, "合计", None, None, None, None, 8000])
    ws2 = wb.create_sheet("主要工日一览表")
    ws2.append(["序号", "编码", "名称", "单位", "数量", "单价", "合价"])
    ws2.append([1, "00010101", "二类人工", "工日", 100, 135, 13500])
    ws2.append([2, "00010102", "机械人工", "工日", 20, 135, 2700])
    ws3 = wb.create_sheet("主要机械台班一览表")
    ws3.append(["序号", "编码", "名称", "单位", "数量", "单价", "合价"])
    ws3.append([1, "99010101", "履带式挖掘机", "台班", 10, 1200, 12000])
    ws4 = wb.create_sheet("表-09")
    ws4.append(["主要材料价格表"])
    ws4.append(["序号", "材料编码", "材料名称", "单位", "数量", "单价", "合价"])
    ws4.append([1, "34110103", "柴油", "kg", 500, 8, 4000])
    return _save(wb)


# ---------------------------------------------------------------- 8. 万元估算表

@shape
def wan_estimate() -> bytes:
    """原型：指标 run11 长兴项目，16 份控制价都没读进来，只用了一张「设备估算表」，表头金额单位
    为万元，旧版泛词「金额」命中「金额（万元）」不换算，算出「单方造价 0.04 元/㎡」（QA:indicator-07）。
    没有国标编码。真值：×10000 后 冷水机组 2×1500000=3000000、冷却塔 2×200000=400000，合计 3400000。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "设备估算表"
    ws.append(["序号", "设备名称", "单位", "数量", "单价（万元）", "金额（万元）"])
    ws.append([1, "冷水机组", "台", 2, 150, 300])
    ws.append([2, "冷却塔", "台", 2, 20, 40])
    ws.append([None, "合计", None, None, None, 340])
    return _save(wb)


@shape
def detail_party_wan() -> bytes:
    """清单明细 + 送审 / 审核两列合价，合价单位万元、单价单位元（有编码有单位，走明细路径不走审核汇总）。
    side=None/submit 时送审合价是 amount，审核合价作为 audit_amount 附带 —— 两列都要 ×10000，
    单价列是元不换算。数字取二进制可精确表示的小数，免得换算尾差掩盖漏乘。
    真值：amount 12500、15000（合计行 2.75 万 = 27500）；audit_amount 11250、13750；price 12.5、30。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "分部分项工程审核对比表"
    ws.append(["序号", "项目编码", "项目名称", "计量单位", "工程量", "综合单价（元）",
               "送审合价（万元）", "审核合价（万元）"])
    ws.append([1, "010101001001", "平整场地", "m2", 1000, 12.5, 1.25, 1.125])
    ws.append([2, "010101002001", "挖一般土方", "m3", 500, 30, 1.5, 1.375])
    ws.append([None, None, "合计", None, None, None, 2.75, 2.5])
    return _save(wb)


# ---------------------------------------------------------------- 9. 无价清单

@shape
def unpriced_bill() -> bytes:
    """原型：R1:ENG-15「室外配套」招标工程量清单（708 行）—— 只有工程量，单价与合价全空。
    旧版：指标 200 + 单方 0.0，投控 200 +「未检出异常」，没有任何告警。真值：3 个叶子行、0 个有价行。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "分部分项工程量清单"
    ws.append(["序号", "项目编码", "项目名称", "项目特征描述", "计量单位", "工程量", "金额（元）", None, None])
    ws.append([None, None, None, None, None, None, "综合单价", "合价", "其中：暂估价"])
    ws.append([1, "040101001001", "挖一般土方", "三类土", "m3", 1200, None, None, None])
    ws.append([2, "040103001001", "回填方", "密实度 93%", "m3", 800, None, None, None])
    ws.append([3, "040501001001", "混凝土管", "DN600", "m", 350, None, None, None])
    return _save(wb)


def priced_bill() -> bytes:
    """有价加无价混合用例里「有价」的那一份（结构同 unpriced_bill，带价）。真值合计 36000+12000=48000。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "分部分项工程量清单"
    ws.append(["序号", "项目编码", "项目名称", "项目特征描述", "计量单位", "工程量", "金额（元）", None, None])
    ws.append([None, None, None, None, None, None, "综合单价", "合价", "其中：暂估价"])
    ws.append([1, "040101001001", "挖一般土方", "三类土", "m3", 1200, 30, 36000, None])
    ws.append([2, "040103001001", "回填方", "密实度 93%", "m3", 800, 15, 12000, None])
    return _save(wb)


# ---------------------------------------------------------------- 10. 样式表损坏

@shape
def broken_stylesheet(variant: str = "missing_name") -> bytes:
    """原型：指标 run03/run04 长兴 1#综合楼 建筑/安装控制价 —— 广联达导出件的 styles.xml 里
    `<cellStyle>` 缺 name 属性，openpyxl 3.1.5 读不了（生产 4372 份 xlsx 中 93 份，QA:indicator-11）。
    variant="missing_name"：删掉 cellStyle 的 name；"invalid_xml"：styles.xml 改成非法 XML。
    真值：010101001001 100×2=200。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "分部分项"
    ws.append(["序号", "项目编码", "项目名称", "计量单位", "工程量", "综合单价", "合价"])
    ws.append([1, "010101001001", "平整场地", "m2", 100, 2, 200])
    raw = _save(wb)
    zin = zipfile.ZipFile(io.BytesIO(raw))
    styles = zin.read("xl/styles.xml").decode("utf-8")
    if variant == "missing_name":
        styles = re.sub(r'(<cellStyle )name="[^"]*" ', r"\1", styles)
    else:
        styles = styles[: len(styles) // 2]     # 截断 → 非法 XML
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "xl/styles.xml":
                data = styles.encode("utf-8")
            zo.writestr(info, data)
    return out.getvalue()


# ---------------------------------------------------------------- 11. 措施费总价项

@shape
def measure_lump_sum() -> bytes:
    """原型：对账 R1「措施费」表 7 行、核减 27,863.17 整块丢失（QA:reconcile-03）：只有
    序号|名称|计算基础|费率|金额，没有编码和单位。真值 7 行总价措施项之和：
    15000+3000+2000+1500+800+700+1000 = 24000（= 合计行）。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "措施费"
    ws.append(["总价措施项目清单与计价表"])
    ws.append(["序号", "项目名称", "计算基础", "费率(%)", "金额(元)", "备注"])
    for i, (n, base, rate, amt) in enumerate([
        ("安全文明施工基本费", "分部分项人工费+机械费", 10.5, 15000),
        ("提前竣工增加费", "分部分项人工费+机械费", 2.1, 3000),
        ("二次搬运费", "分部分项人工费+机械费", 1.4, 2000),
        ("冬雨季施工增加费", "分部分项人工费+机械费", 1.05, 1500),
        ("行车、行人干扰增加费", "分部分项人工费+机械费", 0.56, 800),
        ("已完工程及设备保护费", "分部分项人工费+机械费", 0.49, 700),
        ("夜间施工增加费", "分部分项人工费+机械费", 0.7, 1000),
    ], start=1):
        ws.append([i, n, base, rate, amt, None])
    ws.append([None, "合计", None, None, 24000, None])
    return _save(wb)


# ---------------------------------------------------------------- 12. S.1 标题行 + 补充码

@shape
def s1_title_rows() -> bytes:
    """原型：投控 B2。分部标题行「S.1脚手架工程（编码：011701）」带 qty=0 和合价，旧版当清单项收下，
    两个单体各一行 → 被判「重复计取」（QA:tk-08）。另含「0104 砌筑工程」分部下的补充码 01B047
    与一条无编码、有单位的清单项，编码列写「分部分项合计」的小计行。

    真值：1#楼 12000 + 3000 + 1500 = 16500；2#楼 8000 + 4500 = 12500；两表合计 29000。
    01B047 带 section「0104 砌筑工程」→ 归 0104；011701001001 → 0117 措施项目。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "1#楼建筑"
    ws.append(["序号", "项目编码", "项目名称", "项目特征", "计量单位", "工程量", "综合单价", "合价"])
    ws.append([None, None, "S.1脚手架工程（编码：011701）", None, "项", 0, None, 12000])
    ws.append([1, "011701001001", "综合脚手架", "檐高 30m 内", "m2", 1000, 12, 12000])
    ws.append([None, None, "0104 砌筑工程", None, None, None, None, None])
    ws.append([2, "01B047", "砌体加固筋", "Φ6.5", "t", 1, 3000, 3000])
    ws.append([3, None, "砖砌检查井", "井深 1.2m", "座", 1, 1500, 1500])
    ws.append([None, "分部分项合计", None, None, "项", 1, None, 16500])
    ws2 = wb.create_sheet("2#楼建筑")
    ws2.append(["序号", "项目编码", "项目名称", "项目特征", "计量单位", "工程量", "综合单价", "合价"])
    ws2.append([None, None, "S.1脚手架工程（编码：011701）", None, "项", 0, None, 8000])
    ws2.append([1, "011701001002", "综合脚手架", "檐高 20m 内", "m2", 800, 10, 8000])
    ws2.append([None, None, "0104 砌筑工程", None, None, None, None, None])
    ws2.append([2, "01B047", "砌体加固筋", "Φ8", "t", 1.5, 3000, 4500])
    ws2.append([None, "分部分项合计", None, None, "项", 1, None, 12500])
    return _save(wb)


# ---------------------------------------------------------------- 13. 表-13 费用表

@shape
def fee_table_13() -> bytes:
    """原型：指标 run07 房建控制价「单位工程招标控制价费用表（表-13）」：招标控制价合计 8,715,564.88，
    含干扰行「其中：安全文明施工费合计」。费用表不是清单表，按标题跳过（WS08-16 另行取它的合计）。
    真值：只计分部分项表 96000 + 30000 = 126000。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "单位工程招标控制价费用表"
    ws.append(["单位工程招标控制价费用表"])
    ws.append(["序号", "费用名称", "计算公式", "金额（元）"])
    ws.append([1, "分部分项工程费", "Σ(清单工程量×综合单价)", 6512345.12])
    ws.append([2, "措施项目费", "", 812345.67])
    ws.append([None, "其中：安全文明施工费合计", "", 312345.67])
    ws.append([3, "规费", "", 123456.78])
    ws.append([4, "税金", "", 1267417.31])
    ws.append([5, "招标控制价合计", "1+2+3+4", 8715564.88])
    ws2 = wb.create_sheet("分部分项清单与计价表")
    ws2.append(["序号", "项目编码", "项目名称", "项目特征描述", "计量单位", "工程量", "综合单价", "合价"])
    ws2.append([1, "010401003001", "实心砖墙", "墙厚240", "m3", 200, 480, 96000])
    ws2.append([2, "010502001001", "矩形柱", "C30", "m3", 50, 600, 30000])
    return _save(wb)


# ---------------------------------------------------------------- 13b. 三层费用表（WS26-h，07-012）

# 两张单位工程表的真值（分）：分部分项 + 措施 + 其他项目 + 规费 + 增值税 = 招标控制价合计
_FEE_UNIT_BIG = 15869272.21      # 12,794,000.00 + 1,043,335.23 + 100,000.00 + 721,568.75 + 1,210,368.23
_FEE_UNIT_SMALL = 3895.44        # 3,000.00 + 250.00 + 0 + 323.80 + 321.64
FEE_UNIT_SUM = 15873167.65       # 两张之和；项目级合计取整到元 = 15,873,168，差 0.35


def _fee_unit_sheet(wb: Workbook, name: str, single_name: str, rows: list, total_text: str) -> None:
    ws = wb.create_sheet(f"表10.2.2-13 单位工程招标控制价费用表【{name}】")
    ws.append(["【表10.2.2-13】"])
    ws.append(["单位工程招标控制价费用表"])
    ws.append([f"工程名称：{single_name}-{name}", None, "标段：", "第  1  页  共  1  页", None])
    ws.append(["序号", "费用名称", "计算公式", "金额（元）", "备注"])
    for r in rows:
        ws.append(list(r))
    # 合计行：金额是带千分位的**文本**，左边一格是公式文字（生产原件 113 张单位工程表全是这个写法）
    ws.append(["招标控制价合计", None, "1+2+3+4+5", total_text, None])


def _fee_list_sheet(wb: Workbook, name: str, single_name: str, rows: list) -> None:
    ws = wb.create_sheet(f"表10.2.2-16 分部分项工程清单与计价表【{name}】")
    ws.append(["【表10.2.2-16】"])
    ws.append(["分部分项工程清单与计价表"])
    ws.append([f"单位（专业）工程名称：{single_name}", None, None, None, "标段：", None, None, None, "第 1 页  共 1 页"])
    ws.append(["序号", "项目编码", "项目名称", "项目特征", "计量单位", "工程量", "金额（元）", None, None, None, None, "备注"])
    ws.append([None, None, None, None, None, None, "综合单价", "合价", "其中", None, None, None])
    ws.append([None, None, None, None, None, None, None, None, "人工费", "机械费", "暂估价", None])
    for i, (code, nm, unit, qty, price, amount) in enumerate(rows, 1):
        ws.append([i, code, nm, "", unit, qty, price, amount, None, None, None, None])
    ws.append(["合   计", None, None, None, None, None, None, sum(r[5] for r in rows), None, None, None, None])


@shape
def fee_three_tier(*, project_total: float | None = 15873168, with_units: bool = True,
                   with_small_unit: bool = True, with_single: bool = False) -> bytes:
    """原型：指标 IND-05「…一标段.xlsx」（缺陷 07-012，总造价只有控制价的 58.2%）。形状照生产原件：

      · 项目级「表10.2.2-12 招标控制价费用表」：表内标题「招标控制价费用表」，表头两行——
        「序号|工程名称|金额(元)|其中：（元）|备注」「暂估价|安全文明施工基本费|规费|税金」，
        合计行名是泛名「合计」，金额在 C 列（**数值，取整到元**），右边还有暂估价、安文费、规费
        三个数值格和一个**文本**的税金格；
      · 单位工程「表10.2.2-13 单位工程招标控制价费用表【…】」：合计行「招标控制价合计」，金额是
        带千分位的文本，左边一格是公式文字「1+2+3+4+5」；含干扰行「其中：安全文明施工费合计」；
      · 单项工程「表10.2.2-13 单项工程招标控制价费用表【…】」（with_single）：合计是不带千分位的文本；
      · 两张分部分项清单表（6 行抬头、3 行表头、末行「合   计」）。

    真值：单位工程 15,869,272.21 + 3,895.44 = 15,873,167.65；项目级合计 15,873,168（差 0.35）；
    清单明细合计 12,794,000.00 + 3,000.00 = 12,797,000.00。
    project_total=None 不放项目级表；with_small_unit=False 少一张 3,895.44 的单位工程表。"""
    wb = Workbook()
    wb.remove(wb.active)
    if project_total is not None:
        ws = wb.create_sheet("表10.2.2-12 招标控制价费用表")
        ws.append(["【表10.2.2-12】"])
        ws.append(["招标控制价费用表"])
        ws.append(["工程名称：某地块房地产项目一标段", None, None, None, None, None, None, "第  1  页 共  1  页"])
        ws.append(["序号", "工程名称", "金额(元)", "其中：（元）", None, None, None, "备注"])
        ws.append([None, None, None, "暂估价", "安全文明施工基本费", "规费", "税金", None])
        ws.append(["1", "地下室", "15873167.65", 100000, 852285.06, 721892.55, "1210689.87", None])
        ws.append(["1.1", "土石方工程", "15869272.21", 100000, 852085.06, 721568.75, "1210368.23", None])
        ws.append(["1.2", "标志标线", 3895.44, None, 200, 323.8, 321.64, None])
        ws.append(["合计", None, project_total, 100000, 852285.06, 721892.55, "1210689.87", None])
    if with_single:
        ws = wb.create_sheet("表10.2.2-13 单项工程招标控制价费用表【地下室】")
        ws.append(["【表10.2.2-13】"])
        ws.append(["单项工程招标控制价费用表"])
        ws.append(["工程名称：地下室", None, "标段：", "第  1  页 共  1  页", None])
        ws.append(["序号", "费用名称", "计算公式", "金额（元）", "备注"])
        ws.append([None, None, None, None, None])
        ws.append(["1", "分部分项工程费", "∑（分部分项工程数量×综合单价）", "12797000.00", None])
        ws.append(["2", "措施项目费", None, 1043585.23, None])
        ws.append(["3", "其他项目费", None, 100000, None])
        ws.append(["4", "规费", "∑计算基数×费率", 721892.55, None])
        ws.append(["5", "增值税", "∑计算基数×费率", "1210689.87", None])
        ws.append(["招标控制价合计", None, "1+2+3+4+5", "15873167.65", None])
    if with_units:
        _fee_unit_sheet(wb, "土石方工程", "地下室", [
            ("1", "分部分项工程费", "∑（分部分项工程数量×综合单价）", "12794000.00", "见表 10.2.2-16"),
            ("1.1", "其中：人工费+机械费", "∑分部分项（人工费+机械费）", 9185210.93, None),
            ("2", "措施项目费", None, 1043335.23, None),
            ("2.1", "施工技术措施项目", "∑（技术措施工程数量×综合单价）", 180995.32, "见表 10.2.2-16"),
            ("2.2", "施工组织措施项目", "按实际发生项之和进行计算", 862339.91, "见表 10.2.2-20"),
            ("2.2.1", "其中：安全文明施工费合计", None, 852085.06, "见表 10.2.2-20"),
            ("3", "其他项目费", None, 100000, None),
            ("3.2", "暂估价", "3.2.1+3.2.2+3.2.3", 100000, "见表 10.2.2-21"),
            ("4", "规费", "计算基数×费率", 721568.75, None),
            ("5", "增值税", "计算基数×费率", 1210368.23, None),
        ], "15,869,272.21")
    _fee_list_sheet(wb, "土石方工程", "地下室", [
        ("010101002001", "挖一般土方", "m3", 1000000, 12.5, 12500000),
        ("010103001001", "回填方", "m3", 20000, 14.7, 294000),
    ])
    if with_units and with_small_unit:
        _fee_unit_sheet(wb, "标志标线", "地下室", [
            ("1", "分部分项工程费", "∑（分部分项工程数量×综合单价）", "3000.00", "见表 10.2.2-16"),
            ("2", "措施项目费", None, 250, None),
            ("2.2.1", "其中：安全文明施工费合计", None, 200, None),
            ("3", "其他项目费", None, None, None),
            ("4", "规费", "计算基数×费率", 323.8, None),
            ("5", "增值税", "计算基数×费率", 321.64, None),
        ], "3,895.44")
    _fee_list_sheet(wb, "标志标线", "地下室", [
        ("011407001001", "墙面喷刷涂料", "m2", 100, 30, 3000),
    ])
    return _save(wb)


# ---------------------------------------------------------------- 13c. 单位工程表 + 专业工程表（WS26-h 评审）

FEE_SPEC_A = 6071300.00          # 5,000,000.00 + 400,000.00 + 20,000.00 + 150,000.00 + 501,300.00
FEE_SPEC_B = 1445340.82          # 1,200,000.50 + 90,000.25 + 0 + 36,000.00 + 119,340.07
FEE_SPEC_UNIT = 7516640.82       # 单位工程那一张的合计 = 两张专业工程表之和；项目级合计取整 7,516,641（差 0.18）


def _fee_gld_sheet(wb: Workbook, no: int, form: str, label: str, name: str, rows: list, total) -> None:
    """广联达「N.表10.2.2-13{单位|专业}工程招标控制价费用表」：表名、表内标题写的是这张表的层级，
    第 3 行抬头的标签却不分层——专业工程表上写的也是「单位工程名称：…」。"""
    ws = wb.create_sheet(f"{no}.表10.2.2-13{form}工程招标控制价费用表")
    ws.append(["【表10.2.2-13】", None, None, None, None])
    ws.append([f"{form}工程招标控制价费用表", None, None, None, None])
    ws.append([f"{label}：{name}", None, "标段：", "第 1 页 共 1 页", None])
    ws.append(["序号", "费用名称", "计算公式", "金额（元）", "备注"])
    for r in rows:
        ws.append(list(r))
    ws.append(["招标控制价合计", None, "1+2+3+4+5", total, None])


def _fee_gld_other_sheet(wb: Workbook, no: int, name: str, total: float) -> None:
    """干扰表：每个专业工程后面跟着的「其他项目清单与计价汇总表」，抬头「单位及专业工程名称：…」，末行「合计」。"""
    ws = wb.create_sheet(f"{no}.表10.2.2-21其他项目清单与计价汇总表")
    ws.append(["【表10.2.2-21】", None, None, None])
    ws.append(["其他项目清单与计价汇总表", None, None, None])
    ws.append([f"单位及专业工程名称：{name}", None, "标段：", "第 1 页  共 1 页"])
    ws.append(["序号", "项目名称", "金额(元)", "备注"])
    ws.append(["1", "暂列金额", total, "明细详见表10.2.2-22"])
    ws.append(["合计", None, total, None])


@shape
def fee_unit_with_specialty(*, with_unit: bool = True, with_project: bool = False, with_b: bool = True,
                            b_total: str = "1445340.82", project_total: float = 7516641,
                            label: str = "单位工程名称") -> bytes:
    """原型：样本库「…一期02标-土建工程.xlsx」（1 张单位工程费用表 + 14 张专业工程费用表）。整改前把专业
    工程表也判成 unit 层（靠的是它抬头里的「单位工程名称」五个字），与单位工程表相加，总造价正好翻倍，
    还标成「取自费用表」。样本库里 78 份文件是这个形状，专业工程费用表是那里最常见的费用表版式。

      · 「1.表10.2.2-13单位工程招标控制价费用表」：抬头「单位工程名称：一期02标」，金额全是**文本**
        （不带千分位），合计 7,516,640.82；
      · 「2.…专业工程招标控制价费用表」(1#-地上工程)：抬头「单位工程名称：一期02标-1#-地上工程」，
        金额是数值，合计 6,071,300.00；
      · 「8.…专业工程招标控制价费用表」(地下室工程)：金额是文本，合计 1,445,340.82；
      · 每个专业工程后面一张「其他项目清单与计价汇总表」（末行「合计」，干扰）和一张清单表；
      · with_project：最前面加「1.表10.2.2-12招标控制价费用表」，泛名「合计」7,516,641（取整到元）。

    真值：总造价 = 单位工程那一张 = 7,516,640.82（不是 15,033,281.64）；清单明细 6,200,000.50。
    with_unit=False 去掉单位工程表；with_b=False 只留一张专业工程表；b_total 改地下室那张的合计；
    label 换成「单位及专业工程名称」是样本库里的另一种抬头（整改前它被判成 project 层）。"""
    wb = Workbook()
    wb.remove(wb.active)
    b_num = float(b_total)
    if with_project:
        ws = wb.create_sheet("1.表10.2.2-12招标控制价费用表")
        ws.append(["【表10.2.2-12】", None, None, None, None, None, None, None])
        ws.append(["招标控制价费用表", None, None, None, None, None, None, None])
        ws.append(["工程名称：某地块项目一期02标-土建工程", None, None, None, "第 1 页  共 1 页", None, None, None])
        ws.append(["序号", "工程名称", "金额(元)", "其中：（元）", None, None, None, "备注"])
        ws.append([None, None, None, "暂估价", "安全文明\n施工基本费", "规费", "税金", None])
        ws.append(["1", "一期02标", FEE_SPEC_UNIT, None, 310000, 186000, 620640.07, None])
        ws.append(["合计", None, project_total, None, 310000, 186000, 620640.07, None])
    if with_unit:
        _fee_gld_sheet(wb, 1, "单位", label, "一期02标", [
            ("1", "分部分项工程费", "∑（分部分项工程数量×综合单价）", "6200000.50", None),
            ("1.1", "其中 人工费+机械费", "∑分部分项（人工费+机械费）", "930000.00", None),
            ("2", "措施项目费", None, "490000.25", None),
            ("2.2.1", "其中：安全文明施工费合计", None, "310000.00", None),
            ("3", "其他项目费", None, "20000.00", None),
            ("4", "规费", "计算基数×费率", "186000.00", None),
            ("5", "增值税", "计算基数×费率", "620640.07", None),
        ], f"{FEE_SPEC_UNIT:.2f}")
    _fee_gld_sheet(wb, 2, "专业", label, "一期02标-1#-地上工程", [
        ("1", "分部分项工程费", "∑（分部分项工程数量×综合单价）", 5000000, "见表 10.2.2-16"),
        ("1.1", "其中 人工费+机械费", "∑分部分项（人工费+机械费）", 750000, None),
        ("2", "措施项目费", None, 400000, None),
        ("2.2.1", "其中：安全文明施工费合计", None, 250000, "见表 10.2.2-20"),
        ("3", "其他项目费", None, 20000, None),
        ("4", "规费", "计算基数×费率", 150000, None),
        ("5", "增值税", "计算基数×费率", 501300, None),
    ], FEE_SPEC_A)
    _fee_list_sheet(wb, "1#地上", "一期02标-1#-地上工程", [
        ("010502001001", "矩形柱", "m3", 10000, 500, 5000000),
    ])
    _fee_gld_other_sheet(wb, 6, "一期02标-1#-地上工程", 20000)
    if with_b:
        _fee_gld_sheet(wb, 8, "专业", label, "一期02标-地下室工程", [
            ("1", "分部分项工程费", "∑（分部分项工程数量×综合单价）", "1200000.50", "见表 10.2.2-16"),
            ("2", "措施项目费", None, "90000.25", None),
            ("2.2.1", "其中：安全文明施工费合计", None, "60000.00", "见表 10.2.2-20"),
            ("3", "其他项目费", None, None, None),
            ("4", "规费", "计算基数×费率", "36000.00", None),
            ("5", "增值税", "计算基数×费率", f"{b_num - 1326000.75:.2f}", None),
        ], b_total)
        _fee_list_sheet(wb, "地下室", "一期02标-地下室工程", [
            ("010101002001", "挖一般土方", "m3", 240000.1, 5, 1200000.5),
        ])
    return _save(wb)


@shape
def fee_table_04_with_12(*, with_04: bool = True) -> bytes:
    """GB50500「表-04 单位工程招标控制价汇总表」+「表-12 其他项目清单与计价汇总表」。

    表-12 的表头也是「金额(元)」、末行也叫「合计」，工程名称里还带「建设项目」四个字 —— 不靠表名里的
    「其他项目」把它排除掉，它会被当成项目级汇总表，「其他项目合计 20,000」就成了总造价的候选。
    真值：控制价合计 172,220.00（表-04）；清单明细 126,000.00。with_04=False 只留表-12 与清单表。"""
    wb = Workbook()
    wb.remove(wb.active)
    if with_04:
        ws = wb.create_sheet("表-04 单位工程招标控制价汇总表")
        ws.append(["单位工程招标控制价汇总表"])
        ws.append(["工程名称：某学校建设项目综合楼", None, "标段：", "第 1 页 共 1 页"])
        ws.append(["序号", "汇总内容", "金额(元)", "其中：暂估价(元)"])
        ws.append([1, "分部分项工程", 126000, None])
        ws.append([2, "措施项目", 8000, None])
        ws.append([2.1, "其中：安全文明施工费", 5000, None])
        ws.append([3, "其他项目", 20000, 5000])
        ws.append([4, "规费", 4000, None])
        ws.append([5, "税金", 14220, None])
        ws.append([None, "招标控制价合计", 172220, 5000])
    ws = wb.create_sheet("表-12 其他项目清单与计价汇总表")
    ws.append(["其他项目清单与计价汇总表"])
    ws.append(["工程名称：某学校建设项目综合楼", None, "标段：", "第 1 页 共 1 页"])
    ws.append(["序号", "项目名称", "金额(元)", "结算金额(元)", "备注"])
    ws.append([1, "暂列金额", 15000, None, "明细详见表-12-1"])
    ws.append([2, "暂估价", 5000, None, None])
    ws.append([2.1, "材料(工程设备)暂估价", None, None, "明细详见表-12-2"])
    ws.append([2.2, "专业工程暂估价", 5000, None, "明细详见表-12-3"])
    ws.append([3, "计日工", None, None, "明细详见表-12-4"])
    ws.append([4, "总承包服务费", None, None, "明细详见表-12-5"])
    ws.append([None, "合计", 20000, None, None])
    ws2 = wb.create_sheet("表-08 分部分项工程和单价措施项目清单与计价表")
    ws2.append(["序号", "项目编码", "项目名称", "项目特征描述", "计量单位", "工程量", "综合单价", "合价"])
    ws2.append([1, "010401003001", "实心砖墙", "墙厚240", "m3", 200, 480, 96000])
    ws2.append([2, "010502001001", "矩形柱", "C30", "m3", 50, 600, 30000])
    return _save(wb)


# ---------------------------------------------------------------- 13d. 费用构成表（WS26-h2，07-005 / 07-010 / 07-012）

# 两张单位工程费用表的六项（元）。真值逐项手写，不从解析器输出反抄：
#   地下室土建：1,200,000.00 + 100,000.25 + 50,000.25 + 500,000.00 + 30,000.10 + 170,000.33 = 2,050,000.93
#   室外道路：     80,000.00 +   4,000.00 +  2,000.00 +  10,000.00 +  2,400.40 +   8,964.04 =   107,364.44
FEE_BD_A = {"fbfx": 1200000.00, "measure_tech": 100000.25, "measure_org": 50000.25, "other": 500000.00,
            "regulatory": 30000.10, "tax": 170000.33}
FEE_BD_B = {"fbfx": 80000.00, "measure_tech": 4000.00, "measure_org": 2000.00, "other": 10000.00,
            "regulatory": 2400.40, "tax": 8964.04}
FEE_BD_TOTAL_A = 2050000.93
FEE_BD_TOTAL_B = 107364.44
FEE_BD_TOTAL = 2157365.37        # 两张之和；项目级合计取整到元 = 2,157,365，差 0.37


@shape
def fee_breakdown_units(*, b_tax_delta: float = 0.0, project_total: float | None = None, tax_label: str = "增值税",
                        list_delta: float = 0.0, with_measure_lists: bool = True,
                        b_discount: float | None = None) -> bytes:
    """原型：指标 IND-05「…一标段.xlsx」的 113 张「表10.2.2-13 单位工程招标控制价费用表」（07-005 问的「真正的
    暂估价在哪」、07-012 问的「少的四成是什么」）。形状照生产原件：

      · 每张单位工程表：「1 分部分项工程费 | 1.1 其中：人工费+机械费 | 2 措施项目费 | 2.1 施工技术措施项目 |
        2.1.1 其中：人工费+机械费 | 2.2 施工组织措施项目 | 2.2.1 其中：安全文明施工基本费 | 3 其他项目费 |
        3.1 暂列金额 | 3.2 暂估价 | 3.2.2 专业工程暂估价 | 4 规费 | 5 增值税 | 招标控制价合计」；
        金额列表头「金额（元）」，有的金额是**文本**，「其他项目费」「暂列金额」的金额格可以是空的，
        合计行的金额是带千分位的文本、左边一格是公式文字「1+2+3+4+5」；
      · 「措施项目费」是技术 + 组织的小计，「其中：…」「专业工程暂估价」都是干扰行：当成一项就对不上合计；
      · 每个单位工程一张分部分项清单表；with_measure_lists 时再各一张「施工技术措施项目清单与计价表」（0117）。

    真值：构成表六项 1,280,000.00 / 104,000.25 / 52,000.25 / 510,000.00（其中暂估价 500,000.00、
    暂列金额 10,000.00）/ 32,400.50 / 178,964.37，合计 2,157,365.37。
    清单明细：分部分项 1,280,000.00；技术措施 104,000.25。

    b_tax_delta 只改室外道路那张的「增值税」格（合计不变）→ 那张各项相加与合计差 |b_tax_delta|；
    project_total 加一张项目级「招标控制价费用表」（泛名「合计」）；tax_label 换成 GB 2013 的「税金」；
    list_delta 加在土建清单表的一行合价上（测试件 B「只改清单行、没改费用表」的形状）；
    b_discount 给室外道路那张加一行「6 让利」并把它算进合计（样本库里 100 多份控制价有这一行）。"""
    wb = Workbook()
    wb.remove(wb.active)
    if project_total is not None:
        ws = wb.create_sheet("表10.2.2-12 招标控制价费用表")
        ws.append(["【表10.2.2-12】"])
        ws.append(["招标控制价费用表"])
        ws.append(["工程名称：某地块一标段", None, None, None, None, None, None, "第  1  页 共  1  页"])
        ws.append(["序号", "工程名称", "金额(元)", "其中：（元）", None, None, None, "备注"])
        ws.append([None, None, None, "暂估价", "安全文明施工基本费", "规费", "税金", None])
        ws.append(["1", "地下室土建工程", "2050000.93", 500000, 40000, 30000.10, "170000.33", None])
        ws.append(["2", "室外道路工程", 107364.44, None, 1500, 2400.40, 8964.04, None])
        ws.append(["合计", None, project_total, 500000, 41500, 32400.50, "178964.37", None])

    def unit_sheet(short: str, name: str, v: dict, extra: dict, total: float, tax_delta: float,
                   discount: float | None = None) -> None:
        ws = wb.create_sheet(f"表10.2.2-13 单位工程招标控制价费用表【{short}】")
        ws.append(["【表10.2.2-13】", None, None, None, None])
        ws.append(["单位工程招标控制价费用表", None, None, None, None])
        ws.append([f"工程名称：{name}", None, "标段：", "第  1  页  共  1  页", None])
        ws.append(["序号", "费用名称", "计算公式", "金额（元）", "备注"])
        ws.append(["1", "分部分项工程费", "∑（分部分项工程数量×综合单价）", f"{v['fbfx']:.2f}", "见表 10.2.2-16"])
        ws.append(["1.1", "其中：人工费+机械费", "∑分部分项（人工费+机械费）", round(v["fbfx"] * 0.25, 2), None])
        ws.append(["2", "措施项目费", None, round(v["measure_tech"] + v["measure_org"], 2), None])
        ws.append(["2.1", "施工技术措施项目", "∑（技术措施工程数量×综合单价）", v["measure_tech"], "见表 10.2.2-18"])
        ws.append(["2.1.1", "其中：人工费+机械费", "∑技措项目（人工费+机械费）", round(v["measure_tech"] * 0.5, 2), None])
        ws.append(["2.2", "施工组织措施项目", "按实际发生项之和进行计算", v["measure_org"], "见表 10.2.2-20"])
        ws.append(["2.2.1", "其中：安全文明施工基本费", None, extra["safety"], "见表 10.2.2-20"])
        ws.append(["3", "其他项目费", None, v["other"] or None, None])
        ws.append(["3.1", "暂列金额", "3.1.1+3.1.2+3.1.3", extra["contingency"], "见表 10.2.2-21"])
        ws.append(["3.2", "暂估价", "3.2.1+3.2.2+3.2.3", extra["provisional"], "见表 10.2.2-21"])
        ws.append(["3.2.2", "专业工程暂估价", "按招标文件规定额度列计", extra["provisional"], "见表 10.2.2-24"])
        ws.append(["4", "规费", "计算基数×费率", v["regulatory"], None])
        ws.append(["5", tax_label, "计算基数×费率", f"{v['tax'] + tax_delta:.2f}", None])
        if discount is not None:
            ws.append(["6", "让利", None, discount, None])
            total += discount
        ws.append(["招标控制价合计", None, "1+2+3+4+5" + ("+6" if discount is not None else ""), f"{total:,.2f}", None])

    unit_sheet("地下室土建", "地下室-地下室土建工程", FEE_BD_A,
               {"safety": 40000.00, "contingency": None, "provisional": 500000}, FEE_BD_TOTAL_A, 0.0)
    _mc_list_sheet(wb, "16", "分部分项工程清单与计价表", "地下室土建", "地下室-地下室土建工程", [
        ("010502001001", "矩形柱", "m3", 1000, 600, 600000 + list_delta),
        ("010505001001", "有梁板", "m3", 1200, 500, 600000),
    ])
    if with_measure_lists:
        _mc_list_sheet(wb, "18", "施工技术措施项目清单与计价表", "地下室土建", "地下室-地下室土建工程", [
            ("011702002001", "矩形柱模板", "m2", 2000.005, 50, 100000.25),
        ])
    unit_sheet("室外道路", "室外-室外道路工程", FEE_BD_B,
               {"safety": 1500.00, "contingency": 10000, "provisional": None}, FEE_BD_TOTAL_B, b_tax_delta, b_discount)
    _mc_list_sheet(wb, "16", "分部分项工程清单与计价表", "室外道路", "室外-室外道路工程", [
        ("040202001001", "路床整形", "m2", 4000, 20, 80000),
    ])
    if with_measure_lists:
        _mc_list_sheet(wb, "18", "施工技术措施项目清单与计价表", "室外道路", "室外-室外道路工程", [
            ("041101001001", "墙面脚手架", "m2", 400, 10, 4000),
        ])
    return _save(wb)


# ---------------------------------------------------------------- 13e. 施工技术措施项目综合单价计算表（WS26-p，07-003）

def _mc_fee_sheet(wb: Workbook, short: str, unit: str, rows: list, total_text: str) -> None:
    """单位工程费用表（表10.2.2-13）：抬头「工程名称：单项-单位」，「2.1 施工技术措施项目」行的金额是数值。"""
    ws = wb.create_sheet(f"表10.2.2-13 单位工程招标控制价费用表【{short}】")
    ws.append(["【表10.2.2-13】", None, None, None, None])
    ws.append(["单位工程招标控制价费用表", None, None, None, None])
    ws.append([f"工程名称：{unit}", None, "标段：", "第  1  页  共  1  页", None])
    ws.append(["序号", "费用名称", "计算公式", "金额（元）", "备注"])
    for r in rows:
        ws.append(list(r))
    ws.append(["招标控制价合计", None, "1+2+3+4+5", total_text, None])


def _mc_list_sheet(wb: Workbook, no: str, title: str, short: str, unit: str, rows: list) -> None:
    """分部分项（或施工技术措施项目）清单与计价表：6 行抬头、3 行表头、末行「合   计」。"""
    ws = wb.create_sheet(f"表10.2.2-{no} {title}【{short}】"[:31])     # Excel 表名最多 31 个字符
    ws.append([f"【表10.2.2-{no}】"] + [None] * 11)
    ws.append([title] + [None] * 11)
    ws.append([f"单位（专业）工程名称：{unit}", None, None, None, "标段：", None, None, None, "第 1 页  共 1 页",
               None, None, None])
    ws.append(["序号", "项目编码", "项目名称", "项目特征", "计量单位", "工程量", "金额（元）", None, None, None, None,
               "备注"])
    ws.append([None, None, None, None, None, None, "综合单价", "合价", "其中", None, None, None])
    ws.append([None, None, None, None, None, None, None, None, "人工费", "机械费", "暂估价", None])
    for i, (code, nm, u, qty, price, amount) in enumerate(rows, 1):
        ws.append([i, code, nm, "", u, qty, price, amount, None, None, None, None])
    ws.append(["合   计", None, None, None, None, None, None, sum(r[5] for r in rows), None, None, None, None])


def _mc_calc_sheet(wb: Workbook, sheet_name: str, form_title: str, unit: str, body: list, total,
                   form_no: str = "17") -> None:
    """表10.2.2-17 的版式：第 4、5 行两行表头（「综合单价（元）」下 6 个子列，最后一个是「小计」），
    清单项行的「清单序号」是整数，下面挂的定额子目行「清单序号」为空，编码写「补」「1-93 换」这类定额号；
    「合计(元)」列是取整到元的数；末行「合   计：」写在清单序号列。

    body 每行：("item", 序号, 编码, 名称, 单位, 数量, 小计, 合计) / ("quota", 定额号, 名称, 单位, 数量, 小计, 合计)
    / ("section", 编码, 名称)。人工费…利润五个子列是凑数的（不参与取数），小计才是综合单价。"""
    ws = wb.create_sheet(sheet_name[:31])
    ws.append([f"【表10.2.2-{form_no}】"] + [None] * 11)
    ws.append([form_title] + [None] * 11)
    ws.append([f"单位（专业）工程名称：{unit}", None, None, "标段：", None, None, None, None, None,
               "第 1 页  共 1 页", None, None])
    ws.append(["清单序号", "项目编码\n(定额编码)", "清单(定额)\n项目名称", "计量单位", "数量", "综合单价（元）",
               None, None, None, None, None, "合计(元)"])
    ws.append([None, None, None, None, None, "人工费", "材料(设备)费", "机械费", "管理费", "利润", "小计", None])
    for r in body:
        if r[0] == "section":
            ws.append([None, r[1], r[2]] + [None] * 9)
            continue
        if r[0] == "item":
            _k, seq, code, nm, u, qty, price, amount = r
        else:
            _k, code, nm, u, qty, price, amount = r
            seq = None
        parts = [round(price * f, 2) for f in (0.2, 0.3, 0.1, 0.1)]
        ws.append([seq, code, nm, u, qty, *parts, round(price - sum(parts), 2), price, amount])
    ws.append(["合   计："] + [None] * 10 + [total])


@shape
def measure_calc_sheets(*, with_fee: bool = True, fee_delta: float = 0.0, total_delta: float = 0.0,
                        with_measure_list: bool = False, with_price_analysis: bool = False,
                        audit_qty: float | None = None) -> bytes:
    """原型：投控 TK-C01 / TK-C07、对账 REC-08、指标 IND-05 的「…一标段.xlsx」（缺陷 07-003）。
    113 张「表10.2.2-17 施工技术措施项目综合单价计算表」整张按标题跳过，625 个措施清单项、3907 万元
    不参与检测、对账、指标。这里照原件的形状脱敏造三个单位工程：

      · 地下室-土石方工程：措施表 2 个清单项，各挂 3 / 2 行定额子目；
          011706002001 排水、降水   1 项 × 180,995.32 = 180,995.32（「合计(元)」列 180,995）
          011702001001 垫层模板     182.94 m2 × 50.63 = 9,262.2522 → 9,262.25（「合计(元)」列 9,262）
        清单项之和 190,257.57；本表合计行 190,257（逐行取整后相加）；
        费用表「2.1 施工技术措施项目」190,257.57（fee_delta 加在这一格上）；
      · 地下室-地下室土建工程：分部标题行带章码（011701 / 011702），清单项里有 Z 码与补充码；
          011701003002  内墙砌筑脚手架         11,745.01 m2 × 5.64  = 66,241.8564 → 66,241.86
          Z011701009001 地下室工具式外墙脚手架  1,191.19 m2 × 28.33 = 33,746.4127 → 33,746.41
          01B004        后浇带模板增加费        120 m × 85.50        = 10,260.00
        之和 110,248.27；费用表 110,248.27；
      · 室外-室外道路工程：措施表一个清单序号都没有（只有「[请输入分部名称]」与空合计行），费用表那一格是空的。

    分部分项清单：12,500 + 2,940（土石方）+ 96,000（土建）+ 60,000（道路）= 171,440.00。
    措施清单项合计 300,505.84；全部计入时清单明细合计 471,945.84。

    with_fee=False：不放费用表（只能和本表合计行勾稽）；total_delta 加在土石方那张措施表的合计行上。
    with_measure_list：土石方另有一张「施工技术措施项目清单与计价表」（名称写成「排水降水」，与计算表
      的「排水、降水」不同，跨表去重挡不住）。
    with_price_analysis：土石方另有一张分部分项的「综合单价计算表」——版式与措施表完全相同，只是标题
      不带「措施」。
    audit_qty：审定侧把垫层模板的数量改成这个值，措施表合计行与费用表那一格同步改（对账「调整」用）。"""
    from decimal import ROUND_HALF_UP, Decimal

    def money(q, p) -> float:
        return float((Decimal(repr(q)) * Decimal(repr(p))).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))

    pad_qty = 182.94 if audit_qty is None else audit_qty
    a_items = [(1, "011706002001", "排水、降水", "项", 1, 180995.32),
               (2, "011702001001", "垫层模板", "m2", pad_qty, 50.63)]
    a_sum = round(sum(money(q, p) for _s, _c, _n, _u, q, p in a_items), 2)
    a_total = sum(round(money(q, p)) for _s, _c, _n, _u, q, p in a_items) + total_delta
    a_body = [
        ("section", None, "降排水费"),
        ("item", *a_items[0], 180995),
        ("quota", "补", "电动单级离心清水泵", "台班", 630, 128.92, 81220),
        ("quota", "1-93 换", "直流深井降水 安、拆深井", "座", 11, 3796.52, 41762),
        ("quota", "1-95", "直流深井降水 使用", "座·d", 990, 58.6, 58014),
        ("section", None, "模板工程"),
        ("item", *a_items[1], round(money(pad_qty, 50.63))),
        ("quota", "5-97 换", "基础垫层 复合木模", "m2", pad_qty, 45.0, round(money(pad_qty, 45.0))),
        ("quota", "5-180", "模板支撑增加费", "m2", pad_qty, 5.63, round(money(pad_qty, 5.63))),
    ]
    b_body = [
        ("section", "011701", "脚手架工程"),
        ("item", 1, "011701003002", "内墙砌筑脚手架", "m2", 11745.01, 5.64, 66242),
        ("quota", "18-45 换", "单项脚手架 内墙脚手架 高度3.6m以上", "m2", 11745.01, 5.64, 66242),
        ("item", 2, "Z011701009001", "地下室工具式外墙脚手架", "m2", 1191.19, 28.33, 33746),
        ("quota", "建18B-1 换", "工具式外墙脚手架 落地式 檐口高度10m以内", "m2", 1191.19, 28.33, 33746),
        ("section", "011702", "混凝土模板及支架（撑）"),
        ("item", 3, "01B004", "后浇带模板增加费", "m", 120, 85.5, 10260),
        ("quota", "补", "后浇带模板增加", "m", 120, 85.5, 10260),
    ]
    tech = "施工技术措施项目综合单价计算表"
    wb = Workbook()
    wb.remove(wb.active)
    units = (("土石方工程", "地下室-土石方工程"), ("地下室土建", "地下室-地下室土建工程"),
             ("室外道路工", "室外-室外道路工程"))
    lists = ([("010101002001", "挖一般土方", "m3", 1000, 12.5, 12500),
              ("010103001001", "回填方", "m3", 200, 14.7, 2940)],
             [("010401003001", "实心砖墙", "m3", 200, 480, 96000)],
             [("040203006001", "沥青混凝土", "m2", 500, 120, 60000)])
    fees = (
        ([("1", "分部分项工程费", "∑（分部分项工程数量×综合单价）", "15440.00", "见表 10.2.2-16"),
          ("1.1", "其中：人工费+机械费", "∑分部分项（人工费+机械费）", 9185.21, None),
          ("2", "措施项目费", None, round(a_sum + 5000, 2), None),
          ("2.1", "施工技术措施项目", "∑（技术措施工程数量×综合单价）", round(a_sum + fee_delta, 2), "见表 10.2.2-16"),
          ("2.1.1", "其中：人工费+机械费", "∑技措项目（人工费+机械费）", 50000.0, None),
          ("2.2", "施工组织措施项目", "按实际发生项之和进行计算", 5000.0, "见表 10.2.2-20"),
          ("3", "其他项目费", None, None, None),
          ("4", "规费", "计算基数×费率", 1000.0, None),
          ("5", "增值税", "计算基数×费率", 10000.0, None)],
         f"{15440 + a_sum + 5000 + 1000 + 10000:,.2f}"),
        ([("1", "分部分项工程费", "∑（分部分项工程数量×综合单价）", "96000.00", "见表 10.2.2-16"),
          ("2", "措施项目费", None, 113248.27, None),
          ("2.1", "施工技术措施项目", "∑（技术措施工程数量×综合单价）", 110248.27, "见表 10.2.2-16"),
          ("2.2", "施工组织措施项目", "按实际发生项之和进行计算", 3000.0, "见表 10.2.2-20"),
          ("4", "规费", "计算基数×费率", 800.0, None),
          ("5", "增值税", "计算基数×费率", 9000.0, None)],
         "219,048.27"),
        ([("1", "分部分项工程费", "∑（分部分项工程数量×综合单价）", "60000.00", "见表 10.2.2-16"),
          ("2", "措施项目费", None, 2000.0, None),
          ("2.1", "施工技术措施项目", "∑（技术措施工程数量×综合单价）", None, "见表 10.2.2-16"),
          ("2.2", "施工组织措施项目", "按实际发生项之和进行计算", 2000.0, "见表 10.2.2-20"),
          ("4", "规费", "计算基数×费率", 500.0, None),
          ("5", "增值税", "计算基数×费率", 5000.0, None)],
         "67,500.00"),
    )
    bodies = (a_body, b_body, [("section", None, "[请输入分部名称]")])
    totals = (a_total, 110248, None)
    for (short, unit), rows, (fee_rows, fee_total), body, total in zip(units, lists, fees, bodies, totals):
        if with_fee:
            _mc_fee_sheet(wb, short, unit, fee_rows, fee_total)
        _mc_list_sheet(wb, "16", "分部分项工程清单与计价表", short, unit, rows)
        if short == "土石方工程" and with_measure_list:
            _mc_list_sheet(wb, "18", "施工技术措施项目清单与计价表", short, unit, [
                ("011706002001", "排水降水", "项", 1, 180995.32, 180995.32),
                ("011702001001", "垫层模板", "m2", pad_qty, 50.63, money(pad_qty, 50.63))])
        if short == "土石方工程" and with_price_analysis:
            _mc_calc_sheet(wb, f"表10.2.2-15 综合单价计算表【{short}】", "分部分项工程综合单价计算表", unit, [
                ("item", 1, "010101002001", "挖一般土方", "m3", 1000, 12.5, 12500),
                ("quota", "1-5", "挖掘机挖土 装车", "m3", 1000, 12.5, 12500),
                ("item", 2, "010103001001", "回填方", "m3", 200, 14.7, 2940),
                ("quota", "1-82", "回填土 夯填", "m3", 200, 14.7, 2940),
            ], 15440, form_no="15")
        _mc_calc_sheet(wb, f"表10.2.2-17 {tech}【{short[:2]}", tech, unit, body, total)
    return _save(wb)


# ---------------------------------------------------------------- 14. 同名 sheet 不同标段

@shape
def same_name_sheets_two_lots() -> list[tuple[bytes, str]]:
    """原型：投控 WS08-08「三份进度款合并检测只剩 44 条、分别检测是 84 条」—— 不同标段的两份文件，
    sheet 同名「分部分项清单与计价表」、同一编码。本分片只要求每个 item 带上 file / file_id，
    身份键改造在 WS08-B。真值：标段一 010101001001 12500；标段二 010101001001 9000。
    """
    out = []
    for lot, qty, price in (("标段一", 1000, 12.5), ("标段二", 600, 15)):
        wb = Workbook()
        ws = wb.active
        ws.title = "分部分项清单与计价表"
        ws.append(["序号", "项目编码", "项目名称", "项目特征", "计量单位", "工程量", "综合单价", "合价"])
        ws.append([1, "010101001001", "平整场地", "三类土", "m2", qty, price, qty * price])
        out.append((_save(wb), f"进度款-{lot}.xlsx"))
    return out


# ---------------------------------------------------------------- 覆盖率用例的辅助形状

@shape
def header_broken_with_codes() -> bytes:
    """原型：投控 C1 浙大第五期 7 张明细表（占合同额 76.4%）被整表静默跳过（QA:tk-02）。
    第一张表正常（合计 40000）；第二张表列头被改坏（认不出表头），但数据行有 12 位编码，
    合计行 60000 → 可识别清单表金额覆盖率 40000 / (40000+60000) = 40% → poor。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "正常明细"
    ws.append(["序号", "项目编码", "项目名称", "计量单位", "工程量", "综合单价", "合价"])
    ws.append([1, "010101001001", "平整场地", "m2", 1000, 20, 20000])
    ws.append([2, "010401003001", "实心砖墙", "m3", 50, 400, 20000])
    ws2 = wb.create_sheet("改坏表头")
    ws2.append(["A", "B", "C", "D", "E", "F", "G"])
    ws2.append([1, "010501001001", "现浇混凝土基础", "m3", 50, 400, 20000])
    ws2.append([2, "010502001001", "矩形柱", "m3", 40, 500, 20000])
    ws2.append([3, "010503001001", "矩形梁", "m3", 40, 500, 20000])
    ws2.append([None, None, "合计", None, None, None, 60000])
    return _save(wb)


@shape
def detail_with_total(total: float = 20000) -> bytes:
    """合计勾稽用：明细 12500 + 7500 = 20000；合计行写 total（传 20200 即偏差 1%）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "分部分项清单与计价表"
    ws.append(["序号", "项目编码", "项目名称", "计量单位", "工程量", "综合单价", "合价"])
    ws.append([1, "010101001001", "平整场地", "m2", 1000, 12.5, 12500])
    ws.append([2, "010401003001", "实心砖墙", "m3", 15, 500, 7500])
    ws.append([None, None, "合计", None, None, None, total])
    return _save(wb)


# ---------------------------------------------------------------- 16. 口径指路与跳表披露（WS26-e，07-009 / 07-003 / 07-004）

def _bill_sheet(ws, header: list, rows: list) -> None:
    ws.append(header)
    for r in rows:
        ws.append(list(r))


@shape
def contract_and_current_columns() -> bytes:
    """进度款表：「合同工程量|合同单价|合同合价|本期工程量|本期合价」两组列，没有累计（07-009 的「累计」那条路）。
    真值：basis=current 本期 10×480 + 6×520 = 7,920；basis=contract 合同 200×480 + 100×520 = 148,000。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "分部分项清单与计价表"
    _bill_sheet(ws, ["序号", "项目编码", "项目名称", "计量单位", "合同工程量", "合同单价", "合同合价",
                     "本期工程量", "本期合价"],
                [(1, "010401003001", "实心砖墙", "m3", 200, 480, 96000, 10, 4800),
                 (2, "010501001001", "现浇混凝土基础", "m3", 100, 520, 52000, 6, 3120)])
    return _save(wb)


@shape
def contract_and_settle_amounts() -> bytes:
    """「合同金额 | 结算金额」并列、都没有期次词（结算 ≠ 本期）：auto 落到合同 / 不分期时两列都是候选 ——
    这是真正的「多组列分不清」，WS26-e 不能把它改说成「口径不存在」。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "分部分项清单与计价表"
    _bill_sheet(ws, ["序号", "项目编码", "项目名称", "计量单位", "工程量", "综合单价", "合同金额", "结算金额"],
                [(1, "010401003001", "实心砖墙", "m3", 200, 480, 96000, 95000)])
    return _save(wb)


@shape
def two_unit_prices_plain_amount() -> bytes:
    """「综合单价 | 全费用单价 | 合价」：都不分期，单价有两列。选「本期」时金额、工程量都没有本期列，单价退到
    合同 / 不分期时两列并存 —— 改前判成「口径有歧义」，真正的原因是本期列不存在。
    （「合同单价 | 单价」那种不行：退到合同口径时 _narrow 取合同那一列，到不了单价歧义这条分支。）"""
    wb = Workbook()
    ws = wb.active
    ws.title = "分部分项清单与计价表"
    _bill_sheet(ws, ["序号", "项目编码", "项目名称", "计量单位", "工程量", "综合单价", "全费用单价", "合价"],
                [(1, "010401003001", "实心砖墙", "m3", 200, 480, 530, 96000)])
    return _save(wb)


@shape
def cumulative_only_columns() -> bytes:
    """只有「累计工程量 | 综合单价 | 累计合价」：auto 找不到本期列，退到合同 / 不分期也没有 → 指路到『累计』。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "分部分项清单与计价表"
    _bill_sheet(ws, ["序号", "项目编码", "项目名称", "计量单位", "累计工程量", "综合单价", "累计合价"],
                [(1, "010401003001", "实心砖墙", "m3", 30, 480, 14400)])
    return _save(wb)


def _unreadable_measure_sheet(wb: Workbook, name: str, unit: str, *, codes: bool = True) -> None:
    """读不了的技术措施表：标题含「措施」+「综合单价分析」，但没有「清单序号」列（表头是「序号」），
    WS26-p 的版式认不出 → 整张按标题跳过。codes=False 时编码格写定额号，没有 12 位清单编码。"""
    ws = wb.create_sheet(name[:31])
    ws.append(["施工技术措施项目综合单价分析表"] + [None] * 6)
    ws.append([f"单位（专业）工程名称：{unit}"] + [None] * 6)
    ws.append(["序号", "项目编码", "项目名称", "计量单位", "数量", "综合单价", "合价"])
    ws.append([1, "011701001001" if codes else "18-1", "综合脚手架", "m2", 100, 30.5, 3050])
    ws.append([2, "011702001001" if codes else "5-97", "基础模板", "m2", 50, 60, 3000])


def _fee_sheet(wb: Workbook, name: str) -> None:
    ws = wb.create_sheet(name[:31])
    ws.append(["单位工程招标控制价费用表", None, None, None])
    ws.append(["序号", "费用名称", "计算公式", "金额（元）"])
    ws.append(["1", "分部分项工程费", None, 20000])
    ws.append(["招标控制价合计", None, None, 20000])


@shape
def unreadable_measure_sheets(*, n_fee: int = 2, n_measure: int = 2, n_measure_no_code: int = 0,
                              cover: bool = False) -> bytes:
    """原型：TK-C03 / TK-C07 的「…一标段.xlsx」在 WS26-p 之前 —— 355 张表里 129 张费用表 + 113 张措施表被
    「按标题跳过」，跳表告警一句话顿号拼 242 个表名（7448 字）；措施表被说成「非清单表」。
    这里：1 张分部分项（12,500 + 7,500 = 20,000）+ n_fee 张费用表 + n_measure 张有清单编码的、
    n_measure_no_code 张没有清单编码的措施分析表（版式 WS26-p 认不出）+ 可选一张封面。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "表-08 分部分项清单与计价表"
    _bill_sheet(ws, ["序号", "项目编码", "项目名称", "计量单位", "工程量", "综合单价", "合价"],
                [(1, "010101001001", "平整场地", "m2", 1000, 12.5, 12500),
                 (2, "010401003001", "实心砖墙", "m3", 15, 500, 7500)])
    if cover:
        wb.create_sheet("封面").append(["某工程招标控制价"])
    for i in range(n_fee):
        _fee_sheet(wb, f"表-13 单位工程费用表{i + 1}")
    for i in range(n_measure):
        _unreadable_measure_sheet(wb, f"措施综合单价分析表{i + 1}", f"单位{i + 1}")
    for i in range(n_measure_no_code):
        _unreadable_measure_sheet(wb, f"措施综合单价分析表N{i + 1}", f"单位N{i + 1}", codes=False)
    return _save(wb)


@shape
def progress_with_measure_calc() -> bytes:
    """进度款工作簿：分部分项表有「本期工程量 | 本期合价」，另有一张 WS26-p 认得出的表10.2.2-17 措施计算表
    （只有不分期的数量、综合单价小计）。选「本期」时措施表没有那个口径，按标题跳过 —— 要黄字说清，
    不能混进「按标题跳过」的灰字里说成非清单表。本期合计 10×480 = 4,800。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "分部分项清单与计价表"
    _bill_sheet(ws, ["序号", "项目编码", "项目名称", "计量单位", "合同工程量", "综合单价", "合同合价",
                     "本期工程量", "本期合价"],
                [(1, "010401003001", "实心砖墙", "m3", 200, 480, 96000, 10, 4800)])
    _mc_calc_sheet(wb, "表10.2.2-17 施工技术措施项目综合单价计算表", "施工技术措施项目综合单价计算表",
                   "地下室-土石方工程", [
                       ("item", 1, "011706002001", "排水、降水", "项", 1, 180995.32, 180995),
                       ("quota", "补", "电动单级离心清水泵", "台班", 630, 128.92, 81220),
                   ], 180995)
    return _save(wb)


# ---------------------------------------------------------------- 接线层用例：落到临时目录

def put(raw: bytes, file_name: str, file_id: str | None = None) -> dict:
    """把夹具字节写到本进程的临时目录，返回请求里的文件引用 ``{path, file_id, file_name}``。

    zt 改动：旧仓写到 FILE_ROOT/_engine_test/ws08/ 并返回带 rel_path 的 FileRef；这里写到
    :func:`tests.costsheet.fixtures.endpoint.scratch_dir`（进程退出时删除）。"""
    from tests.costsheet.fixtures.endpoint import scratch_dir
    path = scratch_dir() / file_name
    path.write_bytes(raw)
    return {"file_id": file_id or file_name, "path": str(path), "file_name": file_name}
