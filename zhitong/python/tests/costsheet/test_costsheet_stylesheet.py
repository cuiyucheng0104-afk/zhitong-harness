# 来源：engine-service/tests/test_costsheet_stylesheet.py（zhitong-cost-ai@991ff5cd）
"""WS08-05 样式表损坏 xlsx 兜底（issue #724）。

部分广联达导出的 xlsx，styles.xml 里的 cellStyle 缺 name 属性，openpyxl 3.1.5 读取报错；
端点返回 422，提示却是「请用广联达导出 .xlsx」—— 出问题的恰恰就是广联达导出件
（生产 4372 份 xlsx 中 93 份，5 个项目；QA:indicator-11、reconcile-15）。

变异验收：删掉「补名」一步 → 告警码断言转红（落到 stylesheet_dropped）；
删掉「去样式表」一步 → 非法 XML 用例转红（读不出来）。
"""
from __future__ import annotations

import pytest

from zt.costsheet import xlsx_tolerant
from zt.costsheet.parser import parse_audit_summary, parse_costsheet
from tests.costsheet.fixtures import costsheet_shapes as shapes


def _codes(out):
    return [w["code"] for w in out["warning_items"]]


def test_plain_openpyxl_really_fails_on_the_fixture():
    """负控：夹具确实是 openpyxl 直接读不了的形状 —— 否则下面的用例什么也没证明。"""
    import io
    from openpyxl import load_workbook
    with pytest.raises(Exception):
        load_workbook(io.BytesIO(shapes.broken_stylesheet("missing_name")), read_only=True, data_only=True)
    with pytest.raises(Exception):
        load_workbook(io.BytesIO(shapes.broken_stylesheet("invalid_xml")), read_only=True, data_only=True)


def test_missing_cellstyle_name_is_patched_and_rows_are_read():
    out = parse_costsheet(shapes.broken_stylesheet("missing_name"), "长兴1#综合楼-建筑控制价.xlsx")
    assert [(i["code"], i["amount"]) for i in out["items"]] == [("010101001001", 200.0)]
    assert "stylesheet_patched" in _codes(out)
    assert "stylesheet_dropped" not in _codes(out)
    w = next(w for w in out["warning_items"] if w["code"] == "stylesheet_patched")
    assert w["level"] == "info"


def test_invalid_stylesheet_xml_is_dropped_and_rows_are_read():
    out = parse_costsheet(shapes.broken_stylesheet("invalid_xml"), "坏样式.xlsx")
    assert [(i["code"], i["amount"]) for i in out["items"]] == [("010101001001", 200.0)]
    w = next(w for w in out["warning_items"] if w["code"] == "stylesheet_dropped")
    assert w["level"] == "warn" and "日期列可能显示为数字" in w["message"]


def test_non_zip_bytes_get_actionable_message_not_gbq_advice():
    out = parse_costsheet(b"this is not a zip", "伪装成xlsx.xlsx")
    assert out["items"] == []
    w = next(w for w in out["warning_items"] if w["code"] == "xlsx_unreadable")
    assert "另存为 .xlsx" in w["message"]
    assert "广联达" not in w["message"], "出问题的就是广联达导出件，别再叫人「用广联达导出」"


def test_tolerant_loader_reason_is_capped_at_300_chars():
    with pytest.raises(xlsx_tolerant.XlsxUnreadable) as ei:
        xlsx_tolerant.load_workbook_tolerant(b"PK" + b"x" * 5000)
    assert len(ei.value.reason) <= 300


def test_patch_cellstyle_names_only_touches_nameless_tags():
    xml = ('<cellStyles><cellStyle name="Normal" xfId="0" builtinId="0"/>'
           '<cellStyle xfId="1"/><x:cellStyle xfId="2"/></cellStyles>')
    out, n = xlsx_tolerant.patch_cellstyle_names(xml)
    assert n == 2
    assert out.count('name="Normal"') == 1
    assert 'name="patched_style_1"' in out and 'name="patched_style_2"' in out


def test_parse_audit_summary_shares_the_fallback():
    """汇总表读法与 parse_costsheet 共用同一个容错读取（K25：容错模块两方共用）。"""
    out = parse_audit_summary(shapes.broken_stylesheet("missing_name"), "坏样式.xlsx")
    assert "stylesheet_patched" in _codes(out)
