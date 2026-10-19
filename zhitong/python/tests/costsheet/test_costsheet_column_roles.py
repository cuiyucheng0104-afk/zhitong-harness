# 来源：engine-service/tests/test_costsheet_column_roles.py（zhitong-cost-ai@991ff5cd）
"""WS08-01 表头块模型与列角色（issue #724）。

每条用例的夹具都来自 tests/fixtures/costsheet_shapes.py（真实版式、脱敏、现场生成），
期望值是夹具构造时手算的真值，不从解析器输出反抄。

变异验收（每项回退后必须转红）：
  _bind 删掉期次过滤          → 浙大形、浙四形、A1 形用例
  审定侧改回 side=None        → 门窗形、审定单用例与 test_reconcile_sides
  投票函数恒返回 contract     → A1 形用例
  _row 恢复用 _f()            → test_reconcile_sides 里 qty 为 None 的断言
  _audit_rows 不乘列 scale    → test_wan_approval_form_* 与 test_reconcile_sides 的万元审定单用例
  _detail_rows audit_amount 不乘 scale → test_wan_audit_amount_on_detail_sheet_is_scaled
"""
from __future__ import annotations

import io
import re

import pytest
from openpyxl import load_workbook

from zt.tk import engine as tk_engine
from zt.costsheet.parser import parse_audit_summary, parse_costsheet
from tests.costsheet.fixtures import costsheet_shapes as shapes


def _by_code(out: dict) -> dict:
    return {it["code"]: it for it in out["items"]}


# ---------------------------------------------------------------- 浙大三行表头

def test_zheda_three_row_header_current_period_binds_current_columns():
    """basis=current：工程量取「本期完成工程量」、金额取「本期产值」、单价取综合单价。
    旧版：qty=投标量 200、amount=合同合价 96000。"""
    out = parse_costsheet(shapes.zheda_three_row_header(), "浙大第五期.xlsx", basis="current")
    d = out["meta"]["sheets_detail"][0]
    assert d["status"] == "used"
    assert d["header_rows"] == [4, 5, 6], "三行表头要整块识别"
    assert d["columns"]["qty"] == "施工单位/本期完成工程量"
    assert d["columns"]["amount"] == "施工单位/本期产值"
    assert d["columns"]["price"] == "投标价(元)/综合单价"
    items = _by_code(out)
    assert (items["010401003001"]["qty"], items["010401003001"]["price"], items["010401003001"]["amount"]) \
        == (30.0, 480.0, 14400.0)
    assert items["010401004001"]["amount"] == 0.0
    assert items["010501001001"]["amount"] == 31200.0
    assert items["010502001001"]["amount"] == 24000.0
    assert out["meta"]["parsed_total"] == 69600.0
    assert out["meta"]["basis_used"] == "current.submit"


def test_zheda_auto_basis_takes_current_and_contract_basis_takes_bid_columns():
    auto = parse_costsheet(shapes.zheda_three_row_header(), "浙大第五期.xlsx")
    assert auto["meta"]["parsed_total"] == 69600.0, "auto：有本期列就取本期（D45）"
    contract = parse_costsheet(shapes.zheda_three_row_header(), "浙大第五期.xlsx", basis="contract")
    assert contract["meta"]["parsed_total"] == 327000.0
    assert _by_code(contract)["010401003001"]["qty"] == 200.0
    cum = parse_costsheet(shapes.zheda_three_row_header(), "浙大第五期.xlsx", basis="cumulative")
    assert cum["meta"]["parsed_total"] == 144600.0


def test_zheda_component_columns_are_never_bound():
    """「其中/人工费|机械费|暂估价」是合价的组成，不能被当成合价（取了它金额就少一个数量级）。"""
    out = parse_costsheet(shapes.zheda_three_row_header(), "浙大第五期.xlsx", basis="contract")
    cols = out["meta"]["sheets_detail"][0]["columns"]
    assert "人工费" not in cols["amount"] and "暂估价" not in cols["amount"]


def test_zheda_values_carry_every_role():
    """values 带该行所有带角色的数值，下游（回源定位 / 基准）不必再读原表。"""
    it = _by_code(parse_costsheet(shapes.zheda_three_row_header(), "x.xlsx"))["010401003001"]
    assert it["values"]["contract.amount"] == 96000.0
    assert it["values"]["current.submit.amount"] == 14400.0
    assert it["values"]["cumulative.submit.amount"] == 38400.0
    assert it["values"]["previous.submit.qty"] == 50.0
    assert it["periods"] == {"qty": "current", "price": "contract", "amount": "current"}


# ---------------------------------------------------------------- 门窗三段对比

def test_door_audit_side_takes_audit_block_with_name_from_same_row():
    """side=audit 取「审核」块的量、价、额，且与名称来自同一行同一块（旧版审定侧取成「原合同」）。"""
    out = parse_costsheet(shapes.door_three_block(), "门窗审核.xlsx", side="audit")
    first = out["items"][0]
    assert (first["code"], first["name"]) == ("LC2020", "铝合金推拉窗")
    assert (first["qty"], first["price"], first["amount"]) == (105.0, 480.0, 50400.0)
    assert out["meta"]["parsed_total"] == 149400.0
    detail = {d["sheet"]: d for d in out["meta"]["sheets_detail"]}
    assert detail["6#楼"]["columns"]["amount"] == "审核/金额"


def test_door_submit_side_takes_submit_block():
    out = parse_costsheet(shapes.door_three_block(), "门窗审核.xlsx", side="submit")
    first = out["items"][0]
    assert (first["qty"], first["price"], first["amount"]) == (110.0, 500.0, 55000.0)
    assert first["audit_amount"] == 50400.0, "送审侧附带同口径的审核金额"
    assert out["meta"]["parsed_total"] == 158400.0


# ---------------------------------------------------------------- 浙四：本月上报 / 本月审核

def test_zhesi_submit_and_audit_take_this_month_output():
    sub = parse_costsheet(shapes.zhesi_code_column_totals(), "浙四第25期.xlsx", side="submit")
    aud = parse_costsheet(shapes.zhesi_code_column_totals(), "浙四第25期.xlsx", side="audit")
    assert [i["amount"] for i in sub["items"]] == [5000.0, 52000.0, 12000.0]
    assert [i["amount"] for i in aud["items"]] == [4800.0, 41600.0, 12000.0]
    assert sub["meta"]["basis_used"] == "current.submit"
    assert aud["meta"]["basis_used"] == "current.audit"


def test_zhesi_quantity_without_current_column_is_none_not_contract_qty():
    """表头没有「本月工程量」：qty 记缺失 None，而不是拿合同量顶上（那会让 self_check 用合同量×价比本月产值）。"""
    sub = parse_costsheet(shapes.zhesi_code_column_totals(), "浙四第25期.xlsx", side="submit")
    assert all(i["qty"] is None for i in sub["items"])


# ---------------------------------------------------------------- A1 市政：期次投票

def test_a1_vote_picks_current_period():
    """合价只有一列、不写期次；逐行 合价 ≈ 本期量 × 单价 → 投票出本期（旧版 qty=投标量 1000）。"""
    out = parse_costsheet(shapes.a1_period_qty(), "A1市政.xlsx")
    items = _by_code(out)
    assert items["040101001001"]["qty"] == 200.0
    assert items["040201004001"]["qty"] == 300.0
    assert out["meta"]["parsed_total"] == 47000.0
    assert out["meta"]["basis_used"] == "current"


def test_a1_fifty_fifty_is_ambiguous_not_guessed():
    """一半行按本期算、一半按累计算 → 判歧义：该表不产出行，出 column_ambiguous（不猜）。"""
    out = parse_costsheet(shapes.a1_period_qty("split"), "A1市政.xlsx")
    assert out["items"] == []
    assert out["meta"]["sheets_detail"][0]["status"] == "ambiguous"
    codes = [w["code"] for w in out["warning_items"]]
    assert "column_ambiguous" in codes
    amb = next(w for w in out["warning_items"] if w["code"] == "column_ambiguous")
    assert amb["level"] == "error" and "合价" in amb["message"]


# ---------------------------------------------------------------- 审定单 / 万元

def test_approval_form_audit_side_takes_approved_amount():
    """审定单放在审定侧时取审定金额（旧版恒取报审金额）。"""
    aud = parse_costsheet(shapes.approval_form(), "结算审定单.xlsx", side="audit")
    assert [i["amount"] for i in aud["items"]] == [950000.0, 480000.0]
    legacy = parse_costsheet(shapes.approval_form(), "结算审定单.xlsx")
    assert [i["amount"] for i in legacy["items"]] == [1000000.0, 500000.0], "side=None 保持旧语义：报审"
    assert [i["deduct_amount"] for i in legacy["items"]] == [50000.0, 20000.0]


def test_wan_yuan_columns_are_scaled():
    out = parse_costsheet(shapes.wan_estimate(), "设备估算表.xlsx")
    assert [i["amount"] for i in out["items"]] == [3000000.0, 400000.0]
    assert [i["price"] for i in out["items"]] == [1500000.0, 200000.0]
    assert "amount_scaled_wan" in [w["code"] for w in out["warning_items"]]


def _scaled_cols(out: dict) -> set:
    """amount_scaled_wan 告警里点名的列。告警说「已换算」的列必须真的落进了 item 字段。"""
    return {re.search(r"列「(.+?)」", w["message"]).group(1)
            for w in out["warning_items"] if w["code"] == "amount_scaled_wan"}


def test_wan_approval_form_scales_amount_fields_and_total_row():
    """审核汇总版式（审定单）走 _audit_rows，不经明细路径：amount / audit_amount / deduct_amount
    与合计行都要 ×10000，并与同一行的 values、parse_audit_summary 的读数一致。
    旧实现 amount=100、合计行 150（两者都没换算所以勾稽不报），却照样发「已乘以 10000 换算为元」。"""
    raw = shapes.approval_form_wan()
    legacy = parse_costsheet(raw, "结算审定单.xlsx")
    assert legacy["meta"]["sheets_detail"][0]["format"] == "audit_summary"
    assert [i["amount"] for i in legacy["items"]] == [1000000.0, 500000.0], "side=None 取报审"
    assert [i["audit_amount"] for i in legacy["items"]] == [950000.0, 480000.0]
    assert [i["deduct_amount"] for i in legacy["items"]] == [50000.0, 20000.0]
    assert all(i["amount"] == i["values"]["submit.amount"] for i in legacy["items"])
    assert legacy["meta"]["sheets_detail"][0]["total_row_amount"] == 1500000.0
    assert _scaled_cols(legacy) == {"送审金额(万元)", "审定金额(万元)", "核减金额(万元)"}

    aud = parse_costsheet(raw, "结算审定单.xlsx", side="audit")
    assert [i["amount"] for i in aud["items"]] == [950000.0, 480000.0]
    assert [i["deduct_amount"] for i in aud["items"]] == [50000.0, 20000.0]
    assert aud["meta"]["sheets_detail"][0]["total_row_amount"] == 1430000.0
    assert _scaled_cols(aud) == {"审定金额(万元)", "核减金额(万元)"}, "送审列没有进字段，不该说它换算了"

    for out in (legacy, aud):
        assert "sheet_total_mismatch" not in [w["code"] for w in out["warning_items"]]
    # 两条读审核汇总表的路径口径一致（parse_audit_summary 一直是换算过的）
    (ps,) = parse_audit_summary(raw, "结算审定单.xlsx")["sheets"]
    items = [r for r in ps["rows"] if r["kind"] == "item"]
    assert [float(r["values"]["audit.amount"]) for r in items] == [i["amount"] for i in aud["items"]]
    assert [float(r["values"]["submit.amount"]) for r in items] == [i["amount"] for i in legacy["items"]]


def test_wan_total_row_mismatch_still_fires_on_audit_summary():
    """负控：合计行与明细不符时照样报 sheet_total_mismatch —— 证明上面「不报」是因为两边都换算了，
    而不是勾稽在审核汇总路径上根本没跑。"""
    wb = load_workbook(io.BytesIO(shapes.approval_form_wan()))
    wb.active.cell(row=5, column=3, value=160)      # 合计行送审 160 万 ≠ 150 万
    buf = io.BytesIO()
    wb.save(buf)
    out = parse_costsheet(buf.getvalue(), "结算审定单.xlsx")
    assert out["meta"]["sheets_detail"][0]["total_row_amount"] == 1600000.0
    assert "sheet_total_mismatch" in [w["code"] for w in out["warning_items"]]


def test_wan_audit_amount_on_detail_sheet_is_scaled():
    """明细版式带送审 / 审核两列万元合价：附带的 audit_amount 也 ×10000（旧实现原样取成 1.12）；
    单价列是元，不换算。"""
    out = parse_costsheet(shapes.detail_party_wan(), "送审审核对比.xlsx")
    d = out["meta"]["sheets_detail"][0]
    assert d["format"] == "detail"
    assert [i["amount"] for i in out["items"]] == [12500.0, 15000.0]
    assert [i["audit_amount"] for i in out["items"]] == [11250.0, 13750.0]
    assert [i["price"] for i in out["items"]] == [12.5, 30.0]
    assert d["total_row_amount"] == 27500.0
    assert _scaled_cols(out) == {"送审合价(万元)", "审核合价(万元)"}


def test_missing_qty_and_price_stay_none():
    """缺失的量与价是 None，不是 0.0（旧版对账把缺失写成 0.0）。"""
    out = parse_costsheet(shapes.approval_form(), "结算审定单.xlsx", side="audit")
    assert all(i["qty"] is None and i["price"] is None for i in out["items"])


# ---------------------------------------------------------------- 投控：口径对了，假阳就没了

def test_tk_zheda_rows_with_zero_current_output_raise_no_big_amount():
    """浙大形：本期为 0 或本期 <2 万的行不出 big_amount；旧版按合同合价报了 4 行。"""
    res = tk_engine.run([(shapes.zheda_three_row_header(), "浙大第五期进度款.xlsx")], None, project_id="P1")
    big = {a["item_code"] for a in res["anomalies"] if a["subtype"] == "big_amount"}
    assert big == {"010501001001", "010502001001"}


def test_tk_a1_self_check_has_no_false_positive():
    """A1 形：投票出本期后 量×价 = 金额，self_check 0 条（旧版投标量×价 ≠ 本期金额，报假阳）。"""
    res = tk_engine.run([(shapes.a1_period_qty(), "A1市政进度款.xlsx")], None, project_id="P1")
    assert [a for a in res["anomalies"] if a["subtype"] == "self_check"] == []


def test_tk_self_check_skipped_when_qty_and_amount_are_different_periods():
    """工程量取的是合同量、金额取的是本期产值时，self_check 不跑（跑了必然全是假阳）。"""
    row = {"project_id": "P", "major": "m", "sheet": "s", "item_code": "010101001001",
           "submit_qty": 1000, "submit_price": 30, "submit_amount": 6000,
           "periods": {"qty": "contract", "price": "contract", "amount": "current"}}
    cfg = tk_engine.load_rules()["rules"]["self_check"]
    assert tk_engine.rule_self_check([row], cfg, 0.5) == []
    same = dict(row, periods={"qty": "current", "price": "contract", "amount": "current"})
    assert len(tk_engine.rule_self_check([same], cfg, 0.5)) == 1, "同期次组照常判"


@pytest.mark.parametrize("basis", ["auto", "current", "contract", "cumulative"])
def test_standard_single_row_header_is_unaffected_by_basis(basis):
    """回归：单行表头、不分期不分当事方的标准表，任何口径都取同一组列。"""
    out = parse_costsheet(shapes.detail_with_total(), "标准.xlsx", basis=basis)
    if basis in ("current", "cumulative"):
        # WS26-e（07-009）：这是「所选口径不存在」，不是「多组列分不清」—— 状态 basis_missing、码 basis_not_found
        d = out["meta"]["sheets_detail"][0]
        assert out["items"] == [] and d["status"] == "basis_missing", \
            "明确要本期/累计而表头没有 → 说出来，不拿不分期列冒充"
        codes = [w["code"] for w in out["warning_items"]]
        assert "basis_not_found" in codes and "column_ambiguous" not in codes
        assert d["basis_available"] == ["plain"] and out["meta"]["basis_available"] == ["plain"]
        w = next(w for w in out["warning_items"] if w["code"] == "basis_not_found")
        assert w["level"] == "error" and "不分期的合价列" in w["message"] and "合同价" not in w["message"]
        assert "『自动』" in w["hint"]
    else:
        assert out["meta"]["parsed_total"] == 20000.0


# ---------------------------------------------------------------- 表头续行的两道挡板

def _wb_bytes(rows: list[list], title: str = "分部分项") -> bytes:
    import io
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = title
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_section_title_right_under_header_is_not_absorbed_into_header():
    """表头下紧跟「二 | 单价措施项目」标题行：它含「单价」二字，若被当成表头续行，
    名称列的路径就成了「项目名称/单价措施项目」而被当成单价列，整表取数全错。"""
    raw = _wb_bytes([
        ["序号", "项目编码", "项目名称", "计量单位", "工程量", "综合单价", "合价"],
        ["二", None, "单价措施项目"],
        [1, "011701001001", "综合脚手架", "m2", 100, 12, 1200],
    ])
    out = parse_costsheet(raw, "x.xlsx")
    d = out["meta"]["sheets_detail"][0]
    assert d["header_rows"] == [1]
    assert d["columns"]["name"] == "项目名称" and d["columns"]["price"] == "综合单价"
    assert [(i["code"], i["amount"], i["section"]) for i in out["items"]] == \
        [("011701001001", 1200.0, "二 单价措施项目")]


def test_month_labelled_groups_30003_shape():
    """30003 进度款形：R1「…|金额（元）|金额（元）|2023年6月上报产值|  |2023年6月审核产值|…」、
    R2「…|综合单价|合价|已完工程量|已完产值（元）|…」。R2 的格子里带数字（「6月」），不能因为
    能抠出数字就当成数据行 —— 它是表头续行。月份组不是「本期」，auto 按合价口径取（D45）。"""
    raw = _wb_bytes([
        ["序号", "项目编码", "项目名称", "计量单位", "工程量", "金额（元）", "金额（元）",
         "2023年6月上报产值", None, "2023年6月审核产值", None],
        ["序号", "项目编码", "项目名称", "计量单位", "工程量", "综合单价", "合价",
         "已完工程量", "已完产值（元）", "已完工程量", "已完产值（元）"],
        ["", "", "0101 土石方工程"],
        [1, "010101002001", "挖一般土方", "m3", "327.70", 4.76, 1559.85, 191.55, 911.78, 182.43, 868.37],
    ])
    out = parse_costsheet(raw, "30003形.xlsx", side="submit")
    d = out["meta"]["sheets_detail"][0]
    assert d["header_rows"] == [1, 2]
    assert d["columns"]["amount"] == "金额(元)/合价"
    assert d["basis_used"] == "plain"
    (it,) = out["items"]
    assert (it["qty"], it["price"], it["amount"]) == (327.7, 4.76, 1559.85)
    assert it["values"]["dated.submit.amount"] == 911.78
    assert it["chapter"] == "0101"


def test_continuation_row_with_month_labels_is_still_header():
    """30003「土建造价汇总表」形：续行里是「6月上报产值（元）| 6月咨询审定产值（元）」——
    格子里带数字，但它们是表头文字。按「能抠出数字就是数」判，续行会被当成数据行，
    月份组的当事方 / 期次全部丢失。"""
    raw = _wb_bytes([
        ["序号", "工程名称", "金额(元)", "工程合同造价（不含安全文明基本费）", None, None],
        [None, None, None, "合同金额(元)", "6月上报产值（元）", "6月咨询审定产值（元）"],
        ["1.1", "地下室土石方工程", 3150635.35, 2886878.83, 2117521.03, 2016686.70],
    ], title="土建造价汇总表")
    out = parse_costsheet(raw, "汇总.xlsx", basis="contract")
    d = out["meta"]["sheets_detail"][0]
    assert d["header_rows"] == [1, 2]
    assert d["columns"]["amount"] == "工程合同造价(不含安全文明基本费)/合同金额(元)"
