# 来源：engine-service/tests/test_estimate_render_notes.py（zhitong-cost-ai@991ff5cd，W4 迁入 zt；改动见「zt 改动」注释）
"""WS10-P06 / P14（issue #761）：交付 xlsx 的生成说明、计费依据、缓存值与模板清理。

P06 改前：覆盖率 10.67%、总建筑面积 0 的表与完整概算外观相同；5 表 I 列（费用计算标准）为空；
平台用 openpyxl `data_only=True` 读 `估概算总表!C47` 得 None（只写公式串、`<v>` 为空）。
P14 改前：交付表里有「构筑物（***）工程」「*x%」「单位（x）」「勘察费设计费」「海棉城市」，
1 表 D16/D18 标成 m³ / m，D50 占比用金额格式，两张清单表是空壳。

端到端走 compute 端点（render=true），读回 xlsx 断言；少量用例直接调渲染器。
"""
from __future__ import annotations

import base64
import json
import re
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from zt.estimate import render as er
from tests.estimate.fixtures.endpoint import HEADERS, client  # zt 改动：旧端点的转接

_FX = json.loads((Path(__file__).parent / "fixtures" / "estimate_gsyz_v1_synthetic.json").read_text("utf-8"))
FULL = _FX["inputs"]
ROOT11 = {k: v for k, v in FULL.items() if k.startswith("经济技术指标!")}
GEN = "2026-09-30 10:00:00"


def _h():
    return HEADERS


def _compute(**body) -> tuple[dict, bytes]:
    r = client.post("/engine/estimate/compute", headers=_h(), json={"render": True, **body})
    assert r.status_code == 200, r.text
    d = r.json()
    return d, base64.b64decode(d["xlsx_base64"])


@pytest.fixture(scope="module")
def draft():
    """Cm_K 形状的低覆盖输入：根输入 + 3 个专项成本位，带来源、画像、生成时间。"""
    inputs = dict(ROOT11, **{c: FULL[c] for c in ("3.专项工程费!C7", "3.专项工程费!C8", "3.专项工程费!C9")})
    src = {c: {"source": "user"} for c in inputs}
    src["经济技术指标!D4"] = {"source": "extract", "file_id": "F1", "file_name": "设计说明.pdf",
                          "quote": "用地面积 12000 平方米"}
    return _compute(inputs=inputs, project_name="仓储项目", input_sources=src, generated_at=GEN,
                    project_profile={"building_function": "warehouse_logistics", "project_city": "zhejiang_other"},
                    area_checks={"gross_building_area": 10.0},
                    acknowledged_checks=[{"code": "area_total_mismatch", "reason": "原文总面积含架空层"}])


@pytest.fixture(scope="module")
def full_ok():
    return _compute(inputs=FULL, project_name="完整项目", generated_at=GEN,
                    input_sources={c: {"source": "user"} for c in FULL}, fee_bases={"中标金额": 6000})


# ---------------------------------------------------------------- 缓存值（data_only 读得到金额）

def test_cached_values_match_totals(full_ok):
    """变异：跳过 inject_cached_values → data_only 读 C47 得 None → 本条转红。"""
    body, xlsx = full_ok
    wb = load_workbook(BytesIO(xlsx), data_only=True)
    for code, t in body["totals"].items():
        sheet, coord = t["cell"].split("!", 1)
        got = wb[sheet][coord].value
        if t["value"] is None:
            assert got in (None, ""), code
        else:
            assert got == pytest.approx(t["value"]), f"{code} {t['cell']}：缓存 {got} ≠ totals {t['value']}"
    # 文本透传公式（总表 B25 = '3.专项工程费'!B6）的缓存值是文本
    assert wb["估概算总表"]["B25"].value == wb["3.专项工程费"]["B6"].value
    # 公式本身一个字节没动
    wf = load_workbook(BytesIO(xlsx))
    assert wf["估概算总表"]["C47"].value.startswith("=")


def test_generated_at_makes_the_bytes_deterministic():
    """传入 generated_at 时同一入参两次渲染**逐字节**相同（zip 时间戳与 docProps 也由它定）。"""
    a = _compute(inputs=FULL, generated_at=GEN)[1]
    b = _compute(inputs=FULL, generated_at=GEN)[1]
    assert a == b


# ---------------------------------------------------------------- 生成说明 / 草稿前缀

def test_notes_sheet_is_last_and_active_for_draft(draft):
    body, xlsx = draft
    assert body["quality"]["review_level"] == "draft"
    wb = load_workbook(BytesIO(xlsx))
    assert wb.sheetnames[-1] == "生成说明" and wb.active.title == "生成说明"
    assert wb["估概算总表"]["A1"].value.startswith("【草稿】")
    assert wb["5.工程建设其他费用"]["B1"].value.startswith("【草稿】")


def test_notes_sheet_says_what_is_missing_and_where_values_came_from(draft):
    _, xlsx = draft
    ws = load_workbook(BytesIO(xlsx))["生成说明"]
    rows = [[c for c in r if c is not None] for r in ws.iter_rows(values_only=True)]
    flat = "\n".join("｜".join(str(c) for c in r) for r in rows)
    assert f"生成时间｜{GEN}" in flat
    assert "口径策略版本｜gsyz_v1.policy.2026-10-06" in flat   # WS29-d 升版
    assert "质量结论｜草稿（draft）" in flat
    assert "coverage_below_draft" in flat
    assert "业态｜仓储物流" in flat                       # 画像取中文名（策略里的选项）
    assert re.search(r"1\.建筑工程费｜29｜0｜0\.0%", flat)   # 覆盖率分表
    assert "单体面积合计与原文总面积对不上" in flat and "原文总面积含架空层" in flat   # 已确认的对账差异
    # 逐格来源：抽取的那一格带文件名与引文，其余手填
    assert "经济技术指标!D4｜用地面积｜AI 抽取（用户确认）｜设计说明.pdf；引文：用地面积 12000 平方米" in flat
    assert "3.专项工程费!C7｜" in flat and "用户手填" in flat


def test_ok_quality_has_no_prefix_and_opens_on_the_first_sheet(full_ok):
    body, xlsx = full_ok
    assert body["quality"]["review_level"] == "ok"
    wb = load_workbook(BytesIO(xlsx))
    assert wb.active.title == "经济技术指标"
    assert wb["估概算总表"]["A1"].value == "估概算总表"
    assert wb.sheetnames[-1] == "生成说明"          # ok 也附生成说明，只是不抢首屏


def test_fee_basis_is_written_into_column_i(full_ok):
    """5 表 I 列「费用计算标准」：文号 + 算式 + 签署类型（改前全空）。"""
    _, xlsx = full_ok
    ws = load_workbook(BytesIO(xlsx))["5.工程建设其他费用"]
    i14 = ws["I14"].value
    assert i14 and "签署：仓主默认放行（非造价员签署）" in i14 and "算式" in i14
    assert "基数「工程费用」" in i14 and "（取自本表工程费用合计）" in i14
    # 没出建议值的也写清为什么（合成数据的工程费用超出勘察设计费原表范围），中文 + 原因码
    assert "本次未出建议值：" in ws["I17"].value and "（rule_not_applicable）" in ws["I17"].value
    # 汇总项写在小计行
    assert "财建〔2016〕504号" in ws["I9"].value


def test_render_without_notes_keeps_the_nine_sheets():
    """老调用方（render()、LibreOffice 核对脚本）不带 notes：不附生成说明，但缓存值照补。"""
    xlsx = er.render(ROOT11)
    wb = load_workbook(BytesIO(xlsx), data_only=True)
    assert "生成说明" not in wb.sheetnames and len(wb.sheetnames) == 9
    assert wb["经济技术指标"]["D5"].value == pytest.approx(sum(
        ROOT11[f"经济技术指标!D{r}"] for r in (7, 8, 10, 11, 12, 13)))


# ---------------------------------------------------------------- P14 交付表清理

def test_d50_ratio_uses_percent_format():
    """估概算总表 D 列是「占总投资比例（%）」：D50（=D47+D48+D49，add_chain）也必须是百分比格式。

    变异：_num_format 不看列头量纲（add_chain 一律金额格式）→ 本条转红。
    """
    wb = load_workbook(BytesIO(er.render(ROOT11)))
    ws = wb["估概算总表"]
    assert ws["D50"].number_format == "0.00%"
    assert ws["D47"].number_format == "0.00%"
    assert ws["C50"].number_format == "#,##0.00"      # 金额列不受影响


_PLACEHOLDER = re.compile(r"\*{3}|\*\s*x\s*%|[（(]\s*x\s*[)）]|勘察费设计费|海棉")


def test_no_placeholder_or_typo_left_in_the_deliverable():
    """变异：删掉模板 static 修正（占位符原样照搬）→ 本条转红。"""
    wb = load_workbook(BytesIO(er.render(ROOT11)))
    left = [f"{ws.title}!{c.coordinate}={c.value!r}"
            for ws in wb.worksheets for row in ws.iter_rows() for c in row
            if isinstance(c.value, str) and _PLACEHOLDER.search(c.value)]
    assert left == []
    assert wb["估概算总表"]["B21"].value == "构筑物工程"
    assert wb["5.工程建设其他费用"]["B17"].value == "勘察设计费"
    assert wb["4.室外及配套工程费"]["B13"].value == "海绵城市工程"
    assert wb["1.建筑工程费"]["D16"].value == "㎡" and wb["1.建筑工程费"]["D18"].value == "㎡"
    assert wb["估概算总表"]["E4"].value == "单位"


def test_placeholder_self_check_fires_when_one_is_left():
    """自检真的扫得到：用户把槽位名写成「构筑物（***）」→ template_placeholder_left（告警，不拦截）。"""
    res = er.render_estimate(ROOT11, slot_labels={"经济技术指标!D13": "构筑物（***）"})
    w = [x for x in res.warnings if x["code"] == "template_placeholder_left"]
    assert len(w) == 1 and "构筑物（***）" in w[0]["message"]
    assert er.render_estimate(ROOT11).warnings == []


def test_reference_sheets_are_marked_and_point_at_the_total_area():
    xlsx = er.render(ROOT11)
    wf = load_workbook(BytesIO(xlsx))
    wv = load_workbook(BytesIO(xlsx), data_only=True)
    for name in ("1.1概算建筑项目清单", "2.1估概算安装项目清单"):
        assert "附表（口径参考，本期未编制数量与金额）" in wf[name]["A1"].value
        assert wf[name]["J2"].value == "建筑面积："
        assert wf[name]["K2"].value == "='经济技术指标'!$D$5"
        assert wv[name]["K2"].value == pytest.approx(wv["经济技术指标"]["D5"].value)


def test_static_fix_is_disclosed_to_the_frontend():
    body = client.get("/engine/estimate/template", headers=_h()).json()
    fix = next(f for f in body["fixes"] if f["id"] == "static_text_fixes")
    cells = {c["cell"] for c in fix["changes"]}
    assert {"估概算总表!B21", "估概算总表!B46", "5.工程建设其他费用!B17", "1.建筑工程费!D16"} <= cells
