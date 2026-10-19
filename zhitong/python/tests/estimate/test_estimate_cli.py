"""概算 CLI：extract_prepare / extract_finalize / suggest / compute / template 的信封、落盘与跨平台路径。"""
from __future__ import annotations

import io
import json
from pathlib import Path

from openpyxl import load_workbook

from zt.estimate import extract as ex
from zt.estimate.cli import cli
from tests.quantity.support import run_cli

FIXTURE = json.loads((Path(__file__).parent / "fixtures" / "estimate_gsyz_v1_synthetic.json").read_text("utf-8"))
SOURCE = ("项目概况：总用地面积 12000 平方米，总建筑面积 30000 平方米。\n"
          "A区地下室建筑面积 3000 平方米。\n1#楼 住宅 5000 平方米\n")
REQUIRED_11 = [
    "经济技术指标!D4", "经济技术指标!D7", "经济技术指标!D8", "经济技术指标!D10",
    "经济技术指标!D11", "经济技术指标!D12", "经济技术指标!D13", "经济技术指标!D15",
    "经济技术指标!D16", "经济技术指标!D19", "经济技术指标!D20",
]


def _call(*argv) -> tuple[int, dict]:
    out = io.StringIO()
    code = cli.main([str(a) for a in argv], stdout=out)
    return code, json.loads(out.getvalue())


def _project(tmp_path: Path, rel: str = "资料/设计 说明/可研.txt", text: str = SOURCE) -> tuple[Path, str]:
    root = tmp_path / "项目" / "2026-007-某工程"
    p = root.joinpath(*rel.split("/"))
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(text, encoding="utf-8")
    return root, rel


def _answer(file_id: str) -> dict:
    return {"files": [{
        "file_id": file_id,
        "buildings": [{"name": "A区地下室", "level": "underground", "area": "3000",
                       "quote": "A区地下室建筑面积 3000 平方米"},
                      {"name": "1#楼", "level": "above_ground", "area": "5000", "quote": "1#楼 住宅 5000 平方米"}],
        "checks": {"gross_building_area": [{"value": "30000", "quote": "总建筑面积 30000 平方米"}]},
        "cells": {"经济技术指标!D4": [{"value": "1.2万", "quote": "总用地面积 12000 平方米"}]},
    }]}


def test_extract_two_halves_through_the_cli(tmp_path):
    root, rel = _project(tmp_path)
    out = tmp_path / "工作" / "2026-10-19-概算"
    code, summary = run_cli("zt.estimate.cli", "extract_prepare", "--root", str(root), "--files", rel,
                            "--project-name", "某工程一标段", "--out", f"{out / 'extract'}/")
    assert code == 0 and summary["status"] == "ok", summary
    prep = json.loads((out / "extract" / "result.json").read_text(encoding="utf-8"))
    data = prep["data"]
    assert data["prompt_version"] == ex.ESTIMATE_PROMPT_VERSION
    assert [f["file_id"] for f in data["files"]] == [rel]
    ins = (out / "extract" / "概算抽取说明.md").read_text(encoding="utf-8")
    body = (out / "extract" / "概算抽取正文-01.txt").read_text(encoding="utf-8")
    assert "严禁输出任何金额" in ins and ex.ESTIMATE_PROMPT_VERSION in ins
    assert body.startswith(f"file_id: {rel}\n") and "总用地面积 12000 平方米" in body

    (out / "answers.json").write_text(json.dumps(_answer(rel), ensure_ascii=False), encoding="utf-8")
    code, summary = run_cli("zt.estimate.cli", "extract_finalize", "--prepare", str(out / "extract" / "result.json"),
                            "--answers", str(out / "answers.json"), "--out", f"{out / 'extract_final'}/")
    assert code == 0 and summary["status"] == "ok", summary
    res = json.loads((out / "extract_final" / "result.json").read_text(encoding="utf-8"))
    d4 = next(c for c in res["data"]["candidates"] if c["cell"] == "经济技术指标!D4")
    assert d4["value"] == 12000.0, "1.2万 由代码乘出来"
    assert d4["confidence"] == 1.0 and d4["needs_manual"] is True
    assert res["data"]["scope"] == {"project_lot": 1, "lot_source": "project_name"}
    assert "scope_whole_site_suspect" in {w["code"] for w in res["warning_items"]}
    # 抽取段零金额字段（spec I-3.1）
    assert "totals" not in res["data"] and "sheets" not in res["data"]


def test_extract_finalize_rejects_changed_source_and_bad_answers(tmp_path):
    root, rel = _project(tmp_path)
    code, prep = _call("extract_prepare", "--root", root, "--files", rel)
    assert code == 0
    (tmp_path / "prep.json").write_text(json.dumps(prep, ensure_ascii=False), encoding="utf-8")
    (tmp_path / "bad.json").write_text('{"files": [{"file_id": "资料/x.txt", "cells": {}}]}', encoding="utf-8")
    code, body = _call("extract_finalize", "--prepare", tmp_path / "prep.json", "--answers", tmp_path / "bad.json")
    assert code == 2 and body["error"]["code"] == "invalid_answers"
    (tmp_path / "shape.json").write_text(json.dumps({"files": [{"file_id": rel, "x": 1}]}), encoding="utf-8")
    code, body = _call("extract_finalize", "--prepare", tmp_path / "prep.json", "--answers", tmp_path / "shape.json")
    assert code == 2 and body["error"]["code"] == "invalid_answers"
    (root / rel).write_text(SOURCE.replace("12000", "13000"), encoding="utf-8")
    (tmp_path / "ok.json").write_text(json.dumps(_answer(rel), ensure_ascii=False), encoding="utf-8")
    code, body = _call("extract_finalize", "--prepare", tmp_path / "prep.json", "--answers", tmp_path / "ok.json")
    assert code == 2 and body["error"]["code"] == "prepare_mismatch"


def test_extract_prepare_all_unreadable_is_exit_2(tmp_path):
    root, rel = _project(tmp_path, "资料/空.txt", "   ")
    code, body = _call("extract_prepare", "--root", root, "--files", rel)
    assert code == 2 and body["error"]["code"] == "no_text" and "旁车" in body["error"]["hint"]


def test_compute_writes_real_formula_workbook(tmp_path):
    req = tmp_path / "req.json"
    req.write_text(json.dumps({"inputs": FIXTURE["inputs"], "rates": FIXTURE["rates"],
                               "generated_at": "2026-10-19 09:00:00"}, ensure_ascii=False), encoding="utf-8")
    out = tmp_path / "工作" / "2026-10-19-概算"
    code, summary = run_cli("zt.estimate.cli", "compute", "--in", str(req), "--out", f"{out}/")
    assert code == 0 and summary["status"] == "ok", summary
    res = json.loads((out / "result.json").read_text(encoding="utf-8"))
    assert res["needs_manual"] is True
    assert all(t["needs_manual"] is True for t in res["data"]["totals"].values())
    assert res["data"]["artifacts"]["xlsx"] == (out / "概算.xlsx").as_posix()
    wb = load_workbook(out / "概算.xlsx")
    formulas = sum(1 for ws in wb.worksheets for row in ws.iter_rows() for c in row
                   if isinstance(c.value, str) and c.value.startswith("="))
    assert formulas >= 499, "写真 Excel 公式，不写静态值"
    # 同一入参（generated_at 固定）两次渲染逐字节一致
    out2 = tmp_path / "again"
    assert _call("compute", "--in", req, "--out", f"{out2}/")[0] == 0
    assert (out2 / "概算.xlsx").read_bytes() == (out / "概算.xlsx").read_bytes()


def test_compute_gates_are_insufficient_in_order(tmp_path):
    def run(inputs, **extra):
        p = tmp_path / "r.json"
        p.write_text(json.dumps({"inputs": inputs, **extra}, ensure_ascii=False), encoding="utf-8")
        return _call("compute", "--in", p, "--no-render")

    code, env = run({})
    assert code == 0 and env["status"] == "insufficient" and env["status_reason"] == "required_inputs_missing"
    assert env["data"]["detail"].startswith("缺少必填输入位：")
    zero_area = {c: 0.0 for c in REQUIRED_11}
    zero_area["经济技术指标!D4"] = 12000.0
    code, env = run(zero_area)
    assert env["status"] == "insufficient" and env["status_reason"] == "area_gate_failed", "面积闸在全零闸之前"
    root = {c: 1000.0 for c in REQUIRED_11}
    code, env = run(root)
    assert env["status"] == "insufficient" and env["status_reason"] == "all_zero_result"
    assert env["warning_items"][0]["level"] == "error"
    code, body = run(dict(root, **{"1.建筑工程费!C7": True}))
    assert code == 2 and body["error"]["code"] == "invalid_request", "bool 不许静默转 1.0"


def test_compute_requires_an_output_path_when_rendering(tmp_path):
    p = tmp_path / "r.json"
    p.write_text(json.dumps({"inputs": FIXTURE["inputs"]}, ensure_ascii=False), encoding="utf-8")
    code, body = _call("compute", "--in", p)
    assert code == 2 and body["error"]["code"] == "no_output"


def test_suggest_uses_the_bundled_sample_baselines_by_default(tmp_path):
    p = tmp_path / "r.json"
    p.write_text(json.dumps({"inputs": {c: 1000.0 for c in REQUIRED_11}}, ensure_ascii=False), encoding="utf-8")
    code, env = _call("suggest", "--in", p)
    assert code == 0 and env["status"] == "ok"
    assert env["warning_items"][0]["code"] == "baselines_default_sample"
    assert env["data"]["suggestions"] and all(s["needs_manual"] is True for s in env["data"]["suggestions"])
    p.write_text(json.dumps({"inputs": {c: 1000.0 for c in REQUIRED_11}, "baselines": []}), encoding="utf-8")
    code, env = _call("suggest", "--in", p)
    assert code == 0 and env["status"] == "empty" and env["status_reason"] == "no_baselines"


def test_template_lists_required_inputs():
    code, env = _call("template")
    assert code == 0 and env["data"]["required_inputs"] == REQUIRED_11
    code, body = _call("template", "--template-code", "../etc")
    assert code == 2 and body["error"]["code"] == "invalid_request"
