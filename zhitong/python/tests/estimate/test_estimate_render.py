# 来源：engine-service/tests/test_estimate_render.py（zhitong-cost-ai@991ff5cd，W4 迁入 zt；改动见「zt 改动」注释）
"""概算 xlsx 渲染器（I-2.6）—— estimate_render 模块测试。

本模块存在的意义：产物必须是**带真公式的活表**，OnlyOffice 打开改 D7 就能全表联动重算。
所以测试的重心是「重开之后公式还在不在、逐字符对不对」，而不是「文件生成了没报错」。

依赖说明：`estimate_engine.py` 与 `gsyz_v1.json` 由并行分支 feat/estimate-engine-w2a
提供。未合并时本文件整体 skip（不是伪造依赖硬跑）。
"""
from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

_TPL = (Path(__file__).resolve().parents[2] / "zt" / "estimate" / "data"
        / "estimate_templates" / "gsyz_v1.json")
pytestmark = pytest.mark.skipif(
    not _TPL.is_file(),
    reason="gsyz_v1.json / estimate_engine.py 由 feat/estimate-engine-w2a 提供，未合并时跳过",
)

if _TPL.is_file():
    from zt.estimate.engine import EstimateTemplateError, load_template
    from zt.estimate.render import render

ROOT_INPUTS = {
    "经济技术指标!D4": 12000.0, "经济技术指标!D7": 5200.0, "经济技术指标!D8": 3100.0,
    "经济技术指标!D10": 8800.0, "经济技术指标!D11": 7600.0, "经济技术指标!D12": 6400.0,
    "经济技术指标!D13": 5900.0, "经济技术指标!D15": 4200.0, "经济技术指标!D16": 3300.0,
    "经济技术指标!D19": 260.0, "经济技术指标!D20": 140.0,
}


@pytest.fixture(scope="module")
def book():
    return load_workbook(BytesIO(render(ROOT_INPUTS, project_name="XX 房屋建筑工程")))


# --------------------------------------------------------------- 真公式


def test_all_499_formulas_written_verbatim(book):
    """验收判据：515 − 16 条 DISPIMG = 499 条公式必须逐字符写进去。

    「逐字符」不是形式主义：OnlyOffice 保存后公式字符串要能原样回读，
    少一个 $ 或把 '1.建筑工程费' 的引号丢了，跨表引用就断了。

    ⚠ O-09 起，会除零的 237 格（unit_cost 224 + ratio 12 + 传播到的 估概算总表!D50）
    写成 `=IFERROR(<原公式体>,"")`。**公式体本身仍逐字符不变**——这条测试守的
    「引用没断」在包装内照旧成立，所以按包装后的期望形态断言，而不是放宽成子串匹配。
    """
    from zt.estimate.engine import error_prone_cells

    tpl = load_template()
    wrapped = error_prone_cells()
    assert len(wrapped) == 237, f"除零闭包应 237 格，实得 {len(wrapped)}"

    def _expect(cell: str, formula: str) -> str:
        return f"=IFERROR({formula[1:]},\"\")" if cell in wrapped else formula

    expected = {f["cell"]: _expect(f["cell"], f["formula"])
                for f in tpl["formulas"] if f["class"] != "dispimg"}
    assert len(tpl["formulas"]) == 515
    assert len(expected) == 499
    # 费率格改引独立单元格（契约 §1.3-3），单列断言，见 test_rate_lives_in_its_own_cell
    rate_cell = tpl["rate_params"]["reserve_rate"]["cell"]
    expected.pop(rate_cell)

    missing, wrong = [], []
    for cell, formula in expected.items():
        sheet, coord = cell.split("!", 1)
        got = book[sheet][coord].value
        if got is None:
            missing.append(cell)
        elif got != formula:
            wrong.append((cell, formula, got))
    assert missing == []
    assert wrong == []


def test_no_formula_was_flattened_into_a_static_value(book):
    """一条都不许退化成静态值——退化了 OnlyOffice 里改参数就不重算了。"""
    tpl = load_template()
    for f in tpl["formulas"]:
        if f["class"] == "dispimg":
            continue
        sheet, coord = f["cell"].split("!", 1)
        v = book[sheet][coord].value
        assert isinstance(v, str) and v.startswith("="), f"{f['cell']} 不是公式：{v!r}"


# --------------------------------------------------------- 空态 #DIV/0! 消除（O-09）


def test_special_works_unit_cost_column_is_wrapped_in_iferror(book):
    """3.专项工程费 F 列（单方指标 =Cn/En*10000）必须以 IFERROR 开头。

    这张表的 42 个数量位无根输入可推导（`concept_base_metric` 里 quantity 基数的
    root_input_ref 全 NULL），只填 11 个根指标时分母恒空 —— 造价员打开第一屏
    就是 47 个 #DIV/0!。分母还没填是**正常空态**，不该长得像错误。
    """
    tpl = load_template()
    f_col = [f for f in tpl["formulas"]
             if f["cell"].startswith("3.专项工程费!F") and f["class"] == "unit_cost"]
    assert len(f_col) >= 40, f"3.专项 F 列单方指标公式只找到 {len(f_col)} 条"
    for f in f_col:
        sheet, coord = f["cell"].split("!", 1)
        v = book[sheet][coord].value
        assert v.startswith('=IFERROR('), f"{f['cell']} 未包 IFERROR：{v!r}"
        assert v.endswith(',"")'), f"{f['cell']} 兜底值不是空串：{v!r}"
        # 公式体逐字符不变 —— 包装不许顺手改引用
        assert f["formula"][1:] in v, f"{f['cell']} 公式体被改动了：{v!r}"


def test_only_division_bearing_formulas_are_wrapped(book):
    """限定包装范围：不含除法的公式**不许**被 IFERROR 吞掉。

    全包会把 #REF!/#NAME? 一起变成空白，那是另一种「误导性成功」。
    """
    from zt.estimate.engine import error_prone_cells

    tpl = load_template()
    wrapped = error_prone_cells()
    rate_cell = tpl["rate_params"]["reserve_rate"]["cell"]
    for f in tpl["formulas"]:
        if f["class"] == "dispimg" or f["cell"] in wrapped or f["cell"] == rate_cell:
            continue
        sheet, coord = f["cell"].split("!", 1)
        assert not book[sheet][coord].value.startswith("=IFERROR("), \
            f"{f['cell']} 不含除法却被包了"


def test_rate_render_cells_come_from_the_template_not_a_constant(book):
    """O-14：费率格号必须由模板下发，才能出现在 GET /estimate/template 的返回体里。

    只活在 estimate_render.RATE_CELLS 常量里的时候，前端拿不到 J46，
    面向用户的文案只好写 rate_params.cell（C46）—— 而 C46 是公式格，
    造价员照着去改就把 `=(C5+C45)*J46` 删了，预备费从此不再联动。
    """
    tpl = load_template()
    p = tpl["rate_params"]["reserve_rate"]
    assert p["render_value_cell"] == "估概算总表!J46"
    assert p["render_label_cell"] == "估概算总表!J45"
    assert p["render_formula"] == "=(C5+C45)*J46"
    assert p["cell"] == "估概算总表!C46" and p["cell"] != p["render_value_cell"]
    # 模板说什么就渲染成什么
    assert book["估概算总表"]["J46"].value == 0.05
    assert book["估概算总表"]["J45"].value == "预备费率"
    assert book["估概算总表"]["C46"].value == "=(C5+C45)*J46"


def test_rate_render_spec_falls_back_to_the_constant_for_old_templates():
    """模板没写 render_* 时回落常量 —— 与 O-14 之前的行为一字不差。"""
    from zt.estimate.render import RATE_CELLS, rate_render_spec

    assert rate_render_spec({}, "reserve_rate") == RATE_CELLS["reserve_rate"]
    assert rate_render_spec({}, "no_such_rate") is None


def test_error_prone_closure_covers_propagation():
    """闭包必须含沿引用传播的下游：IFERROR 的兜底是**文本**空串，
    `=D47+D48+D49` 拿到文本会变成 #VALUE! —— 不跟着包等于把 #DIV/0! 换成 #VALUE!。"""
    from zt.estimate.engine import error_prone_cells

    wrapped = error_prone_cells()
    assert "估概算总表!D47" in wrapped          # ratio，种子
    assert "估概算总表!D50" in wrapped          # =D47+D48+D49，传播来的
    assert "估概算总表!C50" not in wrapped      # =C47+C48+C49，纯金额链，不含除法


def test_written_formula_count_is_exactly_499(book):
    """模板公式 499 条（515 − 16 条 DISPIMG）原样写入；另外只多 WS10-P14 的两格：
    两张清单表第 2 行引用本概算总建筑面积（K2）。多出任何别的公式都要红。"""
    found = [
        f"{ws.title}!{c.coordinate}"
        for ws in book.worksheets
        for row in ws.iter_rows()
        for c in row
        if isinstance(c.value, str) and c.value.startswith("=")
    ]
    extra = sorted(x for x in found if x.endswith("!K2"))
    assert extra == ["1.1概算建筑项目清单!K2", "2.1估概算安装项目清单!K2"]
    assert len(found) - len(extra) == 499


def test_dispimg_cells_are_skipped(book):
    """16 张定额图不写进 xlsx：openpyxl 写图必须 Pillow，venv 未装，加依赖需过决策。
    图片改为挂知识库 + 结果页侧栏引用。"""
    tpl = load_template()
    imgs = [f["cell"] for f in tpl["formulas"] if f["class"] == "dispimg"]
    assert len(imgs) == 16
    for cell in imgs:
        sheet, coord = cell.split("!", 1)
        assert book[sheet][coord].value is None, f"{cell} 不该有内容"


# zt 改动：依赖清单是 zhitong/python/pyproject.toml
_REQUIREMENTS = Path(__file__).resolve().parents[2] / "pyproject.toml"
_PILLOW_RE = re.compile(r"^\s*\"?(pillow|pillow-simd)\b", re.IGNORECASE | re.MULTILINE)


def test_pillow_is_still_absent_from_the_dependency_manifest():
    """这条不是测代码，是把「为什么不写图」的前提钉住：Pillow 一旦进了依赖清单，
    这条会红，届时再决定要不要把 16 张图写回 xlsx。

    ⚠ 判据必须是**依赖清单**，不是运行环境。原实现写的是
    `with pytest.raises(ImportError): __import__("PIL")` —— 那断言的是"这台机器上
    没装 Pillow"。实测任何全局 site-packages 里有 Pillow 的开发机跑 pytest 都会红
    这一条（本机复现：`1 failed ... DID NOT RAISE ImportError`），于是"全绿"这个
    信号在本地长期不可用；而 CI 容器里恰好没装，所以它一直是绿的 —— 一条只在
    别人机器上红的哨兵，红了也没人信。它真正想守的是清单有没有新增 Pillow，
    那件事从来没被破坏过。
    """
    text = _REQUIREMENTS.read_text(encoding="utf-8")
    assert _PILLOW_RE.search(text) is None, (
        "Pillow 进了 pyproject.toml —— 回来决定要不要把 16 张定额图写回 xlsx"
        "（见 test_dispimg_cells_are_skipped）")


# --------------------------------------------------------------- 模板修正


def test_template_fixes_are_baked_into_the_output(book):
    """三处原表缺陷修正必须体现在产物里，不能只在 warning 里说说。"""
    tpl = load_template()
    for fix in tpl["fixes"]:
        if fix["kind"] != "formula":
            continue
        for ch in fix["changes"]:
            sheet, coord = ch["cell"].split("!", 1)
            assert book[sheet][coord].value == ch["after"], ch["cell"]
            assert book[sheet][coord].value != ch["before"]


def test_f30_points_at_the_right_row(book):
    """F30 下拉复制漏改：原表指 E35（体育专项，与 F28 重复），应指 E46（交通专项）。
    原表两处数值恰好相同，缓存值一致、肉眼不可见，是真缺陷。"""
    assert book["估概算总表"]["F30"].value == "='3.专项工程费'!E46"


def test_sentinel_rows_are_inside_their_parent_sum(book):
    """12 处哨兵行纳入 SUM：此后在哨兵位插入新行会自动进小计。"""
    assert book["2.安装工程费"]["C6"].value == "=SUM(C7:C13)"
    assert book["估概算总表"]["C34"].value == "=SUM(C35:C44)"
    assert book["5.工程建设其他费用"]["C5"].value == "=C6+C9+SUM(C14:C31)"


# --------------------------------------------------------------- 费率独立格


def test_rate_lives_in_its_own_cell(book):
    """spec I-2.6：费率写独立单元格供 C46 引用（原表是公式内字面量 0.05）。
    这样造价员在 OnlyOffice 里改 J46 就能重算，不必去改公式。"""
    ws = book["估概算总表"]
    assert ws["C46"].value == "=(C5+C45)*J46"
    assert ws["J46"].value == 0.05
    assert ws["J45"].value == "预备费率"
    assert ws["J46"].number_format == "0.00%"
    assert "0.05" not in (ws["C46"].value or "")   # 字面量费率必须消失


def test_rate_cell_landing_spot_is_free(book):
    """J 列落点依据：估概算总表 I/J 列零公式、零输入位、I46 只有静态文本 3%~8%。"""
    tpl = load_template()
    assert not [f for f in tpl["formulas"]
                if f["cell"].startswith(("估概算总表!I", "估概算总表!J"))]
    assert not [i for i in tpl["inputs"] if i["cell"].startswith("估概算总表!")]
    assert book["估概算总表"]["I46"].value == "3%~8%"


def test_one_percent_rate_passes_no_hard_gate():
    """铁律 7：预备费率可调不得写死，1% 必须能出表（区间校验只在求值器出 warning）。"""
    wb = load_workbook(BytesIO(render(ROOT_INPUTS, rates={"reserve_rate": 0.01})))
    assert wb["估概算总表"]["J46"].value == 0.01
    assert wb["估概算总表"]["C46"].value == "=(C5+C45)*J46"


def test_non_numeric_rate_is_rejected():
    with pytest.raises(EstimateTemplateError):
        render(ROOT_INPUTS, rates={"reserve_rate": "5%"})
    with pytest.raises(EstimateTemplateError):
        render(ROOT_INPUTS, rates={"reserve_rate": True})   # bool 不是数值


# --------------------------------------------------------------- 输入位


def test_supplied_inputs_are_written_as_numbers(book):
    ws = book["经济技术指标"]
    assert ws["D4"].value == 12000.0
    assert ws["D7"].value == 5200.0
    assert ws["D20"].value == 140.0


def test_unfilled_inputs_are_left_empty_not_zero(book):
    """留空 ≠ 写 0：与求值器「空格按 0」语义一致，也让造价员一眼看见哪些没填。
    写成 0 会让一张半成品概算表看起来像「已确认为零」。"""
    tpl = load_template()
    supplied = set(ROOT_INPUTS)
    blanks = [i["cell"] for i in tpl["inputs"] if i["cell"] not in supplied]
    assert len(blanks) == 200
    for cell in blanks:
        sheet, coord = cell.split("!", 1)
        assert book[sheet][coord].value is None, cell


def test_undeclared_input_keys_are_ignored():
    """渲染器只写声明位；未声明键由求值器负责报 unknown_input，这里不重复报也不乱写。"""
    wb = load_workbook(BytesIO(render(dict(ROOT_INPUTS, **{"经济技术指标!ZZ99": 7.0}))))
    assert wb["经济技术指标"]["ZZ99"].value is None


def test_dollar_signs_in_input_keys_are_normalized():
    wb = load_workbook(BytesIO(render({"经济技术指标!$D$4": 999.0})))
    assert wb["经济技术指标"]["D4"].value == 999.0


# --------------------------------------------------------------- 结构元数据


def test_sheets_kept_in_template_order(book):
    tpl = load_template()
    assert book.sheetnames == [s["name"] for s in tpl["sheets"]]
    assert len(book.sheetnames) == 9


def test_merges_freeze_and_widths_are_restored(book):
    tpl = load_template()
    for s in tpl["sheets"]:
        ws = book[s["name"]]
        assert {str(r) for r in ws.merged_cells.ranges} == set(s["merges"]), s["name"]
        assert (ws.freeze_panes or None) == (s["freeze"] or None), s["name"]
        for col, width in s["col_widths"].items():
            assert ws.column_dimensions[col].width == pytest.approx(width), f"{s['name']}!{col}"


def test_outline_is_force_expanded(book):
    """原表把大量明细行存成 hidden=true，直接还原会让造价员打开就看见一张空壳表。
    实测模板 193 行带 hidden=true —— 全部强制展开。"""
    tpl = load_template()
    hidden_in_template = sum(
        1 for s in tpl["sheets"] for r in s["rows"]
        if (r.get("outline") or {}).get("hidden")
    )
    assert hidden_in_template == 193
    for s in tpl["sheets"]:
        assert s["force_expand_outline"] is True
        ws = book[s["name"]]
        for r in s["rows"]:
            if r.get("outline"):
                dim = ws.row_dimensions[r["row"]]
                assert dim.hidden is False, f"{s['name']} 第 {r['row']} 行仍是折叠隐藏态"
                assert dim.outlineLevel == r["outline"]["level"]


def test_summary_rows_sit_above_their_details(book):
    """原表小计行在明细行之上（31 条 SUM/composite 公式父行行号全部 < 区间起点），
    summaryBelow 必须为 False，否则大纲折叠按钮画在错误的一侧。"""
    for ws in book.worksheets:
        assert ws.sheet_properties.outlinePr.summaryBelow is False, ws.title


# --------------------------------------------------------------- 文本与防注入


def test_seq_is_written_as_text_not_collapsed_number(book):
    """seq 存 varchar：「1.10」当数字会塌缩成 1.1，与真 1.1 重号（种子 DDL 已记在案的事故）。"""
    c = book["2.1估概算安装项目清单"]["A15"]
    assert c.value == "1.10"
    assert isinstance(c.value, str)
    assert book["估概算总表"]["A5"].value == "一"   # 中文序号照抄


def test_static_labels_are_written(book):
    assert book["1.建筑工程费"]["B7"].value == "土石方工程"
    # WS10-P14：模板 static 修正后不再是占位符「*x%」
    assert book["估概算总表"]["B46"].value == "工程预备费（一+二）×预备费率"
    assert book["1.建筑工程费"]["B14"].value == "……"   # 哨兵行占位保留


def test_project_name_is_written_into_row_2(book):
    """⚠ 契约说「替换 {{PROJECT_NAME}} 占位」，但实测模板里没有这个占位——
    9 张表的 title/project 行 static 全空（脱敏时清掉了工程名）。故为写入而非替换。"""
    for name in book.sheetnames:
        assert book[name]["A2"].value == "XX 房屋建筑工程", name


def test_sheet_title_row_is_not_left_nameless(book):
    """模板 9 张表里 8 张的标题行是空的（脱敏时把表名一起清了）。空着交付是一张无名表，
    所以回填 sheet 名——sheet 自己的身份，不是编造内容；模板自带标题的那张不覆盖也不重复。"""
    assert book["估概算总表"]["A1"].value == "估概算总表"
    assert book["1.建筑工程费"]["A1"].value == "1.建筑工程费"
    # 5.工程建设其他费用 的标题在模板里就是 B1 静态文本，A1 保持空，避免重复
    assert book["5.工程建设其他费用"]["B1"].value == "工程建设其他费分析表"
    assert book["5.工程建设其他费用"]["A1"].value is None


def test_project_name_none_leaves_row_2_empty():
    wb = load_workbook(BytesIO(render(ROOT_INPUTS)))
    assert wb["估概算总表"]["A2"].value is None


def test_project_name_formula_injection_is_neutralized():
    """工程名是用户输入。= 开头的文本会被 Excel/WPS 当公式执行，必须加单引号前缀。"""
    wb = load_workbook(BytesIO(render(ROOT_INPUTS, project_name="=cmd|'/c calc'!A1")))
    v = wb["估概算总表"]["A2"].value
    assert v == "'=cmd|'/c calc'!A1"
    assert not v.startswith("=")


# --------------------------------------------------------------- 其他


def test_render_is_deterministic_in_content():
    """同样入参两次渲染，单元格内容必须完全一致（zip 时间戳不比，只比内容）。"""
    a = load_workbook(BytesIO(render(ROOT_INPUTS, project_name="P")))
    b = load_workbook(BytesIO(render(ROOT_INPUTS, project_name="P")))
    for name in a.sheetnames:
        va = [[c.value for c in row] for row in a[name].iter_rows()]
        vb = [[c.value for c in row] for row in b[name].iter_rows()]
        assert va == vb, name


def test_unknown_template_raises():
    with pytest.raises(EstimateTemplateError):
        render(ROOT_INPUTS, template_code="no_such_template")
