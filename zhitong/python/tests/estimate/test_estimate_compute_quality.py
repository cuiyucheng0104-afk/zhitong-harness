# 来源：engine-service/tests/test_estimate_compute_quality.py（zhitong-cost-ai@991ff5cd，W4 迁入 zt；改动见「zt 改动」注释）
"""WS10-P06（issue #761）：compute 的质量结论（review_level）与采纳值过期判定。

改前覆盖率 10.67% 的表与完整概算外观相同：compute 响应、engine_task.result 都没有质量结论，
交付 xlsx 也看不出来。现在 compute 出 `quality{review_level: ok|needs_review|draft, reasons[]}`
（字段名按统一约定 C1：非任务状态叫 review_level；**不新增任务状态**，铁律 5）。

三个分支都要可达（§5.2：用例名覆盖整个概念，夹具就得穷尽）：
  draft         —— 只填 11 个根输入 + 几个成本位（覆盖率 < 30%），且 1 / 2 表一格没填；
  needs_review  —— 满覆盖，但采纳过单样本参考值 / 未申报来源 / 费率采纳值过期 …；
  ok            —— 满覆盖、全部申报为手填、土地费 / 占地都有。
阈值来自模板策略 review_thresholds（D50），用例不另写数字。
"""
from __future__ import annotations

import json
from pathlib import Path

import pytest

from zt.estimate import fee_rules as fr
from tests.estimate.fixtures.endpoint import HEADERS, client  # zt 改动：旧端点的转接

_FX = json.loads((Path(__file__).parent / "fixtures" / "estimate_gsyz_v1_synthetic.json").read_text("utf-8"))
FULL = _FX["inputs"]                       # 11 个根输入 + 150 个成本位全填（合成数据）
ROOT11 = {k: v for k, v in FULL.items() if k.startswith("经济技术指标!")}
C14 = "5.工程建设其他费用!C14"             # 工程监理费（费率规则 input 型、以工程费用为基数）


def _h():
    return HEADERS


def _all_user(inputs: dict) -> dict:
    return {c: {"source": "user"} for c in inputs}


def _compute(**body) -> dict:
    body.setdefault("render", False)
    r = client.post("/engine/estimate/compute", headers=_h(), json=body)
    assert r.status_code == 200, r.text
    return r.json()


def _codes(q: dict) -> list[str]:
    return [x["code"] for x in q["reasons"]]


# ---------------------------------------------------------------- 三个分支

def test_low_coverage_is_draft():
    """只填根输入 + 3 个 3 表成本位：覆盖率 2% 且 1 / 2 表一格没填 → draft。"""
    inputs = dict(ROOT11)
    for c in ("3.专项工程费!C7", "3.专项工程费!C8", "3.专项工程费!C9"):
        inputs[c] = FULL[c]
    q = _compute(inputs=inputs, input_sources=_all_user(inputs))["quality"]
    assert q["review_level"] == "draft"
    assert "coverage_below_draft" in _codes(q)
    assert _codes(q).count("sheet_unfilled") == 2      # 1.建筑工程费、2.安装工程费 各一条


def test_full_coverage_all_declared_is_ok():
    q = _compute(inputs=FULL, input_sources=_all_user(FULL))["quality"]
    assert q == {"review_level": "ok", "reasons": []}


def test_adopted_reference_value_needs_review():
    src = _all_user(FULL)
    src["1.建筑工程费!C7"] = {"source": "baseline_reference", "refs": ["GSYZ-F1-R7"],
                            "semantic_key": "1.建筑工程费/basement/土石方工程"}
    q = _compute(inputs=FULL, input_sources=src)["quality"]
    assert q["review_level"] == "needs_review" and _codes(q) == ["reference_values_adopted"]


def test_undeclared_sources_need_review():
    """老调用方 / 直接打 API：没申报来源 → needs_review（未申报来源）；只申报一部分 → 同码，数出几个。"""
    q = _compute(inputs=FULL)["quality"]
    assert q["review_level"] == "needs_review" and _codes(q) == ["input_sources_undeclared"]
    assert "未申报任何输入位的来源" in q["reasons"][0]["message"]     # 整体没申报与漏几格分开说
    src = _all_user(FULL)
    for c in list(src)[:5]:
        src.pop(c)
    q2 = _compute(inputs=FULL, input_sources=src)["quality"]
    assert _codes(q2) == ["input_sources_undeclared"] and "5 个" in q2["reasons"][0]["message"]


def test_land_cost_missing_and_footprint_zero_need_review():
    inputs = dict(FULL)
    inputs.pop("5.工程建设其他费用!C15")
    inputs["经济技术指标!D15"] = 0
    inputs["经济技术指标!D16"] = 0
    q = _compute(inputs=inputs, input_sources=_all_user(inputs))["quality"]
    assert q["review_level"] == "needs_review"
    assert {"land_cost_unfilled", "footprint_missing"} <= set(_codes(q))


def test_draft_wins_over_needs_review_but_keeps_both_reasons():
    inputs = dict(ROOT11)
    q = _compute(inputs=dict(inputs, **{"3.专项工程费!C7": 100.0}))["quality"]
    assert q["review_level"] == "draft"
    assert "input_sources_undeclared" in _codes(q) and "coverage_below_draft" in _codes(q)


def test_reserve_rate_1_percent_still_computes_whatever_the_quality():
    """铁律 7：质量结论不是闸 —— 1% 预备费率照样算出、200。"""
    body = _compute(inputs=ROOT11 | {"3.专项工程费!C7": 100.0}, rates={"reserve_rate": 0.01})
    assert body["quality"]["review_level"] == "draft"
    assert body["rates_used"]["reserve_rate"] == 0.01
    assert body["totals"]["reserve_fee"]["value"] is not None


# ---------------------------------------------------------------- 采纳的费率值过期

def _supervision_now(project_cost: float) -> float:
    items = {i["item_id"]: i for i in fr.suggest_all({"工程费用": project_cost}, {})["items"]}
    v = items["supervision_fee"]["suggested"]
    assert v is not None, "前置：监理费在该基数下要出值，否则下面比的是 None"
    return v


def test_adopted_fee_value_that_no_longer_matches_is_stale():
    """C14 申报为「费率规则建议（采纳）」，值与按本次工程费用重算的差 2% → fee_value_stale + needs_review。

    变异：删 stale 判定 → 本条转红。对照：值与重算一致（差 0）不报。
    """
    base = _compute(inputs=FULL, input_sources=_all_user(FULL))
    now = _supervision_now(base["totals"]["project_cost"]["value"])

    src = _all_user(FULL)
    src[C14] = {"source": "fee_rules"}
    same = _compute(inputs=dict(FULL, **{C14: now}), input_sources=src)
    assert "fee_value_stale" not in [w["code"] for w in same["warnings"]]
    # 采纳了仓主放行的费率 → 仍是需复核（这是一条独立理由，不是 stale）
    assert _codes(same["quality"]) == ["owner_released_fee_adopted"]

    stale = _compute(inputs=dict(FULL, **{C14: round(now * 1.02, 4)}), input_sources=src)
    w = [w for w in stale["warnings"] if w["code"] == "fee_value_stale"]
    assert len(w) == 1 and w[0]["cell"] == C14
    assert "fee_value_stale" in _codes(stale["quality"])


def test_stale_check_only_looks_at_cells_declared_as_fee_rules():
    """用户自己手填的监理费与规则不同是正常的 —— 不报过期。"""
    body = _compute(inputs=dict(FULL, **{C14: 1.0}), input_sources=_all_user(FULL))
    assert "fee_value_stale" not in [w["code"] for w in body["warnings"]]


# ---------------------------------------------------------------- WS29-n：来源如实

_S1_FEE_CELLS = [f"5.工程建设其他费用!{c}" for c in ("C14", "C16", "C17", "C19", "C28")]
_REF14 = [k for k in FULL if k.startswith("1.建筑工程费!")][:14]


def _reason(q: dict, code: str) -> str:
    hit = [r["message"] for r in q["reasons"] if r["code"] == code]
    assert len(hit) == 1, f"{code} 应恰好一条：{q['reasons']}"
    return hit[0]


def test_fee_rules_and_reference_values_are_counted_separately():
    """08-014 S1：用户采纳 14 条单样本参考值 + 5 条费率建议 →「采纳了 14 条单样本参考值」+ 另一条写 5 条费率建议。

    改前前端把这 5 格申报成 baseline_reference，理由写「采纳了 19 条」、费率那条从来不出现。
    变异：把 fee_rules 也数进参考值 → 第一条断言写 19 → 本条转红。
    """
    assert len(_REF14) == 14
    src = _all_user(FULL)
    for c in _REF14:
        src[c] = {"source": "baseline_reference", "refs": ["GSYZ-F1-R7"]}
    for c in _S1_FEE_CELLS:
        src[c] = {"source": "fee_rules"}
    q = _compute(inputs=FULL, input_sources=src)["quality"]
    assert q["review_level"] == "needs_review"
    assert "采纳了 14 条单样本参考值" in _reason(q, "reference_values_adopted")
    assert "采纳了 5 条费率建议值" in _reason(q, "owner_released_fee_adopted")


def test_baseline_prefill_and_slot_empty_are_declared_and_need_no_review():
    """新来源种类：系统预填的多样本基准值、归位补 0 的空槽位 —— 引擎认（不 422）、算已申报、不触发需复核。

    变异：把 baseline_prefill 也数进「采纳了 N 条单样本参考值」→ review_level 变 needs_review → 本条转红。
    """
    inputs = dict(FULL, **{"经济技术指标!D8": 0})
    src = _all_user(inputs)
    src["经济技术指标!D8"] = {"source": "slot_empty"}
    for c in _REF14:
        src[c] = {"source": "baseline_prefill"}
    q = _compute(inputs=inputs, input_sources=src)["quality"]
    assert q == {"review_level": "ok", "reasons": []}, q


def test_new_source_kinds_render_their_chinese_names():
    """交付表「生成说明 → 逐格来源」：slot_empty 写「单体归位：该槽位无单体」，不再落成「用户手填」（E38）；
    费率格写「费率规则建议（用户采纳）」；预填的多样本基准值写「多样本基准值（系统预填）」。"""
    import base64
    from io import BytesIO

    from openpyxl import load_workbook

    inputs = dict(FULL, **{"经济技术指标!D8": 0})
    src = _all_user(inputs)
    src["经济技术指标!D8"] = {"source": "slot_empty"}
    src[C14] = {"source": "fee_rules"}
    src["1.建筑工程费!C7"] = {"source": "baseline_prefill"}
    body = _compute(inputs=inputs, input_sources=src, render=True, generated_at="2026-10-03 10:00:00")
    ws = load_workbook(BytesIO(base64.b64decode(body["xlsx_base64"])))["生成说明"]
    by_cell = {r[0]: r for r in ws.iter_rows(values_only=True) if r and isinstance(r[0], str) and "!" in r[0]}
    assert by_cell["经济技术指标!D8"][2] == "单体归位：该槽位无单体"
    assert by_cell[C14][2] == "费率规则建议（用户采纳）"
    assert by_cell["1.建筑工程费!C7"][2] == "多样本基准值（系统预填）"


def test_footprint_derived_from_density_needs_review_and_is_named_in_the_notes():
    """WS29-d（D-R118）：占地按 用地面积 × 建筑密度 推算、用户点「填入推算值」写入 → 来源 derived，
    结果至少需复核（footprint_derived），交付表逐格来源写「按建筑密度推算（用户点选）」。

    变异：judge_quality 不看 derived（推算值当成原文数值）→ review_level 变 ok → 本条转红。
    """
    import base64
    from io import BytesIO

    from openpyxl import load_workbook

    inputs = dict(FULL, **{"经济技术指标!D15": 13135.11})
    src = _all_user(inputs)
    src["经济技术指标!D15"] = {"source": "derived", "quote": "建筑密度 30% × 用地面积 43783.7"}
    body = _compute(inputs=inputs, input_sources=src, render=True, generated_at="2026-10-06 10:00:00")
    q = body["quality"]
    assert q["review_level"] == "needs_review" and _codes(q) == ["footprint_derived"], q
    assert "D15" in q["reasons"][0]["message"] and "不是原文数值" in q["reasons"][0]["message"]
    ws = load_workbook(BytesIO(base64.b64decode(body["xlsx_base64"])))["生成说明"]
    by_cell = {r[0]: r for r in ws.iter_rows(values_only=True) if r and isinstance(r[0], str) and "!" in r[0]}
    assert by_cell["经济技术指标!D15"][2] == "按建筑密度推算（用户点选）"


# ---------------------------------------------------------------- 新入参的 422

@pytest.mark.parametrize("extra,needle", [
    ({"input_sources": {"1.建筑工程费!C7": {"source": "guess"}}}, "source"),
    ({"input_sources": {"C7": {"source": "user"}}}, "表名!格号"),
    ({"fee_bases": {"工程费用": True}}, "fee_bases"),
    ({"fee_attrs": {"project_city": 1}}, "fee_attrs"),
    ({"project_profile": {"building_function": "仓储\n=cmd"}}, "project_profile"),
    ({"generated_at": "昨天下午"}, "generated_at"),
])
def test_bad_output_context_is_422(extra, needle):
    r = client.post("/engine/estimate/compute", headers=_h(),
                    json={"inputs": FULL, "render": False, **extra})
    assert r.status_code == 422, r.text
    assert needle in r.json()["detail"]


def test_compute_response_carries_aggregate_checks():
    """P11 在 compute 复用同一函数：合成数据的 C10-C13 远超小计 → mismatch 进响应与告警。"""
    body = _compute(inputs=FULL, input_sources=_all_user(FULL), fee_bases={})
    assert any(c["status"] == "mismatch" for c in body["aggregate_checks"])
    assert "fee_aggregate_mismatch" in [w["code"] for w in body["warnings"]]


# ---------------------------------------------------------------- WS29-q：整表没填点名、参照造价、指标表单体名

from zt.estimate import checks as _ec  # noqa: E402
from zt.estimate import policy as _ep  # noqa: E402
from zt.estimate import quality as _eq  # noqa: E402

_POL = _ep.load_policy("gsyz_v1")


def _without_sheet(sheet: str) -> dict:
    cells = {c for s, c in _ec.cost_input_cells("gsyz_v1") if s == sheet}
    assert len(cells) > 10, f"前置：{sheet} 要有成本位，实为 {len(cells)}"
    return {k: v for k, v in FULL.items() if k not in cells}


def test_special_sheet_unfilled_is_needs_review_and_names_the_sheet():
    """08-014 第 6 项（E39）：专项工程费 42 格全空、其它表有值 → 改前 ok（横幅写「出数完整」）。

    U18：仍是需复核，理由里点名「3.专项工程费」0/42；不降成草稿。
    变异：删掉 sheet_unfilled_review 那段 → review_level 回到 ok → 本条转红。
    """
    inputs = _without_sheet("3.专项工程费")
    q = _compute(inputs=inputs, input_sources=_all_user(inputs))["quality"]
    assert q["review_level"] == "needs_review", q
    msg = _reason(q, "sheet_unfilled_review")
    assert "「3.专项工程费」0/42" in msg and "按 0 计入" in msg
    for other in ("1.建筑工程费", "2.安装工程费", "4.室外及配套工程费", "5.工程建设其他费用"):
        assert other not in msg, f"有值的表不许被点名：{other}"


def test_building_sheet_unfilled_is_still_draft_not_review():
    """建筑工程费 0 格 → 仍是草稿（原规则不变），不重复记成需复核的整表未填。"""
    inputs = _without_sheet("1.建筑工程费")
    q = _compute(inputs=inputs, input_sources=_all_user(inputs))["quality"]
    assert q["review_level"] == "draft" and "sheet_unfilled" in _codes(q)
    assert "sheet_unfilled_review" not in _codes(q)


def test_two_review_sheets_unfilled_are_named_in_one_reason():
    inputs = {k: v for k, v in _without_sheet("3.专项工程费").items()
              if k not in {c for s, c in _ec.cost_input_cells("gsyz_v1") if s == "4.室外及配套工程费"}}
    q = _compute(inputs=inputs, input_sources=_all_user(inputs))["quality"]
    msg = _reason(q, "sheet_unfilled_review")
    assert "「3.专项工程费」0/42" in msg and "「4.室外及配套工程费」0/18" in msg


def test_b8_derived_footprint_keeps_empty_sheet_and_reference_reasons_together():
    """B8 d+q：推算占地、整表空缺与参照偏离同时存在时三种复核理由都要保留，参照不改合计。"""
    inputs = dict(_without_sheet("3.专项工程费"), **{"经济技术指标!D15": 13135.11})
    src = _all_user(inputs)
    src["经济技术指标!D15"] = {"source": "derived", "quote": "建筑密度 30% × 用地面积 43783.7"}
    plain = _compute(inputs=inputs, input_sources=src)
    body = _compute(inputs=inputs, input_sources=src,
                    project_profile={"reference_cost": {"amount": round(_static(plain) / 1.65, 2)}})
    q = body["quality"]
    assert q["review_level"] == "needs_review"
    assert {"footprint_derived", "sheet_unfilled_review", "reference_cost_deviation"} <= set(_codes(q))
    assert "不是原文数值" in _reason(q, "footprint_derived")
    assert "3.专项工程费" in _reason(q, "sheet_unfilled_review")
    assert "1.65 倍" in _reason(q, "reference_cost_deviation")
    assert body["totals"] == plain["totals"]


def test_reference_comparison_ratio_and_band():
    """U19：参照 31,837.59（招标控制价）、静态总投资 52,495.91 → 1.65 倍，记需复核；参照 50,000 → 1.05 倍，不记。

    变异：比值阈值写反（> lo / < hi）→ 1.65 不报、1.05 报 → 本条转红。下限一侧也要够得到：0.6 倍要报。
    """
    hit = _eq.reference_comparison({"amount": 31837.59, "note": "招标控制价"}, 52495.91, _POL)
    assert hit["deviates"] is True and hit["ratio"] == round(52495.91 / 31837.59, 4)
    assert hit["message"] == "建设静态总投资 52,495.91 万元，是参照造价 31,837.59 万元（招标控制价）的 1.65 倍"
    ok = _eq.reference_comparison({"amount": 50000}, 52495.91, _POL)
    assert ok["deviates"] is False and ok["message"].endswith("是参照造价 50,000.00 万元的 1.05 倍")
    low = _eq.reference_comparison({"amount": 100000}, 60000, _POL)
    assert low["deviates"] is True and "0.60 倍" in low["message"]
    assert _eq.reference_comparison(None, 52495.91, _POL) is None
    assert _eq.reference_comparison({"amount": 0}, 52495.91, _POL) is None


def _static(body: dict) -> float:
    return body["totals"]["static_total"]["value"]


def test_reference_cost_flows_through_compute_without_touching_totals():
    """compute 收 project_profile.reference_cost：偏离 → reference_cost_deviation + needs_review；
    不偏离 → 只带比值不记理由；不填 → 没有这一项、也不报错。三种情况 totals 逐项相同（参照不进公式）。"""
    src = _all_user(FULL)
    plain = _compute(inputs=FULL, input_sources=src)
    assert plain["quality"] == {"review_level": "ok", "reasons": []}
    static = _static(plain)

    far = _compute(inputs=FULL, input_sources=src,
                   project_profile={"reference_cost": {"amount": round(static / 1.65, 2), "note": "招标控制价"}})
    assert far["quality"]["review_level"] == "needs_review"
    msg = _reason(far["quality"], "reference_cost_deviation")
    assert "1.65 倍" in msg and "（招标控制价）" in msg and "不拦" in msg
    assert far["quality"]["reference_comparison"]["deviates"] is True

    near = _compute(inputs=FULL, input_sources=src,
                    project_profile={"reference_cost": {"amount": round(static / 1.05, 2)}})
    assert near["quality"]["review_level"] == "ok" and near["quality"]["reasons"] == []
    rc = near["quality"]["reference_comparison"]
    assert rc["deviates"] is False and "1.05 倍" in rc["message"]

    for other in (far, near):
        assert other["totals"] == plain["totals"], "参照造价不许进任何公式"
        assert other["sheets"] == plain["sheets"]


@pytest.mark.parametrize("ref,needle", [
    ({"amount": True}, "reference_cost.amount"),
    ({"amount": -1}, "reference_cost.amount"),
    ({"amount": "31837"}, "reference_cost.amount"),
    ({"amount": 100, "note": "招标\n控制价"}, "reference_cost.note"),
    ("31837.59", "reference_cost"),
])
def test_bad_reference_cost_is_422(ref, needle):
    r = client.post("/engine/estimate/compute", headers=_h(),
                    json={"inputs": FULL, "render": False, "project_profile": {"reference_cost": ref}})
    assert r.status_code == 422, r.text
    assert needle in r.json()["detail"]


def test_indicator_rows_use_slot_labels():
    """08-014 第 5 项：结果页「经济技术指标」表的单体行仍是模板名，xlsx 里已改。indicators 套用 slot_labels。

    变异：不套 → 3.1 行仍是「A区地下室」→ 本条转红。没起名的槽位与非槽位行不动。
    """
    labels = {"经济技术指标!D7": "地下室", "经济技术指标!D10": "住宅楼（19 栋）"}
    body = _compute(inputs=FULL, input_sources=_all_user(FULL), slot_labels=labels)
    by_seq = {i["seq"]: i for i in body["indicators"]}
    assert by_seq["3.1"]["label"] == "地下室"
    assert by_seq["4.1"]["label"] == "住宅楼（19 栋）"
    assert by_seq["3.2"]["label"] == "B区地下室" and by_seq["4.2"]["label"] == "2#楼"
    assert by_seq["1"]["label"] == "用地面积"
    plain = _compute(inputs=FULL, input_sources=_all_user(FULL))
    assert [i["value"] for i in body["indicators"]] == [i["value"] for i in plain["indicators"]]
    assert {i["seq"]: i["label"] for i in plain["indicators"]}["3.1"] == "A区地下室"


def test_notes_sheet_writes_reference_comparison():
    """交付表「生成说明」页写参照对比一行；画像里不出现 reference_cost 的对象原文。"""
    import base64
    from io import BytesIO

    from openpyxl import load_workbook

    src = _all_user(FULL)
    static = _static(_compute(inputs=FULL, input_sources=src))
    body = _compute(inputs=FULL, input_sources=src, render=True, generated_at="2026-10-06 10:00:00",
                    project_profile={"project_city": "zhejiang_other",
                                     "reference_cost": {"amount": round(static / 1.65, 2), "note": "招标控制价"}})
    ws = load_workbook(BytesIO(base64.b64decode(body["xlsx_base64"])))["生成说明"]
    rows = [r for r in ws.iter_rows(values_only=True) if r and r[0]]
    by_head = {r[0]: r for r in rows}
    assert "■ 参照造价对比" in by_head
    assert by_head["参照造价"][3] == "招标控制价"
    assert "1.65 倍" in by_head["参照对比"][3] and "不参与计算" in by_head["参照对比"][3]
    assert not any("reference_cost" in str(c) for r in rows for c in r if c), "画像里不许写对象原文"
    plain = _compute(inputs=FULL, input_sources=src, render=True, generated_at="2026-10-06 10:00:00")
    ws2 = load_workbook(BytesIO(base64.b64decode(plain["xlsx_base64"])))["生成说明"]
    assert not any(r and r[0] == "■ 参照造价对比" for r in ws2.iter_rows(values_only=True)), "没填参照就不出这一节"
