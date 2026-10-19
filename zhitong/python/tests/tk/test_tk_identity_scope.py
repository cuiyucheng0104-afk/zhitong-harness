# 来源：engine-service/tests/test_tk_identity_scope.py（zhitong-cost-ai@991ff5cd）
"""WS08-08 / WS08-09（issue #760）：投控异常身份含文件与行 + 结果回源定位。

修的现场：
  · 身份键与折叠键只有 project|major|sheet|item_code —— 三份进度款合并检测只剩 44 条，分别检测是 84 条；
    两份不同合同的清单行共用同一个异常 uid，复核结论会串到别的期次；
  · 同一个键被多条规则命中时 `>` 比较、先到先得，结果取决于输入顺序；
  · 异常丢了 row 与文件：C1 的 416 条异常只对应 116 个名称，回不到原表。
"""
from __future__ import annotations

import io
import random

from openpyxl import Workbook

from zt.tk import engine as tk_engine

HEAD = ["序号", "项目编码", "项目名称", "项目特征描述", "计量单位", "工程量", "综合单价", "合价"]


def _book(sheet: str, rows: list[list], unit_project: str | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    if unit_project:
        ws.append([f"单位工程名称：{unit_project}", None, None, None, None, None, None, "第 1 页"])
    ws.append(HEAD)
    for i, r in enumerate(rows, 1):
        ws.append([i] + list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# 两个标段：sheet 同名、同一编码、金额都过大额线（20000）—— 旧身份键下两份文件的这一行折成一条
LOT1 = _book("分部分项清单与计价表", [["010101001001", "平整场地", "", "m2", 1000, 30, 30000],
                                      ["010401003001", "实心砖墙", "", "m3", 100, 480, 48000]])
LOT2 = _book("分部分项清单与计价表", [["010101001001", "平整场地", "", "m2", 900, 30, 27000],
                                      ["010502001001", "矩形柱", "", "m3", 50, 600, 30000]])


def _run(files):
    return tk_engine.run(files, None, project_id="P")


def test_merged_detection_count_equals_sum_of_separate_detections():
    one = _run([(LOT1, "标段一进度款.xlsx", "F1")])["anomalies"]
    two = _run([(LOT2, "标段二进度款.xlsx", "F2")])["anomalies"]
    both = _run([(LOT1, "标段一进度款.xlsx", "F1"), (LOT2, "标段二进度款.xlsx", "F2")])["anomalies"]
    assert len(one) == 2 and len(two) == 2, "夹具失效：每份文件应各有 2 条大额异常"
    assert len(both) == len(one) + len(two), (
        f"合并检测 {len(both)} 条 ≠ 分别检测 {len(one)} + {len(two)} 条 —— 两份文件的同码行被折成了一条")
    assert {a["anomaly_id"] for a in both} == {a["anomaly_id"] for a in one + two}


def test_same_code_in_two_files_gets_two_uids_and_rerun_is_stable():
    both = _run([(LOT1, "标段一进度款.xlsx", "F1"), (LOT2, "标段二进度款.xlsx", "F2")])["anomalies"]
    same_code = [a for a in both if a["item_code"] == "010101001001"]
    assert len(same_code) == 2 and len({a["anomaly_id"] for a in same_code}) == 2
    again = _run([(LOT1, "标段一进度款.xlsx", "F1"), (LOT2, "标段二进度款.xlsx", "F2")])["anomalies"]
    assert sorted(a["anomaly_id"] for a in both) == sorted(a["anomaly_id"] for a in again), "同输入跑两次 uid 必须相同"


def test_file_key_falls_back_to_file_name_without_file_id():
    a = _run([(LOT1, "标段一进度款.xlsx")])["anomalies"]
    b = _run([(LOT1, "标段一进度款-副本.xlsx")])["anomalies"]
    assert {x["anomaly_id"] for x in a}.isdisjoint({x["anomaly_id"] for x in b})


def test_two_rows_same_code_same_sheet_are_two_identities():
    raw = _book("分部分项", [["010101001001", "平整场地", "", "m2", 1000, 30, 30000],
                            ["010101001001", "平整场地", "", "m2", 1000, 30, 30000]])
    anomalies = _run([(raw, "第5期.xlsx", "F1")])["anomalies"]
    dup = [a for a in anomalies if a["subtype"] == "duplicate"]
    assert len({a["anomaly_id"] for a in dup}) == 2 and sorted(a["row"] for a in dup) == [2, 3]


def test_shuffled_input_gives_identical_result():
    rows = []
    for f in ("F1", "F2"):
        for i in range(6):
            rows.append({"project_id": "P", "file_id": f, "source_file": f + ".xlsx", "sheet": "S", "row": i + 2,
                         "item_code": f"01010100{i:04d}", "item_name": f"项{i}", "item_feature": "", "unit": "m3",
                         "submit_qty": 10, "submit_price": 3000 + i, "submit_amount": 29000 + i * 1000})
    cfg = tk_engine.load_rules()
    expect = [(a["anomaly_id"], a["subtype"]) for a in tk_engine.detect(rows, cfg)]
    for seed in range(5):
        shuffled = rows[:]
        random.Random(seed).shuffle(shuffled)
        got = [(a["anomaly_id"], a["subtype"]) for a in tk_engine.detect(shuffled, cfg)]
        assert got == expect


def test_equal_confidence_tie_is_broken_by_severity_then_rule_name():
    """self_check 与 new_item 置信度同为 0.55、severity 同为送审金额 → 规则名升序，new_item 胜出。
    旧版严格大于、先到先得：self_check 先进 hits，永远是它，与置信度表是否平手无关。"""
    cfg = tk_engine.load_rules()
    assert cfg["confidence"]["self_check"] == cfg["confidence"]["new_item"], "夹具失效：两条规则置信度不再平手"
    base = [{"project_id": "P", "item_code": f"01050100100{i}", "item_name": f"基准{i}", "unit": "m3",
             "base_qty": 1, "base_price": 1, "base_amount": 1} for i in range(2)]
    aligned = [{"project_id": "P", "file_id": "F", "sheet": "S", "row": 10 + i, "item_code": f"01050100100{i}",
                "item_name": f"基准{i}", "item_feature": "", "unit": "m3", "submit_qty": 1, "submit_price": 1,
                "submit_amount": 1} for i in range(2)]
    target = {"project_id": "P", "file_id": "F", "sheet": "S", "row": 2, "item_code": "010999999001",
              "item_name": "签证增补", "item_feature": "", "unit": "m3",
              "submit_qty": 10, "submit_price": 100, "submit_amount": 9000}
    hits = [a for a in tk_engine.detect([target] + aligned, cfg, base) if a["row"] == 2]
    assert len(hits) == 1 and hits[0]["subtype"] == "new_item"


def test_every_anomaly_carries_file_sheet_row_and_unit_project():
    raw = _book("土建", [["010101001001", "平整场地", "", "m2", 1000, 30, 30000]], unit_project="6#楼【建筑】")
    anomalies = _run([(raw, "第5期.xlsx", "F-6")])["anomalies"]
    assert anomalies
    for a in anomalies:
        assert a["file_id"] == "F-6" and a["source_file"] == "第5期.xlsx"
        assert a["sheet"] == "土建" and a["row"] == 3
        assert a["unit_project"] == "6#楼【建筑】"
        assert a["locator"]["row"] == 3 and a["locator"]["file_id"] == "F-6"
        assert "第5期.xlsx·土建·第3行" in a["detail"]


def test_duplicate_grouping_uses_unit_project_not_major():
    """同一份文件里两栋楼（两张表，单位工程不同）的同一条清单不是重复计取。
    两栋的工程量不同（金额不同）—— 金额相同的话 costsheet 的跨表去重会先把第二行吃掉，
    用例就打不到分组键这一层了。"""
    wb = Workbook()
    for i, up in enumerate(("6#楼", "7#楼")):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = up
        ws.append([f"单位工程名称：{up}"])
        ws.append(HEAD)
        ws.append([1, "010101001001", "平整场地", "", "m2", 100 + i * 10, 30, (100 + i * 10) * 30])
    buf = io.BytesIO()
    wb.save(buf)
    anomalies = _run([(buf.getvalue(), "第5期.xlsx", "F1")])["anomalies"]
    assert [a for a in anomalies if a["subtype"] in ("duplicate", "code_repeat")] == []
