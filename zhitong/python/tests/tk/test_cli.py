"""python -m zt.tk.cli detect / review / progress。"""
from __future__ import annotations

import json

from openpyxl import load_workbook

from tests.costsheet.cli_helpers import run_cli
from tests.costsheet.fixtures.golden import FIXTURES


def _detect(out):
    code, summary = run_cli("zt.tk.cli", "detect", "--files", str(FIXTURES / "送审-异常样例.xlsx"),
                            "--project-id", "2026-007", "--out", f"{out}/")
    assert code == 0 and summary["status"] == "ok"
    return json.loads((out / "result.json").read_text(encoding="utf-8"))


def test_detect_review_progress_round_trip(tmp_path):
    out = tmp_path / "工作" / "2026-10-19-投控"
    result = _detect(out)
    anomalies = result["data"]["anomalies"]
    assert anomalies and all(a["review_status"] == "pending" and a["project_id"] == "2026-007" for a in anomalies)
    assert result["data"]["rules_version"] == "v1-20260711"
    wb = load_workbook(out / "投控异常.xlsx", read_only=True)
    assert len(list(wb["异常清单"].iter_rows())) == 1 + len(anomalies)

    log = out / "复核.jsonl"
    aid = anomalies[0]["anomaly_id"]
    code, env = run_cli("zt.tk.cli", "review", "--result", str(out / "result.json"), "--log", str(log),
                        "--anomaly-id", aid, "--state", "dismissed", "--reason", "同价分行，非重复计取",
                        "--reviewer", "张工")
    assert code == 0 and env["data"]["record"]["from_state"] == "pending"
    assert env["data"]["progress"]["dismissed"] == 1
    code, env = run_cli("zt.tk.cli", "review", "--result", str(out / "result.json"), "--log", str(log),
                        "--anomaly-id", aid, "--state", "confirmed", "--reason", "复核后改判")
    assert env["data"]["record"]["from_state"] == "dismissed"
    code, env = run_cli("zt.tk.cli", "progress", "--result", str(out / "result.json"), "--log", str(log))
    p = env["data"]["progress"]
    assert (p["confirmed"], p["dismissed"], p["pending"]) == (1, 0, len(anomalies) - 1)
    assert aid not in p["pending_ids"]
    assert len(log.read_text(encoding="utf-8").splitlines()) == 2, "只追加"


def test_review_unknown_anomaly_is_an_error(tmp_path):
    out = tmp_path / "投控"
    _detect(out)
    code, err = run_cli("zt.tk.cli", "review", "--result", str(out / "result.json"), "--log", str(out / "复核.jsonl"),
                        "--anomaly-id", "nope", "--state", "confirmed", "--reason", "x")
    assert code == 2 and err["error"]["code"] == "anomaly_not_found"
    assert not (out / "复核.jsonl").exists()


def test_detect_with_custom_rules(tmp_path):
    import yaml
    from zt.tk.engine import RULES_PATH
    cfg = yaml.safe_load(RULES_PATH.read_text(encoding="utf-8"))
    cfg["version"] = "test-rules"
    cfg["rules"]["big_amount"]["enabled"] = False
    rules = tmp_path / "rules.yaml"
    rules.write_text(yaml.safe_dump(cfg, allow_unicode=True), encoding="utf-8")
    code, env = run_cli("zt.tk.cli", "detect", "--files", str(FIXTURES / "送审-异常样例.xlsx"), "--rules", str(rules))
    assert code == 0 and env["data"]["rules_version"] == "test-rules"
    assert env["data"]["rules_skipped"]["big_amount"] == "rule_disabled"
    assert not [a for a in env["data"]["anomalies"] if a["type"] == "big_amount"]


def test_detect_missing_rules_file(tmp_path):
    code, err = run_cli("zt.tk.cli", "detect", "--files", str(FIXTURES / "送审-异常样例.xlsx"),
                        "--rules", str(tmp_path / "无.yaml"))
    assert code == 2 and err["error"]["code"] == "input_not_found"
