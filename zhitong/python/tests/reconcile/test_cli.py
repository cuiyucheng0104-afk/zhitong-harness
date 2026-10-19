"""python -m zt.reconcile.cli run：stdout 一个 JSON；--out 目录写 result.json 与 核对结果.xlsx。"""
from __future__ import annotations

import json

from openpyxl import load_workbook

from tests.costsheet.cli_helpers import run_cli
from tests.costsheet.fixtures.golden import FIXTURES


def test_run_writes_result_and_workbook(tmp_path):
    out = tmp_path / "工作" / "2026-10-19-核对"
    out.mkdir(parents=True)
    code, summary = run_cli("zt.reconcile.cli", "run", "--submitted", str(FIXTURES / "送审-计价表.xlsx"),
                            "--approved", str(FIXTURES / "审定-计价表.xlsx"), "--out", f"{out}/")
    assert code == 0 and summary["status"] == "ok"
    result = json.loads((out / "result.json").read_text(encoding="utf-8"))
    assert result["data"]["total_diff"] == 18500.0 and result["needs_manual"] is True
    xlsx = out / "核对结果.xlsx"
    assert result["data"]["artifacts"]["xlsx"] == xlsx.as_posix()
    wb = load_workbook(xlsx, read_only=True)
    assert wb.sheetnames == ["汇总", "逐行核对", "按章汇总"]
    rows = list(wb["逐行核对"].iter_rows(values_only=True))
    assert rows[0][0] == "状态" and len(rows) == 1 + len(result["data"]["rows"])
    summary_rows = dict(r[:2] for r in wb["汇总"].iter_rows(min_row=2, values_only=True))
    assert summary_rows["核减额(正=审减)"] == 18500.0


def test_run_without_out_prints_full_envelope_and_no_workbook(tmp_path):
    code, env = run_cli("zt.reconcile.cli", "run", "--submitted", str(FIXTURES / "送审-计价表.xlsx"),
                        "--approved", str(FIXTURES / "审定-计价表.xlsx"), "--basis", "contract", cwd=tmp_path)
    assert code == 0 and env["data"]["summary"]["basis"] == "contract"
    assert "artifacts" not in env["data"] and not list(tmp_path.iterdir())


def test_request_file_and_flags(tmp_path):
    req = tmp_path / "request.json"
    req.write_text(json.dumps({"submitted_file": str(FIXTURES / "送审-计价表.xlsx"),
                               "approved_file": {"path": str(FIXTURES / "审定-计价表.xlsx"), "file_id": "资料/审定.xlsx"}},
                              ensure_ascii=False), encoding="utf-8")
    code, env = run_cli("zt.reconcile.cli", "run", "--in", str(req))
    assert code == 0 and env["data"]["file_coverage"][1]["file_id"] == "资料/审定.xlsx"


def test_xls_is_insufficient_not_an_error(tmp_path):
    code, env = run_cli("zt.reconcile.cli", "run", "--submitted", str(tmp_path / "送审.xls"),
                        "--approved", str(FIXTURES / "审定-计价表.xlsx"), "--out", f"{tmp_path}/")
    assert code == 0 and (env["status"], env["status_reason"]) == ("insufficient", "xls_unsupported")
    assert not (tmp_path / "核对结果.xlsx").exists()


def test_missing_file_exits_2_with_error_json(tmp_path):
    code, out = run_cli("zt.reconcile.cli", "run", "--submitted", str(tmp_path / "无.xlsx"),
                        "--approved", str(FIXTURES / "审定-计价表.xlsx"))
    assert code == 2 and out["error"]["code"] == "input_not_found"


def test_bad_basis_exits_2():
    code, out = run_cli("zt.reconcile.cli", "run", "--submitted", "a.xlsx", "--approved", "b.xlsx",
                        "--basis", "本期")
    assert code == 2 and out["error"]["code"] == "usage"


def test_unparseable_file_is_insufficient(tmp_path):
    code, env = run_cli("zt.reconcile.cli", "run", "--submitted", str(FIXTURES / "不可解析-样例.xlsx"),
                        "--approved", str(FIXTURES / "审定-计价表.xlsx"))
    assert code == 0 and env["status"] == "insufficient"
    assert "送审文件" in env["data"]["detail"]["message"]
