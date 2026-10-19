# 来源：engine-service/tests/test_reconcile_same_file.py（zhitong-cost-ai@991ff5cd）
"""WS26-b（issue #800，spec P05 / D-R78）：对账两侧取到同一份文件的同一列金额时拦截。

修的现场：送审、审定都选同一份招标控制价（单列合价表）→ 200、全部「一致」、核减额 0、没有任何告警。
旧护栏 amounts_identical_suspect 的前提是「表头有 ≥2 组金额列」，单列合价表不触发。

判定放在引擎（只有引擎知道两侧各取了哪一列）：两份原始字节 sha256 相同 → 逐表比**金额列下标**。
  全部计入的表两侧同列           → ParseFailed（same_file_same_columns），端点 422
  部分表同列                     → 这些表两侧都不参与，error 告警 same_file_sheets_excluded
  全部表两侧不同列（审核对比表） → 正常出结果，info 告警 same_file_two_sides
两份不同文件、所有行逐分相同     → warn 告警 all_rows_identical（可能合法，只提示）

夹具按生产版式造（本项目「…一标段.xlsx」的表10.2.2-16：表头三行 金额(元) → 综合单价 | 合价 | 其中 →
人工费 | 机械费 | 暂估价；同簿带费用表与措施综合单价计算表，它们按标题跳过、不算「计入的表」）。
名称与金额全部自编，期望值由夹具构造直接算出，不从引擎输出反抄。
"""
from __future__ import annotations

import io

import pytest
from openpyxl import Workbook

from tests.costsheet.fixtures.endpoint import HEADERS, client
from zt.reconcile import engine as reconcile_engine
from zt.costsheet.parser import parse_costsheet
from tests.costsheet.fixtures import costsheet_shapes as shapes


SHEET_EARTH = "表10.2.2-16 分部分项工程清单与计价表【土石方工程】"
SHEET_PIT = "表10.2.2-16 分部分项工程清单与计价表【基坑支护】"
# 单列合价表：H 列（下标 7）是「金额(元)/合价」
AMOUNT_IDX = 7

EARTH_ROWS = [  # 编码, 名称, 单位, 量, 价, 合价
    ("010101001001", "平整场地", "m2", 1000, 2.5, 2500.0),
    ("010101002001", "挖一般土方", "m3", 2000, 12.0, 24000.0),
    ("010103001001", "回填方", "m3", 800, 20.0, 16000.0),
]                                                     # 合计 42500
PIT_ROWS = [
    ("010202001001", "地下连续墙", "m3", 100, 800.0, 80000.0),
    ("010202009001", "喷射混凝土", "m2", 500, 60.0, 30000.0),
]                                                     # 合计 110000


def _save(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _single_amount_sheet(ws, unit_project: str, rows: list, stamp: str) -> None:
    """表10.2.2-16 同形：三行表头，只有一列合价。"""
    ws.append(["分部分项工程清单与计价表"])
    ws.append([f"单位(专业)工程名称：{unit_project}", None, None, None, None, None, None, None, None,
               f"编制日期 {stamp}", "第1页 共1页"])
    ws.append(["序号", "项目编码", "项目名称", "项目特征", "计量单位", "工程量", "金额(元)", None, None, None, None,
               "备注"])
    ws.append([None, None, None, None, None, None, "综合单价", "合价", "其中", None, None, None])
    ws.append([None, None, None, None, None, None, None, None, "人工费", "机械费", "暂估价", None])
    for i, (code, name, unit, qty, price, amount) in enumerate(rows, 1):
        ws.append([i, code, name, "详见图纸", unit, qty, price, amount, round(amount * 0.2, 2),
                   round(amount * 0.1, 2), None, None])
    ws.append([None, None, "合计", None, None, None, None, sum(r[5] for r in rows), None, None, None, None])


def control_price(stamp: str = "2026-09-01", pit_rows: list | None = None) -> bytes:
    """招标控制价：两张表10.2.2-16（单列合价）+ 一张费用表 + 一张措施综合单价计算表（后两张按标题跳过）。
    stamp 只写进表头上方的「编制日期」格 —— 清单行完全相同、字节不同，即「两份不同的文件」。
    真值：土石方 42500 + 基坑支护 110000 = 152500。"""
    wb = Workbook()
    fee = wb.active
    fee.title = "表10.2.2-13 单位工程招标控制价费用表【土石方工程】"
    fee.append(["单位工程招标控制价费用表"])
    fee.append(["序号", "费用名称", "金额(元)"])
    fee.append([1, "分部分项工程费", 42500])
    fee.append([None, "合计", 42500])
    _single_amount_sheet(wb.create_sheet(SHEET_EARTH), "土石方工程", EARTH_ROWS, stamp)
    _single_amount_sheet(wb.create_sheet(SHEET_PIT), "基坑支护", PIT_ROWS if pit_rows is None else pit_rows, stamp)
    ms = wb.create_sheet("表10.2.2-17 施工技术措施项目综合单价计算表【土石方】")
    ms.append(["施工技术措施项目综合单价计算表"])
    ms.append(["序号", "项目编码", "项目名称", "计量单位", "数量", "综合单价", "合计"])
    ms.append([1, "011706002001", "排水、降水", "项", 1, 18000, 18000])
    return _save(wb)


SHEET_AUDIT = "分部分项工程结算审核对比表【土石方工程】"
AUDIT_ROWS = [  # 编码, 名称, 单位, 送审量, 送审价, 送审合价, 审定量, 审定价, 审定合价
    ("010101002001", "挖一般土方", "m3", 1000, 12.5, 12500.0, 900, 12.5, 11250.0),    # 核减 1250
    ("010103001001", "回填方", "m3", 800, 20.0, 16000.0, 800, 18.0, 14400.0),         # 核减 1600
    ("010101004001", "挖基坑土方", "m3", 200, 30.0, 6000.0, 200, 30.0, 6000.0),       # 核减 0
]                                                     # 送审 34500，审定 31650，核减 2850


def _audit_compare_sheet(ws) -> None:
    """明细版式的审核对比表：同一张表里送审、审定两组「工程量 | 金额(元) → 综合单价 | 合价」，三行表头。
    送审合价在 H 列（下标 7），审定合价在 K 列（下标 10）。"""
    ws.append(["分部分项工程结算审核对比表"])
    ws.append(["单位(专业)工程名称：土石方工程"])
    ws.append(["序号", "项目编码", "项目名称", "项目特征", "计量单位", "送审", None, None, "审定", None, None,
               "核减金额(元)"])
    ws.append([None, None, None, None, None, "工程量", "金额(元)", None, "工程量", "金额(元)", None, None])
    ws.append([None, None, None, None, None, None, "综合单价", "合价", None, "综合单价", "合价", None])
    for i, (code, name, unit, sq, sp, sa, aq, ap, aa) in enumerate(AUDIT_ROWS, 1):
        ws.append([i, code, name, "详见图纸", unit, sq, sp, sa, aq, ap, aa, sa - aa])
    ws.append([None, None, "合计", None, None, None, None, 34500.0, None, None, 31650.0, 2850.0])


def audit_compare_detail() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_AUDIT
    _audit_compare_sheet(ws)
    return _save(wb)


def audit_summary() -> bytes:
    """审核汇总版式（costsheet._audit_rows）：各专业一行，报审 | 审核 | 核减 三组金额，没有编码/单位/量价。
    报审在 C 列（下标 2）、审核在 D 列（下标 3）。真值：报审 1500000，审核 1430000，核减 70000。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "进度款审核汇总表"
    ws.append(["工程进度款审核汇总表"])
    ws.append(["序号", "工程名称", "报审金额(元)", "审核金额(元)", "核减金额(元)", "备注"])
    ws.append([1, "土建工程", 1000000, 950000, 50000, None])
    ws.append([2, "安装工程", 500000, 480000, 20000, None])
    ws.append([None, "合计", 1500000, 1430000, 70000, None])
    return _save(wb)


def qty_check_table() -> bytes:
    """工程量核对表：送审工程量 | 审定工程量 | 综合单价，**没有合价列** —— 金额由 量×价 补出。
    两侧的「金额列表头文本」都是空（没有金额列），取数的列却不同（E、G 对 F、G）。
    真值：送审 100×50 + 40×200 = 13000；审定 90×50 + 40×200 = 12500；核减 500。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "工程量核对表"
    ws.append(["序号", "项目编码", "项目名称", "计量单位", "送审工程量", "审定工程量", "综合单价"])
    ws.append([1, "010401003001", "实心砖墙", "m3", 100, 90, 50])
    ws.append([2, "010502001001", "矩形柱", "m3", 40, 40, 200])
    return _save(wb)


def mixed_workbook() -> bytes:
    """一份工作簿：一张审核对比表（两组列）+ 一张单列合价表（基坑支护，审核没出对比表）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_AUDIT
    _audit_compare_sheet(ws)
    _single_amount_sheet(wb.create_sheet(SHEET_PIT), "基坑支护", PIT_ROWS, "2026-09-01")
    return _save(wb)


def _run(raw_a: bytes, raw_b: bytes, basis: str = "auto", names=("送审.xlsx", "审定.xlsx")) -> dict:
    return reconcile_engine.run(raw_a, names[0], raw_b, names[1], basis=basis)


def _codes(res: dict) -> list[str]:
    return [w["code"] for w in res["warning_items"]]


def _item(res: dict, code: str) -> dict:
    return next(w for w in res["warning_items"] if w["code"] == code)


# ================================================================ 同一份文件、同一列 → 拦

def test_same_single_amount_workbook_on_both_sides_is_rejected():
    raw = control_price()
    with pytest.raises(reconcile_engine.ParseFailed) as ei:
        _run(raw, raw)
    assert ei.value.reason_code == "same_file_same_columns"
    msg = str(ei.value)
    assert "同一份文件" in msg and "同一列" in msg
    assert "请分别选择送审文件和审定文件" in msg


def test_same_content_under_two_file_names_is_rejected_too():
    """用户把同一份文件传了两遍、起了两个名字：fileId 与文件名都不同，内容（sha256）相同。"""
    raw = control_price()
    with pytest.raises(reconcile_engine.ParseFailed) as ei:
        reconcile_engine.run(raw, "一标段-送审.xlsx", bytes(raw), "一标段-审定.xlsx",
                             submitted_file_id="F-SUB", approved_file_id="F-APP")
    assert ei.value.reason_code == "same_file_same_columns"


def test_endpoint_returns_422_for_same_file_and_for_renamed_copy():
    raw = control_price()
    a = shapes.put(raw, "ws26b-一标段.xlsx", "F1")
    r = client.post("/engine/reconcile/run", headers=HEADERS, json={"submitted_file": a, "approved_file": a})
    assert r.status_code == 422, r.text[:300]
    detail = r.json()["detail"]
    # WS26-e：detail 是对象，机器码随它透给 Java（改前只有这句话本身）
    assert detail["reason_code"] == "same_file_same_columns"
    assert "同一份文件" in detail["message"] and "同一列" in detail["message"]
    assert "rows" not in r.json(), "被拦下的请求不出结果表"

    b = shapes.put(raw, "ws26b-一标段-副本.xlsx", "F2")
    r2 = client.post("/engine/reconcile/run", headers=HEADERS, json={"submitted_file": a, "approved_file": b})
    assert r2.status_code == 422, r2.text[:300]
    assert "同一份文件" in r2.json()["detail"]["message"]


# ================================================================ 同一份文件、两组列 → 放行

def test_detail_audit_compare_table_on_both_sides_gives_the_column_difference():
    """明细版式的审核对比表（送审列 + 审定列）两侧选同一份是正确用法：核减额 = 两列之差。"""
    raw = audit_compare_detail()
    res = _run(raw, raw)
    s = res["summary"]
    assert (s["a_total"], s["b_total"]) == (34500.0, 31650.0)
    assert res["total_diff"] == 2850.0
    by = {r["item_code"]: r["diff"] for r in res["rows"]}
    assert by == {"010101002001": 1250.0, "010103001001": 1600.0, "010101004001": 0.0}
    w = _item(res, "same_file_two_sides")
    assert w["level"] == "info"
    assert "H 列" in w["message"] and "K 列" in w["message"], "写明两侧各取了哪一列"
    assert "same_file_sheets_excluded" not in _codes(res)
    assert w["message"] not in res["warnings"], "info 级不进旧的字符串告警通道"


def test_audit_summary_layout_on_both_sides_is_not_blocked():
    """审核汇总版式（_audit_rows）：送审侧取报审金额、审定侧取审核金额。"""
    raw = audit_summary()
    res = _run(raw, raw)
    assert (res["summary"]["a_total"], res["summary"]["b_total"]) == (1500000.0, 1430000.0)
    assert res["total_diff"] == 70000.0
    w = _item(res, "same_file_two_sides")
    assert "C 列" in w["message"] and "D 列" in w["message"]


def test_amount_derived_from_qty_times_price_compares_the_qty_and_price_columns():
    """没有合价列、金额由 量×价 补出：两侧「金额列表头文本」同为空，取数的列不同 —— 不被误拦。
    （比表头文本的实现在这里会把「都没有金额列」当成同一列。）"""
    raw = qty_check_table()
    res = _run(raw, raw)
    assert (res["summary"]["a_total"], res["summary"]["b_total"]) == (13000.0, 12500.0)
    assert res["total_diff"] == 500.0
    w = _item(res, "same_file_two_sides")
    assert "E、G 列" in w["message"] and "F、G 列" in w["message"]


def test_same_file_two_sided_all_equal_is_not_told_to_check_the_approved_file():
    """同一份工程量核对表、审定量与送审量恰好全相同：核减 0 是这张表自己的结论。
    all_rows_identical 的话术是「请确认审定文件选对了」，只对两份不同的文件有意义。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "工程量核对表"
    ws.append(["序号", "项目编码", "项目名称", "计量单位", "送审工程量", "审定工程量", "综合单价"])
    ws.append([1, "010401003001", "实心砖墙", "m3", 100, 100, 50])
    ws.append([2, "010502001001", "矩形柱", "m3", 40, 40, 200])
    raw = _save(wb)
    res = _run(raw, raw)
    assert res["total_diff"] == 0.0 and {r["status"] for r in res["rows"]} == {"一致"}
    assert "same_file_two_sides" in _codes(res) and "all_rows_identical" not in _codes(res)


def test_same_header_text_at_different_column_index_is_not_same_column():
    """比的是下标，不是表头文本：两组金额列合并格展开后都叫「金额(元)/合价」，列号不同就不是同一列。

    当前解析器按表头路径定角色，文本相同的两列必然同角色 → 歧义、整表不计入，所以这个形状今天
    从真实工作簿走不到；这里直接喂表状态，钉的是「判定只看 amount_col_idx」这条不变式 ——
    防以后解析器学会从别处（合并格范围、表名）定当事方时，判定悄悄退回比文本。"""
    def meta(idx: int) -> dict:
        return {"sheets_detail": [{"sheet": "对比表", "status": "used", "basis_label": "不分期",
                                   "columns": {"amount": "金额(元)/合价", "code": "项目编码"},
                                   "amount_col_idx": [idx]}]}
    same, split = reconcile_engine._same_file_sheets(meta(7), meta(10))
    assert same == [] and [s["sheet"] for s in split] == ["对比表"]
    # 对照：下标相同才是同一列
    same2, split2 = reconcile_engine._same_file_sheets(meta(7), meta(7))
    assert [s["sheet"] for s in same2] == ["对比表"] and split2 == []


# ================================================================ 同一份文件、部分表同列 → 剔除并告警

def test_mixed_workbook_drops_the_single_column_sheet_on_both_sides():
    raw = mixed_workbook()
    res = _run(raw, raw)
    w = _item(res, "same_file_sheets_excluded")
    assert w["level"] == "error" and SHEET_PIT in w["message"]
    assert w["message"] in res["warnings"]
    sheets = {r["submitted_sheet"] for r in res["rows"]} | {r["approved_sheet"] for r in res["rows"]}
    assert sheets == {SHEET_AUDIT}, "单列那张表两侧都不参与"
    assert len(res["rows"]) == 3
    # 核减额只来自审核对比表；被剔除的 110000 没有混进任何一侧的合计
    assert (res["summary"]["a_total"], res["summary"]["b_total"]) == (34500.0, 31650.0)
    assert res["total_diff"] == 2850.0
    assert "same_file_two_sides" in _codes(res)


def test_mixed_workbook_endpoint_status_is_partial():
    f = shapes.put(mixed_workbook(), "ws26b-混合.xlsx", "FM")
    r = client.post("/engine/reconcile/run", headers=HEADERS, json={"submitted_file": f, "approved_file": f})
    assert r.status_code == 200, r.text[:300]
    body = r.json()
    assert body["status"] == "partial" and body["status_reason"] == "items_not_counted"
    assert body["total_diff"] == 2850.0


def test_messages_stay_short_when_many_sheets_are_involved():
    """生产工作簿是上百张同版式的表：分列的表按「送审列 / 审定列」合并计数，被剔除的表名最多写 10 个。"""
    wb = Workbook()
    wb.remove(wb.active)
    for i in range(12):
        _audit_compare_sheet(wb.create_sheet(f"审核对比表【{i + 1}#楼】"))
    for i in range(12):
        rows = [(f"0102020010{i:02d}", f"地下连续墙{i}", "m3", 100, 800.0, 80000.0 + i)]
        _single_amount_sheet(wb.create_sheet(f"清单与计价表【{i + 1}#楼】"), f"{i + 1}#楼", rows, "2026-09-01")
    raw = _save(wb)
    res = _run(raw, raw)
    two = _item(res, "same_file_two_sides")["message"]
    assert "12 张表送审取 H 列" in two and two.count("H 列") == 1 and len(two) < 100
    exc = _item(res, "same_file_sheets_excluded")["message"]
    assert "其中 12 张表" in exc and exc.count("「清单与计价表【") == 10 and exc.endswith("等 12 张")
    assert {r["submitted_sheet"] for r in res["rows"]} == {"审核对比表【1#楼】"}, "12 张单列表的行一行都不在结果里"


# ================================================================ 两份不同文件

def test_two_different_files_all_identical_gets_a_warn_not_a_block():
    """两份不同的文件（另一次导出）、清单行逐分相同：可能是合法结果（审定无核减），只提示。"""
    a, b = control_price("2026-09-01"), control_price("2026-09-20")
    assert a != b, "夹具前提：两份文件字节不同"
    res = _run(a, b)
    assert res["total_diff"] == 0.0 and len(res["rows"]) == 5
    w = _item(res, "all_rows_identical")
    assert w["level"] == "warn"
    assert "5 行全部一致" in w["message"] and "核减额 0" in w["message"]
    assert w["message"] in res["warnings"]
    assert not {"same_file_two_sides", "same_file_sheets_excluded", "amounts_identical_suspect"} & set(_codes(res))


def test_two_different_files_with_a_real_difference_have_no_identical_warning():
    """负控：有一行真的调了价，或审定少了一行 → 不是「全部一致」。"""
    cut = [PIT_ROWS[0], ("010202009001", "喷射混凝土", "m2", 500, 55.0, 27500.0)]
    res = _run(control_price("2026-09-01"), control_price("2026-09-20", pit_rows=cut))
    assert res["total_diff"] == 2500.0
    assert "all_rows_identical" not in _codes(res)

    dropped = _run(control_price("2026-09-01"), control_price("2026-09-20", pit_rows=PIT_ROWS[:1]))
    assert dropped["summary"]["count_only_a"] == 1 and dropped["total_diff"] == 30000.0
    assert "all_rows_identical" not in _codes(dropped), "有「仅送审」行时核减额不是 0，不能说全部一致"


def test_multi_amount_header_guard_is_not_doubled_by_the_identical_warning():
    """两份不同的进度款文件都按合同价取：原有的 amounts_identical_suspect（error）照旧，
    不再叠一条 all_rows_identical 说同一件事。"""
    raw = shapes.zheda_three_row_header()
    res = _run(raw, shapes.resaved(raw), basis="contract")
    codes = _codes(res)
    assert "amounts_identical_suspect" in codes and "all_rows_identical" not in codes


def test_same_file_with_basis_that_lands_both_sides_on_one_column_is_rejected():
    """表头有送审列，但所选口径下没有：进度款文件两侧都按合同价取 → 同一列，照拦。"""
    raw = shapes.zheda_three_row_header()
    with pytest.raises(reconcile_engine.ParseFailed) as ei:
        _run(raw, raw, basis="contract")
    assert ei.value.reason_code == "same_file_same_columns"


# ================================================================ 表状态里的 amount_col_idx

def test_sheet_detail_records_the_amount_column_index():
    out = parse_costsheet(control_price(), "控制价.xlsx", side="submit")
    by = {d["sheet"]: d for d in out["meta"]["sheets_detail"]}
    assert by[SHEET_EARTH]["amount_col_idx"] == [AMOUNT_IDX]
    assert by[SHEET_PIT]["amount_col_idx"] == [AMOUNT_IDX]
    skipped = [d for d in out["meta"]["sheets_detail"] if d["status"] == "skipped_by_title"]
    assert len(skipped) == 2 and all(d["amount_col_idx"] == [] for d in skipped), "没取数的表不记列"

    raw = audit_compare_detail()
    sub = parse_costsheet(raw, "对比.xlsx", side="submit")["meta"]["sheets_detail"][0]
    aud = parse_costsheet(raw, "对比.xlsx", side="audit")["meta"]["sheets_detail"][0]
    assert (sub["amount_col_idx"], aud["amount_col_idx"]) == ([7], [10])

    raw = audit_summary()
    sub = parse_costsheet(raw, "汇总.xlsx", side="submit")["meta"]["sheets_detail"][0]
    aud = parse_costsheet(raw, "汇总.xlsx", side="audit")["meta"]["sheets_detail"][0]
    assert sub["format"] == aud["format"] == "audit_summary"
    assert (sub["amount_col_idx"], aud["amount_col_idx"]) == ([2], [3])

    raw = qty_check_table()
    sub = parse_costsheet(raw, "核对.xlsx", side="submit")["meta"]["sheets_detail"][0]
    aud = parse_costsheet(raw, "核对.xlsx", side="audit")["meta"]["sheets_detail"][0]
    assert (sub["amount_col_idx"], aud["amount_col_idx"]) == ([4, 6], [5, 6]), "量×价补出的记量、价两列"


def test_blank_amount_cells_filled_by_qty_times_price_add_those_columns():
    """合价列绑定了，但有的行合价格是空的（公式没缓存）→ 那几行按 量×价 补：实际取数的列是三列。
    两侧合价列相同、工程量列不同时，补出来的金额两侧并不相同，不能判成「同一列」。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "工程量核对表"
    ws.append(["序号", "项目编码", "项目名称", "计量单位", "送审工程量", "审定工程量", "综合单价", "合价"])
    ws.append([1, "010401003001", "实心砖墙", "m3", 100, 90, 50, None])
    ws.append([2, "010502001001", "矩形柱", "m3", 40, 40, 200, 8000])
    raw = _save(wb)
    sub = parse_costsheet(raw, "核对.xlsx", side="submit")["meta"]["sheets_detail"][0]
    aud = parse_costsheet(raw, "核对.xlsx", side="audit")["meta"]["sheets_detail"][0]
    assert (sub["amount_col_idx"], aud["amount_col_idx"]) == ([7, 4, 6], [7, 5, 6])
    res = _run(raw, raw)
    assert res["total_diff"] == 500.0, "100×50 − 90×50"
    assert "same_file_two_sides" in _codes(res)
