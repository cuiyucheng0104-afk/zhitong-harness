"""zt.report 子命令端到端：真实进程、中文与空格路径、Windows 写法的相对路径、结果契约与退出码。

一条完整链路：context → facts → templates → section_prepare → （agent 正文）→ section_finalize → export → fill_template。
旧仓 report/facts、report/export 端点的流程用例（读文件、模板要金额而没读到、时限、导出）改写在这里。
"""
from __future__ import annotations

import io
import json
import os
import subprocess
import sys
import zipfile
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from openpyxl import Workbook

from tests.report._support import DATA_DIR, FIXTURES, SKILL_DIR
from zt.common import ZtError
from zt.report import service
from zt.report.templates import load_templates

PYTHON_ROOT = Path(__file__).resolve().parents[2]
SAMPLE = FIXTURES / "report_facts" / "审定单-样例.xlsx"

EXPLAIN = (
    "一、项目概况\n本工程为某某学校新建工程，建设单位为某某教育局，施工单位为某某建设有限公司。\n"
    "二、审核依据\n施工合同、竣工图纸及现行计价规则。\n"
    "三、审核结果\n送审金额为321721元，审定金额为319786元。\n"
)


def run(*args: str, cwd: Path | None = None, script: bool = False) -> tuple[int, Any]:
    """跑 ``python -m zt.report.cli`` 或 skill 脚本；stdout 必须只有一行 JSON。"""
    env = {**os.environ, "PYTHONPATH": str(PYTHON_ROOT), "PYTHONIOENCODING": "ascii"}
    env.pop("ZT_REPORT_DATA", None)
    cmd = [sys.executable, str(SKILL_DIR / "scripts" / "report.py")] if script else [sys.executable, "-m", "zt.report.cli"]
    proc = subprocess.run([*cmd, *args], capture_output=True, cwd=cwd or PYTHON_ROOT, env=env, check=False)
    lines = proc.stdout.decode("utf-8").splitlines()
    assert len(lines) == 1, (lines, proc.stderr.decode("utf-8", "replace"))
    return proc.returncode, json.loads(lines[0])


def _result(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


@pytest.fixture
def project(tmp_path: Path) -> Path:
    """合成项目：目录名带中文与空格；一份说明、一份审定单、一份重复件、一份读不出的格式。"""
    root = tmp_path / "项目" / "2026-001-示例 学校"
    (root / "资料" / "重复 件").mkdir(parents=True)
    (root / "资料" / "审核说明.txt").write_text(EXPLAIN, encoding="utf-8")
    (root / "资料" / "重复 件" / "审核说明-副本.txt").write_text(EXPLAIN, encoding="utf-8")
    (root / "资料" / "审定单.xlsx").write_bytes(SAMPLE.read_bytes())
    (root / "资料" / "计价.GBQ7").write_bytes(b"\x00\x01binary")
    return root


def _context(root: Path, out: Path) -> dict:
    code, summary = run("context", "--root", str(root), "--files", "资料\\审核说明.txt", "资料/审定单.xlsx",
                        "资料/重复 件/审核说明-副本.txt", "资料/计价.GBQ7", "资料/缺失.pdf", "--out", f"{out}{os.sep}")
    assert code == 0, summary
    return _result(out / "result.json")


# ---------------------------------------------------------------- context


def test_context_numbers_by_path_dedups_and_reports_unreadable(project: Path, tmp_path: Path):
    ctx = _context(project, tmp_path / "ctx")
    assert ctx["status"] == "partial" and ctx["needs_manual"] is True
    files = ctx["data"]["files"]
    # 按规范后的相对路径排序：资料/审定单.xlsx < 资料/审核说明.txt；Windows 写法的反斜杠已规范成 /
    assert [(f["index"], f["rel_path"]) for f in files] == [(1, "资料/审定单.xlsx"), (2, "资料/审核说明.txt")]
    assert files[1]["text"] == EXPLAIN.strip() and files[1]["text_source"] == "text_layer"
    dups = ctx["data"]["duplicates"]
    assert [(d["rel_path"], d["duplicate_of"]) for d in dups] == [("资料/重复 件/审核说明-副本.txt", "资料/审核说明.txt")]
    unread = {u["rel_path"]: u for u in ctx["data"]["unreadable"]}
    assert unread["资料/缺失.pdf"]["extract_status"] == "missing"
    assert unread["资料/计价.GBQ7"]["extract_status"] == "unsupported"
    assert "导出为 Excel" in unread["资料/计价.GBQ7"]["reason"]
    assert ctx["coverage"]["files_total"] == 5 and ctx["coverage"]["files_used"] == 2
    assert {w["code"] for w in ctx["warning_items"]} >= {"unreadable_files", "duplicate_files"}


@pytest.mark.parametrize("bad", ["../外面.txt", "/etc/passwd", "C:\\Windows\\x.txt", "\\\\server\\share\\a.txt"])
def test_context_rejects_paths_outside_the_project(project: Path, bad: str):
    code, out = run("context", "--root", str(project), "--files", bad)
    assert code == 2 and out["error"]["code"] == "invalid_path"


def test_context_with_nothing_readable_is_empty(project: Path):
    code, out = run("context", "--root", str(project), "--files", "资料/计价.GBQ7")
    assert code == 0 and out["status"] == "empty" and out["status_reason"] == "no_readable_sources"


def test_context_needs_files_and_root(project: Path, tmp_path: Path):
    assert run("context", "--root", str(project))[1]["error"]["code"] == "invalid_request"
    assert run("context", "--root", str(tmp_path / "没有"), "--files", "a.txt")[1]["error"]["code"] == "root_not_found"


# ---------------------------------------------------------------- facts


def test_facts_from_context_reads_the_sample_with_cells(project: Path, tmp_path: Path):
    _context(project, tmp_path / "ctx")
    code, _s = run("facts", "--context", str(tmp_path / "ctx" / "result.json"), "--template", "jiesuan_shenhe_bg",
                   "--out", str(tmp_path / "facts" / "result.json"))
    assert code == 0
    out = _result(tmp_path / "facts" / "result.json")
    assert out["status"] == "ok"
    got = [(f["id"], f["key"], f["value"], f["sources"][0]["sheet"], f["sources"][0]["cell"]) for f in out["data"]["facts"]]
    assert got == [("F1", "submitted_total", "321721.00", "汇总表", "C7"),
                   ("F2", "approved_total", "319786.00", "汇总表", "D7"),
                   ("F3", "reduction", "1935.00", "汇总表", "E7"),
                   ("F4", "net_reduction", "1935.00", "汇总表", "F7")]
    assert all(f["needs_manual"] is True for f in out["data"]["facts"])
    assert out["data"]["needed_keys"][:2] == ["contract_price", "submitted_total"]
    assert out["coverage"]["files_total"] == 1 and out["coverage"]["files_used"] == 1


def test_template_needs_amounts_but_none_were_read_says_so(project: Path):
    code, out = run("facts", "--root", str(project), "--files", "资料/审核说明.txt", "--template", "jiesuan_shenhe_bg")
    assert code == 0 and out["status"] == "empty" and out["status_reason"] == "no_spreadsheet_files"
    w = next(x for x in out["warning_items"] if x["code"] == "facts_none")
    assert "所选资料里没有读到金额" in w["message"]


def test_template_without_amount_sections_has_no_such_notice(project: Path):
    code, out = run("facts", "--root", str(project), "--files", "资料/审核说明.txt", "--template", "guzhi_shuoming")
    assert code == 0 and out["data"]["needed_keys"] == [] and out["warning_items"] == []


def test_xls_spreadsheet_is_skipped_as_partial(project: Path):
    (project / "资料" / "旧审定单.xls").write_bytes(b"\xd0\xcf\x11\xe0legacy")
    refs = service.source_refs(["资料/审定单.xlsx", "资料/旧审定单.xls"])
    env = service.build_facts(project, refs)
    assert env.status.value == "partial"
    assert [(s["file_name"], s["reason"]) for s in env.coverage["skipped"]] == [("旧审定单.xls", "旧版 .xls 不能直接读取")]
    assert len(env.data["facts"]) == 4, "读出的照常返回"


def test_deadline_shows_up_as_a_warning(project: Path):
    env = service.build_facts(project, service.source_refs(["资料/审定单.xlsx"]), deadline_ms=0)
    assert env.status.value == "partial"
    w = next(x for x in env.warning_items if x.code == "facts_timeout")
    assert "《审定单.xlsx》" in w.message and w.level.value == "warn"


# ---------------------------------------------------------------- templates


def test_templates_lists_24_and_one_template_with_claims_and_docx(tmp_path: Path):
    code, out = run("templates")
    assert code == 0 and len(out["data"]["templates"]) == 24
    assert len(out["data"]["docx_templates"]) == 19
    code, out = run("templates", "--template", "jiesuan_shenhe_bg", script=True)
    assert code == 0
    d = out["data"]
    assert [s["code"] for s in d["sections"]][:2] == ["daoyu", "gaikuang"]
    assert "gen_prompt" not in d["sections"][0] and "rag_query" not in d["sections"][0]
    assert {c["key"] for c in d["claims_to_confirm"]} >= {"full_process_audit", "site_survey", "field_measure"}
    assert "竣工结算审核报告（房建）.docx" in d["docx_templates"]
    code, out = run("templates", "--template", "nope")
    assert code == 2 and out["error"]["code"] == "template_not_found"
    code, out = run("templates", "--data-dir", str(tmp_path))
    assert code == 2 and out["error"]["code"] == "data_dir_invalid"


# ---------------------------------------------------------------- prepare / finalize / export


def _prepare(tmp_path: Path, section: str, *extra: str) -> dict:
    out = tmp_path / "章节" / section
    code, summary = run("section_prepare", "--context", str(tmp_path / "ctx" / "result.json"),
                        "--facts", str(tmp_path / "facts" / "result.json"), "--template", "jiesuan_shenhe_bg",
                        "--section", section, "--project-name", "7.某某学校新建工程（软件+清单）",
                        "--out", f"{out}{os.sep}", *extra)
    assert code == 0, summary
    return _result(out / "result.json")


def _finalize(tmp_path: Path, section: str, draft: str, *extra: str) -> dict:
    d = tmp_path / "章节" / section
    (d / "draft.md").write_text(draft, encoding="utf-8")
    code, summary = run("section_finalize", "--prepare", str(d / "result.json"), "--draft", str(d / "draft.md"),
                        "--facts", str(tmp_path / "facts" / "result.json"), "--out", str(d / "final.json"), *extra)
    assert code == 0, summary
    return _result(d / "final.json")


@pytest.fixture
def pipeline(project: Path, tmp_path: Path) -> Path:
    _context(project, tmp_path / "ctx")
    code, _ = run("facts", "--context", str(tmp_path / "ctx" / "result.json"), "--out", str(tmp_path / "facts" / "result.json"))
    assert code == 0
    return tmp_path


def test_prepare_writes_prompt_md_without_fact_values(pipeline: Path):
    prep = _prepare(pipeline, "jieguo")
    assert prep["status"] == "ok"
    d = prep["data"]
    prompt = (pipeline / "章节" / "jieguo" / "prompt.md").read_text(encoding="utf-8")
    assert d["prompt_file"].endswith("prompt.md")
    assert d["messages"][0]["content"] in prompt and d["messages"][1]["content"] in prompt
    assert "〔F1〕送审金额（来源《审定单.xlsx》汇总表!C7）" in prompt
    assert "项目名称：某某学校新建工程\n" in prompt, "开头排序号与资料里没有的括注去掉"
    user = d["messages"][1]["content"]
    for v in ("321721.00", "319786.00"):
        assert v not in user
    assert d["section_fact_ids"] == ["F1", "F2", "F3", "F4"]
    assert [e["index"] for e in d["evidence"]] == [1, 2]
    listed = service.template_summary(load_templates(DATA_DIR), "jiesuan_shenhe_bg").data["sections"]
    assert d["prompt_version"] == next(s["prompt_version"] for s in listed if s["code"] == "jieguo")


def test_finalize_substitutes_facts_flags_unbacked_and_claims(pipeline: Path):
    _prepare(pipeline, "jieguo")
    out = _finalize(pipeline, "jieguo", "五、审核结果\n本工程送审总造价为〔F1〕，审核定案总造价为〔F2〕[资料1]。"
                                        "审定金额较送审减少1936元[资料2]。经现场踏勘，按审定单定案[资料1]。")
    d = out["data"]
    assert out["status"] == "partial" and out["needs_manual"] is True
    assert d["content"].startswith("本工程送审总造价为321721元，审核定案总造价为319786元[资料1]。"), "复述的章标题删掉"
    assert "1936元【待核：资料中未找到此数】" in d["content"]
    assert "现场踏勘【待核：「现场踏勘」未经确认、资料中也未见】" in d["content"]
    assert {f["kind"] for f in d["flags"]} == {"unbacked_number", "unconfirmed_claim"}
    refs = [r for r in d["amount_refs"] if r.get("origin") == "fact"]
    assert [(r["fact_id"], r["cell"]) for r in refs] == [("F1", "C7"), ("F2", "D7")]
    assert all(r["needs_manual"] is True for r in d["amount_refs"])
    assert [c["index"] for c in d["cited_files"]] == [1, 2]
    assert d["label"] == "含待核 2 处"


def test_confirmed_claim_is_not_flagged(pipeline: Path):
    _prepare(pipeline, "daoyu", "--confirmed", "site_survey", "unknown_key")
    prep = _result(pipeline / "章节" / "daoyu" / "result.json")
    assert prep["data"]["confirmed_claims"] == ["site_survey", "unknown_key"]
    assert prep["data"]["confirmed_in_section"] == ["现场踏勘"]
    w = next(w for w in prep["warning_items"] if w["code"] == "unknown_claims")
    assert "unknown_key" in w["message"] and "site_survey" not in w["message"]
    assert "- 现场踏勘" in prep["data"]["messages"][1]["content"]
    out = _finalize(pipeline, "daoyu", "受委托，我公司对某某学校新建工程的结算进行了审核，审核中进行了现场踏勘。"
                                       "建设单位为某某教育局[资料2]。")
    assert "unconfirmed_claim" not in {f["kind"] for f in out["data"]["flags"]}


def test_claim_confirmed_for_the_report_is_honoured_in_other_sections(pipeline: Path):
    """确认事项整份报告问一次：jieguo 章没有 site_survey 这条 claim，用户确认过的「现场踏勘」写进来也不标。"""
    _prepare(pipeline, "jieguo", "--confirmed", "site_survey")
    prep = _result(pipeline / "章节" / "jieguo" / "result.json")
    assert prep["data"]["confirmed_in_section"] == []
    assert not any(w["code"] == "unknown_claims" for w in prep["warning_items"])
    out = _finalize(pipeline, "jieguo", "经现场踏勘，送审总造价为〔F1〕[资料1]。审定总造价为〔F2〕[资料1]。")
    assert "unconfirmed_claim" not in {f["kind"] for f in out["data"]["flags"]}


def test_truncated_draft_is_cut_and_marked(pipeline: Path):
    _prepare(pipeline, "jieguo")
    out = _finalize(pipeline, "jieguo", "本工程送审总造价为〔F1〕[资料1]。审定总造价为〔F2〕[资料1]。安装工程部分主要差异在于给排水管道的",
                    "--truncated")
    assert out["status"] in ("partial", "insufficient")
    assert out["data"]["content"].endswith("【待补：本章内容超出篇幅上限，后面没有写完】")
    assert "给排水管道的" not in out["data"]["content"]


def test_finalize_errors_are_structured(pipeline: Path):
    _prepare(pipeline, "jieguo")
    d = pipeline / "章节" / "jieguo"
    (d / "draft.md").write_text("  \n", encoding="utf-8")
    code, out = run("section_finalize", "--prepare", str(d / "result.json"), "--draft", str(d / "draft.md"),
                    "--facts", str(pipeline / "facts" / "result.json"))
    assert code == 2 and out["error"]["code"] == "empty_draft"
    (d / "draft.md").write_text("送审总造价为〔F1〕[资料1]。" * 3, encoding="utf-8")
    code, out = run("section_finalize", "--prepare", str(d / "result.json"), "--draft", str(d / "draft.md"))
    assert code == 2 and out["error"]["code"] == "facts_mismatch", "prepare 给了〔F1〕…，finalize 没给同一份 facts"
    code, out = run("section_finalize", "--prepare", str(pipeline / "ctx" / "result.json"), "--draft", str(d / "draft.md"))
    assert code == 2 and out["error"]["code"] == "invalid_prepare"
    code, out = run("section_prepare", "--context", str(pipeline / "ctx" / "result.json"), "--template",
                    "jiesuan_shenhe_bg", "--section", "nope")
    assert code == 2 and out["error"]["code"] == "section_not_found"


def test_section_without_readable_sources_is_empty_and_needs_no_draft(project: Path, tmp_path: Path):
    code, _ = run("context", "--root", str(project), "--files", "资料/计价.GBQ7", "--out", str(tmp_path / "ctx" / "result.json"))
    assert code == 0
    code, s = run("section_prepare", "--context", str(tmp_path / "ctx" / "result.json"), "--template",
                  "jiesuan_shenhe_bg", "--section", "jieguo", "--out", str(tmp_path / "s" / "result.json"))
    assert code == 0 and s["status"] == "empty"
    assert not (tmp_path / "s" / "prompt.md").exists(), "本章不写：不给 agent 提示词"
    code, out = run("section_finalize", "--prepare", str(tmp_path / "s" / "result.json"))
    assert code == 0 and out["status"] == "empty" and out["status_reason"] == "no_readable_sources"
    assert out["data"]["content"] == service.EMPTY_SECTION_CONTENT


def test_export_builds_docx_with_disclaimer_appendix_and_section_order(pipeline: Path, project: Path):
    _prepare(pipeline, "gaikuang")
    _finalize(pipeline, "gaikuang", "本工程为某某学校新建工程，建设单位为某某教育局，施工单位为某某建设有限公司[资料2]。")
    _prepare(pipeline, "jieguo")
    _finalize(pipeline, "jieguo", "本工程送审总造价为〔F1〕，审核定案总造价为〔F2〕[资料1]。【待补：核增金额】")
    # 目录名不补位时按名字排序是 gaikuang < jieguo；SKILL 要求三位 seq 前缀
    docx = project / "成果" / "结算审核报告-初稿.docx"
    code, out = run("export", "--sections-dir", str(pipeline / "章节"), "--docx", str(docx), "--title", "结算审核报告",
                    "--context", str(pipeline / "ctx" / "result.json"), "--facts", str(pipeline / "facts" / "result.json"))
    assert code == 0, out
    assert out["status"] in ("ok", "partial") and docx.is_file()
    doc = Document(str(docx))
    texts = [p.text for p in doc.paragraphs]
    assert texts[0] == "结算审核报告"
    assert texts.index("一、项目概况") < texts.index("五、审核结果")
    assert "资料来源与完成度" in texts
    assert any(t.startswith("来源文件：") and "审定单.xlsx" in t for t in texts)
    assert any(t.startswith("无法读取的资料：") and "计价.GBQ7" in t for t in texts)
    assert any("核增金额" in t for t in texts), "待补清单"
    header = zipfile.ZipFile(io.BytesIO(docx.read_bytes())).read("word/header1.xml").decode("utf-8")
    assert "AI 初稿，金额未经人工核对，不得直接出具" in header
    code, out = run("export", "--sections-dir", str(pipeline / "没有"), "--docx", str(docx))
    assert code == 2 and out["error"]["code"] == "input_not_found"
    code, out = run("export", "--sections", str(pipeline / "ctx" / "result.json"), "--docx", str(docx))
    assert code == 2 and out["error"]["code"] == "invalid_section"


# ---------------------------------------------------------------- fill_template


def test_fill_template_replaces_and_reports_missing(project: Path, tmp_path: Path):
    values = tmp_path / "取值.json"
    values.write_text(json.dumps({"工程名称": "某某学校新建工程", "建设单位": "某某教育局", "没用的": "x"},
                                 ensure_ascii=False), encoding="utf-8")
    out_docx = project / "成果" / "竣工结算审核报告.docx"
    code, out = run("fill_template", "--template-docx", "竣工结算审核报告（房建）.docx", "--values", str(values),
                    "--docx", str(out_docx), script=True)
    assert code == 0 and out["status"] == "partial"
    d = out["data"]
    assert set(d["replaced"]) == {"工程名称", "建设单位"} and d["unused_values"] == ["没用的"]
    assert len(d["missing"]) == 53 and d["needs_manual"] is True
    xml = "".join(zipfile.ZipFile(out_docx).read(n).decode("utf-8") for n in ("word/document.xml", "word/header1.xml"))
    assert "某某学校新建工程" in xml and "{工程名称}" not in xml and "###施工单位###" in xml
    code, out = run("fill_template", "--template-docx", "没有.docx", "--docx", str(out_docx))
    assert code == 2 and out["error"]["code"] == "template_not_found"


def test_service_rejects_bad_requests():
    with pytest.raises(ZtError) as e:
        service.source_refs([{"rel_path": "a.txt", "kind": "other"}])
    assert e.value.code == "invalid_request"
    with pytest.raises(ZtError) as e:
        service.facts_list({"data": {"facts": [{"id": "F1"}]}})
    assert e.value.code == "invalid_facts"
    with pytest.raises(ZtError) as e:
        service.context_data({"x": 1})
    assert e.value.code == "invalid_context"


def test_xlsx_built_in_memory_is_read_by_the_grid_channel(project: Path):
    """审定单通道只认合计行那一格：合成一份两张表的审定单，汇总级表名的那张胜出。"""
    wb = Workbook()
    wb.remove(wb.active)
    for title, sub, app in (("汇总表", 500, 480), ("明细", 600, 580)):
        ws = wb.create_sheet(title)
        ws.append(["序号", "项目名称", "送审金额(元)", "审定金额(元)"])
        ws.append(["1", "分项", sub, app])
        ws.append(["合计", None, sub, app])
    buf = io.BytesIO()
    wb.save(buf)
    (project / "资料" / "两表.xlsx").write_bytes(buf.getvalue())
    env = service.build_facts(project, service.source_refs(["资料/两表.xlsx"]))
    assert [(f["key"], f["value"], f["sources"][0]["sheet"]) for f in env.data["facts"]] == [
        ("submitted_total", "500.00", "汇总表"), ("approved_total", "480.00", "汇总表")]


def test_context_caps_long_text_and_passes_hidden_sheets(project: Path):
    """旧 report/context 的 P04 / P12：单份正文按上限截断标 truncated（source_capped 告警）；隐藏工作表不进正文、名单透传。"""
    wb = Workbook()
    wb.active.title = "汇总"
    wb.active.append(["送审金额", 100])
    hidden = wb.create_sheet("示例")
    hidden.append(["示例金额", 999999])
    hidden.sheet_state = "hidden"
    buf = io.BytesIO()
    wb.save(buf)
    (project / "资料" / "带隐藏表.xlsx").write_bytes(buf.getvalue())
    env = service.build_context(project, service.source_refs(["资料/审核说明.txt", "资料/带隐藏表.xlsx"]), max_chars=40)
    files = {f["rel_path"]: f for f in env.data["files"]}
    assert env.status.value == "partial"
    assert files["资料/审核说明.txt"]["truncated"] is True and len(files["资料/审核说明.txt"]["text"]) <= 40
    assert any(w.code == "source_capped" for w in env.warning_items)
    sheet_file = files["资料/带隐藏表.xlsx"]
    assert "999999" not in sheet_file["text"] and sheet_file["hidden_sheets_skipped"] == ["示例"]
    assert env.data["hidden_sheets"] == ["《带隐藏表.xlsx》示例"]
