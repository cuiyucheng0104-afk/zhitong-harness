"""正式 docx 模板：占位符扫描与替换（拆在几个 run 里的、嵌套表格、页眉页脚、循环标签），登记清单与 docx 一致。"""
from __future__ import annotations

import hashlib
import json
from pathlib import Path

import pytest
from docx import Document

from tests.report._support import DATA_DIR
from zt.common import ZtError
from zt.report.docx_template import MANIFEST_FILE, build_manifest, load_manifest, render, scan

DOCX_DIR = DATA_DIR / "docx-templates"


def _synthetic(path: Path) -> Path:
    doc = Document()
    p = doc.add_paragraph()
    for piece in ("工程名称：{工", "程名", "称}，建设单位：", "{建设单位}", "。"):
        p.add_run(piece)
    p.runs[0].bold = True
    outer = doc.add_table(rows=1, cols=2)
    outer.cell(0, 0).text = "施工单位"
    inner = outer.cell(0, 1).add_table(rows=1, cols=1)
    inner.cell(0, 0).text = "{施工单位}（盖章）"
    doc.add_paragraph("{#循环}{/循环} 与 【待补：审核程序（只写实际执行过的）】")
    doc.sections[0].header.paragraphs[0].text = "{工程名称} 审核报告"
    doc.sections[0].footer.paragraphs[0].text = "第 1 页"
    doc.save(str(path))
    return path


def test_scan_finds_placeholders_across_runs_tables_and_header(tmp_path: Path):
    sc = scan(_synthetic(tmp_path / "t.docx"))
    assert sc.placeholders == ["工程名称", "建设单位", "施工单位"]
    assert sc.unsupported == ["#循环", "/循环"]
    assert sc.todo_count == 1


def test_render_replaces_split_runs_keeps_formatting_and_marks_missing(tmp_path: Path):
    src = _synthetic(tmp_path / "t.docx")
    res = render(src, {"工程名称": "某某学校新建工程", "建设单位": "某某教育局", "多余": "x", "空": None},
                 tmp_path / "out" / "o.docx")
    assert res.replaced == ["工程名称", "建设单位"] and res.occurrences == 4
    assert res.missing == ["施工单位"] and res.unsupported == ["#循环", "/循环"]
    assert res.unused_values == ["多余"]
    out = Document(str(tmp_path / "out" / "o.docx"))
    first = out.paragraphs[0]
    assert first.text == "工程名称：某某学校新建工程，建设单位：某某教育局。"
    assert first.runs[0].bold, "第一个 run 的格式留着"
    inner = out.tables[0].cell(0, 1).tables[0].cell(0, 0).text
    assert inner == "###施工单位###（盖章）"
    assert out.sections[0].header.paragraphs[0].text == "某某学校新建工程 审核报告"
    assert "{#循环}{/循环}" in out.paragraphs[1].text, "循环标签原样保留"


def test_render_unreadable_template_is_a_structured_error(tmp_path: Path):
    bad = tmp_path / "坏.docx"
    bad.write_bytes(b"not a zip")
    with pytest.raises(ZtError) as e:
        render(bad, {}, tmp_path / "o.docx")
    assert e.value.code == "docx_unreadable"


def test_manifest_is_in_sync_with_the_docx_files():
    """登记清单由 build_manifest 重算写出，不手工登记：与目录里的 docx 逐字段一致。"""
    assert load_manifest(DOCX_DIR) == build_manifest(DOCX_DIR)


def test_manifest_matches_the_old_expansion_pack_registry():
    """扩充包 16 份与旧仓登记清单逐份相同（sha256、字节数、占位符集合）；合计 19 份，占位符并集 77 个。"""
    m = json.loads((DOCX_DIR / MANIFEST_FILE).read_text(encoding="utf-8"))
    old = {t["upload_name"]: t for t in json.loads(
        (DOCX_DIR / "source-manifest-扩充包.json").read_text(encoding="utf-8"))["templates"]}
    assert m["count"] == 19 and len(old) == 16
    for t in m["templates"]:
        assert hashlib.sha256((DOCX_DIR / t["file"]).read_bytes()).hexdigest() == t["sha256"]
        if t["pack"] == "扩充包":
            o = old[t["file"]]
            assert (t["sha256"], t["size_bytes"], t["placeholders"]) == (o["sha256"], o["size_bytes"],
                                                                         sorted(o["placeholders"]))
    assert sum(1 for t in m["templates"] if t["pack"] == "起步包") == 3
    assert len(m["placeholder_union"]) == 77
    assert "概算建安造价" not in m["placeholder_union"], "种子里两行同名，模板不用它"


def test_mapping_only_names_real_templates_and_files():
    from tests.report._support import templates

    codes = {t["template_code"] for t in templates()}
    m = json.loads((DOCX_DIR / MANIFEST_FILE).read_text(encoding="utf-8"))
    for t in m["templates"]:
        assert set(t["report_templates"]) <= codes, t["file"]
    assert sum(1 for t in m["templates"] if t["report_templates"]) == 16
