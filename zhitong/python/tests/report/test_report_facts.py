# 来源：engine-service/tests/test_report_facts.py（zhitong-cost-ai@991ff5c，W3 搬入；纯函数用例照搬，HTTP 端点用例改写进 test_cli.py）
"""WS24-j（issue #1002；spec WS24 §3.8、§4 WS24-j）：报告的金额事实与〔Fn〕替换。

修的现场：运行 D「五、审核结果」模型自编了〔F1〕…〔F20〕（RM-C26），WS24-g 之后全换成【待补：金额】——止血了，
但资料里写着的送审 / 审定 / 核减金额、清单 xlsx 项目级费用表 C134 的控制价，仍然一个都填不进报告；模型照抄的
金额只核对「资料里出现过」，不核对「是不是这个意思的数」。没读到也不提示。

这里钉：
- 通道一（审定单）：夹具 `fixtures/report_facts/审定单-样例.xlsx` → 4 条事实带工作表与单元格；`read_grid_fields` 不传新参数时
  输出与改前逐项相同（净核减不出值），字段抽取的调用点源码里没有 `include_net_reduction`；同一文件三张表的取舍。
- 通道二（项目级费用表）：用 WS26-h 的那组夹具（`costsheet_shapes.fee_three_tier`），恰好 1 条、来源是合计行 C 列那一格、
  check 照搬 WS26-h 的三个数；采用单位之和 / 勾稽不过都不出事实；单位、单项工程表本身不出事实；补的三个字段不改指标。
- 替换与护栏：单位不重复、冲突、越章、自编编号、转写另标只在声明了金额事实的章节生效、提示词里没有数值、
  needs_manual 四态恒真（含「全部金额都来自 facts」的 ok 夹具）。
- 流程：facts 子命令的同类用例在 tests/report/test_cli.py（zt 改动：旧的 HTTP 端点用例改写到那里）。

期望值都由夹具构造手算，不从读取器的输出反抄。
"""
from __future__ import annotations

import io
import re
import time
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook

from tests.costsheet.fixtures import costsheet_shapes as shapes
from zt.costsheet import parser as cs
from zt.report import facts as rf
from zt.report.approval_grid import read_grid_fields
from zt.report.compose import CorpusFile, SectionSpec, build_messages, select_files_for_section

SAMPLE = Path(__file__).parent / "fixtures" / "report_facts" / "审定单-样例.xlsx"
ZT = Path(__file__).resolve().parents[2] / "zt"


def _collect(*named: tuple[str, bytes], **kw) -> rf.FactsResult:
    return rf.collect_facts([(f"id{i}", name, (lambda raw=raw: raw)) for i, (name, raw) in enumerate(named, 1)],
                            **kw)


def _save(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _audit_sheet(wb: Workbook, title: str, submitted, approved) -> None:
    ws = wb.create_sheet(title)
    ws.append(["序号", "项目名称", "送审金额(元)", "审定金额(元)"])
    ws.append(["1", "分项", submitted, approved])
    ws.append(["合计", None, submitted, approved])


# ================================================================ 通道一：审定单 / 汇总表

def test_sample_approval_sheet_gives_four_facts_with_sheet_and_cell():
    res = _collect((SAMPLE.name, SAMPLE.read_bytes()))
    got = {f["key"]: (f["value"], f["display"], f["unit"], f["sources"][0]["sheet"], f["sources"][0]["cell"])
           for f in res.facts}
    assert got == {
        "submitted_total": ("321721.00", "321721", "元", "汇总表", "C7"),
        "approved_total": ("319786.00", "319786", "元", "汇总表", "D7"),
        "reduction": ("1935.00", "1935", "元", "汇总表", "E7"),
        "net_reduction": ("1935.00", "1935", "元", "汇总表", "F7"),
    }
    assert [f["id"] for f in res.facts] == ["F1", "F2", "F3", "F4"]
    assert not any(f["conflict"] for f in res.facts)
    assert all(f["sources"][0]["file_name"] == SAMPLE.name for f in res.facts)


def test_shared_grid_reader_default_is_unchanged_net_reduction_only_on_request():
    """WS07 的字段抽取不传新参数：输出里没有「净核减」，与显式传 False 逐项相同；传 True 才多出那一格。"""
    raw = SAMPLE.read_bytes()
    default = read_grid_fields(raw, "xlsx", "a.xlsx")
    off = read_grid_fields(raw, "xlsx", "a.xlsx", include_net_reduction=False)
    on = read_grid_fields(raw, "xlsx", "a.xlsx", include_net_reduction=True)
    assert default.candidates == off.candidates and default.notes == off.notes
    assert "净核减" not in {c["field"] for c in default.candidates}
    extra = [c for c in on.candidates if c not in default.candidates]
    assert [(c["field"], c["evidence"]["cell"], c["amount_yuan"]) for c in extra] == [("净核减", "F7", "1935.00")]


def test_only_report_facts_passes_include_net_reduction():
    """源码棘轮（zt 改动：旧仓棘轮查字段抽取的两个调用点，字段抽取已废弃没有迁入）：zt 里传
    ``include_net_reduction=True`` 的只有报告金额事实这一处 —— 别处传了，净核减就会以「核减金额」的身份冒出来。"""
    users = sorted(p.relative_to(ZT).as_posix() for p in ZT.rglob("*.py")
                   if "include_net_reduction=True" in p.read_text(encoding="utf-8"))
    assert users == ["report/facts.py"]


def test_summary_sheet_wins_inside_one_workbook():
    """「汇总表」「土建」「安装」各有一个送审合计且不同 → 取汇总表的，不记冲突。"""
    wb = Workbook()
    wb.remove(wb.active)
    _audit_sheet(wb, "汇总表", 300000, 290000)
    _audit_sheet(wb, "土建", 200000, 195000)
    _audit_sheet(wb, "安装", 100000, 95000)
    res = _collect(("审定单.xlsx", _save(wb)))
    by = {f["key"]: f for f in res.facts}
    assert (by["submitted_total"]["value"], by["submitted_total"]["conflict"]) == ("300000.00", False)
    assert by["submitted_total"]["sources"] == [
        {"file_id": "id1", "file_name": "审定单.xlsx", "sheet": "汇总表", "cell": "C3"}]
    assert (by["approved_total"]["value"], by["approved_total"]["conflict"]) == ("290000.00", False)


def test_no_summary_sheet_means_conflict_not_a_pick():
    """三张表都不在汇总级词表里 → conflict=true，三个值都留着（不取第一个）。"""
    wb = Workbook()
    wb.remove(wb.active)
    _audit_sheet(wb, "一标段", 300000, 290000)
    _audit_sheet(wb, "土建", 200000, 195000)
    _audit_sheet(wb, "安装", 100000, 95000)
    res = _collect(("审定单.xlsx", _save(wb)))
    f = next(x for x in res.facts if x["key"] == "submitted_total")
    assert f["conflict"] is True and f["value"] is None
    assert sorted(c["value"] for c in f["candidates"]) == ["100000.00", "200000.00", "300000.00"]
    assert any("程序不取舍" in n for n in res.notes)


def test_same_key_different_values_across_files_is_a_conflict_and_both_kept():
    wb = Workbook()
    wb.remove(wb.active)
    _audit_sheet(wb, "汇总表", 321722, 319786)
    res = _collect(("审定单.xlsx", SAMPLE.read_bytes()), ("审定单（改）.xlsx", _save(wb)))
    f = next(x for x in res.facts if x["key"] == "submitted_total")
    assert f["conflict"] is True
    assert sorted((c["value"], c["source"]["file_name"]) for c in f["candidates"]) == [
        ("321721.00", "审定单.xlsx"), ("321722.00", "审定单（改）.xlsx")]
    same = next(x for x in res.facts if x["key"] == "approved_total")
    assert same["conflict"] is False and len(same["sources"]) == 2, "值相同的不算冲突，来源都留着"


# ================================================================ 通道二：项目级费用表（WS26-h 的夹具）

def test_fee_channel_gives_exactly_one_control_price_from_the_project_total_cell():
    """项目级 1 张 + 单位工程 2 张，单位工程合计是文本，项目级合计行叫「合计」，比单位之和大 0.35。"""
    res = _collect(("控制价.xlsx", shapes.fee_three_tier()))
    assert len(res.facts) == 1
    f = res.facts[0]
    assert (f["key"], f["value"], f["display"], f["unit"], f["conflict"]) == (
        "control_price_total", "15873168.00", "15873168", "元", False)
    assert f["sources"] == [{"file_id": "id1", "file_name": "控制价.xlsx",
                             "sheet": "表10.2.2-12 招标控制价费用表", "cell": "C9"}], "合计行的 C 列，不是规费那一列"
    c = f["check"]
    assert (c["p"], c["u"], c["n_unit"], c["diff"]) == ("15873168.00", "15873167.65", 2, "0.35")
    assert "差 0.35 元" in c["text"]


def test_unit_sum_is_never_a_fact():
    """只有单位工程表、没有项目级表 → WS26-h 采用单位之和（指标照用），报告不出事实（D-R60 不派生）。"""
    raw = shapes.fee_three_tier(project_total=None)
    assert cs.parse_costsheet(raw, "控制价.xlsx")["meta"]["fee_total_check"]["source"] == "unit_sum", "夹具前提"
    res = _collect(("控制价.xlsx", raw))
    assert res.facts == []
    assert any("程序不相加" in n for n in res.notes), res.notes


def test_project_total_two_yuan_off_gives_no_fact_and_carries_the_reason():
    res = _collect(("控制价.xlsx", shapes.fee_three_tier(project_total=15873167.65 + 2)))
    assert res.facts == []
    assert any("差 2.00 元" in n and "对不上" in n for n in res.notes), res.notes


def test_unit_and_single_sheets_never_produce_facts():
    """单项、单位工程表各有「招标控制价合计」：它们本身不出任何事实，只用来勾稽。"""
    res = _collect(("控制价.xlsx", shapes.fee_three_tier(with_single=True)))
    assert [(f["key"], f["sources"][0]["sheet"]) for f in res.facts] == [
        ("control_price_total", "表10.2.2-12 招标控制价费用表")]
    assert res.facts[0]["check"]["n_single"] == 1


def test_fee_hits_carry_cell_amount_str_and_display():
    """WS26-h 命中项只增三个字段：数值格常规写法、文本格照抄（含千分位）、Decimal 串不经 float。"""
    hits = cs.parse_costsheet(shapes.fee_three_tier(), "控制价.xlsx")["meta"]["fee_totals"]
    got = sorted((h["tier"], h["cell"], h["amount_str"], h["display"]) for h in hits)
    assert got == [("project", "C9", "15873168.00", "15873168"),
                   ("unit", "D11", "3895.44", "3,895.44"),
                   ("unit", "D15", "15869272.21", "15,869,272.21")]


def test_indicator_output_is_identical_with_and_without_the_three_new_fields(monkeypatch):
    """只增字段的回归：把三个新字段从命中项里拿掉，指标的输出逐项相同（指标不读它们）。"""
    from zt.indicator import engine as indicator_engine

    raw = shapes.fee_three_tier()
    with_fields = indicator_engine.run([(raw, "控制价.xlsx", "f1")], None, gross_area=1000.0)
    real = cs.scan_fee_sheet

    def stripped(title, rows):
        out = real(title, rows)
        out["hits"] = [{k: v for k, v in h.items() if k not in ("cell", "amount_str", "display")} for h in out["hits"]]
        return out

    monkeypatch.setattr(cs, "scan_fee_sheet", stripped)
    assert "cell" not in cs.parse_costsheet(raw, "控制价.xlsx")["meta"]["fee_totals"][0], "替身生效了"
    without = indicator_engine.run([(raw, "控制价.xlsx", "f1")], None, gross_area=1000.0)
    assert with_fields == without


def _big_workbook() -> bytes:
    """355 张表的合成工作簿（照 §1.1 的生产原件分布）：项目级 1 + 单项 15 + 单位工程 113 + 清单 113 + 材料表 113。"""
    wb = Workbook()
    wb.remove(wb.active)
    unit_cents = [100_000_000 + i * 123_456 for i in range(113)]
    singles = [unit_cents[i::15] for i in range(15)]
    total = sum(unit_cents)
    proj = wb.create_sheet("表10.2.2-12 招标控制价费用表")
    proj.append(["【表10.2.2-12】"])
    proj.append(["招标控制价费用表"])
    proj.append(["序号", "工程名称", "金额(元)", "备注"])
    for i, s in enumerate(singles, 1):
        proj.append([i, f"单项{i}", sum(s) / 100, None])
    proj.append(["合计", None, round(total / 100), None])
    for i, s in enumerate(singles, 1):
        ws = wb.create_sheet(f"10.2.2-13 单项工程费用表S{i}")
        ws.append(["单项工程招标控制价费用表"])
        ws.append(["序号", "费用名称", "计算公式", "金额（元）", "备注"])
        ws.append(["招标控制价合计", None, "1+2+3+4+5", f"{sum(s) / 100:.2f}", None])
    for i, c in enumerate(unit_cents):
        ws = wb.create_sheet(f"10.2.2-13 单位工程费用表U{i}")
        ws.append(["单位工程招标控制价费用表"])
        ws.append([f"工程名称：单项{i % 15 + 1}-U{i}", None, "标段：", None, None])
        ws.append(["序号", "费用名称", "计算公式", "金额（元）", "备注"])
        ws.append(["1", "分部分项工程费", None, c / 100 * 0.8, None])
        ws.append(["招标控制价合计", None, "1+2+3+4+5", f"{c / 100:,.2f}", None])
        lst = wb.create_sheet(f"10.2.2-16 清单与计价表U{i}")
        lst.append(["分部分项工程清单与计价表"])
        lst.append(["序号", "项目编码", "项目名称", "项目特征", "计量单位", "工程量", "综合单价", "合价"])
        for r in range(1, 21):
            lst.append([r, f"0101010{r:02d}001", f"项目{r}", "", "m3", 10 * r, 12.5, 125 * r])
        mat = wb.create_sheet(f"10.2.2-19 主要材料价格表U{i}")
        mat.append(["主要材料价格表"])
        mat.append(["序号", "名称", "单位", "单价"])
        mat.append([1, "钢筋", "t", 4200])
    assert len(wb.sheetnames) == 355
    return _save(wb)


def test_355_sheet_workbook_reads_facts_within_30_seconds():
    raw = _big_workbook()
    t0 = time.monotonic()
    res = _collect(("一标段.xlsx", raw))
    elapsed = time.monotonic() - t0
    assert elapsed < 30, f"facts 用时 {elapsed:.1f} 秒"
    assert [f["key"] for f in res.facts] == ["control_price_total"], res.notes
    assert (res.facts[0]["check"]["n_unit"], res.facts[0]["check"]["n_single"]) == (113, 15)


# ================================================================ 替换

F_SUB = {"id": "F1", "key": "submitted_total", "label": "送审金额", "value": "321721.00", "display": "321721",
         "unit": "元", "conflict": False, "candidates": [], "check": None,
         "sources": [{"file_id": "x1", "file_name": "审定单-样例.xlsx", "sheet": "汇总表", "cell": "C7"}]}
F_APP = dict(F_SUB, id="F2", key="approved_total", label="审定金额", value="319786.00", display="319786",
             sources=[{"file_id": "x1", "file_name": "审定单-样例.xlsx", "sheet": "汇总表", "cell": "D7"}])
F_CTRL = {"id": "F3", "key": "control_price_total", "label": "招标控制价", "value": "318375872.00",
          "display": "318375872", "unit": "元", "conflict": False, "candidates": [],
          "sources": [{"file_id": "c1", "file_name": "一标段.xlsx", "sheet": "表10.2.2-12 招标控制价费用表", "cell": "C134"}],
          "check": {"p": "318375872.00", "u": "318375871.65", "n_unit": 113, "s": None, "n_single": 15, "z": None,
                    "n_specialty": 0, "diff": "0.35",
                    "text": "项目级合计 318,375,872.00，113 张单位工程费用表之和 318,375,871.65，差 0.35 元"}}
F_CONFLICT = {"id": "F4", "key": "approved_total", "label": "审定金额", "value": None, "display": None, "unit": None,
              "conflict": True, "check": None,
              "sources": [{"file_id": "a", "file_name": "审定单A.xlsx", "sheet": "汇总表", "cell": "D7"},
                          {"file_id": "b", "file_name": "审定单B.xlsx", "sheet": "汇总【终版】", "cell": "D9"}],
              "candidates": [
                  {"value": "319786.00", "display": "319786", "unit": "元",
                   "source": {"file_id": "a", "file_name": "审定单A.xlsx", "sheet": "汇总表", "cell": "D7"}},
                  {"value": "319886.00", "display": "319,886.00", "unit": "元",
                   "source": {"file_id": "b", "file_name": "审定单B.xlsx", "sheet": "汇总【终版】", "cell": "D9"}}]}


@pytest.mark.parametrize("raw, want", [
    ("送审金额为〔F1〕元", "送审金额为321721元"),
    ("送审金额为〔F1〕。", "送审金额为321721元。"),
    ("送审金额为〔F1〕 元。", "送审金额为321721元。"),
    ("送审金额为〔F1〕万元。", "送审金额为321721元。"),   # 单位以程序取到的为准，不拼成「321721万元」
])
def test_known_allowed_fact_becomes_cell_display_plus_unit_once(raw, want):
    text, flags, refs = rf.substitute_facts(raw, [F_SUB], {"submitted_total"})
    assert text == want
    assert flags == []
    assert refs == [{"raw": "321721元", "value": "321721.00", "file_ids": ["x1"], "origin": "fact", "fact_id": "F1",
                     "key": "submitted_total", "label": "送审金额", "file_name": "审定单-样例.xlsx",
                     "sheet": "汇总表", "cell": "C7"}]


def test_fact_with_check_puts_check_in_the_ref_and_no_marker_in_the_text():
    text, flags, refs = rf.substitute_facts("招标控制价总价为〔F3〕。", [F_CTRL], {"control_price_total"})
    assert text == "招标控制价总价为318375872元。"
    assert "【待核" not in text and flags == []
    assert refs[0]["check"]["text"].endswith("差 0.35 元") and refs[0]["cell"] == "C134"


def test_conflicting_fact_lists_every_value_and_is_marked():
    text, flags, refs = rf.substitute_facts("审定金额为〔F4〕元。", [F_CONFLICT], {"approved_total"})
    assert text.startswith("审定金额为【待核：金额冲突，")
    assert "《审定单A.xlsx》汇总表!D7 为 319786元" in text and "《审定单B.xlsx》汇总（终版）!D9 为 319,886.00元" in text
    assert re.fullmatch(r"审定金额为【待核[^】\n]*】。", text), "表名里的【】不许把标记提前闭合"
    assert [f.kind for f in flags] == ["fact_conflict"] and refs == []


def test_fact_not_allowed_in_this_section_is_out_of_scope():
    text, flags, refs = rf.substitute_facts("送审金额为〔F1〕元。", [F_SUB], {"control_price_total"})
    assert text == "送审金额为【待核：本章不应出现此金额】。"
    assert [f.kind for f in flags] == ["out_of_scope_amount"] and refs == []


def test_unknown_placeholder_becomes_todo():
    text, flags, _ = rf.substitute_facts("核减〔F9〕元。", [F_SUB], {"submitted_total"})
    assert text == "核减【待补：金额】。"
    assert [f.kind for f in flags] == ["bad_placeholder"]


# ================================================================ 提示词

def _selection(text: str):
    return select_files_for_section([CorpusFile(index=1, file_id="d1", file_name="说明.docx", text=text)], "审核")


def test_prompt_lists_allowed_facts_without_their_values():
    spec = SectionSpec(code="jieguo", title="五、审核结果", gen_prompt="写审核结果",
                       facts_allowed=("submitted_total", "approved_total"))
    msgs = build_messages(spec, _selection("本次审核依据审定单进行。"), facts=[F_SUB, F_APP],
                          project_name="7.合成项目（软件+清单）")
    blob = "\n".join(m["content"] for m in msgs)
    # B8 R3F-e + j：金额事实块与项目名防护同时生效，金额不得泄入提示词。
    assert "【本项目】\n项目名称：合成项目\n工程名称以资料原文为准，项目名称只作口径提示" in blob
    assert "软件+清单" not in blob
    assert "〔F1〕送审金额（来源《审定单-样例.xlsx》汇总表!C7）" in blob
    assert "〔F2〕审定金额" in blob
    for v in ("321721", "319786"):
        assert v not in blob, f"提示词里出现了事实的数值 {v}（铁律 2：模型拿不到数）"


def test_no_facts_no_placeholder_talk_in_the_prompt():
    spec = SectionSpec(code="jieguo", title="五、审核结果", gen_prompt="写审核结果",
                       facts_allowed=("submitted_total",))
    blob = "\n".join(m["content"] for m in build_messages(spec, _selection("本次审核依据审定单进行。"), facts=[]))
    assert "〔F" not in blob and "金额事实" not in blob


def test_prompt_version_moves_with_facts_allowed():
    from zt.report.compose import prompt_version

    a = SectionSpec(code="c", title="t", gen_prompt="g")
    b = SectionSpec(code="c", title="t", gen_prompt="g", facts_allowed=("control_price_total",))
    assert prompt_version(a) != prompt_version(b)


def test_oversize_file_is_not_read_into_memory():
    """按 stat 判 20 MB：超了不调读字节的函数（Java 不知道文件多大，全靠这里挡）。"""
    def boom():
        raise AssertionError("超大的文件不许读进内存")
    res = rf.collect_facts([("a", "大.xlsx", boom)], size=lambda fid: rf.FACTS_MAX_FILE_BYTES + 1)
    assert res.files_read == 0 and res.skipped == [("a", "大.xlsx", "超过 20 MB")]
    assert any("超过 20 MB" in n for n in res.notes)


def test_deadline_leaves_unread_files_in_notes():
    ticks = iter([0.0, 0.0, 61.0, 61.0])
    res = rf.collect_facts([("a", "A.xlsx", SAMPLE.read_bytes), ("b", "B.xlsx", SAMPLE.read_bytes)],
                           clock=lambda: next(ticks))
    assert res.files_read == 1 and res.timed_out == ["B.xlsx"]
    assert "金额事实读取超时：B.xlsx" in res.notes



def test_needed_notice_counts_only_the_keys_the_template_needs():
    assert rf.needed_notice([F_CTRL], ["submitted_total"]) is not None, "读到的是控制价，模板要的是送审金额"
    assert rf.needed_notice([F_SUB], ["submitted_total"]) is None
    assert rf.needed_notice([], []) is None
    assert Decimal(F_CTRL["value"]) == Decimal("318375872")
