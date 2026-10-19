"""生成合成样例项目夹具 ``sample-project/``（全部虚构，不含真实客户文件）。

用法（在 zhitong/python 下）::

    uv run python tests/inventory/fixtures/make_sample_project.py

场景：结算审核项目（咨询类型 20002、阶段 审定），资料基本齐全，故意留出三种结果：

- 缺件：没有「投标报价文件」（重要项）、「造价指标分析」（可选项）；
- 待确认：``变更资料/签证单汇总.xlsx`` 只命中裸词「签证」（弱证据）→ 变更签证费用待确认；
- 干扰：施工日志（黑名单，不是造价资料）、``INDEX.md`` 与 ``.ocr.txt`` 旁车（系统生成物）、``工作/`` 下的中间产物。

文件内容是几行虚构数字，PDF 正文用英文（pypdf 不嵌中文字体）；DWG 是占位字节，不是真实图纸。
"""
from __future__ import annotations

import io
import shutil
import sys
from pathlib import Path

import docx
from openpyxl import Workbook
from pypdf import PdfWriter

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parents[2]))  # zhitong/python，供 import tests.*
from tests.textextract import ws04_fixtures as fx  # noqa: E402

OUT = HERE / "sample-project"

PROJECT_DOC = """---
项目编号: 2026-001
工程名称: 示例住院楼
咨询类型: 结算审核
阶段: 审定
业态: 医疗
规模: 12000 平方米（虚构）
别名: [合同号 HT-DEMO-01]
状态: 进行中
---
# 示例住院楼

合成样例项目，全部内容虚构，仅用于测试。

## 基本信息

| 项 | 值 |
|---|---|
| 项目编号 | 2026-001 |
| 咨询类型 | 结算审核 |
| 阶段 | 审定 |

## 关键文件指引

| 资料 | 位置 |
|---|---|
| 施工合同 | 资料/1 合同及协议/施工合同及补充协议.docx |
| 结算送审书 | 资料/4 结算资料/结算送审书.pdf |

## 已出成果

- 成果/结算审核报告.docx
- 成果/结算审定单.xlsx
- 成果/审核汇总表.xlsx
"""


def docx_file(path: Path, *paragraphs: str) -> None:
    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(path)


def xlsx_file(path: Path, rows: list[list[object]]) -> None:
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def pdf_file(path: Path, *kinds_and_text: str) -> None:
    """kinds: 'text' 现造文字页，'image' 扫描页。"""
    w = PdfWriter()
    for i, kind in enumerate(kinds_and_text, start=1):
        if kind == "text":
            fx.add_text_page(w, f"Demo document page {i} contract amount 1234567.89 yuan, fictional sample")
        else:
            fx.add_image_page(w)
    buf = io.BytesIO()
    w.write(buf)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(buf.getvalue())


def main() -> None:
    if OUT.exists():
        shutil.rmtree(OUT)
    OUT.mkdir(parents=True)
    (OUT / "项目说明.md").write_text(PROJECT_DOC, encoding="utf-8", newline="\n")

    ziliao = OUT / "资料"
    docx_file(ziliao / "1 合同及协议" / "施工合同及补充协议.docx", "施工合同（合成样例）", "合同价款 1234567.89 元（虚构）")
    pdf_file(ziliao / "2 招标资料" / "招标文件.pdf", "text", "text")
    xlsx_file(ziliao / "2 招标资料" / "工程量清单.xlsx", [["序号", "项目", "工程量"], [1, "土方开挖", 100.5]])
    pdf_file(ziliao / "3 竣工资料" / "竣工图" / "竣工图总说明.pdf", "text")
    pdf_file(ziliao / "4 结算资料" / "结算送审书.pdf", "text", "text", "text")
    xlsx_file(ziliao / "5 变更资料" / "签证单汇总.xlsx", [["序号", "签证事项", "金额"], [1, "地基处理", 2000.0]])
    docx_file(ziliao / "6 其他" / "施工日志.docx", "2026-03-01 晴 浇筑混凝土（虚构）")
    pdf_file(ziliao / "6 其他" / "扫描件0001.pdf", "image", "image")
    (ziliao / "6 其他" / "扫描件0001.pdf.ocr.txt").write_text(
        fx.paged_sidecar([("ok", ["某示例工程变更联系单 金额 2000 元（虚构）"]), ("ok", ["第二页 审批意见 同意"])],
                         name="扫描件0001.pdf"),
        encoding="utf-8", newline="\n")
    (ziliao / "6 其他" / "INDEX.md").write_text(
        "<!-- generated-by: scripts/build_source_index.py -->\n# 目录\n", encoding="utf-8", newline="\n")

    tuzhi = OUT / "图纸"
    tuzhi.mkdir(parents=True)
    (tuzhi / "建筑施工图.dwg").write_bytes(b"AC1032 synthetic placeholder, not a real drawing\n")

    chengguo = OUT / "成果"
    docx_file(chengguo / "结算审核报告.docx", "结算审核报告（合成样例）")
    xlsx_file(chengguo / "结算审定单.xlsx", [["项目", "审定金额"], ["合计", 1200000.0]])
    xlsx_file(chengguo / "审核汇总表.xlsx", [["单位工程", "送审", "审定"], ["土建", 1300000.0, 1200000.0]])

    work = OUT / "工作" / "2026-10-19-核对"
    work.mkdir(parents=True)
    (work / "核对结果.json").write_text('{"note": "中间产物，盘点不看"}\n', encoding="utf-8", newline="\n")


if __name__ == "__main__":
    main()
