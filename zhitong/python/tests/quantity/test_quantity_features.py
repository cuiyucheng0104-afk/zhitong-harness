# 来源：engine-service/tests/test_quantity_features.py（zhitong-cost-ai@991ff5cd，W4 迁入 zt；改动见「zt 改动」注释）
"""项目特征摘录（QTY-06 / QTY-14）：槽位号的键形状、逐字校验、两步走后特征不丢。

zt 改动：旧用例验「命中 LLM 缓存后〔AI摘录〕不变成待补」（issue #823：首跑整数键、回放字符串键）。
zt 里摘录回答由 agent 写进 answers 文件、经 JSON 往返交给 finalize —— 同一个键形状问题换成「answers 文件
往返后取值不变」；同一份 answers 重跑 finalize 特征逐字相同。
"""
from __future__ import annotations

import json
from io import BytesIO

from openpyxl import load_workbook

from zt.quantity import features as feat
from zt.quantity import pipeline as ql
from tests.quantity.fixtures.shapes import FakeAgent, keyword_picker, txt_file
from tests.quantity.support import run

# 通用砌体做法用语，不含项目标识
TASKS = [
    feat.FeatureTask("i1", "010402001", "砌块墙",
                     {1: "砌块品种及规格、强度等级", 4: "砂浆强度等级"},
                     ["200厚混凝土加气块B07 A5.0，M7.5专用砂浆砌筑"]),
    feat.FeatureTask("i2", "011102003", "块料楼地面",
                     {5: "结合层种类及强度等级", 6: "面层材料品种及规格", 8: "勾缝材料种类"},
                     ["10厚防滑地砖面层", "20厚1:3干硬性水泥砂浆结合层"]),
]
# i2 的 6 号槽：改写过的值（过不了原文校验）；8 号槽：没答
ANSWER = {"f": [
    {"k": "i1", "v": {"1": "混凝土加气块B07 A5.0", "4": "M7.5"}},
    {"k": "i2", "v": {"5": "1:3干硬性水泥砂浆", "6": "防滑地砖（300×300）"}},
]}

EXPECTED = {
    "i1": {1: "〔AI摘录〕混凝土加气块B07 A5.0", 4: "〔AI摘录〕M7.5"},
    "i2": {5: "〔AI摘录〕1:3干硬性水泥砂浆", 6: feat.PENDING_UNVERIFIED, 8: feat.PENDING},
}


def _fill(answer) -> dict:
    got = feat.parse_feature_answer(answer)
    return {t.key: feat.fill_values(t, got.get(t.key, {})) for t in TASKS}


def test_answers_file_round_trip_keeps_ai_extracted_values():
    # 夹具必须同时有：通过校验的、没通过校验的、没答的 —— 三种回填在 JSON 往返后原样保住
    direct = _fill(ANSWER)
    replay = _fill(json.loads(json.dumps(ANSWER, ensure_ascii=False)))
    assert direct == replay == EXPECTED
    assert sum(v.startswith(feat.AI_PREFIX) for d in replay.values() for v in d.values()) == 3


def test_parse_feature_answer_emits_string_slot_keys():
    out = feat.parse_feature_answer({"f": [{"k": "i1", "v": {"1": "甲", " 02 ": " 乙 ", "x": "丙", "3": 7}}]})
    # 「02」归一成「2」（同一个槽）；非数字键、非字符串值丢弃
    assert out == {"i1": {"1": "甲", "2": "乙"}}
    assert json.loads(json.dumps(out)) == out
    assert feat.parse_feature_answer([{"k": "i1", "v": {"1": "甲"}}]) == {"i1": {"1": "甲"}}, "直接给数组也认"
    assert feat.parse_feature_answer({"f": "不是数组"}) == {} and feat.parse_feature_answer(None) == {}


def test_fill_values_accepts_both_key_types():
    """按字符串取，兼容整数键（直接递字典的调用方）。"""
    task = TASKS[0]
    assert feat.fill_values(task, {"1": "混凝土加气块B07 A5.0", "4": "M7.5"}) == EXPECTED["i1"]
    assert feat.fill_values(task, {1: "混凝土加气块B07 A5.0", 4: "M7.5"}) == EXPECTED["i1"]
    assert feat.fill_values(task, None) == {1: feat.PENDING, 4: feat.PENDING}
    assert feat.fill_values(task, {}) == {1: feat.PENDING, 4: feat.PENDING}


def test_task_json_carries_slots_and_sources():
    row = feat.task_json(TASKS[1])
    assert row == {"k": "i2", "编码": "011102003", "名称": "块料楼地面",
                   "槽位": {"5": "结合层种类及强度等级", "6": "面层材料品种及规格", "8": "勾缝材料种类"},
                   "原文": ["10厚防滑地砖面层", "20厚1:3干硬性水泥砂浆结合层"]}


# ============================================================ 整条流水线跑两遍

FLOOR_TXT = (
    "楼面1：防滑地砖楼面（用于卫生间）\n"
    "1. 8厚防滑地砖面层\n"
    "2. 20厚1:3干硬性水泥砂浆结合层\n"
    "楼面2：玻化砖楼面（用于大厅）\n"
    "1. 10厚800×800玻化砖面层\n"
    "2. 30厚1:3干硬性水泥砂浆结合层\n"
)


def _verbatim_features(item):
    """假 agent 的特征回答：只摘原文里有的字（过得了逐字校验），所以一定产出〔AI摘录〕。"""
    src = "".join(item["原文"])
    vals = {}
    for n, label in item["槽位"].items():
        if "面层材料品种" in label:
            vals[n] = "防滑地砖" if "防滑地砖" in src else "800×800玻化砖"
        elif "结合层种类" in label:
            vals[n] = "1:3干硬性水泥砂浆"
    return vals


def _d_column(xlsx: bytes) -> list:
    ws = load_workbook(BytesIO(xlsx))["分部分项工程量清单"]
    return [row[3].value for row in ws.iter_rows()]


def test_pipeline_second_finalize_with_saved_answers_keeps_features():
    fake = FakeAgent(keyword_picker([(("砖",), "main", "011102003")]), features=_verbatim_features)
    files = [txt_file(FLOOR_TXT)]
    first = run(files, agent=fake)
    assert fake.feature_calls, "夹具要走到特征摘录"
    n_ai = sum(i["features"].count(feat.AI_PREFIX) for i in first["items"])
    assert n_ai >= 4, [i["features"] for i in first["items"]]
    assert not first["feature_tasks"], "摘录作答后不再有待摘录任务"

    saved = json.loads(json.dumps(first["answers"], ensure_ascii=False))
    second = ql.finalize(files, saved, prepare_fingerprint=first["prepare_fingerprint"])
    assert len(second["items"]) == len(first["items"]) >= 2
    for a, b in zip(first["items"], second["items"]):
        assert b["code"] == a["code"]
        assert b["features"] == a["features"], a["code"]
    from zt.quantity.render import render_xlsx
    # 交付物：xlsx「项目特征描述」列逐格相同（生产基线是 299 对 0、132 行不同）
    assert _d_column(render_xlsx(second)) == _d_column(first["xlsx_bytes"])
