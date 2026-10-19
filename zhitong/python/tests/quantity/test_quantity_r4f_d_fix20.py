# 来源：engine-service/tests/test_quantity_r4f_d_fix20.py（zhitong-cost-ai@991ff5cd，W4 迁入 zt；改动见「zt 改动」注释）
"""fix20：材料指称、量化和零序数通过真实DOCX/API/XLSX验证；外部模型固定。"""
import base64
import json
import os
import tempfile
from io import BytesIO
from pathlib import Path
import pytest
from openpyxl import load_workbook
from zt.quantity import pipeline as ql
from zt.quantity import render as qr
from tests.quantity import test_quantity_r4f_d_fix18 as h
from tests.quantity.support import run


def product(monkeypatch, cid, body, tail, anchor, expected, crosscell=False):
    # 删除材料回指/量化/零序数识别将让取消负例漏入正式表；扩大格范围会让crosscell正控变待核。
    xml = h._cad([["外墙甲", body, tail]]) if crosscell else h._p(h._r("外墙甲："+body+"。"+tail))
    data, name, ext = h._docx(xml, name="结构验证.docx")
    traces=[]
    original=ql.decide
    def trace(seg,ans,ctx):
        result=original(seg,ans,ctx)
        row,role=result
        traces.append({'id':seg.id,'text':seg.text,'location':seg.location,'dest':row.dest,'reason_code':row.reason_code,'suggested':row.suggested,'llm_code':row.llm_code})
        return result
    monkeypatch.setattr(ql,'decide',trace)
    def pick(seg): return ('main',h.W) if anchor in seg['原文'] else ('base','')
    h._fake(monkeypatch,pick)
    pipeline=run([(data,name,ext)],agent=h._AGENT)
    traces.clear()
    h._fake(monkeypatch,pick)
    response=h._api(data,name,'fix20-'+cid)
    output = os.environ.get('FIX20_OUT')
    root = Path(output)/cid if output else Path(tempfile.mkdtemp(prefix='fix20-products-'))/cid
    root.mkdir(parents=True,exist_ok=False)
    (root/'input.docx').write_bytes(data)
    api=response.json()
    (root/'api.json').write_text(json.dumps(api,ensure_ascii=False,indent=2))
    (root/'traces.json').write_text(json.dumps(traces,ensure_ascii=False,indent=2))
    assert response.status_code==200,api
    (root/'api.xlsx').write_bytes(base64.b64decode(api['xlsx_base64']))
    (root/'pipeline.xlsx').write_bytes(pipeline['xlsx_bytes'])
    books=[load_workbook(BytesIO(pipeline['xlsx_bytes'])),load_workbook(root/'api.xlsx')]
    targets=[t for t in traces if anchor in t['text']]
    assert len(targets)==1,targets
    target=targets[0]
    assert books[1].sheetnames==[qr.SHEET_LIST,qr.SHEET_PENDING,qr.SHEET_COVERAGE,'说明']
    readback={s:list(books[1][s].values) for s in books[1].sheetnames}
    (root/'xlsx-readback.json').write_text(json.dumps(readback,ensure_ascii=False,indent=2))
    for sheet in books[1].sheetnames[:3]:assert list(books[0][sheet].values)==readback[sheet]
    assert pipeline['unmatched']==api['unmatched']
    cov=[r for r in readback[qr.SHEET_COVERAGE][4:] if r[0]==target['id']]
    assert len(cov)==1 and cov[0][7]==target['text'] and cov[0][2]==target['location']
    for row in readback[qr.SHEET_LIST]:
        if isinstance(row[0],int):assert all(v is None for v in row[5:12]),row
    assert target['dest']==('pending_review' if expected=='pending' else 'item'),target
    if expected=='pending':
        assert target['reason_code']=='semantic_mismatch'
        assert target['llm_code']==h.W and target['suggested'].startswith('建议编码：011001003 保温隔热墙面（m²）')
        assert '待人工确认' in target['suggested']
        assert any(w['code']=='quantity_no_formal_items' for w in api['warning_items'])
        assert any(r[4]==target['text'] for r in readback[qr.SHEET_PENDING][1:]),readback[qr.SHEET_PENDING]
    else:
        assert any(item['code'].startswith(h.W) and item['unit']=='m²' and any(s['quote']==target['text'] and s['location']==target['location'] for s in item['sources']) for item in api['items'])
    return root


FORMS=[
    '80厚泡沫玻璃板（密度120kg/m³）保温层',
    '80厚泡沫玻璃板，作为保温层',
    '80厚泡沫玻璃板',
    '泡沫玻璃板保温层（80厚）',
    '80厚泡沫玻璃板粘贴保温层',
]
@pytest.mark.parametrize('form',FORMS)
@pytest.mark.parametrize('tail,expected',[('泡沫玻璃板取消','pending'),('泡沫玻璃板保留','formal')])
def test_material_reference_survives_description_structure(monkeypatch,form,tail,expected):
    product(monkeypatch,'material-'+str(FORMS.index(form))+'-'+expected,'1)20厚砂浆找平层；2)'+form+'；3)5厚抗裂砂浆保护层',tail,'80厚',expected)

@pytest.mark.parametrize('quant',['每一种构造层','每一个材料层','每一类施工层','各类构造层','逐个构造层'])
@pytest.mark.parametrize('verb,expected',[('取消','pending'),('保留','formal')])
def test_quantifier_classifier_and_modifier_are_compositional(monkeypatch,quant,verb,expected):
    product(monkeypatch,'quant-'+quant+'-'+verb,'1)80厚岩棉板保温层；2)5厚抗裂砂浆保护层',quant+verb,'80厚',expected)

@pytest.mark.parametrize('number',['零','0'])
@pytest.mark.parametrize('crosscell',[False,True])
def test_zero_ordinal_is_wide_only_in_same_cell(monkeypatch,number,crosscell):
    product(monkeypatch,'zero-'+number+'-'+str(crosscell),'1)20厚砂浆找平层；2)80厚岩棉板保温层；3)5厚抗裂砂浆保护层','第'+number+'层不做','80厚','formal' if crosscell else 'pending',crosscell)

@pytest.mark.parametrize('tail',['泡沫玻璃板取消','每一种构造层取消'])
def test_reference_does_not_cross_source_cell(monkeypatch,tail):
    product(monkeypatch,'cell-'+tail,'1)20厚砂浆找平层；2)80厚泡沫玻璃板（密度120kg/m³）保温层；3)5厚抗裂砂浆保护层',tail,'80厚','formal',True)
