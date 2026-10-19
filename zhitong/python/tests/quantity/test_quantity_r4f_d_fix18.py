# 来源：engine-service/tests/test_quantity_r4f_d_fix18.py（zhitong-cost-ai@991ff5cd，W4 迁入 zt；改动见「zt 改动」注释）
"""R4F-d fix18（R4-006，P1 存疑闸）：review18a / 18b / 18c 三份评审的落实（用户 2026-10-12 四项决定与协调者 2026-10-13
十条裁定不变；存疑闸只降级，不升级、不改码、不出工程量）。

  决定 1 / 裁定 4  本层原文后句只收窄「纯序数、明确指向别的（存在的）层」的句子；带回指 / 自指 / 量化 / 补集 / 相对序数、
                   序数越界、单位是 条 / 道 / 步 的一律取宽（review18a R18A-01、review18b R18B-02、review18c R18C-01 / 05）。
  裁定 4           另一层另一句：补集（除第N层外 / 其余 / 剩下的 / 别的 / 另外的 / 另一层）、相对序数（第N层以下 / 起 / 上方
                   一层 / 的下一层）取宽或算本层；量化语素（全已 / 全被 / 全予以 / 整幢 / 每一层 / 逐层 / 层层 / 一切 / 一同 …）、
                   方位语素（之下 / 此下 / 下部 / 以后 / 底下 / 上部 / 上方 / 前部 / 先前 …）按语素类认；反方向计数（由内向外 /
                   自下起 / 从底部算起 / 由底至顶）按层数换算；首道 / 头道 / 头一层 / 前两层 取宽；层数按层号数
                   （review18a R18A-02 / 05 / 06 / 07）。
  裁定 3           段首层号按语素类（罗马数字、天干地支、大写数字、方括号层号、多级编号、圈码字母 / 汉字、Word 自动编号）、
                   项目符号按 Unicode 类别、标签可带括号 / 空格 / 较长、方括号标签；兄弟层片段右侧是标点时用真实右邻判
                   叉号 / 斜杠 / 破折号（review18a R18A-03 / 04、review18c R18C-04）。
  裁定 2           去除类前字按语素类补齐（刷 冲 锤 炸 掰 钻 …）；「掉」只在缺陷语境里按后字排除（review18a R18A-08）。
  裁定 6           「除湿区外 / 除草坪外」不是整词「除湿 / 除草」，照「除…外」算（review18c R18C-07）。
  裁定 1           objects 只管标注 / 谓语用法、整段原文里找；固定写法照放；「无缝」补 objects；after_gap
                   （review18b R18B-01 / 03 / 04、review18c R18C-03）。
  裁定 5           无表头 CAD 行：规格格前面紧挨的主材格并进本层格（review18c R18C-02）。
  决定 1           首层层号粘在做法名称冒号后面（「外墙C19：1）…」）时认出层号；名称批注按说明文字取宽（review18c R18C-05 / 06）。
  决定 2           触发清单里每个汉字的繁体 / 异体按 ICU 补全（review18c R18C-10）。
"""
from __future__ import annotations

import io
import json
import zipfile
from io import BytesIO
from xml.sax.saxutils import escape

import pytest
from openpyxl import load_workbook

from zt.quantity import pipeline as ql
from zt.quantity import render as qr
from zt.quantity import rules
from zt.quantity import segment as qs
from tests.quantity.fixtures.shapes import FakeAgent
from tests.quantity.support import run

W, P = "011001003", "011404001"
FILE = "单元-R4F-d-fix18.docx"
FILE_ID = "synthetic-fix18-unit"
OL = "5厚聚合物抗裂砂浆面层"
LV = "20厚1:3水泥砂浆找平层"
PT = "刷外墙涂料两遍"


def gate(text, code=W):
    return rules.doubt_gate(code, text)


def tokens(v):
    return [t.token for t in v.gating] if v is not None else []


def links(v):
    return [t.link for t in v.gating] if v is not None else []


def practice(segs):
    return [s for s in segs if s.kind == qs.PRACTICE]


def pieces(seg):
    return [(p.start, p.end, p.role) for p in seg.host.pieces]


def decided(seg, code=W):
    """quantity_list._doubt_of 的同一顺序：本行原文（收窄后）→ doubt_gate；没有命中再看所在格 / 段 → doubt_host。"""
    text = seg.scope_text()
    if seg.host is not None:
        text = rules.doubt_own_scope(text, seg.host.text, pieces(seg), seg.host_piece)
    row = rules.doubt_gate(code, text)
    hv = None
    if row is None and seg.host is not None:
        hv = rules.doubt_host(code, seg.host.text, seg.host.where, pieces(seg), seg.host_piece)
    return row, hv


def pending(seg, code=W):
    row, hv = decided(seg, code)
    return row is not None or hv is not None


def text_seg(cell, probe):
    segs = [s for s in practice(qs.segment_text(cell, 0, "x.txt")) if probe in s.text]
    assert len(segs) == 1, ("夹具：探针要恰好切出一片", cell, [s.text for s in practice(qs.segment_text(cell, 0, "x.txt"))])
    return segs[0]


def _docx(body, name=FILE, numbering=None):
    xml = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document '
           'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>' + body + '</w:body></w:document>')
    ct = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/'
          'package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.'
          'relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/'
          'document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
          + ('<Override PartName="/word/numbering.xml" ContentType="application/vnd.openxmlformats-officedocument.'
             'wordprocessingml.numbering+xml"/>' if numbering else '') + '</Types>')
    rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.'
            'org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/'
            'officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml", ct)
        z.writestr("_rels/.rels", rels)
        z.writestr("word/document.xml", xml)
        if numbering:
            z.writestr("word/_rels/document.xml.rels",
                       '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.'
                       'openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.'
                       'openxmlformats.org/officeDocument/2006/relationships/numbering" Target="numbering.xml"/>'
                       '</Relationships>')
            z.writestr("word/numbering.xml", numbering)
    return buf.getvalue(), name, "docx"


def _p(*runs, num=0):
    ppr = f'<w:pPr><w:numPr><w:ilvl w:val="0"/><w:numId w:val="{num}"/></w:numPr></w:pPr>' if num else ""
    return "<w:p>" + ppr + "".join(runs) + "</w:p>"


def _r(text):
    return f'<w:r><w:t xml:space="preserve">{escape(text)}</w:t></w:r>'


def _tc(*paras):
    return "<w:tc>" + "".join(paras) + "</w:tc>"


def _tr(*cells):
    return "<w:tr>" + "".join(cells) + "</w:tr>"


def _tbl(*rows):
    return "<w:tbl>" + "".join(rows) + "</w:tbl>"


def _cad(rows):
    return _tbl(*[_tr(*[_tc(_p(_r(c))) for c in row]) for row in rows])


def _numbering(fmt, lvl, start):
    return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:numbering xmlns:w="http://schemas.openxmlformats.'
            'org/wordprocessingml/2006/main"><w:abstractNum w:abstractNumId="0"><w:lvl w:ilvl="0">'
            f'<w:start w:val="{start}"/><w:numFmt w:val="{fmt}"/><w:lvlText w:val="{escape(lvl)}"/>'
            '<w:suff w:val="nothing"/></w:lvl></w:abstractNum><w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num>'
            '</w:numbering>')


_AGENT: "FakeAgent | None" = None


class _Resp:
    def __init__(self, status_code, body):
        self.status_code = status_code
        self._body = body

    def json(self):
        return self._body


def _api(data, name, stem):
    """zt 改动：旧版经 FastAPI 端点；现在在本进程里走 CLI 三步（prepare → 假 agent 作答 → finalize → render），
    文件落在临时项目目录，返回旧响应体的形状（finalize 的 data + warning_items + xlsx_base64）。"""
    import base64
    import tempfile
    from pathlib import Path

    from zt.quantity.cli import cli

    root = Path(tempfile.mkdtemp(prefix=f"zt-qty-{stem}-"))
    rel = f"资料/{name}"
    (root / "资料").mkdir()
    (root / rel).write_bytes(data)

    def call(*argv):
        out = io.StringIO()
        code = cli.main(list(argv), stdout=out)
        return code, json.loads(out.getvalue())

    code, _ = call("prepare", "--root", str(root), "--files", rel, "--out", str(root / "prep") + "/")
    if code != 0:
        return _Resp(422, _)
    prep = json.loads((root / "prep" / "result.json").read_text(encoding="utf-8"))
    (root / "answers.json").write_text(json.dumps(_AGENT.answer(prep["data"]), ensure_ascii=False), encoding="utf-8")
    code, summary = call("finalize", "--prepare", str(root / "prep" / "result.json"), "--answers",
                         str(root / "answers.json"), "--out", str(root / "fin") + "/")
    if code != 0:
        return _Resp(422, summary)
    code, _ = call("render", "--result", str(root / "fin" / "result.json"), "--xlsx", str(root / "清单.xlsx"))
    assert code == 0, _
    env = json.loads((root / "fin" / "result.json").read_text(encoding="utf-8"))
    body = {**env["data"], "status": env["status"], "warning_items": env["warning_items"],
            "xlsx_base64": base64.b64encode((root / "清单.xlsx").read_bytes()).decode("ascii")}
    return _Resp(200, body)


def _fake(monkeypatch, pick):
    global _AGENT
    _AGENT = FakeAgent(pick)
    return _AGENT


# ====================================================================== 本层原文后句（决定 1 + 裁定 4）
@pytest.mark.parametrize("tail,narrowed", [
    ("第1层取消", True), ("第1层改为真石漆", True), ("第一层不做", True), ("第1层取消，砂浆另计", True),
    ("第1项取消", True), ("倒数第二层取消", True),
    # 拿不准：取宽
    ("第1层取消，本层同样取消", False), ("第1层及本层均不做", False), ("第1层与此层都改为XPS板", False),
    ("第1层撤销，该层一并撤销", False), ("第一层作废，这一层亦作废", False), ("第1层亦取消", False),
    ("第1层取消，连同本层", False), ("第1层撤销，自身也撤销", False), ("第1层取消，当前这层也取消", False),
    ("除第1层外均取消", False), ("第1层以外全部改为XPS板", False), ("第1层保留，其余取消", False),
    ("第1层照旧，剩余的不做", False), ("第1层照做，其他的取消", False), ("第1层起全部取消", False),
    ("第1层之后的取消", False), ("第1层以下均取消", False), ("第1层至本层均取消", False),
    ("第1層取消，本層亦取消", False), ("第5层取消", False), ("第一道取消", False), ("详见总说明第7条，取消", False),
    ("第1层取消，两层都取消", False), ("第1层与其余各层均删除", False), ("第2项取消", False),
])
def test_own_text_later_sentence_is_narrowed_only_when_it_is_a_pure_ordinal_elsewhere(tail, narrowed):
    cell = "1、" + LV + "。2、8101厚岩棉板保温层。" + tail
    seg = text_seg(cell, "8101厚")
    assert tail in seg.text, "夹具：那句要落进本层自己的原文"
    out = rules.doubt_own_scope(seg.scope_text(), seg.host.text, pieces(seg), seg.host_piece)
    assert (out != seg.scope_text()) is narrowed, (tail, out)
    assert len(out) == len(seg.scope_text())
    assert pending(seg) is (not narrowed or rules.doubt_gate(W, out) is not None)


def test_own_text_trigger_kept_in_the_pending_reason_when_not_narrowed():
    seg = text_seg("1、" + LV + "。2、8102厚岩棉板保温层。第1层取消，本层亦取消", "8102厚")
    row, _ = decided(seg)
    assert row is not None and tokens(row) == ["取消", "取消"]


# ====================================================================== 另一层另一句（裁定 4：补集 / 相对 / 量化 / 方位 / 序数）
@pytest.mark.parametrize("cell,probe,label", [
    ("1、" + LV + "。除第1层外均取消。2、8111厚岩棉板保温层", "8111厚", "序数指层"),
    ("1)" + PT + "。第1层保留，其余均不做；2)8112厚岩棉板保温层", "8112厚", "全称"),
    ("1、" + PT + "，第1层照旧，剩下的取消。2、8113厚岩棉板保温层", "8113厚", "全称"),
    ("1、" + PT + "。第1层照做，别的不做。2、8114厚岩棉板保温层", "8114厚", "全称"),
    ("1、" + PT + "。从第1层起都取消。2、8115厚岩棉板保温层", "8115厚", "序数指层"),
    ("1、8116厚岩棉板保温层，2、" + OL + "。第2层上方的取消", "8116厚", "序数指层"),
    ("1、8117厚岩棉板保温层，2、" + OL + "。另一层取消", "8117厚", "补集"),
    ("1、8118厚岩棉板保温层。2、" + OL + "。全已取消", "8118厚", "全称"),
    ("1、8119厚岩棉板保温层。2、" + OL + "。全予以取消", "8119厚", "全称"),
    ("1、8120厚岩棉板保温层。2、" + OL + "。整幢不做", "8120厚", "全称"),
    ("1、8121厚岩棉板保温层。2、" + OL + "。每一层都取消", "8121厚", ("全称", "回指")),
    ("1、8122厚岩棉板保温层。2、" + OL + "。层层取消", "8122厚", "全称"),
    ("1、" + PT + "。逐层取消。2、8123厚岩棉板保温层", "8123厚", "全称"),
    ("1、" + PT + "，之下取消。2、8124厚岩棉板保温层", "8124厚", "前指"),
    ("1、" + PT + "。下部取消。2、8125厚岩棉板保温层", "8125厚", "前指"),
    ("1、" + PT + "。以后的作废。2、8126厚岩棉板保温层", "8126厚", "前指"),
    ("1、8127厚岩棉板保温层。2、" + OL + "。上部取消", "8127厚", "回指"),
    ("1、8128厚岩棉板保温层。2、" + OL + "。先前的取消", "8128厚", "回指"),
    ("1、" + PT + "，2、" + OL + "，3、8129厚岩棉板保温层，4、" + LV + "。由内向外第2层取消", "8129厚", "序数指层"),
    ("1、" + PT + "，2、" + OL + "，3、8130厚岩棉板保温层，4、" + LV + "。从底部算起第2层取消", "8130厚", "序数指层"),
    ("1、8131厚岩棉板保温层，2、" + OL + "，3、" + LV + "。首道取消", "8131厚", "序数指层"),
    ("1、" + PT + "，2、" + OL + "，3、8132厚岩棉板保温层，8厚胶粘剂粘贴，4、" + LV + "。倒数第二层取消", "8132厚", "序数指层"),
    ("1、8133厚岩棉板保温层，2、" + OL + "，3、" + LV + "。第2层上方一层取消", "8133厚", "序数指层"),
    ("1、" + PT + "，2、8134厚岩棉板保温层，3、" + OL + "。前两层取消", "8134厚", ("序数指层", "回指")),
    # 序数点了别的层、句里另有补集说法（其余部分 / 任何别的）：本层在补集里（句子在本层之前，回指 / 量化都够不着）
    ("1、" + PT + "。第1层保留，其他部分取消。2、8135厚岩棉板保温层", "8135厚", "补集"),
    ("1、" + PT + "。第1层照做，别处取消。2、8136厚岩棉板保温层", "8136厚", "补集"),
    # 指向别的层 / 方向不对 / 同形的非指代：不算
    ("1、" + LV + "。第1层取消。2、8141厚岩棉板保温层", "8141厚", None),
    ("1、8142厚岩棉板保温层，2、" + OL + "，3、" + PT + "，4、" + LV + "。倒数第一层取消", "8142厚", None),
    ("1、" + PT + "，2、" + OL + "，3、8143厚岩棉板保温层，8厚胶粘剂粘贴，4、" + LV + "。倒数第一层取消", "8143厚", None),
    ("1、8144厚岩棉板保温层。2、" + OL + "。以下取消。3、" + LV, "8144厚", None),
    ("1、" + PT + "。上部取消。2、8145厚岩棉板保温层", "8145厚", None),
    ("1、8146厚岩棉板保温层。2、" + OL + "。涂刷均匀，不得漏涂", "8146厚", None),
    ("1、8147厚岩棉板保温层。2、" + OL + "。安全网不得拆除", "8147厚", None),
    ("1、8148厚岩棉板保温层。2、" + OL + "，全粘，不得空鼓", "8148厚", None),
    # 方位词 / 一同 是时间、部位、方式用法（不在分句开头）：不指别的层
    ("1、" + LV + "，终凝以后不得踩踏。2、8149厚岩棉板保温层", "8149厚", None),
    ("1、" + LV + "，墙体下部不得空鼓。2、8150厚岩棉板保温层", "8150厚", None),
    ("1、8152厚岩棉板保温层。2、" + OL + "，与找平层一同施工，不得漏涂", "8152厚", None),
    ("1、8153厚岩棉板保温层。2、" + OL + "。一同取消", "8153厚", ("全称", "回指")),
    ("1、8154厚岩棉板保温层，2、" + OL + "，3、" + LV + "。第一层和最后一层取消", "8154厚", "序数指层"),
    ("1、" + PT + "，2、8155厚岩棉板保温层，3、" + LV + "。第一层和最后一层取消", "8155厚", None),
])
def test_other_layer_sentence_links_by_morpheme_class(cell, probe, label):
    seg = text_seg(cell, probe)
    row, hv = decided(seg)
    assert row is None, (cell, tokens(row))
    if label is None:
        assert hv is None, (cell, tokens(hv), links(hv))
    else:
        want = label if isinstance(label, tuple) else (label,)
        assert hv is not None and any(x.startswith(want) for x in links(hv)), (cell, links(hv))


def test_ordinal_reader_conversions_relatives_and_synonyms():
    assert rules._doubt_ordinal("由内向外第2层取消", 4)[0] == frozenset({3})
    assert rules._doubt_ordinal("从里往外第二层不做", 4)[0] == frozenset({3})
    assert rules._doubt_ordinal("自下起第二层作废", 4)[0] == frozenset({3})
    assert rules._doubt_ordinal("从底部算起第2层取消", 4)[0] == frozenset({3})
    assert rules._doubt_ordinal("由底至顶第2层取消", 4)[0] == frozenset({3})
    assert rules._doubt_ordinal("一至二层取消")[0] == frozenset({1, 2})
    assert rules._doubt_ordinal("第壹层取消")[0] == frozenset({1})
    for first in ("首道取消", "头道不做", "头一层取消", "头层取消", "最先一层取消", "最上面一层取消", "最外层取消"):
        assert rules._doubt_ordinal(first, 4)[0] == frozenset({1}), first
    for last in ("最下层取消", "最底层取消", "最后一层取消", "末层取消", "尾层取消", "最内层取消"):
        assert rules._doubt_ordinal(last, 4)[0] == frozenset({4}), last
        assert rules._doubt_ordinal(last)[0] is None, last                 # 层数认不出：取宽
    for wide in ("第1层以下均取消", "第1层起取消", "第2层上方一层取消", "第1层的下一层取消", "第3层往上数第2层取消",
                 "前两层取消", "后两层取消", "由内向外第2层取消", "首层取消", "底层取消", "顶层取消"):
        od = rules._doubt_ordinal(wide)
        assert od is not None and od[0] is None, (wide, od)


def test_layer_total_counts_layer_numbers_when_a_layer_is_split_by_thickness():
    cell = "1、" + PT + "，2、" + OL + "，3、8151厚岩棉板保温层，8厚胶粘剂粘贴，4、" + LV
    seg = text_seg(cell, "8151厚")
    assert rules._doubt_layer_total(seg.host.text, pieces(seg)) == 4


# ====================================================================== 裁定 3：段首标记、兄弟层右邻
@pytest.mark.parametrize("prefix", [
    "Ⅱ、", "Ⅲ）", "Ⅳ.", "ⅲ）", "IV、", "iii)", "（Ⅱ）", "(iv)", "甲、", "乙）", "（丁）", "寅、", "贰、", "（叁）", "【2】", "[3]",
    "〔2〕", "〈2〉", "1.1、", "（2.1）", "2.1.3、", "Ⓑ", "⑵", "⒝", "㈡", "㊁", "·", "•", "◦", "♦", "❖", "✦", "*", "＊",
    "保温层（外墙）：", "外墙保温层 A型：", "外墙外保温防火隔离带构造做法补充说明：", "保温层(A级): ", "【保温层】", "[保温]",
    "1、保温层（A级）：", "Ⅱ、保温层：", "甲、● ", "（三）保温：", "二、☆ ", "第Ⅱ层：", "B、", "(c)", "外墙H04：●",
    "No.1 ", "＜1＞", "§1、", "⬥", "🔸", "♢",
])
@pytest.mark.parametrize("mark", ["×", "x"])
def test_head_marks_by_morpheme_class(prefix, mark):
    v = gate(prefix + mark + "8161厚岩棉板保温层")
    assert v is not None and tokens(v) == [mark], (prefix, mark, tokens(v))
    assert gate(prefix + "8161厚岩棉板保温层") is None


@pytest.mark.parametrize("text", ["（3+3）×1层8162厚岩棉板保温层", "8163厚岩棉板保温层，规格600×1200",
                                  "8164厚岩棉板保温层（1.2厚）×2层", "8165厚岩棉板保温层，分格600 × 600"])
def test_specification_crosses_are_not_heads(text):
    assert gate(text) is None, (text, tokens(gate(text)))


@pytest.mark.parametrize("cell,mark", [
    ("1、8171厚岩棉板保温层，2、" + OL + "×，3、" + LV, "×"), ("1、" + PT + "✖，2、8172厚岩棉板保温层", "✖"),
    ("1、8173厚岩棉板保温层，2、" + OL + "／，3、" + LV, "／"), ("1、8174厚岩棉板保温层，2、" + OL + "——，3、" + LV, "——"),
    ("1、8175厚岩棉板保温层，2、" + OL + "×、3、" + LV, "×"),
])
def test_sibling_trailing_symbol_uses_the_real_right_neighbour(cell, mark):
    seg = text_seg(cell, "厚岩棉板保温层")
    row, hv = decided(seg)
    assert row is None and hv is not None and tokens(hv) == [mark] and links(hv)[0].startswith("同句"), (cell, tokens(hv))


def test_sibling_specification_cross_is_not_a_mark():
    seg = text_seg("1、8176厚岩棉板保温层，2、" + OL + "，规格600×600，3、" + LV, "8176厚")
    assert decided(seg) == (None, None)


@pytest.mark.parametrize("fmt,lvl,start", [("upperRoman", "%1、", 2), ("ideographTraditional", "%1、", 2),
                                           ("ideographZodiac", "%1、", 3), ("bullet", "o", 1), ("upperRoman", "（%1）", 4)])
def test_word_auto_numbering_head(fmt, lvl, start):
    body = _p(_r("外墙818：独立保温墙")) + _p(_r("×8181厚岩棉板保温层"), num=1)
    data, name, ext = _docx(body, numbering=_numbering(fmt, lvl, start))
    seg = [s for s in practice(qs.segment_file(data, name, ext, "", 0).segments) if "8181厚" in s.text][0]
    row, _ = decided(seg)
    assert row is not None and tokens(row) == ["×"], (fmt, seg.text, tokens(row))


# ====================================================================== 裁定 2：去除类前字、「掉」的缺陷语境
@pytest.mark.parametrize("tail,hit", [
    ("刷掉灰层", "掉"), ("冲掉砂层", "掉"), ("锤掉块料", "掉"), ("炸掉块石", "掉"), ("掰掉棱角", "掉"), ("冲除浮浆", "除"),
    ("刷除浮灰", "除"), ("钻除浮浆", "除"), ("清理掉旧漆", "掉"), ("处理掉旧漆", "掉"), ("冲走浮砂", "冲走"), ("刷去浮灰", "刷去"),
    ("处理掉灰层", "掉"), ("清理掉灰渣", "掉"), ("弄掉灰层", "掉"),
    ("缺棱掉角的砖挑出", None), ("防止掉落", None), ("表面起砂掉灰时返工", None), ("基层除尘后施工", None),
    ("已剝開檢查", None), ("撬开检查后复位", None),
    ("施工前除湿", None), ("换气扇洞口周边加强", None), ("吊顶带换气功能", None), ("换气百叶另详", None),
    # fix19：「换热计算」是计算术语，照算（不触发「换」），整词进 next_no_words；「换热固复合聚苯板」「换气凝胶毡」是换 + 材料，照算触发
    ("换热计算详节能专篇", None), ("换热固复合聚苯板", "换"), ("换气凝胶毡", "换"),
])
def test_removal_verb_class_and_diao_defect_context(tail, hit):
    v = gate("8201厚岩棉板保温层，" + tail)
    if hit is None:
        assert v is None, (tail, tokens(v))
    else:
        assert v is not None and hit in tokens(v), (tail, tokens(v))


# ====================================================================== 裁定 6：「除…外」——排除词不是整词时照配对
@pytest.mark.parametrize("tail,hit", [
    ("除湿区外均做", True), ("除草坪区外均做", True), ("除灰色区域外均做", True), ("除尘器间外均做", True),
    ("除渣场以外均做", True), ("施工前除湿，室外部分另行防护", False), ("基层除尘后施工，户外平台同做", False),
    ("铁件除油后刷防锈漆，室外同做", False),
])
def test_chu_wai_pairs_when_the_exclusion_is_not_the_whole_word(tail, hit):
    v = gate("8211厚岩棉板保温层，" + tail)
    assert (v is not None and "除" in tokens(v)) is hit, (tail, tokens(v))


# ====================================================================== 决定 2：繁体 / 异体补全
@pytest.mark.parametrize("tail", ["彆用", "颳除浮浆", "釦項", "挖齣", "改為擠塑聚苯板", "刪去", "本層棄用"])
def test_traditional_variants_of_inventory_chars_fold(tail):
    assert gate("8221厚岩棉板保温层，" + tail) is not None, tail


@pytest.mark.parametrize("tail", ["賸下的取消", "儘數取消"])
def test_traditional_quantifiers_fold(tail):
    seg = text_seg("1、8222厚岩棉板保温层。2、" + OL + "。" + tail, "8222厚")
    row, hv = decided(seg)
    assert row is None and hv is not None and links(hv)[0].startswith("全称"), (tail, links(hv))


# ====================================================================== 裁定 1：放行名单的 objects / after_gap / 无缝
@pytest.mark.parametrize("text,code,pend", [
    ("外墙面砖勾缝（无缝）", "Z011201003", True), ("1.5厚铝合金变形缝盖板，本层无缝", "010903004", True),
    ("8231厚铝合金框格栅装饰面层。本层无框", "011205001", True), ("8厚釉面砖面层；本层无釉", "011204003", True),
    ("8厚釉面砖面层。本层无釉面砖", "011204003", True), ("20厚砂浆粘结找平层；局部无粘结", "Z011101006", True),
    ("全玻无框玻璃幕墙，12厚钢化玻璃面板，无框玻璃肋支承", "Z011206006", False),
    ("12厚全玻无框玻璃幕墙，主体钢框架支承", "Z011206006", False),
    ("120厚C35无粘结预应力混凝土楼板，钢筋粘结长度详结施", "010502013", False),
    ("120厚C35无粘结预应力混凝土楼板，无粘结筋采用钢绞线", "010502013", False),
    ("10厚无釉防滑地砖面层", "011102003", False), ("12厚钢化玻璃无框门，门框不锈钢包边", "Z010805004", False),
    ("压入无碱玻纤网格布一层，网格布耐碱强力保留率不小于90%", "011001003", False),
    ("地坪漆面层（无缝）", "011101001", False), ("无缝钢管支架间距600", "011001003", False),
    ("A级不燃，本层取消", "011001003", True),
])
def test_safe_term_objects_guard_only_annotation_and_predicate_uses(text, code, pend):
    v = gate(text, code)
    assert (v is not None) is pend, (text, tokens(v))


def test_safe_list_stays_within_forty_entries_with_reasons():
    entries = rules.load_rules().get("doubt_safe_terms") or []
    assert len(entries) <= 40 and all(e.get("reason") for e in entries)
    terms = {t for e in entries for t in e.get("terms") or ()}
    assert not terms & {"不同", "不填料", "不带扶手", "不采用", "不做", "不设", "取消", "无此项"}
    assert any(e.get("objects") == ["缝"] for e in entries if "无缝" in (e.get("terms") or ()))


# ====================================================================== 裁定 5：主材格在规格格前面
@pytest.mark.parametrize("rows,probe,pend", [
    ([["1", "加气混凝土砌块墙（取消）", "8241厚，M5混合砂浆砌筑"]], "8241厚", True),
    ([["2", "外墙面砖饰面（本层取消）", "8242厚1:2水泥砂浆粘贴"]], "8242厚", True),
    ([["1", "8243厚岩棉板保温层", "砂浆禁止暴晒"]], "8243厚", False),
    ([["1", "8244厚岩棉板保温层", "XPS板另行采购，本层不用"]], "8244厚", False),
])
def test_cad_row_main_material_cell_before_the_spec_cell(rows, probe, pend):
    data, name, ext = _docx(_cad([["外墙824：独立保温墙"]] + rows))
    seg = [s for s in practice(qs.segment_file(data, name, ext, "", 0).segments) if probe in s.text][0]
    assert pending(seg) is pend, (seg.text, seg.scope_text())


# ====================================================================== 决定 1：首层层号粘在名称后、名称批注取宽
def test_first_layer_number_after_the_name_colon_is_recognised():
    cell = "外墙C19：1）8251厚1:3水泥砂浆找平层；2）" + OL + "；3）" + PT + "。倒数第一层取消"
    seg = text_seg(cell, "8251厚")
    assert rules._doubt_piece_no(seg.host.text, seg.host.pieces[seg.host_piece].start,
                                 seg.host.pieces[seg.host_piece].end) == 1
    assert decided(seg, "Z011101006") == (None, None)


def test_name_annotation_glued_to_the_first_layer_takes_wide_for_later_layers():
    seg = text_seg("外墙C03（取消）：1、60厚岩棉板保温层；2、8252厚抗裂砂浆面层", "8252厚")
    row, hv = decided(seg, "Z011201001")
    assert row is None and hv is not None and tokens(hv) == ["取消"] and links(hv)[0].startswith("取宽")


def test_first_layer_description_glued_before_a_colon_is_not_a_name():
    # 首层片「8260厚岩棉板保温层（取消），构造：1）…」冒号前是一层的描述（带厚度），不是做法名称：它的批注只管这一层，
    # 另一句的第 2 层不受影响（决定 1）
    cell = "8260厚岩棉板保温层（取消），构造：1）" + OL + "；2）刷8261号色外墙涂料两遍"
    seg = text_seg(cell, "8261号")
    assert decided(seg, P) == (None, None)


def test_layer_description_before_a_colon_is_not_a_name():
    # 「2)5厚…面层，以下取消：3)…」是第 2 层的描述 + 前指，不是做法名称：第 1 层不受影响
    seg = text_seg("1)8253厚岩棉板保温层；2)5厚聚合物抗裂砂浆面层，以下取消：3)10厚面砖层", "8253厚")
    assert decided(seg) == (None, None)


# ====================================================================== 管线 + API：待核行格式与只降级
def test_pipeline_and_api_downgrade_the_new_shapes_only(monkeypatch):
    rows = [_tr(_tc(_p(_r("■ 外墙826"))), _tc(_p(_r("独立保温墙"))), _tc(_p(_r("适用部位")))),
            _tr(_tc(_p(_r("1)8261厚岩棉板保温层"), _r("。第1层取消，本层亦取消"))), _tc(_p(_r("合成墙体")))),
            _tr(_tc(_p(_r("■ 外墙827"))), _tc(_p(_r("对照保温墙"))), _tc(_p(_r("适用部位")))),
            _tr(_tc(_p(_r("1)8271厚岩棉板保温层"))), _tc(_p(_r("合成墙体"))))]
    data, name, ext = _docx(_tbl(*rows))

    def pick(seg):
        return ("main", W) if "厚岩棉板保温层" in seg["原文"] else ("main", "")

    _fake(monkeypatch, pick)
    out = run([(data, name, ext)], agent=_AGENT)
    wb = load_workbook(BytesIO(out["xlsx_bytes"]))
    assert wb.sheetnames[:3] == [qr.SHEET_LIST, qr.SHEET_PENDING, qr.SHEET_COVERAGE]
    pend = [u for u in out["unmatched"] if "8261厚" in u["text"]]
    assert len(pend) == 1
    u = pend[0]
    assert (u["kind"], u["reason_code"], u["llm_guessed_code"]) == ("pending_review", "semantic_mismatch", W)
    assert "同句含否定/替代/取消表述" in u["reason"] and u["suggested"].startswith(f"建议编码：{W} ")
    assert "待人工确认是否采用" in u["suggested"]
    formal = [s["quote"] for i in out["items"] for s in i["sources"]]
    assert "8271厚岩棉板保温层" in formal and not any("8261厚" in q for q in formal)
    for item in out["items"]:
        assert not any(item.get(k) for k in ("quantity", "amount", "price", "total"))
    _fake(monkeypatch, pick)
    resp = _api(data, name, "pipeline-api")
    assert resp.status_code == 200
    body = resp.json()
    assert [(u["text"], u["reason_code"], u["llm_guessed_code"], u["suggested"]) for u in body["unmatched"]] == \
        [(u["text"], u["reason_code"], u["llm_guessed_code"], u["suggested"]) for u in out["unmatched"]]
    assert [(i["code"], [s["quote"] for s in i["sources"]]) for i in body["items"]] == \
        [(i["code"], [s["quote"] for s in i["sources"]]) for i in out["items"]]
