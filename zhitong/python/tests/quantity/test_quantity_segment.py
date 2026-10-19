# 来源：engine-service/tests/test_quantity_segment.py（zhitong-cost-ai@991ff5cd，W4 迁入 zt；改动见「zt 改动」注释）
"""工程量清单生成：确定性切片（WS11 QTY-03 / QTY-13，issue #763）。零 LLM。"""
from __future__ import annotations

import io
import zipfile

import pytest

from zt.quantity import rules
from zt.quantity import segment as qs
from zt.quantity.segment import EXCLUSION, HEADING, OTHER, PRACTICE
from tests.quantity.fixtures.shapes import (K2_DERUST, TITLED_PRACTICES, docx_bytes, docx_p, docx_table, k2_docx,
                                            near_title_docx, r1_cad_workbook, r9_workbook, titled_docx)


def _practice(segs):
    return [s for s in segs if s.kind == PRACTICE]


# ============================================================ 网格（r9 形）

def test_grid_sheet_r9_shape():
    fs = qs.segment_file(r9_workbook(), "做法表.xlsx", "xlsx", "", 0)
    assert fs.mode == "grid"
    segs = _practice(fs.segments)
    floor = [s for s in segs if s.group == "楼地面做法" and s.part == "地下汽车库"]
    # 合并单元格：部位 B3:B7 纵向合并，5 层都带「地下汽车库」、同一个做法
    assert [s.row for s in floor] == [3, 4, 5, 6, 7], "真实 Excel 行号，不是部位序号"
    assert [s.layer_no for s in floor] == [1, 2, 3, 4, 5], "层号来自层号列"
    assert len({s.practice_key for s in floor}) == 1
    assert floor[0].location == "做法表!D3（第3行·楼地面做法第1层）"
    assert floor[0].cell == "D3" and floor[0].sheet == "做法表"
    assert floor[0].category == "楼地面"
    # 新的合并块 = 新做法
    toilet = [s for s in segs if s.part == "卫生间"]
    assert toilet and {s.practice_key for s in toilet}.isdisjoint({floor[0].practice_key})
    assert toilet[0].location == "做法表!D8（第8行·楼地面做法第1层）"
    # 无层号的续行并入上一层
    skirting = [s for s in segs if s.group == "踢脚做法"]
    assert [s.layer_no for s in skirting] == [1, 2]
    assert skirting[1].text == "12厚1:3水泥砂浆打底扫毛 （踢脚高 150）"
    assert skirting[0].category == "踢脚" and skirting[0].practice == "地下汽车库·踢脚做法"
    # 隐藏行、隐藏工作表跳过并回报
    assert not any("隐藏行" in s.text for s in fs.segments)
    assert fs.hidden_sheets == ["作废"]
    assert any(code == "quantity_hidden_sheets" for code, _ in fs.warnings)
    # 表外格（备注列）按文本模式，不是做法
    remark = [s for s in fs.segments if "颜色见室内设计" in s.text]
    assert remark and remark[0].kind == OTHER and remark[0].cell == "I3"


def test_xlsx_location_is_real_cell():
    fs = qs.segment_file(r9_workbook(), "做法表.xlsx", "xlsx", "", 0)
    for s in _practice(fs.segments):
        assert s.location.startswith(f"做法表!{s.cell}（第{s.row}行")
        assert s.cell.endswith(str(s.row))


def test_hidden_column_skipped():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["B1"], ws["C1"] = "编号", "楼地面做法", "作废做法"
    ws["A2"], ws["B2"], ws["C2"] = "地面1", "20厚1:2水泥砂浆面层", "30厚作废砂浆"
    ws.column_dimensions["C"].hidden = True
    buf = io.BytesIO()
    wb.save(buf)
    fs = qs.segment_file(buf.getvalue(), "a.xlsx", "xlsx", "", 0)
    assert [s.text for s in _practice(fs.segments)] == ["20厚1:2水泥砂浆面层"]


# ============================================================ CAD 导出（r1 形）

def test_cad_export_r1_shape():
    fs = qs.segment_file(r1_cad_workbook(), "CAD做法.xlsx", "xlsx", "", 0)
    segs = _practice(fs.segments)
    roof = [s for s in segs if s.practice.startswith("屋面1")]
    assert [s.layer_no for s in roof] == [1, 2, 3, 4, 5, 6, 7], "一格两层要拆开；「3.3厚」按第 3 层读"
    assert roof[0].text == "40厚C20细石混凝土保护层" and roof[1].text == "10厚低标号砂浆隔离层"
    assert roof[0].row == roof[1].row == 3, "同一格拆出的两层位置都指向那一格"
    assert roof[2].text == "3厚SBS改性沥青防水卷材"
    assert roof[0].part == "主楼屋面" and roof[0].category == "屋面"
    ceiling = [s for s in segs if s.practice.startswith("顶棚1")]
    assert [s.text for s in ceiling] == ["9.5厚纸面石膏板", "U型轻钢龙骨中距600", "φ8吊杆"], \
        "「1.9.5厚」是第 1 层 9.5 厚，不是多级标题；拆进相邻两格的名称行要拼起来"
    assert ceiling[0].location == "Sheet1!A14（第14行）"
    walls = [s for s in segs if s.practice.startswith("内墙1")]
    assert len(walls) == 3 and walls[0].category == "墙面"


# ============================================================ 文本模式

TEXT = """[p1]
建筑设计说明
（四）屋面工程
屋面1：不上人屋面（用于：主楼屋面）
1. 40厚C20细石混凝土保护层
2. 10厚低标号砂浆隔离层
（内掺抗裂纤维）
3. 3厚SBS改性沥青防水卷材
[p2]
4.3 楼地面
地面1：地砖地面
1、10厚防滑地砖面层，20厚1:3干硬性水泥砂浆结合层
2、素水泥浆一道
盖边柱、道路两侧柱、盖板顶棚涂料等盖板装修计入IV标。
（五）桥梁工程
桥面铺装：100厚C40混凝土
外墙涂料颜色详见立面图
"""


def test_text_mode_page_anchor_and_clause():
    segs = qs.segment_text(TEXT, 0, "说明.pdf")
    pr = _practice(segs)
    roof = [s for s in pr if s.practice.startswith("屋面1")]
    assert [s.page for s in roof] == [1, 1, 1]
    assert roof[0].location == "p1 第4行（抽取文本行）"
    assert roof[1].text == "10厚低标号砂浆隔离层 （内掺抗裂纤维）", "无层号续行并入上一层"
    assert roof[0].heading == "屋面工程" and roof[0].part == "主楼屋面"
    floor = [s for s in pr if s.practice.startswith("地面1")]
    assert [s.text for s in floor] == ["10厚防滑地砖面层", "20厚1:3干硬性水泥砂浆结合层", "素水泥浆一道"], \
        "同一行两个「厚度+材料」拆成两层"
    assert floor[0].location == "p2 第3行（抽取文本行）" and floor[0].page == 2
    ex = [s for s in segs if s.kind == EXCLUSION]
    assert len(ex) == 1 and ex[0].exclusions[0][0] == "lot" and ex[0].exclusions[0][2] == 4
    heads = [s.text for s in segs if s.kind == HEADING]
    assert heads == ["（四）屋面工程", "4.3 楼地面", "（五）桥梁工程"]
    bridge = [s for s in pr if "C40" in s.text]
    assert bridge and bridge[0].section_domain == rules.MUNICIPAL
    assert [s.kind for s in segs if "颜色详见" in s.text] == [OTHER], "块外散行没有厚度不算做法"


# 阿拉伯序号节标题 + 编号层（PR #766 评审阻断项）：节标题行夹在编号层中间，层号从 1 重起。
# 改前「7.」「8.」被当成第 7、8 层吞进第一个做法，后面所有层都挂在「屋面做法」下。
ARABIC_SECTIONS = """6{sep}屋面做法
1. 3厚SBS改性沥青防水卷材
2. 20厚1:3水泥砂浆找平层
7{sep}楼面做法（卫生间）
1. 10厚防滑地砖面层
2. 20厚1:3干硬性水泥砂浆结合层
3. 3厚SBS改性沥青防水卷材
4. 20厚1:3水泥砂浆找平层
8{sep}外墙做法
1. 50厚挤塑聚苯板保温层
2. 20厚1:3水泥砂浆找平层
"""


def _blocks(segs):
    """按 practice_key 分组：[(做法名, 类别, [层号…])]，保持出现顺序。"""
    order: list[str] = []
    by: dict[str, list] = {}
    for s in _practice(segs):
        if s.practice_key not in by:
            order.append(s.practice_key)
            by[s.practice_key] = []
        by[s.practice_key].append(s)
    return [(by[k][0].practice, by[k][0].category, [s.layer_no for s in by[k]]) for k in order]


@pytest.mark.parametrize("sep", [".", "、", "．"])
def test_text_mode_arabic_section_headings_split_practices(sep):
    segs = qs.segment_text(ARABIC_SECTIONS.format(sep=sep), 0, "说明.txt")
    assert _blocks(segs) == [
        ("屋面做法", "屋面", [1, 2]),
        ("楼面做法（卫生间）", "楼地面", [1, 2, 3, 4]),
        ("外墙做法", "墙面", [1, 2]),
    ]
    assert [s.text for s in segs if s.kind == HEADING] == [
        f"6{sep}屋面做法", f"7{sep}楼面做法（卫生间）", f"8{sep}外墙做法"]
    assert not [s for s in segs if s.kind == OTHER], "节标题不再被当成第 7、8 层"


def test_text_mode_arabic_heading_with_suffix_equal_to_next_layer():
    # 节号恰好等于「下一层」的层号：带「做法 / 工程」尾缀的直接判节标题（含材料词「保温」也一样）
    segs = qs.segment_text("2.屋面做法\n1. 3厚SBS改性沥青防水卷材\n2. 20厚1:3水泥砂浆找平层\n"
                           "3.外墙保温做法\n1. 50厚挤塑聚苯板保温层\n2. 20厚1:3水泥砂浆找平层\n", 0, "a.txt")
    assert _blocks(segs) == [("屋面做法", "屋面", [1, 2]), ("外墙保温做法", "墙面", [1, 2])]
    assert [s.text for s in segs if s.kind == HEADING] == ["2.屋面做法", "3.外墙保温做法"]


def test_text_mode_arabic_heading_confirmed_by_layer_restart():
    # 没有尾缀的「3.外墙面」层号接得上，先按层收下；等下一行层号回到 1 才确认它是节标题
    segs = qs.segment_text("2.屋面做法\n1. 3厚SBS改性沥青防水卷材\n2. 20厚1:3水泥砂浆找平层\n"
                           "3.外墙面\n1. 50厚挤塑聚苯板保温层\n2. 20厚1:3水泥砂浆找平层\n", 0, "a.txt")
    assert _blocks(segs) == [("屋面做法", "屋面", [1, 2]), ("外墙面", "墙面", [1, 2])]
    assert [s.text for s in segs if s.kind == HEADING] == ["2.屋面做法", "3.外墙面"]
    wall = [s for s in segs if s.text == "3.外墙面"][0]
    assert wall.practice_key == "" and wall.layer_no is None and wall.category == ""


def test_text_mode_layer_number_restart_opens_new_practice():
    # 没有任何节标题行，只有层号从 1 重起（与网格模式「同一部位下层号重新从 1 起」对齐）
    segs = qs.segment_text("（六）屋面工程\n1. 3厚SBS改性沥青防水卷材\n2. 20厚1:3水泥砂浆找平层\n"
                           "1. 4厚SBS改性沥青防水卷材\n2. 30厚1:3水泥砂浆找平层\n", 0, "a.txt")
    assert _blocks(segs) == [("屋面工程", "屋面", [1, 2]), ("屋面工程", "屋面", [1, 2])]


def test_text_mode_arabic_heading_does_not_eat_real_layers():
    # 反向对照：块内编号跳号但正文是层（有材料词 / 施工动作）的仍是层；「3.3厚」照旧按第 3 层读
    segs = qs.segment_text("地面1：地砖地面\n1. 10厚防滑地砖面层\n2. 20厚1:3干硬性水泥砂浆结合层\n"
                           "4. 素水泥浆一道\n6、刷基层处理剂一遍\n7. 详见05J1工程做法\n"
                           "8.3厚SBS改性沥青防水卷材\n", 0, "a.txt")
    assert not [s for s in segs if s.kind == HEADING]
    assert _blocks(segs) == [("地面1 地砖地面", "楼地面", [1, 2, 4, 8])]
    assert [s.text for s in segs if s.kind == OTHER] == ["刷基层处理剂一遍", "详见05J1工程做法"]
    # 层号接得上、正文带类别词但后面没有「层号回到 1」：仍是本做法的最后一层，不是节标题
    segs = qs.segment_text("屋面1：不上人屋面\n1. 3厚SBS改性沥青防水卷材\n2. 20厚1:3水泥砂浆找平层\n"
                           "3. 现浇屋面板\n", 0, "a.txt")
    assert not [s for s in segs if s.kind == HEADING]
    assert [s.text for s in segs if s.kind == OTHER] == ["现浇屋面板"]


# 做法里嵌套的子序号（PR #766 第二轮评审阻断项）：「2. 防水层：」下面按视觉换行抽出的「①…②…」「1、…」「1）…」。
# 上一轮「行首层号回到 1 就新开做法」把 ① 当成新做法，而且新做法的类别、做法名、部位全是空串 ——
# 流水线上卷材 / 涂膜被静默编成楼(地)面防水（status=ok，不进待核）。
NESTED_SUB = ("屋面1：不上人屋面（用于主楼屋面）\n1. 保护层：40厚C20细石混凝土\n2. 防水层：\n"
              "{a}3厚SBS改性沥青防水卷材\n{b}2厚聚氨酯防水涂膜\n3. 找平层：20厚1:3水泥砂浆\n")
ROOF1 = ("屋面", "屋面1 不上人屋面（用于主楼屋面）", "主楼屋面")


@pytest.mark.parametrize("a,b,sub_no", [
    ("①", "②", 2), ("1、", "2、", 2), ("1）", "2）", 2), ("1)", "2)", 2),
    # 「(1)」不是层号标记，原样留在原文里、没有层号；改前改后都不受影响，放在这里当对照
    ("(1)", "(2)", None), ("（1）", "（2）", None),
])
def test_text_mode_nested_sub_numbers_keep_practice_context(a, b, sub_no):
    pr = _practice(qs.segment_text(NESTED_SUB.format(a=a, b=b), 0, "说明.txt"))
    assert len(pr) == 5
    assert len({s.practice_key for s in pr}) == 1, "子序号是第 2 层「防水层」的子项，不是新做法"
    for s in pr:
        assert (s.category, s.practice, s.part) == ROOF1, f"子层丢了所在做法的上下文：{s.text}"
    assert [s.layer_no for s in pr] == [1, 2, sub_no, sub_no, 3], "子项挂在父层号下，不推进本块层号"
    assert pr[2].text.endswith("3厚SBS改性沥青防水卷材") and pr[3].text.endswith("2厚聚氨酯防水涂膜")


def test_text_mode_sub_numbers_under_dun_style_layers():
    # 标记样式按本块首层定，不是写死「N. 是一级」：首层用「1、」时，「①」仍是子项
    pr = _practice(qs.segment_text("屋面1：不上人屋面（用于主楼屋面）\n1、40厚C20细石混凝土保护层\n2、防水层：\n"
                                   "①3厚SBS改性沥青防水卷材\n②2厚聚氨酯防水涂膜\n3、20厚1:3水泥砂浆找平层\n",
                                   0, "a.txt"))
    assert len({s.practice_key for s in pr}) == 1
    assert {(s.category, s.practice, s.part) for s in pr} == {ROOF1}
    assert [s.layer_no for s in pr] == [1, 2, 2, 2, 3]


def test_text_mode_same_style_restart_inherits_practice_context():
    # 同样式的层号回到 1、前面又没有节标题：新做法（另起 practice_key，同码不在做法内合并），
    # 做法名、类别、部位沿用当前做法 —— 与网格模式「同一部位下层号重新从 1 起」一致，不清空
    segs = qs.segment_text("屋面1：不上人屋面（用于主楼屋面）\n1. 3厚SBS改性沥青防水卷材\n2. 防水层：\n"
                           "①2厚聚氨酯防水涂膜\n②1.5厚JS防水涂料\n3. 20厚1:3水泥砂浆找平层\n"
                           "1. 4厚SBS改性沥青防水卷材\n2. 30厚1:3水泥砂浆找平层\n", 0, "a.txt")
    assert _blocks(segs) == [(ROOF1[1], "屋面", [1, 2, 2, 2, 3]), (ROOF1[1], "屋面", [1, 2])]
    assert {s.part for s in _practice(segs)} == {"主楼屋面"}


def test_text_mode_same_style_nested_numbers_keep_context():
    # 子序号与一级层同样式（「2. 防水层：」下接「1. …2. …」）：切片分不出是子项还是重起，按重起另开做法；
    # 但上下文沿用，不会像上一轮那样清成空串
    segs = qs.segment_text(NESTED_SUB.format(a="1. ", b="2. "), 0, "a.txt")
    assert _blocks(segs) == [(ROOF1[1], "屋面", [1, 2]), (ROOF1[1], "屋面", [1, 2, 3])]
    assert {(s.category, s.practice, s.part) for s in _practice(segs)} == {ROOF1}


def test_sidecar_page_anchor_format():
    segs = qs.segment_text("## p3\n地面1：水泥地面\n1. 20厚1:2水泥砂浆面层\n", 0, "扫描.pdf")
    pr = _practice(segs)
    assert pr[0].page == 3 and pr[0].location == "p3 第2行（抽取文本行）"


@pytest.mark.parametrize("line,expected", [
    ("1.20厚C20细石混凝土面层 2.素水泥浆一道 3.1.5厚聚氨酯防水涂料",
     [(1, "20厚C20细石混凝土面层"), (2, "素水泥浆一道"), (3, "1.5厚聚氨酯防水涂料")]),
    ("1.5厚JS防水涂料", [(None, "1.5厚JS防水涂料")]),
    ("2.0厚聚氨酯防水涂料", [(None, "2.0厚聚氨酯防水涂料")]),
    ("① 10厚地砖 ② 20厚1:3干硬性水泥砂浆结合层", [(1, "10厚地砖"), (2, "20厚1:3干硬性水泥砂浆结合层")]),
    ("20厚1:3水泥砂浆找平层，1.5厚聚氨酯防水涂料，40厚C20细石混凝土保护层",
     [(None, "20厚1:3水泥砂浆找平层"), (None, "1.5厚聚氨酯防水涂料"), (None, "40厚C20细石混凝土保护层")]),
    ("3. 3厚SBS改性沥青防水卷材", [(3, "3厚SBS改性沥青防水卷材")]),
])
def test_split_layers(line, expected):
    assert qs.split_layers(line) == expected


def test_split_layers_uses_expected_next_layer_for_ambiguous_decimal():
    assert qs.split_layers("3.3厚SBS改性沥青防水卷材") == [(None, "3.3厚SBS改性沥青防水卷材")]
    assert qs.split_layers("3.3厚SBS改性沥青防水卷材", expect_next=3) == [(3, "3厚SBS改性沥青防水卷材")]


# ============================================================ docx（WS04 块级 API）

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _p(t: str) -> str:
    return f"<w:p><w:r><w:t>{t}</w:t></w:r></w:p>"


def _tr(*cells: str) -> str:
    return "<w:tr>" + "".join(f"<w:tc>{_p(c)}</w:tc>" for c in cells) + "</w:tr>"


def _docx(body: str) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>')
        z.writestr("word/document.xml", f'<w:document xmlns:w="{_W}"><w:body>{body}</w:body></w:document>')
    return buf.getvalue()


def test_docx_table_and_paragraph_locations():
    body = (_p("工程做法说明") + _p("屋面防水：2.0厚聚氨酯防水涂料")
            + "<w:tbl>" + _tr("编号", "名称", "构造做法") + _tr("地面1", "地砖地面", "1.10厚地砖面层")
            + _tr("", "", "2.20厚1:3干硬性水泥砂浆结合层") + _tr("楼面2", "水泥楼面", "1.20厚1:2水泥砂浆面层")
            + "</w:tbl>")
    fs = qs.segment_file(_docx(body), "说明.docx", "docx", "", 0)
    assert fs.mode == "docx"
    pr = _practice(fs.segments)
    note = [s for s in pr if "聚氨酯" in s.text]
    assert note and note[0].location == "段落#2" and note[0].paragraph == 2
    tbl = [s for s in pr if s.location.startswith("表1")]
    assert [s.location for s in tbl] == ["表1 第2行第3列（构造做法第1层）", "表1 第3行第3列（构造做法第2层）",
                                         "表1 第4行第3列（构造做法第1层）"]
    assert tbl[0].practice_key == tbl[1].practice_key != tbl[2].practice_key, "空编号格 = 同一做法的续行"
    assert tbl[0].practice == "地面1 地砖地面"


def test_docx_unparseable_falls_back_to_lines_with_warning():
    fs = qs.segment_file(b"not a zip", "坏.docx", "docx", "地面1：水泥地面\n1. 20厚1:2水泥砂浆面层\n", 0)
    assert fs.mode == "text"
    assert [c for c, _ in fs.warnings] == ["quantity_docx_line_mode"]
    assert _practice(fs.segments)


def test_assign_ids_is_global_and_ordered():
    a = qs.segment_file(None, "a.txt", "txt", "地面1：水泥\n1. 20厚1:2水泥砂浆面层\n", 0)
    b = qs.segment_file(None, "b.txt", "txt", "地面1：水泥\n1. 30厚1:2水泥砂浆面层\n", 1)
    segs = qs.assign_ids([a, b])
    assert [s.id for s in segs] == [f"s{i}" for i in range(1, len(segs) + 1)]
    assert [s.file_index for s in _practice(segs)] == [0, 1]


# ============================================================ WS29-f 先切层、后判排除（issue #835）
# 生产实测的做法表没有表头：一行各格拼成一句走文本模式。改前整行先判排除句，命中就整行变成一个
# exclusion 片段 —— 一层写了「自理」，同一行的抗裂砂浆、保温板、防水涂料一起从覆盖核算里消失。

def _code(seg):
    return seg.inline.code if seg.inline else ""


# 第 4 层带着材料词（素水泥浆）。只写「基层墙体」四个字的层词表认不出、会落成 other，那是既有行为（归 k2），
# 不在这条夹具要钉的范围里
ROW_OWNER = ("1)面层瓷砖用户自理（详二次装修） 2)5厚抗裂砂浆压入一道玻纤网格布 3)1.5厚Ⅱ型聚合物水泥防水涂料 "
             "4)基层墙体（混凝土墙梁柱处素水泥浆甩毛）")
ROW_NEGATION = ("1)无机涂料饰面 2)满刮2厚柔性耐水腻子两遍分层找平 3)6厚DP15预拌砂浆找平 4)15厚DP15预拌砂浆分层打底 "
                "5)基层墙体（混凝土墙梁柱处素水泥浆甩毛）")
PART_NEGATION = "楼梯间、楼梯间前室、地上变配电室、消控中心、计量间、门卫等不需二次装修的非配水房间"


def _headerless_docx():
    return docx_bytes(docx_table(
        ("■ 内墙5", "面砖墙面", "30厚", "适用部位"),
        (ROW_OWNER, "住宅厨房、卫生间"),
        ("■ 内墙9", "无机涂料墙面", "25厚", "适用部位"),
        (ROW_NEGATION, PART_NEGATION),
    ) + docx_p("注：本表上注明详二次装修设计部位的材料做法仅为暂定，规格材质最终由装修设计确定。"))


def test_headerless_row_sliced_before_exclusion():
    fs = qs.segment_file(_headerless_docx(), "做法表.docx", "docx", "", 0)
    assert fs.mode == "docx"
    assert not [s for s in fs.segments if s.kind == EXCLUSION], "这份表里没有任何全局排除句（基线认出 8 条，全错）"
    row = [s for s in fs.segments if s.location.startswith("表1 第2行")]
    pr = _practice(row)
    # WS29-k1 起「适用部位」栏的文字是部位，不再接在最后一层后面
    assert [s.text for s in pr] == ["面层瓷砖用户自理（详二次装修）", "5厚抗裂砂浆压入一道玻纤网格布",
                                   "1.5厚Ⅱ型聚合物水泥防水涂料", "基层墙体（混凝土墙梁柱处素水泥浆甩毛）"], \
        "一行四层都是 practice 片段：自理的那一层也要进覆盖核算"
    assert [s.layer_no for s in pr] == [1, 2, 3, 4]
    assert {s.part for s in pr} == {"住宅厨房、卫生间"}
    assert [_code(s) for s in pr] == [rules.OWNER_INLINE, "", "", ""], "自理只管第 1 层"
    assert pr[0].inline.also == "详二次装修"


def test_headerless_row_with_negation_keeps_every_layer():
    # 适用部位「…等不需二次装修的非配水房间」。改前它接在最后一层后面、整行成了排除句、五层全部消失；
    # WS29-k1 起它是这个做法的部位，不进层正文
    fs = qs.segment_file(_headerless_docx(), "做法表.docx", "docx", "", 0)
    row = _practice([s for s in fs.segments if s.location.startswith("表1 第4行")])
    assert [s.layer_no for s in row] == [1, 2, 3, 4, 5]
    assert row[0].text == "无机涂料饰面" and row[-1].text == "基层墙体（混凝土墙梁柱处素水泥浆甩毛）"
    assert {s.part for s in row} == {PART_NEGATION}
    assert [_code(s) for s in row] == [""] * 5, "「不需二次装修」的部位说明不触发任何排除或标记"


@pytest.mark.parametrize("line", [ROW_OWNER, ROW_NEGATION + " " + PART_NEGATION])
def test_text_mode_layer_row_never_becomes_exclusion(line):
    # 同样的行走纯文本（pdf / txt）路径
    segs = qs.segment_text(line, 0, "做法表.txt")
    assert [s.kind for s in segs if s.kind == EXCLUSION] == []
    assert len(_practice(segs)) >= 4


def test_note_paragraph_sentences_are_not_layers():
    # 说明段落、表注：不成层，也不是排除句式 → 不产生排除句，原样留作 other / heading
    segs = qs.segment_text("注：本表上注明详二次装修设计部位的材料做法仅为暂定，规格材质最终由装修设计确定。\n"
                           "13.13 二次装修室内空间应考虑隔声降噪措施。\n17.3 厨房设备见二次装修。\n", 0, "说明.txt")
    assert [s.kind for s in segs] == [OTHER, HEADING, HEADING]


def test_inline_marker_on_continuation_line():
    # 一层分两行，第二行只有「（用户自理）」：必须在续行并进上一层之后再识别
    segs = qs.segment_text("楼面2：卫生间楼面\n1. 预留防滑地砖装修面层\n（用户自理）\n"
                           "2. 20厚1:3干硬性水泥砂浆结合层\n", 0, "a.txt")
    pr = _practice(segs)
    assert [s.text for s in pr] == ["预留防滑地砖装修面层 （用户自理）", "20厚1:3干硬性水泥砂浆结合层"]
    assert [_code(s) for s in pr] == [rules.OWNER_INLINE, ""]


def test_grid_practice_column_inline_marker_and_continuation():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["B1"], ws["C1"], ws["D1"] = "编号", "部位", "层次", "楼地面做法"
    rows = [("楼面2", "卫生间", 1, "预留防滑地砖装修面层"), ("", "", None, "（用户自理）"),
            ("", "", 2, "30厚C25细石混凝土找平找坡（用户自理）"), ("", "", 3, "1.5厚聚合物水泥防水涂料（详精装修设计）"),
            ("", "", 4, "面层为二次装修"), ("", "", 5, "面层瓷砖用户自理或（详精装修设计）"),
            ("", "", 6, "防滑地砖面层由业主自理"), ("", "", 7, "20厚1:3水泥砂浆找平层"),
            ("楼面3", "走廊", 1, "外廊真石漆面层计入IV标"), ("", "", 2, "20厚1:3水泥砂浆找平层")]
    for i, (code, part, no, text) in enumerate(rows, start=2):
        for col, v in zip("ABCD", (code, part, no, text)):
            if v not in ("", None):
                ws[f"{col}{i}"] = v
    buf = io.BytesIO()
    wb.save(buf)
    fs = qs.segment_file(buf.getvalue(), "做法表.xlsx", "xlsx", "", 0)
    toilet = [s for s in _practice(fs.segments) if s.part == "卫生间"]
    assert [s.layer_no for s in toilet] == [1, 2, 3, 4, 5, 6, 7], "带标记的层仍是 practice 片段"
    assert toilet[0].text == "预留防滑地砖装修面层 （用户自理）"
    # 第 6 层「…由业主自理」是全局 owner 句式认得的写法：写在做法列里也只管这一层，不拿「防滑地砖面层」去全局匹配
    assert [_code(s) for s in toilet] == [rules.OWNER_INLINE, rules.OWNER_INLINE, rules.FITOUT_PROVISIONAL,
                                          rules.SECONDARY_INLINE, rules.OWNER_AMBIGUOUS, rules.OWNER_INLINE, ""]
    # 标段句写在做法列里仍是全局排除句（不变）
    ex = [s for s in fs.segments if s.kind == EXCLUSION]
    assert [(s.text, s.exclusions[0][0], s.exclusions[0][2]) for s in ex] == [("外廊真石漆面层计入IV标", "lot", 4)]


def test_marked_layer_without_lexicon_word_is_still_practice():
    # 「装修吊顶（详二次装修图纸）」词表里没有材料词，也没有厚度：带着标记就是做法层，要进覆盖核算
    segs = qs.segment_text("1)装修吊顶（详二次装修图纸） 2)基层清理 3)钢筋混凝土板底", 0, "a.txt")
    pr = _practice(segs)
    assert pr[0].text == "装修吊顶（详二次装修图纸）" and _code(pr[0]) == rules.FITOUT_PROVISIONAL
    # 反向对照：说明条款「3、厨房设备见二次装修」不是层，不能因为带「见二次装修」就变成做法片段
    segs = qs.segment_text("（四）装修说明\n3、厨房设备见二次装修\n", 0, "a.txt")
    assert _practice(segs) == []


def test_fitout_reference_on_practice_name_marks_every_layer():
    segs = qs.segment_text("顶棚3：装修吊顶（详二次装修图纸）\n1. U型轻钢龙骨中距600\n2. 9.5厚纸面石膏板\n"
                           "顶棚4：乳胶漆顶棚\n1. 刷白色乳胶漆两遍\n2. 满刮腻子两遍\n", 0, "a.txt")
    pr = _practice(segs)
    assert [(_code(s), s.inline.where if s.inline else "") for s in pr] == [
        (rules.FITOUT_PROVISIONAL, "practice"), (rules.FITOUT_PROVISIONAL, "practice"), ("", ""), ("", "")]
    # 章节标题不是做法名称：标题里的「见二次装修」不下传给下面的层
    segs = qs.segment_text("17.3 厨房设备见二次装修。\n1. 20厚1:3水泥砂浆找平层\n2. 10厚防滑地砖面层\n", 0, "a.txt")
    assert [_code(s) for s in _practice(segs)] == ["", ""]


def test_numbered_clause_outside_practice_stays_global():
    # 只带一个层号、又不像做法层的说明条款：仍按全局排除句判（QTY-11 不回归）
    segs = qs.segment_text("（四）装修说明\n1、盖边柱、道路两侧柱等盖板装修计入IV标。\n"
                           "2、卫生间洁具由业主另行发包\n3、楼地面面层均为二次装修。\n"
                           "4、屋面保温不在本次设计范围\n", 0, "说明.txt")
    ex = [s for s in segs if s.kind == EXCLUSION]
    assert [s.exclusions[0][0] for s in ex] == ["lot", "owner", "secondary", "not_in_scope"]
    assert ex[2].exclusions[0][1] == ["楼地面面层"]
    assert _practice(segs) == []


def test_layer_row_with_owner_sentence_is_sliced_first():
    # 「…由业主自理」是全局 owner 句式认得的写法。整行先判排除的话，这一行会整个变成一条排除句（practice 为 0），
    # 还拿「防滑地砖面层」去全局匹配别的做法。先切层：三层都在，自理只管第 1 层
    segs = qs.segment_text("1)防滑地砖面层由业主自理 2)20厚1:3干硬性水泥砂浆结合层 3)1.5厚聚合物水泥防水涂料", 0, "a.txt")
    assert [s for s in segs if s.kind == EXCLUSION] == []
    pr = _practice(segs)
    assert [s.layer_no for s in pr] == [1, 2, 3]
    assert [_code(s) for s in pr] == [rules.OWNER_INLINE, "", ""]


def test_numbered_layer_under_practice_name_is_inline_not_global():
    # 做法名称下面的编号行是这个做法的层：「业主自理」只管这一层，不拿「防滑地砖面层」去全局匹配
    segs = qs.segment_text("楼面2：卫生间楼面\n1. 防滑地砖面层由业主自理\n2. 20厚1:3干硬性水泥砂浆结合层\n", 0, "a.txt")
    assert [s for s in segs if s.kind == EXCLUSION] == []
    assert [_code(s) for s in _practice(segs)] == [rules.OWNER_INLINE, ""]


def test_rejected_sentence_is_not_merged_into_previous_layer():
    # 「注：面层为二次装修」对象只有泛指词，不成排除句。它是一句说明，不能当续行并进上一层 ——
    # 并进去上一层就凭空多了一个「为二次装修」的标记
    segs = qs.segment_text("地面1：地砖地面\n1. 10厚防滑地砖面层\n注：面层为二次装修\n"
                           "2. 20厚1:3干硬性水泥砂浆结合层\n", 0, "a.txt")
    pr = _practice(segs)
    assert [s.text for s in pr] == ["10厚防滑地砖面层", "20厚1:3干硬性水泥砂浆结合层"]
    assert [_code(s) for s in pr] == ["", ""]
    assert [s.text for s in segs if s.kind == OTHER] == ["注：面层为二次装修"]
    assert [s for s in segs if s.kind == EXCLUSION] == []


def test_lot_sentence_inside_layer_row_still_global():
    segs = qs.segment_text("1)喷真石漆两遍 2)盖边柱真石漆计入IV标 3)20厚1:3水泥砂浆找平层", 0, "a.txt")
    ex = [s for s in segs if s.kind == EXCLUSION]
    assert [(s.text, s.exclusions[0][0]) for s in ex] == [("盖边柱真石漆计入IV标", "lot")]
    assert [s.text for s in _practice(segs)] == ["喷真石漆两遍", "20厚1:3水泥砂浆找平层"]


# —— 做法名称块后面的编号表注（PR #839 评审阻断项）——
# 「判成层，所以不整行判全局排除句」和「真的按层处理了」必须是同一个结论。改前名称块里只带一个层号的行
# 一律判成层；层号对不上块里的下一层时，它却落进节标题分支 —— 既不成排除句、也不成层，哪一页都查不到，
# 另一个做法里该排掉的层进了正式清单（相对 dev 的 QTY-11 回归）。

_NAMED = "楼面2：卫生间楼面\n1. 10厚防滑地砖面层\n2. 20厚1:3水泥砂浆找平层\n"
_OWNER_NOTE, _SECONDARY_NOTE = "卫生间洁具、橱柜由业主另行发包。", "住宅户内墙面均为二次装修。"


def _global(segs):
    return [(s.text, s.exclusions[0][0], s.exclusions[0][1]) for s in segs if s.kind == EXCLUSION]


@pytest.mark.parametrize("first,second", [
    ("注：\n1. ", "2. "),       # 「注：」单独一行，表注从 1 重新编号
    ("注：1. ", "2. "),         # 「注：」与第 1 条写在同一行
    ("注：\n1、", "2、"),       # 表注的标记样式与层不同（「1、」对「1.」）
    ("1. ", "2. "),             # 没有「注：」，做法块后面直接从 1 重新编号
])
def test_numbered_table_notes_after_named_block_stay_global(first, second):
    text = _NAMED + first + _OWNER_NOTE + "\n" + second + _SECONDARY_NOTE + "\n"
    segs = qs.segment_text(text, 0, "a.txt")
    lines = text.splitlines()
    assert _global(segs) == [
        (next(x for x in lines if _OWNER_NOTE in x), "owner", ["卫生间洁具", "橱柜"]),
        (lines[-1], "secondary", ["住宅户内墙面"]),
    ]
    pr = _practice(segs)
    assert [(s.layer_no, _code(s)) for s in pr] == [(1, ""), (2, "")], "表注不是层，也不给哪一层挂标记"
    assert not [s for s in segs if s.kind == HEADING], "表注不能落成节标题（落成标题就哪一页都查不到）"


@pytest.mark.parametrize("text", [
    # 做法只有 1 层：第 2 条表注的号恰好等于块里等着的下一层
    "顶棚2：乳胶漆顶棚\n1. 刷白色乳胶漆两遍\n注：\n1. " + _OWNER_NOTE + "\n2. " + _SECONDARY_NOTE + "\n",
    # 第 1 条不是排除句、层号回到 1 另起了一段：第 2 条的号同样接得上
    _NAMED + "注：\n1. 未尽事宜详见国家规范\n2. " + _SECONDARY_NOTE + "\n",
    _NAMED + "备注： 1. 未尽事宜详见国家规范\n2. " + _SECONDARY_NOTE + "\n",
])
def test_table_note_whose_number_fits_the_block_is_still_global(text):
    # 「注：」之后是表注条款，层号恰好接得上也不是这个做法的层
    segs = qs.segment_text(text, 0, "a.txt")
    assert ("2. " + _SECONDARY_NOTE, "secondary", ["住宅户内墙面"]) in _global(segs)
    assert not [s.text for s in _practice(segs) if "二次装修" in s.text]


def test_long_clause_out_of_sequence_under_practice_name_is_global():
    # 超过节标题长度上限的条款落不进标题分支。改前按层收下（层号回到 1、另起一段），发包句变成「只管这一层」
    clause = "1、卫生间洁具、橱柜、淋浴房及其全部五金配件均由业主另行发包。"
    segs = qs.segment_text(_NAMED + clause + "\n", 0, "a.txt")
    assert [(t, k) for t, k, _o in _global(segs)] == [(clause, "owner")]
    assert "卫生间洁具" in _global(segs)[0][2]
    assert [s.text for s in _practice(segs)] == ["10厚防滑地砖面层", "20厚1:3水泥砂浆找平层"]


def test_layer_row_that_never_reaches_layers_is_still_scanned():
    # 带厚度和材料词，判成层；开头的「屋面40」又让它走了做法名称行分支，名称后面切不出层。
    # 改前这句话不产出任何片段
    owner = "屋面40厚细石混凝土保护层由业主另行发包"
    segs = qs.segment_text("楼面2：卫生间楼面\n1. 10厚防滑地砖面层\n" + owner + "\n", 0, "a.txt")
    assert [(t, k) for t, k, _o in _global(segs)] == [(owner, "owner")]
    # 一行两个层号，判成层；又短得像节标题，走了标题分支。改前整行变成一个 heading
    row = "1、洁具业主自理 2、橱柜业主自理"
    segs = qs.segment_text("（四）装修说明\n" + row + "\n", 0, "a.txt")
    assert [(t, k) for t, k, _o in _global(segs)] == [(row, "owner")]
    assert [s.text for s in segs if s.kind == HEADING] == ["（四）装修说明"]


def test_inline_note_between_layers_keeps_named_layers():
    # 反向对照：夹在两层之间的一句「注：…」不是表注列表的开头，后面的编号行仍是这个做法的层
    segs = qs.segment_text("顶棚3：轻钢龙骨吊顶\n1. U型轻钢龙骨中距600\n注：龙骨间距按厂家要求\n"
                           "2. 装修吊顶（详二次装修图纸）\n", 0, "a.txt")
    assert [(s.layer_no, _code(s)) for s in _practice(segs)] == [(1, ""), (2, rules.FITOUT_PROVISIONAL)]


def test_layers_after_bare_note_header_keep_practice_context():
    # 反向对照：「注：」不结束做法块。结束了，后面真正的层会丢掉做法名称、部位和类别，被编成别类的码
    segs = qs.segment_text("地面1：地砖地面（用于门厅）\n1. 10厚防滑地砖面层\n注：\n"
                           "2. 20厚1:3干硬性水泥砂浆结合层\n", 0, "a.txt")
    pr = _practice(segs)
    assert [s.layer_no for s in pr] == [1, 2] and len({s.practice_key for s in pr}) == 1
    assert {(s.practice, s.part, s.category) for s in pr} == {("地面1 地砖地面（用于门厅）", "门厅", "楼地面")}


def test_inline_secondary_layer_keeps_marker_after_continuation():
    # 续行并进上一层之后标记还在。改前「二次装修」后面只认句末和标点，并完续行这一层就照常送模型
    segs = qs.segment_text("地面1：地砖地面\n1. 面层为二次装修\n详见装修图\n"
                           "2. 20厚1:3干硬性水泥砂浆结合层\n", 0, "a.txt")
    assert [(s.text, _code(s)) for s in _practice(segs)] == [
        ("面层为二次装修 详见装修图", rules.SECONDARY_INLINE), ("20厚1:3干硬性水泥砂浆结合层", "")]


# ============================================================ 标题行 + 构造层行（WS29-k1，issue #891）

FLOOR9, WALL9, CEIL2, STEP2 = ("楼面9 花岗岩楼面1（燃烧性能A级）", "内墙9 无机涂料内墙面1（燃烧性能A级）",
                               "顶棚2 石膏腻子顶棚2（燃烧性能A级）", "其它2 花岗石台阶")


def _titled():
    return qs.segment_file(titled_docx(), "做法表.docx", "docx", "", 0)


def _by_practice(segs):
    out: dict[str, list] = {}
    for s in _practice(segs):
        out.setdefault(s.practice, []).append(s)
    return out


def test_title_rows_produce_no_segment():
    fs = _titled()
    assert fs.mode == "docx"
    # 标题行在第 1、3、5、7 行：一个片段都不出（基线：4 个标题行成了 practice，或按续行并进上一层）
    assert sorted({s.row for s in fs.segments}) == [2, 4, 6, 8]
    assert not [s.text for s in fs.segments if "■" in s.text or "适用部位" in s.text]
    assert list(_by_practice(fs.segments)) == [FLOOR9, WALL9, CEIL2, STEP2], "做法名 = 编号 + 名称，不带总厚度"
    assert len({s.practice_key for s in _practice(fs.segments)}) == 4


def test_titled_layers_carry_name_part_category():
    by = _by_practice(_titled().segments)
    floor = by[FLOOR9]
    assert [s.layer_no for s in floor] == [1, 2, 3, 4]
    assert {(s.part, s.category) for s in floor} == {("电梯厅、公共走道", "楼地面")}
    # 每个做法的部位取自它自己那一行的「适用部位」栏
    for name, (_title, body) in zip((FLOOR9, WALL9, CEIL2, STEP2), TITLED_PRACTICES):
        assert {s.part for s in by[name]} == {body[-1]}, name
    assert {s.category for s in by[WALL9]} == {"墙面"} and {s.category for s in by[CEIL2]} == {"顶棚"}
    assert {s.category for s in by[STEP2]} == {"台阶"}, "编号头「其它」不是部位词：类别按名称判"


def test_last_layer_has_no_part_text_nor_next_title():
    by = _by_practice(_titled().segments)
    texts = [s.text for v in by.values() for s in v]
    for title, body in TITLED_PRACTICES:
        assert not [t for t in texts if body[-1] in t], f"部位文字不进层正文：{body[-1]}"
        assert not [t for t in texts if title[1] in t], f"标题行的文字不进任何一层：{title[1]}"
    assert by[FLOOR9][-1].text == "钢筋混凝土楼板"
    # 下一个标题行（顶棚2）没有厚度：改前它按续行并进内墙9 的第 3 层
    assert [s.text for s in by[WALL9]] == ["无机涂料（一底两面）饰面", "满刮2厚柔性耐水腻子两遍分层找平", "6厚预拌砂浆找平"]
    assert by[CEIL2][-1].text == "钢筋混凝土楼板清理"


def test_negation_in_part_triggers_no_exclusion():
    segs = _titled().segments
    assert not [s for s in segs if s.kind == EXCLUSION]
    wall = _by_practice(segs)[WALL9]
    assert "不需二次装修" in wall[0].part
    assert [_code(s) for s in wall] == ["", "", ""], "部位是做法的属性、不是句子：不排除，也不打标记"


def test_titled_locations_use_real_columns():
    by = _by_practice(_titled().segments)
    assert {s.location for s in by[FLOOR9]} == {"表1 第2行第1列"}
    assert {s.part_location for s in by[FLOOR9]} == {"表1 第2行第2列（适用部位）"}
    # 编号格纵向合并时，构造层行首格是空的：层在第 2 列、部位在第 4 列
    fs = qs.segment_file(docx_bytes(docx_table(
        ("■ 楼面1", "地砖楼面", "50厚", "适用部位"),
        ("", "1)10厚防滑地砖面层 2)20厚1:3干硬性水泥砂浆结合层", "", "卫生间"))), "a.docx", "docx", "", 0)
    pr = _practice(fs.segments)
    assert [(s.location, s.part, s.part_location) for s in pr] == [
        ("表1 第2行第2列", "卫生间", "表1 第2行第4列（适用部位）")] * 2


def test_titled_xlsx_merged_layers_heading_row_and_part_over_rows():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["■ 楼面9", "花岗岩楼面1（燃烧性能A级）", "50厚", "适用部位"])
    ws.append(["1.20厚花岗岩面层\n2.30厚干硬性水泥砂浆结合层", None, None, "电梯厅\n公共走道"])
    ws.append([None, None, None, "地下室门厅"])                      # 部位写到下一行
    ws.append(["二、内墙"])                  # 不写「做法」：合并的「…做法」整行会被认成列组表头（改前就如此）
    ws.append(["■ 内墙9", "无机涂料内墙面1", "16.5厚", "适用部位"])
    ws.append(["1.无机涂料饰面 2.6厚预拌砂浆找平", None, None, "楼梯间"])
    ws.merge_cells("A2:C2")
    ws.merge_cells("A4:D4")
    ws.merge_cells("A6:C6")
    buf = io.BytesIO()
    wb.save(buf)
    fs = qs.segment_file(buf.getvalue(), "做法表.xlsx", "xlsx", "", 0)
    by = _by_practice(fs.segments)
    assert list(by) == [FLOOR9, "内墙9 无机涂料内墙面1"]
    assert [(s.text, s.location) for s in by[FLOOR9]] == [("20厚花岗岩面层", "Sheet!A2（第2行）"),
                                                       ("30厚干硬性水泥砂浆结合层", "Sheet!A2（第2行）")]
    assert {s.part for s in by[FLOOR9]} == {"电梯厅、公共走道、地下室门厅"}, "多行部位用顿号连起来，前面的层也回填"
    assert {s.part_location for s in by[FLOOR9]} == {"Sheet!D2（第2行·适用部位）"}
    assert [s.kind for s in fs.segments if s.row == 4] == [HEADING], "整行一格的章节标题结束上一个做法"
    assert {(s.part, s.category) for s in by["内墙9 无机涂料内墙面1"]} == {("楼梯间", "墙面")}


def test_single_cell_row_ends_titled_practice():
    # docx 整行一格（「配件」「二、内墙做法」）：不是构造层行，上一个做法到此结束。不结束的话「配件」会按
    # 无层号的续行并进坡道的最后一层
    fs = qs.segment_file(docx_bytes(docx_table(
        ("■ 坡道1", "石材坡道1", "50厚", "适用部位"),
        ("1)20厚毛面花岗石板，干水泥擦缝 2)30厚干硬性水泥砂浆结合层", "非机动车坡道"),
        ("配件",),
        ("■ 楼面8", "水泥砂浆楼面1", "30厚", "适用部位"),
        ("1)30厚1:2.5水泥砂浆找平层 2)钢筋混凝土楼板", "用于人防楼梯间"))), "a.docx", "docx", "", 0)
    by = _by_practice(fs.segments)
    assert [s.text for s in by["坡道1 石材坡道1"]] == ["20厚毛面花岗石板，干水泥擦缝", "30厚干硬性水泥砂浆结合层"]
    assert [(s.kind, s.text) for s in fs.segments if s.row == 3] == [(OTHER, "配件")]
    assert {(s.part, s.category) for s in by["楼面8 水泥砂浆楼面1"]} == {("用于人防楼梯间", "楼地面")}

def test_long_part_truncated_for_display_but_read_in_full():
    long_part = "用于非机动车库、库房、地下汽车库，除水泵房，风机房等无噪声的地下设备用房，管道及桥架增加与顶板同色的金属防锈漆两道，燃烧性能不低于A级，" \
                "面层颜色详精装修设计"
    assert len(long_part) > 60 and long_part.index("详精装修设计") > 60
    fs = qs.segment_file(docx_bytes(docx_table(
        ("■ 顶棚1", "防霉涂料顶棚", "8厚", "适用部位"),
        ("1)喷无机防霉涂料两道 2)满刮耐水腻子两遍", long_part))), "a.docx", "docx", "", 0)
    pr = _practice(fs.segments)
    assert {s.part for s in pr} == {long_part[:60]}, "部位截到 60 字"
    assert {s.part_full for s in pr} == {long_part}
    # 截掉的那一段里有精装修引用：各层照样带「做法暂定」，不因为截断悄悄丢掉（铁律 8）
    assert {(s.inline.code, s.inline.where) for s in pr} == {(rules.FITOUT_PROVISIONAL, "part")}

def test_text_mode_practice_name_allows_leading_symbol():
    segs = qs.segment_text("■ 楼面9 花岗岩楼面1（燃烧性能A级）\n1)20厚花岗岩面层 2)30厚干硬性水泥砂浆结合层\n"
                           "■ 地防1 地下室底板\n1)4厚SBS改性沥青防水卷材 2)钢筋混凝土底板\n", 0, "做法.txt")
    by = _by_practice(segs)
    assert list(by) == [FLOOR9, "地防1 地下室底板"]
    assert not [s for s in segs if "■" in s.text]


def _snapshot(fs):
    return [[s.kind, s.text, s.location, s.layer_no, s.practice, s.part, s.category,
             s.inline.code if s.inline else None] for s in fs.segments]


def _before():
    import json
    from pathlib import Path

    return json.loads((Path(__file__).parent / "fixtures" / "ws29k1_segments_before.json")
                      .read_text(encoding="utf-8"))


@pytest.mark.parametrize("name,content", [
    # 回归一：没有部位列头的无表头表（一行拆进几格），外加三种差一点像标题行的行
    ("near_title.docx", near_title_docx), ("r1_cad.xlsx", r1_cad_workbook),
    # 回归二：有表头的做法表
    ("r9.xlsx", r9_workbook),
    ("设计说明-案例.xlsx", lambda: (__import__("pathlib").Path(__file__).parent / "fixtures"
                                / "设计说明-案例.xlsx").read_bytes()),
])
def test_tables_without_title_rows_segment_exactly_as_before(name, content):
    ext = name.rsplit(".", 1)[1]
    assert _snapshot(qs.segment_file(content(), name, ext, "", 0)) == _before()[name]



# ============================================================ 漏切层看得见、类别拆开（WS29-k2，issue #913）

def _block_lines(segs):
    return [(s.text, s.block_practice) for s in segs if s.kind == OTHER and s.block_practice is not None]


def test_paint_layers_become_practice_and_derusting_is_a_block_line():
    fs = qs.segment_file(k2_docx(), "做法表.docx", "docx", "", 0)
    paint = [s for s in _practice(fs.segments) if s.practice == "漆1 外露铁件油漆"]
    # 改前四层都是 other：整个做法不在任何一页（E31，表1 第102行）
    assert [(s.layer_no, s.text) for s in paint] == [(1, "氟碳金属面漆（颜色另详）"), (2, "氟碳金属底漆"),
                                                      (3, "刷专用防锈漆")]
    assert {s.part for s in paint} == {"外露铁件"}
    assert _block_lines(fs.segments) == [(K2_DERUST, "漆1 外露铁件油漆")]
    derust = next(s for s in fs.segments if s.text == K2_DERUST)
    assert derust.location == "表1 第10行第1列" and derust.practice_key == paint[0].practice_key


def test_block_line_in_named_text_block():
    segs = qs.segment_text("外墙5：外露铁件油漆\n1. 清理基层，除锈\n2. 刷专用防锈漆\n3. 氟碳金属面漆\n", 0, "a.txt")
    assert [s.text for s in _practice(segs)] == ["刷专用防锈漆", "氟碳金属面漆"]
    assert _block_lines(segs) == [("清理基层，除锈", "外墙5 外露铁件油漆")]


def test_numbered_clauses_outside_named_blocks_are_not_block_lines():
    # 说明文字里一串「1、…2、…」也会开一个块。那里认不出的条款不是漏掉的层（建施说明实测 45 行规范名称）
    segs = qs.segment_text("1、本工程设计依据国家现行规范及当地规定执行，图中尺寸以毫米计。\n"
                           "2、屋面防水采用3厚SBS改性沥青防水卷材\n"
                           "3、门窗立樘位置除注明者外均居墙中，施工时以现场核对为准\n", 0, "说明.txt")
    assert len(_practice(segs)) == 1
    others = [s for s in segs if s.kind == OTHER]
    assert len(others) == 2, "夹具要产出块里的 other，否则本用例是空比空"
    assert _block_lines(segs) == []


def test_block_line_in_grid_practice_column():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["编号", "部位", "顶棚做法", None])
    ws.merge_cells("C1:D1")
    ws.append(["顶棚1", "车库", 1, "刷白色乳胶漆两遍"])
    ws.append([None, None, 2, "清理基层"])
    buf = io.BytesIO()
    wb.save(buf)
    fs = qs.segment_file(buf.getvalue(), "做法表.xlsx", "xlsx", "", 0)
    assert [s.text for s in _practice(fs.segments)] == ["刷白色乳胶漆两遍"]
    assert _block_lines(fs.segments) == [("清理基层", "顶棚1")]


def test_titled_practice_without_any_layer_leaves_a_trace():
    # 标题行后面紧跟下一个标题行、标题行在表尾：改前这两个做法整个消失，不出片段也没有提示（铁律 8）
    fs = qs.segment_file(docx_bytes(docx_table(
        ("■ 楼面1", "地砖楼面", "50厚", "适用部位"),
        ("■ 楼面2", "水泥砂浆楼面", "20厚", "适用部位"),
        ("1)20厚1:2水泥砂浆面层 2)钢筋混凝土楼板", "门厅"),
        ("■ 配件1", "车库配件", "适用部位"),
    )), "做法表.docx", "docx", "", 0)
    assert [s.practice for s in _practice(fs.segments)] == ["楼面2 水泥砂浆楼面"] * 2, "标题行仍不出做法片段"
    assert [(s.text, s.block_practice, s.location) for s in fs.segments if s.kind == OTHER] == [
        ("■ 楼面1 地砖楼面 50厚 适用部位", "楼面1 地砖楼面", "表1 第1行第1列"),
        ("■ 配件1 车库配件 适用部位", "配件1 车库配件", "表1 第4行第1列")]


def test_titled_practice_with_only_other_lines_gets_no_extra_trace():
    fs = qs.segment_file(docx_bytes(docx_table(
        ("■ 配件1", "车库配件", "适用部位"), ("1)车轮挡：见图集 2)减速带：见图集", "车库"))), "a.docx", "docx", "", 0)
    assert _block_lines(fs.segments) == [("车轮挡：见图集", "配件1 车库配件"), ("减速带：见图集", "配件1 车库配件")]


def test_step_ramp_and_apron_are_separate_categories():
    # 改前三者并成「台阶坡道」：坡道的每一层上下文都带「台阶」，一致性闸认不出坡道给了台阶的码（E30）
    segs = qs.segment_text("坡道1：石材坡道\n1. 20厚花岗石板\n台阶1：花岗石台阶\n1. 30厚花岗石板\n"
                           "散水1：混凝土散水\n1. 60厚C20细石混凝土面层\n", 0, "a.txt")
    assert [(s.practice, s.category) for s in _practice(segs)] == [
        ("坡道1 石材坡道", "坡道"), ("台阶1 花岗石台阶", "台阶"), ("散水1 混凝土散水", "散水")]
    by = {s.practice: s for s in _practice(qs.segment_file(k2_docx(), "a.docx", "docx", "", 0).segments)}
    assert "台阶" not in by["坡道2 细石混凝土坡道"].context()
    assert "坡道" not in by["其它2 花岗石台阶"].context()
