# 来源：engine-service/tests/test_quantity_render.py（zhitong-cost-ai@991ff5cd，W4 迁入 zt；改动见「zt 改动」注释）
"""工程量清单生成：xlsx 四页、编码依据声明、项目特征（WS11 QTY-06/08/10/14，issue #763）。"""
from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook

from zt import code_registry
from zt.quantity import features as feat
from zt.quantity import pipeline as ql
from zt.quantity import render as qr
from tests.quantity.fixtures.shapes import (R8_PRACTICES_TXT, R9_PICK, FakeAgent, docx_file, docx_p, docx_table,
                                            keyword_picker, r9_workbook, txt_file)
from tests.quantity.support import run


def _run(monkeypatch, fake, files, **kw):
    return run(files, agent=fake, **kw)


def _wb(out):
    return load_workbook(BytesIO(out["xlsx_bytes"]))


# ============================================================ QTY-08 四页 + 待核页

def test_workbook_has_four_sheets(monkeypatch):
    out = _run(monkeypatch, FakeAgent(R9_PICK), [(r9_workbook(), "做法表.xlsx", "xlsx")])
    assert _wb(out).sheetnames == ["分部分项工程量清单", "待核与未匹配", "覆盖核算", "说明"]


def test_unmatched_sheet_rows_match_response(monkeypatch):
    pick = keyword_picker([(("地砖",), "main", "011102003"), (("垫层",), "main", "999999999"),
                           (("细石混凝土面层",), "main", "011101001"), (("罩面",), "main", "011105001")])
    out = _run(monkeypatch, FakeAgent(pick), [(r9_workbook(), "做法表.xlsx", "xlsx")])
    assert out["unmatched"], "夹具必须能产生待核 / 未匹配（否则本用例是空比空）"
    ws = _wb(out)["待核与未匹配"]
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    assert len(rows) == len(out["unmatched"])
    for row, u in zip(rows, out["unmatched"]):
        assert row[2] == u["reason_code"] and row[4] == u["text"] and row[6] == u["location"]
    kinds = {u["reason_code"] for u in out["unmatched"]}
    assert {"code_not_in_registry", "semantic_mismatch"} <= kinds


def test_coverage_sheet_sums(monkeypatch):
    out = _run(monkeypatch, FakeAgent(R9_PICK), [(r9_workbook(), "做法表.xlsx", "xlsx")])
    ws = _wb(out)["覆盖核算"]
    n = out["coverage_detail"]["practice_segments"]
    assert f"做法片段共 {n} 个；去向合计 {n} 个" in ws.cell(row=1, column=1).value
    dests = [ws.cell(row=r, column=9).value for r in range(5, ws.max_row + 1)]
    assert len(dests) == n and None not in dests and "（缺去向）" not in dests


# ============================================================ WS29-f 新原因与标记在三页上看得到（issue #835）

def test_scope_reasons_and_marks_on_sheets(monkeypatch):
    note = "注：本表上注明详二次装修设计部位的材料做法仅为暂定，规格材质最终由装修设计确定。"
    doc = docx_file(docx_p("卫生间洁具由业主另行发包") + docx_table(
        ("1)预留防滑地砖装修面层（用户自理） 2)面层瓷砖用户自理或（详精装修设计） "
         "3)1.5厚聚合物水泥防水涂料（详精装修设计） 4)20厚1:3水泥砂浆找平层", "住宅卫生间"),
        ("1)装修吊顶（详二次装修图纸） 2)钢筋混凝土板底", "门厅"),
    ) + docx_p(note))
    pick = keyword_picker([(("防水涂料",), "main", "Z010904002"), (("找平",), "main", "Z011101006")])
    out = _run(monkeypatch, FakeAgent(pick), [doc])
    wb = _wb(out)
    # —— 待核与未匹配：新原因码、去向、引用原句、「做法暂定」
    pend = {r[4]: r for r in ([c.value for c in row] for row in wb["待核与未匹配"].iter_rows(min_row=2))}
    own = pend["预留防滑地砖装修面层（用户自理）"]
    assert own[1] == "范围排除" and own[2] == "excluded_owner_inline" and "预留防滑地砖装修面层（用户自理）" in own[3]
    amb = pend["面层瓷砖用户自理或（详精装修设计）"]
    assert amb[1] == "待核" and amb[2] == "owner_inline_ambiguous" and "请人工确认" in amb[3]
    ceil = pend["装修吊顶（详二次装修图纸）"]
    assert ceil[1] == "待核" and ceil[2] == "role_conflict" and ceil[3].endswith("【原文注明详精装修设计，做法暂定】")
    # —— 覆盖核算：每一层一行，带标记的层在说明列注明
    cov = {r[7]: r for r in ([c.value for c in row] for row in wb["覆盖核算"].iter_rows(min_row=5))}
    assert len(cov) == out["coverage_detail"]["practice_segments"] == 6
    assert cov["预留防滑地砖装修面层（用户自理）"][8] == "范围排除"
    assert cov["1.5厚聚合物水泥防水涂料（详精装修设计）"][8] == "正式清单项"
    assert cov["1.5厚聚合物水泥防水涂料（详精装修设计）"][10] == "【原文注明详精装修设计，做法暂定】"
    assert cov["20厚1:3水泥砂浆找平层 住宅卫生间"][10] in ("", None)
    # —— 清单页：来源依据列带「做法暂定」
    ws = wb["分部分项工程量清单"]
    m = {ws.cell(row=r, column=2).value: ws.cell(row=r, column=13).value for r in range(4, ws.max_row + 1)}
    assert "做法暂定" in m["Z010904002001"] and "做法暂定" not in m["Z011101006001"]
    # —— 说明页：做法暂定的层数、排除句清单、范围相关原文
    notes = {wb["说明"].cell(row=r, column=1).value: wb["说明"].cell(row=r, column=2).value
             for r in range(1, wb["说明"].max_row + 1)}
    assert notes["做法暂定的层"].startswith("2 个做法层")
    assert notes["排除句清单"].startswith("全局排除句 1 条") and "排掉 1 层" in notes["排除句清单"]
    assert "原句：卫生间洁具由业主另行发包" in notes["排除句 e1"] and "排掉 0 个片段" in notes["排除句 e1"]
    assert notes["范围相关原文"].startswith("1 句")
    assert note in notes["范围相关原文 1"] and "没有据此排除" in notes["范围相关原文 1"]
    assert any(v and "已按暂定做法列项" in v and "仅为暂定" in v for k, v in notes.items() if k.startswith("告警"))


def test_notes_sheet_lists_block_lines_and_states_reproducibility():
    """WS29-k2：「说明」页的「块内未识别行」清单逐行列出原句、位置、所在做法；整个做法一层都没认出的单独点名。"""
    cov = {"practice_segments": 0, "by_destination": {}, "block_lines_unrecognized": [
        {"text": "清理基层，除锈", "file_name": "a.docx", "location": "表1 第10行第1列", "practice": "漆1 外露铁件油漆",
         "practice_has_layers": True},
        {"text": "车轮挡：见图集", "file_name": "a.docx", "location": "表3 第74行第1列", "practice": "配件1 车库配件",
         "practice_has_layers": False}]}
    # zt 改动：render_xlsx 的输入是 finalize 结果
    wb = load_workbook(BytesIO(qr.render_xlsx({"items": [], "coverage_detail": cov,
                                               "meta": {"prompt_version": ql.QUANTITY_PROMPT_VERSION}})))
    col = {c[0].value: c[1].value for c in wb["说明"].iter_rows()}
    assert col["块内未识别行"].startswith("2 行") and col["块内未识别行"].endswith("一层都没有认出：配件1 车库配件")
    assert col["块内未识别行 1"] == "a.docx 表1 第10行第1列 | 做法：漆1 外露铁件油漆 | 原文：清理基层，除锈"
    assert col["块内未识别行 2"] == "a.docx 表3 第74行第1列 | 做法：配件1 车库配件（这个做法一层都没有认出） | 原文：车轮挡：见图集"
    assert "片段编号" in col["复现性"] and "同一份 answers 重跑 finalize 结果逐字相同" in col["复现性"]


# ============================================================ QTY-10 编码依据

def test_sheet_declares_code_basis(monkeypatch):
    out = _run(monkeypatch, FakeAgent(R9_PICK), [(r9_workbook(), "做法表.xlsx", "xlsx")])
    ws = _wb(out)["分部分项工程量清单"]
    assert ws.cell(row=1, column=1).value == "分部分项工程量清单（框架 · AI 初稿 · 未含工程量）"
    basis = ws.cell(row=2, column=1).value
    assert "zj2025" in basis and "490 条" in basis and out["registry_fingerprint"][:12] in basis
    assert "2013" in basis and "同号异义" in basis
    assert ws.cell(row=3, column=2).value == "项目编码" and ws.cell(row=3, column=6).value == "工程量"
    notes = {ws2.cell(row=r, column=1).value: ws2.cell(row=r, column=2).value
             for ws2 in [_wb(out)["说明"]] for r in range(1, ws2.max_row + 1)}
    assert notes["提示词版本"] == ql.QUANTITY_PROMPT_VERSION and notes["码表版本"] == "zj2025"
    # zt 改动：不再记模型名；记片段指纹与作答覆盖
    assert notes["片段指纹"] == out["prepare_fingerprint"]
    assert notes["分类回答"].startswith(f"{out['prepare']['segments_to_classify']} / ")


def test_cross_version_collision_warning(monkeypatch):
    fake = FakeAgent(R9_PICK)
    out = _run(monkeypatch, fake, [(r9_workbook(), "做法表.xlsx", "xlsx")])
    assert not any(w["code"] == "quantity_cross_version" for w in out["warning_items"]), "只有一个版本时不提示"
    assert all(not i["notes"] for i in out["items"])
    reg = dict(code_registry.load_registry("zj2025"))
    fake2013 = {"011102003": {**reg["011102003"], "name": "块料楼地面（2013 口径）"},
                "011101002": dict(reg["011101002"])}          # 同号同名：不提示
    monkeypatch.setattr(ql, "other_registries", lambda v: {"gb2013": fake2013})
    out = run([(r9_workbook(), "做法表.xlsx", "xlsx")], agent=FakeAgent(R9_PICK))
    tile = next(i for i in out["items"] if i["base_code"] == "011102003")
    assert tile["notes"] and "gb2013" in tile["notes"][0] and "同号异义" in tile["notes"][0]
    assert "同号异义" in tile["source"], "M 列跟着文件走"
    floor = next(i for i in out["items"] if i["base_code"] == "011101002")
    assert not floor["notes"]
    assert any(w["code"] == "quantity_cross_version" for w in out["warning_items"])


# ============================================================ QTY-06 / QTY-14 项目特征

def test_features_fill_verbatim_only(monkeypatch):
    def features(item):
        vals = {}
        for n, label in item["槽位"].items():
            if "面层材料品种" in label:
                vals[n] = "防滑地砖"                 # 原文子串 → 通过
            elif "结合层种类" in label:
                vals[n] = "1:3水泥砂浆（干硬性）"      # 改写过 → 不通过
        return vals

    txt = "楼面1：防滑地砖楼面（用于卫生间）\n1. 8厚防滑地砖面层\n2. 20厚1:3干硬性水泥砂浆结合层\n"
    out = _run(monkeypatch, FakeAgent(keyword_picker([(("地砖",), "main", "011102003")]), features=features),
               [txt_file(txt)])
    d = out["items"][0]["features"]
    assert "面层材料品种及规格：〔AI摘录〕防滑地砖" in d
    assert "结合层种类及强度等级：待补（AI 摘录未通过原文校验）" in d
    assert "勾缝材料种类：待补" in d


def test_d_column_lists_all_layers(monkeypatch):
    fake = FakeAgent(keyword_picker([(("砖",), "main", "011102003")]))
    out = _run(monkeypatch, fake, [txt_file(R8_PRACTICES_TXT, "r8.txt")])
    items = [i for i in out["items"] if i["base_code"] == "011102003"]
    assert len(items) == 4
    ds = [i["features"] for i in items]
    assert len(set(ds)) == 4, "r8 的 4 条 011102003 在 D 列必须能区分"
    for d, part in zip(ds, ["卫生间", "大厅", "走廊", "茶室"]):
        assert f"部位：{part}" in d
        assert "构造做法（原文，自上而下）：①" in d and "②" in d
    template = code_registry.load_registry("zj2025")["011102003"]["features_text"]
    assert all(d != template for d in ds)


def test_deterministic_slots(monkeypatch):
    out = _run(monkeypatch, FakeAgent(keyword_picker([(("地砖",), "main", "011102003")])),
               [txt_file("楼面1：防滑地砖楼面（用于卫生间）\n1. 8厚防滑地砖面层\n")])
    d = out["items"][0]["features"]
    assert d.startswith("1.建筑做法名称：楼面1 防滑地砖楼面（用于卫生间）\n")
    slots = feat.parse_slots("1.土类别：详见地勘报告\n2.挖方部位：")
    vals = {s.no: feat.deterministic_value(s.label, "某做法", "基坑") for s in slots}
    assert vals == {1: None, 2: "基坑"}
    assert slots[0].preset == "详见地勘报告"


def test_slot_numbering_consecutive_for_all_registry_codes():
    reg = code_registry.load_registry("zj2025")
    assert len(reg) == 490
    for code, hit in reg.items():
        slots = feat.parse_slots(hit["features_text"])
        assert [s.no for s in slots] == list(range(1, len(slots) + 1)), code
        text = feat.render(slots, {}, "", [])
        nums = [int(line.split(".", 1)[0]) for line in text.split("\n") if line[:1].isdigit()]
        assert nums == list(range(1, len(slots) + 1)), code
    assert len(feat.parse_slots(reg["011102003"]["features_text"])) == 9


@pytest.mark.parametrize("raw,labels", [
    # 从 1 重新起编、跳号、一行并列三个槽 —— 渲染一律从 1 连续
    ("1.建筑做法名称：\n1.防水膜品种\n3.涂膜厚度及遍数", ["建筑做法名称", "防水膜品种", "涂膜厚度及遍数"]),
    ("1.结构形式 、跨度 2.杆件钢材品种及规格 3.节点形式 、连接方式", ["结构形式 、跨度", "杆件钢材品种及规格", "节点形式 、连接方式"]),
    ("1.构件名称\n3.井、池净空尺寸\n3.垫层材料种类、厚度", ["构件名称", "井、池净空尺寸", "垫层材料种类、厚度"]),
])
def test_slot_renumbering_on_broken_templates(raw, labels):
    slots = feat.parse_slots(raw)
    assert [s.label for s in slots] == labels, "标签文字逐字保留"
    rendered = feat.render(slots, {}, "", [])
    assert [line.split("：")[0] for line in rendered.split("\n")[:len(labels)]] == \
        [f"{i}.{lbl}" for i, lbl in enumerate(labels, start=1)]


def test_verify_verbatim():
    assert feat.verify_verbatim("1:3水泥砂浆", ["20厚1:3水泥砂浆找平层"])
    assert feat.verify_verbatim("1:3 水泥砂浆", ["20厚1:3水泥砂浆找平层"]), "空白差异不算改写"
    assert not feat.verify_verbatim("M10水泥砂浆", ["20厚1:3水泥砂浆找平层"])
    assert not feat.verify_verbatim("", ["x"])


# ============================================================ 旧用例保留：公式注入

def test_render_xlsx_escapes_formula_injection_in_source_column():
    items = [{"code": "010103001001", "base_code": "010103001", "name": "平整场地", "unit": "m²",
              "features": "x", "section": "土石方工程", "source": "=1+1"}]
    ws = load_workbook(BytesIO(qr.render_xlsx({"items": items}))).active
    assert ws.cell(row=qr.HEADER_ROW + 2, column=13).value == "'=1+1"
    assert ws.cell(row=qr.HEADER_ROW + 2, column=6).value is None, "F 列（工程量）必须留空"
