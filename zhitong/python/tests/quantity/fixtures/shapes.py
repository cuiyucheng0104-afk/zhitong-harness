# 来源：engine-service/tests/fixtures/quantity_shapes.py（zhitong-cost-ai@991ff5cd，W4 搬入 zt；改动见「zt 改动」注释）
"""工程量清单生成测试的版式夹具（WS11，issue #763）。不是用例文件（不以 test_ 开头）。

生产 QA 的输入（r1 CAD 导出、r8、r9 做法表、建筑说明）是客户资料，不入仓（C19 / D7）。
这里按它们的**版式**手工造出等价形状，内容全部是通用做法描述，没有客户信息：

- r9 形：房间做法表 —— 编号 / 部位列纵向合并，每个做法列组 = 「层号列 + 层正文列」横向合并表头；
  有续行（无层号的第二行说明）、有隐藏行、有隐藏工作表。
- r1 形：CAD 导出 —— 没有表头，一行一格，名称行后跟编号层，一格里塞多层。
- r8 形：同一编码（块料楼地面）四个做法，规格各不相同。
"""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook

R9_ROOMS = [
    ("地面1", "地下汽车库", [
        (1, "20厚C25细石混凝土面层，随打随抹光"),
        (2, "素水泥浆一道（内掺建筑胶）"),
        (3, "150厚C20混凝土垫层"),
        (4, "300厚级配碎石垫层"),
        (5, "素土夯实"),
    ]),
    ("楼面2", "卫生间", [
        (1, "10厚防滑地砖面层"),
        (2, "20厚1:3干硬性水泥砂浆结合层"),
        (3, "1.5厚聚氨酯防水涂料，四周上翻300高"),
        (4, "20厚1:3水泥砂浆找平层"),
        (5, "钢筋混凝土楼板"),
    ]),
]


def r9_workbook(extra_rooms: int = 0, *, hidden_row: bool = True, hidden_sheet: bool = True) -> bytes:
    """房间做法表。表头第 2 行，「楼地面做法」「踢脚做法」「顶棚做法」三个列组各占两列（层号 + 正文）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "做法表"
    ws["A1"] = "室内装修做法表"
    ws.merge_cells("A1:H1")
    ws["A2"], ws["B2"] = "编号", "部位"
    ws["C2"] = "楼地面做法"
    ws.merge_cells("C2:D2")
    ws["E2"] = "踢脚做法"
    ws.merge_cells("E2:F2")
    ws["G2"] = "顶棚做法"
    ws.merge_cells("G2:H2")
    ws["I2"] = "备注"
    row = 3
    rooms = list(R9_ROOMS)
    for k in range(extra_rooms):
        rooms.append((f"楼面{k + 10}", f"办公室{k + 1}", [
            (1, f"{8 + k % 3}厚地砖面层"),
            (2, "20厚1:3干硬性水泥砂浆结合层"),
            (3, "素水泥浆一道"),
            (4, f"{30 + k}厚C20细石混凝土找平层"),
            (5, "钢筋混凝土楼板"),
            (6, "界面剂一道"),
        ]))
    for code, part, layers in rooms:
        start = row
        for no, text in layers:
            ws.cell(row=row, column=3, value=no)
            ws.cell(row=row, column=4, value=text)
            row += 1
        ws.cell(row=start, column=1, value=code)
        ws.cell(row=start, column=2, value=part)
        if row - 1 > start:
            ws.merge_cells(start_row=start, start_column=1, end_row=row - 1, end_column=1)
            ws.merge_cells(start_row=start, start_column=2, end_row=row - 1, end_column=2)
        if code == "地面1":
            # 踢脚：两层，第二层是无层号的续行
            ws.cell(row=start, column=5, value=1)
            ws.cell(row=start, column=6, value="10厚1:2水泥砂浆罩面压实赶光")
            ws.cell(row=start + 1, column=5, value=2)
            ws.cell(row=start + 1, column=6, value="12厚1:3水泥砂浆打底扫毛")
            ws.cell(row=start + 2, column=6, value="（踢脚高 150）")
            ws.cell(row=start, column=7, value=1)
            ws.cell(row=start, column=8, value="刷白色乳胶漆两遍")
            ws.cell(row=start + 1, column=7, value=2)
            ws.cell(row=start + 1, column=8, value="满刮腻子两遍")
            ws.cell(row=start, column=9, value="顶棚乳胶漆颜色见室内设计")
    if hidden_row:
        ws.cell(row=row, column=4, value="30厚隐藏行不该出现的做法")
        ws.row_dimensions[row].hidden = True
        row += 1
    if hidden_sheet:
        hs = wb.create_sheet("作废")
        hs["A1"] = "20厚作废做法水泥砂浆"
        hs.sheet_state = "hidden"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def r1_cad_workbook() -> bytes:
    """CAD 导出：没有表头，一行一格；名称行 + 编号层；一格里塞多层；有相邻格拆开的一行。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    lines = [
        "工程做法",
        "屋面1：不上人屋面（用于主楼屋面）",
        "1.40厚C20细石混凝土保护层 2.10厚低标号砂浆隔离层",
        "3.3厚SBS改性沥青防水卷材",
        "4.20厚1:3水泥砂浆找平层",
        "5.最薄处30厚LC5.0轻集料混凝土找坡层",
        "6.60厚挤塑聚苯板保温层",
        "7.钢筋混凝土屋面板",
        "内墙1：乳胶漆墙面",
        "1.刷乳胶漆两遍",
        "2.满刮腻子两遍",
        "3.10厚1:1:6混合砂浆打底",
    ]
    for i, t in enumerate(lines, start=1):
        ws.cell(row=i, column=1, value=t)
    # 被拆进相邻两格的一行
    ws.cell(row=len(lines) + 1, column=1, value="顶棚1：")
    ws.cell(row=len(lines) + 1, column=2, value="纸面石膏板吊顶")
    ws.cell(row=len(lines) + 2, column=1, value="1.9.5厚纸面石膏板")
    ws.cell(row=len(lines) + 3, column=1, value="2.U型轻钢龙骨中距600")
    ws.cell(row=len(lines) + 4, column=1, value="3.φ8吊杆")
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


R8_PRACTICES_TXT = (
    "楼面1：防滑地砖楼面（用于卫生间）\n"
    "1. 8厚防滑地砖面层\n2. 20厚1:3干硬性水泥砂浆结合层\n"
    "楼面2：玻化砖楼面（用于大厅）\n"
    "1. 10厚800×800玻化砖面层\n2. 30厚1:3干硬性水泥砂浆结合层\n"
    "楼面3：地砖楼面（用于走廊）\n"
    "1. 10厚600×600地砖面层\n2. 20厚1:3干硬性水泥砂浆结合层\n"
    "楼面4：仿古砖楼面（用于茶室）\n"
    "1. 12厚仿古地砖面层\n2. 25厚1:3干硬性水泥砂浆结合层\n"
)


# ============================================================ 假 agent（按 prepare 的片段作答）
# zt 改动：旧夹具是打桩 llm_client.chat 的 FakeLLM（按 prompt 里的片段 JSON 作答）；zt 里模型就是 agent，
# 测试改成按 prepare 给出的批次逐片段作答，再交给 finalize。pick / drop / features 的含义不变。


class FakeAgent:
    """可配置的假 agent：pick(片段 dict) → (role, code[, ex])。

    - drop：永远不回答的片段原文关键字（模拟漏答）；
    - features：特征摘录的回答函数 task_dict → {槽号: 值}；缺省不作答（全部待补）；
    - batches：每次作答时看到的批次（prepare 结果里的 batches）。

    pick 返回的 ex 写批内序号（"e1" = 本批列出的第一条排除句），作答时换成全局编号，与旧夹具口径一致；
    批内没有这个序号时原样交出（finalize 判为不存在的编号）。
    """

    def __init__(self, pick, *, drop=(), features=None):
        self.pick = pick
        self.drop = tuple(drop)
        self.features = features
        self.batches: list[dict] = []
        self.feature_calls: list[list[dict]] = []

    def answer(self, prep: dict) -> dict:
        rows = []
        for b in prep["batches"]:
            self.batches.append(b)
            local = {f"e{i}": e["id"] for i, e in enumerate(b["exclusions"], start=1)}
            for s in b["segments"]:
                if any(d in s.get("原文", "") for d in self.drop):
                    continue
                got = self.pick(s)
                role, code = got[0], got[1]
                ex = got[2] if len(got) > 2 else ""
                rows.append({"id": s["id"], "role": role, "c": code, "ex": local.get(ex, ex), "why": ""})
        return {"r": rows}

    def answer_features(self, tasks: list[dict]) -> list[dict]:
        self.feature_calls.append(tasks)
        return [{"k": t["k"], "v": self.features(t) if self.features else {}} for t in tasks]

    @property
    def segments_seen(self) -> list[dict]:
        return [s for b in self.batches for s in b["segments"]]


def keyword_picker(table, default=("aux", "")):
    """[(关键字元组, role, code[, ex])]：原文或做法名包含全部关键字的第一条生效。"""
    def pick(s):
        hay = s.get("原文", "") + " " + s.get("做法", "") + " " + s.get("类别", "")
        for row in table:
            keys = row[0]
            if all(k in hay for k in keys):
                return tuple(row[1:])
        return default
    return pick


R9_PICK = keyword_picker([
    (("地砖",), "main", "011102003"),
    (("细石混凝土面层",), "main", "011101002"),
    (("垫层",), "main", "010501002"),
    (("防水涂料",), "main", "Z010904002"),
    (("找平",), "main", "Z011101006"),
    (("罩面",), "main", "011105001"),
    (("乳胶漆",), "main", "011404002"),
])


def txt_file(text: str, name: str = "做法说明.txt"):
    return (text.encode("utf-8"), name, "txt")


# ============================================================ 无表头的做法表（WS29-f，issue #835）
# 生产实测的版式：docx 表格没有「做法」表头，每个做法两行 —— 标题行 4 格「■ 楼面9｜名称｜厚度｜适用部位」，
# 构造层行 2 格「1)… 2)… 3)…｜适用部位的文字」。引擎把一行各格拼成一句按文本模式处理。
# 句式取自通用构造做法用语，不含项目标识。

import io as _io  # noqa: E402
import zipfile as _zipfile  # noqa: E402

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def docx_p(text: str) -> str:
    return f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"


def docx_table(*rows: tuple) -> str:
    return "<w:tbl>" + "".join(
        "<w:tr>" + "".join(f"<w:tc>{docx_p(c)}</w:tc>" for c in cells) + "</w:tr>" for cells in rows) + "</w:tbl>"


def docx_bytes(body: str) -> bytes:
    buf = _io.BytesIO()
    with _zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml", '<?xml version="1.0"?>'
                   '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>')
        z.writestr("word/document.xml", f'<w:document xmlns:w="{_W}"><w:body>{body}</w:body></w:document>')
    return buf.getvalue()


def docx_file(body: str, name: str = "做法表.docx"):
    return (docx_bytes(body), name, "docx")


# 13 处层内自理里的 8 种句式（其余是重复出现）
OWNER_INLINE_LAYERS = [
    "预留防滑地砖装修面层（用户自理）",
    "预留50厚装修面层（用户自理）",
    "20厚预留装修面层(用户自理)",
    "30厚C25细石混凝土找平找坡（用户自理）",
    "轻质混凝土回填并找坡（同层排水区域）（用户自理）",
    "面层用户自理",
    "40厚C25细石混凝土保护层（内配钢丝网，置于中上部）(用户自理)",
    "最薄处40厚C25细石混凝土找平找坡，坡向地漏；（用户自理）",
]

# 读着带「二次装修」、但都不是排除句的 7 句：三句否定、无谓词两句、引用一句、表注一句
NOT_EXCLUSION_SENTENCES = [
    "楼梯间、楼梯间前室、地上变配电室、消控中心、计量间、门卫等不需二次装修的非配水房间",
    "地下室车库、自行车库、地下室楼梯间、楼梯间前室等不需要二次装修的墙面",
    "商业网点、物业用房、居家养老用房等无二次装修部分",
    "住宅二次装修时，不得破坏承重结构的墙体，钢筋混凝土梁、柱、楼板，不得改变设备系统及建筑外形。",
    "景观及二次装修的材料及做法均应符合现行的防火规范要求。",
    "厨房设备见二次装修。",
    "注：本表上注明详二次装修设计部位的材料做法仅为暂定，规格材质最终由装修设计确定。",
]


# ============================================================ 标题行 + 构造层行（WS29-k1，issue #891）
# E27 的版式：4 个做法、8 行。标题行「■ 编号｜名称｜总厚度｜适用部位」（总厚度可能空着、也可能整格没有），
# 构造层行「1)… 2)…｜适用部位的文字」。第 2 个做法的部位是「…等不需二次装修的非配水房间」（与 WS29-f 的
# 规则一起验），第 3 个做法的标题行没有厚度（改前它按续行并进上一个做法的最后一层），第 4 个做法的编号头
# 「其它」不是部位词、标题行只有 3 格。句式取自通用构造做法用语，不含项目标识。

TITLED_PRACTICES = [
    (("■ 楼面9", "花岗岩楼面1（燃烧性能A级）", "50厚", "适用部位"),
     ("1)20厚花岗岩面层，干水泥擦缝 2)30厚干硬性水泥砂浆结合层 3)素水泥浆一道（内掺建筑胶） 4)钢筋混凝土楼板",
      "电梯厅、公共走道")),
    (("■ 内墙9", "无机涂料内墙面1（燃烧性能A级）", "16.5厚", "适用部位"),
     ("1)无机涂料（一底两面）饰面 2)满刮2厚柔性耐水腻子两遍分层找平 3)6厚预拌砂浆找平 4)基层墙体",
      "楼梯间、楼梯间前室、消控中心、计量间、门卫等不需二次装修的非配水房间")),
    (("■ 顶棚2", "石膏腻子顶棚2（燃烧性能A级）", "", "适用部位"),
     ("1)满刮专业石膏腻子批白两遍，四周下翻100mm 2)板底刷素界面剂一道 3)钢筋混凝土楼板清理",
      "用于住宅套内除厨、卫外的房间")),
    (("■ 其它2", "花岗石台阶", "适用部位"),
     ("1)30厚花岗石板铺面，灌水泥浆擦缝 2)30厚1:4干硬性水泥砂浆结合层 3)素水泥浆一道 4)100厚C15混凝土垫层",
      "一层入口台阶")),
]


def titled_table_rows() -> list[tuple]:
    return [row for title, body in TITLED_PRACTICES for row in (title, body)]


def titled_docx() -> bytes:
    return docx_bytes(docx_table(*titled_table_rows()))


# 回归夹具：没有部位列头的无表头表（CAD 导出那种一行拆进几格），外加三种「差一点像标题行」的行 ——
# 「适用部位：…」带着内容（不是列头）、首格不是做法编号的列头行「编号｜名称｜部位」、标题行没有「适用部位」格。
# 本片之后它们逐片段与改前相同（快照取自改前的 origin/dev）。都不带前导符号：带「■」的名称行按本片的约定
# 在文本模式里认成做法名称，那是有意的变化，另有用例。
NEAR_TITLE_ROWS = [
    ("工程做法",),
    ("楼面1：", "防滑地砖楼面"),
    ("1.10厚防滑地砖面层 2.20厚1:3干硬性水泥砂浆结合层", "适用部位：卫生间"),
    ("编号", "名称", "部位"),
    ("楼面2", "块料楼面", "用于门厅"),
    ("1.10厚600×600地砖面层", "2.30厚1:3干硬性水泥砂浆结合层"),
    ("内墙1", "乳胶漆墙面", "20厚"),
    ("1.刷乳胶漆两遍", "2.满刮腻子两遍", "3.10厚1:1:6混合砂浆打底"),
]


def near_title_docx() -> bytes:
    return docx_bytes(docx_table(*NEAR_TITLE_ROWS))


# ============================================================ 错码闸、漏切层（WS29-k2，issue #913）
# E30 / E31 的版式（同上：标题行 + 构造层行）：坡道、台阶、楼梯踏步、防水砂浆、外露铁件油漆各一个做法。
# 外露铁件四层：三层漆（改前一个材料词都不中，整组落成 other、哪一页都查不到）和一层「清理基层，除锈…」
#（仍不是做法层，由「块内未识别行」交代）。句式取自通用构造做法用语，不含项目标识。

K2_PRACTICES = [
    (("■ 坡道2", "细石混凝土坡道", "90厚", "适用部位"),
     ("1)90厚C20细石混凝土随捣随抹，内配钢筋网片 2)钢筋混凝土现浇板", "地下汽车坡道")),
    (("■ 其它2", "花岗石台阶", "", "适用部位"),
     ("1)30厚花岗石板铺面，灌水泥浆擦缝 2)素水泥浆一道 3)100厚C15混凝土垫层", "一层入口台阶")),
    (("■ 楼面12", "水泥砂浆楼梯面", "20厚", "适用部位"),
     ("1)20厚1:2水泥砂浆面层，踏步做防滑条 2)素水泥浆一道 3)钢筋混凝土楼梯板", "楼梯踏步、休息平台")),
    (("■ 地防3", "砂浆防水层", "20厚", "适用部位"),
     ("1)20厚1：2水泥防水砂浆 2)钢筋混凝土底板", "地下室集水坑")),
    (("■ 漆1", "外露铁件油漆", "", "适用部位"),
     ("1)氟碳金属面漆（颜色另详） 2)氟碳金属底漆 3)刷专用防锈漆 4)清理基层，除锈等级不应低于Sa2.5或St3级",
      "外露铁件")),
]
K2_DERUST = "清理基层，除锈等级不应低于Sa2.5或St3级"

K2_PICK = keyword_picker([
    (("细石混凝土随捣",), "main", "011107001"),     # 坡道给了台阶的码（生产 6 条）
    (("现浇板",), "base", ""),
    (("花岗石板",), "main", "011107002"),
    (("垫层",), "main", "010501002"),
    (("水泥砂浆面层",), "main", "011101001"),       # 楼梯踏步给了楼地面整体面层的码
    (("楼梯板",), "base", ""),
    (("防水砂浆",), "main", "010905002"),           # 防水砂浆给了基础涂膜防水的码
    (("底板",), "base", ""),
    (("面漆",), "main", "011402001"),
    (("底漆",), "aux", ""),
    (("防锈漆",), "aux", ""),
    (("素水泥浆",), "aux", ""),
])


def k2_docx() -> bytes:
    return docx_bytes(docx_table(*[row for title, body in K2_PRACTICES for row in (title, body)]))


# 搭在前面的另一份文件：自己有做法层，还有一条全局排除句，对象词「卫生间洁具」在 k2_docx 里不出现。
# 两个条件缺一个，「提示词只取决于本批」的两条分支（批内编号、排除句按批筛选）就有一条够不到
OTHER_FILE_TXT = "卫生间洁具由业主另行发包\n地面1：地砖地面\n1. 10厚防滑地砖面层\n2. 20厚1:3水泥砂浆找平层\n"
