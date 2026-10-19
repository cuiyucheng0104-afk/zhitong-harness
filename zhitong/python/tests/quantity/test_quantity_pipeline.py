# 来源：engine-service/tests/test_quantity_pipeline.py（zhitong-cost-ai@991ff5cd，W4 迁入 zt；改动见「zt 改动」注释）
"""工程量清单生成流水线（WS11，issue #763）：QTY-01/02/03/04/05/07/09/11/12/13/15。

zt 改动：旧用例只打桩 ``quantity_list.chat``（假模型按 prompt 里的片段 JSON 作答）；zt 里分类由 agent 作答，
用例改成 prepare → 假 agent（:class:`tests.quantity.fixtures.shapes.FakeAgent`）按批次作答 → finalize → render，
其余全部真跑：切片、装批、闸、聚合、覆盖核算、xlsx。只验 LLM 调用机制的用例（max_tokens、截断二分、解析重试、
进程内缓存、token 用量）在 zt 里没有对应物，改写成对应的确定性契约或标 skip 写明原因。
"""
from __future__ import annotations

import hashlib
import json
from io import BytesIO

import pytest
from openpyxl import load_workbook

from zt.quantity import pipeline as ql
from zt.quantity import prompt as qp
from zt.quantity import segment as segmod
from zt.quantity.segment import PRACTICE
from zt.textextract.model import DocText, failed_doc
from tests.quantity.fixtures.shapes import (
    K2_DERUST,
    K2_PICK,
    NOT_EXCLUSION_SENTENCES,
    OTHER_FILE_TXT,
    OWNER_INLINE_LAYERS,
    R9_PICK,
    TITLED_PRACTICES,
    FakeAgent,
    docx_file,
    docx_p,
    docx_table,
    k2_docx,
    keyword_picker,
    r9_workbook,
    titled_docx,
    txt_file,
)
from tests.quantity.support import run

NO_LLM = "zt 不调用模型：分类由 agent 作答，旧用例验的是 LLM 调用机制（{}），没有对应物"


def _xlsx(out) -> bytes:
    return out["xlsx_bytes"]


def _run(monkeypatch, fake, files, **kw):
    return run(files, agent=fake, **kw)


FLOOR_TXT = (
    "地面1：地砖地面（用于门厅）\n"
    "1. 10厚防滑地砖面层\n"
    "2. 20厚1:3干硬性水泥砂浆结合层\n"
    "3. 20厚1:3水泥砂浆找平层\n"
)
FLOOR_PICK = keyword_picker([
    (("地砖",), "main", "011102003"),
    (("找平",), "main", "Z011101006"),
])


# ============================================================ QTY-01 装批与片段数闸

@pytest.mark.skip(reason=NO_LLM.format("每次调用显式 max_tokens"))
def test_every_llm_call_passes_max_tokens():
    pass


def test_large_table_split_into_batches_of_at_most_40(monkeypatch):
    # zt 改动：旧用例验截断二分；现在验 prepare 按 ≤40 片段 / ≤6000 字装批，全部片段都有去向
    fake = FakeAgent(R9_PICK)
    out = _run(monkeypatch, fake, [(r9_workbook(extra_rooms=13), "做法表.xlsx", "xlsx")])
    sizes = [len(b["segments"]) for b in fake.batches]
    assert len(sizes) > 1 and max(sizes) <= 40 and sum(sizes) == out["prepare"]["segments_to_classify"]
    assert out["status"] == "ok", out["warnings"]
    assert out["coverage_detail"]["by_destination"]["llm_failed"] == 0
    assert out["items"]


@pytest.mark.skip(reason=NO_LLM.format("单片段仍截断时标 llm_truncated"))
def test_single_segment_still_truncated_marks_llm_failed():
    pass


def test_missing_ids_marked_llm_missing_and_partial(monkeypatch):
    # zt 改动：旧版漏答补问一次；现在 agent 的 answers 里缺的片段直接记「模型未完成」，结果 partial
    fake = FakeAgent(FLOOR_PICK, drop=("找平",))
    out = _run(monkeypatch, fake, [txt_file(FLOOR_TXT)])
    failed = [u for u in out["unmatched"] if u["kind"] == "llm_failed"]
    assert [u["reason_code"] for u in failed] == ["llm_missing"]
    assert "找平" in failed[0]["text"]
    assert out["status"] == "partial"
    assert any(w["code"] == "quantity_llm_failed" and w["level"] == "warn" for w in out["warning_items"])


@pytest.mark.skip(reason=NO_LLM.format("回答不是 JSON 时同批重试"))
def test_unparseable_answer_retried_once():
    pass


def test_input_gate_counts_segments(monkeypatch):
    many = "".join(f"地面{i}：水泥地面\n1. {10 + i % 5}厚1:2水泥砂浆面层\n" for i in range(900))
    fake = FakeAgent(keyword_picker([(("水泥砂浆面层",), "main", "011101001")]))
    with pytest.raises(ql.InputTooLarge) as ei:
        run([txt_file(many, "大说明.txt")], agent=fake)
    assert ei.value.reason_code == "too_many_segments"
    assert "900" in ei.value.detail and "拆分" in ei.value.detail and "「大说明.txt」900 个" in ei.value.detail
    assert fake.batches == [], "超限必须在交给 agent 之前拦下"
    # 反向对照：400 个片段照常通过
    few = "".join(f"地面{i}：水泥地面\n1. {10 + i % 5}厚1:2水泥砂浆面层\n" for i in range(400))
    out = run([txt_file(few, "小说明.txt")], agent=fake)
    assert out["coverage_detail"]["practice_segments"] == 400


def test_no_answers_at_all_is_answers_missing(monkeypatch):
    # zt 改动：旧版全部批次失败 → 502（带 usage）；现在一个回答都没有 → AnswersMissing（CLI 退出码 2）
    fake = FakeAgent(FLOOR_PICK, drop=("厚",))
    with pytest.raises(ql.AnswersMissing) as ei:
        run([txt_file(FLOOR_TXT)], agent=fake)
    assert "3 个待分类片段都没有回答" in ei.value.detail
    assert ei.value.coverage["by_destination"]["llm_failed"] == 3


# ============================================================ QTY-02 码表注入与语义闸

def test_prompt_contains_full_registry(monkeypatch):
    # zt 改动：码表进 prepare 写出的「清单分类说明.md」（旧版进 system prompt）
    fake = FakeAgent(FLOOR_PICK)
    out = _run(monkeypatch, fake, [txt_file(FLOOR_TXT)])
    system = ql.instructions_text(out["prepare"])
    from zt import code_registry
    reg = code_registry.load_registry("zj2025")
    missing = [c for c in reg if f"{c}|{reg[c]['name']}|{reg[c]['unit']}" not in system]
    assert not missing, f"码表有 {len(missing)} 条没进分类说明：{missing[:5]}"
    assert "011105001|水泥砂浆踢脚线|m" in system
    assert "## 0111 楼地面装饰工程" in system
    assert "水泥砂浆 + 楼地面/找平 → 011101001" not in system, "旧锚点表会把找平层推向 011101001"
    assert "易混淆对照" in system
    assert out["prompt_version"] in system


def test_gate_never_rewrites_code(monkeypatch):
    txt = "踢脚1：地砖踢脚\n1. 10厚地砖踢脚\n2. 8厚1:2水泥砂浆打底\n"
    fake = FakeAgent(keyword_picker([(("地砖",), "main", "Z011201001")]))   # QA r9：踢脚判成墙面抹灰
    # fix17（协调者裁定 8）：有可展示待核行 → 照常出表 + 告警，不再 no_items
    out = run([txt_file(txt)], agent=fake)
    assert out["items"] == [] and out["status"] == "ok"
    assert any(w["code"] == "quantity_no_formal_items" for w in out["warning_items"])
    pend = [u for u in out["unmatched"] if u["reason_code"] == "semantic_mismatch"]
    assert len(pend) == 1 and pend[0]["kind"] == "pending_review"
    assert pend[0]["llm_guessed_code"] == "Z011201001", "闸只降级，编码原样保留给人看"
    # 主项降级后，它的打底层没有主项可并 → 同样待核（不许悄悄消失）
    assert [u["reason_code"] for u in out["unmatched"] if u is not pend[0]] == ["aux_without_main"]
    assert "011105" in (pend[0]["suggested"] or ""), "建议栏写规则期望的码表名称"


def test_semantic_gate_is_called_on_the_pipeline(monkeypatch):
    """反向对照：同一输入换成正确编码就进正式清单（闸不是恒降级）。"""
    txt = "踢脚1：地砖踢脚\n1. 10厚地砖踢脚\n"
    fake = FakeAgent(keyword_picker([(("地砖",), "main", "011105003")]))
    out = _run(monkeypatch, fake, [txt_file(txt)])
    assert [i["base_code"] for i in out["items"]] == ["011105003"]
    assert out["items"][0]["unit"] == "m"


# ============================================================ QTY-03 覆盖核算

def test_coverage_invariant_when_llm_drops_segments(monkeypatch):
    fake = FakeAgent(R9_PICK, drop=("垫层", "找平"))
    out = _run(monkeypatch, fake, [(r9_workbook(), "做法表.xlsx", "xlsx")])
    cov = out["coverage_detail"]
    assert cov["practice_segments"] == sum(cov["by_destination"].values())
    assert cov["by_destination"]["llm_failed"] == 3, "漏答的 3 层（两层垫层 + 找平层）必须有去向"
    ws = load_workbook(BytesIO(_xlsx(out)))["覆盖核算"]
    ids = [ws.cell(row=r, column=1).value for r in range(5, ws.max_row + 1)]
    assert len(ids) == cov["practice_segments"], "覆盖核算页每个片段一行"


ARABIC_SECTIONS_TXT = (
    "6.屋面做法\n1. 3厚SBS改性沥青防水卷材\n2. 20厚1:3水泥砂浆找平层\n"
    "7.楼面做法（卫生间）\n1. 10厚防滑地砖面层\n2. 20厚1:3干硬性水泥砂浆结合层\n"
    "3. 3厚SBS改性沥青防水卷材\n4. 20厚1:3水泥砂浆找平层\n"
    "8.外墙做法\n1. 50厚挤塑聚苯板保温层\n2. 20厚1:3水泥砂浆找平层\n"
)


def _pick_by_context(s):
    """按片段所在做法作答（真模型也只能靠「做法 / 类别」区分同一句找平层该用平面还是立面码）。"""
    t, ctx = s.get("原文", ""), s.get("做法", "") + " " + s.get("类别", "")
    if "SBS" in t:
        return "main", ("Z010904001" if "楼地面" in ctx else "Z010902001")
    if "挤塑" in t:
        return "main", "011001003"
    if "地砖" in t:
        return "main", "011102003"
    if "找平" in t:
        return "main", ("Z011201004" if "墙面" in ctx else "Z011101006")
    return "aux", ""


@pytest.mark.parametrize("sep", [".", "、"])
def test_arabic_section_headings_no_cross_section_merge(monkeypatch, sep):
    """PR #766 评审阻断项：「6.屋面做法 / 7.楼面做法 / 8.外墙做法」+ 编号层。

    改前三节塌成一个「屋面做法」：三处找平层被做法内同码合并成一条、两处卷材合成一条「屋面卷材防水」、
    外墙保温被 roof_insulation 按「屋面」降为待核 —— 而覆盖不变式照样成立（6+1+1=8），核算页看不出来。
    """
    txt = ARABIC_SECTIONS_TXT.replace("6.", f"6{sep}").replace("7.", f"7{sep}").replace("8.", f"8{sep}")
    out = _run(monkeypatch, FakeAgent(_pick_by_context), [txt_file(txt)])
    got = sorted((i["code"], i["practice"], len(i["sources"])) for i in out["items"])
    assert got == sorted([
        ("Z010902001001", "屋面做法", 1),
        ("Z010904001001", "楼面做法（卫生间）", 1),        # 楼面卷材不再并进「屋面卷材防水」
        # 屋面、楼面两处找平同码同规格：按 QTY-09 跨来源去重合成一项，做法名与来源两处都列出（看得见）
        ("Z011101006001", "屋面做法、楼面做法（卫生间）", 2),
        ("011102003001", "楼面做法（卫生间）", 2),          # 地砖 + 并入的结合层
        ("011001003001", "外墙做法", 1),
        ("Z011201004001", "外墙做法", 1),                   # 外墙立面找平单独成项
    ]), got
    level = next(i for i in out["items"] if i["base_code"] == "Z011101006")
    assert [s["evidence"]["heading"] for s in level["sources"]] == ["屋面做法", "楼面做法（卫生间）"]
    assert not out["unmatched"], "外墙保温按正确编码作答，不该被屋面规则降为待核"
    by = out["coverage_detail"]["by_destination"]
    assert (by["item"], by["aux_merged"], by["pending_review"]) == (7, 1, 0)


def test_arabic_sections_same_code_different_spec_not_merged(monkeypatch):
    # 做法内同码合并不看规格：改前「20厚」「30厚」两处找平落在同一个做法里，被合成一条
    txt = ARABIC_SECTIONS_TXT.replace("4. 20厚1:3水泥砂浆找平层", "4. 30厚1:3水泥砂浆找平层")
    out = _run(monkeypatch, FakeAgent(_pick_by_context), [txt_file(txt)])
    level = sorted((i["code"], i["practice"], [s["quote"] for s in i["sources"]])
                   for i in out["items"] if i["base_code"] == "Z011101006")
    assert level == [("Z011101006001", "屋面做法", ["20厚1:3水泥砂浆找平层"]),
                     ("Z011101006002", "楼面做法（卫生间）", ["30厚1:3水泥砂浆找平层"])], level


def _pick_by_category(s):
    """模拟真模型：同一句卷材 / 涂膜，靠「类别 / 做法 / 部位」区分屋面防水和楼(地)面防水；上下文空着就当楼地面。"""
    t = s.get("原文", "")
    roof = "屋面" in " ".join(str(s.get(k, "")) for k in ("类别", "做法", "部位"))
    if "SBS" in t:
        return "main", ("Z010902001" if roof else "Z010904001")
    if "涂膜" in t or "涂料" in t:
        return "main", ("Z010902002" if roof else "Z010904002")
    if "细石混凝土" in t:
        return "main", ("Z010902004" if roof else "011101002")
    if "找平" in t:
        return "main", "Z011101006"
    return "aux", ""


ROOF1 = ("屋面1 不上人屋面（用于主楼屋面）", "主楼屋面", "屋面")


def _assert_roof1_items(out):
    codes = sorted(i["base_code"] for i in out["items"])
    assert not {"Z010904001", "Z010904002"} & set(codes), f"屋面卷材 / 涂膜被静默编成楼(地)面防水：{codes}"
    assert {"Z010902001", "Z010902002"} <= set(codes), codes
    for i in out["items"]:
        assert (i["practice"], i["part"], i["category"]) == ROOF1, (i["code"], i["practice"], i["part"])
        # D 列由代码填的两个槽（做法名称、部位）不能是空的
        assert f"建筑做法名称：{ROOF1[0]}" in i["features"] and f"部位：{ROOF1[1]}" in i["features"], i["features"]


@pytest.mark.parametrize("a,b", [("①", "②"), ("1、", "2、"), ("1）", "2）"), ("(1)", "(2)")])
def test_nested_sub_numbers_not_silently_miscoded(monkeypatch, a, b):
    """PR #766 第二轮评审阻断项：「2. 防水层：」下的 ①② 子层。

    改后（上一轮）子层被当成新做法、上下文全空，按类别作答的模型把卷材 / 涂膜编成 Z010904001 / 002
    楼(地)面防水 —— status=ok、不进待核，覆盖核算照样平账。
    """
    txt = ("屋面1：不上人屋面（用于主楼屋面）\n1. 保护层：40厚C20细石混凝土\n2. 防水层：\n"
           f"{a}3厚SBS改性沥青防水卷材\n{b}2厚聚氨酯防水涂膜\n3. 找平层：20厚1:3水泥砂浆\n")
    out = _run(monkeypatch, FakeAgent(_pick_by_category), [txt_file(txt)])
    _assert_roof1_items(out)
    sbs = next(i for i in out["items"] if i["base_code"] == "Z010902001")
    assert sbs["sources"][0]["quote"].endswith("3厚SBS改性沥青防水卷材")


def test_layer_restart_without_heading_keeps_practice_context(monkeypatch):
    # 层号回到 1、前面没有节标题：新做法沿用当前做法的名称 / 类别 / 部位（上一轮清空成空串，涂膜被编成楼地面）
    txt = ("屋面1：不上人屋面（用于主楼屋面）\n1. 3厚SBS改性沥青防水卷材\n2. 20厚1:3水泥砂浆找平层\n"
           "1. 2厚聚氨酯防水涂膜\n2. 30厚1:3水泥砂浆找平层\n")
    out = _run(monkeypatch, FakeAgent(_pick_by_category), [txt_file(txt)])
    _assert_roof1_items(out)
    assert out["status"] == "ok" and not out["unmatched"]


def test_invariant_breach_raises(monkeypatch):
    seg = segmod.Segment(kind=PRACTICE, text="x", file_index=0, file_name="a", location="", id="s1")
    with pytest.raises(ql.CoverageInvariantError):
        ql.check_invariant([seg], {})


# ============================================================ QTY-04 可复现

def test_rerun_with_same_answers_is_identical(monkeypatch):
    # zt 改动：旧版靠进程内缓存复现；现在 prepare / finalize 全确定性，同一份 answers 重跑逐字相同
    import random

    rnd = random.Random(7)

    def pick(s):
        return "main", rnd.choice(["011102003", "Z011101006", "011101001"])

    fake = FakeAgent(pick)
    files = [txt_file(FLOOR_TXT)]
    first = _run(monkeypatch, fake, files)
    second = ql.finalize(files, first["answers"], prepare_fingerprint=first["prepare_fingerprint"])
    assert [i["code"] for i in second["items"]] == [i["code"] for i in first["items"]]
    assert [u["reason_code"] for u in second["unmatched"]] == [u["reason_code"] for u in first["unmatched"]]
    assert ql.prepare(files)["batches"] == first["prepare"]["batches"]
    wb1 = load_workbook(BytesIO(_xlsx(first)))
    from zt.quantity.render import render_xlsx
    wb2 = load_workbook(BytesIO(render_xlsx(second)))
    for name in ("分部分项工程量清单", "待核与未匹配", "覆盖核算"):
        v1 = [[c.value for c in row] for row in wb1[name].iter_rows()]
        v2 = [[c.value for c in row] for row in wb2[name].iter_rows()]
        assert v1 == v2, name


def test_prepare_fingerprint_sensitivity(monkeypatch):
    # zt 改动：旧版验缓存键（prompt_version / 码表指纹 / 片段内容 / max_tokens 变了要重新调用）；
    # 现在同样的四类变化必须改变 prepare_fingerprint，finalize 拿旧指纹就拒绝（PrepareMismatch）
    files = [txt_file(FLOOR_TXT)]
    base = ql.prepare(files)["prepare_fingerprint"]
    assert ql.prepare(files)["prepare_fingerprint"] == base, "同输入同指纹"
    changed = ql.prepare([txt_file(FLOOR_TXT.replace("10厚", "12厚"))])["prepare_fingerprint"]
    assert changed != base, "片段内容变了指纹必须变"
    assert ql.prepare(files, config=ql.QuantityConfig(batch_max_segments=2))["prepare_fingerprint"] != base, "装批参数变了"
    monkeypatch.setattr(ql, "QUANTITY_PROMPT_VERSION", "qty-test-bump")
    assert ql.prepare(files)["prepare_fingerprint"] != base, "prompt_version 变了指纹必须变"
    monkeypatch.undo()
    qp.registry_fingerprint.cache_clear()
    monkeypatch.setattr(qp, "registry_fingerprint", lambda v: "0" * 64)
    assert ql.prepare(files)["prepare_fingerprint"] != base, "码表指纹变了指纹必须变"
    with pytest.raises(ql.PrepareMismatch):
        ql.finalize(files, {"r": []}, prepare_fingerprint=base)


# 提示词指纹：改了提示词（分类说明、片段格式、易混淆对照、条件候选、特征摘录说明）而不升版本号 → 红。
# 升版本时在这里追加一行新版本的指纹，历史行不许改。
PROMPT_PINS = {
    "qty-2026-09-30.1": "480d92e358aeab5b4aa4b388f2b520c3b166ce882c3690f2ab708f9383ffedde",
    "qty-2026-10-05.1": "8f65337885344f3f5e01ce7c3192f05f9b14398eda9ed25619225343e08fde43",
    "qty-2026-10-10.1": "ea3bf81504b4dab81eadbc4b0b6a7d7a571bc83fa692565a7d058adf17827814",
    # zt：改为 agent 作答（全局片段编号、answers 文件、特征 k 用清单项编码）。
    "qty-2026-10-19.zt1": "915c3d38e3f36692a2f78a5cc1cae01513dcd597ba5bb12045c9a255f4538c8b",
}


def test_prompt_version_pinned():
    # zt 改动：校验 prompt_version 常量与提示词指纹成对登记
    digest = hashlib.sha256(qp.prompt_material().encode("utf-8")).hexdigest()
    assert qp.QUANTITY_PROMPT_VERSION == ql.QUANTITY_PROMPT_VERSION == "qty-2026-10-19.zt1"
    assert qp.QUANTITY_PROMPT_VERSION in PROMPT_PINS, "升了版本号就在 PROMPT_PINS 里登记新指纹"
    assert PROMPT_PINS[qp.QUANTITY_PROMPT_VERSION] == digest == qp.prompt_fingerprint(), (
        f"提示词变了（sha256={digest}），但 QUANTITY_PROMPT_VERSION 没升")


def test_segmentation_deterministic():
    a = segmod.segment_file(r9_workbook(), "做法表.xlsx", "xlsx", "", 0)
    b = segmod.segment_file(r9_workbook(), "做法表.xlsx", "xlsx", "", 0)
    ids_a = [(s.id, s.text, s.location) for s in segmod.assign_ids([a])]
    ids_b = [(s.id, s.text, s.location) for s in segmod.assign_ids([b])]
    assert ids_a == ids_b and ids_a


def test_response_records_versions(monkeypatch):
    out = _run(monkeypatch, FakeAgent(FLOOR_PICK), [txt_file(FLOOR_TXT)])
    assert out["prompt_version"] == ql.QUANTITY_PROMPT_VERSION == out["prepare"]["prompt_version"]
    assert out["prepare_fingerprint"] == out["prepare"]["prepare_fingerprint"]
    assert len(out["registry_fingerprint"]) == 64 and out["code_version"] == "zj2025"


# ============================================================ QTY-05 空结果

def test_empty_llm_result_is_422_no_items(monkeypatch):
    fake = FakeAgent(lambda s: ("main", ""))
    with pytest.raises(ql.QuantityRejected) as ei:
        run([txt_file(FLOOR_TXT)], agent=fake)
    e = ei.value
    assert e.reason_code == "no_items"
    # 面层、找平层没拿到编码 → 未匹配；结合层是辅层，主项都没了 → 待核（aux_without_main）
    assert e.coverage["practice_segments"] == 3
    assert e.coverage["by_destination"]["unmatched"] == 2 and e.coverage["by_destination"]["pending_review"] == 1
    assert sorted(u["reason_code"] for u in e.unmatched) == ["aux_without_main", "llm_no_code", "llm_no_code"]


def test_no_practice_segments_is_422(monkeypatch):
    fake = FakeAgent(FLOOR_PICK)
    with pytest.raises(ql.QuantityRejected) as ei:
        run([txt_file("本工程为某某办公楼，建筑面积 5000 平方米。\n设计依据详见总说明。\n")], agent=fake)
    assert ei.value.reason_code == "no_practice"
    assert "0101–0118" in ei.value.detail
    assert fake.batches == [], "没有做法片段就不交给 agent"


# ============================================================ QTY-07 专业范围闸

HVAC = ("空调风管采用镀锌钢板，风机盘管吊装，排烟风机设于屋顶，新风机组、配电箱、桥架、电缆、照明灯具、"
        "插座、弱电综合布线、监控摄像机、消火栓、喷淋喷头、报警阀；风管保温采用30厚橡塑保温板。\n") * 3
ROAD = ("道路路基填方压实度≥96%，路面结构：4厚细粒式沥青混凝土，6厚中粒式沥青混凝土，"
        "20厚水泥稳定碎石基层，底基层级配碎石；车行道、人行道侧石平石，桥梁桥台桥墩箱梁伸缩缝，雨水口、边沟。\n") * 3


def test_municipal_file_rejected_before_llm(monkeypatch):
    fake = FakeAgent(FLOOR_PICK)
    with pytest.raises(ql.QuantityRejected) as ei:
        run([txt_file(ROAD, "路基说明.txt")], agent=fake)
    assert ei.value.reason_code == "out_of_scope" and "市政" in ei.value.detail
    assert fake.batches == []


def test_installation_file_rejected(monkeypatch):
    fake = FakeAgent(FLOOR_PICK)
    with pytest.raises(ql.QuantityRejected) as ei:
        run([txt_file(HVAC, "暖通说明.txt")], agent=fake)
    assert ei.value.reason_code == "out_of_scope" and "安装" in ei.value.detail
    assert fake.batches == []


def test_mixed_batch_skips_out_of_scope_file(monkeypatch):
    fake = FakeAgent(FLOOR_PICK)
    out = _run(monkeypatch, fake, [txt_file(HVAC, "暖通说明.txt"), txt_file(FLOOR_TXT, "建筑做法.txt")])
    asked_files = {b["file_name"] for b in fake.batches}
    assert asked_files == {"建筑做法.txt"}, "越界文件的片段不交给 agent"
    assert any(w["code"] == "quantity_file_out_of_scope" for w in out["warning_items"])
    assert out["coverage_detail"]["by_destination"]["out_of_scope"] >= 1
    assert out["items"]


def test_bridge_section_segments_out_of_scope(monkeypatch):
    txt = FLOOR_TXT + "（五）桥梁工程\n桥面铺装：100厚C40防水混凝土\n"
    fake = FakeAgent(keyword_picker([(("地砖",), "main", "011102003"), (("找平",), "main", "Z011101006"),
                                   (("C40",), "main", "010502030")]))
    out = _run(monkeypatch, fake, [txt_file(txt)])
    asked = [s["原文"] for s in fake.segments_seen]
    assert not any("C40" in a for a in asked)
    outs = [u for u in out["unmatched"] if u["kind"] == "out_of_scope"]
    assert len(outs) == 1 and outs[0]["reason_code"] == "section_out_of_scope"


# ============================================================ QTY-09 粒度

def test_same_practice_same_code_merged(monkeypatch):
    txt = ("顶棚1：纸面石膏板吊顶（用于办公室）\n1. 9.5厚纸面石膏板面层\n2. 50系列轻钢龙骨吊顶\n"
           "3. φ8吊杆吊顶\n4. 刷防火涂料吊顶面层\n")
    fake = FakeAgent(lambda s: ("main", "Z011302001"))
    out = _run(monkeypatch, fake, [txt_file(txt)])
    assert [i["code"] for i in out["items"]] == ["Z011302001001"], "一个吊顶只出一条"
    assert len(out["items"][0]["sources"]) == 4


def test_skirting_face_and_base_layers_merged(monkeypatch):
    txt = "踢脚1：水泥踢脚（高150）\n1. 10厚1:2水泥砂浆罩面压实赶光\n2. 12厚1:3水泥砂浆打底扫毛\n"
    fake = FakeAgent(lambda s: ("main", "011105001"))
    out = _run(monkeypatch, fake, [txt_file(txt)])
    assert [i["base_code"] for i in out["items"]] == ["011105001"]
    roles = sorted(src["role"] for src in out["items"][0]["sources"])
    assert roles == ["aux", "main"], "打底层并进罩面主项"


def test_general_notes_and_practice_table_dedup(monkeypatch):
    notes = "（四）屋面工程\n屋面防水：2.0厚聚氨酯防水涂料\n"
    table = "屋面1：上人屋面\n1. 2厚聚氨酯防水涂料\n"
    fake = FakeAgent(keyword_picker([(("聚氨酯",), "main", "Z010902002")]))
    out = _run(monkeypatch, fake, [txt_file(notes, "总说明.txt"), txt_file(table, "做法表.txt")])
    assert [i["code"] for i in out["items"]] == ["Z010902002001"], "同码同规格（2.0 = 2）只出一条"
    files = {s["file_name"] for s in out["items"][0]["sources"]}
    assert files == {"总说明.txt", "做法表.txt"}, "M 列列出两处来源"
    assert out["items"][0]["source"].count("原文") == 2


def test_different_practices_not_merged(monkeypatch):
    txt = "地面1：水泥地面\n1. 20厚1:2水泥砂浆面层\n地面2：水泥地面\n1. 30厚1:2水泥砂浆面层\n"
    fake = FakeAgent(lambda s: ("main", "011101001"))
    out = _run(monkeypatch, fake, [txt_file(txt)])
    assert [i["code"] for i in out["items"]] == ["011101001001", "011101001002"]


def test_aux_with_standalone_word_goes_pending(monkeypatch):
    txt = "楼面1：防水楼面\n1. 10厚地砖面层\n2. 1.5厚聚氨酯防水涂料\n"
    fake = FakeAgent(keyword_picker([(("地砖",), "main", "011102003"), (("防水",), "aux", "")]))
    out = _run(monkeypatch, fake, [txt_file(txt)])
    pend = [u for u in out["unmatched"] if u["kind"] == "pending_review"]
    assert [u["reason_code"] for u in pend] == ["role_conflict"]


def test_aux_only_layer_forced_aux(monkeypatch):
    txt = "内墙1：乳胶漆墙面\n1. 刷乳胶漆两遍\n2. 满刮腻子两遍\n"
    fake = FakeAgent(lambda s: ("main", "011404001"))    # 模型把腻子也判成主项
    out = _run(monkeypatch, fake, [txt_file(txt)])
    assert [i["base_code"] for i in out["items"]] == ["011404001"]
    assert out["coverage_detail"]["by_destination"]["aux_merged"] == 1


# ============================================================ QTY-11 标段排除

LOT_TXT = ("（四）装修说明\n盖边柱、道路两侧柱、盖板顶棚涂料等盖板装修计入IV标。\n"
           "外墙2：盖边柱真石漆\n1. 喷真石漆两遍\n"
           "内墙1：乳胶漆墙面\n1. 刷乳胶漆两遍\n")
LOT_PICK = keyword_picker([(("真石漆",), "main", "011404001"), (("乳胶漆",), "main", "011404001")])


@pytest.mark.parametrize("project,dest", [
    ("某站房工程III标", "excluded_scope"),
    ("某站房工程Ⅳ标段", None),
    ("某站房工程", "pending_review"),
])
def test_lot_exclusion_moves_items_out(monkeypatch, project, dest):
    fake = FakeAgent(LOT_PICK)
    out = _run(monkeypatch, fake, [txt_file(LOT_TXT, "建筑说明.txt")], project_name=project)
    col = [u for u in out["unmatched"] if "真石漆" in u["text"]]
    if dest is None:
        assert not col, "排除句说计入 IV 标，本项目就是 IV 标 → 照常列项"
        assert any("真石漆" in s["quote"] for i in out["items"] for s in i["sources"])
    else:
        assert len(col) == 1 and col[0]["kind"] == dest
        assert "盖边柱" in col[0]["reason"] and "计入IV标" in col[0]["reason"], "引用排除句原文"
        assert "p" not in col[0]["location"] or "行" in col[0]["location"]
    # 与排除句无关的内墙照常列项
    assert any("乳胶漆两遍" in s["quote"] for i in out["items"] for s in i["sources"])


def test_llm_flagged_exclusion_without_code_match_goes_pending(monkeypatch):
    fake = FakeAgent(keyword_picker([(("真石漆",), "main", "011404001"),
                                   (("乳胶漆",), "main", "011404001", "e1"),
                                   (("地砖",), "main", "011102003"), (("找平",), "main", "Z011101006")]))
    out = _run(monkeypatch, fake, [txt_file(LOT_TXT + FLOOR_TXT)], project_name="某工程III标")
    by_text = {u["text"]: u["reason_code"] for u in out["unmatched"]}
    assert by_text.get("刷乳胶漆两遍") == "llm_exclusion_unconfirmed"
    assert by_text.get("喷真石漆两遍") == "excluded_lot"


def test_llm_unknown_exclusion_id_ignored_with_warning(monkeypatch):
    fake = FakeAgent(keyword_picker([(("地砖",), "main", "011102003", "e9"), (("找平",), "main", "Z011101006")]))
    out = _run(monkeypatch, fake, [txt_file(FLOOR_TXT)])
    assert any(w["code"] == "quantity_llm_unknown_exclusion" for w in out["warning_items"])
    assert "011102003" in [i["base_code"] for i in out["items"]]


# ============================================================ QTY-12 排序

def test_interleaved_sections_sorted_once(monkeypatch):
    txt = ("墙1：外墙\n1. 20厚1:3水泥砂浆面层外墙抹灰\n"
           "地面1：地砖\n1. 10厚地砖面层\n"
           "屋面1：屋面\n1. 3厚SBS防水卷材\n"
           "地面2：水泥\n1. 20厚1:2水泥砂浆面层\n"
           "屋面2：屋面\n1. 4厚SBS防水卷材\n"
           "地面3：地砖\n1. 8厚地砖面层\n")
    fake = FakeAgent(keyword_picker([(("外墙",), "main", "Z011201001"), (("SBS",), "main", "Z010902001"),
                                   (("地砖",), "main", "011102003"), (("水泥砂浆面层",), "main", "011101001")]))
    out = _run(monkeypatch, fake, [txt_file(txt)])
    codes = [i["code"] for i in out["items"]]
    assert codes == ["Z010902001001", "Z010902001002", "011101001001", "011102003001", "011102003002",
                     "Z011201001001"], codes
    ws = load_workbook(BytesIO(_xlsx(out))).active
    titles = [ws.cell(row=r, column=3).value for r in range(4, ws.max_row + 1)
              if ws.cell(row=r, column=1).value is None and ws.cell(row=r, column=3).value]
    sections = [t for t in titles if not str(t).startswith("注")]
    assert len(sections) == len(set(sections)) == 3, "分部标题行数 = 不同分部数"


# ============================================================ QTY-13 位置由代码产生

def test_pipeline_ignores_llm_location(monkeypatch):
    class Liar(FakeAgent):
        def answer(self, prep):
            obj = super().answer(prep)
            for r in obj.get("r", []):
                r["source_location"] = "Sheet1 第1行"
                r["source_quote"] = "模型编的摘录"
            return obj

    out = _run(monkeypatch, Liar(R9_PICK), [(r9_workbook(), "做法表.xlsx", "xlsx")])
    locs = [s["location"] for i in out["items"] for s in i["sources"]]
    assert locs and all(l.startswith("做法表!") for l in locs)
    assert not any("Sheet1 第1行" in l for l in locs)
    assert not any("模型编的摘录" in s["quote"] for i in out["items"] for s in i["sources"])
    floor = next(i for i in out["items"] if i["base_code"] == "011101002")
    assert floor["sources"][0]["location"] == "做法表!D3（第3行·楼地面做法第1层）"
    assert floor["sources"][0]["evidence"]["cell"] == "D3" and floor["sources"][0]["evidence"]["row"] == 3


# ============================================================ QTY-06 特征摘录两步走

def test_features_answered_on_second_finalize(monkeypatch):
    # zt 改动：旧用例验「每一次调用的 token 都进账」；现在特征摘录是 finalize 给任务、agent 作答、再 finalize
    fake = FakeAgent(FLOOR_PICK)
    files = [txt_file(FLOOR_TXT)]
    first = _run(monkeypatch, fake, files)
    tasks = first["feature_tasks"]
    assert tasks and {t["k"] for t in tasks} == {i["code"] for i in first["items"]}
    assert any(w["code"] == "quantity_features_skipped" for w in first["warning_items"])
    tile = next(t for t in tasks if t["编码"] == "011102003")
    slot = next(n for n, lbl in tile["槽位"].items() if "规格" in lbl or "品种" in lbl)
    answers = {**first["answers"], "f": [{"k": tile["k"], "v": {slot: "10厚防滑地砖"}},
                                          {"k": "999999999001", "v": {"1": "x"}}]}
    second = ql.finalize(files, answers, prepare_fingerprint=first["prepare_fingerprint"])
    item = next(i for i in second["items"] if i["code"] == tile["k"])
    assert "〔AI摘录〕10厚防滑地砖" in item["features"]
    assert tile["k"] not in {t["k"] for t in second["feature_tasks"]}
    assert any(w["code"] == "quantity_features_unknown_items" for w in second["warning_items"])


# ============================================================ WS29-f 范围规则（issue #835）
# 生产实测（做法表没有表头，一行各格拼成一句走文本模式）：「范围排除」22 条全错，3 个做法整行消失；
# 13 处「用户自理」散在正式清单、未匹配、待核里；条文号「13」当对象词排掉了全部住宅外墙保温。

SCOPE_PICK = keyword_picker([
    (("花岗岩",), "main", "011102001"), (("地砖",), "main", "011102003"),
    (("防水涂料",), "main", "Z010904002"), (("找平",), "main", "Z011101006"),
    (("无机涂料",), "main", "011404001"), (("防霉涂料",), "main", "011404001"),
    (("真石漆",), "main", "011404001"), (("保温",), "main", "011001003"),
    (("石膏板",), "main", "Z011302001"),
])
TABLE_NOTE = "注：本表上注明详二次装修设计部位的材料做法仅为暂定，规格材质最终由装修设计确定。"


def _owner_docx():
    """两个做法、8 种层内自理句式，各做法另有两层正常的层（夹具必须同时有这两类层）。"""
    a, b = OWNER_INLINE_LAYERS[:4], OWNER_INLINE_LAYERS[4:]
    row = lambda layers: " ".join(f"{i})" + t for i, t in enumerate(layers, start=1))  # noqa: E731
    return docx_file(docx_table(
        ("■ 楼面2", "卫生间楼面", "50厚", "适用部位"),
        (row(a + ["1.5厚聚合物水泥防水涂料", "20厚1:3水泥砂浆找平层"]), "住宅卫生间"),
        ("■ 楼面3", "阳台楼面", "50厚", "适用部位"),
        (row(b + ["20厚花岗岩面层", "20厚1:3水泥砂浆找平层"]), "住宅阳台"),
    ))


def _sent_texts(fake) -> list[str]:
    return [s["原文"] for s in fake.segments_seen]


def _excl(out, scope=None):
    return [e for e in out["coverage_detail"]["exclusions"] if scope is None or e["scope"] == scope]


def test_inline_owner_layers_excluded_with_quote_and_not_sent_to_llm(monkeypatch):
    fake = FakeAgent(SCOPE_PICK)
    out = _run(monkeypatch, fake, [_owner_docx()])
    inline = _excl(out, "inline")
    assert [e["sentence"] for e in inline] == OWNER_INLINE_LAYERS, "8 层各占一条，原句逐字"
    assert {e["reason_code"] for e in inline} == {"excluded_owner_inline"} and {e["kind"] for e in inline} == {"owner"}
    assert all(e["hits"] == 1 and e["location"].startswith("表1 第") for e in inline)
    assert _excl(out, "global") == [], "层内自理不产生全局排除句"
    rows = [u for u in out["unmatched"] if u["reason_code"] == "excluded_owner_inline"]
    assert [u["text"] for u in rows] == OWNER_INLINE_LAYERS and {u["kind"] for u in rows} == {"excluded_scope"}
    for u in rows:
        assert u["text"] in u["reason"] and "用户自理" in u["reason"], "范围排除要引用原句"
    assert not [t for t in _sent_texts(fake) if "自理" in t], "自理的层不送模型"
    # 同一做法里另一层照常分类（全局误伤的反向对照）
    quotes = [s["quote"] for i in out["items"] for s in i["sources"]]
    assert any("20厚花岗岩面层" in q for q in quotes) and any("1.5厚聚合物水泥防水涂料" in q for q in quotes)
    assert not any("自理" in q for q in quotes), "正式清单里搜「用户自理」0 处（基线 3 处）"
    cd = out["coverage_detail"]
    assert cd["by_destination"]["excluded_scope"] == 8
    assert cd["practice_segments"] == sum(cd["by_destination"].values()), "带层内标记的层计入覆盖核算"
    ws = load_workbook(BytesIO(_xlsx(out)))
    listed = [str(c.value) for row in ws["分部分项工程量清单"].iter_rows() for c in row if c.value]
    assert listed and not [v for v in listed if "自理" in v]
    cov = [[c.value for c in row] for row in ws["覆盖核算"].iter_rows(min_row=5)]
    assert len(cov) == cd["practice_segments"]
    for layer in OWNER_INLINE_LAYERS:
        hit = [r for r in cov if r[7] == layer]
        assert len(hit) == 1 and hit[0][8] == "范围排除", f"被排除的层在覆盖核算页查得到：{layer}"


def test_inline_marker_priority_on_pipeline(monkeypatch):
    doc = docx_file(docx_table(
        ("1)面层瓷砖用户自理（详二次装修） 2)5厚抗裂砂浆压入一道玻纤网格布 3)1.5厚Ⅱ型聚合物水泥防水涂料", "住宅厨房"),
        ("1)面层瓷砖用户自理或（详精装修设计） 2)1.5厚Ⅱ型聚合物水泥防水涂料 3)20厚1:3水泥砂浆找平层", "住宅卫生间"),
        ("1)面层为二次装修 2)20厚1:3水泥砂浆找平层", "商业"),
        ("1)防滑地砖面层由业主自理 2)20厚花岗岩面层 3)20厚1:3水泥砂浆找平层", "门厅"),
    ))
    out = _run(monkeypatch, FakeAgent(SCOPE_PICK), [doc])
    by = {u["text"]: u for u in out["unmatched"]}
    # 「…由业主自理」写在层里只管这一层：同一行的花岗岩面层照常列项，也不拿「防滑地砖面层」去全局匹配
    assert by["防滑地砖面层由业主自理"]["reason_code"] == "excluded_owner_inline"
    assert any("20厚花岗岩面层" in s["quote"] for i in out["items"] for s in i["sources"])
    own = by["面层瓷砖用户自理（详二次装修）"]
    assert (own["kind"], own["reason_code"]) == ("excluded_scope", "excluded_owner_inline")
    assert "原文同时注明「详二次装修」" in own["reason"]
    amb = [u for u in out["unmatched"] if u["reason_code"] == "owner_inline_ambiguous"]
    assert [(u["kind"], u["text"]) for u in amb] == [("pending_review", "面层瓷砖用户自理或（详精装修设计）")]
    sec = by["面层为二次装修"]
    assert (sec["kind"], sec["reason_code"]) == ("excluded_scope", "excluded_secondary_inline")
    assert [(e["kind"], e["sentence"]) for e in _excl(out, "inline")] == [
        ("owner", "面层瓷砖用户自理（详二次装修）"), ("secondary", "面层为二次装修"),
        ("owner", "防滑地砖面层由业主自理")], "待核的那一层不算排除"
    assert _excl(out, "global") == []
    assert len([i for i in out["items"] if i["base_code"] == "Z010904002"]) >= 1, "同一行的防水涂料照常列项"


def test_fitout_reference_layers_listed_with_provisional_mark(monkeypatch):
    marked = ["防滑地砖面层（详精装修设计）", "30厚DS15干硬性预拌砂浆结合层（详精装修设计）",
              "1.5厚Ⅱ型聚合物水泥防水涂料（详精装修设计）"]
    doc = docx_file(docx_table(
        (" ".join(f"{i})" + t for i, t in enumerate(marked + ["20厚1:3水泥砂浆找平层"], start=1)), "住宅卫生间"),
    ) + docx_p(TABLE_NOTE))
    fake = FakeAgent(SCOPE_PICK)
    out = _run(monkeypatch, fake, [doc])
    assert set(marked) <= set(_sent_texts(fake)), "带精装修引用的层照常送模型分类"
    cd = out["coverage_detail"]
    assert cd["by_destination"] == {**{d: 0 for d in ql.DESTINATIONS}, "item": 3, "aux_merged": 1}, "去向不变"
    assert cd["provisional_segments"] == 3
    warn = [w for w in out["warning_items"] if w["code"] == "quantity_fitout_provisional"]
    assert len(warn) == 1 and warn[0]["level"] == "info" and warn[0]["message"].startswith("3 个做法层")
    assert "仅为暂定" in warn[0]["message"], "表注原文作为依据引用在告警里"
    notes = {i["base_code"]: i["notes"] for i in out["items"]}
    assert "做法暂定" in notes["011102003"][0] and "本条 2 层" in notes["011102003"][0], "地砖 + 并入的结合层"
    assert "做法暂定" in notes["Z010904002"][0]
    assert notes["Z011101006"] == [], "没有标记的找平层不带「做法暂定」"
    wb = load_workbook(BytesIO(_xlsx(out)))
    m_col = [wb["分部分项工程量清单"].cell(row=r, column=13).value or "" for r in range(4, 12)]
    assert sum("做法暂定" in v for v in m_col) == 2
    cov = [[c.value for c in row] for row in wb["覆盖核算"].iter_rows(min_row=5)]
    assert [("做法暂定" in (r[10] or "")) for r in cov] == [True, True, True, False]
    assert _excl(out) == [] and not out["unmatched"]


def test_fitout_reference_on_practice_name_marks_all_layers(monkeypatch):
    txt = ("顶棚3：装修吊顶（详二次装修图纸）\n1. U型轻钢龙骨中距600\n2. 9.5厚纸面石膏板\n"
           "顶棚4：石膏板吊顶\n1. U型轻钢龙骨中距600\n2. 12厚纸面石膏板\n")
    out = _run(monkeypatch, FakeAgent(SCOPE_PICK), [txt_file(txt)])
    assert out["coverage_detail"]["provisional_segments"] == 2
    marked = [i for i in out["items"] if i["notes"]]
    assert len(marked) == 1 and "顶棚3" in marked[0]["practice"] and "本条 2 层" in marked[0]["notes"][0]


def _three_practices_docx():
    """基线里整行消失的三个做法：适用部位写着「不需 / 不需要 / 无 二次装修」。"""
    neg = NOT_EXCLUSION_SENTENCES[:3]
    return docx_file(docx_table(
        ("■ 内墙9", "无机涂料墙面", "25厚", "适用部位"),
        ("1)无机涂料（一底两面）饰面 2)满刮2厚柔性耐水腻子两遍分层找平 3)6厚DP15预拌砂浆找平 "
         "4)基层墙体（混凝土墙梁柱处素水泥浆甩毛）", neg[0]),
        ("■ 内墙12", "防霉涂料墙面", "15厚", "适用部位"),
        ("1)无机防霉涂料（一底两面）饰面 2)批防潮腻子两道 3)12厚DP15预拌砂浆分层打底找平 "
         "4)基层墙体（混凝土墙梁柱处素水泥浆甩毛）", neg[1]),
        ("■ 顶棚2", "腻子顶棚", "5厚", "适用部位"),
        ("1)满刮专业石膏腻子批白两遍 2)板底刷素界面剂一道 3)钢筋混凝土楼板清理（孔洞用M20聚合物砂浆修补平整）",
         neg[2]),
    ) + docx_p(TABLE_NOTE) + "".join(docx_p(f"{n} {t}") for n, t in zip(
        ("13.13", "15.18", "17.3"), (NOT_EXCLUSION_SENTENCES[3], NOT_EXCLUSION_SENTENCES[4],
                                    NOT_EXCLUSION_SENTENCES[5]))))


def test_sentences_that_only_mention_fitout_exclude_nothing(monkeypatch):
    fake = FakeAgent(SCOPE_PICK)
    out = _run(monkeypatch, fake, [_three_practices_docx()])
    cd = out["coverage_detail"]
    assert _excl(out) == [] and cd["by_destination"]["excluded_scope"] == 0, "基线：8 条排除句、22 条误排"
    sent = _sent_texts(fake)
    for layer in ("无机涂料（一底两面）饰面", "无机防霉涂料（一底两面）饰面", "满刮专业石膏腻子批白两遍"):
        assert layer in sent, f"基线里整行消失的做法回到片段里：{layer}"
    # 标题行「■ 内墙12 防霉涂料墙面 15厚」带材料词和厚度，仍会被当成一层（归 WS29-k1），所以不数条数，只看来源
    quotes = [s["quote"] for i in out["items"] if i["base_code"] == "011404001" for s in i["sources"]]
    for layer in ("无机涂料（一底两面）饰面", "无机防霉涂料（一底两面）饰面"):
        assert layer in quotes, f"内墙无机涂料、地下室防霉涂料两个做法都列项：{layer}"
    everywhere = [u["text"] for u in out["unmatched"]] + [s["quote"] for i in out["items"] for s in i["sources"]]
    assert "满刮专业石膏腻子批白两遍" in everywhere, "顶棚腻子在清单页或待核页有对应条目"
    related = [x["sentence"] for x in cd["scope_related"]]
    for s in NOT_EXCLUSION_SENTENCES:
        assert any(s in r for r in related), f"读到了但没据此排除的句子要列进「范围相关原文」：{s}"
    assert {x["reason"] for x in cd["scope_related"]} == {"not_exclusion"}
    notes = load_workbook(BytesIO(_xlsx(out)))["说明"]
    col = {notes.cell(row=r, column=1).value: notes.cell(row=r, column=2).value for r in range(1, notes.max_row + 1)}
    assert col["范围相关原文"].startswith(f"{len(related)} 句")
    assert sum(1 for k in col if k.startswith("范围相关原文 ")) == len(related)
    assert any("不需二次装修的非配水房间" in (v or "") for k, v in col.items() if k.startswith("范围相关原文 "))


def test_generic_object_sentence_warns_and_excludes_nothing(monkeypatch):
    doc = docx_file(docx_p("注：面层为二次装修") + docx_p("面层详二次装修")
                    + docx_table(("编号", "名称", "构造做法"), ("楼面1", "块料楼面", "1.10厚防滑地砖面层"),
                                 ("", "", "2.20厚1:3水泥砂浆找平层")))
    out = _run(monkeypatch, FakeAgent(SCOPE_PICK), [doc])
    assert _excl(out) == [], "「面层」不能当全局对象词"
    assert [i["base_code"] for i in out["items"]] == ["Z011101006", "011102003"], "带「面层」的地砖层照常列项"
    warn = [w for w in out["warning_items"] if w["code"] == "quantity_exclusion_object_too_generic"]
    assert len(warn) == 1 and warn[0]["level"] == "warn" and warn[0]["message"].startswith("1 句")
    rel = {x["sentence"]: x["reason"] for x in out["coverage_detail"]["scope_related"]}
    assert rel == {"注：面层为二次装修": "too_generic", "面层详二次装修": "not_exclusion"}


def test_clause_number_never_excludes_layers_containing_it(monkeypatch):
    # 做法表 + 建施说明一起勾选（ZJ-B08）：改前对象词「13」排掉全部带「13」的片段，包括住宅外墙保温
    table = docx_file(docx_table(
        ("编号", "名称", "构造做法"),
        ("外墙1", "1#2#3#12#13#15#住宅外墙", "1.30mm厚无机轻集料保温砂浆"),
        ("", "", "2.13厚DP15预拌砂浆找平")), "做法表.docx")
    notes = docx_file("".join(docx_p(t) for t in (
        "13.13 二次装修室内空间应考虑隔声降噪措施。", "15.18 景观及二次装修的材料及做法均应符合现行的防火规范要求。",
        "17.3 厨房设备见二次装修。", "23.3 住宅二次装修时，不得破坏承重结构的墙体。")), "建施说明.docx")
    out = _run(monkeypatch, FakeAgent(SCOPE_PICK), [table, notes])
    assert _excl(out) == []
    assert {i["base_code"] for i in out["items"]} == {"011001003", "Z011101006"}, "保温砂浆那一层照常列项"
    assert out["coverage_detail"]["by_destination"]["excluded_scope"] == 0
    assert len(out["coverage_detail"]["scope_related"]) == 4


# —— 跨表、跨文件的全局排除（QTY-11 不回归的关键夹具）：排除句写在说明段落，做法在另一张表 / 另一个文件 ——

def _practice_table(code: str, name: str, layer: str) -> str:
    return docx_table(("编号", "名称", "构造做法"), (code, name, "1." + layer),
                      ("内墙1", "乳胶漆墙面", "1.刷无机涂料两遍"))


@pytest.mark.parametrize("case", ["lot_other_file", "not_in_scope_other_table", "owner_other_table"])
def test_global_exclusion_reaches_other_table_and_other_file(monkeypatch, case):
    if case == "lot_other_file":
        files = [txt_file("（四）装修说明\n盖边柱、道路两侧柱等盖板装修计入IV标。\n", "建筑说明.txt"),
                 docx_file(_practice_table("外墙2", "盖边柱真石漆", "喷真石漆两遍"))]
        code, layer, obj = "excluded_lot", "喷真石漆两遍", "盖边柱"
    elif case == "not_in_scope_other_table":
        files = [docx_file(docx_p("屋面保温不在本次设计范围。")
                           + docx_table(("编号", "名称", "构造做法"), ("楼面1", "地砖楼面", "1.10厚防滑地砖面层"))
                           + docx_p("以下为屋面做法")
                           + _practice_table("屋面1", "上人屋面保温", "60厚挤塑聚苯板保温层"))]
        code, layer, obj = "excluded_not_in_scope", "60厚挤塑聚苯板保温层", "屋面保温"
    else:
        files = [docx_file(docx_p("卫生间洁具由业主另行发包")
                           + docx_table(("编号", "名称", "构造做法"), ("楼面1", "地砖楼面", "1.10厚防滑地砖面层"))
                           + docx_p("以下为卫生间做法")
                           + _practice_table("楼面5", "卫生间洁具安装基层", "20厚1:3水泥砂浆找平层"))]
        code, layer, obj = "excluded_owner", "20厚1:3水泥砂浆找平层", "卫生间洁具"
    out = _run(monkeypatch, FakeAgent(SCOPE_PICK), files, project_name="某站房工程III标")
    hit = [u for u in out["unmatched"] if u["text"] == layer]
    assert len(hit) == 1 and (hit[0]["kind"], hit[0]["reason_code"]) == ("excluded_scope", code), out["unmatched"]
    assert obj in hit[0]["reason"]
    glob = _excl(out, "global")
    assert len(glob) == 1 and glob[0]["hits"] == 1 and obj in glob[0]["objects"] and glob[0]["id"] == "e1"
    assert _excl(out, "inline") == []
    assert any("刷无机涂料两遍" in s["quote"] for i in out["items"] for s in i["sources"]), "无关的内墙照常列项"


def test_global_exclusion_with_zero_hits_is_still_listed(monkeypatch):
    doc = docx_file(docx_p("卫生间洁具由业主另行发包")
                    + docx_table(("编号", "名称", "构造做法"), ("楼面1", "地砖楼面", "1.10厚防滑地砖面层")))
    out = _run(monkeypatch, FakeAgent(SCOPE_PICK), [doc])
    glob = _excl(out, "global")
    assert [(e["id"], e["kind"], e["objects"], e["hits"]) for e in glob] == [("e1", "owner", ["卫生间洁具"], 0)]
    notes = load_workbook(BytesIO(_xlsx(out)))["说明"]
    col = {notes.cell(row=r, column=1).value: notes.cell(row=r, column=2).value for r in range(1, notes.max_row + 1)}
    assert col["排除句清单"].startswith("全局排除句 1 条")
    assert "卫生间洁具由业主另行发包" in col["排除句 e1"] and "排掉 0 个片段" in col["排除句 e1"]
    assert "对象词：卫生间洁具" in col["排除句 e1"] and "段落#1" in col["排除句 e1"]


# —— 做法名称块后面的编号表注（PR #839 评审阻断项）。文本模式的做法表后面跟「注：1. … 2. …」是常见版式 ——

TABLE_NOTE_DOC = (
    "楼面2：卫生间楼面\n1. 10厚防滑地砖面层\n2. 20厚1:3水泥砂浆找平层\n"
    "{notes}"
    "（五）其他做法\n内墙3：住宅户内墙面\n1. 刷无机涂料两遍\n"
    "楼面5：卫生间洁具安装基层\n1. 20厚1:3水泥砂浆找平层\n")
TABLE_NOTES = {
    "header_line": "注：\n1. 卫生间洁具、橱柜由业主另行发包。\n2. 住宅户内墙面均为二次装修。\n",
    "header_inline": "注：1. 卫生间洁具、橱柜由业主另行发包。\n2. 住宅户内墙面均为二次装修。\n",
    "other_mark": "注：\n1、卫生间洁具、橱柜由业主另行发包。\n2、住宅户内墙面均为二次装修。\n",
    "no_header": "1. 卫生间洁具、橱柜由业主另行发包。\n2. 住宅户内墙面均为二次装修。\n",
}


@pytest.mark.parametrize("form", sorted(TABLE_NOTES))
def test_numbered_table_notes_after_named_block_exclude_globally(monkeypatch, form):
    # 改前（header_line）：第 1 句落成节标题，排除句清单、范围相关原文、未匹配里都查不到，楼面5 的找平层进了
    # 正式清单；（header_inline）：第 2 句不再全局，内墙3 的涂料进了正式清单，「范围相关原文」还说它不是排除句
    out = _run(monkeypatch, FakeAgent(SCOPE_PICK), [txt_file(TABLE_NOTE_DOC.format(notes=TABLE_NOTES[form]))])
    glob = _excl(out, "global")
    assert [(e["kind"], e["objects"], e["hits"]) for e in glob] == [
        ("owner", ["卫生间洁具", "橱柜"], 1), ("secondary", ["住宅户内墙面"], 1)]
    assert _excl(out, "inline") == [], "表注不是哪个做法的层"
    excluded = [(u["text"], u["reason_code"]) for u in out["unmatched"] if u["kind"] == "excluded_scope"]
    assert sorted(excluded) == [("20厚1:3水泥砂浆找平层", "excluded_owner"), ("刷无机涂料两遍", "excluded_secondary")]
    quotes = [s["quote"] for i in out["items"] for s in i["sources"]]
    assert any("10厚防滑地砖面层" in q for q in quotes), "表注上面那个做法照常列项"
    assert len([q for q in quotes if "找平" in q]) == 1, "楼面2 的找平层在，楼面5 的被排掉"
    assert not any("无机涂料" in q for q in quotes)
    assert out["coverage_detail"]["scope_related"] == [], "两句都是排除句，不该列成「读到了但没有据此排除」"
    cd = out["coverage_detail"]
    assert cd["by_destination"]["excluded_scope"] == 2 and cd["practice_segments"] == sum(cd["by_destination"].values())


def test_all_layers_inline_excluded_is_no_items_with_ledger(monkeypatch):
    fake = FakeAgent(SCOPE_PICK)
    with pytest.raises(ql.QuantityRejected) as ei:
        run([txt_file("楼面2：卫生间楼面\n1. 预留防滑地砖装修面层（用户自理）\n2. 面层用户自理\n")], agent=fake)
    e = ei.value
    assert e.reason_code == "no_items", "认出了做法层，只是都划出范围：不能报成「没有识别到任何工程做法」"
    assert e.coverage["practice_segments"] == 2 and e.coverage["by_destination"]["excluded_scope"] == 2
    assert [u["reason_code"] for u in e.unmatched] == ["excluded_owner_inline"] * 2
    assert fake.batches == [], "没有可分类的层，不交给 agent"


# ============================================================ 标题行 + 构造层行（WS29-k1，issue #891）

TITLED_PICK = keyword_picker([
    (("花岗岩",), "main", "011102001"), (("花岗石",), "main", "011107001"), (("无机涂料",), "main", "011404001"),
    (("石膏腻子",), "main", "011406003"),
])


def test_titled_table_items_have_part_and_no_title_rows(monkeypatch):
    fake = FakeAgent(TITLED_PICK)
    out = _run(monkeypatch, fake, [(titled_docx(), "做法表.docx", "docx")])
    sent = _sent_texts(fake)
    assert sent and not [t for t in sent if "■" in t or "适用部位" in t], "标题行不送模型分类"
    rows = fake.segments_seen
    assert {r["部位"] for r in rows if "花岗岩面层" in r["原文"]} == {"电梯厅、公共走道"}, "部位随片段送模型"
    everywhere = [str(v) for i in out["items"] for v in (i["source"], i["features"])] + \
                 [u["text"] for u in out["unmatched"]]
    assert not [v for v in everywhere if "■" in v], "清单页与待核页都没有标题行生成的条目（基线 21～26 条）"
    assert out["items"] and all(i["part"] for i in out["items"]), "每条清单项都有部位（基线 141/142 待补）"
    assert not [i for i in out["items"] if "部位：待补" in i["features"]]
    wall = [i for i in out["items"] if i["base_code"] == "011404001"]
    assert len(wall) == 1 and wall[0]["practice"] == "内墙9 无机涂料内墙面1（燃烧性能A级）"
    assert {s["location"] for s in wall[0]["sources"]} == {"表1 第4行第1列"}, "来源指向内墙9 的构造层行，不是下一个标题行"
    cd = out["coverage_detail"]
    assert cd["practices"] == 4
    assert cd["practice_segments"] == sum(cd["by_destination"].values())
    assert _excl(out) == [] and cd["by_destination"]["excluded_scope"] == 0, "「不需二次装修」的部位不触发排除"
    part = [x for x in cd["scope_related"] if "不需二次装修" in x["sentence"]]
    assert [(x["location"], x["reason"]) for x in part] == [("表1 第4行第2列（适用部位）", "not_exclusion")]
    assert "适用部位文字" in part[0]["note"]
    notes = load_workbook(BytesIO(_xlsx(out)))["说明"]
    col = {notes.cell(row=r, column=1).value: notes.cell(row=r, column=2).value for r in range(1, notes.max_row + 1)}
    assert col["识别出的做法"].startswith("识别出做法 4 个")
    cov = [[c.value for c in row] for row in load_workbook(BytesIO(_xlsx(out)))["覆盖核算"].iter_rows(min_row=5)]
    assert len(cov) == cd["practice_segments"] and {r[4] for r in cov} == {b[-1] for _t, b in TITLED_PRACTICES}



# ============================================================ 错码闸、漏切层、分类稳定（WS29-k2，issue #913）

K2_FILE = "建施完整材料做法表.docx"


def _k2(name: str = K2_FILE):
    return (k2_docx(), name, "docx")


def test_wrong_codes_for_ramp_stair_and_waterproof_mortar_go_pending(monkeypatch):
    out = _run(monkeypatch, FakeAgent(K2_PICK), [_k2()])
    pend = {u["text"]: u for u in out["unmatched"] if u["reason_code"] == "semantic_mismatch"}
    ramp = pend["90厚C20细石混凝土随捣随抹，内配钢筋网片"]
    assert ramp["kind"] == "pending_review" and ramp["llm_guessed_code"] == "011107001"
    assert ramp["suggested"] == "规则期望：010502031 坡道（m²）" and "坡道不是台阶" in ramp["reason"]
    stair = pend["20厚1:2水泥砂浆面层，踏步做防滑条"]
    assert stair["llm_guessed_code"] == "011101001" and "011106001 水泥砂浆楼梯" in stair["suggested"]
    mortar = pend["20厚1：2水泥防水砂浆"]
    assert mortar["llm_guessed_code"] == "010905002"
    assert "010903003 墙面砂浆防水" in mortar["suggested"] and "Z010904003" in mortar["suggested"]
    codes = {i["base_code"] for i in out["items"]}
    assert not codes & {"011107001", "011101001", "010905002"}, "三条错码都不进正式清单"
    assert "011107002" in codes, "台阶给台阶的码照常列项（反例不被误伤）"
    ws = load_workbook(BytesIO(_xlsx(out)))["待核与未匹配"]
    sugg = {r[4]: r[8] for r in ([c.value for c in row] for row in ws.iter_rows(min_row=2))}
    assert sugg["90厚C20细石混凝土随捣随抹，内配钢筋网片"] == "规则期望：010502031 坡道（m²）"


def test_paint_layers_in_coverage_and_derusting_listed_on_notes(monkeypatch):
    out = _run(monkeypatch, FakeAgent(K2_PICK), [_k2()])
    cd = out["coverage_detail"]
    wb = load_workbook(BytesIO(_xlsx(out)))
    cov = {r[7]: r for r in ([c.value for c in row] for row in wb["覆盖核算"].iter_rows(min_row=5))}
    for layer, dest in (("氟碳金属面漆（颜色另详）", "正式清单项"), ("氟碳金属底漆", "并入主项特征"),
                        ("刷专用防锈漆", "并入主项特征")):
        assert cov[layer][8] == dest, layer
    assert K2_DERUST not in cov, "「除锈」不是做法层，不进覆盖核算"
    assert cd["block_lines_unrecognized"] == [{"text": K2_DERUST, "file_name": K2_FILE, "location": "表1 第10行第1列",
                                               "practice": "漆1 外露铁件油漆", "practice_has_layers": True}]
    warn = [w for w in out["warning_items"] if w["code"] == "quantity_block_lines_unrecognized"]
    assert len(warn) == 1 and warn[0]["level"] == "info" and "1 行" in warn[0]["message"]
    col = {wb["说明"].cell(row=r, column=1).value: wb["说明"].cell(row=r, column=2).value
           for r in range(1, wb["说明"].max_row + 1)}
    assert col["块内未识别行"].startswith("1 行")
    assert col["块内未识别行 1"] == f"{K2_FILE} 表1 第10行第1列 | 做法：漆1 外露铁件油漆 | 原文：{K2_DERUST}"
    assert "片段编号" in col["复现性"] and "answers" in col["复现性"]


def test_block_lines_name_practices_with_no_layer_at_all(monkeypatch):
    doc = docx_file(docx_table(("■ 楼面1", "地砖楼面", "50厚", "适用部位"),
                               ("1)10厚防滑地砖面层 2)20厚1:3水泥砂浆找平层", "门厅"),
                               ("■ 配件1", "车库配件", "适用部位"), ("1)车轮挡：见图集", "车库")))
    out = _run(monkeypatch, FakeAgent(FLOOR_PICK), [doc])
    assert [(x["text"], x["practice_has_layers"]) for x in out["coverage_detail"]["block_lines_unrecognized"]] == [
        ("车轮挡：见图集", False)]
    warn = next(w for w in out["warning_items"] if w["code"] == "quantity_block_lines_unrecognized")
    assert "1 个做法一层都没有认出" in warn["message"]
    col = {c[0].value: c[1].value for c in load_workbook(BytesIO(_xlsx(out)))["说明"].iter_rows()}
    assert "一层都没有认出：配件1 车库配件" in col["块内未识别行"]
    assert "（这个做法一层都没有认出）" in col["块内未识别行 1"]


def test_no_block_lines_still_says_zero_on_notes(monkeypatch):
    out = _run(monkeypatch, FakeAgent(FLOOR_PICK), [txt_file(FLOOR_TXT)])
    assert out["coverage_detail"]["block_lines_unrecognized"] == []
    assert not [w for w in out["warning_items"] if w["code"] == "quantity_block_lines_unrecognized"]
    col = {c[0].value: c[1].value for c in load_workbook(BytesIO(_xlsx(out)))["说明"].iter_rows()}
    assert col["块内未识别行"].startswith("0 行")


def _file_prompts(fake, name):
    return [b for b in fake.batches if b["file_name"] == name]


def test_same_file_same_batches_alone_or_after_a_file_with_an_exclusion(monkeypatch):
    """E32：同一份做法表单独生成、与另一份文件一起生成（那份在前、带一条全局排除句），做法表各批的片段与排除句相同。

    zt 改动：片段用全局编号（agent 直接按编号作答，不再批内重编号），所以只比片段内容；排除句仍只列与本批有关的。"""
    alone = FakeAgent(K2_PICK)
    _run(monkeypatch, alone, [_k2()])
    both = FakeAgent(K2_PICK)
    out = _run(monkeypatch, both, [txt_file(OTHER_FILE_TXT, "建施说明.txt"), _k2()])
    glob = [e for e in out["coverage_detail"]["exclusions"] if e["scope"] == "global"]
    assert [(e["id"], e["objects"]) for e in glob] == [("e1", ["卫生间洁具"])]
    assert out["coverage_detail"]["by_file"][0]["segments_total"] > 0, "前一份文件占了全局编号"
    strip = lambda bs: [([{k: v for k, v in s.items() if k != "id"} for s in b["segments"]], b["exclusions"])  # noqa: E731
                        for b in bs]
    a, b = _file_prompts(alone, K2_FILE), _file_prompts(both, K2_FILE)
    assert a and strip(a) == strip(b)
    assert b[0]["exclusions"] == [], "无关的排除句不列进本批"
    assert b[0]["segments"][0]["id"] != "s1", "全局编号：前一份文件的片段在前"


@pytest.mark.skip(reason=NO_LLM.format("进程内缓存命中后编码稳定"))
def test_second_run_with_another_file_hits_cache_and_keeps_codes():
    pass


def test_relevant_exclusion_listed_under_batch_local_id(monkeypatch):
    lot_txt = "盖边柱、道路两侧柱等盖板装修计入IV标。\n外墙2：盖边柱真石漆\n1. 喷真石漆两遍\n内墙1：乳胶漆墙面\n1. 刷乳胶漆两遍\n"
    fake = FakeAgent(keyword_picker([(("真石漆",), "main", "011404001"), (("乳胶漆",), "main", "011404001", "e1"),
                                   (("地砖",), "main", "011102003"), (("找平",), "main", "Z011101006")]))
    out = _run(monkeypatch, fake, [txt_file(OTHER_FILE_TXT, "建施说明.txt"), txt_file(lot_txt, "外墙说明.txt")],
               project_name="某工程III标")
    glob = {e["id"]: e for e in out["coverage_detail"]["exclusions"] if e["scope"] == "global"}
    assert glob["e2"]["objects"][0] == "盖边柱", "全局编号：前一份文件的排除句是 e1，这一条是 e2"
    batch = _file_prompts(fake, "外墙说明.txt")[0]
    assert batch["exclusions"] == [{"id": "e2", "sentence": "盖边柱、道路两侧柱等盖板装修计入IV标。"}], \
        "对象词出现在本批片段里：本批只列这一条，用全局编号 e2"
    assert "卫生间洁具" not in json.dumps(batch, ensure_ascii=False)
    by_text = {u["text"]: u for u in out["unmatched"]}
    assert by_text["喷真石漆两遍"]["reason_code"] == "excluded_lot"
    # 假 agent 答的批内第 1 条换成全局 e2，不是前一份文件的 e1
    assert by_text["刷乳胶漆两遍"]["reason_code"] == "llm_exclusion_unconfirmed"
    assert "（e2）" in by_text["刷乳胶漆两遍"]["reason"]


def test_exclusion_id_not_in_this_batch_is_ignored_even_if_it_exists_globally(monkeypatch):
    # 本批提示词里没有排除句；模型答「e1」，全局恰好有 e1（前一份文件的）。不能拿它当前一份文件的 e1 用
    fake = FakeAgent(keyword_picker([(("面漆",), "main", "011402001", "e1"), (("底漆",), "aux", ""),
                                   (("防锈漆",), "aux", "")]))
    out = _run(monkeypatch, fake, [txt_file(OTHER_FILE_TXT, "建施说明.txt"),
                                   txt_file("外墙5：外露铁件油漆\n1. 氟碳金属面漆\n2. 刷专用防锈漆\n", "铁件.txt")])
    assert "011402001" in {i["base_code"] for i in out["items"]}
    assert not [u for u in out["unmatched"] if u["reason_code"] == "llm_exclusion_unconfirmed"]
    assert any(w["code"] == "quantity_llm_unknown_exclusion" for w in out["warning_items"])


def test_answers_in_shuffled_order_map_back_to_their_segments(monkeypatch):
    class Reversed(FakeAgent):
        def answer(self, prep):
            obj = super().answer(prep)
            obj["r"] = list(reversed(obj["r"]))
            return obj

    fake = Reversed(lambda seg: (FLOOR_PICK if "地砖" in seg["原文"] else K2_PICK)(seg))
    out = _run(monkeypatch, fake, [txt_file(OTHER_FILE_TXT, "建施说明.txt"), _k2()])
    got = {s["quote"]: i["base_code"] for i in out["items"] for s in i["sources"] if s["role"] == "main"}
    assert got["30厚花岗石板铺面，灌水泥浆擦缝"] == "011107002"
    assert got["100厚C15混凝土垫层"] == "010501002"
    assert got["氟碳金属面漆（颜色另详）"] == "011402001"
    assert got["10厚防滑地砖面层"] == "011102003"
    assert not [w for w in out["warning_items"] if w["code"] == "quantity_llm_unknown_ids"]


# ============================================================ WS29-p 同内容的文件只读一份（issue #970，E13，K-R43）
# 生产 ZJ-B09：子目录与根目录各一份字节相同的做法表，两份都切片、都送模型 —— 片段 770（单份 385）、
# 27 万 token、待核 135 条（单份 59），「来源依据」同一个位置写两遍。夹具必须是「不同文件名、相同字节」：
# 同名的话按文件名去重也能过，测不出「按内容不按文件名」。

K2_COPY = "建施完整材料做法表（副本）.docx"


def _dup_warns(out):
    return [w for w in out["warning_items"] if w["code"] == "quantity_duplicate_file"]


def _single_k2(monkeypatch):
    fake = FakeAgent(K2_PICK)
    out = _run(monkeypatch, fake, [_k2()])
    return out, len(fake.segments_seen)


def test_same_bytes_different_names_read_once(monkeypatch):
    single, single_calls = _single_k2(monkeypatch)
    fake = FakeAgent(K2_PICK)
    out = _run(monkeypatch, fake, [_k2(), _k2(K2_COPY)])
    cd = out["coverage_detail"]
    # 片段数、去向、待核条数都与单份相同（基线：770 对 385，待核 135 对 59）
    assert cd["practice_segments"] == single["coverage_detail"]["practice_segments"] > 0
    assert cd["by_destination"] == single["coverage_detail"]["by_destination"]
    assert len(out["unmatched"]) == len(single["unmatched"])
    # 交给 agent 的片段与单份同量：副本一个片段都没送
    assert len(fake.segments_seen) == single_calls and not _file_prompts(fake, K2_COPY)
    warn = _dup_warns(out)
    assert len(warn) == 1 and warn[0]["level"] == "info"
    assert warn[0]["message"] == f"「{K2_FILE}」与「{K2_COPY}」内容相同，只按「{K2_FILE}」生成"
    by_file = {f["file_name"]: f for f in cd["by_file"]}
    assert set(by_file) == {K2_FILE, K2_COPY}
    assert by_file[K2_COPY]["duplicate_of"] == {"file_id": None, "file_name": K2_FILE, "display_path": None}
    assert by_file[K2_COPY]["practice_segments"] == 0 and by_file[K2_FILE].get("duplicate_of") is None
    # 清单每条的来源只写一处（基线：127 条把同一个位置写了两遍）
    for it in out["items"]:
        places = [(s["file_name"], s["location"], s["quote"]) for s in it["sources"]]
        assert len(places) == len(set(places)), it["code"]
        assert {s["file_name"] for s in it["sources"]} == {K2_FILE}
        assert K2_COPY not in it["source"]
    assert [i["code"] for i in out["items"]] == [i["code"] for i in single["items"]]
    env = out["coverage"]
    assert (env["files_total"], env["files_used"]) == (2, 1)
    assert [(s["file_name"], s["stage"]) for s in env["skipped"]] == [(K2_COPY, "dedupe")]
    wb = load_workbook(BytesIO(_xlsx(out)))
    col = {wb["说明"].cell(row=r, column=1).value: wb["说明"].cell(row=r, column=2).value
           for r in range(1, wb["说明"].max_row + 1)}
    assert col[f"输入文件：{K2_COPY}"] == f"与「{K2_FILE}」内容相同，没有单独读取，只按那一份生成"


def test_same_name_different_content_not_merged(monkeypatch):
    other = FLOOR_TXT.replace("10厚防滑地砖面层", "8厚防滑地砖面层")
    out = _run(monkeypatch, FakeAgent(FLOOR_PICK), [txt_file(FLOOR_TXT, "做法说明.txt"), txt_file(other, "做法说明.txt")])
    assert not _dup_warns(out)
    assert out["coverage_detail"]["practice_segments"] == 6
    assert [f.get("duplicate_of") for f in out["coverage_detail"]["by_file"]] == [None, None]
    assert out["coverage"]["files_used"] == 2


def test_two_of_three_identical_keeps_two(monkeypatch):
    floor_alone = _run(monkeypatch, FakeAgent(FLOOR_PICK), [txt_file(FLOOR_TXT)])["coverage_detail"]["practice_segments"]
    single, _ = _single_k2(monkeypatch)
    out = _run(monkeypatch, FakeAgent(lambda s: (FLOOR_PICK if "地砖" in s["原文"] else K2_PICK)(s)),
               [_k2(), txt_file(FLOOR_TXT, "地面做法.txt"), _k2(K2_COPY)])
    cd = out["coverage_detail"]
    assert cd["practice_segments"] == single["coverage_detail"]["practice_segments"] + floor_alone
    assert [(f["file_name"], (f.get("duplicate_of") or {}).get("file_name")) for f in cd["by_file"]] == [
        (K2_FILE, None), ("地面做法.txt", None), (K2_COPY, K2_FILE)]
    warn = _dup_warns(out)
    assert len(warn) == 1 and "地面做法.txt" not in warn[0]["message"]
    assert out["coverage"]["files_used"] == 2


def test_same_name_copies_in_two_folders_named_by_path(monkeypatch):
    """生产的形状：同名、两个目录。有 display_path 就用路径点名；没有就写「第 N 份」，不出现「『A』与『A』」。"""
    refs = [{"file_id": "f1", "display_path": "/1.资料/做法表.docx"}, {"file_id": "f2", "display_path": "/做法表.docx"}]
    out = _run(monkeypatch, FakeAgent(K2_PICK), [_k2("做法表.docx"), _k2("做法表.docx")], file_refs=refs)
    assert _dup_warns(out)[0]["message"] == "「/1.资料/做法表.docx」与「/做法表.docx」内容相同，只按「/1.资料/做法表.docx」生成"
    assert out["coverage_detail"]["by_file"][1]["duplicate_of"] == {
        "file_id": "f1", "file_name": "做法表.docx", "display_path": "/1.资料/做法表.docx"}
    out = _run(monkeypatch, FakeAgent(K2_PICK), [_k2("做法表.docx"), _k2("做法表.docx")])
    assert _dup_warns(out)[0]["message"] == (
        "「做法表.docx（第 1 份）」与「做法表.docx（第 2 份）」内容相同，只按「做法表.docx（第 1 份）」生成")


def test_unreadable_first_copy_falls_back_to_the_copy_with_a_sidecar(monkeypatch):
    """字节相同的扫描件，只有后一份旁边有 OCR 旁车：不能因为「保留先出现的」就读不到正文。"""
    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(width=595, height=842)
    buf = BytesIO()
    w.write(buf)
    blank = buf.getvalue()
    asked: list[int] = []

    # zt 改动：PDF 读取回调返回 DocText（文字层 + 旁车，由 zt.textextract.load_document 提供）
    def loader(idx, content, file_name):
        asked.append(idx)
        return DocText(text=FLOOR_TXT, text_source="ocr_sidecar") if idx == 1 else failed_doc("no_text", "扫描页")

    out = run([(blank, "扫描说明.pdf", "pdf"), (blank, "扫描说明-旁车.pdf", "pdf")], pdf_loader=loader, agent=FakeAgent(FLOOR_PICK))
    assert asked == [0, 1] * 2, "prepare 与 finalize 各读一遍（zt 改动：finalize 重跑确定性前半段）"
    assert out["coverage_detail"]["practice_segments"] == 3
    assert not [w for w in out["warning_items"] if w["code"] == "quantity_file_skipped_no_text"]
    assert _dup_warns(out)[0]["message"] == "「扫描说明-旁车.pdf」与「扫描说明.pdf」内容相同，只按「扫描说明-旁车.pdf」生成"
    assert {f["file_name"]: (f.get("duplicate_of") or {}).get("file_name") for f in out["coverage_detail"]["by_file"]} == {
        "扫描说明.pdf": "扫描说明-旁车.pdf", "扫描说明-旁车.pdf": None}


def test_empty_files_are_never_same_content(monkeypatch):
    """0 字节的 sha256 全都一样。两份空 PDF 各有自己的旁车（按下标给不同正文）→ 两份都要读，不能并成「同一份」。"""
    other = "地面2：水泥地面（用于车库）\n1. 20厚1:2水泥砂浆面层\n2. 100厚C15混凝土垫层\n"

    def loader(idx, content, file_name):
        return DocText(text=(FLOOR_TXT if idx == 0 else other), text_source="ocr_sidecar")

    out = run([(b"", "说明一.pdf", "pdf"), (b"", "说明二.pdf", "pdf")], pdf_loader=loader, agent=FakeAgent(FLOOR_PICK))
    assert not _dup_warns(out)
    assert out["coverage_detail"]["practice_segments"] == 5
    assert out["coverage"]["files_used"] == 2


def test_three_identical_copies_one_warning_naming_both_dropped(monkeypatch):
    out = _run(monkeypatch, FakeAgent(K2_PICK), [_k2(), _k2(K2_COPY), _k2("做法表R.docx")])
    assert [w["message"] for w in _dup_warns(out)] == [
        f"「{K2_FILE}」与「{K2_COPY}」、「做法表R.docx」内容相同，只按「{K2_FILE}」生成"]
    assert [(f.get("duplicate_of") or {}).get("file_name") for f in out["coverage_detail"]["by_file"]] == [
        None, K2_FILE, K2_FILE]


def test_same_content_groups_keep_input_order_and_isolate_unhashed():
    from zt.quantity import same_content as sc

    assert sc.groups(["a", None, "b", "a", None, "b"]) == [[0, 3], [1], [2, 5], [4]]
    assert sc.groups(["x", "y"]) == [[0], [1]], "没有重复时与逐份处理一模一样"
    assert sc.content_key(b"") is None and sc.content_key(b"x") == sc.content_key(b"x") != sc.content_key(b"y")
    assert sc.labels(["a.docx", "b.docx", "a.docx"]) == ["a.docx（第 1 份）", "b.docx", "a.docx（第 3 份）"]
