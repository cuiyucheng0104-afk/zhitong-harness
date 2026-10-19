# 来源：engine-service/tests/test_quantity_list.py（zhitong-cost-ai@991ff5cd，W4 迁入 zt；改动见「zt 改动」注释）
"""清单 CLI（prepare / finalize / render）：入参、结果信封、拒绝与失败的映射、跨平台路径。

zt 改动：旧用例测 FastAPI 端点（鉴权、HTTP 状态码、usage）。zt 没有 HTTP 层：成功是信封（ok / partial），
输入齐全但没有结果是 ``empty`` + status_reason，调用错误与「没有回答」是退出码 2 + 错误 JSON。鉴权、usage、
DeepSeek 凭据缺失、模型输出不可解析这几类端点用例没有对应物，不迁。
"""
from __future__ import annotations

import io
import json
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

from zt.quantity import prompt as qp
from zt.quantity.cli import cli
from tests.quantity.fixtures.shapes import FakeAgent, keyword_picker
from tests.quantity.support import run_cli

WALL_ROWS = ["内墙2：蒸压粉煤灰砖内墙（用于地下室内墙）", "1. 200mm厚MU20蒸压粉煤灰砖，M7.5水泥砂浆砌筑"]
WALL_PICK = keyword_picker([(("粉煤灰砖",), "main", "010401003")])


def _xlsx(root: Path, rows: list[str], rel: str = "资料/设计说明-测试.xlsx") -> str:
    """在项目目录里落一份设计说明 xlsx（一行一格，没有表头 → 逐行处理），返回相对路径。"""
    wb = Workbook()
    ws = wb.active
    for i, t in enumerate(rows, start=1):
        ws.cell(row=i, column=1, value=t)
    p = root.joinpath(*rel.split("/"))
    p.parent.mkdir(parents=True, exist_ok=True)
    wb.save(p)
    return rel


def _call(*argv) -> tuple[int, dict]:
    out = io.StringIO()
    code = cli.main([str(a) for a in argv], stdout=out)
    return code, json.loads(out.getvalue())


def _flow(root: Path, rels: list[str], agent: FakeAgent, *extra) -> tuple[int, dict]:
    """prepare → 假 agent 作答 → finalize；返回 finalize 的 (退出码, 完整信封或错误)。"""
    code, prep = _call("prepare", "--root", root, "--files", *rels, *extra)
    if code != 0 or prep["status"] != "ok":
        return code, prep
    (root / "answers.json").write_text(json.dumps(agent.answer(prep["data"]), ensure_ascii=False), encoding="utf-8")
    (root / "prep.json").write_text(json.dumps(prep, ensure_ascii=False), encoding="utf-8")
    return _call("finalize", "--prepare", root / "prep.json", "--answers", root / "answers.json")


def test_prepare_requires_files():
    code, body = _call("prepare")
    assert code == 2 and body["error"]["code"] == "no_files"


def test_rejects_unknown_code_version(tmp_path):
    rel = _xlsx(tmp_path, WALL_ROWS)
    code, body = _call("prepare", "--root", tmp_path, "--files", rel, "--code-version", "not-a-version")
    assert code == 2 and body["error"]["code"] == "unknown_code_version" and "zj2025" in body["error"]["hint"]


def test_happy_path_writes_instructions_batches_and_workbook(tmp_path):
    rel = _xlsx(tmp_path, WALL_ROWS)
    out = tmp_path / "工作" / "2026-10-19-清单"
    code, summary = run_cli("zt.quantity.cli", "prepare", "--root", str(tmp_path), "--files", rel,
                            "--project-name", "某工程III标", "--eng-category", "房屋建筑", "--out", f"{out / 'prepare'}/")
    assert code == 0 and summary["status"] == "ok"
    prep = json.loads((out / "prepare" / "result.json").read_text(encoding="utf-8"))
    data = prep["data"]
    assert data["prompt_version"] == qp.QUANTITY_PROMPT_VERSION and data["code_version"] == "zj2025"
    assert data["project_lot"] == 3 and data["segments_to_classify"] == 1
    instructions = (out / "prepare" / "清单分类说明.md").read_text(encoding="utf-8")
    batches = (out / "prepare" / "清单分类批次.md").read_text(encoding="utf-8")
    assert "010401003|多孔砖墙|" in instructions and qp.QUANTITY_PROMPT_VERSION in instructions
    assert "3 标" in batches and "房屋建筑" in batches and '"id": "s1"' in batches
    assert data["artifacts"] == {"instructions": (out / "prepare" / "清单分类说明.md").as_posix(),
                                 "batches": (out / "prepare" / "清单分类批次.md").as_posix()}

    answers = FakeAgent(WALL_PICK).answer(data)
    (out / "answers.json").write_text(json.dumps(answers, ensure_ascii=False), encoding="utf-8")
    code, summary = run_cli("zt.quantity.cli", "finalize", "--prepare", str(out / "prepare" / "result.json"),
                            "--answers", str(out / "answers.json"), "--out", f"{out / 'finalize'}/")
    assert code == 0 and summary["status"] == "ok", summary
    fin = json.loads((out / "finalize" / "result.json").read_text(encoding="utf-8"))
    assert fin["needs_manual"] is True
    items = fin["data"]["items"]
    assert [i["code"] for i in items] == ["010401003001"] and items[0]["name"] == "多孔砖墙", "名称来自码表"
    assert items[0]["needs_manual"] is True
    src = items[0]["sources"][0]
    assert src["evidence"]["rel_path"] == rel and src["evidence"]["cell"] == "A2"
    assert fin["coverage"]["files_total"] == 1 and fin["coverage"]["files_used"] == 1
    cd = fin["data"]["coverage_detail"]
    assert cd["practice_segments"] == sum(cd["by_destination"].values()) == 1 and cd["project_lot"] == 3

    code, summary = run_cli("zt.quantity.cli", "render", "--result", str(out / "finalize" / "result.json"),
                            "--xlsx", str(out / "清单框架.xlsx"))
    assert code == 0 and summary["data"]["xlsx"] == (out / "清单框架.xlsx").as_posix()
    wb = load_workbook(out / "清单框架.xlsx")
    assert wb.sheetnames == ["分部分项工程量清单", "待核与未匹配", "覆盖核算", "说明"]
    assert wb.active.cell(row=3, column=2).value == "项目编码"


def test_windows_style_relative_path_and_chinese_space_dir(tmp_path):
    root = tmp_path / "项目" / "2026-007-某 工程"
    rel = _xlsx(root, WALL_ROWS, "资料/子目录 一/设计说明.xlsx")
    code, fin = _flow(root, [rel.replace("/", "\\")], FakeAgent(WALL_PICK))
    assert code == 0 and fin["status"] == "ok"
    assert fin["data"]["items"][0]["sources"][0]["evidence"]["rel_path"] == rel, "反斜杠规范成 /"
    code, body = _call("prepare", "--root", root, "--files", "../外面.xlsx")
    assert code == 2 and body["error"]["code"] == "invalid_path"


def test_scope_exclusion_trace_survives_the_envelope(tmp_path):
    rows = ["楼面2：卫生间楼面", "1. 预留防滑地砖装修面层（用户自理）", "2. 1.5厚聚合物水泥防水涂料（详精装修设计）",
            "3. 20厚1:3水泥砂浆找平层", "注：面层为二次装修", "卫生间洁具由业主另行发包"]
    agent = FakeAgent(keyword_picker([(("防水涂料",), "main", "Z010904002"), (("找平",), "main", "Z011101006")]))
    code, fin = _flow(tmp_path, [_xlsx(tmp_path, rows, "资料/做法-范围.xlsx")], agent)
    assert code == 0, fin
    cd = fin["data"]["coverage_detail"]
    assert cd["practice_segments"] == sum(cd["by_destination"].values()) == 3
    assert [(e["scope"], e["kind"], e["hits"]) for e in cd["exclusions"]] == [("global", "owner", 0),
                                                                             ("inline", "owner", 1)]
    assert cd["provisional_segments"] == 1
    assert [(x["sentence"], x["reason"]) for x in cd["scope_related"]] == [("注：面层为二次装修", "too_generic")]
    codes = {w["code"] for w in fin["warning_items"]}
    assert {"quantity_fitout_provisional", "quantity_exclusion_object_too_generic"} <= codes


def test_same_name_copies_in_two_folders_read_once(tmp_path):
    a = _xlsx(tmp_path, WALL_ROWS, "资料/1.资料/做法表.xlsx")
    b = "资料/做法表.xlsx"
    (tmp_path / "资料" / "做法表.xlsx").write_bytes((tmp_path / "资料" / "1.资料" / "做法表.xlsx").read_bytes())
    agent = FakeAgent(WALL_PICK)
    code, fin = _flow(tmp_path, [a, b], agent)
    assert code == 0, fin
    cd = fin["data"]["coverage_detail"]
    assert cd["practice_segments"] == 1 and len(agent.batches) == 1
    assert [(f.get("duplicate_of") or {}).get("display_path") for f in cd["by_file"]] == [None, a]
    dup = [w for w in fin["warning_items"] if w["code"] == "quantity_duplicate_file"]
    assert [w["message"] for w in dup] == [f"「{a}」与「{b}」内容相同，只按「{a}」生成"]
    assert (fin["coverage"]["files_total"], fin["coverage"]["files_used"]) == (2, 1)


def test_partial_when_answers_miss_a_segment(tmp_path):
    rows = WALL_ROWS + ["地面1：水泥地面", "1. 20厚1:2水泥砂浆面层"]
    code, fin = _flow(tmp_path, [_xlsx(tmp_path, rows)], FakeAgent(WALL_PICK, drop=("水泥砂浆面层",)))
    assert code == 0 and fin["status"] == "partial" and fin["status_reason"] == "answers_incomplete"
    assert fin["coverage"]
    assert any(w["level"] == "warn" and w["code"] == "quantity_llm_failed" for w in fin["warning_items"])
    assert [u["kind"] for u in fin["data"]["unmatched"]] == ["llm_failed"]


def test_no_text_is_exit_2(tmp_path):
    (tmp_path / "资料").mkdir()
    (tmp_path / "资料" / "空文件.txt").write_bytes(b"   ")
    code, body = _call("prepare", "--root", tmp_path, "--files", "资料/空文件.txt")
    assert code == 2 and body["error"]["code"] == "no_text" and "read_image" in body["error"]["hint"]


def test_empty_result_is_empty_envelope_not_ok(tmp_path):
    code, fin = _flow(tmp_path, [_xlsx(tmp_path, WALL_ROWS)], FakeAgent(lambda s: ("main", "")))
    assert code == 0 and fin["status"] == "empty" and fin["status_reason"] == "no_items"
    assert "没有形成任何正式清单项" in fin["data"]["detail"]
    assert fin["data"]["unmatched"] and fin["data"]["coverage_detail"]["practice_segments"] == 1


def test_out_of_scope_is_empty_before_agent(tmp_path):
    rows = ["空调风管采用镀锌钢板，风机盘管吊装，排烟风机、配电箱、桥架、电缆、照明灯具、插座、监控摄像机"] * 4
    code, prep = _call("prepare", "--root", tmp_path, "--files", _xlsx(tmp_path, rows, "资料/暖通.xlsx"))
    assert code == 0 and prep["status"] == "empty" and prep["status_reason"] == "out_of_scope"
    assert "batches" not in prep["data"]


def test_no_answers_and_stale_prepare_are_exit_2(tmp_path):
    rel = _xlsx(tmp_path, WALL_ROWS)
    code, prep = _call("prepare", "--root", tmp_path, "--files", rel)
    (tmp_path / "prep.json").write_text(json.dumps(prep, ensure_ascii=False), encoding="utf-8")
    code, body = _call("finalize", "--prepare", tmp_path / "prep.json")
    assert code == 2 and body["error"]["code"] == "no_answers"
    (tmp_path / "empty.json").write_text('{"r": []}', encoding="utf-8")
    code, body = _call("finalize", "--prepare", tmp_path / "prep.json", "--answers", tmp_path / "empty.json")
    assert code == 2 and body["error"]["code"] == "answers_missing"
    _xlsx(tmp_path, WALL_ROWS[:1] + ["1. 240mm厚MU20蒸压粉煤灰砖，M7.5水泥砂浆砌筑"])   # 改了原件
    answers = FakeAgent(WALL_PICK).answer(prep["data"])
    (tmp_path / "answers.json").write_text(json.dumps(answers, ensure_ascii=False), encoding="utf-8")
    code, body = _call("finalize", "--prepare", tmp_path / "prep.json", "--answers", tmp_path / "answers.json")
    assert code == 2 and body["error"]["code"] == "prepare_mismatch"


def test_render_requires_finalize_result(tmp_path):
    (tmp_path / "x.json").write_text('{"status": "ok", "data": {}}', encoding="utf-8")
    code, body = _call("render", "--result", tmp_path / "x.json", "--xlsx", tmp_path / "a.xlsx")
    assert code == 2 and body["error"]["code"] == "invalid_result"


def test_render_partial_stays_partial(tmp_path):
    rows = WALL_ROWS + ["地面1：水泥地面", "1. 20厚1:2水泥砂浆面层"]
    code, fin = _flow(tmp_path, [_xlsx(tmp_path, rows)], FakeAgent(WALL_PICK, drop=("水泥砂浆面层",)))
    (tmp_path / "fin.json").write_text(json.dumps(fin, ensure_ascii=False), encoding="utf-8")
    code, out = _call("render", "--result", tmp_path / "fin.json", "--out", f"{tmp_path / 'xlsx'}/")
    assert code == 0 and out["status"] == "partial"
    wb = load_workbook(BytesIO((tmp_path / "xlsx" / "清单框架.xlsx").read_bytes()))
    assert any("部分完成" in str(c.value) for row in wb.active.iter_rows() for c in row if c.value)
